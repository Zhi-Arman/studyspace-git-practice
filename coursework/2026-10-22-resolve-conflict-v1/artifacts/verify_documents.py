from pathlib import Path
from datetime import datetime
from zoneinfo import ZoneInfo
from zipfile import ZipFile
import hashlib
import json
import re
import subprocess
from docx import Document
from lxml import etree

ROOT=Path(__file__).resolve().parents[1]
path=ROOT/'StudySpace_解决合并冲突_中俄双语_v1.docx'
doc=Document(path)
content=json.loads((ROOT/'artifacts/report-content.json').read_text())
compact=lambda s:re.sub(r'\s+','',s)
flat=compact(''.join(p.text for p in doc.paragraphs)+''.join(c.text for t in doc.tables for row in t.rows for c in row.cells))
expected=[]
for page in content:
    expected.extend([page['zh'],page['ru']])
    for b in page['blocks']:
        if b['type']=='pair':expected.extend([b['zh'],b['ru']])
        elif b['type']=='code':expected.append(b['text'])
        else:
            expected.extend(b['heads']);expected.extend(x for row in b['rows'] for x in row)
assert all(compact(x) in flat for x in expected)
with ZipFile(path) as z:
    assert z.testzip() is None
    for name in z.namelist():
        if name.endswith('.xml'):etree.fromstring(z.read(name))
    assert b'<w:pBdr' not in z.read('word/document.xml')
    assert b'<w:pBdr' not in z.read('word/styles.xml')
md=(ROOT/'README.md').read_text()
for target in re.findall(r'\]\(([^)]+)\)',md):
    if not target.startswith(('https://','http://')):assert (ROOT/target).exists(),target
assert len(re.findall(r'^\| RS-0[1-6] \|',md,re.M))==6
E=json.loads((ROOT/'artifacts/execution.json').read_text())
L=json.loads((ROOT/'artifacts/local-resolution.json').read_text())
assert E['status']=='CONFLICT_RESOLVED_BUILD_PASS_MERGED'
assert all(c['passed'] for c in E['checks']) and all(c['passed'] for c in L['checks'])
assert E['build_exit_code']==0 and E['build_warnings']==0
assert Path(E['build_html']).exists() and 'build succeeded.' in (ROOT/'artifacts/sphinx-build.log').read_text()
assert 'WARNING:' not in (ROOT/'artifacts/sphinx-build.log').read_text()
B=Path(E['workspaces']['B']['path'])
assert (B/'docs/build-check.md').read_text()==E['resolved_command']
assert not re.search(r'^(<<<<<<<|=======|>>>>>>>)',(B/'docs/build-check.md').read_text(),re.M)
assert subprocess.check_output(['git','status','--porcelain=v1','--untracked-files=all'],cwd=B,text=True)==''
assert not (B/'.git/MERGE_HEAD').exists()
assert (ROOT/'artifacts/conflict-before.md').read_text().startswith('<<<<<<< HEAD')
assert E['independent_approval']=='NOT_VERIFIED' and E['business_tests']=='NOT_RUN'
images=sorted((ROOT/'artifacts/render').glob('page-*.png'))
result=dict(checked_at=datetime.now(ZoneInfo('Asia/Shanghai')).isoformat(),docx_sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
            zip_xml_integrity='PASS',bilingual_content_parity='PASS',source_text_items=len(expected),markdown_links='PASS',
            six_task_items_present=True,resolved_file_has_no_markers=True,real_build='PASS',warnings=0,
            worktree='CLEAN',previous_conflict_snapshot_preserved=True,pages=len(images),all_pages_visually_reviewed=False)
(ROOT/'artifacts/document-verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(result,ensure_ascii=False,indent=2))
