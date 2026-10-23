from pathlib import Path
import json
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
E=json.loads((ROOT/'artifacts/execution.json').read_text())
L=json.loads((ROOT/'artifacts/local-resolution.json').read_text())
TITLE='StudySpace_解决合并冲突_中俄双语_v1'
pages=[]
def page(zh,ru):
    blocks=[];pages.append(dict(zh=zh,ru=ru,blocks=blocks));return blocks
def pair(b,zh,ru):b.append(dict(type='pair',zh=zh,ru=ru))
def code(b,text):b.append(dict(type='code',text=text))
def table(b,heads,rows,widths):b.append(dict(type='table',heads=heads,rows=rows,widths=widths))

b=page('StudySpace 解决合并冲突','StudySpace Разрешение конфликта')
pair(b,'SS-GIT-RESOLVE-01 · v1.0 · 2026-10-22 · Asia/Shanghai · 文档 DRAFT',
       'SS-GIT-RESOLVE-01 · v1.0 · 22.10.2026 · Asia/Shanghai · Документ DRAFT')
pair(b,'已在上一题保留的 B 工作区解决冲突，同时保留 -W 与 -n。严格文档构建成功，随后暂存文件、完成本地合并并合并 PR #3。A、B 的 main 均与线上一致，工作区干净，所有练习分支保留。',
       'Конфликт разрешён в сохранённом каталоге B; сохранены -W и -n. Строгая сборка документации прошла, затем файлы добавлены в индекс, завершено локальное слияние и слит PR #3. main A и B совпадает с GitHub, деревья чистые; все учебные ветки сохранены.')
code(b,E['pr_url'])
table(b,['编号\nКод','图片要求\nТребование на фото','实际结果\nРезультат'],[
['RS-01','打开文件并找到冲突标记\nОткройте docs/build-check.md и найдите маркеры','核对原始标记及索引三阶段，已保存快照\nМаркеры и три стадии проверены; снимок сохранён'],
['RS-02','阅读两版并确定最终内容\nПрочитайте оба варианта и согласуйте итоговую формулировку','按自动化演示决定保留 -W -n\nВ демонстрации выбрано сохранение -W -n'],
['RS-03','删除冲突标记\nУдалите маркеры конфликта','最终文件仅一行命令，无标记\nОдна строка команды, маркеров нет'],
['RS-04','检查 diff 并构建文档\nПроверьте diff и соберите документацию','diff 检查通过；实际构建退出 0、警告 0\nDiff проверен; сборка: код 0, предупреждений 0'],
['RS-05','将文件加入索引\nДобавьте файл в индекс','git add 后未合并索引为空\nПосле git add неслитых записей нет'],
['RS-06','完成合并并检查历史\nЗавершите слияние и проверьте историю','双父本地合并 '+E['resolution_commit'][:7]+'；PR #3 已合并\nЛокальное слияние с двумя родителями '+E['resolution_commit'][:7]+'; PR #3 слит'],
],[1.8,7.5,8.1])
pair(b,'沿用用户已授权的单账号自动化角色演示，实际账号为 Zhi-Arman。保留两项检查是本次自动化决策，真人协商、本人执行与独立批准仍为 NOT_VERIFIED。',
       'Использована разрешённая пользователем автоматическая демонстрация аккаунтом Zhi-Arman. Сохранение двух проверок — решение автоматизации; реальное согласование студентов, личное выполнение и независимое одобрение — NOT_VERIFIED.')

b=page('2 最终命令与实际文档构建','2 Итоговая команда и реальная сборка')
pair(b,'A 的 -W 将警告视为构建失败，B 的 -n 启用更严格的引用检查。两者可以同时保留，因此将冲突内容合并为以下唯一一行。',
       '-W из версии A делает предупреждения ошибками сборки; -n из версии B включает более строгую проверку ссылок. Оба параметра сохранены в единственной строке ниже.')
code(b,E['resolved_command'].strip())
pair(b,'共享仓库原先只有练习文档，没有完整构建配置。为实际执行这行命令，本轮补齐最小入口、配置和已有课程的哈希依赖锁；未修改应用业务代码。',
       'Изначально общий репозиторий содержал учебные документы без полной конфигурации сборки. Для реального запуска добавлены минимальный вход, конфигурация и хешированный lock-файл зависимостей из курса; бизнес-код не изменялся.')
table(b,['文件\nФайл','修改用途\nНазначение'],[
['docs/build-check.md','冲突解决后保留 -W -n\nРазрешение конфликта с сохранением -W -n'],
['docs/conf.py','Sphinx 与 MyST 配置\nКонфигурация Sphinx и MyST'],
['docs/index.md','展示当前用户指南、管理员指南和最终命令\nТекущие руководства и итоговая команда'],
['requirements-docs.txt','复制课程现有的哈希锁，精确安装文档依赖\nСуществующий lock курса с проверкой хешей'],
['.gitignore','忽略 .venv、site 与 Python 缓存\nИсключение .venv, site и кеша Python'],
],[5.7,11.7])
pair(b,'在 B 仓库根目录创建独立虚拟环境，按哈希安装依赖并通过 uv pip check。实际版本为 Python 3.13.12、Sphinx 8.2.3、MyST 4.0.1。离线缓存缺包后，使用相同锁文件联网安装。',
       'В корне B создана отдельная виртуальная среда, зависимости установлены с проверкой хешей; uv pip check пройден. Версии: Python 3.13.12, Sphinx 8.2.3, MyST 4.0.1. После нехватки офлайн-кеша тот же lock-файл установлен по сети.')
pair(b,'实际执行完整最终命令，退出码 0，警告数 0，生成 site/index.html。HTML 内容检查确认页面含最终命令及当前用户、管理员指南。构建发生在 git add 和合并提交之前。',
       'Выполнена точная итоговая команда: код возврата 0, предупреждений 0, создан site/index.html. Проверка HTML подтвердила команду и текущие руководства. Сборка выполнена до git add и коммита слияния.')
pair(b,'验证范围是本仓库的文档构建与内容；未验证浏览器外观，应用业务测试仍为 NOT_RUN，未宣称教师验收通过。',
       'Проверены сборка и содержание документации репозитория. Внешний вид в браузере не проверялся; бизнес-тесты приложения остаются NOT_RUN, приёмка преподавателем не заявляется.')

b=page('3 合并历史与交付位置','3 История слияния и материалы')
pair(b,f'执行时间 {E["executed_at"]}。本地 {E["local_checks"]}/{E["local_checks"]} 项检查与发布 {len(E["checks"])}/{len(E["checks"])} 项检查通过。最终线上树与已实际构建的提交树完全一致。',
       f'Время выполнения {E["executed_at"]}. Пройдены {E["local_checks"]}/{E["local_checks"]} локальных и {len(E["checks"])}/{len(E["checks"])} проверок публикации. Итоговое дерево GitHub полностью совпадает с реально собранным деревом.')
table(b,['记录\nЗапись','提交或结果\nКоммит или результат'],[
['原 A 合并\nПрежнее слияние A',E['base_main']],
['原 B 提交\nПрежний коммит B',E['b_original']],
['本地冲突解决\nЛокальное разрешение',E['resolution_commit']],
['最终 PR 合并\nИтоговое слияние PR',E['final_merge']],
],[5.5,11.9])
pair(b,'本地解决提交保留原 A 合并与原 B 提交两个父节点。通过普通快进推送更新 docs/conflict-b 后合并 PR #3；没有强推或删除分支。A/B 均已同步 main，未合并索引为空，MERGE_HEAD 已清除。',
       'Локальный коммит разрешения сохраняет двух родителей: прежние A и B. docs/conflict-b обновлена обычным fast-forward push, затем PR #3 слит. Принудительный push и удаление веток не применялись. A/B синхронизированы с main; неслитых записей и MERGE_HEAD нет.')
pair(b,'以下为本轮报告与证据目录。实际解决文件和生成的网站沿用上一题的 author-b 工作区；上一题原始冲突快照仍保存在其 artifacts 中。',
       'Ниже каталог отчёта и свидетельств. Разрешённый файл и сайт находятся в прежнем author-b; исходный снимок конфликта предыдущего задания сохранён в его artifacts.')
code(b,str(ROOT)+'\n\nREADME.md\n'+TITLE+'.docx\nartifacts/conflict-before.md\nartifacts/sphinx-build.log\nartifacts/local-resolution.json\nartifacts/history-final.txt\nartifacts/execution.json')
pair(b,'可在实际 B 工作区运行 git status、git log --graph --oneline --decorate，以及最终命令复核。构建产物与虚拟环境均已忽略，未上传 GitHub。',
       'В фактическом каталоге B можно проверить git status, git log --graph --oneline --decorate и итоговую команду. Результат сборки и виртуальная среда исключены из Git и не отправлены на GitHub.')

doc=Document();s=doc.sections[0]
s.page_width=Cm(21);s.page_height=Cm(29.7);s.top_margin=Cm(1.6);s.bottom_margin=Cm(1.6);s.left_margin=Cm(1.8);s.right_margin=Cm(1.8)
for name in ('Normal','Title','Subtitle','Heading 1'):
    sty=doc.styles[name];sty.font.name='Arial';sty.font.color.rgb=RGBColor(0,0,0)
    sty._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Hiragino Sans GB')
n=doc.styles['Normal'];n.font.size=Pt(10.5);n.paragraph_format.space_after=Pt(4);n.paragraph_format.line_spacing=1.03
doc.styles['Title'].font.size=Pt(20);doc.styles['Subtitle'].font.size=Pt(13);doc.styles['Heading 1'].font.size=Pt(16)
cs=doc.styles.add_style('Code',1);cs.font.size=Pt(8.5);cs.font.name='Courier New';cs.paragraph_format.space_after=Pt(6);cs.paragraph_format.line_spacing=1
doc.core_properties.title='StudySpace 解决合并冲突 Разрешение конфликта';doc.core_properties.author='StudySpace'
def paragraph(text,style=None,ru=False):
    p=doc.add_paragraph(style=style)
    if ru:p.paragraph_format.space_after=Pt(6)
    for i,line in enumerate(text.split('\n')):
        if i:p.add_run().add_break()
        p.add_run(line)
    return p
md=[]
for pi,pag in enumerate(pages):
    h=paragraph(pag['zh'],'Title' if pi==0 else 'Heading 1')
    if pi:h.paragraph_format.page_break_before=True
    paragraph(pag['ru'],'Subtitle');md.extend([('#' if pi==0 else '##')+' '+pag['zh'],'',pag['ru'],''])
    for block in pag['blocks']:
        if block['type']=='pair':
            paragraph(block['zh']);paragraph(block['ru'],ru=True);md.extend([block['zh'],'',block['ru'],''])
        elif block['type']=='code':
            paragraph(block['text'],'Code');md.extend(['```text',block['text'],'```',''])
        else:
            t=doc.add_table(rows=1,cols=len(block['heads']));t.style='Table Grid';t.autofit=False
            for j,width in enumerate(block['widths']):t.columns[j].width=Cm(width)
            for i,values in enumerate([block['heads']]+block['rows']):
                row=t.rows[0] if i==0 else t.add_row();row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
                if i==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
                for j,value in enumerate(values):
                    cell=row.cells[j];cell.width=Cm(block['widths'][j]);cell.text=value
                    for p in cell.paragraphs:
                        p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1
                        for r in p.runs:r.font.size=Pt(9.5);r.bold=(i==0)
                    if i==0:
                        sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EEF4');cell._tc.get_or_add_tcPr().append(sh)
            doc.add_paragraph().paragraph_format.space_after=Pt(0)
            md.append('| '+' | '.join(x.replace('\n','<br>') for x in block['heads'])+' |');md.append('| '+' | '.join('---' for _ in block['heads'])+' |')
            md.extend('| '+' | '.join(x.replace('\n','<br>') for x in row)+' |' for row in block['rows']);md.append('')
for root in (doc._element,doc.styles.element):
    for border in root.xpath('.//w:pBdr'):border.getparent().remove(border)
for p in list(doc.paragraphs)+[p for t in doc.tables for row in t.rows for c in row.cells for p in c.paragraphs]:
    for r in p.runs:
        rf=r._element.get_or_add_rPr().get_or_add_rFonts()
        for attr in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):rf.attrib.pop(qn('w:'+attr),None)
        for key in ('ascii','hAnsi'):rf.set(qn('w:'+key),'Courier New' if p.style.name=='Code' else 'Arial')
        rf.set(qn('w:eastAsia'),'Hiragino Sans GB')
doc.save(ROOT/(TITLE+'.docx'))

md.extend(['## 文件入口 Входные файлы','',f'[双语 Word Двуязычный Word]({TITLE}.docx)',
           '[已合并 PR #3 Слитый PR #3]('+E['pr_url']+')',
           '[完整检查记录 Проверки](artifacts/execution.json)',
           '[构建日志 Журнал сборки](artifacts/sphinx-build.log)',
           '[最终历史 История](artifacts/history-final.txt)',
           '[已解决文件 Разрешённый файл](../2026-10-22-controlled-conflict-v1/author-b/docs/build-check.md)',
           '[构建站点 Сайт](../2026-10-22-controlled-conflict-v1/author-b/site/index.html)',''])
(ROOT/'README.md').write_text('\n'.join(md))
(ROOT/'artifacts/report-content.json').write_text(json.dumps(pages,ensure_ascii=False,indent=2)+'\n')
previous=ROOT.parent/'2026-10-22-controlled-conflict-v1/artifacts/fonts.conf'
(ROOT/'artifacts/fonts.conf').write_text(previous.read_text().replace(str(previous.parent/'font-cache'),str(ROOT/'artifacts/font-cache')))
print(ROOT/(TITLE+'.docx'))

