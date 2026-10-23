from pathlib import Path
from datetime import datetime, timezone
from zipfile import ZipFile
import json
import hashlib
import re

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

ROOT = Path('/Users/zhi/Documents/软件生命周期/studyspace')
OUT = ROOT / 'coursework/2026-10-15-environments-protection-v1'
DATA = OUT / 'data'
DATA.mkdir(parents=True, exist_ok=True)
STEM = 'StudySpace_环境与防护表_中俄双语_v1'
DOCX = OUT / f'{STEM}.docx'
SOURCE_PATHS = [
    'AGENTS.md', 'compose.yaml', '.env.example', '.gitignore',
    'frontend/vite.config.ts', 'backend/app/config.py', 'backend/app/main.py',
    'backend/tests/conftest.py', 'backend/tests/test_operations.py', 'backend/tests/test_config.py',
    'scripts/dev.sh', 'scripts/check.sh', 'docs/zh-CN/04_System_Design.md',
    'docs/ru-RU/04_System_Design.md', 'docs/zh-CN/05_Test_Plan.md',
    'docs/ru-RU/05_Test_Plan.md', 'docs/data/project-decisions.json',
    'docs/decisions/zh-CN/0004-docker-https.md', 'docs/decisions/ru-RU/0004-docker-https.md',
    'coursework/2026-10-15-deployment-nodes-v1/StudySpace_节点与制品表_中俄双语_v1.docx',
]
fingerprints = {p: hashlib.sha256((ROOT / p).read_bytes()).hexdigest() for p in SOURCE_PATHS}
headers = [
    ('环境', 'Среда'),
    ('用途', 'Назначение'),
    ('与生产环境的差异', 'Отличия от продакшена'),
    ('访问权限与敏感凭据', 'Доступ и секреты'),
    ('备份与监控', 'Резервное копирование и мониторинг'),
]
rows = [
    [
        ('开发\ndevelopment', 'Разработка'),
        ('编写和调试前后端，运行骨架功能。业务功能仍待实现。', 'Разработка и отладка клиента и сервера, запуск каркаса. Бизнес-функции ещё предстоят.'),
        ('SQLite 与本地 HTTP；Vite 热更新，端口 5173，API 8000。本地 Compose 候选入口为\n8080，未配置生产 HTTPS。', 'SQLite и локальный HTTP; Vite с обновлением, порт 5173, API 8000. Локальный вариант Compose использует вход 8080 без производственного HTTPS.'),
        ('本机开发者访问，服务绑定 127.0.0.1；本地 .env 不入 Git。生产凭据不用于开发；VITE_* 不存密钥。', 'Доступ разработчика на своём ПК; сервисы привязаны к 127.0.0.1. Файл .env вне Git; производственные секреты не используются, ключи не помещаются в VITE_*.'),
        ('建议按需保存或重建开发样例数据；日志及现有健康、就绪接口用于排障。自动备份和告警尚未配置。', 'Предлагается сохранять или пересоздавать учебные данные по необходимости. Журналы и имеющиеся health/ready служат диагностике. Автоматические копии и оповещения не настроены.'),
    ],
    [
        ('测试\ntest', 'Тест'),
        ('用 pytest、Vitest 验证骨架；计划验证权限、预约冲突、迁移及恢复。业务用例仍为 NOT_RUN。', 'pytest и Vitest проверяют каркас; планируются проверки прав, конфликтов брони, миграции и восстановления. Бизнес-кейсы остаются NOT_RUN.'),
        ('当前测试夹具使用临时\nSQLite 和隔离配置；计划使用合成数据。PostgreSQL 集成、并发验证及独立测试部署待实施。', 'Текущие фикстуры используют временную SQLite и изолированные настройки; планируются синтетические данные. Интеграция и конкуренция PostgreSQL и отдельное тестовое размещение ещё предстоят.'),
        ('当前夹具不读取项目 .env，并清除相关环境变量。独立测试环境拟仅向团队开放，使用独立测试凭据，不复制生产密码或学生证明影像。', 'Фикстуры не читают .env проекта и очищают связанные переменные. Отдельная среда планируется только для команды со своими секретами, без производственных паролей и изображений документов студентов.'),
        ('临时测试库可重建；计划保存脱敏测试报告。已有 health/ready 和失败返回测试；PostgreSQL 备份恢复与告警验证待实施。', 'Временная БД пересоздаётся; планируется сохранять обезличенные отчёты. Есть тесты health/ready и ответа при сбое; восстановление PostgreSQL и оповещения ещё предстоит проверить.'),
    ],
    [
        ('生产\nproduction\n目标方案', 'Продакшен\nЦелевая схема'),
        ('正式为学生和管理员提供预约服务。目标 APP_ENV 为 production；尚未部署。', 'Рабочий сервис бронирования для студентов и администраторов. Целевой APP_ENV — production; размещение ещё не выполнено.'),
        ('本列比较的基准。目标采用 Docker Compose、Caddy HTTPS、Nginx 静态前端、FastAPI 和内网 PostgreSQL，使用持久数据卷；配置待实施。', 'База сравнения. Цель: Docker Compose, HTTPS Caddy, статика Nginx, FastAPI и PostgreSQL во внутренней сети с постоянными томами; конфигурация ещё предстоит.'),
        ('计划按最小权限控制：学生仅见本人数据，管理员按职责授权；主机运维访问受限，数据库不公开。生产密钥独立，通过环境变量注入，不入 Git 或前端。', 'Планируются минимальные права: студент видит своё, администратор действует по роли; доступ к хосту ограничен, БД не публична. Секреты отдельные, через переменные среды, вне Git и клиента.'),
        ('已选策略待实施：每日备份，保留最近 7 份；恢复后先到期清理再开放。计划监控 HTTPS、健康/就绪、数据库、磁盘和备份失败，配置告警并演练恢复。', 'Выбрано, ещё не реализовано: ежедневные копии, последние 7; после восстановления сначала очистка по срокам. Планируются контроль HTTPS, health/ready, БД, диска и сбоев копирования, оповещения и тренировка восстановления.'),
    ],
]
intro = (
    '本表对应任务 1б，比较 StudySpace 开发、测试和生产环境。现有开发配置与隔离测试代码已经建立；生产为目标方案，部署、备份恢复及告警按实施状态标注。',
    'Таблица задания 1б сравнивает разработку, тест и продакшен StudySpace. Конфигурация разработки и код изолированных тестов имеются; продакшен остаётся целевой схемой, размещение, восстановление и оповещения отмечены по состоянию реализации.'
)
notes = [
    ('状态说明', 'Пояснение статусов',
     '当前 compose.yaml 仍是 development 本地 HTTP 候选，未包含 Caddy 或 PostgreSQL；独立测试部署和生产部署待实施，生产机位与域名由 OPEN-010 确认。表内“计划”“建议”是方案，不表示配置已生效。',
     'Текущий compose.yaml — локальный вариант development с HTTP без Caddy и PostgreSQL. Отдельные тестовое и производственное размещения ещё предстоят; хост и домен определяются в OPEN-010. «Планируется» и «предлагается» обозначают схему, а не действующую настройку.'),
    ('健康检查边界', 'Границы проверок состояния',
     '/api/v1/health 仅检查进程，/api/v1/ready 仅检查数据库连接；Compose 已声明就绪检查，但不等于容器运行验证或业务验收。认证、预约等业务尚未实现，业务用例保持 NOT_RUN。',
     '/api/v1/health проверяет процесс, /api/v1/ready — соединение с БД. Проверка готовности объявлена в Compose, но это не подтверждает запуск контейнеров или бизнес-приёмку. Аутентификация, бронь и другие бизнес-функции ещё не реализованы; кейсы остаются NOT_RUN.'),
]

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(0.55)
section.left_margin = section.right_margin = Inches(0.6)
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor(0,0,0)
    rpr = style.element.get_or_add_rPr()
    rf = OxmlElement('w:rFonts')
    for a in ['ascii', 'hAnsi', 'cs']:
        rf.set(qn('w:'+a), 'Arial')
    rf.set(qn('w:eastAsia'), 'PingFang SC')
    rpr.insert(0, rf)
for style in doc.styles:
    for b in list(style.element.iter(qn('w:pBdr'))):
        b.getparent().remove(b)
doc.styles['Normal'].font.size = Pt(10.5)
doc.styles['Normal'].paragraph_format.line_spacing = 1.05
doc.styles['Normal'].paragraph_format.space_after = Pt(4)
doc.styles['Title'].font.size = Pt(18)
doc.styles['Title'].paragraph_format.space_after = Pt(4)
doc.styles['Heading 1'].font.size = Pt(12)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(8)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(4)

def run(p, text, size=10.5, bold=False, color='000000', lang='zh-CN'):
    r = p.add_run(text)
    r.font.name = 'Arial'
    r.font.size = Pt(size)
    r.bold = bold
    r.font.color.rgb = RGBColor.from_string(color)
    rpr=r._element.get_or_add_rPr()
    rf=rpr.find(qn('w:rFonts'))
    rf.set(qn('w:eastAsia'), 'PingFang SC')
    rf.set(qn('w:hAnsi'), 'Arial')
    le=OxmlElement('w:lang')
    le.set(qn('w:val'), lang)
    le.set(qn('w:eastAsia'), 'zh-CN')
    rpr.append(le)
    return r

def para(text, size=10.5, style=None, lang='zh-CN'):
    p=doc.add_paragraph(style=style)
    run(p,text,size=size,lang=lang)
    return p

def pair(zh,ru,style=None,size=10.5):
    p=para(zh,size,style)
    p.paragraph_format.keep_with_next=True
    p=para(ru,size,style,'ru-RU')
    p.paragraph_format.keep_together=True
    return p

pair('StudySpace 环境与防护表', 'StudySpace Среды и защита', 'Title', 18)
para('任务 1б / Задание 1б  |  SS-ENV-T1B-01  |  v1  |  2026-10-15  |  DRAFT',9)
pair(*intro)

table=doc.add_table(rows=1,cols=5)
table.alignment=WD_TABLE_ALIGNMENT.CENTER
table.autofit=False
widths=[1.00,1.30,1.55,1.64,1.81]
for col,w in zip(table.columns,widths):
    col.width=Inches(w)
pr=table._tbl.tblPr
borders=OxmlElement('w:tblBorders')
for edge in ['top','left','bottom','right','insideH','insideV']:
    e=OxmlElement('w:'+edge)
    for k,v in [('val','single'),('sz','5'),('color','D9D9D9')]:
        e.set(qn('w:'+k),v)
    borders.append(e)
pr.append(borders)
repeat=OxmlElement('w:tblHeader')
table.rows[0]._tr.get_or_add_trPr().append(repeat)

def cell_text(cell, val, ci, header=False):
    cell.width=Inches(widths[ci])
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cp=cell._tc.get_or_add_tcPr()
    mar=OxmlElement('w:tcMar')
    for edge,w in [('top',95),('bottom',95),('left',80),('right',80)]:
        e=OxmlElement('w:'+edge)
        e.set(qn('w:w'),str(w)); e.set(qn('w:type'),'dxa'); mar.append(e)
    cp.append(mar)
    for j,t in enumerate(val):
        p=cell.paragraphs[0] if j==0 else cell.add_paragraph()
        p.paragraph_format.space_after=Pt(4 if j==0 else 0)
        p.paragraph_format.line_spacing=1.03
        p.paragraph_format.keep_with_next=False
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER if header or ci==0 else WD_ALIGN_PARAGRAPH.LEFT
        run(p,t,9.5,header,'FFFFFF' if header else '000000','zh-CN' if j==0 else 'ru-RU')

for ci,val in enumerate(headers):
    c=table.rows[0].cells[ci]
    cell_text(c,val,ci,True)
    shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),'29485D'); c._tc.get_or_add_tcPr().append(shd)
for ri,vals in enumerate(rows):
    row=table.add_row()
    no_split=OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(no_split)
    for ci,val in enumerate(vals):
        c=row.cells[ci]; cell_text(c,val,ci)
        shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),'F0F4F7' if ri%2==0 else 'FFFFFF'); c._tc.get_or_add_tcPr().append(shd)

for zh,ru,zt,rt in notes:
    pair(zh,ru,'Heading 1',12)
    pair(zt,rt)
pair('项目依据', 'Источники проекта','Heading 1',12)
pair('系统设计 DOC-04、测试计划 DOC-05、部署决定 ADR-0004 与 docs/data/project-decisions.json。', 'Проект DOC-04, план тестирования DOC-05, решение ADR-0004 и docs/data/project-decisions.json.',size=9)
para('配置与代码 / Конфигурации и код\ncompose.yaml; .env.example; .gitignore; backend/tests/conftest.py;\nbackend/app/main.py; scripts/check.sh.',size=9)
para('项目目录 / Каталог проекта\n/Users/zhi/Documents/软件生命周期/studyspace',size=9)
doc.core_properties.title='StudySpace 环境与防护表 Среды и защита'
doc.core_properties.subject='任务 1б / Задание 1б'
doc.core_properties.author='StudySpace'
doc.save(DOCX)

content={'doc_id':'SS-ENV-T1B-01','version':'v1','date':'2026-10-15','status':'DRAFT','intro':intro,'headers':headers,'rows':rows,'notes':notes,'sources':fingerprints}
(DATA/'content.json').write_text(json.dumps(content,ensure_ascii=False,indent=2),encoding='utf-8')
font_config=(ROOT/'coursework/2026-10-15-deployment-nodes-v1/data/fonts.conf').read_text().replace('studyspace-nodes-fontcache','studyspace-env-fontcache')
(DATA/'fonts.conf').write_text(font_config,encoding='utf-8')
md=['# StudySpace 环境与防护表','# StudySpace Среды и защита','SS-ENV-T1B-01 · v1 · 2026-10-15 · DRAFT',*intro]
md += ['| '+' | '.join('<br>'.join(h) for h in headers)+' |','| '+' | '.join(['---']*5)+' |']
md += ['| '+' | '.join('<br>'.join(v).replace('\n','<br>') for v in row)+' |' for row in rows]
for zh,ru,zt,rt in notes:
    md += [f'## {zh} / {ru}',zt,rt]
(OUT/f'{STEM}.md').write_text('\n\n'.join(md),encoding='utf-8')
(OUT/'README.md').write_text(f'''# StudySpace 环境与防护 Среды и защита

SS-ENV-T1B-01 · v1 · 2026-10-15 · DRAFT

交付 / Результат [{STEM}.docx]({STEM}.docx)

课堂五列表格，包含开发、测试、生产三种环境；同一单元格中文在上、俄语在下，原生 Word 表格可编辑。
Пять полей задания для разработки, теста и продакшена; китайский над русским в одной ячейке, таблица Word редактируется.

内容及来源 / Содержимое и источники data/content.json
文档检查 / Проверка документа data/verification.json
''',encoding='utf-8')

rd=Document(DOCX)
with ZipFile(DOCX) as z:
    xml=etree.fromstring(z.read('word/document.xml'))
    styles=etree.fromstring(z.read('word/styles.xml'))
    files=z.namelist()
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
text='\n'.join(p.text for p in rd.paragraphs)+'\n'+'\n'.join(c.text for row in rd.tables[0].rows for c in row.cells)
tests={
 'five_columns_three_environments':len(rd.tables)==1 and len(rd.tables[0].rows)==4 and len(rd.tables[0].columns)==5,
 'bilingual_cells_match_source':all(rd.tables[0].cell(ri,ci).text=='\n'.join(v) for ri,vals in enumerate([headers,*rows]) for ci,v in enumerate(vals)),
 'chinese_above_russian':all(len(c.paragraphs)==2 and re.search('[\u4e00-\u9fff]',c.paragraphs[0].text) and re.search('[А-Яа-яЁё]',c.paragraphs[1].text) for row in rd.tables[0].rows for c in row.cells),
 'three_environments_present':all(x in text for x in ['Разработка','Тест','Продакшен','development','test','production']),
 'isolated_test_database_supported_by_fixture':'tmp_path' in (ROOT/'backend/tests/conftest.py').read_text() and '_env_file=None' in (ROOT/'backend/tests/conftest.py').read_text(),
 'current_compose_still_development':'APP_ENV: development' in (ROOT/'compose.yaml').read_text() and '未包含 Caddy 或 PostgreSQL' in text,
 'backup_policy_preserved':'最近 7 份' in text and '每日备份' in text and '先到期清理再开放' in text,
 'production_and_business_statuses_preserved':all(x in text for x in ['尚未部署','NOT_RUN','OPEN-010','待实施']),
 'no_real_secret_values_used':all(x not in text for x in ['CHANGE_ME','postgresql+psycopg://','BEGIN PRIVATE KEY']),
 'repeated_table_header':bool(xml.xpath('//w:tbl/w:tr[1]/w:trPr/w:tblHeader',namespaces=ns)),
 'rows_kept_together':len(xml.xpath('//w:tbl/w:tr/w:trPr/w:cantSplit',namespaces=ns))==3,
 'no_fixed_row_heights':not xml.xpath('//w:trHeight',namespaces=ns),
 'native_editable_table':not xml.xpath('//w:drawing',namespaces=ns),
 'no_title_border':not styles.xpath('//w:style[@w:styleId="Title"]//w:pBdr',namespaces=ns),
 'no_headers_footers':not any(re.fullmatch(r'word/(header|footer)\d+\.xml',f) for f in files),
 'source_files_and_previous_delivery_unchanged':all(hashlib.sha256((ROOT/p).read_bytes()).hexdigest()==sha for p,sha in fingerprints.items()),
}
report={'document':str(DOCX),'checked_at_utc':datetime.now(timezone.utc).isoformat(),'execution_environment':'Local macOS; bundled Codex Python; document checks only','checks':tests,'all_pass':all(tests.values()),'visual_review':'PENDING','business_tests':'NOT_RUN','source_fingerprints':fingerprints,'docx_sha256':hashlib.sha256(DOCX.read_bytes()).hexdigest()}
(DATA/'verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'document':str(DOCX),'checks':len(tests),'all_pass':report['all_pass'],'failed':[k for k,v in tests.items() if not v]},ensure_ascii=False))
if not report['all_pass']:
    raise SystemExit(1)
