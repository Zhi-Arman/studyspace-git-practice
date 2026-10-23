from __future__ import annotations

import hashlib
import json
import re
import sys
from pathlib import Path
from zipfile import ZipFile

# Pure-Python validation packages already installed in the project.  The
# interpreter and document/image libraries come from the bundled runtime.
sys.path.append('/Users/zhi/Documents/软件生命周期/studyspace/backend/.venv/lib/python3.13/site-packages')
import yaml
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BASE = Path(__file__).resolve().parents[1]
WORK = BASE / 'work'
DOCX = BASE / '读书俱乐部_组件图时序图与API说明_中俄双语.docx'
YAML = BASE / 'book-club.openapi.yaml'

def bi(zh, ru):
    return zh + '\n' + ru

COMPONENTS = [
    ('C1', bi('Web 界面', 'Веб-интерфейс'), bi('为参与者和组织者提供注册、目录、报名、评价及管理页面。', 'Предоставляет участникам и организатору страницы регистрации, каталога, записи, отзывов и управления.')),
    ('C2', bi('用户服务', 'Сервис пользователей'), bi('处理用户注册、登录和用户资料读取。', 'Обрабатывает регистрацию, вход и получение данных пользователей.')),
    ('C3', bi('图书目录服务', 'Сервис каталога'), bi('提供图书浏览并支持组织者添加图书。', 'Предоставляет просмотр каталога и добавление книг организатором.')),
    ('C4', bi('活动服务', 'Сервис встреч'), bi('管理活动创建、活动列表、名额检查和参与者报名。', 'Управляет созданием встреч, их списком, проверкой мест и записью участников.')),
    ('C5', bi('评价服务', 'Сервис отзывов'), bi('保存和展示参与者对图书的评价。', 'Сохраняет и показывает отзывы участников о книгах.')),
    ('C6', bi('通知服务', 'Сервис уведомлений'), bi('发送报名确认邮件并按计划发送活动提醒。', 'Отправляет подтверждения записи по электронной почте и плановые напоминания о встречах.')),
    ('C7', bi('数据库', 'База данных'), bi('持久保存用户、图书、活动、报名、评价和通知任务。', 'Постоянно хранит пользователей, книги, встречи, записи, отзывы и задания уведомлений.')),
]
INTERFACES = [
    ('IWeb', 'C1 → 用户 / Пользователь', bi('参与者和组织者的网页入口', 'Веб-страницы участника и организатора')),
    ('IUsers', 'C2 → C1', bi('注册 登录 用户资料', 'Регистрация вход данные пользователя')),
    ('ICatalog', 'C3 → C1', bi('图书列表 添加图书', 'Список книг добавление книги')),
    ('IMeetings', 'C4 → C1', bi('活动列表 创建活动 报名', 'Список встреч создание встречи запись')),
    ('IReviews', 'C5 → C1', bi('读取和提交评价', 'Чтение и отправка отзывов')),
    ('INotifications', 'C6 → C4', bi('受理确认邮件和提醒任务', 'Приём заданий подтверждения и напоминания')),
    ('IData', 'C7 → C2 C3 C4 C5 C6', bi('SQL 查询 写入及事务', 'SQL запросы запись и транзакции')),
]

def schema_ref(name):
    return {'$ref': '#/components/schemas/' + name}

meeting = {'id': 12, 'bookId': 5, 'title': '读书讨论会 / Обсуждение книги', 'startsAt': '2026-11-07T06:00:00Z', 'capacity': 20, 'remainingSeats': 3}
registration = {'id': 301, 'meetingId': 12, 'participantId': 7, 'createdAt': '2026-11-01T08:30:00Z'}
def err(code, zh, ru):
    return {'code': code, 'message': zh + ' / ' + ru}
errors = {
    'invalidRequest': err('INVALID_REQUEST', '请求参数无效', 'Некорректные параметры запроса'),
    'meetingNotFound': err('MEETING_NOT_FOUND', '活动不存在', 'Встреча не найдена'),
    'alreadyRegistered': err('ALREADY_REGISTERED', '参与者已经报名', 'Участник уже записан'),
    'meetingFull': err('MEETING_FULL', '活动名额已满', 'Свободных мест нет'),
}
def json_content(schema, example=None, examples=None):
    media = {'schema': schema}
    if example is not None: media['example'] = example
    if examples: media['examples'] = examples
    return {'application/json': media}
def error_response(description, which):
    return {'description': description, 'content': json_content(schema_ref('Error'), examples={k: {'summary': v[0], 'value': errors[k]} for k,v in which.items()})}
def prop(t, desc, **kwargs):
    return {'type': t, 'description': desc, **kwargs}

SPEC = {
    'openapi': '3.0.3',
    'info': {'title': '读书俱乐部 API / API книжного клуба', 'version': '1.0.0', 'description': bi('获取活动列表并为已注册参与者创建报名。时间戳使用带时区的 RFC 3339 格式；from 按 UTC 日期包含边界筛选。', 'Получение списка встреч и создание записи для зарегистрированного участника. Метки времени имеют формат RFC 3339 с часовым поясом; from задаёт включительную нижнюю границу даты в UTC.')},
    'tags': [{'name': 'Meetings', 'description': bi('活动和报名', 'Встречи и запись')}],
    'paths': {
        '/meetings': {'get': {
            'tags': ['Meetings'], 'operationId': 'listMeetings', 'summary': '获取活动列表 / Получить список встреч',
            'description': bi('from 可选；传入时只返回 startsAt 的 UTC 日期大于或等于 from 的活动，按 startsAt 和 id 升序排列。不传则返回全部活动，无结果返回空数组。', 'Параметр from необязателен: возвращаются встречи, у которых дата startsAt в UTC не раньше from, по возрастанию startsAt и id. Без параметра возвращаются все встречи; при отсутствии результатов возвращается пустой массив.'),
            'parameters': [{'name': 'from', 'in': 'query', 'required': False, 'description': bi('起始日期，包含该日，采用 UTC 日期。', 'Начальная дата включительно по UTC.'), 'schema': {'type': 'string', 'format': 'date'}, 'example': '2026-11-01'}],
            'responses': {
                '200': {'description': '活动列表，允许为空 / Список встреч, возможно пустой', 'content': json_content({'type': 'array', 'items': schema_ref('Meeting')}, examples={'meetings': {'summary': '活动列表 / Список встреч', 'value': [meeting]}, 'empty': {'summary': '无符合条件的活动 / Подходящих встреч нет', 'value': []}})},
                '400': error_response('日期格式无效 / Некорректный формат даты', {'invalidRequest': ('参数无效 / Некорректный параметр',)}),
            },
        }},
        '/meetings/{id}/registrations': {'post': {
            'tags': ['Meetings'], 'operationId': 'registerForMeeting', 'summary': '报名活动 / Записаться на встречу',
            'description': bi('为 participantId 指定的已注册参与者报名。先检查活动是否存在，再检查重复报名和剩余名额；检查与保存必须在同一事务中完成。提交成功后受理确认邮件任务；邮件送达不属于 201 的含义。重复报名和名额已满均返回 409，并由 code 区分。', 'Записывает зарегистрированного участника participantId. Сначала проверяется существование встречи, затем повторная запись и наличие мест; проверки и сохранение выполняются в одной транзакции. После фиксации принимается задание подтверждения по почте; ответ 201 не означает доставку письма. Повторная запись и отсутствие мест дают 409 с разными code.'),
            'parameters': [{'name': 'id', 'in': 'path', 'required': True, 'description': bi('活动编号。', 'Идентификатор встречи.'), 'schema': {'type': 'integer', 'minimum': 1}, 'example': 12}],
            'requestBody': {'required': True, 'description': bi('报名参与者编号；参与者已注册。', 'Идентификатор записываемого участника; участник уже зарегистрирован.'), 'content': json_content(schema_ref('RegistrationRequest'), {'participantId': 7})},
            'responses': {
                '201': {'description': '报名记录已创建 / Запись создана', 'content': json_content(schema_ref('Registration'), registration)},
                '400': error_response('路径或请求体无效 / Некорректный путь или тело запроса', {'invalidRequest': ('请求无效 / Некорректный запрос',)}),
                '404': error_response('活动不存在 / Встреча не найдена', {'meetingNotFound': ('活动不存在 / Встреча не найдена',)}),
                '409': error_response('报名冲突 / Конфликт записи', {'alreadyRegistered': ('已经报名 / Участник уже записан',), 'meetingFull': ('名额已满 / Свободных мест нет',)}),
            },
        }},
    },
    'components': {'schemas': {
        'Meeting': {'type': 'object', 'description': bi('活动及当前可用名额。', 'Встреча и текущее число свободных мест.'), 'additionalProperties': False, 'required': list(meeting), 'properties': {
            'id': prop('integer', '活动编号 / Идентификатор встречи', minimum=1),
            'bookId': prop('integer', '讨论图书编号 / Идентификатор обсуждаемой книги', minimum=1),
            'title': prop('string', '活动名称 / Название встречи', minLength=1),
            'startsAt': prop('string', '开始时间 / Время начала', format='date-time'),
            'capacity': prop('integer', '总名额 / Вместимость', minimum=1),
            'remainingSeats': prop('integer', '剩余名额，不大于 capacity / Свободные места, не больше capacity', minimum=0),
        }, 'example': meeting},
        'RegistrationRequest': {'type': 'object', 'description': bi('报名请求。', 'Запрос записи.'), 'additionalProperties': False, 'required': ['participantId'], 'properties': {'participantId': prop('integer', '已注册参与者编号 / Идентификатор зарегистрированного участника', minimum=1)}, 'example': {'participantId': 7}},
        'Registration': {'type': 'object', 'description': bi('已保存的报名记录。', 'Сохранённая запись.'), 'additionalProperties': False, 'required': list(registration), 'properties': {
            'id': prop('integer', '报名编号 / Идентификатор записи', minimum=1),
            'meetingId': prop('integer', '活动编号 / Идентификатор встречи', minimum=1),
            'participantId': prop('integer', '参与者编号 / Идентификатор участника', minimum=1),
            'createdAt': prop('string', '创建时间 / Время создания', format='date-time'),
        }, 'example': registration},
        'Error': {'type': 'object', 'description': bi('统一错误结构；客户端依据 code 判断错误类型。', 'Единая структура ошибки; клиент определяет её тип по code.'), 'additionalProperties': False, 'required': ['code', 'message'], 'properties': {
            'code': prop('string', '稳定的错误代码 / Стабильный код ошибки', enum=[e['code'] for e in errors.values()]),
            'message': prop('string', '可读错误说明 / Понятное описание ошибки', minLength=1),
        }, 'example': errors['alreadyRegistered']},
    }},
}

# High-resolution diagram drawing.  Text is kept out of route intersections.
INK='#19324A'; BLUE='#255F8A'; GRAY='#667584'; LIGHT='#EAF2F8'; BORDER='#9CAEBB'
ZH='/System/Library/Fonts/Hiragino Sans GB.ttc'
RU='/System/Library/Fonts/Supplemental/Arial.ttf'
RUB='/System/Library/Fonts/Supplemental/Arial Bold.ttf'

class Canvas:
    def __init__(self,w,h):
        self.im=Image.new('RGB',(w,h),'white'); self.d=ImageDraw.Draw(self.im)
    def font(self,s,zh=False,bold=False):
        return ImageFont.truetype(ZH if zh else RUB if bold else RU,s)
    def text(self,xy,txt,size=30,fill=INK,anchor='mm',bold=False):
        chunks=re.findall(r'[\u3400-\u9fff]+|[^\u3400-\u9fff]+',txt)
        runs=[(chunk,self.font(size,bool(re.search('[\u3400-\u9fff]',chunk)),bold)) for chunk in chunks]
        widths=[self.d.textlength(chunk,font=font) for chunk,font in runs]
        left=xy[0]-sum(widths)/2
        for (chunk,font),width in zip(runs,widths):
            self.d.text((left+width/2,xy[1]),chunk,font=font,fill=fill,anchor=anchor)
            left+=width
    def lines(self,x,y,lines,size=30,gap=36,fill=INK):
        for i,line in enumerate(lines): self.text((x,y+i*gap),line,size,fill)
    def rect(self,box,fill=LIGHT,outline=BORDER,width=3):
        self.d.rectangle(box,fill=fill,outline=outline,width=width)
    def line(self,pts,fill=BLUE,width=3,dash=False):
        for a,b in zip(pts,pts[1:]):
            if dash:
                import math
                dx,dy=b[0]-a[0],b[1]-a[1]; length=math.hypot(dx,dy)
                for p in range(0,int(length),15):
                    q=min(p+8,length)
                    self.d.line([(a[0]+dx*p/length,a[1]+dy*p/length),(a[0]+dx*q/length,a[1]+dy*q/length)],fill=fill,width=width)
            else: self.d.line([a,b],fill=fill,width=width)
    def arrow(self,pts,dash=False,color=BLUE,filled=False):
        import math
        self.line(pts,color,3,dash)
        a,b=pts[-2:]; ang=math.atan2(b[1]-a[1],b[0]-a[0]); n=14
        head=[(b[0]-n*math.cos(ang-.45),b[1]-n*math.sin(ang-.45)),b,(b[0]-n*math.cos(ang+.45),b[1]-n*math.sin(ang+.45))]
        if filled: self.d.polygon(head,fill=color)
        else: self.d.line(head,fill=color,width=3)
    def save(self,name): self.im.save(WORK/name)

def component_diagram():
    c=Canvas(2600,1350)
    c.rect((20,12,2580,1320),'white',BORDER,3)
    c.lines(320,52,['«subsystem»','读书俱乐部系统','Система книжного клуба'],29,40)
    def port(x,y): c.rect((x-10,y-10,x+10,y+10),'white',INK,2)
    def assembly(x,y,provider_edge,label):
        # A provided lollipop meets a required socket; provider lies below.
        c.line([(x,y+13),(x,provider_edge)],BLUE,3)
        c.d.ellipse((x-13,y-13,x+13,y+13),fill='white',outline=BLUE,width=3)
        c.d.arc((x-22,y-22,x+22,y+22),180,360,fill=BLUE,width=3)
        if label:
            label_width=c.d.textlength(label,font=c.font(27))
            c.text((x+30+label_width/2,y+4),label,27,BLUE)
        port(x,provider_edge)
    # UML component notation with two small protruding component rectangles.
    def box(x,y,w,h,cid,zh,ru,interface):
        c.rect((x,y,x+w,y+h))
        c.rect((x+w-48,y+16,x+w-20,y+80),'white',INK,2)
        c.rect((x+w-60,y+27,x+w-35,y+42),'white',INK,2)
        c.rect((x+w-60,y+55,x+w-35,y+70),'white',INK,2)
        c.lines(x+w/2,y+30,['«component»',cid+'  '+zh,ru],29,43)
        c.text((x+w/2,y+h-27),'provides '+interface,27,BLUE)
    box(930,35,740,185,'C1','Web 界面','Веб-интерфейс','IWeb')
    cards=[(65,'C2','用户服务','Сервис пользователей','IUsers'),(705,'C3','图书目录服务','Сервис каталога','ICatalog'),(1345,'C4','活动服务','Сервис встреч','IMeetings'),(1985,'C5','评价服务','Сервис отзывов','IReviews')]
    for x,cid,zh,ru,interface in cards: box(x,385,550,185,cid,zh,ru,interface)
    for j,(x,cid,zh,ru,intf) in enumerate(cards):
        dest=x+275; src=1050+j*165; yy=245+j*23
        c.line([(src,220),(src,yy),(dest,yy),(dest,328)],BLUE,3)
        port(src,220)
        assembly(dest,350,385,intf)
    box(1800,780,735,185,'C6','通知服务','Сервис уведомлений','INotifications')
    c.line([(1620,570),(1620,680),(2167,680),(2167,723)],BLUE,3)
    port(1620,570)
    assembly(2167,745,780,'INotifications')
    box(330,1095,1240,190,'C7','数据库','База данных','IData')
    # Separate IData ports keep dependencies visible without joining lines.
    for x,target,y in [(340,460,720),(980,760,785),(1510,1060,850)]:
        c.line([(x,570),(x,y),(target,y),(target,1033)],BLUE,3)
        port(x,570)
        assembly(target,1055,1095,'')
        c.text((target+65,y+30),'IData',27,GRAY)
    c.line([(2450,570),(2450,710),(2570,710),(2570,1020),(1360,1020),(1360,1033)],BLUE,3)
    port(2450,570); assembly(1360,1055,1095,'')
    c.text((2350,1065),'IData',27,GRAY)
    c.line([(1800,895),(1670,895),(1670,985),(1480,985),(1480,1033)],BLUE,3)
    port(1800,895); assembly(1480,1055,1095,'')
    c.text((1580,945),'IData',27,GRAY)
    c.save('component.png')

ACTORS=[('参与者','Участник'),('Web 界面','Веб-интерфейс'),('活动服务','Сервис встреч'),('数据库','База данных'),('通知服务','Сервис уведомлений')]
XS=[160,680,1200,1720,2240]
def sequence_base(height):
    c=Canvas(2400,height)
    c.rect((2,2,2397,height-3),'white',BORDER,3)
    c.rect((2,2,725,40),'#F1F5F8',BORDER,2)
    c.text((365,21),'sd  报名活动 / Запись на встречу',25,bold=True)
    for x,(zh,ru) in zip(XS,ACTORS):
        c.rect((x-150,48,x+150,158))
        c.lines(x,80,[zh,ru],27,39)
        c.line([(x,158),(x,height-25)],GRAY,2,True)
    return c

def message(c,a,b,y,n,zh,ru,ret=False):
    x1,x2=XS[a],XS[b]
    c.arrow([(x1,y),(x2,y)],ret,GRAY if ret else BLUE,filled=not ret)
    mid=(x1+x2)/2
    # An opaque label background prevents lifelines from running through labels.
    for yy,txt in [(y-52,str(n)+'  '+zh),(y-23,ru)]:
        f=c.font(26,bool(re.search('[\u3400-\u9fff]',txt)))
        bb=c.d.textbbox((mid,yy),txt,font=f,anchor='mm')
        c.d.rectangle((bb[0]-7,bb[1]-3,bb[2]+7,bb[3]+3),fill='white')
        c.text((mid,yy),txt,26)

MESSAGES=[
    (0,1,1,'点击报名','Нажать «Записаться»',False),
    (1,2,2,'POST 报名请求','POST /meetings/{id}/registrations',False),
    (2,3,3,'开启事务并锁定活动','BEGIN; SELECT meeting FOR UPDATE',False),
    (3,2,4,'返回活动与名额','Вернуть встречу и места',True),
    (2,3,5,'检查该参与者是否已报名','Проверить запись участника',False),
    (3,2,6,'返回未报名且有名额','Записи нет, места есть',True),
    (2,3,7,'保存报名和通知任务并提交','INSERT registration + mail job; COMMIT',False),
    (3,2,8,'返回已提交的报名记录','Вернуть сохранённую запись',True),
    (2,4,9,'请求处理确认邮件任务','Обработать задание подтверждения',False),
    (4,2,10,'返回任务已受理','Задание принято',True),
    (2,1,11,'返回 201 和报名记录','201 Created + Registration',True),
    (1,0,12,'显示报名成功','Показать подтверждение записи',True),
]
def success_diagram():
    c=sequence_base(1430)
    # A UML critical combined fragment spans service/database transaction work.
    c.rect((1025,370,1910,1015),'white','#B9C7D3',2)
    c.rect((1025,370,1180,405),'#EEF3F6','#B9C7D3',2)
    c.text((1100,388),'critical',23,bold=True)
    c.text((1515,390),'检查与保存 / Проверки и запись',23)
    for x in XS[2:4]: c.line([(x,405),(x,1015)],GRAY,2,True)
    # Execution specifications (activation bars), as in the supplied reference.
    for actor,start,end in [(0,245,1375),(1,245,1375),(2,345,1275),(3,475,575),(3,675,775),(3,875,975),(4,1075,1175)]:
        x=XS[actor]; c.rect((x-8,start,x+8,end),'#D3E7F4',BLUE,2)
    for idx,row in enumerate(MESSAGES):
        message(c,row[0],row[1],245+idx*100+(30 if idx>=2 else 0),row[2],row[3],row[4],row[5])
    c.save('sequence-success.png')

def error_diagram():
    c=sequence_base(1460)
    c.rect((30,195,2370,1410),'white',BORDER,3)
    c.rect((30,195,145,238),'#EEF3F6',BORDER,2)
    c.text((85,216),'alt',25,bold=True)
    for x in XS: c.line([(x,240),(x,1400)],GRAY,2,True)
    groups=[(260,'活动不存在 / Встреча не найдена','404 MEETING_NOT_FOUND','显示活动不存在','Показать отсутствие встречи'),
            (655,'活动存在且已报名 / Встреча есть, участник записан','409 ALREADY_REGISTERED','显示已经报名','Показать повторную запись'),
            (1050,'活动存在且未报名且无名额 / Встреча есть, записи нет, мест нет','409 MEETING_FULL','显示名额已满','Показать отсутствие мест')]
    for idx,(y,title,code,zh,ru) in enumerate(groups):
        if idx: c.line([(30,y-35),(2370,y-35)],BORDER,2,True)
        c.text((1200,y), '['+title+']',27,bold=True)
        for actor,start,end in [(2,y+65,y+205),(1,y+205,y+310)]:
            x=XS[actor]; c.rect((x-8,start,x+8,end),'#D3E7F4',BLUE,2)
        message(c,3,2,y+100,'E'+str(idx+1),'返回对应检查结果','Вернуть результат проверки',True)
        message(c,2,1,y+205,'E'+str(idx+1),code,code,True)
        message(c,1,0,y+310,'E'+str(idx+1),zh,ru,True)
    c.save('sequence-errors.png')

doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11)
sec.top_margin=Inches(.62); sec.bottom_margin=Inches(.62)
sec.left_margin=Inches(.7); sec.right_margin=Inches(.7)
for sty in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    s=doc.styles[sty]
    s.font.name='Arial'; s.font.color.rgb=RGBColor.from_string('000000')
    rf=s.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in list(rf.attrib): del rf.attrib[attr]
    for attr,val in [('ascii','Arial'),('hAnsi','Arial'),('eastAsia','Arial Unicode MS'),('cs','Arial')]:
        rf.set(qn('w:'+attr),val)
    for border in s.element.xpath('./w:pPr/w:pBdr'): border.getparent().remove(border)
doc.styles['Normal'].font.size=Pt(10.5)
doc.styles['Normal'].paragraph_format.space_after=Pt(5)
doc.styles['Normal'].paragraph_format.line_spacing=1.08
for name,size in [('Title',21),('Subtitle',12),('Heading 1',15),('Heading 2',12)]:
    doc.styles[name].font.size=Pt(size)
    doc.styles[name].paragraph_format.space_after=Pt(7)
    doc.styles[name].paragraph_format.space_before=Pt(8 if 'Heading' in name else 0)
doc.core_properties.title='读书俱乐部架构与 API / Архитектура и API книжного клуба'
doc.core_properties.subject='组件图 时序图 OpenAPI / Компоненты последовательность OpenAPI'
doc.core_properties.author=''

INVENTORY=[]
def p(txt,style=None):
    INVENTORY.append(txt)
    return doc.add_paragraph(txt,style)
def heading(zh,ru,level=1): return p(bi(zh,ru),'Heading '+str(level))
def pair(zh,ru): return p(bi(zh,ru))
def table(headers,rows,widths,size=9.5):
    t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for cell,w in zip(t.columns,widths): cell.width=Inches(w)
    for j,h in enumerate(headers): t.cell(0,j).text=h
    for row in rows:
        cells=t.add_row().cells
        for j,txt in enumerate(row): cells[j].text=str(txt)
    for i,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
        if i==0: trpr.append(OxmlElement('w:tblHeader'))
        for j,cell in enumerate(row.cells):
            cell.width=Inches(widths[j]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr(); shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'244C69' if i==0 else 'EFF4F8' if i%2 else 'FFFFFF'); tcpr.append(shade)
            mar=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side); el.set(qn('w:w'),'80'); el.set(qn('w:type'),'dxa'); mar.append(el)
            tcpr.append(mar)
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'5'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
            tcpr.append(borders)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(2); para.paragraph_format.space_before=Pt(1); para.paragraph_format.line_spacing=1.02
                if j==0: para.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:
                    run.font.size=Pt(size); run.bold=i==0
                    if i==0: run.font.color.rgb=RGBColor(255,255,255)
    for row in [headers]+list(rows):
        for txt in row: INVENTORY.append(str(txt))
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t
def newpage(landscape=False):
    s=doc.add_section(WD_SECTION_START.NEW_PAGE)
    s.orientation=WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    s.page_width=Inches(11 if landscape else 8.5); s.page_height=Inches(8.5 if landscape else 11)
    s.top_margin=Inches(.55 if landscape else .62); s.bottom_margin=Inches(.55 if landscape else .62)
    s.left_margin=Inches(.65 if landscape else .7); s.right_margin=Inches(.65 if landscape else .7)
def figure(name,width,alt):
    para=doc.add_paragraph(); para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    inline=para.add_run().add_picture(str(WORK/name),width=Inches(width))
    inline._inline.docPr.set('descr',alt)
def code(txt):
    para=p(txt)
    para.paragraph_format.line_spacing=1.0
    para.paragraph_format.keep_together=True
    for r in para.runs: r.font.name='Courier New'; r.font.size=Pt(9)
    return para

def create_doc():
    p('读书俱乐部系统架构与接口设计','Title')
    p('Архитектура и проектирование API книжного клуба','Subtitle')
    p('版本 / Версия 1.0.0     日期 / Дата 2026-10-20     状态 / Статус DRAFT',None)
    pair('系统支持参与者注册、浏览图书、报名活动和发表评价；组织者添加图书并创建活动，系统发送邮件提醒。本文给出组件架构、报名时序及两个活动接口。', 'Система поддерживает регистрацию участников, просмотр книг, запись на встречи и отзывы; организатор добавляет книги и создаёт встречи, система отправляет напоминания по почте. Ниже представлены компоненты, последовательность записи и два API встреч.')
    heading('1 组件职责','1 Назначение компонентов')
    table([bi('编号','Код'),bi('组件','Компонент'),bi('用途','Назначение')],COMPONENTS,[.45,1.78,4.87],9.5)
    pair('组件表示逻辑职责，可部署在同一个服务器应用中；网站通过服务接口访问数据。', 'Компоненты обозначают логические обязанности и могут работать в одном серверном приложении; сайт обращается к данным через интерфейсы сервисов.')

    newpage(True)
    heading('1 组件图及接口','1 Диаграмма компонентов и интерфейсы')
    figure('component.png',9.7,'中俄双语 UML 组件图，7 个组件，依赖箭头由调用方指向提供接口的组件 / Диаграмма 7 компонентов и их зависимостей')
    pair('图 1  外框为系统边界；圆形为提供接口，半圆为需求接口，小方框为端口，配对连接为接口装配。provides 表示提供接口；IData 由 C2 至 C6 使用。', 'Рисунок 1  Рамка — граница системы; круг — предоставляемый интерфейс, полукруг — требуемый, квадрат — порт, парное соединение — сборка интерфейсов. provides означает предоставление; IData используют C2–C6.')

    newpage(False)
    heading('1 接口与依赖说明','1 Интерфейсы и зависимости')
    table([bi('接口','Интерфейс'),bi('提供方到使用方','Поставщик и потребитель'),bi('操作职责','Операции')],INTERFACES,[1.05,1.6,4.45],10)
    pair('Web 界面依赖用户、目录、活动和评价四个服务；活动服务依赖通知服务与数据库；用户、目录、评价和通知服务也依赖数据库。', 'Веб-интерфейс зависит от сервисов пользователей, каталога, встреч и отзывов; сервис встреч зависит от уведомлений и базы данных; сервисы пользователей, каталога, отзывов и уведомлений также зависят от базы данных.')
    newpage(False)
    heading('2 报名场景与规则','2 Сценарий записи и правила')
    pair('前置条件：参与者已注册并已选择活动，客户端持有 participantId；Web 界面提交活动 id 和参与者编号。', 'Предусловия: участник зарегистрирован и выбрал встречу, клиент располагает participantId; веб-интерфейс передаёт id встречи и идентификатор участника.')
    pair('成功条件：活动存在、参与者尚未报名且有剩余名额。活动服务在同一事务中锁定活动、检查重复和名额、创建报名及邮件任务；数据库对 (meetingId, participantId) 设置唯一约束。', 'Условия успеха: встреча существует, участник ещё не записан и есть места. Сервис в одной транзакции блокирует встречу, проверяет повтор и места, создаёт запись и задание письма; база задаёт уникальность (meetingId, participantId).')
    pair('剩余名额按 capacity 减去当前报名数计算，并在锁定活动后使用事务内最新数据检查。', 'Свободные места равны capacity минус текущее число записей; после блокировки встречи проверяются актуальные данные внутри транзакции.')
    pair('提交后通知服务受理邮件任务。报名成功表示记录已保存；邮件实际送达由通知服务处理，失败时保留任务重试，不能把邮件受理当作送达。', 'После фиксации сервис уведомлений принимает задание. Успех записи означает сохранение данных; доставкой занимается сервис уведомлений, при сбое задание сохраняется для повтора, приём задания не равен доставке.')
    pair('错误优先级：活动不存在返回 404；已报名返回 409 ALREADY_REGISTERED；否则名额已满返回 409 MEETING_FULL。失败时结束事务，不新增报名和通知任务。', 'Приоритет ошибок: отсутствующая встреча даёт 404; повторная запись — 409 ALREADY_REGISTERED; иначе отсутствие мест — 409 MEETING_FULL. При отказе транзакция завершается без новой записи и задания уведомления.')
    pair('时序图读法：时间自上而下，sd 表示交互框架；竖向虚线为生命线，细长矩形为激活条；实线实心箭头表示同步调用，虚线开放箭头表示返回。', 'Чтение диаграмм: время идёт сверху вниз, sd задаёт рамку взаимодействия; вертикальный пунктир — линия жизни, узкий прямоугольник — активация; сплошная стрелка с закрашенным концом — синхронный вызов, пунктир с открытым концом — ответ.')

    newpage(True)
    heading('2 报名成功时序图','2 Диаграмма успешной записи')
    figure('sequence-success.png',9.7,'五条生命线和十二步报名成功消息；数据库事务包含检查、保存和提交 / Пять линий жизни и двенадцать сообщений успешной записи')
    pair('图 2a  报名成功的 12 条消息；critical 表示检查和保存构成原子事务，最后返回 201 和报名记录。', 'Рисунок 2a  Двенадцать сообщений успешной записи; critical объединяет проверки и сохранение в атомарную транзакцию, итоговый ответ — 201 и запись.')

    newpage(True)
    heading('2 报名失败替代分支','2 Альтернативные ветви отказа')
    figure('sequence-errors.png',9.1,'报名失败 alt 组合片段：404 活动不存在，409 已报名，409 无名额 / Альтернативы 404 и двух вариантов 409')
    pair('图 2b  承接图 2a 的检查阶段，alt 中三个条件互斥；进入任一分支即终止成功流程，不执行保存和发送通知。通知服务生命线保留，但没有消息。', 'Рисунок 2b  Продолжение проверок рисунка 2a; три условия alt взаимоисключающие. Любая ветвь завершает сценарий без сохранения и уведомления. Линия сервиса уведомлений остаётся без сообщений.')

    newpage(False)
    heading('3 OpenAPI 接口定义','3 Описание API в OpenAPI')
    pair('规范版本为 OpenAPI 3.0.3，接口版本为 1.0.0；请求及响应使用 application/json。独立文件 book-club.openapi.yaml 包含完整定义和双语说明。', 'Версия спецификации — OpenAPI 3.0.3, версия API — 1.0.0; запросы и ответы используют application/json. Полное определение с двуязычными описаниями находится в book-club.openapi.yaml.')
    heading('3 1 获取活动列表','3 1 Получение списка встреч',2)
    code('GET /meetings?from=2026-11-01')
    table([bi('参数及位置','Параметр и место'),bi('类型及必填性','Тип и обязательность'),bi('含义','Значение')],[('from · query',bi('string date\n可选','Необязателен'),bi('按 UTC 日期从该日开始筛选，包含该日。','Отбор с указанной даты UTC включительно.'))],[1.35,1.6,4.15],10)
    pair('不传 from 返回全部活动；结果按 startsAt 和 id 升序排列。没有符合条件的活动时返回 200 和 []，列表为空不返回 404。from 日期格式错误返回 400 INVALID_REQUEST。', 'Без from возвращаются все встречи, по возрастанию startsAt и id. Если подходящих встреч нет, возвращаются 200 и [], а не 404. Некорректная дата from даёт 400 INVALID_REQUEST.')
    heading('3 2 创建活动报名','3 2 Создание записи на встречу',2)
    code('POST /meetings/12/registrations\nContent-Type: application/json\n\n{"participantId": 7}')
    table([bi('输入','Вход'),bi('类型及必填性','Тип и обязательность'),bi('含义','Значение')],[('id · path',bi('integer ≥ 1\n必填','Обязателен'),bi('活动编号，对应路径中的 {id}。','Идентификатор встречи из {id}.')),('participantId · body',bi('integer ≥ 1\n必填','Обязателен'),bi('已注册参与者编号，请求体为 RegistrationRequest。','Идентификатор зарегистрированного участника; тело RegistrationRequest.'))],[1.35,1.6,4.15],10)
    pair('设计约定：from 使用 UTC 日期；请求体必须提供 participantId；错误响应增加 400 处理格式或类型错误；无名额采用 409 MEETING_FULL。', 'Проектные допущения: from использует дату UTC; тело содержит participantId; добавлен ответ 400 для ошибок формата или типа; отсутствие мест даёт 409 MEETING_FULL.')

    newpage(False)
    heading('3 响应与数据结构','3 Ответы и структуры данных')
    table([bi('接口和状态','API и статус'),bi('结果','Результат')],[
        ('GET · 200',bi('返回 Meeting 数组，允许为空。','Массив Meeting, возможно пустой.')),
        ('POST · 201',bi('返回新建的 Registration，报名已保存。','Новая Registration; запись сохранена.')),
        ('POST · 404',bi('MEETING_NOT_FOUND，活动不存在。','MEETING_NOT_FOUND, встреча не найдена.')),
        ('POST · 409',bi('ALREADY_REGISTERED 已报名；MEETING_FULL 无名额。','ALREADY_REGISTERED — повтор; MEETING_FULL — мест нет.')),
        ('GET POST · 400',bi('INVALID_REQUEST，参数或请求体无效。','INVALID_REQUEST, неверные параметры или тело.')),
    ],[1.45,5.65],9.5)
    table([bi('模式','Схема'),bi('必填字段及类型','Обязательные поля и типы')],[
        ('Meeting',bi('id、bookId：正整数；title：非空字符串；startsAt：date-time；capacity：正整数；remainingSeats：非负整数且不大于 capacity。','id, bookId: положительные целые; title: непустая строка; startsAt: date-time; capacity: положительное целое; remainingSeats: целое от 0 до capacity.')),
        ('RegistrationRequest',bi('participantId：正整数。','participantId: положительное целое.')),
        ('Registration',bi('id、meetingId、participantId：正整数；createdAt：date-time。','id, meetingId, participantId: положительные целые; createdAt: date-time.')),
        ('Error',bi('code：枚举错误代码；message：非空错误说明。','code: код из перечисления; message: непустое описание.')),
    ],[1.45,5.65],9.5)
    pair('所有模式拒绝未声明字段。时间戳包含时区；示例使用 UTC 的 Z 后缀。', 'Все схемы запрещают необъявленные поля. Метки времени содержат часовой пояс; примеры используют суффикс Z для UTC.')
    newpage(False)
    heading('3 JSON 响应示例','3 Примеры ответов JSON')
    p(bi('200 活动列表示例','Пример списка встреч с кодом 200'),'Heading 2')
    code(json.dumps([meeting],ensure_ascii=False,indent=2))
    p(bi('201 成功响应示例','Пример успешного ответа 201'),'Heading 2')
    code(json.dumps(registration,ensure_ascii=False,indent=2))
    p(bi('409 重复报名响应示例','Пример ответа 409 при повторной записи'),'Heading 2')
    code(json.dumps(errors['alreadyRegistered'],ensure_ascii=False,indent=2))
    pair('活动不存在时使用 MEETING_NOT_FOUND；无名额时使用 MEETING_FULL。完整示例和模式约束见配套 YAML。', 'Для отсутствующей встречи используется MEETING_NOT_FOUND, для отсутствия мест — MEETING_FULL. Полные примеры и ограничения схем приведены в YAML.')
    doc.save(DOCX)

def main():
    component_diagram(); success_diagram(); error_diagram()
    YAML.write_text(yaml.safe_dump(SPEC,allow_unicode=True,sort_keys=False,width=110),encoding='utf-8')
    create_doc()
    inventory={'components':COMPONENTS,'interfaces':INTERFACES,'success_messages':MESSAGES,'text':INVENTORY,'source_coverage':{
        'image1':'OpenAPI 格式参考；采用 3.0.3；不额外创建 /books 接口。',
        'image2':'三个任务；7 个组件、接口和依赖；5 个时序参与者、消息与返回、无名额分支；两个接口及 200/201/404/409。',
        'image3':'读书俱乐部业务背景和三项正式任务。',
        'image4':'补充组件图样式参考；系统边界、组件图标、端口及提供与需求接口装配；不引入商城业务。',
        'image5':'补充时序图样式参考；sd 框、对象头、生命线、激活条、编号消息和虚线返回；不添加无业务依据的循环或销毁。',
    }}
    (BASE/'data/content.json').write_text(json.dumps(inventory,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'docx':str(DOCX),'yaml':str(YAML),'components':7,'embedded_figures':3,'sections':len(doc.sections),'tables':len(doc.tables)},ensure_ascii=False))

if __name__=='__main__': main()
