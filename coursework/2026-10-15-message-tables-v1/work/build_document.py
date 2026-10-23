from pathlib import Path
from datetime import datetime, timezone
from zipfile import ZipFile
import hashlib
import json
import re
import shutil
from lxml import etree
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path('/Users/zhi/Documents/软件生命周期/studyspace')
OUT=ROOT/'coursework/2026-10-15-message-tables-v1'
DATA=OUT/'data'
WORK=OUT/'work'
for p in [OUT,DATA,WORK]:p.mkdir(parents=True,exist_ok=True)
BASE=ROOT/'coursework/2026-10-15-sequence-diagrams-v1'
SOURCE=BASE/'data/content.json'
source=json.loads(SOURCE.read_text())
STEM='StudySpace_任务三消息表_中俄双语_v1'
DOCX=OUT/(STEM+'.docx')
sources=[SOURCE,BASE/'StudySpace_双场景时序图与说明_中俄双语_v1.docx',BASE/'StudySpace_双场景时序图_中俄双语_v1.drawio',ROOT/'docs/zh-CN/04_System_Design.md',ROOT/'docs/ru-RU/04_System_Design.md',ROOT/'docs/data/project-decisions.json',ROOT/'AGENTS.md']
fingerprints={str(p):hashlib.sha256(p.read_bytes()).hexdigest() for p in sources}
photo=Path('/var/folders/ff/kqw394c93dd7krl0z2hcqqdh0000gn/T/codex-clipboard-d46572de-1dfa-4a34-8d82-a4e98b16fc3e.jpg')
shutil.copy2(photo,DATA/'classroom-task-3.jpg')

# Every source arrow has one row; data and result specify the same proposed design.
BOOK={
'1':(('seat_id；D；带时区 start/end','seat_id, D; start/end с часовым поясом'),('界面接收座位与时段','Интерфейс принимает выбор')),
'2':(('会话；seat_id；D；start/end','Сессия, seat_id, D, start/end'),('提交预约申请，进入服务端校验','Заявка передана на проверку')),
'3':(('会话标识；所需 STUDENT 角色','Сессия; требуемая роль STUDENT'),('授权上下文见消息 4','Контекст доступа в сообщении 4')),
'4':(('服务端 user_id；角色与权限','Серверный user_id; роль и права'),('预约组件取得可信操作者身份','Получен доверенный автор заявки')),
'5':(('user_id；room_id；seat_id；tx','user_id, room_id, seat_id, tx'),('按 User → Room → Seat 加锁','Блокировки User → Room → Seat')),
'6':(('共享事务 tx；锁后实际 now','Общая tx; фактическое now после блокировок'),('后续校验使用同一事务与时刻','Дальше используются те же tx и время')),
'7':(('tx；user_id；账号当前状态','tx, user_id; текущее состояние аккаунта'),('复核启用、STUDENT、VERIFIED；见 8','Проверить активность, STUDENT, VERIFIED; см. 8')),
'8':(('资格是否通过；失败原因','Результат допуска; причина отказа'),('记录资格结果，供最终分支判断','Сохранить допуск для выбора ветки')),
'9':(('tx；seat_id；D；start/end；now','tx, seat_id, D, start/end, now'),('检查启停、开放、粒度、日期与时长','Проверить ресурс, часы, шаг, дату и срок')),
'10':(('资源与时间是否合法；原因码','Допустимость ресурса и времени; причина'),('记录校验结果，供最终分支判断','Сохранить результат для выбора ветки')),
'11':(('tx；user_id；seat_id；[start,end)','tx, user_id, seat_id, [start,end)'),('读取同学生或同座位的未取消重叠','Найти неотменённые пересечения студента или места')),
'12':(('学生重叠／座位重叠／无重叠','Пересечение студента, места либо отсутствие'),('校验通过时：冲突走 b，无冲突走 c','При верных проверках: конфликт — b, иначе c')),
'13a':(('tx；资格或资源时间校验失败','tx; отказ допуска либо неверный ресурс/время'),('回滚，不创建预约','Откат без создания записи')),
'14a':(('403：资格不符；422：资源或时间非法','403: нет допуска; 422: неверный ресурс или время'),('界面展示对应原因','Интерфейс показывает причину')),
'13b':(('tx；已确认的学生或座位重叠','tx; установленное пересечение студента или места'),('回滚，座位不被本申请占用','Откат, заявка не занимает место')),
'14b':(('HTTP 409；STUDENT_OVERLAP\n或 SEAT_OVERLAP','HTTP 409; STUDENT_OVERLAP\nлибо SEAT_OVERLAP'),('界面按冲突类型提示重新选择','Предложить выбор по типу конфликта')),
'15b':(('冲突类型；重新选择提示','Тип конфликта; предложение нового выбора'),('学生可更换时段后重新查询','Студент выбирает другой интервал')),
'13c':(('tx；user_id；单个 seat_id；D；start/end','tx, user_id, один seat_id, D, start/end'),('生成一条待提交预约；见 14c','Создана одна незафиксированная запись; см. 14c')),
'14c':(('booking_id；未提交状态','booking_id; ещё не зафиксировано'),('不能据此提前向界面报告成功','Рано возвращать успех интерфейсу')),
'15c':(('tx；booking_id；actor_id；创建动作','tx, booking_id, actor_id; создание'),('审计组件沿用原事务','Аудит использует исходную tx')),
'16c':(('tx；预约 ID；操作者；event_at；动作','tx, ID записи, автор, event_at, действие'),('事件待提交，无独立 COMMIT','Событие ожидает общего COMMIT')),
'17c':(('同一 tx 中的预约与创建事件','Запись и событие в одной tx'),('原子提交；确认见 18c','Атомарная фиксация; ответ в 18c')),
'18c':(('事务提交成功的结果','Результат успешного COMMIT'),('此后才可返回创建成功','Теперь допустим ответ об успехе')),
'19c':(('HTTP 201；booking_id；seat；start/end','HTTP 201, booking_id, seat, start/end'),('界面取得已提交预约','Интерфейс получает сохранённую запись')),
'20c':(('预约编号；座位；确认时段','ID брони, место, подтверждённое время'),('学生看到预约确认','Студент видит подтверждение')),
}
REVIEW={
'1':(('待审核目标 user_id','user_id проверяемого студента'),('界面请求待审核记录','Интерфейс запрашивает запись')),
'2':(('管理员会话；PENDING 过滤；目标 user_id','Сессия администратора, PENDING, user_id'),('请求必要审核信息；见 7','Запрошены сведения; ответ в 7')),
'3':(('管理员会话；所需管理权限','Сессия; требуемые права администратора'),('授权身份与权限见 4','Автор и права в сообщении 4')),
'4':(('管理员 ID；角色；当前管理权限','ID администратора, роль и права'),('授权后读取目标身份','После авторизации читать данные')),
'5':(('目标 user_id；必要字段范围','user_id; минимальный набор полей'),('读取身份、当前状态和版本；见 6','Прочитать данные, статус, версию; см. 6')),
'6':(('必要学号、姓名；PENDING；revision','Необходимые номер, имя; PENDING, revision'),('形成待审核身份快照','Получен снимок для проверки')),
'7':(('必要身份信息；状态；revision','Необходимые данные, статус, revision'),('界面接收审核资料','Интерфейс получает сведения')),
'8':(('学号、姓名；PENDING；快照版本','Номер, имя, PENDING, версия снимка'),('管理员看到待审核内容','Администратор видит запись')),
'9':(('必要学号、姓名；学校认可渠道 IF-04','Необходимые номер, имя; признанный канал IF-04'),('人工比对，未调用自动学校 API','Ручная сверка без автоматического API')),
'10':(('相符／不符／无法核实；来源依据','Совпадение, несоответствие либо отсутствие проверки; основание'),('管理员取得结果；缺来源不得批准','Получен итог; без источника подтверждение запрещено')),
'11':(('拟定决定；来源类型、引用；revision','Решение, тип и ссылка источника, revision'),('界面整理审核申请','Интерфейс формирует заявку')),
'12':(('user_id；decision；source_type/ref；revision','user_id, decision, source_type/ref, revision'),('仅提交申请，尚未变更状态','Заявка отправлена, статус ещё не изменён')),
'13':(('保存时的管理员会话','Сессия администратора при сохранении'),('再次核对当前管理员身份；见 14','Повторная проверка автора; см. 14')),
'14':(('当前管理员 ID 与会话身份','Текущий ID и личность администратора'),('继续事务内的最终权限复核','Продолжить проверку прав в tx')),
'15':(('tx；相关账号 ID；固定锁序','tx, ID аккаунтов, единый порядок блокировок'),('取得锁，准备核对当前状态','Взять блокировки для сверки состояния')),
'16':(('tx；当前权限；当前状态；revision','tx; текущие права, статус, revision'),('按权限、来源及快照选择 a/b/c/d','Выбрать a/b/c/d по правам, источнику и снимку')),
'17a':(('tx；无权限或来源不可核验','tx; нет прав либо источник не проверен'),('回滚，审核状态保持不变','Откат, статус не изменяется')),
'18a':(('HTTP 403 或 422；对应原因码','HTTP 403 либо 422; код причины'),('不授予 VERIFIED；缺来源保持 PENDING','Без VERIFIED; без источника остаётся PENDING')),
'17b':(('tx；权限来源有效，但身份或状态变化','tx; права и источник верны, данные или статус изменились'),('回滚，拒绝保存过期决定','Откат, устаревшее решение не сохранено')),
'18b':(('HTTP 409；STALE_REVIEW','HTTP 409; STALE_REVIEW'),('界面要求刷新后重新审核','Обновить сведения и повторить проверку')),
'17c':(('tx；user_id；decision=VERIFIED；来源引用','tx, user_id, decision=VERIFIED, ссылка источника'),('可信核验相符，写入待提交状态','Сведения совпали; статус ещё не зафиксирован')),
'17d':(('tx；user_id；decision=REJECTED；来源引用','tx, user_id, decision=REJECTED, ссылка источника'),('可信核验确实不符，写入待提交状态','Достоверное несоответствие; статус не зафиксирован')),
'18':(('tx；user_id；审核人；event_at；决定；来源引用','tx, user_id, автор, event_at, решение, ссылка'),('仅 c/d 执行；准备最少必要事件','Только после c/d; подготовить минимальное событие')),
'19':(('tx；审核事件必要字段','tx; необходимые поля события проверки'),('不复制学号、姓名或证明影像','Без копий номера, имени и изображений')),
'20':(('同一 tx 的状态、决定与审核事件','Статус, решение и событие одной tx'),('原子提交；结果见 21','Атомарная фиксация; ответ в 21')),
'21':(('事务提交成功的结果','Результат успешного COMMIT'),('此后才可返回审核保存成功','Теперь допустим ответ о сохранении')),
'22':(('HTTP 200；VERIFIED 或 REJECTED','HTTP 200; VERIFIED либо REJECTED'),('界面展示已保存决定；拒绝不授予预约资格','Показать решение; REJECTED не даёт допуск')),
}
PARTICIPANTS={
'学生':('学生','Студент'),'管理员':('管理员','Администратор'),
'CMP-001':('CMP-001 界面','Интерфейс'),
'CMP-002':('CMP-002 认证','Доступ'),
'CMP-003':('CMP-003 资格','Проверка вуза'),
'CMP-004':('CMP-004 资源','Ресурсы'),
'CMP-006':('CMP-006 预约','Бронирование'),
'CMP-008':('CMP-008 审计','Аудит'),
'CMP-009':('CMP-009 数据','Данные'),
'CMP-010':('CMP-010 来源','Источник вуза'),
}
SCENARIOS=[
    {'id':'US-002','zh':'学生创建单座位预约','ru':'Бронирование одного места студентом','map':BOOK,'panels':source['panels'][:2]},
    {'id':'US-004','zh':'管理员审核校内身份','ru':'Проверка принадлежности к вузу','map':REVIEW,'panels':source['panels'][2:]},
]
HEADERS=[('序号','№'),('发送方','От'),('接收方','Кому'),('操作或消息','Операция / сообщение'),('数据','Данные'),('响应或结果','Ответ / результат')]
for s in SCENARIOS:
    s['messages']=[m for p in s['panels'] for m in p['messages']]
    assert {m['number'] for m in s['messages']}==set(s['map'])
    s['rows']=[[(m['number'],),PARTICIPANTS[m['src']],PARTICIPANTS[m['dst']],(m['zh'],m['ru']),*s['map'][m['number']]] for m in s['messages']]

doc=Document()
sec=doc.sections[0]
sec.orientation=WD_ORIENT.LANDSCAPE
sec.page_width,sec.page_height=Inches(11),Inches(8.5)
sec.top_margin=sec.bottom_margin=Inches(.48)
sec.left_margin=sec.right_margin=Inches(.55)
for name in ['Normal','Title','Heading 1','Heading 2']:
    st=doc.styles[name]
    st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
    rf=st.element.get_or_add_rPr().get_or_add_rFonts()
    for k,v in [('ascii','Arial'),('hAnsi','Arial'),('eastAsia','Heiti SC')]:rf.set(qn('w:'+k),v)
    for el in list(st.element.iter(qn('w:pBdr'))):el.getparent().remove(el)
doc.styles['Normal'].font.size=Pt(10.5)
doc.styles['Normal'].paragraph_format.space_after=Pt(4)
doc.styles['Normal'].paragraph_format.line_spacing=1.04
doc.styles['Title'].font.size=Pt(18)
doc.styles['Title'].paragraph_format.space_after=Pt(3)
doc.styles['Heading 1'].font.size=Pt(13)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(8)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(3)
MD=[]
def run(p,text,size=None,bold=False,color='000000',lang='zh-CN'):
    r=p.add_run(text);r.font.name='Arial';r.bold=bold;r.font.color.rgb=RGBColor.from_string(color)
    if size:r.font.size=Pt(size)
    rp=r._element.get_or_add_rPr();rp.rFonts.set(qn('w:eastAsia'),'Heiti SC')
    le=OxmlElement('w:lang');le.set(qn('w:val'),lang);rp.append(le)
    return r
def pair(zh,ru,style=None,size=None):
    p=doc.add_paragraph(style=style);run(p,zh,size);p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(2)
    p=doc.add_paragraph(style=style);run(p,ru,size,lang='ru-RU');p.paragraph_format.keep_together=True;p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(4)
    MD.extend([zh,ru,''])
    return p
def add_table(s):
    t=doc.add_table(rows=1,cols=6);t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
    widths=[.46,1.12,1.12,2.72,2.48,2.0]
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    bd=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge)
        for k,v in [('val','single'),('sz','5'),('color','D9D9D9')]:e.set(qn('w:'+k),v)
        bd.append(e)
    t._tbl.tblPr.append(bd)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for ri,rowvals in enumerate([HEADERS,*s['rows']]):
        row=t.rows[0] if ri==0 else t.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for ci,(c,texts) in enumerate(zip(row.cells,rowvals)):
            c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
            for edge,val in [('top',85),('bottom',85),('left',75),('right',75)]:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),str(val));e.set(qn('w:type'),'dxa');mar.append(e)
            pr.append(mar)
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'29485D' if ri==0 else ('F1F5F8' if ri%2 else 'FFFFFF'));pr.append(sh)
            for li,text in enumerate(texts):
                p=c.paragraphs[0] if li==0 else c.add_paragraph()
                p.paragraph_format.space_after=Pt(3 if li==0 and len(texts)>1 else 0)
                p.paragraph_format.line_spacing=1.02
                p.paragraph_format.keep_with_next=False
                nowrap=OxmlElement('w:wordWrap');nowrap.set(qn('w:val'),'0');p._p.get_or_add_pPr().append(nowrap)
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER if ri==0 or ci==0 else WD_ALIGN_PARAGRAPH.LEFT
                run(p,text,9.5,ri==0,'FFFFFF' if ri==0 else '000000','zh-CN' if li==0 else 'ru-RU')
    MD.append('| '+' | '.join('<br>'.join(h) for h in HEADERS)+' |')
    MD.append('| '+' | '.join('---' for _ in HEADERS)+' |')
    MD.extend('| '+' | '.join('<br>'.join(c) for c in row)+' |' for row in s['rows'])
    MD.append('')
    return t

pair('StudySpace 任务三消息表','StudySpace Задание 3 Таблицы сообщений','Title')
p=doc.add_paragraph();run(p,'SS-SEQ-MSG-01  |  v1  |  2026-10-15  |  DRAFT',9)
pair('六列消息表对应 SS-SEQ-01 v1 的全部 52 个箭头，含返回消息；中文在上、俄语在下，结果均为设计预期，业务测试 NOT_RUN。',
     'Шесть колонок охватывают все 52 стрелки SS-SEQ-01 v1, включая возвраты. Китайский над русским; результаты проектные, бизнес-тесты NOT_RUN.',size=10)
pair('1 US-002 学生创建单座位预约','1 US-002 Бронирование одного места студентом','Heading 1')
pair('a 为资格或资源时间失败；b 为校验通过但存在学生或座位重叠；c 为全部通过且无重叠。三分支互斥，只有 c 写入预约和事件。',
     'a — нет допуска либо неверны ресурс/время; b — проверки пройдены, но есть пересечение; c — проверки пройдены без пересечений. Ветки взаимоисключающие, только c создаёт запись и событие.',size=9.5)
add_table(SCENARIOS[0])
pair('2 US-004 管理员审核校内身份','2 US-004 Проверка принадлежности к вузу','Heading 1')
pair('a 为无权限或来源不可核验；b 为权限与来源有效但快照变化；有效且快照一致时，c 为可信核验相符，d 为真实核验不符。步骤 18–22 仅在 c 或 d 后执行；CMP-010 由人工查询，OPEN-003 仍待确认。',
     'a — нет прав либо источник не проверен; b — права и источник верны, но снимок изменился. При верном снимке c означает совпадение, d — достоверное несоответствие. Шаги 18–22 выполняются только после c/d; CMP-010 проверяется вручную, OPEN-003 не решён.',size=9.5)
add_table(SCENARIOS[1])
pair('3 共用约束与追溯','3 Общие ограничения и трассируемость','Heading 1')
pair('返回行的最后一列说明接收后的效果，不代表再次响应。表中 API、错误码与 revision 沿用拟定契约，尚未实现。',
     'В строке возврата последний столбец описывает эффект получения, а не новый ответ. API, коды ошибок и revision — предлагаемые, ещё не реализованные контракты.',size=9.5)
pair('US-002 沿用 [start,end) 与 DEC-016：同学生或同座位重叠拒绝，相邻与多条非重叠预约允许。资格、锁后时刻、资源时间、冲突检查和写入属于同一服务端事务。\nPostgreSQL 锁序为 User → Room → Seat；SQLite 使用 BEGIN IMMEDIATE。',
     'US-002 следует [start,end) и DEC-016: пересечения студента или места запрещены, соседние и несколько непересекающихся записей допустимы. Допуск, время после блокировок, ресурсы, конфликты и запись проверяются в одной tx. PostgreSQL: User → Room → Seat; SQLite: BEGIN IMMEDIATE.',size=9.5)
pair('US-004 无可信来源不授予 VERIFIED，本场景保持 PENDING；只有真实核验不符才走 REJECTED。审核事件不复制原始学号、姓名或证明影像，状态与事件一起提交。锁超时与未知数据库异常不能一律返回 409；失败应回滚。',
     'US-004 без доверенного источника не даёт VERIFIED, здесь остаётся PENDING; REJECTED требует достоверного несоответствия. Событие не копирует номер, имя или изображения; статус и событие фиксируются вместе. Тайм-аут и неизвестные сбои БД не всегда 409; ошибка требует отката.',size=9.5)
pair('依据为课堂任务三照片、SS-SEQ-01 v1 时序图、既有 CMP 组件与 DOC-04 0.3.0。\nUS-002：SP-FR-008、SP-NFR-004；US-004：SP-FR-022、SP-IF-002、SP-DR-001。\n执行环境为本地 macOS 与 Codex 内置 Python、文档渲染工具。',
     'Основания: фото задания 3, схемы SS-SEQ-01 v1, прежние CMP и DOC-04 0.3.0.\nUS-002: SP-FR-008, SP-NFR-004; US-004: SP-FR-022, SP-IF-002, SP-DR-001.\nСреда: локальная macOS, встроенные Python и рендеринг Codex.',size=9)

doc.core_properties.title='StudySpace 任务三消息表 Таблицы сообщений'
doc.core_properties.subject='US-002 US-004 六列消息表 / Таблицы сообщений из шести колонок'
doc.core_properties.author='StudySpace'
doc.save(DOCX)
(OUT/(STEM+'.md')).write_text('\n\n'.join(MD),encoding='utf-8')
saved={'doc_id':'SS-SEQ-MSG-01','version':'v1','status':'DRAFT','date':'2026-10-15','business_tests':'NOT_RUN','headers':HEADERS,'scenarios':[{k:v for k,v in s.items() if k not in {'map','panels'}} for s in SCENARIOS],'source_fingerprints':fingerprints}
(DATA/'content.json').write_text(json.dumps(saved,ensure_ascii=False,indent=2),encoding='utf-8')
shutil.copy2(BASE/'data/fonts.conf',DATA/'fonts.conf')
rd=Document(DOCX)
with ZipFile(DOCX) as z:
    xml=etree.fromstring(z.read('word/document.xml'));names=z.namelist()
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
checks={
 'two_scenarios_two_tables':len(rd.tables)==2,
 'six_original_columns':all(len(t.columns)==6 for t in rd.tables),
 'headers_match_photo':all(t.cell(0,i).text=='\n'.join(h) for t in rd.tables for i,h in enumerate(HEADERS)),
 'all_source_arrows_covered_once':all(len(t.rows)-1==len(s['messages']) and [t.cell(i+1,0).text for i in range(len(s['messages']))]==[m['number'] for m in s['messages']] for t,s in zip(rd.tables,SCENARIOS)),
 'messages_and_directions_match_source':all(t.cell(i+1,1).text=='\n'.join(PARTICIPANTS[m['src']]) and t.cell(i+1,2).text=='\n'.join(PARTICIPANTS[m['dst']]) and t.cell(i+1,3).text=='\n'.join((m['zh'],m['ru'])) for t,s in zip(rd.tables,SCENARIOS) for i,m in enumerate(s['messages'])),
 'all_data_and_results_filled':all(all(c.text.strip() for c in row.cells) for t in rd.tables for row in t.rows),
 'chinese_above_russian':all(len(c.paragraphs)==2 and (re.search('[А-Яа-яЁё]',c.paragraphs[1].text) or re.sub(r'[\s；，、;,]','',c.paragraphs[0].text)==re.sub(r'[\s；，、;,]','',c.paragraphs[1].text)) for t in rd.tables for row in t.rows[1:] for c in row.cells[1:]),
 'stable_branch_ids_kept':all(v in [row.cells[0].text for row in rd.tables[0].rows] for v in ['13a','14a','13b','14b','13c','20c']) and all(v in [row.cells[0].text for row in rd.tables[1].rows] for v in ['17a','18a','17b','18b','17c','17d','22']),
 'no_added_component_id':all(m['src'] in PARTICIPANTS and m['dst'] in PARTICIPANTS for s in SCENARIOS for m in s['messages']),
 'repeating_headers':len(xml.xpath('//w:tbl/w:tr[1]/w:trPr/w:tblHeader',namespaces=ns))==2,
 'rows_do_not_split':len(xml.xpath('//w:tr/w:trPr/w:cantSplit',namespaces=ns))==sum(len(t.rows) for t in rd.tables),
 'no_fixed_row_heights':not xml.xpath('//w:trHeight',namespaces=ns),
 'tables_are_editable_not_images':len(rd.inline_shapes)==0,
 'all_sources_preserved':all(hashlib.sha256(Path(p).read_bytes()).hexdigest()==h for p,h in fingerprints.items()),
 'no_headers_or_footers':not any(re.fullmatch(r'word/(header|footer)\d+\.xml',n) for n in names),
 'manual_source_and_not_run_retained':all(v in '\n'.join(p.text for p in rd.paragraphs) for v in ['OPEN-003','NOT_RUN','人工查询','拟定契约']),
}
report={'document':str(DOCX),'checks':checks,'check_count':len(checks),'all_pass':all(checks.values()),'visual_review':'PENDING','row_counts':{s['id']:len(s['messages']) for s in SCENARIOS},'business_tests':'NOT_RUN','source_fingerprints':fingerprints,'docx_sha256':hashlib.sha256(DOCX.read_bytes()).hexdigest(),'timestamp_utc':datetime.now(timezone.utc).isoformat(),'environment':'Local macOS; bundled Codex Python; document validation only'}
(DATA/'verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
(OUT/'README.md').write_text(f'''# StudySpace 任务三消息表\n# StudySpace Задание 3 Таблицы сообщений\n\nSS-SEQ-MSG-01 · v1 · 2026-10-15 · DRAFT\n\n[{STEM}.docx]({STEM}.docx)\n\n按课堂六列模板逐条覆盖既有两场景时序图，中文在上，俄语在下。\nШесть колонок покрывают все стрелки двух сценариев; китайский над русским.\n\nUS-002 {len(SCENARIOS[0]['messages'])} 条；US-004 {len(SCENARIOS[1]['messages'])} 条。\nUS-002 {len(SCENARIOS[0]['messages'])} сообщений; US-004 {len(SCENARIOS[1]['messages'])} сообщений.\n\n内容 data/content.json；检查 data/verification.json；业务测试 NOT_RUN。\nСодержание data/content.json; проверка data/verification.json; бизнес-тесты NOT_RUN.\n''',encoding='utf-8')
print(json.dumps({'document':str(DOCX),'rows':report['row_counts'],'checks':len(checks),'all_pass':report['all_pass'],'failed':[k for k,v in checks.items() if not v]},ensure_ascii=False))
if not report['all_pass']:raise SystemExit(1)
