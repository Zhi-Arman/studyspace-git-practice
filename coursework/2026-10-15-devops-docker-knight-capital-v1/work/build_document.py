#!/usr/bin/env python3
"""Build the bilingual classroom handout from the checked content inventory."""

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


BASE = Path(__file__).resolve().parent.parent
DATA = json.loads((BASE / "work" / "content.json").read_text(encoding="utf-8"))
OUTPUT = BASE / "DevOps_课堂三项任务_中俄双语_v1.docx"


def set_font(run, *, size=9, bold=False, color="263244", font="Arial"):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:eastAsia"), "Heiti SC")
    if any("\u4e00" <= c <= "\u9fff" for c in run.text):
        lang = rpr.find(qn("w:lang"))
        if lang is None:
            lang = OxmlElement("w:lang")
            rpr.append(lang)
        lang.set(qn("w:val"), "zh-CN")
    return run


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=85, start=90, bottom=85, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = qn(f"w:{side}")
        el = mar.find(tag)
        if el is None:
            el = OxmlElement(f"w:{side}")
            mar.append(el)
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")


def set_no_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:tblHeader"))


def pair(container, zh, ru, *, size=9, before=0, after=4, bold=False):
    p1 = container.add_paragraph()
    p1.paragraph_format.space_before = Pt(before)
    p1.paragraph_format.space_after = Pt(1)
    p1.paragraph_format.line_spacing = 1.12
    set_font(p1.add_run(zh), size=size, bold=bold)
    p2 = container.add_paragraph()
    p2.paragraph_format.space_after = Pt(after)
    p2.paragraph_format.line_spacing = 1.12
    set_font(p2.add_run(ru), size=size - 0.3, color="4B5A6D")
    return p1, p2


def fill_pair_cell(cell, zh, ru, *, size=8.6, bold=False):
    cell.text = ""
    p1 = cell.paragraphs[0]
    p1.paragraph_format.space_after = Pt(1)
    p1.paragraph_format.line_spacing = 1.05
    set_font(p1.add_run(zh), size=size, bold=bold)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.05
    set_font(p2.add_run(ru), size=size - 0.3, color="536174")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def set_widths(table, widths):
    table.autofit = False
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Cm(width)


doc = Document()
section = doc.sections[0]
section.page_width = Cm(29.7)
section.page_height = Cm(21)
section.top_margin = Cm(1.35)
section.bottom_margin = Cm(1.25)
section.left_margin = Cm(1.25)
section.right_margin = Cm(1.25)
section.header_distance = Cm(0.6)
section.footer_distance = Cm(0.6)

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"].font.size = Pt(9)
styles["Normal"].paragraph_format.space_after = Pt(3)
normal_rpr = styles["Normal"]._element.get_or_add_rPr()
normal_rfonts = normal_rpr.rFonts
normal_rfonts.set(qn("w:eastAsia"), "Heiti SC")
for name, size in (("Title", 18), ("Heading 1", 13), ("Heading 2", 10.5)):
    style = styles[name]
    style.font.name = "Arial"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(26, 58, 93)
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Heiti SC")
    style.paragraph_format.space_before = Pt(10)
    style.paragraph_format.space_after = Pt(4)

p = doc.add_paragraph(style="Title")
p.paragraph_format.space_before = Pt(0)
set_font(p.add_run(DATA["title_zh"]), size=17.5, bold=True, color="17375E")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(4)
set_font(p.add_run(DATA["title_ru"]), size=12, bold=True, color="365878")
pair(doc, f"版本：{DATA['version']}　日期：{DATA['date']}　状态：{DATA['status_zh']}",
     f"Версия: {DATA['version']}　Дата: {DATA['date']}　Статус: {DATA['status_ru']}", size=8.5, after=3)
pair(doc, DATA["source_note_zh"], DATA["source_note_ru"], size=8.7, after=2)

doc.add_heading("1. 手工发布流程与自动化判断 / Ручной выпуск и автоматизация", level=1)
pair(doc,
     "题目：对十项人工操作逐项判断是否可自动化、应放在构建/测试/部署的哪个阶段，并简要解释。",
     "Задание: для каждой из десяти ручных операций определить возможность автоматизации, этап сборки/тестирования/развёртывания и кратко обосновать выбор.",
     size=8.8, after=5)

ops = DATA["operations"]
table = doc.add_table(rows=1, cols=4)
table.style = "Table Grid"
headers = [("序号", "№"), ("操作", "Операция"), ("自动化 / 阶段", "Автоматизация / этап"), ("理由", "Обоснование")]
for cell, (zh, ru) in zip(table.rows[0].cells, headers):
    fill_pair_cell(cell, zh, ru, size=8.7, bold=True)
    set_cell_shading(cell, "DDEAF5")
set_repeat_header(table.rows[0])
set_no_split(table.rows[0])
for idx, item in enumerate(ops):
    row = table.add_row()
    set_no_split(row)
    row.cells[0].text = str(item["n"])
    row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_cell_margins(row.cells[0])
    fill_pair_cell(row.cells[1], item["zh"], item["ru"], size=8.45)
    fill_pair_cell(row.cells[2], f"{item['auto_zh']} · {item['phase_zh']}",
                   f"{item['auto_ru']} · {item['phase_ru']}", size=8.3)
    fill_pair_cell(row.cells[3], item["why_zh"], item["why_ru"], size=8.25)
    if idx % 2:
        for cell in row.cells:
            set_cell_shading(cell, "F5F8FB")
set_widths(table, [1.0, 6.0, 4.2, 16.0])
pair(doc,
     "结论：代码获取、构建和重复性检查适合自动执行；生产放行、敏感配置变更及事故决策保留明确的人为责任。",
     "Вывод: получение кода, сборка и повторяемые проверки подходят для автоматизации; выпуск в production, изменение чувствительных настроек и решения при инцидентах сохраняют ответственность человека.",
     size=8.7, before=5, after=2)

doc.add_heading("2. Python Web 服务的 Dockerfile / Dockerfile для веб-сервиса Python", level=1)
pair(doc,
     "题目条件：Python 3.12；项目根目录含 requirements.txt 与 app.py；命令 python app.py；端口 8000；Dockerfile 应包含 EXPOSE。",
     "Условия: Python 3.12; в корне проекта находятся requirements.txt и app.py; запуск — python app.py; порт 8000; в Dockerfile требуется EXPOSE.",
     size=8.8, after=5)
code_table = doc.add_table(rows=1, cols=1)
code_table.style = "Table Grid"
code_cell = code_table.cell(0, 0)
code_cell.text = ""
set_cell_shading(code_cell, "EFF4F8")
set_cell_margins(code_cell, top=140, bottom=140, start=220, end=140)
for i, line in enumerate(DATA["dockerfile"]):
    p = code_cell.paragraphs[0] if i == 0 else code_cell.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    set_font(p.add_run(line), size=9.2, color="172E46", font="Menlo")
set_no_split(code_table.rows[0])
set_widths(code_table, [27.2])
for explanation in DATA["docker_explanations"]:
    pair(doc, explanation["zh"], explanation["ru"], size=8.6, after=2)
pair(doc,
     "验证范围：这里只给出 Dockerfile 文本。照片没有提供 app.py 和 requirements.txt 的实际文件，因此没有执行镜像构建或运行测试。",
     "Объём проверки: здесь приведён только текст Dockerfile. Файлы app.py и requirements.txt на фотографии не предоставлены, поэтому сборка образа и запуск не выполнялись.",
     size=8.3, before=2, after=2)

doc.add_heading("3. Knight Capital 事故与 DevOps 实践 / Кейс Knight Capital и практики DevOps", level=1)
pair(doc,
     "题目：至少提出三项可能避免 Knight Capital 故障的 DevOps 实践，并解释其作用。",
     "Задание: назвать не менее трёх практик DevOps, которые могли бы предотвратить сбой Knight Capital, и объяснить действие каждой.",
     size=8.8, after=4)
pair(doc,
     "课堂材料：Knight Capital 在八台服务器上手工部署新版交易程序，其中一台漏装新版，残留的旧代码意外运行。照片称，45 分钟内出现大量无用交易，损失约 4.4 亿美元，公司不久后不再独立存在。材料提示关注人工操作、自动检查缺失、版本不同步和步骤执行控制。",
     "Материал занятия: новую торговую программу Knight Capital вручную развёртывали на восьми серверах. На одном сервере новую версию забыли установить, и старый код неожиданно сработал. По слайду, за 45 минут произошло множество ненужных сделок, потери составили около 440 млн долларов, а вскоре компания перестала быть самостоятельной. Подсказка слайда: ручные операции, отсутствие автоматической проверки, рассинхронизация версий и контроль шагов.",
     size=8.6, after=5)
pair(doc,
     "分析边界：下表针对课堂材料描述的故障链提出可能的预防和止损措施；任何单项措施都不能保证事故绝不会发生。",
     "Граница анализа: ниже предложены меры предотвращения и ограничения ущерба для описанной на занятии цепочки сбоя; ни одна отдельная мера не гарантирует полной защиты.",
     size=8.4, after=4)

practices = DATA["practices"]
table = doc.add_table(rows=1, cols=4)
table.style = "Table Grid"
headers = [("序号", "№"), ("实践", "Практика"), ("作用", "Роль"), ("如何帮助", "Как помогает")]
for cell, (zh, ru) in zip(table.rows[0].cells, headers):
    fill_pair_cell(cell, zh, ru, size=8.7, bold=True)
    set_cell_shading(cell, "DDEAF5")
set_repeat_header(table.rows[0])
set_no_split(table.rows[0])
for idx, item in enumerate(practices):
    row = table.add_row()
    set_no_split(row)
    row.cells[0].text = str(item["n"])
    row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_cell_margins(row.cells[0])
    fill_pair_cell(row.cells[1], item["name_zh"], item["name_ru"], size=8.4)
    fill_pair_cell(row.cells[2], item["kind_zh"], item["kind_ru"], size=8.3)
    fill_pair_cell(row.cells[3], item["how_zh"], item["how_ru"], size=8.2)
    if idx % 2:
        for cell in row.cells:
            set_cell_shading(cell, "F5F8FB")
set_widths(table, [1.0, 5.6, 3.0, 17.6])

doc.add_heading("资料与核验状态 / Источники и статус проверки", level=2)
pair(doc,
     "参考线索：美国证券交易委员会（SEC），In the Matter of Knight Capital Americas LLC，Release No. 34-70694（2013）。",
     "Ориентир для проверки: Комиссия по ценным бумагам и биржам США (SEC), In the Matter of Knight Capital Americas LLC, Release No. 34-70694 (2013).",
     size=8.3, after=2)
pair(doc, DATA["research_note_zh"], DATA["research_note_ru"], size=8.3, after=0)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_font(footer.add_run("DevOps 课堂任务 / Практическая работа по DevOps  ·  v1.0"), size=7.5, color="6A7785")

doc.save(OUTPUT)
print(OUTPUT)
