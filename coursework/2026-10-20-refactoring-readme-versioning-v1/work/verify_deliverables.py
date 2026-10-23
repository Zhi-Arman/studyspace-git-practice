"""Verify coverage and generated document text without replacing visual QA."""

import ast
from collections import Counter
import hashlib
import json
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
STEM = "代码重构与软件维护实践_中俄双语_v1"
doc = Document(ROOT / f"{STEM}.docx")
md = (ROOT / f"{STEM}.md").read_text()
inventory = json.loads((ROOT / "work/content.json").read_text())
doc_text = [p.text for p in doc.paragraphs]
for table in doc.tables:
    doc_text.extend(cell.text for row in table.rows for cell in row.cells)
actual = Counter(doc_text)
expected = []
in_code = False
for line in md.splitlines():
    if line.startswith("```"):
        in_code = not in_code
        continue
    if not line.strip() or line == "<!-- pagebreak -->":
        continue
    if in_code:
        expected.append(line)
    elif line.startswith("|"):
        cells = [x.strip() for x in line.strip("|").split("|")]
        if not all(set(x) <= set(" -:") for x in cells):
            expected.extend(x.replace("`", "") for x in cells)
    else:
        expected.append(line.lstrip("# ").replace("`", ""))
missing = list((Counter(expected) - actual).elements())
employee_tree = ast.parse((ROOT / "code/employees_refactored.py").read_text())
test_tree = ast.parse((ROOT / "code/test_employees.py").read_text())
tests = [n for n in ast.walk(test_tree) if isinstance(n, ast.FunctionDef) and n.name.startswith("test_")]
functions = [n for n in ast.walk(employee_tree) if isinstance(n, ast.FunctionDef)]
readme = (ROOT / "tax-calculator/README.md").read_text()
revision = json.loads((ROOT / "artifacts/employee-revision-results.json").read_text())
backup = json.loads((ROOT / "work/revision-backup.json").read_text())
before_md = (Path(backup["backup_dir"]) / f"{STEM}.md").read_text()
before_prefix, before_tail = before_md.split("## 3 项目 README / README проекта", 1)
before_suffix = before_tail.split("## 4 版本号判断 / Выбор номера версии", 1)[1]
after_prefix, after_tail = md.split("## 3 重构后的测试与结果 / Тесты и результаты после рефакторинга", 1)
after_suffix = after_tail.split("## 4 版本号判断 / Выбор номера версии", 1)[1]
checks = {
    "four_source_tasks": len(inventory["tasks"]) == 4,
    "both_images_in_inventory": len(inventory["source_images"]) == 2,
    "docx_matches_every_markdown_item": not missing,
    "all_three_functions_have_docstrings": len(functions) == 3 and all(ast.get_docstring(n) for n in functions),
    "exactly_two_employee_tests": len(tests) == 2,
    "readme_five_required_sections": sum(line.startswith("## ") for line in readme.splitlines()) == 5 and all(s in readme for s in inventory["readme_sections"]),
    "version_answers_present": all(v in md for v in inventory["tasks"][3]["answers"]),
    "version_table_is_editable": len(doc.tables[0].rows) == 4 and len(doc.tables[0].columns) == 3,
    "evidence_table_is_editable": len(doc.tables[1].rows) == 5 and len(doc.tables[1].columns) == 3,
    "before_tests_passed": "Ran 2 tests" in (ROOT / "artifacts/employee-before.log").read_text() and (ROOT / "artifacts/employee-before.log").read_text().rstrip().endswith("OK"),
    "after_tests_passed": "Ran 2 tests" in (ROOT / "artifacts/employee-after.log").read_text() and (ROOT / "artifacts/employee-after.log").read_text().rstrip().endswith("OK"),
    "tax_tests_passed": "Ran 4 tests" in (ROOT / "artifacts/tax-tests.log").read_text() and (ROOT / "artifacts/tax-tests.log").read_text().rstrip().endswith("OK"),
    "example_output_verified": (ROOT / "artifacts/tax-example.log").read_text().strip() == "Tax: 1300.00",
    "readme_cli_behavior_verified": len(json.loads((ROOT / "artifacts/tax-cli-checks.json").read_text())) == 4,
    "no_decorative_paragraph_borders": not list(doc.styles.element.iter(qn("w:pBdr"))) and not list(doc.element.iter(qn("w:pBdr"))),
    "no_pdf_in_delivery_folder": not list(ROOT.glob("*.pdf")),
    "requested_section_replaced": "## 3 项目 README" not in md and "### 3.3 修改前后说明" in md,
    "content_outside_section_three_preserved": before_prefix == after_prefix and before_suffix == after_suffix,
    "application_code_unchanged": all(hashlib.sha256((ROOT / name).read_bytes()).hexdigest() == digest for name, digest in backup["code_sha256"].items()),
    "fresh_before_and_after_tests_passed": all(run["exit_code"] == 0 and run["passed"] == run["total"] == 2 for run in revision["test_runs"]),
    "reported_actual_values_verified": all(case["expected"] == case["original_actual"] == case["refactored_actual"] and case["status"] == "PASS" for case in revision["cases"]),
    "actual_after_test_log_embedded": (ROOT / "artifacts/employee-revision-after.log").read_text().strip() in md,
    "six_paired_change_descriptions": all(f"{i}. " in md.split("### 3.3 修改前后说明", 1)[1].split("### 3.4", 1)[0] for i in range(1, 7)),
}
render = Path("/private/tmp/refactoring-coursework-render-project-edit-final")
pdf = next(render.glob("*.pdf"))
pages = PdfReader(pdf).pages
extracted = "\n".join(p.extract_text() for p in pages)
checks["rendered_chinese_present"] = "代码重构" in extracted and "测试" in extracted
checks["one_preview_image_per_page"] = len(list(render.glob("page-*.png"))) == len(pages)
result = {
    "date": "2026-10-20",
    "scope": "independent classroom examples",
    "checks": checks,
    "all_structural_checks_pass": all(checks.values()),
    "missing_text": missing,
    "pages_rendered": len(pages),
    "visual_review": "PENDING",
    "docx_sha256": hashlib.sha256((ROOT / f"{STEM}.docx").read_bytes()).hexdigest(),
    "sources_compared": inventory["source_images"],
    "source_coverage": [{"id": t["id"], "outputs": t["outputs"]} for t in inventory["tasks"]],
}
(ROOT / "artifacts/verification.json").write_text(json.dumps(result, ensure_ascii=False, indent=2))
print(json.dumps({"checks": len(checks), "all_structural_checks_pass": result["all_structural_checks_pass"], "pages_rendered": len(pages), "missing_text": missing}, ensure_ascii=False))
assert all(checks.values()), [k for k, v in checks.items() if not v]
