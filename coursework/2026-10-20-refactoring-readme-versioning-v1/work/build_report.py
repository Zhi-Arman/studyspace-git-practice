"""Build paired coursework Markdown and DOCX from one content inventory."""

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
STEM = "代码重构与软件维护实践_中俄双语_v1"
inventory = json.loads((ROOT / "work/content.json").read_text())
refactored = (ROOT / "code/employees_refactored.py").read_text().rstrip()
tests = (ROOT / "code/test_employees.py").read_text().rstrip()
revision_section = (ROOT / "work/refactoring_results_section.md").read_text().strip()
after_log = (ROOT / "artifacts/employee-revision-after.log").read_text().strip()
revision_section = revision_section.replace("{{AFTER_LOG}}", after_log)

report = f'''# 代码重构与软件维护实践

Практика рефакторинга и сопровождения ПО

日期 / Дата: 2026-10-20    版本 / Версия: 1.0    状态 / Статус: DRAFT

本作业完成四项任务，含代码、测试、README 和版本判断；两个回归测试在重构前后均通过。

Выполнены четыре задания: код, тесты, README и версии. Два регрессионных теста проходят до и после изменений.

## 题目要求 / Условия задания

1. 为函数和变量起清晰名称，按 DRY 原则消除重复，并补充 docstring。

Переименовать функции и переменные, устранить дублирование по DRY, добавить docstring.

2. 解释重构前测试的作用，并编写两个简单测试。

Объяснить, зачем нужны тесты перед рефакторингом, и написать два простых теста.

3. 为 Python 控制台税额计算器编写 README，包含说明、安装、运行、示例和测试。

Написать README калькулятора налога: описание, установка, запуск, пример, тесты.

4. 从版本 1.4.2 分别判断修复舍入错误、新增兼容报表、旧配置失效三种改动的版本号。

От 1.4.2 выбрать версии для исправления округления, нового отчёта и несовместимых настроек.

## 1 数据约定与重构思路 / Данные и подход

每条记录为五项元组：名、姓、部门、未使用休假天数、病假天数。f 使用下标 3，g 使用下标 4；均要求部门相同且对应天数大于 0，返回以一个空格连接的姓名列表。

Каждая запись — кортеж из пяти полей: имя, фамилия, отдел, неиспользованные отпускные дни, больничные дни. f проверяет поле 3, g — поле 4. Отдел должен совпадать, число дней должно быть больше нуля; результат — список имён и фамилий с одним пробелом.

原代码问题是命名不清、筛选与姓名拼接重复，以及数字下标含义不明。重构提取公共筛选函数、命名字段常量，使用清晰业务名称，并通过别名保留 f 和 g。返回顺序、重复姓名、空列表和严格大于零的条件保持不变；不修改输入，也不额外增加数据校验规则。

Устранены неясные имена, повторение логики и числовые индексы без пояснений. Общая функция и константы обеспечивают DRY; псевдонимы сохраняют вызовы f и g. Порядок, повторы, пустой результат и условие больше нуля сохранены. Вход не изменяется; новые правила проверки не добавлены.

<!-- pagebreak -->

## 1 重构后的代码 / Код после рефакторинга

文件 / Файл: code/employees_refactored.py。公共函数说明了输入、筛选规则和返回值；两个业务函数复用同一个筛选实现。

Docstring описывают вход, условие отбора и результат. Обе функции используют одну общую реализацию отбора.

```python
{refactored}
```

<!-- pagebreak -->

## 2 重构前的测试 / Тесты перед рефакторингом

测试先固定原有可观察行为，作为重构的参照。修改后运行同一组测试，可以发现字段选错、条件变化、顺序或姓名格式变化等回归。两个测试覆盖休假与病假两个分支，并包含零值、负值、其他部门和重复姓名；有限测试不能证明所有输入都正确。

Тесты сначала фиксируют наблюдаемое поведение исходной программы. Повторный запуск выявляет регрессии: неверное поле, изменение условия, порядка или формата имени. Два теста проверяют отпуск и больничные, включая нули, отрицательные значения, другой отдел и повтор имени. Конечный набор тестов не доказывает корректность для всех входов.

文件 / Файл: code/test_employees.py。EMPLOYEE_MODULE 仅用于选择原代码或重构代码，不改变断言。

EMPLOYEE_MODULE выбирает исходный или изменённый модуль; проверки остаются одинаковыми.

```python
{tests}
```

<!-- pagebreak -->

{revision_section}

<!-- pagebreak -->

## 4 版本号判断 / Выбор номера версии

采用 MAJOR.MINOR.PATCH。以下三项分别以 1.4.2 为起点；将配置文件格式视为用户依赖的公开兼容性约定。

Используется MAJOR.MINOR.PATCH. Каждое изменение отдельно отсчитывается от 1.4.2; формат настроек рассматривается как публичный контракт совместимости.

| 改动 / Изменение | 新版本 / Версия | 理由 / Обоснование |
| --- | --- | --- |
| 修复舍入错误 / Исправлена ошибка округления | 1.4.3 | 兼容的错误修复，递增 PATCH。 / Совместимое исправление, увеличивается PATCH. |
| 增加新报表，原报表仍可用 / Добавлен новый отчёт, прежние работают | 1.5.0 | 新增兼容功能，递增 MINOR，PATCH 清零。 / Новая совместимая функция: MINOR + 1, PATCH = 0. |
| 修改配置格式，旧文件不能使用 / Новый формат настроек, старые файлы не работают | 2.0.0 | 破坏兼容性，递增 MAJOR，其余清零。 / Нарушение совместимости: MAJOR + 1, остальные части = 0. |

依据 / Источник: Semantic Versioning 2.0.0，第 6 至 8 条 / пункты 6–8。
https://semver.org/lang/zh-CN/

若三项按题目顺序连续发布，依次也是 1.4.3 → 1.5.0 → 2.0.0；若三项合并为一次发布，则直接为 2.0.0。

При трёх последовательных выпусках в указанном порядке: 1.4.3 → 1.5.0 → 2.0.0. Если все изменения входят в один выпуск, его версия — 2.0.0.

<!-- pagebreak -->

## 5 实际运行记录 / Фактическая проверка

执行日期为 2026-10-20，环境为 macOS 和 Python 3.12.14。以下结果来自所附独立课堂示例。

Проверка выполнена 2026-10-20 в macOS, Python 3.12.14. Результаты относятся к приложенным учебным примерам.

| 验证 / Проверка | 结果 / Результат | 日志 / Журнал |
| --- | --- | --- |
| 原代码，两个员工筛选测试 / Исходный код, два теста | 2 / 2 PASS | artifacts/employee-before.log |
| 重构代码，同样两个测试 / Тот же набор после рефакторинга | 2 / 2 PASS | artifacts/employee-after.log |
| 税额计算器 / Калькулятор налога | 4 / 4 PASS | artifacts/tax-tests.log |
| 命令行示例 / Пример CLI | Tax: 1300.00 | artifacts/tax-example.log |

复现命令 / Команды воспроизведения

在 code 文件夹中 / В папке code:

```bash
EMPLOYEE_MODULE=employees_original python3 -m unittest -v test_employees.py
python3 -m unittest -v test_employees.py
```

在 tax-calculator 文件夹中 / В папке tax-calculator:

```bash
python3 -m unittest -v test_tax_calculator.py
python3 tax_calculator.py --amount 10000 --rate 13
```
'''

(ROOT / f"{STEM}.md").write_text(report, encoding="utf-8")

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(0.65)
section.left_margin = section.right_margin = Inches(0.75)

for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"]:
    style = doc.styles[name]
    style.font.name = "Liberation Sans"
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Songti SC")
normal = doc.styles["Normal"]
normal.font.size = Pt(10.5)
normal.paragraph_format.line_spacing = 1.12
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.widow_control = True
doc.styles["Title"].font.size = Pt(22)
doc.styles["Title"].paragraph_format.space_after = Pt(8)
for name, size in [("Heading 1", 15), ("Heading 2", 12.5), ("Heading 3", 11)]:
    doc.styles[name].font.size = Pt(size)
    doc.styles[name].paragraph_format.space_before = Pt(10)
    doc.styles[name].paragraph_format.space_after = Pt(5)
    doc.styles[name].paragraph_format.keep_with_next = True
code_style = doc.styles.add_style("Code", 1)
code_style.font.name = "Liberation Mono"
code_style.font.size = Pt(9)
code_style.font.color.rgb = RGBColor(0, 0, 0)
code_style.paragraph_format.line_spacing = Pt(10.3)
code_style.paragraph_format.space_after = Pt(0)
code_style.paragraph_format.keep_together = True
code_style.paragraph_format.widow_control = False


def rich_text(paragraph, text):
    for index, part in enumerate(text.split("`")):
        run = paragraph.add_run(part)
        if index % 2:
            run.font.name = "Liberation Mono"
            run.font.size = Pt(9)


def table(rows):
    result = doc.add_table(rows=1, cols=len(rows[0]))
    result.autofit = False
    widths = [2.65, 0.95, 3.4] if len(rows[0]) == 3 else [7 / len(rows[0])] * len(rows[0])
    for col, width in zip(result.columns, widths):
        col.width = Inches(width)
    for ri, row in enumerate(rows):
        cells = result.rows[0].cells if ri == 0 else result.add_row().cells
        for ci, (cell, text) in enumerate(zip(cells, row)):
            cell.width = Inches(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.space_before = Pt(4)
            rich_text(paragraph, text)
            for run in paragraph.runs:
                run.font.size = Pt(9.5)
                if ri == 0:
                    run.bold = True
            props = cell._tc.get_or_add_tcPr()
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "E7EDF4" if ri == 0 else "FFFFFF")
            props.append(shading)
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right"):
                tag = OxmlElement(f"w:{edge}")
                for key, val in [("val", "single"), ("sz", "4"), ("color", "D9D9D9")]:
                    tag.set(qn(f"w:{key}"), val)
                borders.append(tag)
            props.append(borders)
            margins = OxmlElement("w:tcMar")
            for edge in ("top", "left", "bottom", "right"):
                tag = OxmlElement(f"w:{edge}")
                tag.set(qn("w:w"), "90")
                tag.set(qn("w:type"), "dxa")
                margins.append(tag)
            props.append(margins)
        tr_pr = result.rows[ri]._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
        if ri == 0:
            tr_pr.append(OxmlElement("w:tblHeader"))
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


lines = report.splitlines()
i = 0
in_code = False
while i < len(lines):
    line = lines[i]
    if line.startswith("```"):
        in_code = not in_code
        if not in_code:
            doc.add_paragraph().paragraph_format.space_after = Pt(1)
    elif in_code:
        doc.add_paragraph(line or " ", style="Code")
    elif line == "<!-- pagebreak -->":
        doc.add_page_break()
    elif line.startswith("|"):
        rows = []
        while i < len(lines) and lines[i].startswith("|"):
            cells = [value.strip() for value in lines[i].strip("|").split("|")]
            if not all(set(value) <= set(" -:") for value in cells):
                rows.append(cells)
            i += 1
        table(rows)
        continue
    elif line.startswith("# "):
        doc.add_paragraph(line[2:], style="Title")
    elif line.startswith("## "):
        doc.add_paragraph(line[3:], style="Heading 1")
    elif line.startswith("### "):
        doc.add_paragraph(line[4:], style="Heading 2")
    elif line.strip():
        paragraph = doc.add_paragraph()
        rich_text(paragraph, line)
    i += 1

doc.core_properties.title = "代码重构与软件维护实践"
doc.core_properties.subject = "Практика рефакторинга и сопровождения ПО"
doc.core_properties.author = ""
# Remove inherited theme choices and decorative paragraph borders.
for fonts in doc.styles.element.iter(qn("w:rFonts")):
    for attr in list(fonts.attrib):
        if "Theme" in attr:
            del fonts.attrib[attr]
for border in list(doc.styles.element.iter(qn("w:pBdr"))):
    border.getparent().remove(border)
for border in list(doc.element.iter(qn("w:pBdr"))):
    border.getparent().remove(border)
doc.save(ROOT / f"{STEM}.docx")
print(ROOT / f"{STEM}.docx")
print("tables:", len(doc.tables), "paragraphs:", len(doc.paragraphs))
