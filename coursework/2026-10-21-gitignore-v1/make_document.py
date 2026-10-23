from pathlib import Path
import json
from zipfile import ZipFile
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE=Path(__file__).resolve().parent
ROOT=HERE.parent.parent
report=json.loads((HERE/'artifacts/verification.json').read_text())
assert report['all_pass']
rules=[
('.venv/','Python 虚拟环境','Виртуальная среда Python'),
('site/','生成的文档网站','Собранный сайт документации'),
('docs/_build/','根目录文档构建输出','Сборка документации в корне'),
('**/bin/','任意层级的编译输出目录','Каталоги сборки bin на любой глубине'),
('**/obj/','任意层级的编译中间文件','Промежуточные файлы obj'),
('TestResults/','测试运行结果目录','Результаты выполнения тестов'),
('*.log','日志文件','Файлы журналов'),
('.env','本地环境变量与配置','Локальные переменные и настройки'),
('.DS_Store','macOS 目录元数据','Метаданные каталогов macOS'),
('Thumbs.db','Windows 缩略图缓存','Кэш миниатюр Windows'),
]
pairs1=[
('已在 StudySpace 根目录补齐图片中的十项忽略规则，保留原有依赖、缓存、数据库等规则。实际生效文件为根目录 .gitignore。',
 'В корневом .gitignore StudySpace обеспечены все десять правил с фото. Прежние правила для зависимостей, кэшей и БД сохранены. Рабочий файл находится в корне проекта.'),
('文件位置 / Путь: /Users/zhi/Documents/软件生命周期/studyspace/.gitignore',
 '规则用于版本管理；项目文件仍保留在本机。 Правила управляют включением в Git; файлы остаются на компьютере.'),
]
pairs2=[
('除图片要求外，新增 **/docs/_build/，使嵌套文档项目的构建产物也被忽略。保留 !.env.example，便于提交不含秘密的配置示例；保留全部原有规则。',
 'Дополнительно добавлено **/docs/_build/ для вложенных проектов документации. Сохранено !.env.example для безопасного образца настроек; все прежние правила оставлены.'),
('使用 git check-ignore -v --no-index -- <路径> 实际检查 37 个路径，37/37 通过。其中 23 个应忽略、14 个应保留。路径探针为匹配示例，没有创建这些测试文件。',
 'Командой git check-ignore -v --no-index -- <путь> проверены 37 путей: 37/37 успешно. Из них 23 исключаются, 14 доступны для добавления. Это примеры путей для проверки правил; соответствующие файлы не создавались.'),
('另外核对已有课程网站：site/index.html 文件仍存在，现由 site/ 规则忽略。Markdown 源文档、OpenAPI、应用源码、依赖锁文件与 .env.example 仍可纳入 Git。',
 'Дополнительно проверен существующий учебный сайт: site/index.html остался на диске и исключается правилом site/. Markdown, OpenAPI, код приложения, lock-файлы и .env.example по-прежнему можно включить в Git.'),
('本次仅修改忽略规则及配套说明，没有删除文件或修改 Git 索引。检查时根仓库已跟踪文件数为 0，无需移出已跟踪文件。一般而言，.gitignore 不会自动取消对既有文件的跟踪。',
 'Изменены правила и пояснения; файлы не удалялись, индекс Git не менялся. При проверке в корневом репозитории было 0 отслеживаемых файлов, удаление из индекса не требовалось. Обычно .gitignore не прекращает отслеживание уже включённых файлов.'),
('原始结果见 artifacts/verification.json；它验证忽略行为，不代表应用业务测试或验收。前一份 README 提交件保留为当时版本，项目根 README 已同步新规则。',
 'Исходные результаты находятся в artifacts/verification.json. Проверка касается игнорирования Git, а не бизнес-тестов или приёмки. Предыдущая учебная версия README сохранена как историческая; корневой README обновлён.'),
]
title='StudySpace Git 忽略规则'
subtitle='Правила игнорирования Git в StudySpace'
meta='SS-GITIGNORE-01 · v1.0 · DRAFT · 2026-10-21 · Asia/Shanghai'
md=f'# {title}\n\n{subtitle}\n\n{meta}\n\n'
md+='## 1 图片要求与实际规则 / Требования фото и рабочие правила\n\n'
for a,b in pairs1:md+=a+'\n\n'+b+'\n\n'
md+='| 规则 Правило | 用途 Назначение |\n| --- | --- |\n'
for rule,zh,ru in rules:md+=f'| `{rule}` | {zh}<br>{ru} |\n'
md+='\n## 2 验证与说明 / Проверка и пояснения\n\n'
for a,b in pairs2:md+=a+'\n\n'+b+'\n\n'
md+='[实际 .gitignore / Рабочий файл](../../.gitignore) · [检查结果 / Результаты](artifacts/verification.json)\n'
(HERE/'README.md').write_text(md)
doc=Document();sec=doc.sections[0]
sec.page_width,sec.page_height=Cm(21),Cm(29.7)
sec.left_margin=sec.right_margin=Cm(1.8)
sec.top_margin=sec.bottom_margin=Cm(1.6)
for name,size in [('Normal',10),('Title',19),('Subtitle',13),('Heading 1',13),('Heading 2',11)]:
 st=doc.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
 st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Hiragino Sans GB')
 st.paragraph_format.space_after=Pt(5);st.paragraph_format.line_spacing=1.0
doc.add_paragraph(title,'Title');doc.add_paragraph(subtitle,'Subtitle');doc.add_paragraph(meta)
doc.add_paragraph('1 图片要求与实际规则','Heading 1');doc.add_paragraph('1 Требования фото и рабочие правила','Heading 2')
def pair(a,b):
 p=doc.add_paragraph(a);p.paragraph_format.keep_with_next=True
 doc.add_paragraph(b)
for a,b in pairs1:pair(a,b)
t=doc.add_table(rows=1,cols=2);t.autofit=False;t.columns[0].width=Cm(3.5);t.columns[1].width=Cm(13.9)
t.rows[0].cells[0].text='规则 Правило';t.rows[0].cells[1].text='用途 Назначение'
t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
for rule,zh,ru in rules:
 cells=t.add_row().cells;cells[0].text=rule;cells[1].text=zh+'\n'+ru
for ri,row in enumerate(t.rows):
 row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
 for ci,c in enumerate(row.cells):
  c.width=Cm(3.5 if ci==0 else 13.9)
  for p in c.paragraphs:
   p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3)
   for r in p.runs:r.font.size=Pt(9);r.bold=ri==0
  pr=c._tc.get_or_add_tcPr();bd=OxmlElement('w:tcBorders')
  for edge in ['top','left','bottom','right']:
   el=OxmlElement('w:'+edge)
   for k,v in [('val','single'),('sz','4'),('color','CCCCCC')]:el.set(qn('w:'+k),v)
   bd.append(el)
  pr.append(bd)
doc.add_page_break()
doc.add_paragraph('2 验证与说明','Heading 1');doc.add_paragraph('2 Проверка и пояснения','Heading 2')
for a,b in pairs2:pair(a,b)
for tree in [doc.styles.element,doc.element]:
 for el in list(tree.iter(qn('w:pBdr'))):el.getparent().remove(el)
dest=HERE/'StudySpace_gitignore_中俄双语_v1.docx';doc.save(dest)
with ZipFile(dest) as z:assert z.testzip() is None
text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for row in t.rows for c in row.cells)
assert all(a in text and b in text for a,b in pairs1+pairs2)
assert all(r in text for r,_,_ in rules)
print(dest)
