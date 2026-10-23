from pathlib import Path
import hashlib
import json
import re
from datetime import datetime, timezone
from zipfile import ZipFile

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

ROOT = Path('/Users/zhi/Documents/软件生命周期/studyspace')
OUT = ROOT / 'coursework/2026-10-15-deployment-nodes-v1'
DATA = OUT / 'data'
DATA.mkdir(parents=True, exist_ok=True)
STEM = 'StudySpace_节点与制品表_中俄双语_v1'
DOCX = OUT / f'{STEM}.docx'

sources = [
    ('S1', ROOT / 'docs/zh-CN/04_System_Design.md', '系统设计 0.3.0', 'Проектирование системы 0.3.0'),
    ('S2', ROOT / 'docs/ru-RU/04_System_Design.md', '系统设计俄文对应版', 'Русская версия проектирования'),
    ('S3', ROOT / 'compose.yaml', '本地容器端口和数据卷配置', 'Локальные порты контейнеров и том данных'),
    ('S4', ROOT / 'frontend/Dockerfile', '前端镜像与静态文件路径', 'Образ клиента и путь статических файлов'),
    ('S5', ROOT / 'frontend/nginx.conf', 'Nginx 静态服务及 API 转发', 'Статика Nginx и проксирование API'),
    ('S6', ROOT / 'backend/Dockerfile', '后端镜像与 Uvicorn 启动', 'Образ сервера и запуск Uvicorn'),
    ('S7', ROOT / 'scripts/dev.sh', '本地启动地址和端口', 'Адреса и порты локального запуска'),
    ('S8', ROOT / 'frontend/vite.config.ts', 'Vite 与开发代理配置', 'Настройки Vite и прокси разработки'),
    ('S9', ROOT / 'backend/app/config.py', 'SQLite 默认文件与运行环境', 'Файл SQLite и среда по умолчанию'),
    ('S10', ROOT / 'backend/app/main.py', '当前基础 API', 'Текущие базовые API'),
    ('S11', ROOT / 'backend/app/database.py', 'SQLAlchemy 基础设施与数据库访问', 'Инфраструктура SQLAlchemy и доступ к БД'),
    ('S12', ROOT / '.env.example', 'PostgreSQL 连接示例', 'Пример подключения PostgreSQL'),
    ('S13', ROOT / 'docs/data/project-decisions.json', '部署决定与待确认事项', 'Решения о размещении и открытые вопросы'),
    ('S14', ROOT / 'docs/decisions/zh-CN/0004-docker-https.md', 'Docker 与 HTTPS 部署决定', 'Решение о Docker и HTTPS'),
    ('S15', ROOT / 'docs/decisions/ru-RU/0004-docker-https.md', '部署决定俄文对应版', 'Русская версия решения о размещении'),
    ('S16', ROOT / 'coursework/2026-10-14-architecture-decisions-v1/StudySpace_两份ADR架构决策记录_中俄双语_v1.md', 'ADR-0001 与 ADR-0002', 'ADR-0001 и ADR-0002'),
    ('S17', ROOT.parent / 'StudySpace_组件交付/StudySpace_组件候选与组件图_中俄双语_v2.md', '统一组件表 SS-ARCH-03 v2', 'Единые компоненты SS-ARCH-03 v2'),
    ('S18', ROOT / 'backend/pyproject.toml', '后端依赖及 PostgreSQL 可选驱动', 'Зависимости сервера и дополнительный драйвер PostgreSQL'),
    ('S19', ROOT / 'AGENTS.md', '项目文档与实现状态规则', 'Правила документов и статусов проекта'),
]
fingerprints = {str(p): hashlib.sha256(p.read_bytes()).hexdigest() for _, p, _, _ in sources}

# Five fields match the classroom table. Each cell is a Chinese/Russian pair.
headers = [
    ('节点', 'Узел'),
    ('类型', 'Тип'),
    ('部署的组件与制品', 'Размещаемые компоненты / артефакты'),
    ('端口与协议', 'Порты и протоколы'),
    ('环境', 'Среда'),
]
rows = [
    [
        ('N-01\n学生或管理员终端', 'Устройство студента или администратора'),
        ('设备\n电脑或手机', 'Устройство\nПК или телефон'),
        ('浏览器执行 CMP-001 界面代码；从前端服务下载 HTML、JavaScript 和 CSS。', 'Браузер исполняет интерфейс CMP-001 и получает HTML, JavaScript и CSS от сервиса клиента.'),
        ('无固定入站服务端口。开发访问 HTTP 5173；本地容器访问 HTTP 8080；目标 HTTPS 443。', 'Постоянный входящий порт не нужен. HTTP 5173 при разработке, HTTP 8080 для локальных контейнеров, целевой HTTPS 443.'),
        ('开发及目标使用端\n目标 HTTPS 待部署', 'Клиент разработки и целевой среды\nHTTPS ещё не развёрнут'),
    ],
    [
        ('N-02\n本地开发工作站', 'Локальная рабочая станция'),
        ('设备\n本地服务宿主', 'Устройство\nХост локальных сервисов'),
        ('Vite 提供 CMP-001；Uvicorn 运行\nFastAPI 与 SQLAlchemy 基础层。SQLite 文件 backend/data/studyspace.db 随后端访问；业务模块待实现。', 'Vite обслуживает CMP-001; Uvicorn запускает FastAPI и основу SQLAlchemy. Сервер обращается к файлу backend/data/studyspace.db SQLite; бизнес-модули ещё не реализованы.'),
        ('127.0.0.1:5173 HTTP；/api 代理至 127.0.0.1:8000，HTTP/JSON。SQLite 使用文件访问，无网络端口。', '127.0.0.1:5173 HTTP; /api передаётся на 127.0.0.1:8000 по HTTP/JSON. SQLite использует файл без сетевого порта.'),
        ('development\n现有启动配置', 'development\nИмеющаяся конфигурация запуска'),
    ],
    [
        ('N-03\n前端静态服务', 'Сервис статического клиента'),
        ('容器\nNginx', 'Контейнер\nNginx'),
        ('CMP-001 的分发制品：frontend/dist 中的 index.html 与 assets；镜像内位于 /usr/share/nginx/html。nginx.conf 配置静态服务和 /api/ 转发。', 'Артефакты CMP-001: index.html и assets из frontend/dist, в образе — /usr/share/nginx/html. nginx.conf задаёт статику и прокси /api/.'),
        ('127.0.0.1:8080 → 80，HTTP。\n/api/ → backend:8000，\nHTTP/JSON。\n目标由 Caddy 提供 HTTPS。', 'Локально 127.0.0.1:8080 → порт 80 контейнера, HTTP. /api/ → backend:8000, HTTP/JSON. В целевой схеме HTTPS предоставляет Caddy.'),
        ('本地容器候选\n配置已存在\n目标部署待实施', 'Локальный вариант контейнеров\nКонфигурация есть\nЦелевое размещение впереди'),
    ],
    [
        ('N-04\n后端应用服务', 'Сервис приложения'),
        ('容器\nPython\nUvicorn', 'Контейнер\nPython\nUvicorn'),
        ('backend/app 与依赖环境；FastAPI、Pydantic、同步 SQLAlchemy。CMP-002 至 CMP-008 为设计中的认证、审核、资源、查询、预约、统计及审计模块；CMP-009 的数据访问基础层已存在。', 'backend/app и зависимости; FastAPI, Pydantic, синхронный SQLAlchemy. CMP-002–CMP-008 проектируют аутентификацию, проверку, ресурсы, поиск, бронь, показатели и аудит; основа доступа CMP-009 уже есть.'),
        ('容器 HTTP 8000；本地绑定 127.0.0.1:8000 → 8000。基础 API 为 /api/v1/health、ready、docs。内部模块为进程内调用。', 'HTTP 8000 контейнера; локально 127.0.0.1:8000 → 8000. Базовые API: /api/v1/health, ready, docs. Вызовы внутренних модулей идут в одном процессе.'),
        ('本地容器候选\nAPP_ENV=\ndevelopment\n业务待实现', 'Локальный вариант контейнеров\nAPP_ENV=\ndevelopment\nБизнес не реализован'),
    ],
    [
        ('N-05\nHTTPS 入口', 'Вход HTTPS'),
        ('目标容器\nCaddy', 'Целевой контейнер\nCaddy'),
        ('目标制品为 Caddy 反向代理配置、证书及其持久存储；对接前端静态服务与 /api 路由。ADR-0004 已选定该技术方向。', 'Целевые артефакты: конфигурация обратного прокси Caddy, сертификаты и постоянное хранилище; доступ к статике клиента и маршрутам /api. Направление выбрано ADR-0004.'),
        ('目标 TCP 80 HTTP；\nTCP 443 HTTPS/TLS。内部转发至前端 80 或后端 8000。具体路由与绑定随生产配置确定。', 'Целевые TCP 80 HTTP и TCP 443 HTTPS/TLS; внутри — порт 80 клиента либо 8000 сервера. Маршруты и привязки определяются производственной конфигурацией.'),
        ('目标部署方案\n尚未执行\nOPEN-010', 'Целевая схема\nНе выполнено\nOPEN-010'),
    ],
    [
        ('N-06\n目标数据库服务', 'Целевой сервис БД'),
        ('目标容器\nPostgreSQL', 'Целевой контейнер\nPostgreSQL'),
        ('承载业务数据库及持久数据卷。CMP-009 的 SQLAlchemy 层仍在 N-04，通过驱动访问此处数据；业务表与迁移待实现。', 'Размещает бизнес-БД и постоянный том. Слой SQLAlchemy CMP-009 остаётся в N-04 и обращается сюда через драйвер; бизнес-таблицы и миграции ещё предстоят.'),
        ('PostgreSQL/TCP，连接示例为 5432；目标为内部网络访问。最终服务名、端口及绑定待生产配置确定。', 'PostgreSQL/TCP, в примере порт 5432; целевой доступ — во внутренней сети. Имя сервиса, порт и привязки задаются при подготовке размещения.'),
        ('目标部署方案\n当前 Compose 未包含\n迁移未执行', 'Целевая схема\nВ текущем Compose отсутствует\nМиграция не выполнена'),
    ],
    [
        ('N-07\n学校核验来源', 'Источник проверки вуза'),
        ('外部依赖\n承载形式待定', 'Внешний источник\nТип уточняется'),
        ('CMP-010：学校确认的名单或平台，为管理员人工核验提供依据；学校负责维护。StudySpace 内部记录必要来源引用及审核决定。', 'CMP-010: подтверждённый вузом реестр или платформа для ручной проверки администратором; поддерживается вузом. StudySpace сохраняет необходимые ссылки и решение.'),
        ('IF-04 为人工核验依赖；具体渠道和端口待学校确认，未定义自动 API。', 'IF-04 — зависимость ручной проверки; канал и порты определяет вуз, автоматический API не определён.'),
        ('学校外部环境\nOPEN-003\nPARTIAL', 'Внешняя среда вуза\nOPEN-003\nPARTIAL'),
    ],
]

intro = (
    'StudySpace 采用 React 浏览器客户端与 FastAPI 模块化单体；开发使用 SQLite，目标部署选择 Docker Compose、Caddy 和 PostgreSQL。下表按课堂五个字段列出节点、制品与通信方式，并保留现有配置和目标方案的状态。',
    'StudySpace использует браузерный клиент React и модульный монолит FastAPI; для разработки — SQLite, для целевого размещения выбраны Docker Compose, Caddy и PostgreSQL. Таблица содержит пять полей задания и сохраняет статусы текущих конфигураций и целевой схемы.'
)
notes = [
    ('节点与组件关系', 'Связь узлов и компонентов',
     'N-02 是直接启动方式，N-03 与 N-04 是另一种本地容器方式；它们不是必须同时运行的多套服务。CMP-001 由 N-03 分发并在 N-01 浏览器执行；CMP-002 至 CMP-009 在同一后端部署边界内，不为各业务模块单独增加服务端口。',
     'N-02 — прямой локальный запуск, N-03 и N-04 — альтернативный запуск в контейнерах; одновременная работа обеих схем не требуется. CMP-001 раздаётся из N-03 и исполняется в браузере N-01; CMP-002–CMP-009 находятся в одной границе сервера без отдельных портов каждого модуля.'),
    ('数据库制品与持久化', 'Артефакты БД и сохранение данных',
     '本地容器的 SQLite 文件为 /app/backend/data/studyspace.db，studyspace-data 卷挂载到 /app/backend/data；它是 N-04 内访问的文件制品，不是独立数据库服务器。目标 PostgreSQL 替代该存储方式，迁移、回滚、并发与恢复须分别验证；Alembic 尚待引入。',
     'Файл SQLite локального контейнера — /app/backend/data/studyspace.db; том studyspace-data подключён к /app/backend/data. Это файловый артефакт N-04, а не отдельный сервер БД. Целевой PostgreSQL заменяет эту схему; миграция, откат, конкуренция и восстановление проверяются отдельно, Alembic ещё предстоит добавить.'),
    ('环境与实施状态', 'Среда и состояние реализации',
     '当前 Compose 为 development 的本地 HTTP 配置，不包含 Caddy 或 PostgreSQL。目标部署状态为 NOT_EXECUTED，服务器、域名与具体生产配置仍受 OPEN-010 约束。现有基础 API 不等于预约业务已完成；容器部署、HTTPS、数据库迁移与业务验收未据本表确认，业务用例保持 NOT_RUN。学校核验来源仍为 OPEN-003，不能据此假定已有 SSO 或自动 API。',
     'Текущий Compose — локальная конфигурация HTTP для development без Caddy и PostgreSQL. Статус целевого размещения NOT_EXECUTED; сервер, домен и производственная конфигурация связаны с OPEN-010. Базовые API не означают готовую бронь; таблица не подтверждает развёртывание контейнеров, HTTPS, миграцию или бизнес-приёмку, бизнес-кейсы остаются NOT_RUN. Источник вуза остаётся OPEN-003; готовые SSO и автоматический API не предполагаются.'),
]

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(0.6)
sec.left_margin = sec.right_margin = Inches(0.6)
normal = doc.styles['Normal']
normal.font.name = 'Arial'
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(4)
normal.paragraph_format.line_spacing = 1.08
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.insert(0, rf)
    for attr in ['ascii', 'hAnsi', 'cs']:
        rf.set(qn(f'w:{attr}'), 'Arial')
    rf.set(qn('w:eastAsia'), 'PingFang SC')
doc.styles['Title'].font.size = Pt(19)
doc.styles['Title'].paragraph_format.space_after = Pt(5)
doc.styles['Heading 1'].font.size = Pt(13)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(11)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(6)
doc.styles['Heading 2'].font.size = Pt(11)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(7)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(3)
# The bundled default template can carry a blue border in its Title style.
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)

def fmt_run(run, size=None, bold=False, color='000000', lang=None):
    run.font.name = 'Arial'
    if size:
        run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.insert(0, rf)
    for a in ['ascii', 'hAnsi', 'cs']:
        rf.set(qn('w:' + a), 'Arial')
    rf.set(qn('w:eastAsia'), 'PingFang SC')
    if lang:
        le = OxmlElement('w:lang')
        le.set(qn('w:val'), lang)
        if lang == 'zh-CN':
            le.set(qn('w:eastAsia'), lang)
        rpr.append(le)

def para(text, style=None, size=None, bold=False, lang=None):
    p = doc.add_paragraph(style=style)
    fmt_run(p.add_run(text), size, bold, lang=lang)
    return p

def pair(zh, ru, style=None):
    p = para(zh, style, lang='zh-CN')
    p.paragraph_format.keep_with_next = True
    p = para(ru, style, lang='ru-RU')
    return p

pair('StudySpace 节点与制品表', 'StudySpace Узлы и артефакты', 'Title')
pair('任务 1  节点与制品', 'Задание 1  Узлы и артефакты')
para('SS-DEP-T1-01  |  v1  |  2026-10-15  |  DRAFT', size=9)
pair(*intro)
pair('1 节点与制品表', '1 Таблица узлов и артефактов', 'Heading 1')

table = doc.add_table(rows=1, cols=5)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = [1.14, 0.91, 2.32, 1.79, 1.14]
for c, width in zip(table.columns, widths):
    c.width = Inches(width)
pr = table._tbl.tblPr
borders = OxmlElement('w:tblBorders')
for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
    el = OxmlElement('w:' + edge)
    el.set(qn('w:val'), 'single')
    el.set(qn('w:sz'), '5')
    el.set(qn('w:color'), 'D9D9D9')
    borders.append(el)
pr.append(borders)
layout = pr.find(qn('w:tblLayout'))
if layout is None:
    layout = OxmlElement('w:tblLayout')
    pr.append(layout)
layout.set(qn('w:type'), 'fixed')
hdr = OxmlElement('w:tblHeader')
table.rows[0]._tr.get_or_add_trPr().append(hdr)

def fill_cell(cell, val, idx, header=False):
    cell.width = Inches(widths[idx])
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cp = cell._tc.get_or_add_tcPr()
    margins = OxmlElement('w:tcMar')
    for edge, amount in [('top', 90), ('bottom', 90), ('left', 80), ('right', 80)]:
        e = OxmlElement('w:' + edge)
        e.set(qn('w:w'), str(amount))
        e.set(qn('w:type'), 'dxa')
        margins.append(e)
    cp.append(margins)
    for n, text in enumerate(val):
        p = cell.paragraphs[0] if n == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(4 if n == 0 else 0)
        p.paragraph_format.line_spacing = 1.04
        p.paragraph_format.widow_control = True
        p.paragraph_format.keep_with_next = False
        word_wrap = OxmlElement('w:wordWrap')
        word_wrap.set(qn('w:val'), '0')
        p._p.get_or_add_pPr().append(word_wrap)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if header or idx in (1, 4) else WD_ALIGN_PARAGRAPH.LEFT
        fmt_run(p.add_run(text), 9.5, bold=header, color='FFFFFF' if header else '000000', lang='zh-CN' if n == 0 else 'ru-RU')

for i, h in enumerate(headers):
    fill_cell(table.rows[0].cells[i], h, i, True)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), '29485D')
    table.rows[0].cells[i]._tc.get_or_add_tcPr().append(shd)
for ri, vals in enumerate(rows):
    row = table.add_row()
    no_split = OxmlElement('w:cantSplit')
    row._tr.get_or_add_trPr().append(no_split)
    for i, val in enumerate(vals):
        fill_cell(row.cells[i], val, i)
        shd = OxmlElement('w:shd')
        shd.set(qn('w:fill'), 'F0F4F7' if ri % 2 == 0 else 'FFFFFF')
        row.cells[i]._tc.get_or_add_tcPr().append(shd)

doc.add_paragraph().paragraph_format.space_after = Pt(1)
pair('2 部署关系与状态说明', '2 Связи размещения и статусы', 'Heading 1')
for n, (zhh, ruh, zh, ru) in enumerate(notes, 1):
    pair(f'2 {n} {zhh}', f'2 {n} {ruh}', 'Heading 2')
    p = pair(zh, ru)
    p.paragraph_format.keep_together = True

pair('3 项目依据', '3 Источники проекта', 'Heading 1')
pair('项目根目录  /Users/zhi/Documents/软件生命周期/studyspace', 'Корневой каталог проекта  /Users/zhi/Documents/软件生命周期/studyspace')
source_groups = [
    ('S1 S2', '系统设计及中俄对应版', 'Проектирование системы на двух языках', 'docs/zh-CN/04_System_Design.md\ndocs/ru-RU/04_System_Design.md'),
    ('S3 S4 S5 S6', '容器制品与通信配置', 'Артефакты и связи контейнеров', 'compose.yaml; frontend/Dockerfile; frontend/nginx.conf; backend/Dockerfile'),
    ('S7 S8 S9 S10 S11 S12 S18', '开发端口与实际代码基础', 'Порты разработки и существующий код', 'scripts/dev.sh; frontend/vite.config.ts; backend/app/config.py; backend/app/main.py; backend/app/database.py; .env.example; backend/pyproject.toml'),
    ('S13 S14 S15', '部署技术选择与外部前提', 'Выбор размещения и внешние условия', 'docs/data/project-decisions.json; docs/decisions/zh-CN/0004-docker-https.md\ndocs/decisions/ru-RU/0004-docker-https.md'),
    ('S16 S17 S19', '既有 ADR 与统一组件编号', 'Прежние ADR и единые ID компонентов', 'coursework/2026-10-14-architecture-decisions-v1/StudySpace_两份ADR架构决策记录_中俄双语_v1.md\n../StudySpace_组件交付/StudySpace_组件候选与组件图_中俄双语_v2.md; AGENTS.md'),
]
for ids, zh, ru, paths in source_groups:
    p = para(f'{ids}  {zh} / {ru}', size=9)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(2)
    p = para(paths, size=8.5)
    p.paragraph_format.space_after = Pt(3)

doc.core_properties.title = 'StudySpace 节点与制品表 Узлы и артефакты'
doc.core_properties.subject = '任务 1 / Задание 1'
doc.core_properties.author = 'StudySpace'
doc.core_properties.keywords = 'StudySpace, 中俄双语, узлы, артефакты, deployment'
doc.save(DOCX)

content = {
    'doc_id': 'SS-DEP-T1-01', 'version': 'v1', 'date': '2026-10-15', 'status': 'DRAFT',
    'intro': intro, 'headers': headers, 'rows': rows, 'notes': notes,
    'sources': [{'id': i, 'path': str(p), 'zh': zh, 'ru': ru, 'sha256': fingerprints[str(p)]} for i, p, zh, ru in sources],
}
(DATA / 'content.json').write_text(json.dumps(content, ensure_ascii=False, indent=2), encoding='utf-8')
(DATA / 'fonts.conf').write_text('''<?xml version="1.0"?>
<!DOCTYPE fontconfig SYSTEM "urn:fontconfig:fonts.dtd">
<fontconfig>
<dir>/System/Library/Fonts</dir><dir>/System/Library/Fonts/Supplemental</dir>
<dir>/System/Library/AssetsV2/com_apple_MobileAsset_Font7</dir>
<dir>/Library/Fonts</dir><dir>/Users/zhi/Library/Fonts</dir>
<cachedir>/private/tmp/studyspace-nodes-fontcache</cachedir>
<match target="pattern"><test name="family"><string>sans-serif</string></test><edit name="family" mode="prepend"><string>Arial</string><string>PingFang SC</string></edit></match>
</fontconfig>
''', encoding='utf-8')
md = ['# StudySpace 节点与制品表', '# StudySpace Узлы и артефакты', 'SS-DEP-T1-01 · v1 · 2026-10-15 · DRAFT', '', *intro, '']
md += ['| ' + ' | '.join('<br>'.join(x) for x in headers) + ' |', '| ' + ' | '.join(['---'] * 5) + ' |']
for row in rows:
    md.append('| ' + ' | '.join('<br>'.join(cell).replace('\n', '<br>') for cell in row) + ' |')
for zh, ru, zht, rut in notes:
    md += ['', f'## {zh} / {ru}', zht, rut]
md += ['', '## 项目依据 Источники проекта']
md += [f'- {i}: {p}' for i, p, _, _ in sources]
(OUT / f'{STEM}.md').write_text('\n\n'.join(md), encoding='utf-8')
(OUT / 'README.md').write_text(f'''# StudySpace 节点与制品 Узлы и артефакты

SS-DEP-T1-01 · v1 · 2026-10-15 · DRAFT

交付 / Результат [{STEM}.docx]({STEM}.docx)

按课堂五列填写 7 个节点或外部依赖；中文在上、俄语在下，Word 表格可编辑。
Заполнены пять полей для 7 узлов или внешних зависимостей; китайский над русским, таблица Word редактируется.

依据现有源文件区分开发配置、本地容器候选和目标部署方案；业务实现和业务测试状态未改变。
По исходным файлам различаются разработка, локальный вариант контейнеров и целевая схема; статусы реализации и бизнес-тестов сохранены.

内容 / Содержимое data/content.json
文档核查 / Проверка документа data/verification.json
渲染检查 / Проверка отрисовки work/render/
''', encoding='utf-8')

read_back = Document(DOCX)
ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
with ZipFile(DOCX) as zf:
    xml = etree.fromstring(zf.read('word/document.xml'))
    names = zf.namelist()
alltext = '\n'.join(p.text for p in read_back.paragraphs) + '\n' + '\n'.join(c.text for r in read_back.tables[0].rows for c in r.cells)
tests = {}
tests['five_columns_and_seven_nodes'] = len(read_back.tables) == 1 and len(read_back.tables[0].columns) == 5 and len(read_back.tables[0].rows) == 8
tests['all_cells_match_bilingual_source'] = all(read_back.tables[0].cell(ri, ci).text == '\n'.join(val) for ri, vals in enumerate([headers, *rows]) for ci, val in enumerate(vals))
tests['all_cells_paired_zh_ru'] = all(len(c.paragraphs) == 2 and re.search('[\u4e00-\u9fff]', c.paragraphs[0].text) and re.search('[А-Яа-яЁё]', c.paragraphs[1].text) for r in read_back.tables[0].rows for c in r.cells)
tests['node_ids_unique'] = len(re.findall(r'\bN-0[1-7]\b', '\n'.join(r.cells[0].text for r in read_back.tables[0].rows))) == 7
tests['components_covered'] = all(f'CMP-{i:03d}' in alltext for i in [1, 2, 8, 9, 10])
tests['existing_local_ports'] = all(x in alltext for x in ['5173', '8080', '8000', '127.0.0.1'])
tests['sqlite_file_not_server'] = '无网络端口' in alltext and 'без сетевого порта' in alltext
tests['target_status_explicit'] = all(x in alltext for x in ['NOT_EXECUTED', 'OPEN-010', 'OPEN-003', 'PARTIAL', 'NOT_RUN'])
tests['postgres_not_in_current_compose'] = '当前 Compose 未包含' in alltext and 'postgres' not in (ROOT / 'compose.yaml').read_text().lower()
tests['current_api_paths_match_source'] = all(x in (ROOT / 'backend/app/main.py').read_text() for x in ['/api/v1/health', '/api/v1/ready', '/api/v1/docs'])
tests['repeating_header'] = bool(xml.xpath('//w:tbl/w:tr[1]/w:trPr/w:tblHeader', namespaces=ns))
tests['rows_do_not_split'] = len(xml.xpath('//w:tbl/w:tr/w:trPr/w:cantSplit', namespaces=ns)) == 7
tests['no_fixed_row_heights'] = not xml.xpath('//w:trHeight', namespaces=ns)
tests['editable_not_screenshot'] = not xml.xpath('//w:drawing', namespaces=ns) and len(xml.xpath('//w:tbl', namespaces=ns)) == 1
tests['no_headers_footers'] = not any(re.fullmatch(r'word/(header|footer)\d+\.xml', name) for name in names)
with ZipFile(DOCX) as zf:
    styles_xml = etree.fromstring(zf.read('word/styles.xml'))
tests['no_title_border'] = not styles_xml.xpath('//w:style[@w:styleId="Title"]//w:pBdr', namespaces=ns)
tests['sources_unchanged'] = all(hashlib.sha256(Path(p).read_bytes()).hexdigest() == sha for p, sha in fingerprints.items())
report = {'document': str(DOCX), 'checked_at_utc': datetime.now(timezone.utc).isoformat(), 'execution_environment': 'Local macOS; bundled Codex Python; static document verification', 'all_pass': all(tests.values()), 'checks': tests, 'visual_review': 'PENDING', 'business_tests': 'NOT_RUN', 'source_fingerprints': fingerprints}
(DATA / 'verification.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({'document': str(DOCX), 'all_pass': report['all_pass'], 'checks': len(tests), 'failed': [k for k,v in tests.items() if not v]}, ensure_ascii=False))
if not report['all_pass']:
    raise SystemExit(1)
