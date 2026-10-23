from pathlib import Path
from datetime import datetime, timezone
from zipfile import ZipFile
import ast
import json
import hashlib
import re
from lxml import etree
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path('/Users/zhi/Documents/软件生命周期/studyspace/coursework/2026-10-15-task1-scrum-v1')
SOURCE = Path('/Users/zhi/Documents/软件生命周期/studyspace/coursework/2026-10-15-scrum-kanban-v1')
OUT = ROOT / '任务一_Scrum冲刺规划_中俄双语_v1.docx'
PAGES = json.loads((SOURCE / 'data/content.json').read_text(encoding='utf-8'))[:5]
PAGES[0]['items'][0]['value'] = {
    'ru': 'Задание 1 Планирование спринта Scrum',
    'zh': '任务一 Scrum冲刺规划'
}
PAGES[0]['items'][2]['value'] = {
    'ru': '1 Продукт и исходные условия',
    'zh': '1 产品背景与已知条件'
}
PAGES.append({'landscape': False, 'items': [
    {'kind': 'h', 'value': {'ru': '2 Источник задания и методическая основа', 'zh': '2 题目来源与方法依据'}},
    {'kind': 'p', 'small': True, 'value': {
        'ru': 'Условие — задание 1 «Планирование спринта Scrum» на предоставленном изображении. Выполнены все пять шагов: роли, цель спринта, оценка историй 1–6, выбор историй в пределах 20 очков и Definition of Done из пяти пунктов.',
        'zh': '题目依据为最初图片中的任务一“Scrum冲刺规划”。已覆盖五项要求：角色分配、一句话冲刺目标、故事1—6估算、总点数不超过20的冲刺选择，以及5项完成定义。'
    }},
    {'kind': 'p', 'small': True, 'value': {
        'ru': '[1] Schwaber K., Sutherland J. The Scrum Guide. 2020. Использованы положения о ролях, Sprint Planning, Sprint Backlog и Definition of Done. Дата обращения 15.10.2026.',
        'zh': '[1] Scrum指南，2020年版。用于核对角色、冲刺规划、冲刺待办与完成定义，查阅日期为2026年10月15日。'
    }},
    {'kind': 'link', 'value': {'label': 'https://scrumguides.org/scrum-guide.html', 'url': 'https://scrumguides.org/scrum-guide.html'}}
]})

# Record the complete bilingual inventory before authoring the Word artifact.
(ROOT / 'data/content.json').write_text(json.dumps(PAGES, ensure_ascii=False, indent=2), encoding='utf-8')
fontconf = (SOURCE / 'data/fonts.conf').read_text().replace('scrum-kanban-fontcache', 'task1-scrum-fontcache')
(ROOT / 'data/fonts.conf').write_text(fontconf)

# Reuse only the formatting helpers, without executing the source builder.
helper_names = {'font', 'section_setup', 'run_text', 'write_value', 'draw_table', 'hyperlink', 'all_strings'}
tree = ast.parse((SOURCE / 'work/build_document.py').read_text(encoding='utf-8'))
helpers = ast.Module(body=[n for n in tree.body if isinstance(n, ast.FunctionDef) and n.name in helper_names], type_ignores=[])
assert {n.name for n in helpers.body} == helper_names
exec(compile(helpers, str(SOURCE / 'work/build_document.py'), 'exec'))

doc = Document()
doc.core_properties.title = '任务一 Scrum冲刺规划'
doc.core_properties.subject = '中俄双语软件生命周期课程实践'
doc.core_properties.author = ''
doc.core_properties.last_modified_by = ''
doc.core_properties.version = '1.0'
doc.core_properties.created = datetime(2026, 10, 15, 0, 0, tzinfo=timezone.utc)
doc.core_properties.modified = datetime(2026, 10, 15, 0, 0, tzinfo=timezone.utc)

for name, size, bold in [('Normal', 11, False), ('Title', 16, True), ('Heading 1', 13.5, True), ('Heading 2', 12, True)]:
    font(doc.styles[name], size, bold)
    pf = doc.styles[name].paragraph_format
    pf.space_after = Pt(6)
    pf.space_before = Pt(0 if name == 'Normal' else 8)
    pf.line_spacing = 1.08
    pf.widow_control = True
for sty in doc.styles:
    for border in sty.element.findall('.//' + qn('w:pBdr')):
        border.getparent().remove(border)

section_setup(doc.sections[0], False)
TABLE_DIMS = []
for pg in PAGES:
    for block in pg['items']:
        kind, value = block['kind'], block['value']
        if kind in ('p', 'h', 'title'):
            p = doc.add_paragraph(style='Title' if kind == 'title' else ('Heading 1' if kind == 'h' else 'Normal'))
            p.paragraph_format.keep_together = True
            if kind == 'h':
                p.paragraph_format.keep_with_next = True
            snap = OxmlElement('w:snapToGrid')
            snap.set(qn('w:val'), '0')
            p._p.get_or_add_pPr().append(snap)
            size = 16 if kind == 'title' else (13.5 if kind == 'h' else (10.5 if block.get('small') else 11))
            write_value(p, value, size=size, bold=kind in ('title', 'h'))
        elif kind == 'table':
            draw_table(value)
        elif kind == 'link':
            hyperlink(value)
        else:
            raise ValueError(kind)

doc.save(OUT)
loaded = Document(OUT)
ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
with ZipFile(OUT) as archive:
    root = etree.fromstring(archive.read('word/document.xml'))
    text = '\n'.join(root.xpath('//w:t/text()', namespaces=ns))
    styles = etree.fromstring(archive.read('word/styles.xml'))
    assert not styles.xpath('//w:pBdr', namespaces=ns)
    assert not any('header' in s or 'footer' in s for s in archive.namelist())

required = list(all_strings(PAGES))
missing = [x for x in required if x not in text]
assert not missing, missing
assert '任务一 Scrum冲刺规划' in text and 'Задание 1 Планирование спринта Scrum' in text
assert '任务二' not in text and 'Kanban' not in text
assert TABLE_DIMS == [[6, 3], [13, 3], [7, 4], [6, 3], [6, 2]]
actual_dims = [[len(t.rows), len(t.columns)] for t in loaded.tables]
assert TABLE_DIMS == actual_dims
estimates = {row[0]: int(row[2]) for row in loaded.tables[2].rows[1:] for row in [[c.text for c in row.cells]]}
assert estimates == {'US01': 3, 'US02': 3, 'US03': 5, 'US04': 5, 'US05': 8, 'US06': 3}
assert all(n in [1, 2, 3, 5, 8, 13] for n in estimates.values())
selected = ['US01', 'US02', 'US03', 'US04', 'US06']
assert sum(estimates[s] for s in selected) == 19 <= 20
dod_heading = next(i for i, x in enumerate(PAGES[4]['items']) if x['kind'] == 'h' and x['value']['ru'].startswith('1 9'))
dod_items = PAGES[4]['items'][dod_heading + 1:dod_heading + 6]
assert len(dod_items) == 5 and all(x['value']['ru'].startswith(str(i+1)) for i,x in enumerate(dod_items))
assert len(root.xpath('//w:tbl/w:tr[1]/w:trPr/w:tblHeader', namespaces=ns)) == 5

report = {
    'artifact': str(OUT),
    'source_task': 'Original first attachment Task 1 Scrum sprint planning',
    'date': '2026-10-15',
    'environment': 'Codex local desktop bundled Python and LibreOffice',
    'status': 'STRUCTURAL_PASS',
    'required_text_items': len(required),
    'missing_text': missing,
    'tables': actual_dims,
    'original_backlog_stories': 12,
    'estimated_stories': estimates,
    'selected_stories': selected,
    'sprint_points': 19,
    'capacity': 20,
    'dod_items': 5,
    'all_five_task_steps_covered': True,
    'bilingual_pairs_present': True,
    'tables_editable': True,
    'visual_review': 'PENDING',
    'sha256': hashlib.sha256(OUT.read_bytes()).hexdigest(),
    'previous_task2_sha256': hashlib.sha256((SOURCE / '任务二_Scrum冲刺规划与Kanban看板设计_中俄双语_v1.docx').read_bytes()).hexdigest()
}
(ROOT / 'data/verification.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(report, ensure_ascii=False, indent=2))
