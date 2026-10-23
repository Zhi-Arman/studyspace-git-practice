from __future__ import annotations

import hashlib
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
PROJECT = ROOT.parents[1]
CONTENT = ROOT / "data/content.json"
OUTPUT = ROOT / "StudySpace_任务一端点列表_中俄双语_v1.docx"
VERIFY = ROOT / "data/verification.json"
data = json.loads(CONTENT.read_text(encoding="utf-8"))


def font(run, *, size=8.4, bold=False, color=None):
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key, val in (("ascii", "Arial"), ("hAnsi", "Arial"), ("eastAsia", "Heiti SC")):
        rfonts.set(qn("w:" + key), val)


def pair(container, value, *, size=8.4, bold=False):
    for lang in ("zh", "ru"):
        p = container.add_paragraph()
        p.paragraph_format.space_after = Pt(1 if lang == "zh" else 0)
        p.paragraph_format.line_spacing = 1.03
        font(p.add_run(value[lang]), size=size, bold=bold)


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def prevent_split(row, repeat=False):
    trpr = row._tr.get_or_add_trPr()
    flag = OxmlElement("w:cantSplit")
    trpr.append(flag)
    if repeat:
        repeat_flag = OxmlElement("w:tblHeader")
        repeat_flag.set(qn("w:val"), "true")
        trpr.append(repeat_flag)


doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width = Inches(11.69)
sec.page_height = Inches(8.27)
sec.left_margin = sec.right_margin = Inches(0.5)
sec.top_margin = Inches(0.45)
sec.bottom_margin = Inches(0.42)
sec.header_distance = sec.footer_distance = Inches(0.2)

normal = doc.styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(8.5)
normal.paragraph_format.space_after = Pt(3)
normal.paragraph_format.line_spacing = 1.05

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(1)
font(p.add_run(data["title"]["zh"]), size=15, bold=True, color=(18, 67, 83))
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(4)
font(p.add_run(data["title"]["ru"]), size=12, bold=True, color=(18, 67, 83))
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(6)
font(p.add_run(f'{data["document_id"]}  ·  {data["version"]}  ·  {data["date"]}  ·  {data["status"]}'), size=8, color=(85, 95, 104))

for lang in ("zh", "ru"):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1 if lang == "zh" else 6)
    font(p.add_run(data["intro"][lang]), size=8.5)

widths = [2.25, 1.80, 2.35, 1.45, 2.80]
table = doc.add_table(rows=1, cols=5)
table.autofit = False
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = "Table Grid"
for i, width in enumerate(widths):
    table.columns[i].width = Inches(width)
for i, item in enumerate(data["headers"]):
    c = table.rows[0].cells[i]
    c.width = Inches(widths[i])
    c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    shade(c, "DCEEF0")
    c.paragraphs[0].text = ""
    pair(c, item, size=8.5, bold=True)
prevent_split(table.rows[0], repeat=True)

for n, item in enumerate(data["rows"]):
    row = table.add_row()
    prevent_split(row)
    values = [item["message"], item["method_path"], item["requirement_story"], item["role"], item["reason"]]
    for i, value in enumerate(values):
        c = row.cells[i]
        c.width = Inches(widths[i])
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if n % 2:
            shade(c, "F5F8F9")
        c.paragraphs[0].text = ""
        if isinstance(value, dict):
            pair(c, value, size=8.1 if i in (0, 2, 4) else 8.4)
        else:
            p = c.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            font(p.add_run(value), size=9, bold=True, color=(21, 72, 93))

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(5)
p.paragraph_format.space_after = Pt(1)
font(p.add_run("说明 / Примечания"), size=9, bold=True, color=(18, 67, 83))
for note in data["notes"]:
    for lang in ("zh", "ru"):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        font(p.add_run(note[lang]), size=8.1)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(4)
p.paragraph_format.space_after = Pt(1)
font(p.add_run("依据 / Источники"), size=8.5, bold=True, color=(18, 67, 83))
for source in data["sources"]:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    font(p.add_run(source), size=7.2, color=(83, 91, 96))

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)

check = Document(OUTPUT)
assert len(check.tables) == 1
assert len(check.tables[0].rows) == 1 + len(data["rows"])
assert all(len(row.cells) == 5 for row in check.tables[0].rows)
all_text = "\n".join(p.text for p in check.paragraphs)
all_text += "\n" + "\n".join(c.text for row in check.tables[0].rows for c in row.cells)
required = [data["title"]["zh"], data["title"]["ru"]]
for key in ("intro",):
    required.extend(data[key].values())
for item in data["headers"] + data["notes"]:
    required.extend(item.values())
for item in data["rows"]:
    required.append(item["method_path"])
    for key in ("message", "requirement_story", "role", "reason"):
        required.extend(item[key].values())
required.extend(data["sources"])
missing = [value for value in required if value not in all_text]
assert not missing, missing
assert all((PROJECT / source).exists() for source in data["sources"])
result = {
    "document": str(OUTPUT),
    "sha256": hashlib.sha256(OUTPUT.read_bytes()).hexdigest(),
    "status": data["status"],
    "endpoint_count": len(data["rows"]),
    "table_dimensions": [len(check.tables[0].rows), len(check.tables[0].columns)],
    "paired_text_checks": len(required),
    "structural_result": "PASS",
    "business_tests": "NOT_RUN"
}
VERIFY.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps(result, ensure_ascii=False, indent=2))
