from pathlib import Path
from datetime import datetime
from zoneinfo import ZoneInfo
from zipfile import ZipFile
import hashlib
import json
import re
import subprocess
from docx import Document
from lxml import etree

ROOT=Path(__file__).resolve().parents[1]
path=ROOT/'StudySpace_可控合并冲突_中俄双语_v1.docx'
doc=Document(path)
content=json.loads((ROOT/'artifacts/report-content.json').read_text())
compact=lambda s:re.sub(r'\s+','',s)
flat=compact(''.join(p.text for p in doc.paragraphs)+''.join(c.text for t in doc.tables for row in t.rows for c in row.cells))
expected=[]
for page in content:
    expected.extend([page['zh'],page['ru']])
    for b in page['blocks']:
        if b['type']=='pair':expected.extend([b['zh'],b['ru']])
        elif b['type']=='code':expected.append(b['text'])
        else:
            expected.extend(b['heads']);expected.extend(x for row in b['rows'] for x in row)
assert all(compact(x) in flat for x in expected)
with ZipFile(path) as z:
    assert z.testzip() is None
    for name in z.namelist():
        if name.endswith('.xml'):etree.fromstring(z.read(name))
    assert b'<w:pBdr' not in z.read('word/document.xml')
    assert b'<w:pBdr' not in z.read('word/styles.xml')
md=(ROOT/'README.md').read_text()
for target in re.findall(r'\]\(([^)]+)\)',md):
    if not target.startswith(('https://','http://')):assert (ROOT/target).exists(),target
assert len(re.findall(r'^\| CF-0[1-6] \|',md,re.M))==6
E=json.loads((ROOT/'artifacts/execution.json').read_text())
P=json.loads((ROOT/'artifacts/preparation.json').read_text())
assert all(c['passed'] for c in E['checks']) and all(c['passed'] for c in P['checks'])
assert E['status']=='EXPECTED_CONFLICT_PRESERVED' and E['merge_exit_code']==1
B=ROOT/'author-b'
def git(*args):return subprocess.check_output(['git',*args],cwd=B,text=True).strip()
assert git('status','--porcelain=v1')=='UU docs/build-check.md'
assert git('rev-parse','MERGE_HEAD')==E['b_head']
assert git('rev-parse','HEAD')==E['a_merge_commit']
assert (B/E['file']).read_text()==E['conflict_text']==(ROOT/'artifacts/conflict-file.txt').read_text()
assert git('merge-base','HEAD','MERGE_HEAD')==E['common_base']
assert all(git('show',f':{i}:'+E['file'])+'\n'==E['stages'][str(i)] for i in (1,2,3))
assert E['build_execution']=='NOT_RUN' and E['conflict_resolution']=='NOT_RUN' and E['merge_abort']=='NOT_RUN'
assert 'NOT_RUN' in md and 'NOT_VERIFIED' in md
images=sorted((ROOT/'artifacts/render-final').glob('page-*.png'))
assert len(images)==3,len(images)
result=dict(checked_at=datetime.now(ZoneInfo('Asia/Shanghai')).isoformat(),docx_sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
            zip_xml_integrity='PASS',bilingual_content_parity='PASS',source_text_items=len(expected),markdown_links='PASS',
            six_task_items_present=True,real_conflict_still_preserved=True,index_three_stages_match=True,
            expected_merge_failure=True,build_status='NOT_RUN',pages=len(images),all_pages_visually_reviewed=False)
(ROOT/'artifacts/document-verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(result,ensure_ascii=False,indent=2))
