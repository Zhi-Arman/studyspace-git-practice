from pathlib import Path
import json
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
E=json.loads((ROOT/'artifacts/execution.json').read_text())
P=json.loads((ROOT/'artifacts/preparation.json').read_text())
TITLE='StudySpace_可控合并冲突_中俄双语_v1'
pages=[]
def page(zh,ru):
    blocks=[];pages.append(dict(zh=zh,ru=ru,blocks=blocks));return blocks
def pair(b,zh,ru):b.append(dict(type='pair',zh=zh,ru=ru))
def code(b,text):b.append(dict(type='code',text=text))
def table(b,heads,rows,widths):b.append(dict(type='table',heads=heads,rows=rows,widths=widths))

b=page('StudySpace 可控合并冲突','StudySpace Контролируемый конфликт')
pair(b,'SS-GIT-CONFLICT-01 · v1.0 · 2026-10-22 · Asia/Shanghai · 文档 DRAFT',
       'SS-GIT-CONFLICT-01 · v1.0 · 22.10.2026 · Asia/Shanghai · Документ DRAFT')
pair(b,'可控冲突已真实复现。A、B 从同一提交分支，仅修改同一行；A 已先合并，B 合并到更新后的 main 时发生内容冲突。B 工作区保留未解决状态，任务达到预期终点。',
       'Контролируемый конфликт реально воспроизведён. A и B созданы от одного коммита и изменяют одну строку. A слит первым; слияние B в обновлённую main вызвало конфликт содержимого. Каталог B сохранён в неразрешённом состоянии; цель задания достигнута.')
code(b,'A PR '+E['a_pr_url']+'\nB PR '+E['b_pr_url'])
table(b,['编号\nКод','图片要求\nТребование на фото','实际结果\nРезультат'],[
['CF-01','创建 docs/build-check.md\nСоздайте учебный файл docs/build-check.md','已在共同基线提交中创建\nСоздан в общем базовом коммите'],
['CF-02','添加一行构建检查命令\nДобавьте одну строку с командой проверки сборки','文件恰好一行 Sphinx 命令\nВ файле ровно одна команда Sphinx'],
['CF-03','A、B 从同一提交创建独立分支\nA и B создают отдельные ветки от одного коммита','共同父提交 '+E['common_base'][:7]+'\nОбщий родитель '+E['common_base'][:7]],
['CF-04','每人只修改这一行\nКаждый меняет только эту строку','A 加 -W；B 加 -n；各自 1 行替换\nA добавил -W, B добавил -n; по одной замене'],
['CF-05','A 先合并自己的分支\nA объединяет свою ветку первым','PR #'+str(E['a_pr_number'])+' 已合并；main 使用 A 命令\nPR #'+str(E['a_pr_number'])+' слит; main содержит команду A'],
['CF-06','B 尝试合并并得到冲突\nB пытается объединить свою ветку и получает конфликт','退出码 1；UU docs/build-check.md；PR #'+str(E['b_pr_number'])+' 冲突\nКод 1; UU docs/build-check.md; PR #'+str(E['b_pr_number'])+' конфликтует'],
],[1.8,7.7,7.9])
pair(b,'沿用用户已授权的单账号自动化角色演示：课堂 A 为支浩然，B 为俞智文；实际 GitHub 账号为 Zhi-Arman。真人本人执行与独立批准为 NOT_VERIFIED。',
       'Использована разрешённая пользователем автоматическая демонстрация одним аккаунтом: учебный A 支浩然, B 俞智文; фактический аккаунт GitHub Zhi-Arman. Личное выполнение и независимое одобрение студентов — NOT_VERIFIED.')

b=page('2 同源分支与真实冲突证据','2 Общая база и свидетельства конфликта')
pair(b,'先在 docs/conflict-base 创建文件并提交，再分别创建 docs/conflict-a 与 docs/conflict-b。两次修改的父提交和 merge-base 均等于共同基线，排除了分支起点不同造成的干扰。',
       'Сначала файл создан и закоммичен в docs/conflict-base, затем созданы docs/conflict-a и docs/conflict-b. Родитель обоих изменений и merge-base совпадают с общей базой; обе ветки имеют одинаковую исходную версию.')
table(b,['版本\nВерсия','文件中的唯一一行\nЕдинственная строка файла'],[
['基线\nБаза',E['commands_text']['baseline'].strip()],
['A\nA',E['commands_text']['A'].strip()],
['B\nB',E['commands_text']['B'].strip()],
],[2.5,14.9])
pair(b,'A 合并后，B 先快进同步 main，再执行以下真实命令。退出码为 1，Git 明确报告 docs/build-check.md 的内容冲突。',
       'После слияния A каталог B получил main через fast-forward и выполнил реальную команду ниже. Код возврата 1; Git сообщил о конфликте содержимого docs/build-check.md.')
code(b,'git merge --no-ff --no-edit docs/conflict-b\n\n'+E['merge_output'].strip())
pair(b,'冲突标记由 Git 自动写入，以下为现场文件原文。HEAD 一侧是 main 上的 A 版本，另一侧是 B 分支版本。',
       'Маркеры конфликта записаны Git автоматически. Ниже точное содержимое файла: сторона HEAD — версия A из main, другая сторона — версия ветки B.')
code(b,E['conflict_text'].strip())
table(b,['索引阶段\nСтадия','来源\nИсточник','已验证内容\nПроверенное содержание'],[
['1','共同祖先\nОбщая база','等于基线命令\nСовпадает с базовой командой'],
['2','ours 即 main 上的 A\nours версия A из main','等于带 -W 的 A 命令\nКоманда A с -W'],
['3','theirs 即 B 分支\ntheirs ветка B','等于带 -n 的 B 命令\nКоманда B с -n'],
],[2.7,6.1,8.6])

b=page('3 验证结果与保留现场','3 Результаты проверки и сохранённое состояние')
pair(b,f'执行时间 {E["executed_at"]}。准备检查 {len(P["checks"])}/{len(P["checks"])}、最终流程检查 {len(E["checks"])}/{len(E["checks"])} 全部通过。退出码 1 是本题预期的冲突结果。',
       f'Время выполнения {E["executed_at"]}. Пройдены {len(P["checks"])}/{len(P["checks"])} подготовительных и {len(E["checks"])}/{len(E["checks"])} итоговых проверок. Код 1 — ожидаемый результат конфликта в этом задании.')
table(b,['对象\nОбъект','保留状态\nСохранённое состояние'],[
['GitHub main','A 合并提交 '+E['a_merge_commit'][:7]+'；未包含 B 的冲突解决\nСлияние A '+E['a_merge_commit'][:7]+'; разрешение конфликта B не добавлено'],
['A PR / B PR','A 已合并；B 保持打开，GitHub mergeable=false\nA слит; B открыт, GitHub mergeable=false'],
['A 工作区\nКаталог A','main 与 GitHub 一致，工作区干净\nmain совпадает с GitHub, дерево чистое'],
['B 工作区\nКаталог B','main 上合并未完成；唯一未合并文件为 docs/build-check.md\nНезавершённое слияние в main; единственный неслитый файл docs/build-check.md'],
['分支与既有文件\nВетки и прежние файлы','基线、A、B、旧 docs/preview 分支均保留；既有共享文档和 StudySpace 根仓库未改\nБаза, A, B и прежняя docs/preview сохранены; прежние документы и корень StudySpace не изменены'],
],[5.0,12.4])
pair(b,'构建命令是本次冲突练习的文本。共享仓库没有完整 Sphinx 配置，构建执行与业务测试均为 NOT_RUN；不能据此声称文档网站构建通过。冲突解决留给后续任务。',
       'Команды сборки используются как текст упражнения. В общем репозитории нет полной конфигурации Sphinx: сборка и бизнес-тесты имеют статус NOT_RUN. Успешная сборка сайта не заявляется; разрешение конфликта оставлено следующему заданию.')
pair(b,'交付根目录如下。实际冲突文件位于 author-b/docs/build-check.md；可以查看文件及 git status，继续下一步前应保留当前合并状态。',
       'Ниже корень материалов. Фактический конфликт находится в author-b/docs/build-check.md. Для просмотра откройте файл и git status; до следующего шага сохраняйте текущее состояние слияния.')
code(b,str(ROOT)+'\n\nREADME.md\n'+TITLE+'.docx\nauthor-b/docs/build-check.md\nartifacts/task.jpg\nartifacts/preparation.json\nartifacts/execution.json\nartifacts/merge-output.txt\nartifacts/status-conflict.txt\nartifacts/index-stages.txt\nartifacts/conflict-file.txt\nartifacts/stage-1.txt  stage-2.txt  stage-3.txt')

doc=Document();s=doc.sections[0]
s.page_width=Cm(21);s.page_height=Cm(29.7);s.top_margin=Cm(1.6);s.bottom_margin=Cm(1.6);s.left_margin=Cm(1.8);s.right_margin=Cm(1.8)
for name in ('Normal','Title','Subtitle','Heading 1'):
    sty=doc.styles[name];sty.font.name='Arial';sty.font.color.rgb=RGBColor(0,0,0)
    sty._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Hiragino Sans GB')
n=doc.styles['Normal'];n.font.size=Pt(10.5);n.paragraph_format.space_after=Pt(4);n.paragraph_format.line_spacing=1.03
doc.styles['Title'].font.size=Pt(20);doc.styles['Subtitle'].font.size=Pt(13);doc.styles['Heading 1'].font.size=Pt(16)
cs=doc.styles.add_style('Code',1);cs.font.size=Pt(8.5);cs.font.name='Courier New';cs.paragraph_format.space_after=Pt(6);cs.paragraph_format.line_spacing=1
doc.core_properties.title='StudySpace 可控合并冲突 Контролируемый конфликт';doc.core_properties.author='StudySpace'
def paragraph(text,style=None,ru=False):
    p=doc.add_paragraph(style=style)
    if ru:p.paragraph_format.space_after=Pt(6)
    for i,line in enumerate(text.split('\n')):
        if i:p.add_run().add_break()
        p.add_run(line)
    return p
md=[]
for pi,pag in enumerate(pages):
    h=paragraph(pag['zh'],'Title' if pi==0 else 'Heading 1')
    if pi:h.paragraph_format.page_break_before=True
    paragraph(pag['ru'],'Subtitle');md.extend([('#' if pi==0 else '##')+' '+pag['zh'],'',pag['ru'],''])
    for block in pag['blocks']:
        if block['type']=='pair':
            paragraph(block['zh']);paragraph(block['ru'],ru=True);md.extend([block['zh'],'',block['ru'],''])
        elif block['type']=='code':
            paragraph(block['text'],'Code');md.extend(['```text',block['text'],'```',''])
        else:
            t=doc.add_table(rows=1,cols=len(block['heads']));t.style='Table Grid';t.autofit=False
            for j,width in enumerate(block['widths']):t.columns[j].width=Cm(width)
            for i,values in enumerate([block['heads']]+block['rows']):
                row=t.rows[0] if i==0 else t.add_row();row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
                if i==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
                for j,value in enumerate(values):
                    cell=row.cells[j];cell.width=Cm(block['widths'][j]);cell.text=value
                    for p in cell.paragraphs:
                        p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1
                        for r in p.runs:r.font.size=Pt(9.5);r.bold=(i==0)
                    if i==0:
                        sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EEF4');cell._tc.get_or_add_tcPr().append(sh)
            doc.add_paragraph().paragraph_format.space_after=Pt(0)
            md.append('| '+' | '.join(x.replace('\n','<br>') for x in block['heads'])+' |');md.append('| '+' | '.join('---' for _ in block['heads'])+' |')
            md.extend('| '+' | '.join(x.replace('\n','<br>') for x in row)+' |' for row in block['rows']);md.append('')
for root in (doc._element,doc.styles.element):
    for border in root.xpath('.//w:pBdr'):border.getparent().remove(border)
for p in list(doc.paragraphs)+[p for t in doc.tables for row in t.rows for c in row.cells for p in c.paragraphs]:
    for r in p.runs:
        rf=r._element.get_or_add_rPr().get_or_add_rFonts()
        for attr in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):rf.attrib.pop(qn('w:'+attr),None)
        for key in ('ascii','hAnsi'):rf.set(qn('w:'+key),'Courier New' if p.style.name=='Code' else 'Arial')
        rf.set(qn('w:eastAsia'),'Hiragino Sans GB')
doc.save(ROOT/(TITLE+'.docx'))
md.extend(['## 文件入口 Входные файлы','',f'[双语 Word Двуязычный Word]({TITLE}.docx)',
           '[A PR 已合并 A PR слит]('+E['a_pr_url']+')',
           '[B PR 冲突 B PR конфликт]('+E['b_pr_url']+')',
           '[实际冲突文件 Фактический конфликт](author-b/docs/build-check.md)',
           '[完整流程检查 Проверки](artifacts/execution.json)',''])
(ROOT/'README.md').write_text('\n'.join(md))
(ROOT/'artifacts/report-content.json').write_text(json.dumps(pages,ensure_ascii=False,indent=2)+'\n')
previous=ROOT.parent/'2026-10-22-first-pr-v1/artifacts/fonts.conf'
(ROOT/'artifacts/fonts.conf').write_text(previous.read_text().replace(str(previous.parent/'font-cache'),str(ROOT/'artifacts/font-cache')))
print(ROOT/(TITLE+'.docx'))
