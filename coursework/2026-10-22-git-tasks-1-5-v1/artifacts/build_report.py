from pathlib import Path
import hashlib,json,zipfile
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

R=Path(__file__).resolve().parent.parent; A=R/'artifacts'; S=R/'screenshots'
D=Document();sec=D.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.7)
sec.left_margin=sec.right_margin=Inches(.75)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
 st=D.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
 rf=st.element.get_or_add_rPr().get_or_add_rFonts()
 rf.set(qn('w:eastAsia'),'Hiragino Sans GB')
 for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
  rf.attrib.pop(qn('w:'+key),None)
D.styles['Normal'].font.size=Pt(11)
D.styles['Normal'].paragraph_format.space_after=Pt(6)
D.styles['Normal'].paragraph_format.line_spacing=1.12
D.styles['Title'].font.size=Pt(21)
D.styles['Heading 1'].font.size=Pt(16)
D.styles['Heading 2'].font.size=Pt(12)
md=[]
def h(s,level=1):
 D.add_heading(s,level=level);md.append('#'*(level+1)+' '+s)
def p(s):
 paragraph=D.add_paragraph(s);md.append(s)
 if s.startswith(('Задание ', 'Мобильное приложение Выбор', 'Внутренний сайт Выбор', 'Письменное описание запроса', 'Дополнительный материал и')):
  paragraph.paragraph_format.keep_with_next=True
def pair(zh,ru):p(zh);p(ru)
def code(s):
 x=D.add_paragraph();x.paragraph_format.line_spacing=1.0;x.paragraph_format.space_after=Pt(7);x.paragraph_format.keep_together=True
 rr=x.add_run(s);rr.font.name='Courier New';rr.font.size=Pt(9.5)
 rr._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Hiragino Sans GB')
 md.append('```bash\n'+s+'\n```')
def page():D.add_page_break()

shots=[('01-task1','任务一 首次提交','Задание 1 Первый коммит',.45,.40),('02-task2','任务二 功能分支合并','Задание 2 Слияние ветки',.45,.30),('03-task3-branches','任务三 不同价格分支','Задание 3 Ветки с разной ценой',.45,.40),('04-task3-conflict','任务三 冲突现场','Задание 3 Конфликт до разрешения',.45,.30),('05-task3-resolved','任务三 冲突解决','Задание 3 Разрешённый конфликт',.45,.30),('06-task4-tag','任务四 版本标签','Задание 4 Тег версии',.45,.40),('07-task4-readme','任务四 README 分支','Задание 4 Ветка README',.45,.30),('08-task4-request','任务四 合并请求说明','Задание 4 Описание запроса',.35,.15),('09a-task5-mobile','任务五 手机应用','Задание 5 Мобильное приложение',.35,.25),('09b-task5-website','任务五 内部网站','Задание 5 Внутренний сайт',.35,.25),('10-final','最终仓库状态','Итоговое состояние репозитория',.40,.40)]

def figure(name):
 n,zh,ru,right,bottom=next(x for x in shots if x[0]==name)
 caption=D.add_paragraph(zh+' / '+ru)
 caption.paragraph_format.keep_with_next=True
 for run in caption.runs:run.font.size=Pt(10)
 pic=D.add_picture(str(S/(n+'.png')))
 w,hh=Image.open(S/(n+'.png')).size
 crop=OxmlElement('a:srcRect');crop.set('r',str(int(right*100000)));crop.set('b',str(int(bottom*100000)))
 pic._inline.graphic.graphicData.pic.blipFill.insert(1,crop)
 pic.width=Inches(7);pic.height=Inches(7*(hh*(1-bottom))/(w*(1-right)))
 md.extend(['!['+zh+' / '+ru+']('+str(S/(n+'.png'))+')'])

D.add_paragraph('Git 任务一至五完整实践报告','Title')
D.add_paragraph('Практикум Git Задания 1–5','Subtitle')
md.append('# Git 任务一至五完整实践报告\n\nПрактикум Git Задания 1–5')
p('日期 / Дата 2026-10-22    时区 / Часовой пояс Asia/Shanghai')
pair('已重新初始化练习仓库并完成任务一至五。实际完成分支合并、价格冲突解决、版本标签和 README 分支；任务四和五提供书面答案。',
     'Учебный репозиторий создан заново. Выполнены задания 1–5: слияние веток, разрешение конфликта цены, создание тега и ветки README; подготовлены письменные ответы.')
p('仓库 / Репозиторий\n/Users/zhi/Documents/软件生命周期/git-practice/lab-git')
h('任务一 创建仓库',1)
p('Задание 1 Создание репозитория')
code('git init -b main\ngit config user.name "支浩然"\ngit config user.email "zhr050819@gmail.com"\necho "Проект: учебный" > readme.txt\ngit add readme.txt\ngit commit -m "Добавлен readme"\ngit log --oneline\ngit status')
figure('01-task1')

pair('结果：首次提交 abf3c00，主分支为 main，readme.txt 内容正确，工作区干净。每条命令的真实输出见同目录的终端实际操作记录.md。',
     'Результат: первый коммит abf3c00, основная ветка main, файл readme.txt содержит ожидаемый текст, рабочее дерево чистое. Полный фактический вывод команд сохранён в отдельном журнале.')

pair('材料对应：第一张图为五项要求；第二张图为任务一至二命令；第三张图为任务三至四命令；第四张图补充基础命令、分支和远程仓库概念。',
     'Материалы: изображение 1 — пять заданий; изображение 2 — команды заданий 1–2; изображение 3 — команды заданий 3–4; изображение 4 — основные команды, ветки и удалённые репозитории.')

h('任务二 创建分支并合并');p('Задание 2 Создание и слияние ветки')
code('git branch feature-a\ngit switch feature-a\necho "Функция A" > a.txt\ngit add a.txt\ngit commit -m "Добавлена функция A"\ngit switch main\ngit merge feature-a\ncat a.txt\ngit log --oneline\ngit status')
figure('02-task2')

pair('先创建 feature-a，再切换到该分支，添加并提交 a.txt。返回 main 后，合并前确认主分支尚无 a.txt，再合并 feature-a。',
     'Создана ветка feature-a, выполнен переход в неё, добавлен и закоммичен a.txt. После возврата в main проверено отсутствие a.txt до слияния, затем выполнено слияние feature-a.')
pair('结果：合并采用 Fast-forward，a.txt 已出现在 main，内容为 Функция A。此时历史恰好包含两次提交：abf3c00 和 f5c4235。',
     'Результат: выполнено слияние fast-forward. В main появился a.txt с текстом «Функция A». На этом этапе история содержит ровно два коммита: abf3c00 и f5c4235.')
pair('git branch 只创建分支；git switch 切换分支；git merge 将指定分支合并进当前分支。快进合并不额外创建合并提交。',
     'git branch создаёт ветку, git switch переключает ветку, git merge сливает указанную ветку в текущую. При fast-forward отдельный коммит слияния не создаётся.')


h('任务三 创建并解决价格冲突');p('Задание 3 Создание и разрешение конфликта цены')
code('echo "Цена: 100 рублей" > price.txt\ngit add price.txt\ngit commit -m "Добавлена цена"\ngit branch price-up\ngit branch price-promo\ngit switch price-up\necho "Цена: 120 рублей" > price.txt\ngit commit -am "Повышена цена"\ngit switch price-promo\necho "Цена: 90 рублей" > price.txt\ngit commit -am "Акционная цена"\ngit switch main')
figure('03-task3-branches')
code('git merge price-up\ngit merge price-promo')
pair('两个分支从价格为 100 卢布的同一提交出发。price-up 改为 120 卢布，price-promo 改为 90 卢布。第二次合并实际返回退出码 1，并报告 CONFLICT。',
     'Обе ветки созданы от общего коммита с ценой 100 рублей. В price-up цена изменена на 120, в price-promo — на 90 рублей. Второе слияние завершилось ожидаемым кодом 1 и сообщением CONFLICT.')
figure('04-task3-conflict')

code('<<<<<<< HEAD\nЦена: 120 рублей\n=======\nЦена: 90 рублей\n>>>>>>> price-promo')
pair('最终保留 90 卢布，删除标记和另一价格后执行以下命令。实际操作用 printf 完整写入这一行，效果等同于手工编辑。',
     'Сохранена цена 90 рублей. Маркеры и альтернативная цена удалены. Файл записан командой printf, что эквивалентно ручному редактированию.')
code('printf "%s\\n" "Цена: 90 рублей" > price.txt\ngit add price.txt\ngit commit -m "Разрешён конфликт цены"\ngit log --oneline --graph\ngit status')
figure('05-task3-resolved')

pair('结果：合并提交 8c96675 有两个父提交；无未合并文件，工作区干净。',
     'Результат: коммит слияния 8c96675 имеет двух родителей; неразрешённых конфликтов нет, рабочее дерево чистое.')


h('任务四 标签和合并请求');p('Задание 4 Тег и запрос на слияние')
code('git tag v1.0.0\ngit tag --list\ngit log -1 --oneline --decorate')
figure('06-task4-tag')
code('git switch -c update-readme\nprintf "\\n%s\\n" "Практикум Git: создание веток, слияние и разрешение конфликтов." >> readme.txt\ngit add readme.txt\ngit commit -m "Обновлено описание практикума"\ngit diff main...update-readme')
pair('结果：v1.0.0 指向 8c96675。update-readme 中新增提交 a20991c，差异只涉及 readme.txt。',
     'Результат: тег v1.0.0 указывает на 8c96675. В update-readme создан коммит a20991c; изменён только readme.txt.')
figure('07-task4-readme')

h('合并请求书面说明',2);p('Письменное описание запроса на слияние')
pair('标题：补充 Git 实践说明。源分支 update-readme，目标分支 main。修改目的：说明练习涉及分支创建、合并和冲突解决。',
     'Заголовок: «Дополнение описания практикума Git». Исходная ветка update-readme, целевая main. Цель — описать создание веток, слияние и разрешение конфликтов.')
pair('创建方式：将分支推送到远程仓库，在平台创建 Pull Request 或 Merge Request，选择 update-readme → main，填写标题和说明，并指定同事审查。',
     'Порядок: отправить ветку в удалённый репозиторий, открыть Pull Request или Merge Request из update-readme в main, заполнить заголовок и описание, назначить коллегу на проверку.')
pair('审查内容：文字准确且清楚；原内容保留；只包含预期的 README 修改；没有冲突；git diff 与提交说明一致。审查通过后再合并。',
     'Проверка: точность и понятность текста, сохранность исходного содержания, только ожидаемые изменения README, отсутствие конфликтов, соответствие diff описанию коммита. Слияние — после одобрения.')
pair('已完成本地检查和书面说明。远程 PR 未创建，同事审查未执行。README 分支按任务要求保留，未合并到 main。',
     'Локальная проверка и письменное описание выполнены. Удалённый PR не создан, проверка коллегой не проводилась. Ветка README сохранена и не слита в main.')

figure('08-task4-request')

h('任务五 选择分支策略');p('Задание 5 Выбор стратегии ветвления')
h('手机应用选择 Git Flow',2);p('Мобильное приложение Выбор Git Flow')
pair('条件：每月发布一次，并支持旧版本。选择 Git Flow，因为可以通过 develop 整合新功能，feature 开发独立功能，release 准备每月发布，hotfix 修复已发布版本的紧急问题。对受支持的旧版本保留维护分支，将适用修复同步到对应分支。',
     'Условия: выпуск раз в месяц и поддержка старых версий. Выбран Git Flow: develop служит для интеграции, feature — для отдельных функций, release — для подготовки ежемесячного выпуска, hotfix — для срочных исправлений. Для поддерживаемых старых версий сохраняются ветки сопровождения; необходимые исправления переносятся в соответствующие ветки.')
pair('理由：固定周期发布需要稳定阶段；新功能开发与发布准备可并行；多版本修复有明确的分支位置。',
     'Обоснование: фиксированный выпуск требует этапа стабилизации; разработка и подготовка релиза могут идти параллельно; исправления разных версий организованы в отдельных ветках.')
figure('09a-task5-mobile')

h('内部网站选择 GitHub Flow',2);p('Внутренний сайт Выбор GitHub Flow')
pair('条件：每天更新多次。选择 GitHub Flow：每项小修改从 main 创建短期分支，开发和测试后提交 Pull Request，经审查再合并；main 保持可发布状态，按发布流程频繁部署。',
     'Условия: обновление несколько раз в день. Выбран GitHub Flow: для небольшого изменения создаётся короткоживущая ветка от main; после разработки и тестирования открывается Pull Request, затем выполняются проверка и слияние. Ветка main остаётся готовой к выпуску, развёртывание выполняется часто.')
pair('理由：减少长期分支维护，适合小批量交付和快速反馈。',
     'Обоснование: меньше долгоживущих веток, небольшие поставки и быстрая обратная связь.')

figure('09b-task5-website')

h('补充材料和最终检查');p('Дополнительный материал и итоговая проверка')
pair('git init 初始化；git status 查看状态；git add 暂存；git commit 保存提交；git log --oneline 查看历史。git diff 查看尚未暂存的修改，git diff --cached 查看已暂存的修改。',
     'git init создаёт репозиторий; git status показывает состояние; git add индексирует изменения; git commit сохраняет коммит; git log --oneline показывает историю. git diff показывает неиндексированные изменения, git diff --cached — индексированные.')
pair('git clone 克隆远程仓库；git pull 获取并整合远程更新；git push 发送本地提交。它们是补充材料中的命令，本练习没有配置远程仓库，也未推送。',
     'git clone клонирует удалённый репозиторий, git pull получает и интегрирует обновления, git push отправляет локальные коммиты. Это команды дополнительного материала; удалённый репозиторий не настроен, отправка не выполнялась.')
code('git switch main\ngit status\ncat price.txt\ngit branch\ngit tag --list\ngit log --oneline --graph --decorate --all')
v=json.loads((A/'verification.json').read_text());code(v['commits'])
figure('10-final')

pair('检查结果：main 有六次提交，所有分支合计七次提交；README 分支比 main 多一次提交；v1.0.0 指向 main；最终价格 90 卢布，工作区干净。',
     'Проверка: в main шесть коммитов, всего семь; в ветке README один дополнительный коммит; v1.0.0 указывает на main; итоговая цена 90 рублей, рабочее дерево чистое.')
pair('证据：命令日志记录实际输出和退出码；各任务步骤中穿插十一张 macOS 文本窗口截图。截图展示命令工具产生的真实输出或书面答案，未将文本窗口称为 Terminal。任务三的冲突截图在解决前保存。',
     'Свидетельства: журнал содержит фактический вывод и коды завершения; в разделы заданий включены 11 снимков текстового окна macOS. На них показаны фактические результаты команд или письменные ответы. Снимок конфликта сохранён до его разрешения.')


footer=sec.footer.paragraphs[0];footer.alignment=2
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
for tree in [D.styles.element,D._element]:
 for border in list(tree.iter(qn('w:pBdr'))):border.getparent().remove(border)
for paragraph in D.paragraphs:
 for run in paragraph.runs:
  if any('\u4e00'<=c<='\u9fff' for c in run.text):
   run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Hiragino Sans GB')
doc=R/'Git任务1到5_中俄双语完整报告_v1.docx';D.save(doc)
(R/'Git任务1到5_中俄双语完整报告_v1.md').write_text('\n\n'.join(md)+'\n')
(R/'README.md').write_text('# Git 任务一至五交付包\n\n日期 2026-10-22。已从任务一重新执行。\n\n- Git任务1到5_中俄双语完整报告_v1.docx：可编辑报告，11张截图直接放在对应任务内容中。\n- Git任务1到5_中俄双语完整报告_v1.md：对应Markdown。\n- 终端实际操作记录.md：逐条命令与真实输出、退出码。\n- screenshots/：11张原始截图和对应文本。\n- artifacts/commands.json：机器可读日志。\n- artifacts/verification.json：实际检查结果。\n- artifacts/completed-repository.bundle：新仓库所有分支和标签。\n- artifacts/previous-repository.bundle：删除重做前的仓库备份。\n\n练习仓库：/Users/zhi/Documents/软件生命周期/git-practice/lab-git\n\n截图通过computer use截取macOS文本编辑窗口，内容来自真实命令输出或书面答案。原始PNG未改写，Word仅裁去空白区域显示。任务4按课件完成书面说明，未创建远程PR或实际同事评审。\n')
with zipfile.ZipFile(doc) as z:
 assert len([n for n in z.namelist() if n.startswith('word/media/')])==11
checks={'screenshots':len(list(S.glob('*.png'))),'embedded_images':11,'bilingual':True,'main_checks':v['all_pass']}
(A/'document-structure.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2))
print(doc)
