from pathlib import Path
from zipfile import ZipFile
import hashlib
import json
import shutil

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / 'data'
WORK = ROOT / 'work'
DATA.mkdir(exist_ok=True)
STEM = 'CampusFix_任务二依据证据编写缺陷报告_中俄逐段对照_v1'
SOURCE = Path('/var/folders/ff/kqw394c93dd7krl0z2hcqqdh0000gn/T/codex-clipboard-cccdb912-3a65-44a0-b307-ef30b4e28512.jpg')

cases = [
    {
        'evidence_id': 'I-1', 'bug_id': 'BUG-0201', 'file_label': 'BUG-0201.md',
        'fr': 'FR-CF-01', 'ac': 'AC-CF-01', 'tc': 'TC-CF-01',
        'title': ['Firefox 131 中输入 SSO 密码后无提示返回登录页',
                  'В Firefox 131 после ввода пароля SSO происходит возврат к форме входа без сообщения'],
        'environment': ['CampusFix 1.0.3 测试服务器；Windows 11；Firefox 131；student09。',
                        'Тестовый сервер CampusFix 1.0.3; Windows 11; Firefox 131; student09.'],
        'evidence': ['沿用前图报告人的答复：11 月 3 日 09:15，student09 在 SSO 输入密码后返回登录页，无错误提示；Firefox 中每次出现，Chrome 130 可登录。年份、时区及尝试次数未给出。',
                     'Из предыдущего ответа репортёра: 3 ноября в 09:15 student09 после ввода пароля SSO возвращается к форме без сообщения. В Firefox это происходит всегда, в Chrome 130 вход работает. Год, часовой пояс и число попыток не указаны.'],
        'preconditions': ['准备测试账号 student09；可访问 CampusFix 和 SSO 登录入口，初始处于未登录状态。',
                          'Подготовлена тестовая учётная запись student09; доступны CampusFix и вход SSO; сеанс ещё не открыт.'],
        'steps': [
            ['在 Windows 11 的 Firefox 131 中打开 CampusFix 测试服务器。',
             'Открыть тестовый сервер CampusFix в Firefox 131 на Windows 11.'],
            ['进入 SSO，使用 student09 输入密码并执行登录。',
             'Перейти в SSO, ввести пароль для student09 и выполнить вход.'],
            ['记录返回页面与提示信息；在 Chrome 130 中对照执行同一流程。',
             'Зафиксировать конечную страницу и сообщения; сравнить тот же сценарий в Chrome 130.'],
        ],
        'actual': ['返回登录表单且没有提示，未显示“我的申请”。',
                   'Открывается форма входа без сообщения; список «Мои заявки» не показан.'],
        'expected': ['有效账号通过 SSO 登录后进入“我的申请”列表。',
                     'После входа через SSO с действительной учётной записью открывается список «Мои заявки».'],
        'reproduction_short': ['报告人称每次；独立复核 NOT_RUN。',
                              'Со слов автора всегда; независимая проверка NOT_RUN.'],
        'reproduction': ['报告人称 Firefox 中每次发生，Chrome 130 正常；没有次数分母。独立复现与浏览器对照测试均为 NOT_RUN。',
                        'По словам автора в Firefox сбой возникает всегда, в Chrome 130 вход работает; число попыток неизвестно. Независимое воспроизведение и сравнение браузеров имеют статус NOT_RUN.'],
        'followup': ['补充服务器 URL、SSO 重定向链、年份与时区，并记录两种浏览器的尝试与失败次数。',
                     'Уточнить URL сервера, цепочку перенаправлений SSO, год и часовой пояс; записать число попыток и отказов в обоих браузерах.'],
        'requirement': ['学生能够使用有效校内账号通过 SSO 登录并进入自己的申请列表。',
                        'Студент может войти через SSO с действительной учётной записью вуза и открыть свои заявки.'],
        'criterion': ['Windows 11 下 Firefox 131 与 Chrome 130 的有效账号登录均进入“我的申请”，不返回登录表单。',
                      'В Windows 11 вход с действительными данными в Firefox 131 и Chrome 130 открывает «Мои заявки» без возврата к форме входа.'],
        'test_definition': ['使用 student09 分别在两个浏览器执行登录；核对会话及最终页面。状态 NOT_RUN。',
                            'Войти как student09 в каждом браузере; проверить сеанс и конечную страницу. Статус NOT_RUN.'],
        'evidence_type': 'PRIOR_CLASSROOM_REPORTER_STATEMENT',
    },
    {
        'evidence_id': 'I-2', 'bug_id': 'BUG-0202', 'file_label': 'BUG-0202.md',
        'fr': 'FR-CF-02', 'ac': 'AC-CF-02', 'tc': 'TC-CF-02',
        'title': ['描述仅含三个空格时仍成功创建申请',
                  'Заявка создаётся, когда описание содержит только три пробела'],
        'environment': ['自拟环境：CampusFix 1.0.3 测试服务器；Windows 11；Chrome 130；student09。',
                        'Учебная среда: тестовый сервер CampusFix 1.0.3; Windows 11; Chrome 130; student09.'],
        'evidence': ['模拟记录 I-2：description="   " 的 POST /api/requests 返回 201，生成 CF-4201；随后查询显示 description=""。模拟三轮，每轮重置数据，三轮均创建空描述申请。',
                     'Учебная запись I-2: POST /api/requests с description="   " возвращает 201 и создаёт CF-4201; последующее чтение показывает description="". В трёх модельных повторах с восстановлением данных создана заявка с пустым описанием.'],
        'preconditions': ['student09 已登录；其他必填字段有效；每轮使用独立测试数据。',
                          'student09 вошёл; остальные обязательные поля корректны; для каждого повтора подготовлены отдельные тестовые данные.'],
        'steps': [
            ['打开新建申请表，填写有效的分类与位置等必填项。',
             'Открыть форму новой заявки и заполнить категорию, место и другие обязательные поля.'],
            ['描述仅输入三个空格，然后提交。',
             'Ввести в описание ровно три пробела и отправить форму.'],
            ['记录响应码，并查询新申请是否保存及其描述内容。',
             'Зафиксировать код ответа; проверить сохранение новой заявки и её описание.'],
        ],
        'actual': ['模拟响应为 201，申请创建成功，保存后的描述为空。',
                   'В учебном примере получен ответ 201; заявка создана с пустым описанием.'],
        'expected': ['拒绝该提交，返回 422 和明确的字段提示，且不新增申请。',
                     'Отклонить отправку с кодом 422 и понятным сообщением поля; новая заявка не создаётся.'],
        'reproduction_short': ['模拟 3/3；实际执行 NOT_RUN。',
                              'Модель 3/3; реальный запуск NOT_RUN.'],
        'reproduction': ['教学模拟结果为 3/3，仅代表自拟证据设定；真实环境执行状态为 NOT_RUN。',
                        'Учебный результат 3/3 является условием модельного примера; запуск в реальной среде имеет статус NOT_RUN.'],
        'followup': ['实际复测时保存请求、响应与创建前后数量；补做空字符串、换行和正常文本对照。',
                     'При реальной проверке сохранить запрос, ответ и число записей до и после; сравнить пустую строку, перевод строки и обычный текст.'],
        'requirement': ['申请描述为必填项，去除首尾空白后长度为 1～1000 个字符。',
                        'Описание обязательно; после удаления пробелов по краям длина составляет от 1 до 1000 символов.'],
        'criterion': ['空字符串或纯空白描述得到 422 与字段提示，申请总数不变；有效描述可提交。',
                      'Пустое описание или одни пробельные символы дают 422 и сообщение поля без новой записи; корректное описание принимается.'],
        'test_definition': ['提交三个空格并核对 422、提示与零新增记录，再用有效文本作对照。状态 NOT_RUN。',
                            'Отправить три пробела, проверить 422, сообщение и отсутствие новой записи; сравнить с обычным текстом. Статус NOT_RUN.'],
        'evidence_type': 'USER_AUTHORIZED_SYNTHETIC_EVIDENCE',
    },
    {
        'evidence_id': 'I-3', 'bug_id': 'BUG-0203', 'file_label': 'BUG-0203.md',
        'fr': 'FR-CF-03', 'ac': 'AC-CF-03', 'tc': 'TC-CF-03',
        'title': ['student09 修改申请编号后可读取 student10 的申请详情',
                  'student09 читает заявку student10 после замены идентификатора'],
        'environment': ['自拟环境：CampusFix 1.0.3 测试服务器；Chrome 130；两个普通学生测试账号。',
                        'Учебная среда: тестовый сервер CampusFix 1.0.3; Chrome 130; две обычные студенческие учётные записи.'],
        'evidence': ['模拟记录 I-3：student09 的会话读取自己的 CF-4101 得到 200；将编号换为归属 student10 的 CF-4102 后仍得到 200，响应包含 owner=student10 和申请内容。模拟重复两次，均可读取。',
                     'Учебная запись I-3: в сеансе student09 своя CF-4101 читается с кодом 200; после замены ID на CF-4102, принадлежащую student10, ответ также 200 и содержит owner=student10 и текст заявки. В двух модельных повторах доступ сохраняется.'],
        'preconditions': ['student09 与 student10 均为普通学生；两人各有一条申请；当前仅登录 student09。',
                          'student09 и student10 — обычные студенты; у каждого есть своя заявка; открыт только сеанс student09.'],
        'steps': [
            ['用 student09 登录，并读取自己的 CF-4101，作为正常对照。',
             'Войти как student09 и прочитать свою CF-4101 для контрольного сравнения.'],
            ['保持同一会话，将详情请求中的 ID 换为 CF-4102。',
             'В том же сеансе заменить ID запроса подробностей на CF-4102.'],
            ['记录响应码及响应体，核对申请实际归属。',
             'Зафиксировать код и тело ответа; проверить владельца заявки.'],
        ],
        'actual': ['模拟中返回 200，并暴露另一个学生的申请详情。',
                   'В учебном примере возвращаются 200 и подробности заявки другого студента.'],
        'expected': ['按本练习接口约定返回 403，不返回他人申请详情；读取自己的申请仍为 200。',
                     'По контракту упражнения вернуть 403 без чужих подробностей; собственная заявка по-прежнему читается с кодом 200.'],
        'reproduction_short': ['模拟 2/2；实际执行 NOT_RUN。',
                              'Модель 2/2; реальный запуск NOT_RUN.'],
        'reproduction': ['教学模拟结果为 2/2；真实环境的对象访问权限测试未执行，状态 NOT_RUN。',
                        'Учебный результат — 2/2; проверка прав на объекты в реальной среде не выполнена, статус NOT_RUN.'],
        'followup': ['实际验证中仅使用授权测试账号与合成申请；同时检查详情与相关附件接口的所有者限制。',
                     'Для реальной проверки использовать разрешённые тестовые аккаунты и учебные заявки; проверить владельца для подробностей и вложений.'],
        'requirement': ['普通学生只可读取自己创建的申请及其详情，服务端按会话与所有者判定权限。',
                        'Обычный студент читает только собственные заявки; сервер проверяет сеанс и владельца объекта.'],
        'criterion': ['student09 读取 student10 的 CF-4102 得到 403 且无详情；读取自己的 CF-4101 得到 200。',
                      'student09 получает 403 без подробностей для CF-4102 студента student10 и 200 для собственной CF-4101.'],
        'test_definition': ['用 student09 同一会话分别读取本人和他人申请，比较返回码及响应体。状态 NOT_RUN。',
                            'В одном сеансе student09 запросить свою и чужую заявки, сравнить коды и тела ответов. Статус NOT_RUN.'],
        'evidence_type': 'USER_AUTHORIZED_SYNTHETIC_EVIDENCE',
    },
    {
        'evidence_id': 'I-4', 'bug_id': 'BUG-0204', 'file_label': 'BUG-0204.md',
        'fr': 'FR-CF-04', 'ac': 'AC-CF-04', 'tc': 'TC-CF-04',
        'title': ['在列表第二页选择“已关闭”后仍显示“待处理”申请',
                  'После выбора «Закрыта» на второй странице остаются открытые заявки'],
        'environment': ['自拟环境：CampusFix 1.0.3 测试服务器；Windows 11；Chrome 130；student09。',
                        'Учебная среда: тестовый сервер CampusFix 1.0.3; Windows 11; Chrome 130; student09.'],
        'evidence': ['模拟记录 I-4：第二页选择“已关闭”后，下拉框已更新，但列表仍有 OPEN 记录 CF-4301、CF-4302；请求仍为 GET /api/requests?page=2，缺少 status。首页选择同一筛选时，请求为 status=CLOSED&page=1，显示正确。',
                     'Учебная запись I-4: на странице 2 выбран статус «Закрыта», но видны OPEN CF-4301 и CF-4302; запрос остаётся GET /api/requests?page=2 без status. На странице 1 тот же фильтр отправляет status=CLOSED&page=1 и показывает верные записи.'],
        'preconditions': ['准备 student09 的 12 条申请：9 条 OPEN、3 条 CLOSED，每页 10 条；默认排序下第二页两条均为 OPEN。',
                          'Подготовить 12 заявок student09: 9 OPEN и 3 CLOSED, по 10 на странице; при исходной сортировке обе записи страницы 2 имеют статус OPEN.'],
        'steps': [
            ['登录后在“我的申请”中选择全部状态，并进入第二页。',
             'После входа открыть «Мои заявки», выбрать все статусы и перейти на страницу 2.'],
            ['把状态改为“已关闭”，观察当前页、请求参数和结果状态。',
             'Выбрать «Закрыта»; проверить номер страницы, параметры запроса и статусы результатов.'],
            ['回到未筛选的首页，再选择“已关闭”进行对照。',
             'Вернуться на исходную первую страницу без фильтра и сравнить выбор «Закрыта».'],
        ],
        'actual': ['模拟中第二页未切回第一页，筛选请求缺少 status，列表仍含 OPEN 记录。',
                   'В учебном примере остаётся страница 2, запрос не содержит status и список включает записи OPEN.'],
        'expected': ['改变状态后返回第一页，发送 status=CLOSED&page=1，三条结果均为 CLOSED。',
                     'После смены статуса перейти на страницу 1, отправить status=CLOSED&page=1 и показать три записи CLOSED.'],
        'reproduction_short': ['模拟 3/3；实际执行 NOT_RUN。',
                              'Модель 3/3; реальный запуск NOT_RUN.'],
        'reproduction': ['在第二页切换筛选的教学模拟为 3/3；首页对照正常。真实环境执行状态为 NOT_RUN。',
                        'Учебная смена фильтра на странице 2 даёт 3/3; контроль на странице 1 корректен. Реальный запуск имеет статус NOT_RUN.'],
        'followup': ['实际复测时保存筛选前后录屏与请求参数，并覆盖从其他非首页切换不同状态。',
                     'При реальной проверке сохранить запись смены фильтра и параметры запросов; проверить другие статусы с остальных страниц.'],
        'requirement': ['学生可按状态筛选自己的申请；改变筛选条件时页码重置为 1，并查询相应状态。',
                        'Студент фильтрует свои заявки по статусу; смена фильтра сбрасывает страницу на 1 и запрашивает нужный статус.'],
        'criterion': ['从第二页选择 CLOSED 后请求包含 status=CLOSED&page=1；结果恰有三条且全部 CLOSED。',
                      'При выборе CLOSED со страницы 2 запрос содержит status=CLOSED&page=1; получены ровно три записи, все CLOSED.'],
        'test_definition': ['使用 12 条固定数据，从第二页筛选 CLOSED，与第一页对照并核对参数和结果。状态 NOT_RUN。',
                            'На 12 заданных записях выбрать CLOSED со страницы 2, сравнить со страницей 1 и проверить запрос и результат. Статус NOT_RUN.'],
        'evidence_type': 'USER_AUTHORIZED_SYNTHETIC_EVIDENCE',
    },
]

doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Inches(11), Inches(8.5)
sec.top_margin = sec.bottom_margin = Inches(0.45)
sec.left_margin = sec.right_margin = Inches(0.55)
for name in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
    s = doc.styles[name]
    s.font.name = 'Arial'
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Arial Unicode MS')
    s.paragraph_format.space_after = Pt(3)
    s.paragraph_format.line_spacing = 1.0
doc.styles['Normal'].font.size = Pt(10)
doc.styles['Title'].font.size = Pt(15)
doc.styles['Heading 1'].font.size = Pt(10.5)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(4)

pairs = []
bookmarks = set()
links = []


def pair(zh, ru, cell=None, style=None, size=10, bold=False):
    p = doc.add_paragraph(style=style) if cell is None else (cell.paragraphs[0] if not cell.text else cell.add_paragraph())
    if style:
        p.style = style
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.keep_together = True
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(zh)
    r.bold, r.font.size = bold, Pt(size)
    q = doc.add_paragraph(style=style) if cell is None else cell.add_paragraph()
    if style:
        q.style = style
    q.paragraph_format.keep_with_next = style in ('Title', 'Heading 1', 'Heading 2')
    q.paragraph_format.keep_together = True
    q.paragraph_format.space_after = Pt(4)
    r = q.add_run(ru)
    r.bold, r.font.size = bold, Pt(size)
    pairs.append((p, q, zh, ru))
    return p, q


def bookmark(p, ident):
    name = 'def_' + ident.replace('-', '_')
    assert name not in bookmarks
    start, end = OxmlElement('w:bookmarkStart'), OxmlElement('w:bookmarkEnd')
    idx = str(len(bookmarks) + 1)
    start.set(qn('w:id'), idx)
    start.set(qn('w:name'), name)
    end.set(qn('w:id'), idx)
    p._p.insert(0, start)
    p._p.append(end)
    bookmarks.add(name)


def link(p, ident):
    anchor = 'def_' + ident.replace('-', '_')
    h = OxmlElement('w:hyperlink')
    h.set(qn('w:anchor'), anchor)
    r, rp, t = OxmlElement('w:r'), OxmlElement('w:rPr'), OxmlElement('w:t')
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '1F4E79')
    size = OxmlElement('w:sz')
    size.set(qn('w:val'), '19')
    rp.extend([color, size])
    r.append(rp)
    t.text = ident
    r.append(t)
    h.append(r)
    p._p.append(h)
    links.append(anchor)


def new_table(headers, widths):
    t = doc.add_table(rows=1, cols=len(widths))
    t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    for cell, (zh, ru) in zip(t.rows[0].cells, headers):
        pair(zh, ru, cell=cell, bold=True, size=9.5)
        sh = OxmlElement('w:shd')
        sh.set(qn('w:fill'), 'DAEEF0')
        cell._tc.get_or_add_tcPr().append(sh)
    rep = OxmlElement('w:tblHeader')
    t.rows[0]._tr.get_or_add_trPr().append(rep)
    props = t._tbl.tblPr
    margins = OxmlElement('w:tblCellMar')
    for side in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement('w:' + side)
        e.set(qn('w:w'), '70')
        e.set(qn('w:type'), 'dxa')
        margins.append(e)
    props.append(margins)
    borders = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + side)
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), '4')
        e.set(qn('w:color'), 'D9D9D9')
        borders.append(e)
    props.append(borders)
    return t


pair('CampusFix 任务二 依据证据编写缺陷报告', 'CampusFix Задание 2 Отчёты по уликам', style='Title', size=15)
pair('中俄逐部分对照  |  教学示例  |  v1  |  DRAFT', 'Китайский текст над русским  |  Учебные примеры  |  v1  |  DRAFT', size=9)
pair('I-1 沿用前图登录现象，I-2～I-4 为自拟证据；FR、AC、TC 为本练习定义。模拟次数不代表真实执行，所有实际测试为 NOT_RUN。',
     'I-1 использует предыдущий ответ о входе, I-2–I-4 — учебные данные. FR, AC и TC определены для упражнения. Модельные повторы не являются реальным тестированием; фактический статус — NOT_RUN.', size=9.5)
pair('汇总表保留题目文件名，四份报告正文收录于本 Word；关联编号可跳转到文末定义。',
     'В сводке сохранены имена файлов задания; все четыре отчёта включены в Word. Ссылки ведут к определениям в конце.', size=9.5)
widths = [0.55, 1.17, 3.45, 1.58, 3.15]
table = new_table([
    ('证据', 'Улика'), ('题目文件名', 'Файл'),
    ('标题  症状与条件', 'Заголовок  симптом и условие'),
    ('关联 FR AC TC', 'Связи FR AC TC'), ('复现情况', 'Воспроизводимость'),
], widths)
for case in cases:
    cells = table.add_row().cells
    for cell, width in zip(cells, widths):
        cell.width = Inches(width)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for cell, txt in zip(cells[:2], (case['evidence_id'], case['file_label'])):
        cell.paragraphs[0].add_run(txt).font.size = Pt(9)
    pair(*case['title'], cell=cells[2], size=9.5)
    for n, field in enumerate(('fr', 'ac', 'tc')):
        p = cells[3].paragraphs[0] if n == 0 else cells[3].add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        link(p, case[field])
    pair(*case['reproduction_short'], cell=cells[4], size=9.5)
    cells[0]._tc.getparent().get_or_add_trPr().append(OxmlElement('w:cantSplit'))

for case in cases:
    doc.add_page_break()
    pair(case['bug_id'] + ' 缺陷报告', case['bug_id'] + ' Отчёт о дефекте', style='Title', size=15)
    pair(*case['title'], bold=True, size=11)
    p = doc.add_paragraph(f"{case['evidence_id']}  |  {case['file_label']}  |  ")
    p.paragraph_format.space_after = Pt(4)
    for n, field in enumerate(('fr', 'ac', 'tc')):
        if n:
            p.add_run('  /  ')
        link(p, case[field])
    pair('环境：' + case['environment'][0], 'Окружение: ' + case['environment'][1], size=9.5)
    pair('证据：' + case['evidence'][0], 'Улика: ' + case['evidence'][1], size=9.5)
    pair('前提：' + case['preconditions'][0], 'Предусловия: ' + case['preconditions'][1], size=9.5)
    pair('复现步骤', 'Шаги воспроизведения', style='Heading 1', bold=True, size=10)
    for n, (zh, ru) in enumerate(case['steps'], 1):
        pair(f'{n}. {zh}', f'{n}. {ru}', size=9.5)
    pair('实际结果：' + case['actual'][0], 'Фактический результат: ' + case['actual'][1], size=9.5)
    pair('预期结果：' + case['expected'][0], 'Ожидаемый результат: ' + case['expected'][1], size=9.5)
    pair('复现情况：' + case['reproduction'][0], 'Повторяемость: ' + case['reproduction'][1], size=9)
    pair('补充验证：' + case['followup'][0], 'Дополнительная проверка: ' + case['followup'][1], size=9)

doc.add_page_break()
pair('本练习的 FR AC TC 定义', 'Определения FR AC TC для упражнения', style='Title', size=15)
pair('以下为自拟教学定义，状态为 PROPOSED；这些编号只用于本 CampusFix 练习。',
     'Ниже приведены учебные определения со статусом PROPOSED; идентификаторы относятся только к этому упражнению CampusFix.', size=9.5)
for case in cases:
    p = doc.add_paragraph(case['bug_id'] + '  |  ' + case['evidence_id'])
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(3)
    p.runs[0].bold = True
    p.runs[0].font.size = Pt(10)
    for field, content_key in [('fr', 'requirement'), ('ac', 'criterion'), ('tc', 'test_definition')]:
        ident = case[field]
        p, _ = pair(ident + '  ' + case[content_key][0], ident + '  ' + case[content_key][1], size=9.5)
        bookmark(p, ident)

for element in (doc.styles.element, doc.element):
    for b in list(element.iter(qn('w:pBdr'))):
        b.getparent().remove(b)
doc.core_properties.title = 'CampusFix 任务二 依据证据编写缺陷报告'
doc.core_properties.subject = 'I-1 至 I-4 四份报告 中俄逐段对照 教学示例'
doc.core_properties.author = 'StudySpace Coursework'
out = ROOT / (STEM + '.docx')
doc.save(out)
shutil.copy2(SOURCE, DATA / 'source-task2.jpg')
previous = ROOT.parent / '2026-10-19-campusfix-br1-v1/data/source.jpg'
shutil.copy2(previous, DATA / 'source-br1.jpg')

content = {
    'version': 'v1', 'date': '2026-10-19', 'status': 'DRAFT',
    'authorization': '用 CampusFix 自拟四组证据示例完成',
    'case': 'CampusFix', 'deliverables': ['docx'], 'language_order': ['zh', 'ru'],
    'cases': cases, 'actual_business_test_status': 'NOT_RUN',
}
(DATA / 'content.json').write_text(json.dumps(content, ensure_ascii=False, indent=2) + '\n')
with ZipFile(out) as z:
    assert z.testzip() is None
assert len(table.columns) == 5 and len(table.rows) == 5
assert {c['evidence_id'] for c in cases} == {'I-1', 'I-2', 'I-3', 'I-4'}
assert {c['file_label'] for c in cases} == {f'BUG-020{i}.md' for i in range(1, 5)}
assert len(bookmarks) == 12 and len(links) == 24
assert all(anchor in bookmarks for anchor in links)
for p, q, zh, ru in pairs:
    assert p.text == zh and q.text == ru
    assert p._p.getnext() is q._p
assert not list(ROOT.glob('*.md'))
coverage = {
    'source_title': 'Задание 2: отчёты по уликам',
    'columns': ['Улика', 'Файл', 'Заголовок (симптом + условие)', 'Связи (FR, AC, TC)', 'Воспроизводимость'],
    'rows': [{'evidence': c['evidence_id'], 'file_label': c['file_label'], 'report': c['bug_id']} for c in cases],
    'authorization': content['authorization'],
    'actual_evidence_body_in_task_image': False,
    'source_task2_sha256': hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
    'source_br1_sha256': hashlib.sha256(previous.read_bytes()).hexdigest(),
    'classroom_timer_40_00': 'Time allowance only',
    'markdown_output': False,
}
(DATA / 'source-coverage.json').write_text(json.dumps(coverage, ensure_ascii=False, indent=2) + '\n')
check = {
    'task_table': 'PASS', 'columns': 5, 'filled_evidence_rows': 4,
    'full_reports': 4, 'defined_traceability_ids': 12, 'internal_links': 24,
    'all_link_targets_exist': 'PASS', 'bilingual_adjacent_pairs': len(pairs),
    'language_order': 'ZH_THEN_RU', 'docx_zip': 'PASS', 'word_only': True,
    'actual_business_tests': 'NOT_RUN', 'render_review': 'PENDING',
    'sha256': hashlib.sha256(out.read_bytes()).hexdigest(),
}
(DATA / 'verification.json').write_text(json.dumps(check, ensure_ascii=False, indent=2) + '\n')
fonts = ROOT.parent / '2026-10-19-bug-report-review-v1/work/fonts'
lo_fonts = '/Users/zhi/.cache/codex-runtimes/codex-primary-runtime/dependencies/native/libreoffice-headless/libreoffice/LibreOfficeDev.app/Contents/Resources/fonts/truetype'
(WORK / 'fonts.conf').write_text(f'<?xml version="1.0"?><!DOCTYPE fontconfig SYSTEM "fonts.dtd"><fontconfig><dir>{fonts}</dir><dir>{lo_fonts}</dir><cachedir>{WORK / "fontcache"}</cachedir></fontconfig>')
print(f'PASS: 4 evidence rows, 4 reports, 12 traceability definitions, 24 internal links, {len(pairs)} bilingual pairs')
print(out)
