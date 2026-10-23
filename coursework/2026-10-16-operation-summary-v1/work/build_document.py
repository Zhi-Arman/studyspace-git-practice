from __future__ import annotations

import hashlib
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
PROJECT = ROOT.parents[1]
CONTENT = ROOT / "data/content.json"
OUTPUT = ROOT / "StudySpace_任务二操作汇总_中俄双语_v1.docx"
VERIFY = ROOT / "data/verification.json"
data = json.loads(CONTENT.read_text(encoding="utf-8"))


def font(run, *, size=9.1, bold=False, color=None):
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key, val in (("ascii", "Arial"), ("hAnsi", "Arial"), ("eastAsia", "Heiti SC")):
        rfonts.set(qn("w:" + key), val)


def pair(container, value, *, size=9.1, bold=False, use_first=False):
    for index, lang in enumerate(("zh", "ru")):
        p = container.paragraphs[0] if use_first and index == 0 else container.add_paragraph()
        p.paragraph_format.space_after = Pt(3 if lang == "zh" else 0)
        p.paragraph_format.line_spacing = 1.08
        font(p.add_run(value[lang]), size=size, bold=bold)


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def cell_padding(cell, top=100, bottom=100, left=105, right=105):
    tcpr = cell._tc.get_or_add_tcPr()
    mar = tcpr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tcpr.append(mar)
    for side, value in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        node = mar.find(qn("w:" + side))
        if node is None:
            node = OxmlElement("w:" + side)
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def prevent_split(row, *, repeat=False):
    trpr = row._tr.get_or_add_trPr()
    trpr.append(OxmlElement("w:cantSplit"))
    if repeat:
        flag = OxmlElement("w:tblHeader")
        flag.set(qn("w:val"), "true")
        trpr.append(flag)


doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width = Inches(16.54)
sec.page_height = Inches(11.69)
sec.left_margin = sec.right_margin = Inches(0.50)
sec.top_margin = Inches(0.50)
sec.bottom_margin = Inches(0.46)
sec.header_distance = sec.footer_distance = Inches(0.22)

normal = doc.styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(9)
normal.paragraph_format.space_after = Pt(3)
normal.paragraph_format.line_spacing = 1.08

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(1)
font(p.add_run(data["title"]["zh"]), size=17, bold=True, color=(18, 67, 83))
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(2)
font(p.add_run(data["title"]["ru"]), size=14, bold=True, color=(18, 67, 83))
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(7)
font(p.add_run(f'{data["document_id"]}  ·  {data["version"]}  ·  {data["date"]}  ·  {data["status"]}'), size=8.5, color=(85, 95, 104))

for lang in ("zh", "ru"):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2 if lang == "zh" else 7)
    font(p.add_run(data["intro"][lang]), size=9.2)

widths = [2.39, 2.13, 2.50, 2.25, 3.18, 3.09]
assert round(sum(widths), 2) == 15.54
table = doc.add_table(rows=1, cols=6)
table.autofit = False
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = "Table Grid"
for index, width in enumerate(widths):
    table.columns[index].width = Inches(width)

for index, header in enumerate(data["headers"]):
    cell = table.rows[0].cells[index]
    cell.width = Inches(widths[index])
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    shade(cell, "DCEEF0")
    cell_padding(cell, top=90, bottom=90)
    pair(cell, header, size=9.3, bold=True, use_first=True)
prevent_split(table.rows[0], repeat=True)

keys = ("operation", "parameters", "body", "success", "errors", "security")
for row_index, item in enumerate(data["rows"]):
    row = table.add_row()
    prevent_split(row)
    for col_index, key in enumerate(keys):
        cell = row.cells[col_index]
        cell.width = Inches(widths[col_index])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        cell_padding(cell)
        if row_index % 2:
            shade(cell, "F4F8F9")
        if key == "operation":
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(3)
            font(p.add_run(item["id"]), size=9, bold=True, color=(18, 67, 83))
            pair(cell, item[key], size=9.2, bold=True)
            p = cell.add_paragraph()
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.space_after = Pt(0)
            font(p.add_run(item["trace"]), size=7.7, color=(85, 95, 104))
        else:
            pair(cell, item[key], size=9.0, use_first=True)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(8)
p.paragraph_format.space_after = Pt(2)
font(p.add_run("说明 / Примечания"), size=10, bold=True, color=(18, 67, 83))
for note in data["notes"]:
    for lang in ("zh", "ru"):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        font(p.add_run(note[lang]), size=8.4)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(5)
p.paragraph_format.space_after = Pt(2)
font(p.add_run("依据 / Источники"), size=9.5, bold=True, color=(18, 67, 83))
for source in data["sources"]:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    font(p.add_run(source), size=7.6, color=(85, 95, 104))

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)

check = Document(OUTPUT)
assert len(check.tables) == 1
assert len(check.tables[0].rows) == 1 + len(data["rows"])
assert all(len(row.cells) == 6 for row in check.tables[0].rows)
all_text = "\n".join(p.text for p in check.paragraphs)
all_text += "\n" + "\n".join(c.text for row in check.tables[0].rows for c in row.cells)
required = list(data["title"].values()) + list(data["intro"].values())
for item in data["headers"] + data["notes"]:
    required.extend(item.values())
for item in data["rows"]:
    required.extend((item["id"], item["trace"]))
    for key in keys:
        required.extend(item[key].values())
required.extend(data["sources"])
missing = [value for value in required if value not in all_text]
assert not missing, missing
assert all((PROJECT / source).exists() for source in data["sources"])
result = {
    "document": str(OUTPUT),
    "sha256": hashlib.sha256(OUTPUT.read_bytes()).hexdigest(),
    "status": data["status"],
    "operation_count": len(data["rows"]),
    "table_dimensions": [len(check.tables[0].rows), len(check.tables[0].columns)],
    "paired_text_checks": len(required),
    "structural_result": "PASS",
    "business_tests": "NOT_RUN",
}
VERIFY.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps(result, ensure_ascii=False, indent=2))
