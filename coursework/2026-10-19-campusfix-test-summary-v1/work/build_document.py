from collections import Counter
from pathlib import Path
from zipfile import ZipFile
import hashlib
import json
import shutil

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
DATA, WORK = ROOT / 'data', ROOT / 'work'
DATA.mkdir(exist_ok=True)
STEM = 'CampusFix_任务五测试报告汇总_中俄逐段对照_v1'
SOURCE = Path('/var/folders/ff/kqw394c93dd7krl0z2hcqqdh0000gn/T/codex-clipboard-90535a1c-5bde-4b23-b4e0-978ecdd0a471.jpg')
PARENT = ROOT.parent
input_paths = {
    'triage': PARENT / '2026-10-19-campusfix-triage-v1/data/content.json',
    'fix_record': PARENT / '2026-10-19-campusfix-fix-confirmation-v1/data/content.json',
    'execution': PARENT / '2026-10-19-campusfix-fix-confirmation-v1/data/execution.json',
    'evidence_cases': PARENT / '2026-10-19-campusfix-evidence-reports-v1/data/content.json',
}
input_hashes = {key: hashlib.sha256(path.read_bytes()).hexdigest() for key, path in input_paths.items()}
inputs = {key: json.loads(path.read_text()) for key, path in input_paths.items()}
triage, fix, execution = inputs['triage'], inputs['fix_record'], inputs['execution']
case_lookup = {case['bug_id']: case for case in inputs['evidence_cases']['cases']}
main_defects = [x for x in triage['rows'] if x['decision'] == 'DEFECT']
duplicates = [x for x in triage['rows'] if x['decision'] == 'DUPLICATE']
rejected = [x for x in triage['rows'] if x['decision'] == 'NOT_A_DEFECT']
demo_id = fix['bug_id']
assert execution['demonstrator_confirmation'] == 'PASS'
assert execution['original_system_confirmation'] == 'NOT_RUN'
for run in execution['runs']:
    assert hashlib.sha256(Path(run['log']).read_bytes()).hexdigest() == run['log_sha256']
assert hashlib.sha256((PARENT / '2026-10-19-campusfix-fix-confirmation-v1/demo/validator.py').read_bytes()).hexdigest() == execution['after_source_sha256']

severity_map = {'S1': 'Critical', 'S2': 'Major', 'S3': 'Minor', 'S4': 'Trivial'}
case_severity = Counter(severity_map[x['severity']] for x in main_defects)
combined_severity = case_severity + Counter({'Minor': 1})
metrics = {
    'scope': 'COURSEWORK_AGGREGATE_WITH_SEPARATE_TEST_OBJECTS',
    'case_registry': {'found': len(main_defects), 'verified_closed': 0, 'open': len(main_defects), 'severity': dict(case_severity)},
    'independent_demonstrator': {'found': 1, 'verified_closed': 1, 'open': 0, 'bug': demo_id, 'severity': 'Minor', 'classification_added_in_task5': True},
    'coursework_combined': {'found': len(main_defects) + 1, 'verified_closed': 1, 'open': len(main_defects), 'severity': dict(combined_severity)},
    'duplicates': len(duplicates), 'non_defects_rejected': len(rejected),
    'severity_label_mapping': severity_map,
    'original_system_tests': 'NOT_RUN',
    'recommendation': 'DEFER_ACCEPTANCE',
    'official_test_plan_3_4_1_provided': False,
    'exit_criteria_status': 'PROPOSED_CLASSROOM_MAPPING_PENDING_OFFICIAL_PLAN',
    'source_hashes': input_hashes,
}
assert metrics['coursework_combined']['found'] == 6
assert metrics['coursework_combined']['verified_closed'] + metrics['coursework_combined']['open'] == 6
assert sum(combined_severity.values()) == 6
assert [combined_severity[x] for x in ('Critical', 'Major', 'Minor', 'Trivial')] == [1, 1, 3, 1]
assert demo_id not in {x['bug'] for x in main_defects}

summary_rows = [
    [
        ['发现的缺陷 不含重复项和非缺陷', 'Найдено дефектов без дубликатов и не дефектов'],
        ['6 项：任务三登记案例 5 项＋独立教学示例 1 项。',
         '6: 5 дефектов реестра задания 3 и 1 в отдельном учебном примере.'],
    ],
    [
        ['严重程度 Critical / Major / Minor / Trivial', 'По серьёзности Critical / Major / Minor / Trivial'],
        ['1 / 1 / 3 / 1；包含 BUG-0206 本次补定为 Minor。',
         '1 / 1 / 3 / 1; для BUG-0206 в этой сводке задан Minor.'],
    ],
    [
        ['已关闭并复核', 'Закрыто и проверено'],
        ['1 项：仅独立示例 BUG-0206，确认记录 CONF-CF-01；CampusFix 案例登记中为 0 项。',
         '1: только BUG-0206 в отдельном примере, запись CONF-CF-01; в реестре CampusFix — 0.'],
    ],
    [
        ['未关闭', 'Открыто'],
        ['5 项：BUG-0201、BUG-0202、BUG-0203、BUG-0204、BUG-0205；均保留待修复或待原系统复核状态。',
         '5: BUG-0201, BUG-0202, BUG-0203, BUG-0204, BUG-0205; ожидают исправления или проверки в исходной системе.'],
    ],
    [
        ['重复项与驳回项', 'Дубликатов и отклонённых'],
        ['共 2 项：I-5 为重复 1 项，归并 BUG-0201；I-6 为非缺陷 1 项。',
         'Всего 2: I-5 — 1 дубликат BUG-0201; I-6 — 1 отклонённый отчёт, не дефект.'],
    ],
    [
        ['对出口准则的影响 测试计划 3.4.1', 'Влияние на критерии выхода Тест-план 3.4.1'],
        ['Critical 与 Major 各有 1 项未关闭，原系统必测仍为 NOT_RUN；拟定出口条件未满足，正式 3.4.1 条款待核对。',
         'Открыты 1 Critical и 1 Major; обязательные тесты исходной системы — NOT_RUN. Предложенные условия выхода не выполнены; официальный п. 3.4.1 требует сверки.'],
    ],
    [
        ['建议 接受 附条件接受 或暂缓', 'Рекомендация Принять С замечаниями Или отложить'],
        ['暂缓验收。先处理越权读取与 Firefox 登录问题，再执行目标版本的必测及回归。',
         'Отложить приёмку. Сначала устранить доступ к чужим данным и сбой входа Firefox, затем выполнить обязательные тесты и регрессию целевой версии.'],
    ],
]

exit_rows = [
    ['EXIT-CF-01',
     ['未关闭的 Critical 与 Major 均为 0。', 'Нет открытых Critical и Major.'],
     ['未满足：BUG-0203 Critical、BUG-0201 Major 均未关闭。',
      'Не выполнено: открыты BUG-0203 Critical и BUG-0201 Major.']],
    ['EXIT-CF-02',
     ['目标系统必测用例全部执行并通过，无 NOT_RUN。',
      'Все обязательные тесты целевой системы выполнены и пройдены; NOT_RUN нет.'],
     ['未满足：TC-CF-01～TC-CF-04 的原系统测试仍为 NOT_RUN。',
      'Не выполнено: TC-CF-01–TC-CF-04 в исходной системе остаются NOT_RUN.']],
    ['EXIT-CF-03',
     ['修复在目标版本中集成，重测与回归证据对应同一对象。',
      'Исправление включено в целевую версию; повторная проверка и регрессия относятся к тому же объекту.'],
     ['独立示例已满足本地确认；原 CampusFix 集成与回归未执行。',
      'Отдельный пример подтверждён локально; интеграция и регрессия исходного CampusFix не выполнены.']],
    ['EXIT-CF-04',
     ['残留风险、验收条件及实际验收决定已记录。',
      'Зафиксированы остаточные риски, условия и фактическое решение о приёмке.'],
     ['待完成：已有残留缺陷；当前只有暂缓建议，无实际验收批准。',
      'Ожидает завершения: есть остаточные дефекты; сейчас дана рекомендация отложить, утверждённой приёмки нет.']],
]

doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Inches(11), Inches(8.5)
sec.top_margin = sec.bottom_margin = Inches(0.43)
sec.left_margin = sec.right_margin = Inches(0.55)
for name in ('Normal', 'Title', 'Heading 1'):
    s = doc.styles[name]
    s.font.name = 'Arial'
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Arial Unicode MS')
    s.paragraph_format.line_spacing = 1.0
    s.paragraph_format.space_after = Pt(3)
doc.styles['Normal'].font.size = Pt(10)
doc.styles['Title'].font.size = Pt(15)
doc.styles['Heading 1'].font.size = Pt(10.5)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(4)
pairs = []


def pair(zh, ru, cell=None, style=None, size=10, bold=False):
    p = doc.add_paragraph(style=style) if cell is None else (cell.paragraphs[0] if not cell.text else cell.add_paragraph())
    if style:
        p.style = style
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.keep_together = True
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(zh)
    r.bold, r.font.size = bold, Pt(size)
    q = doc.add_paragraph(style=style) if cell is None else cell.add_paragraph()
    if style:
        q.style = style
    q.paragraph_format.keep_with_next = style in ('Title', 'Heading 1')
    q.paragraph_format.keep_together = True
    q.paragraph_format.space_before = Pt(0)
    q.paragraph_format.space_after = Pt(3)
    r = q.add_run(ru)
    r.bold, r.font.size = bold, Pt(size)
    pairs.append((p, q, zh, ru))


def table(headers, widths, values, size=9.5):
    t = doc.add_table(rows=1, cols=len(widths))
    t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    for cell, text in zip(t.rows[0].cells, headers):
        pair(*text, cell=cell, bold=True, size=size)
        sh = OxmlElement('w:shd')
        sh.set(qn('w:fill'), 'DAEEF0')
        cell._tc.get_or_add_tcPr().append(sh)
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for values_row in values:
        cells = t.add_row().cells
        for cell, value in zip(cells, values_row):
            if isinstance(value, str):
                cell.paragraphs[0].add_run(value).font.size = Pt(size)
            else:
                pair(*value, cell=cell, size=size)
    for row in t.rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    props = t._tbl.tblPr
    mar = OxmlElement('w:tblCellMar')
    for side in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement('w:' + side)
        e.set(qn('w:w'), '60')
        e.set(qn('w:type'), 'dxa')
        mar.append(e)
    props.append(mar)
    borders = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + side)
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), '4')
        e.set(qn('w:color'), 'D9D9D9')
        borders.append(e)
    props.append(borders)
    return t


pair('CampusFix 任务五 测试报告汇总', 'CampusFix Задание 5 Сводка для отчёта о тестировании', style='Title', size=15)
pair('课程过程汇总  |  2026-10-19  |  v1  |  DRAFT', 'Учебная сводка  |  2026-10-19  |  v1  |  DRAFT', size=9)
pair('本页合计两种测试对象：CampusFix 教学案例与独立 Python 示例。示例的已关闭结果只归属于示例，不计为原 CampusFix 的修复。',
     'Включены два объекта: учебный CampusFix и отдельный пример Python. Закрытие в примере относится только к нему и не подтверждает исправление исходного CampusFix.', size=9)
summary_table = table([['指标', 'Показатель'], ['数值', 'Значение']], [3.25, 6.65], summary_rows, size=9.5)

doc.add_page_break()
pair('缺陷明细与统计核对', 'Реестр дефектов и проверка подсчёта', style='Title', size=15)
pair('任务三的严重程度顺序映射为 S1=Critical、S2=Major、S3=Minor、S4=Trivial；原判断不变。BUG-0206 属局部长度校验问题，本次补定为 S3/Minor。',
     'Соответствие шкале задания 3: S1=Critical, S2=Major, S3=Minor, S4=Trivial; прежние оценки сохранены. Для локальной ошибки длины BUG-0206 в этой сводке задан S3/Minor.', size=9.5)
registry_values = []
for r in main_defects:
    if r['bug'] in case_lookup:
        title = case_lookup[r['bug']]['title']
    else:
        title = ['提交按钮俄语拼写错误', 'Опечатка в русской надписи кнопки отправки']
    registry_values.append([
        r['bug'], ['任务三案例', 'Реестр задания 3'], severity_map[r['severity']],
        ['未关闭 待修复或复核', 'Открыт Ждёт исправления или проверки'], title,
    ])
registry_values.append([
    demo_id, ['独立 Python 示例', 'Отдельный пример Python'], 'Minor',
    ['示例已关闭并复核', 'Закрыт и проверен в примере'],
    ['1001 字符被错误接受，已修复并通过本地回归。',
     'Ошибочное принятие 1001 символа исправлено; локальная регрессия пройдена.'],
])
registry_table = table([
    ['编号', 'ID'], ['对象', 'Объект'], ['严重程度', 'Серьёзность'],
    ['当前状态', 'Статус'], ['问题', 'Проблема'],
], [0.85, 1.30, 1.16, 1.90, 4.69], registry_values, size=9.2)
pair('核对：CampusFix 案例 5=0 已关闭＋5 未关闭；独立示例 1=1 已关闭＋0 未关闭；课程合计 6=1＋5。',
     'Проверка: CampusFix — 5=0 закрытых+5 открытых; отдельный пример — 1=1 закрытый+0 открытых; учебный итог — 6=1+5.', size=9.5)
pair('排除项：I-5 只引用 BUG-0201，不重复计数；I-6 符合预期行为，不计缺陷。',
     'Исключения: I-5 ссылается на BUG-0201 и не добавляет дефект; I-6 соответствует ожидаемому поведению и не учитывается как дефект.', size=9.5)

doc.add_page_break()
pair('出口条件与验收建议', 'Критерии выхода и рекомендация по приёмке', style='Title', size=15)
pair('以下 EXIT-CF-01～EXIT-CF-04 为本练习拟定条件，状态 PROPOSED；正式验收时须与课件引用的测试计划 3.4.1 核对。',
     'EXIT-CF-01–EXIT-CF-04 — предложенные условия упражнения со статусом PROPOSED; для официальной приёмки их нужно сверить с п. 3.4.1 тест-плана.', size=9.5)
exit_table = table([['编号', 'ID'], ['拟定条件', 'Предложенное условие'], ['影响与当前判断', 'Влияние и оценка']],
                   [1.0, 4.4, 4.5], exit_rows, size=9.5)
pair('实际测试证据', 'Реальные свидетельства тестирования', style='Heading 1', bold=True, size=10.5)
pair('CONF-CF-01 记录：独立示例修复前为 17/18 通过，修复后 1001 字符回归 1/1、全部测试 18/18 通过。回归用例包含在这 18 个测试中。',
     'CONF-CF-01: до исправления в примере прошли 17/18; после исправления регрессия на 1001 символе — 1/1, весь набор — 18/18. Регрессионный тест входит в этот набор из 18.', size=9.5)
pair('建议暂缓验收', 'Рекомендуется отложить приёмку', style='Heading 1', bold=True, size=10.5)
pair('先处理 BUG-0203 的越权读取和 BUG-0201 的 Firefox 登录问题，再在原 CampusFix 目标版本执行必测与回归；其余残留问题按风险安排。',
     'Сначала устранить доступ к чужим данным BUG-0203 и вход Firefox BUG-0201, затем выполнить обязательные тесты и регрессию исходной целевой версии CampusFix; остальные дефекты планировать по риску.', size=9.5)

for element in (doc.styles.element, doc.element):
    for border in list(element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
doc.core_properties.title = 'CampusFix 任务五 测试报告汇总'
doc.core_properties.subject = '缺陷统计 出口条件 暂缓验收建议 中俄逐段对照'
doc.core_properties.author = 'StudySpace Coursework'
out = ROOT / (STEM + '.docx')
doc.save(out)
shutil.copy2(SOURCE, DATA / 'source-task5.jpg')

(DATA / 'metrics.json').write_text(json.dumps(metrics, ensure_ascii=False, indent=2) + '\n')
content = {
    'version': 'v1', 'status': 'DRAFT', 'date': '2026-10-19',
    'summary_rows': summary_rows, 'defect_registry': registry_values,
    'proposed_exit_criteria': exit_rows, 'metrics': metrics,
    'language_order': ['zh', 'ru'], 'deliverables': ['docx'],
}
(DATA / 'content.json').write_text(json.dumps(content, ensure_ascii=False, indent=2) + '\n')
with ZipFile(out) as z:
    assert z.testzip() is None
assert len(summary_table.rows) == 8 and len(summary_table.columns) == 2
assert len(registry_table.rows) == 7
assert len(exit_table.rows) == 5
assert len({x[0] for x in registry_values}) == 6
assert sum(x['severity'] == 'S1' for x in main_defects) == 1
for p, q, zh, ru in pairs:
    assert p.text == zh and q.text == ru and p._p.getnext() is q._p
for key, path in input_paths.items():
    assert hashlib.sha256(path.read_bytes()).hexdigest() == input_hashes[key]
assert not list(ROOT.glob('*.md'))
coverage = {
    'source_title': 'Задание 5: сводка для отчёта о тестировании',
    'headers': ['Показатель', 'Значение'],
    'rows': [
        'Найдено дефектов (без дубликатов и не дефектов)',
        'По серьёзности: Critical / Major / Minor / Trivial',
        'Закрыто и проверено', 'Открыто', 'Дубликатов и отклонённых',
        'Влияние на критерии выхода (тест-план 3.4.1)',
        'Рекомендация (принять / принять с замечаниями / отложить)',
    ],
    'source_sha256': hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
    'all_seven_rows_covered': True,
    'separate_test_objects': True,
    'official_plan_not_substituted_with_studyspace_plan': True,
    'timer_10_00': 'Classroom time allowance only',
}
(DATA / 'source-coverage.json').write_text(json.dumps(coverage, ensure_ascii=False, indent=2) + '\n')
check = {
    'template_rows': 7, 'template_columns': 2,
    'arithmetic': 'PASS', 'severity_totals': 'PASS',
    'duplicate_exclusion': 'PASS', 'test_scope_separation': 'PASS',
    'source_log_hashes': 'PASS', 'earlier_deliverables_unchanged': 'PASS',
    'bilingual_pairs': len(pairs), 'zh_above_ru': 'PASS',
    'docx_zip': 'PASS', 'word_only': True, 'render_review': 'PENDING',
    'sha256': hashlib.sha256(out.read_bytes()).hexdigest(),
}
(DATA / 'verification.json').write_text(json.dumps(check, ensure_ascii=False, indent=2) + '\n')
fonts = PARENT / '2026-10-19-bug-report-review-v1/work/fonts'
lo_fonts = '/Users/zhi/.cache/codex-runtimes/codex-primary-runtime/dependencies/native/libreoffice-headless/libreoffice/LibreOfficeDev.app/Contents/Resources/fonts/truetype'
(WORK / 'fonts.conf').write_text(f'<?xml version="1.0"?><!DOCTYPE fontconfig SYSTEM "fonts.dtd"><fontconfig><dir>{fonts}</dir><dir>{lo_fonts}</dir><cachedir>{WORK / "fontcache"}</cachedir></fontconfig>')
print(f'PASS: 7 summary rows, 6 unique coursework defects, 1 local verified closure, 5 open, severity 1/1/3/1, {len(pairs)} bilingual pairs')
print(out)
