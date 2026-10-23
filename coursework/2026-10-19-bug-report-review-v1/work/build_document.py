from pathlib import Path
from zipfile import ZipFile
import json
import hashlib
import shutil

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / 'work'
DATA = ROOT / 'data'
DATA.mkdir(exist_ok=True)
STEM = 'StudySpace_任务一缺陷报告审查_中俄双语示例_v1'

reports = [
    {
        'id': 'BR-1',
        'source_ru': 'Не могу забронировать место. Кнопка не работает.',
        'source_zh': '我无法预约座位，按钮不起作用。',
        'missing_ru': 'Нет точного места сбоя, шагов, данных бронирования и окружения. Не указаны состояние учётной записи, ожидаемая реакция кнопки, фактическое поведение и повторяемость.',
        'missing_zh': '缺少出错页面、操作步骤、预约数据和运行环境；未说明账号状态、按钮的预期反应、实际表现及复现频率。',
        'questions_ru': [
            'На какой странице и какая кнопка не работает? Что сделали до нажатия?',
            'Каковы роль, статус аккаунта и проверка принадлежности к вузу? Не сообщайте пароль.',
            'Какие место, дата, время и часовой пояс выбраны? Есть ли другие бронирования?',
            'Что ожидали и что произошло: кнопка неактивна, нет реакции или показана ошибка? Приложите снимок и текст ошибки.',
            'Каковы версия приложения, ОС и браузер? В скольких попытках повторяется сбой?',
        ],
        'questions_zh': [
            '在哪个页面、哪个按钮失效？点击前执行了哪些步骤？',
            '账号角色、启用状态和学校核验状态是什么？请勿提供密码。',
            '选择了哪个座位、日期、时段及时区？是否已有其他预约？',
            '预期应怎样，实际是按钮禁用、无响应还是出现错误？请附截图和错误原文。',
            '应用、操作系统及浏览器版本是什么？尝试几次、复现几次？',
        ],
    },
    {
        'id': 'BR-2',
        'source_ru': 'После отмены бронь всё ещё видна. Это ошибка.',
        'source_zh': '取消后仍然能看到预约，这是错误。',
        'missing_ru': 'Не указан экран и статус записи: видимость в истории сама по себе не доказывает дефект. Нет основания ожидаемого результата, точных шагов отмены, подтверждения операции и данных о доступности места.',
        'missing_zh': '未说明记录出现在哪个页面、处于什么状态；历史记录仍可见本身不足以判定缺陷。缺少预期结果依据、取消步骤、操作确认及座位是否释放的信息。',
        'questions_ru': [
            'Где осталась запись: в активных бронированиях или истории? Какой статус показан?',
            'Какой результат ожидается по требованию или критерию приёмки? Укажите его ID.',
            'Как отменяли бронь и какое подтверждение получили? Укажите ID брони и время операции с часовым поясом.',
            'Что видно после обновления страницы и нового входа? Освободились ли место и интервал?',
            'Пришлите снимки до и после, версию приложения и обезличенные ответ API или журнал, если доступны.',
        ],
        'questions_zh': [
            '记录留在“有效预约”还是“历史记录”中？显示什么状态？',
            '哪条需求或验收准则规定了预期行为？请提供编号。',
            '如何执行取消、收到什么确认？请提供预约编号、操作时间及时区。',
            '刷新页面和重新登录后怎样？对应座位及时段是否已释放？',
            '请提供操作前后截图、应用版本，以及可获得的脱敏 API 响应或日志。',
        ],
    },
    {
        'id': 'BR-3',
        'source_ru': 'Вечером список мест очень долго загружается. Срочно исправьте.',
        'source_zh': '晚上座位列表加载非常慢，请立即修复。',
        'missing_ru': '«Очень долго» и «вечером» не измерены. Нет порога из требования, интервала измерения, фильтров, объёма данных, нагрузки, окружения и сетевых условий. Срочность не обоснована влиянием на пользователей.',
        'missing_zh': '“非常慢”和“晚上”缺少量化定义。缺少需求中的性能阈值、计时起止点、筛选条件、数据规模、负载、运行环境和网络情况；紧急程度缺少用户影响依据。',
        'questions_ru': [
            'Каковы дата, точное время с часовым поясом, страница и выбранные фильтры?',
            'Сколько секунд занимает загрузка и от какого действия до какого события измерено? Каковы число замеров, минимум, медиана и максимум?',
            'Какой допустимый порог и по какому требованию? Блокируется ли бронирование, сколько пользователей затронуто?',
            'Каковы версия приложения, устройство, ОС, браузер, сеть и наличие VPN? Повторяется ли в другой сети?',
            'Каковы объём выборки и одновременная нагрузка, если известны? Приложите запись экрана и обезличенные сетевые тайминги или журнал.',
        ],
        'questions_zh': [
            '具体日期、带时区的时间、页面及筛选条件是什么？',
            '加载用了几秒？从哪个操作计时到哪个事件？测量几次，最小值、中位数及最大值是多少？',
            '依据哪条需求，允许的耗时是多少？是否阻断预约、影响多少用户？',
            '应用版本、设备、系统、浏览器、网络及 VPN 情况是什么？换网络后能否复现？',
            '如已知，查询数据量和同时访问量是多少？请附录屏及脱敏的网络耗时记录或日志。',
        ],
    },
]

meta = {
    'version': 'v1', 'date': '2026-10-19', 'status': 'DRAFT',
    'source_type': 'USER_AUTHORIZED_SYNTHETIC_CLASSROOM_EXAMPLES',
    'source_image': 'codex-clipboard-29744f44-891a-4ee8-aff3-6975c63ca59e.jpg',
    'reports': reports,
    'conclusion_ru': 'Сначала уточнить данные у автора. До получения проверяемых шагов, ожидаемого и фактического результата, окружения и свидетельств нельзя утверждать воспроизводимость, причину, серьёзность или факт исправления.',
    'conclusion_zh': '先向报告人补齐信息。获得可验证的步骤、预期与实际结果、环境和证据之前，不应断言问题可复现、根因、严重程度或已经修复。',
}
(DATA / 'content.json').write_text(json.dumps(meta, ensure_ascii=False, indent=2) + '\n')

doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width = Inches(11)
sec.page_height = Inches(8.5)
sec.top_margin = sec.bottom_margin = Inches(0.45)
sec.left_margin = sec.right_margin = Inches(0.55)

for name in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Arial Unicode MS')
    style.paragraph_format.space_after = Pt(3)
    style.paragraph_format.line_spacing = 1.02

doc.styles['Normal'].font.size = Pt(10.5)
doc.styles['Title'].font.size = Pt(15)
doc.styles['Heading 1'].font.size = Pt(11)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(6)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(3)


def para(text, style=None, first=None, bold=False, pair_top=False, size=10.5):
    p = first if first is not None else doc.add_paragraph(style=style)
    if first is not None and style:
        p.style = style
    p.paragraph_format.space_after = Pt(2 if pair_top else 5)
    p.paragraph_format.line_spacing = 1.02
    p.paragraph_format.keep_with_next = pair_top
    p.paragraph_format.keep_together = True
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    return p


def pair(zh, ru, cell=None, bold=False, size=10.5):
    if cell is None:
        first = None
    else:
        first = cell.paragraphs[0] if not cell.text else cell.add_paragraph()
    para(zh, first=first, bold=bold, pair_top=True, size=size)
    ru_p = doc.add_paragraph() if cell is None else cell.add_paragraph()
    para(ru, first=ru_p, bold=bold, size=size)


def style_table(table):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [0.64, 3.16, 6.10]
    for col, width in zip(table.columns, widths):
        col.width = Inches(width)
    for row in table.rows:
        no_split = OxmlElement('w:cantSplit')
        row._tr.get_or_add_trPr().append(no_split)
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    for cell in table.rows[0].cells:
        shade = OxmlElement('w:shd')
        shade.set(qn('w:fill'), 'DAEEF0')
        cell._tc.get_or_add_tcPr().append(shade)
    repeat = OxmlElement('w:tblHeader')
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    props = table._tbl.tblPr
    margins = OxmlElement('w:tblCellMar')
    for side in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement('w:' + side)
        e.set(qn('w:w'), '75')
        e.set(qn('w:type'), 'dxa')
        margins.append(e)
    props.append(margins)
    borders = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + side)
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), '4')
        e.set(qn('w:color'), 'A8B7BA')
        borders.append(e)
    props.append(borders)


for index, item in enumerate(reports):
    if index:
        doc.add_page_break()
    para('StudySpace 缺陷报告审查练习', style='Title', pair_top=True, size=15)
    para('StudySpace Анализ отчётов о дефектах', style='Title', size=15)
    pair('任务一 找出问题并提出澄清问题', 'Задание 1 Что не так и что спросить', bold=True)
    pair('课堂示例  |  v1  |  2026-10-19  |  DRAFT', 'Учебные примеры  |  v1  |  2026-10-19  |  DRAFT', size=9)
    para(item['id'] + ' 报告描述', style='Heading 1', pair_top=True, bold=True, size=11)
    para(item['id'] + ' Описание отчёта', style='Heading 1', bold=True, size=11)
    pair(item['source_zh'], item['source_ru'])
    table = doc.add_table(rows=2, cols=3)
    for cell, zh, ru in zip(table.rows[0].cells, ['报告', '缺少什么', '向报告人提出的问题'], ['Отчёт', 'Чего не хватает', 'Вопросы репортёру']):
        pair(zh, ru, cell=cell, bold=True)
    cells = table.rows[1].cells
    para(item['id'], first=cells[0].paragraphs[0], bold=True)
    pair(item['missing_zh'], item['missing_ru'], cell=cells[1])
    for n, (zh, ru) in enumerate(zip(item['questions_zh'], item['questions_ru']), 1):
        pair(f'{n}. {zh}', f'{n}. {ru}', cell=cells[2])
    style_table(table)

for element in (doc.styles.element, doc.element):
    for border in list(element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
doc.core_properties.title = 'StudySpace 缺陷报告审查练习 Анализ отчётов о дефектах'
doc.core_properties.subject = '逐部分中俄对照 中文在上 俄语在下'
doc.core_properties.author = 'StudySpace'
doc.save(ROOT / (STEM + '.docx'))

# The answer is delivered only as a Word document.
md_path = ROOT / (STEM + '.md')
if md_path.exists():
    md_path.unlink()

with ZipFile(ROOT / (STEM + '.docx')) as z:
    assert z.testzip() is None
assert len(doc.tables) == 3
assert all(len(t.rows) == 2 and len(t.columns) == 3 for t in doc.tables)
for table, item in zip(doc.tables, reports):
    missing = table.rows[1].cells[1].paragraphs
    assert missing[0].text == item['missing_zh']
    assert missing[1].text == item['missing_ru']
    question_paras = table.rows[1].cells[2].paragraphs
    assert len(question_paras) == 10
    for i, (zh, ru) in enumerate(zip(item['questions_zh'], item['questions_ru'])):
        assert question_paras[2 * i].text == f'{i + 1}. {zh}'
        assert question_paras[2 * i + 1].text == f'{i + 1}. {ru}'
body = [p.text for p in doc.paragraphs if p.text]
for item in reports:
    n = body.index(item['source_zh'])
    assert body[n + 1] == item['source_ru']
meta['layout'] = 'PART_BY_PART_ZH_ABOVE_RU'
meta['deliverable_formats'] = ['docx']
(DATA / 'content.json').write_text(json.dumps(meta, ensure_ascii=False, indent=2) + '\n')
check = {
    'source_image_columns': ['Отчёт', 'Чего не хватает', 'Вопросы репортёру'],
    'report_ids': [item['id'] for item in reports],
    'tables': 3, 'rows_per_table_including_header': 2, 'columns_per_table': 3,
    'questions_per_report_per_language': 5,
    'language_order': ['zh', 'ru'], 'pairing': 'PART_BY_PART',
    'paragraph_pair_order_check': 'PASS',
    'source_type': meta['source_type'],
    'docx_zip': 'PASS', 'bilingual_content_parity': 'PASS',
    'render_review': 'PENDING', 'business_testing': 'NOT_RUN',
    'deliverable_formats': ['docx'], 'markdown_answer_removed': not md_path.exists(),
    'sha256': {'docx': hashlib.sha256((ROOT / (STEM + '.docx')).read_bytes()).hexdigest()},
}
(DATA / 'verification.json').write_text(json.dumps(check, ensure_ascii=False, indent=2) + '\n')
print('PASS: three paired bilingual tables; Chinese above Russian for every description, gap and question; DOCX only')
print(ROOT / (STEM + '.docx'))
