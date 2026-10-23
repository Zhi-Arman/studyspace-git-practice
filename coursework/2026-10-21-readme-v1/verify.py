from pathlib import Path
import json, re, hashlib
from datetime import datetime
from zoneinfo import ZoneInfo
from zipfile import ZipFile
from docx import Document
from docx.oxml.ns import qn

HERE=Path(__file__).resolve().parent
ROOT=HERE.parent.parent
ART=HERE/'artifacts'
checks=[]
def check(name, passed, **extra):
 checks.append({'check':name,'passed':bool(passed),**extra})
md=(HERE/'README.md').read_text()
rootmd=(ROOT/'README.md').read_text()
blocks=json.loads((ART/'content.json').read_text())
for path in (ROOT/'README.md', HERE/'README.md'):
 missing=[];total=0
 for target in re.findall(r'\[[^]]+\]\(([^)]+)\)',path.read_text()):
  if '://' in target or target.startswith('#'):continue
  total+=1
  if not (path.parent/target.split('#')[0]).exists():missing.append(target)
 check(path.relative_to(ROOT).as_posix()+' local links',not missing,total=total,missing=missing)
docpath=HERE/'StudySpace_README_中俄双语_v1.docx'
with ZipFile(docpath) as z:check('DOCX ZIP integrity',z.testzip() is None)
doc=Document(docpath)
txt='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
for b in blocks:
 if b['kind']=='heading': check(b['zh']+' Chinese Russian headings',b['zh'] in txt and b['ru'] in txt)
 elif b['kind']=='pair':check('paired paragraph '+b['zh'][:12],b['zh'] in txt and b['ru'] in txt)
 elif b['kind']=='code':check('identical code '+b['text'].splitlines()[0],b['text'] in txt)
 else:
  for row in b['rows']:
   for cell in row:
    for target in re.findall(r'\[[^]]+\]\(([^)]+)\)',cell):check('Word path '+target,target in txt)
    plain=re.sub(r'\[([^]]+)\]\(([^)]+)\)',lambda m:m[1]+'\n'+m[2],cell)
    check('table content '+row[0].splitlines()[0],plain in txt)
check('exactly six bilingual sections',len(re.findall(r'^## [1-6] ',md,re.M))==6)
check('two editable tables',len(doc.tables)==2)
check('no decorative borders',not list(doc.element.iter(qn('w:pBdr'))))
check('business NOT_RUN preserved','NOT_RUN' in md and 'NOT_IMPLEMENTED' in md)
check('six root topics',all(s in rootmd for s in ['项目与边界','需求 架构 API 与测试入口','如何打开和构建文档','分支与团队流程','已验证的内容与仍在设计的内容','仓库中不保存什么']))
check('Sphinx warnings empty',(ART/'sphinx-warnings.log').stat().st_size==0)
commands=[
 ['backend/.venv/bin/python','scripts/check_docs.py'],
 ['backend/.venv/bin/python','scripts/generate_registers.py','--check'],
 ['backend/.venv/bin/python','scripts/check_api_contract.py'],
 ['coursework/2026-10-21-docs-as-code-environment-v1/08-sphinx-site/project-docs/.venv/bin/python','-m','sphinx','-b','html','-n','-W','--keep-going','-E','-a','-N','-w','coursework/2026-10-21-readme-v1/artifacts/sphinx-warnings.log','coursework/2026-10-21-docs-as-code-environment-v1/08-sphinx-site/project-docs/docs','coursework/2026-10-21-readme-v1/artifacts/site']]
for i,command in enumerate(commands):
 p=ART/f'check-{i}.json';data=json.loads(p.read_text());data['command']=command;data['cwd']=str(ROOT)
 p.write_text(json.dumps(data,ensure_ascii=False,indent=2))
 check('actual check '+str(i),data['value']['exit_code']==0)
check('dependency locks retained',(ROOT/'frontend/package-lock.json').exists() and (ROOT/'backend/uv.lock').exists())
report={'checked_at':datetime.now(ZoneInfo('Asia/Shanghai')).isoformat(),'checks':checks,'passed':sum(x['passed'] for x in checks),'total':len(checks),'all_pass':all(x['passed'] for x in checks),'business_tests':'NOT_RUN','docx_sha256':hashlib.sha256(docpath.read_bytes()).hexdigest()}
(ART/'verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
print(json.dumps({k:v for k,v in report.items() if k!='checks'},ensure_ascii=False))
for item in checks:
 if not item['passed']:print(item)
raise SystemExit(0 if report['all_pass'] else 1)
