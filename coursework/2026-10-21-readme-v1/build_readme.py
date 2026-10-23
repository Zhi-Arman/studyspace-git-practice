from pathlib import Path
import json, re, hashlib, platform, shutil
from datetime import datetime
from zoneinfo import ZoneInfo
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent.parent
ART = HERE / 'artifacts'
SITE = 'coursework/2026-10-21-docs-as-code-environment-v1/08-sphinx-site/project-docs'
TITLE = 'StudySpace 项目 README'
SUBTITLE = 'README учебного проекта StudySpace'
META = 'SS-README-01 · v1.0 · DRAFT · 2026-10-21 · Asia/Shanghai'
blocks = []
def heading(zh, ru): blocks.append({'kind':'heading','zh':zh,'ru':ru})
def pair(zh, ru): blocks.append({'kind':'pair','zh':zh,'ru':ru})
def code(text): blocks.append({'kind':'code','text':text})
def table(headers, rows): blocks.append({'kind':'table','headers':headers,'rows':rows})
def link(label, path): return f'[{label}]({path})'

heading('1 学习项目的用途', '1 Назначение учебного проекта')
pair('StudySpace 是大学软件工程课程项目，面向学生和学校管理人员，练习需求、架构、API、测试和文档协作。MVP 目标是查询和预约教室、自习室中的单个座位，预约前验证校内身份。',
     'StudySpace — учебный проект по программной инженерии для студентов и сотрудников вуза. Он охватывает требования, архитектуру, API, тесты и совместную документацию. Цель MVP — поиск и бронирование отдельного места с предварительной проверкой принадлежности к вузу.')
pair('应用版本 0.1.0，需求版本 0.3.0。当前已有 React 与 FastAPI 骨架、三语首页、健康与数据库就绪接口；认证、校内身份核验、资源管理、预约业务仍待实现。正式文档为 DRAFT，需求为 PROPOSED，业务用例为 NOT_RUN。',
     'Приложение 0.1.0, требования 0.3.0. Имеются каркас React и FastAPI, главная страница на трёх языках, API работоспособности и готовности БД. Аутентификация, проверка вуза, ресурсы и бронирование ещё не реализованы. Документы — DRAFT, требования — PROPOSED, бизнес-тесты — NOT_RUN.')

heading('2 需求 架构 API 与测试的位置', '2 Расположение требований архитектуры API и тестов')
pair('下表路径均相对于 StudySpace 根目录。本提交件位于 coursework/2026-10-21-readme-v1；Markdown 中的链接可直接打开对应文件。',
     'Пути ниже заданы от корня StudySpace. Эта работа находится в coursework/2026-10-21-readme-v1; ссылки Markdown открывают соответствующие файлы.')
table(['资料 Материал','根目录相对路径 Путь от корня'],[
 ['总入口与长期规则\nОбщий вход и правила',link('项目 README','README.md')+'\n'+link('AGENTS','AGENTS.md')+'\n'+link('文档目录 Каталог','docs/README.md')],
 ['需求与澄清\nТребования и уточнения',link('中文 SRS','docs/zh-CN/03_SRS_IEEE830.md')+'\n'+link('Русская SRS','docs/ru-RU/03_SRS_IEEE830.md')+'\n'+link('需求源 Реестр','docs/data/requirements.json')+'\n'+link('决定 Решения','docs/data/project-decisions.json')],
 ['架构与 ADR\nАрхитектура и ADR',link('中文设计','docs/zh-CN/04_System_Design.md')+'\n'+link('Русский проект','docs/ru-RU/04_System_Design.md')+'\n'+link('架构决策 Решения','docs/decisions/README.md')],
 ['API 契约\nКонтракт API',link('OpenAPI YAML','docs/api/openapi.yaml')+'\n'+link('中文契约','docs/zh-CN/17_API_Contract.md')+'\n'+link('Русский контракт','docs/ru-RU/17_API_Contract.md')],
 ['测试设计\nПроектирование тестов',link('中文计划','docs/zh-CN/05_Test_Plan.md')+'\n'+link('Русский план','docs/ru-RU/05_Test_Plan.md')+'\n'+link('中文用例','docs/zh-CN/06_Test_Cases.md')+'\n'+link('Русские сценарии','docs/ru-RU/06_Test_Cases.md')+'\n'+link('用例源 Сценарии','docs/data/test-cases.json')],
 ['可运行检查与追溯\nПроверки и трассируемость',link('后端测试 Тесты backend','backend/tests/')+'\n'+link('文档测试 Тесты документов','tests/')+'\n'+link('前端测试 Тест frontend','frontend/src/App.test.tsx')+'\n'+link('中文追溯','docs/zh-CN/07_Traceability_Matrix.md')+'\n'+link('Русская трассируемость','docs/ru-RU/07_Traceability_Matrix.md')],
 ['课程文档网站\nУчебный сайт документации',link('网站项目 Проект сайта',SITE+'/README.md')+'\n'+link('HTML 首页 Главная',SITE+'/site/index.html')],
])

heading('3 如何打开和构建文档', '3 Как открыть и собрать документацию')
pair('阅读源码：在 VS Code 打开项目根目录，打开 README.md 或 docs 下的 Markdown，按 Cmd+Shift+V 预览；也可直接打开上表 HTML 首页。网站是独立的课程文档项目，记录已有应用骨架，其 materials 是课程快照；最新正式需求以根目录 docs 为准。',
     'Для чтения открыть корень в VS Code, выбрать README.md либо Markdown из docs и нажать Cmd+Shift+V; можно открыть HTML из таблицы. Сайт — отдельный учебный проект о каркасе приложения, а materials содержит учебные снимки. Актуальные официальные требования находятся в корневом docs.')
pair('构建网站：从根目录进入以下目录。已有 .venv 时直接构建；新环境先用 Python 3.13 创建 .venv，再按带哈希的锁文件安装。首次安装需要网络。',
     'Сборка сайта: перейти из корня в каталог ниже. При наличии .venv сразу выполнить сборку; в новой среде создать .venv на Python 3.13 и установить зависимости по lock-файлу с хешами. Первая установка требует сети.')
code('cd "'+SITE+'"\n# 仅新环境执行 / Только для новой среды\npython3.13 -m venv .venv\n.venv/bin/python -m pip install --require-hashes -r requirements-docs.txt\n# 已有环境从此开始 / Для готовой среды начать здесь\nbash build-site.sh\n.venv/bin/python -m http.server 8765 --bind 127.0.0.1 --directory site')
pair('然后打开 http://127.0.0.1:8765/，按 Ctrl+C 停止静态服务器。build-site.sh 使用 Sphinx 的 -n -W --keep-going -E -a 严格构建，HTML 写入 site/，日志写入该文档项目的 artifacts/。若 8765 已占用，选择空闲端口。',
     'Затем открыть http://127.0.0.1:8765/; Ctrl+C останавливает сервер. build-site.sh запускает строгую сборку Sphinx с -n -W --keep-going -E -a. HTML записывается в site/, журналы — в artifacts/ проекта документации. Если 8765 занят, выбрать свободный порт.')
pair('正式文档一致性检查：在 StudySpace 根目录、应用依赖已安装后执行下列命令。它们检查结构、登记表同步和静态 API 契约；根目录没有统一 Sphinx 构建命令。应用安装与启动方法见根 README 第 3 节。',
     'Проверки официальных документов выполняются из корня StudySpace после установки зависимостей приложения. Они проверяют структуру, синхронизацию реестров и статический контракт API; единой сборки Sphinx в корне нет. Установка и запуск приложения описаны в разделе 3 корневого README.')
code('backend/.venv/bin/python scripts/check_docs.py\nbackend/.venv/bin/python scripts/generate_registers.py --check\nbackend/.venv/bin/python scripts/check_api_contract.py')

heading('4 分支名称与团队工作顺序', '4 Ветки и порядок командной работы')
pair('main 为集成主线；feature/<topic> 用于功能，fix/<topic> 用于缺陷，docs/<topic> 用于文档。例如 docs/readme。它们是约定，不表示所有分支已创建。本地根仓库当前为尚无首次提交的 main。',
     'main — основная линия интеграции; feature/<topic> — функции, fix/<topic> — исправления, docs/<topic> — документация, например docs/readme. Это соглашение, а не перечень созданных веток. Локальный корневой репозиторий находится на main без первого коммита.')
pair('工作顺序：明确任务和关联 ID → 在已建立提交基线后创建主题分支 → 小范围修改并同步中俄文档 → 运行相关检查并记录真实结果 → 提交小批量变更 → 在远程协作已获授权并配置后开 PR → 由另一成员评审 → 经授权合并回 main。',
     'Порядок: определить задачу и ID → после появления базового коммита создать тематическую ветку → внести небольшие изменения и синхронно обновить оба языка → выполнить проверки и записать фактический результат → создать небольшой коммит → после разрешения и настройки удалённой работы открыть PR → получить ревью другого участника → с разрешения слить в main.')
pair('提交前缀为 feat:、fix:、docs:、test:、refactor:、chore:。示例 docs: clarify README navigation。需求变更须同步 JSON 来源、测试设计和追溯；未经授权不推送、合并或部署。',
     'Префиксы коммитов: feat:, fix:, docs:, test:, refactor:, chore:. Пример: docs: clarify README navigation. Изменение требований сопровождается обновлением JSON, проекта тестов и трассируемости. Отправка, слияние и развёртывание требуют разрешения.')
pair('角色：支浩然为分析员与技术任务书作者；俞智文为文档编辑；曹浩然为测试员；吕嘉伟为项目经理。开发全员协作，主题轮换要记录交接；另一成员的评审不能由自动检查替代。',
     'Роли: 支浩然 — аналитик и автор ТЗ; 俞智文 — редактор; 曹浩然 — тестировщик; 吕嘉伟 — менеджер. Разработка совместная; смена роли по теме фиксируется передачей. Автоматические проверки не заменяют ревью другого участника.')

heading('5 已验证的内容与仍在设计的内容', '5 Что проверено и что ещё проектируется')
pair('2026-10-21 在本机实际复核如下。执行命令和原始输出保存于本提交目录 artifacts/check-0.json 至 check-3.json；网站复构建输出放在本目录 artifacts/site，保留原课程网站不变。',
     'Ниже фактические проверки на этом компьютере 2026-10-21. Команды и исходный вывод сохранены в artifacts/check-0.json–check-3.json этой работы. Повторная сборка сайта записана в artifacts/site этой работы; прежний сайт сохранён.')
table(['验证对象 Объект','结果与边界 Результат и границы'],[
 ['正式文档结构\nСтруктура документов','PASS：中俄元数据、章节和 ID、需求与测试引用一致。\nPASS: совпадают метаданные, разделы, ID и ссылки.'],
 ['登记表与 API\nРеестры и API','PASS：生成表与 JSON 一致；API 契约静态检查 110 项通过。\nPASS: реестры совпадают с JSON; 110 статических проверок API пройдены.'],
 ['文档网站构建\nСборка сайта','PASS：Sphinx 严格构建退出码 0，警告日志 0 字节。\nPASS: строгая сборка Sphinx, код 0, журнал предупреждений 0 байт.'],
 ['业务功能与验收\nБизнес-функции и приёмка','NOT_IMPLEMENTED / NOT_RUN：认证、核验、资源和预约尚未实现；未进行真实业务验收。\nNOT_IMPLEMENTED / NOT_RUN: аутентификация, проверка вуза, ресурсы и бронирование не реализованы; бизнес-приёмка не проведена.'],
 ['本轮未执行\nСейчас не выполнялось','应用完整检查 scripts/check.sh、容器部署、真实学校核验、HTTPS、迁移与并发验收。\nПолная проверка приложения scripts/check.sh, контейнерное развёртывание, проверка вуза, HTTPS, миграция и конкурентная приёмка.'],
])
pair('健康接口可连接、文档构建成功或静态契约通过，均不能把业务 TC 改为 PASS。SRS 与正式文档仍为 DRAFT，需求仍为 PROPOSED；翻译语义、团队评审和教师批准须另有证据。',
     'Доступность health API, сборка документации и статические проверки не переводят бизнес-TC в PASS. SRS и официальные документы остаются DRAFT, требования — PROPOSED. Смысл перевода, командное ревью и одобрение преподавателя требуют отдельных свидетельств.')
pair('未完成的外部前提为 OPEN-003 学校真实核验来源、OPEN-008 课程具体条件、OPEN-010 部署位置与域名。提醒、签到、违约和使用率为未来范围 FR-018 至 FR-021。',
     'Незавершённые внешние условия: OPEN-003 — источник проверки вуза, OPEN-008 — условия курса, OPEN-010 — сервер и домен. Напоминания, присутствие, нарушения и загрузка относятся к будущим FR-018–FR-021.')

heading('6 仓库中不保存什么', '6 Что не хранится в репозитории')
pair('禁止提交真实密码、API 密钥、访问令牌、私钥、生产 .env、学生身份材料和真实个人数据、数据库与备份、生产日志及带敏感信息的截图。可提交不含真实秘密的 .env.example 与明确标记的合成样例。提交前人工复核，不能仅依赖忽略规则。',
     'Не коммитить реальные пароли, API-ключи, токены, закрытые ключи, рабочий .env, документы студентов и персональные данные, БД и резервные копии, производственные журналы и снимки с секретами. Допустимы .env.example без секретов и явно обозначенные синтетические примеры. Перед коммитом нужна ручная проверка; одного ignore недостаточно.')
pair('现有 .gitignore 排除 .env 及其变体（保留 .env.example）、.venv/、node_modules/、dist/、缓存、*.db 与 *.sqlite* 数据、backend/data/、*.log、.DS_Store、根 artifacts/ 与 .gstack/ 等。依赖环境和应用构建产物应按锁文件重建。',
     'Текущий .gitignore исключает .env и варианты, сохраняя .env.example, а также .venv/, node_modules/, dist/, кэши, БД *.db и *.sqlite*, backend/data/, *.log, .DS_Store, корневой artifacts/ и .gstack/. Среды и сборки приложения воспроизводятся по lock-файлам.')
pair('课程提交件有例外：本项目已保留脱敏后的课程证据、DOCX 和静态网站作为交付材料，不应把所有 artifacts/ 或 site/ 都称为已忽略。仅保留确有提交用途且不含敏感数据的文件，渲染中间 PNG/PDF 不作为本次提交正文。',
     'Для учебной сдачи есть исключения: проект сохраняет обезличенные свидетельства, DOCX и статический сайт. Нельзя считать все artifacts/ или site/ автоматически исключёнными. Сохраняются только нужные для сдачи файлы без чувствительных данных; промежуточные PNG/PDF не входят в основной результат этой работы.')

# Root README retains its existing project instructions and gains concrete coverage.
root_readme = ROOT / 'README.md'
before = root_readme.read_text()
backup = ART / 'root-readme-before.md'
if not backup.exists(): backup.write_text(before)
before = backup.read_text()
def mdblocks(items, prefix=''):
    def transform(s):
        return re.sub(r'\[([^]]+)\]\(([^)]+)\)',lambda m:f'[{m[1]}]({prefix}{m[2]})',s).replace('\n','<br>')
    out=[]
    for b in items:
        if b['kind']=='heading': out.append('## '+b['zh']+' / '+b['ru'])
        elif b['kind']=='pair': out.extend([re.sub(r'(feature|fix|docs)/<topic>',lambda m:'`'+m[0]+'`',b[k]) for k in ('zh','ru')])
        elif b['kind']=='code': out.append('```bash\n'+b['text']+'\n```')
        else:
            out.append('| '+' | '.join(b['headers'])+' |\n| '+' | '.join(['---']*len(b['headers']))+' |\n'+'\n'.join('| '+' | '.join(transform(c) for c in row)+' |' for row in b['rows']))
    return '\n\n'.join(out)+'\n'
sections=[]
for b in blocks:
    if b['kind']=='heading': sections.append([])
    sections[-1].append(b)
new = before.replace('Updated: **2026-10-16**','Updated: **2026-10-21**')
new = new.replace('当前Docker守护进程未运行，容器构建/启动尚未验证。','初始化时Docker守护进程未运行；本轮文档工作未复测容器构建/启动。')
new = new.replace('При инициализации Docker-демон не запущен; сборка/запуск контейнеров не проверены.','При инициализации Docker-демон не был запущен; в этой работе сборка и запуск контейнеров повторно не проверялись.')
nav = '\n### 5.1 需求 架构 API 与测试入口 / Входы в требования архитектуру API и тесты\n\n'+mdblocks(sections[1][1:])
nav += '\n### 5.2 分支与团队流程 / Ветки и командная работа\n\n'+mdblocks(sections[3][1:])
new = new.replace('## 6. 目录导航 / Навигация по каталогам',nav+'\n## 6. 目录导航 / Навигация по каталогам')
supplement=''
for n,idx in [(7,2),(8,4),(9,5)]:
    h=sections[idx][0]
    supplement += f'## {n}. '+re.sub(r'^\d+ ','',h['zh'])+' / '+re.sub(r'^\d+ ','',h['ru'])+'\n\n'+mdblocks(sections[idx][1:])+'\n'
# Evidence paths are local to the coursework folder, clarify in root view.
supplement=supplement.replace('本提交目录 artifacts/check-0.json 至 check-3.json','coursework/2026-10-21-readme-v1/artifacts/check-0.json 至 check-3.json')
supplement=supplement.replace('本目录 artifacts/site','coursework/2026-10-21-readme-v1/artifacts/site')
supplement=supplement.replace('artifacts/check-0.json–check-3.json этой работы','coursework/2026-10-21-readme-v1/artifacts/check-0.json–check-3.json')
supplement=supplement.replace('artifacts/site этой работы','coursework/2026-10-21-readme-v1/artifacts/site')
supplement += '课堂六项 README 提交版 / Учебная версия README по шести пунктам: [Markdown](coursework/2026-10-21-readme-v1/README.md) · [Word](coursework/2026-10-21-readme-v1/StudySpace_README_中俄双语_v1.docx)。\n\n'
new = new.replace('依据 / Источники:',supplement+'依据 / Источники:')
root_readme.write_text(new)

# Root-relative paths remain literal in the visible table; only Markdown targets gain ../../.
md = '# '+TITLE+'\n\n'+SUBTITLE+'\n\n'+META+'\n\n'+mdblocks(blocks,'../../')
(HERE/'README.md').write_text(md)
(ART/'content.json').write_text(json.dumps(blocks,ensure_ascii=False,indent=2))
shutil.copy2('/var/folders/ff/kqw394c93dd7krl0z2hcqqdh0000gn/T/codex-clipboard-d521e186-e181-416d-9391-3bab850d8e93.png',ART/'task.png')

doc=Document(); sec=doc.sections[0]
sec.page_width,sec.page_height=Cm(21),Cm(29.7)
sec.left_margin=sec.right_margin=Cm(1.8)
sec.top_margin,sec.bottom_margin=Cm(1.6),Cm(1.6)
for name,size in [('Normal',10),('Title',20),('Subtitle',13),('Heading 1',13),('Heading 2',11)]:
    st=doc.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
    st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Hiragino Sans GB')
    st.paragraph_format.space_after=Pt(4);st.paragraph_format.line_spacing=1.05
    if name.startswith('Heading'): st.paragraph_format.keep_with_next=True;st.paragraph_format.space_before=Pt(9)
doc.add_paragraph(TITLE,'Title');doc.add_paragraph(SUBTITLE,'Subtitle');doc.add_paragraph(META)

def rich(p,txt):
    # Include visible root-relative targets beside labels so the printed copy remains usable.
    txt=re.sub(r'\[([^]]+)\]\(([^)]+)\)',lambda m:m[1]+'\n'+m[2],txt)
    for j,line in enumerate(txt.splitlines()):
        if j:p.add_run().add_break()
        r=p.add_run(line)
        if '/' in line or line.endswith('.md') or line.endswith('.json'):r.font.size=Pt(8)

for b in blocks:
    kind=b['kind']
    if kind=='heading':
        # Keep each section coherent; two navigation pages are allowed naturally.
        if b['zh'][0] in ('3','4','5','6'):doc.add_page_break()
        doc.add_paragraph(b['zh'],'Heading 1'); p=doc.add_paragraph(b['ru'],'Heading 2');p.paragraph_format.space_before=Pt(0)
    elif kind=='pair':
        p=doc.add_paragraph(b['zh']);p.paragraph_format.keep_with_next=True
        doc.add_paragraph(b['ru'])
    elif kind=='code':
        for i,line in enumerate(b['text'].splitlines()):
            p=doc.add_paragraph(line);p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=1
            p.paragraph_format.keep_with_next=i<len(b['text'].splitlines())-1
            for r in p.runs:r.font.name='Courier New';r.font.size=Pt(8)
    else:
        t=doc.add_table(rows=1,cols=2);t.autofit=False;t.columns[0].width=Cm(4.3);t.columns[1].width=Cm(13.1)
        for i,h in enumerate(b['headers']):t.rows[0].cells[i].text=h
        t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for row in b['rows']:
            cells=t.add_row().cells
            for i,txt in enumerate(row):rich(cells[i].paragraphs[0],txt)
        for ri,row in enumerate(t.rows):
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for ci,c in enumerate(row.cells):
                c.width=Cm(4.3 if ci==0 else 13.1)
                pr=c._tc.get_or_add_tcPr(); borders=OxmlElement('w:tcBorders')
                for edge in ('top','left','bottom','right'):
                    el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D6D6D6');borders.append(el)
                pr.append(borders)
                if ri==0:
                    el=OxmlElement('w:shd');el.set(qn('w:fill'),'EFEFEF');pr.append(el)
                for p in c.paragraphs:
                    p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3);p.paragraph_format.line_spacing=1
                    for r in p.runs:
                        if r.font.size is None:r.font.size=Pt(9)
                        r.bold=ri==0
        doc.add_paragraph()
for tree in (doc.styles.element,doc.element):
    for el in list(tree.iter(qn('w:pBdr'))):el.getparent().remove(el)
out=HERE/'StudySpace_README_中俄双语_v1.docx';doc.save(out)
print(json.dumps({'root_readme':str(root_readme),'markdown':str(HERE/'README.md'),'docx':str(out),'sections':len(sections)},ensure_ascii=False))
