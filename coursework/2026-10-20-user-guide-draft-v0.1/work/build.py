from pathlib import Path
from datetime import datetime, timezone
import hashlib
import json
import re
from zipfile import ZipFile

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = Path(__file__).resolve().parents[1]
WORK = BASE / 'work'
OUT = BASE / 'VSCode_Markdown操作指南初稿_中俄双语_v0.1.docx'
OLD = BASE.parent / '2026-10-20-markdown-user-guide-v1'
AUTHORS = '支浩然、俞智文、曹浩然、吕嘉伟'
SOURCES = [
    ('https://code.visualstudio.com/docs/editing/getting-started/editor-tutorial#_open-a-folder-in-vs-code', 'VS Code Editor tutorial'),
    ('https://code.visualstudio.com/docs/languages/markdown#_markdown-preview', 'VS Code Markdown preview'),
]

DATA = {
 'zh': {
  'title': 'VS Code Markdown 操作指南初稿',
  'subtitle': '中文版',
  'meta': [
   '版本：0.1    状态：初稿    日期：2026-10-20',
   '作者：StudySpace 课程小组（' + AUTHORS + '）',
   '适用环境：macOS 27.0.1，VS Code 1.142.0，英文界面，默认快捷键，UTF-8 编码。',
  ],
  'intro': '本指南面向第一次使用 Markdown 的学生。Markdown 是用普通文本标记标题和段落的格式。先创建并保存文件，再打开预览。使用 VS Code 内置功能即可，无需安装扩展；菜单和命令保留英文原名。',
  'labels': ['目的', '前提条件', '操作步骤', '预期结果', '常见错误及处理'],
  'tasks': [
   {
    'id': 'A', 'heading': '操作 A 创建并保存 Markdown 文件',
    'goal': '在学习文件夹中创建 user-guide.md，保存标题和正文，并确认文件能够重新打开。',
    'pre': '已安装并启动 VS Code；已通过 Finder 创建空文件夹 markdown-practice，知道其位置并有写入权限；其中尚无同名文件。',
    'steps': [
     '选择 File > Open Folder...，在文件夹选择窗口中选中 markdown-practice 并确认打开。',
     '若出现工作区信任提示，仅对自己创建且内容已知的学习文件夹选择 Yes, I trust the authors。确认左侧 Explorer 显示 markdown-practice。',
     '在 Explorer 中右键单击该文件夹，选择 New File...；输入 user-guide.md 并按 Enter。保留 .md 扩展名，它用于识别 Markdown 文件。',
     '在编辑区输入下方示例，第一行的 # 后保留一个空格，标题与正文之间空一行。只输入示例内容。',
     '按 Cmd+S 保存。确认标签页的未保存圆点消失，并检查 Explorer 中的文件名确为 user-guide.md。',
     '关闭该文件标签页，再单击 Explorer 中的 user-guide.md 重新打开。确认标题和正文仍在；随后保持文件打开，继续操作 B。',
    ],
    'sample': '# 我的第一份 Markdown 文档\n\n这是我的学习笔记。',
    'result': 'markdown-practice 中存在已保存的 user-guide.md；重新打开后仍能看到相同的标题和正文。',
    'errors': [
     '文件名成了 user-guide.md.txt：在 Explorer 中右键单击文件并选择 Rename，将其改为 user-guide.md。若出现同名文件提示，先查看已有文件，勿覆盖。',
     '出现权限错误或无法保存：选择自己有写入权限的学习文件夹重新操作；已有输入时先复制文本，防止丢失。',
     '左侧没有 Explorer：按 Cmd+Shift+E 显示文件列表，再确认已打开正确的文件夹。',
    ],
   },
   {
    'id': 'B', 'heading': '操作 B 打开并检查 Markdown 预览',
    'goal': '在编辑区旁显示排版后的内容，并通过修改正文确认预览会同步更新。',
    'pre': '已完成操作 A；user-guide.md 已在 VS Code 中打开，标题和正文已保存。',
    'steps': [
     '单击 user-guide.md 的编辑区，使其成为活动文件。确认右下角语言模式为 Markdown。',
     '按 Cmd+Shift+P 打开命令面板。输入并选择 Markdown: Open Preview to the Side，然后按 Enter。',
     '确认编辑区旁出现预览窗格。比较两侧内容：源文保留 #，预览中的第一行应显示为较大的标题，正文显示为普通段落。',
     '回到源文，在正文末尾添加“ 预览成功。”并按 Cmd+S。确认相同文字出现在预览中；预览只显示效果，修改应在源文中完成。',
     '关闭预览后，按第 2 步重新打开。确认新增文字仍在，且文件标签页没有未保存圆点。',
    ],
    'result': '源文与预览并排显示；标题、正文和新增文字排版正确，修改后的内容已保存。',
    'errors': [
     '找不到预览命令：重新选中 .md 文件；若语言模式为 Plain Text，单击它并选择 Markdown，再打开命令面板。',
     '标题仍显示为普通文字：在源文中把第一行改为“# 标题”，确保 # 后有空格，且该行没有置于三反引号代码块内。',
     '预览内容没有变化：确认修改的是当前预览对应的文件，保存后关闭并重新打开该文件的预览。',
    ],
   },
  ],
  'sources_title': '参考资料',
  'sources_note': '界面名称核对日期：2026-10-20。',
 },
 'ru': {
  'title': 'Руководство по Markdown в VS Code',
  'subtitle': 'Русская версия',
  'meta': [
   'Версия: 0.1    Статус: черновик    Дата: 20.10.2026',
   'Авторы: учебная группа StudySpace (' + AUTHORS + ').',
   'Среда: macOS 27.0.1, VS Code 1.142.0, английский интерфейс, стандартные сочетания клавиш, кодировка UTF-8.',
  ],
  'intro': 'Руководство предназначено для студентов, которые впервые используют Markdown. Этот формат позволяет обозначать заголовки и абзацы в обычном тексте. Сначала создайте и сохраните файл, затем откройте предпросмотр. Достаточно встроенных функций VS Code; расширения не нужны. Названия меню и команд приведены на английском языке.',
  'labels': ['Цель', 'Предусловия', 'Порядок действий', 'Ожидаемый результат', 'Ошибки и их устранение'],
  'tasks': [
   {
    'id': 'A', 'heading': 'Инструкция A Создание и сохранение файла Markdown',
    'goal': 'Создать user-guide.md в учебной папке, сохранить заголовок и текст и проверить повторное открытие файла.',
    'pre': 'VS Code установлен и запущен. Пустая папка markdown-practice создана через Finder; её расположение известно и запись в неё разрешена. Одноимённого файла в папке пока нет.',
    'steps': [
     'Выберите «File > Open Folder...». В окне выбора папки найдите markdown-practice и подтвердите открытие.',
     'Если появится запрос доверия, выберите Yes, I trust the authors только для собственной учебной папки с известным содержимым. Убедитесь, что слева в Explorer отображается markdown-practice.',
     'В Explorer щёлкните правой кнопкой по этой папке и выберите «New File...». Введите user-guide.md и нажмите Enter. Сохраните расширение .md: оно обозначает файл Markdown.',
     'Введите в редакторе пример ниже. После # в первой строке оставьте один пробел, между заголовком и абзацем оставьте пустую строку. Вводите только содержимое примера.',
     'Нажмите Cmd+S. Убедитесь, что круглый индикатор несохранённых изменений на вкладке исчез, а в Explorer указан файл user-guide.md.',
     'Закройте вкладку файла и снова откройте user-guide.md из Explorer. Проверьте, что заголовок и абзац сохранились. Оставьте файл открытым для инструкции B.',
    ],
    'sample': '# Мой первый документ Markdown\n\nЭто мои учебные заметки.',
    'result': 'В markdown-practice существует сохранённый user-guide.md. После повторного открытия в нём видны прежние заголовок и абзац.',
    'errors': [
     'Имя файла получилось user-guide.md.txt: щёлкните по файлу правой кнопкой в Explorer, выберите Rename и замените имя на user-guide.md. При сообщении о совпадении имён сначала проверьте существующий файл, не перезаписывайте его.',
     'Нет разрешения на запись или сохранение не выполняется: выберите собственную доступную для записи учебную папку и повторите действия. Если текст уже введён, сначала скопируйте его, чтобы не потерять.',
     'Explorer не виден: нажмите Cmd+Shift+E, чтобы показать список файлов, и проверьте, что открыта нужная папка.',
    ],
   },
   {
    'id': 'B', 'heading': 'Инструкция B Открытие и проверка предпросмотра Markdown',
    'goal': 'Показать оформленный документ рядом с редактором и проверить обновление предпросмотра после изменения текста.',
    'pre': 'Инструкция A выполнена. Файл user-guide.md открыт в VS Code, заголовок и абзац сохранены.',
    'steps': [
     'Щёлкните в редакторе user-guide.md, чтобы сделать его активным. Убедитесь, что в правом нижнем углу указан режим языка Markdown.',
     'Нажмите Cmd+Shift+P для открытия палитры команд. Введите и выберите Markdown: Open Preview to the Side, затем нажмите Enter.',
     'Убедитесь, что рядом с редактором появилась панель предпросмотра. Сравните содержимое: в исходном тексте сохраняется #, а в предпросмотре первая строка показана крупным заголовком; основной текст отображается обычным абзацем.',
     'Вернитесь к исходному тексту, добавьте в конец абзаца « Предпросмотр работает.» и нажмите Cmd+S. Проверьте появление этих слов в предпросмотре. Он показывает оформление; редактировать следует исходный текст.',
     'Закройте предпросмотр и откройте его снова по шагу 2. Убедитесь, что добавленный текст остался, а на вкладке файла нет индикатора несохранённых изменений.',
    ],
    'result': 'Исходный текст и предпросмотр показаны рядом. Заголовок, абзац и добавленные слова оформлены правильно; изменения сохранены.',
    'errors': [
     'Команда предпросмотра не находится: снова выберите файл .md. Если указан режим Plain Text, нажмите на него и выберите Markdown, затем откройте палитру команд.',
     'Заголовок выглядит как обычный текст: измените первую строку на «# Заголовок», оставьте пробел после # и убедитесь, что строка не помещена в блок кода с тройными обратными кавычками.',
     'Предпросмотр не меняется: проверьте, что редактируете соответствующий ему файл. Сохраните изменения, закройте и заново откройте предпросмотр этого файла.',
    ],
   },
  ],
  'sources_title': 'Источники',
  'sources_note': 'Названия элементов интерфейса проверены 20.10.2026.',
 },
}

WORK.mkdir(parents=True, exist_ok=True)
preserved = {}
for p in [OLD / 'VSCode_Markdown初学者操作指南_中俄双语.docx', OLD / 'user-guide.md']:
    preserved[str(p)] = hashlib.sha256(p.read_bytes()).hexdigest()
(WORK / 'content.json').write_text(json.dumps(DATA, ensure_ascii=False, indent=2), encoding='utf-8')

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.65)
sec.left_margin = sec.right_margin = Inches(.72)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    st = doc.styles[name]
    st.font.name = 'Arial'
    st.font.color.rgb = RGBColor(0, 0, 0)
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.insert(0, rf)
    for k, value in [('ascii', 'Arial'), ('hAnsi', 'Arial'), ('eastAsia', 'Arial Unicode MS')]:
        rf.set(qn('w:' + k), value)
normal = doc.styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = 1.08
normal.paragraph_format.space_after = Pt(4)
normal.paragraph_format.widow_control = True
doc.styles['Title'].font.size = Pt(19)
doc.styles['Title'].paragraph_format.space_after = Pt(4)
doc.styles['Subtitle'].font.size = Pt(11)
doc.styles['Subtitle'].paragraph_format.space_after = Pt(8)
doc.styles['Heading 1'].font.size = Pt(13)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(10)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(5)
doc.styles['Heading 2'].font.size = Pt(11)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(6)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(3)

def para(text='', *, style=None, size=None, bold=False, before=0, after=None, keep=False, indent=False):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    if after is not None: p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = keep
    p.paragraph_format.keep_together = True
    if current == 'zh':
        p.paragraph_format.line_spacing = Pt(16)
    if indent:
        p.paragraph_format.left_indent = Inches(.18)
        p.paragraph_format.first_line_indent = Inches(-.18)
    r = p.add_run(text)
    r.bold = bold
    if size: r.font.size = Pt(size)
    return p

def labeled(label, text):
    p = para()
    p.add_run(label + ('：' if current == 'zh' else ': ')).bold = True
    p.add_run(text)
    return p

def hyperlink(label, url):
    p = para(after=2)
    h = OxmlElement('w:hyperlink')
    rel = doc.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    h.set(qn('r:id'), rel)
    r = OxmlElement('w:r')
    pr = OxmlElement('w:rPr')
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '19'); pr.append(sz)
    r.append(pr)
    t = OxmlElement('w:t'); t.text = label; r.append(t)
    h.append(r); p._p.append(h)

for current in ['zh', 'ru']:
    data = DATA[current]
    title = para(data['title'], style='Title')
    title.paragraph_format.line_spacing = 1.0
    title.paragraph_format.page_break_before = current == 'ru'
    para(data['subtitle'], style='Subtitle')
    for line in data['meta']: para(line, size=9.5, after=3)
    para(data['intro'], before=5, after=6)
    for ti, task in enumerate(data['tasks']):
        head = para(task['heading'], style='Heading 1')
        head.paragraph_format.page_break_before = ti == 1
        head.paragraph_format.line_spacing = 1.0
        labeled(data['labels'][0], task['goal'])
        labeled(data['labels'][1], task['pre'])
        para(data['labels'][2], style='Heading 2')
        for num, step in enumerate(task['steps'], 1):
            para(f'{num}. {step}', indent=True)
            if task['id'] == 'A' and num == 4:
                p = para(task['sample'], after=6)
                p.paragraph_format.left_indent = Inches(.18)
        labeled(data['labels'][3], task['result'])
        para(data['labels'][4], style='Heading 2')
        for num, err in enumerate(task['errors'], 1):
            para(f'{num}. {err}', indent=True)
    para(data['sources_title'], style='Heading 2')
    para(data['sources_note'], size=9.5, after=3)
    for idx, (url, label) in enumerate(SOURCES, 1):
        hyperlink(f'[{idx}] {label}', url)

for border in list(doc.styles.element.xpath('.//w:pBdr')) + list(doc.element.xpath('.//w:pBdr')):
    border.getparent().remove(border)
for grid in list(doc.element.xpath('.//w:docGrid')):
    grid.getparent().remove(grid)
for p in doc.paragraphs:
    snap = OxmlElement('w:snapToGrid'); snap.set(qn('w:val'), '0')
    p._p.get_or_add_pPr().append(snap)

props = doc.core_properties
props.title = 'VS Code Markdown 操作指南初稿 / Руководство по Markdown в VS Code'
props.subject = '两项初学者操作说明 / Две инструкции для начинающих'
props.author = 'StudySpace 课程小组: ' + AUTHORS
props.version = '0.1'
props.content_status = 'DRAFT'
props.created = props.modified = datetime(2026, 10, 20, 9, 0, tzinfo=timezone.utc)
props.keywords = 'Markdown, VS Code, 中文, русский, v0.1'
doc.save(OUT)

text = '\n'.join(p.text for p in Document(OUT).paragraphs)
ru_parts = [DATA['ru']['intro']]
for task in DATA['ru']['tasks']:
    ru_parts.extend([task['goal'], task['pre'], *task['steps'], task.get('sample', ''), task['result'], *task['errors']])
ru_word_count = len(re.findall(r'\S+', '\n'.join(ru_parts)))
checks = {
    'exactly_two_instructions_per_language': all(len(d['tasks']) == 2 for d in DATA.values()),
    'all_five_sections_per_instruction': all(all(task.get(k) for k in ['goal','pre','steps','result','errors']) for d in DATA.values() for task in d['tasks']),
    'six_and_five_steps_in_both_versions': all([len(t['steps']) for t in d['tasks']] == [6,5] for d in DATA.values()),
    'three_error_solutions_per_instruction': all(len(t['errors']) == 3 for d in DATA.values() for t in d['tasks']),
    'version_date_author_environment_present': all(k in text for k in ['0.1','2026-10-20','20.10.2026',AUTHORS,'27.0.1','1.142.0','UTF-8']),
    'russian_body_500_to_800_words': 500 <= ru_word_count <= 800,
    'all_authored_paragraphs_present': all(s in text for d in DATA.values() for t in d['tasks'] for s in [t['goal'],t['pre'],*t['steps'],t['result'],*t['errors']]),
    'original_files_preserved': all(hashlib.sha256(Path(p).read_bytes()).hexdigest() == h for p,h in preserved.items()),
    'no_placeholder_text': not any(s in text for s in ['TODO','[作者]','待填写','待补充','ФИО']),
    'core_version_and_author': props.version == '0.1' and AUTHORS in props.author,
}
with ZipFile(OUT) as z:
    checks['zip_integrity'] = z.testzip() is None
    checks['four_clickable_source_links'] = z.read('word/document.xml').count(b'<w:hyperlink ') == 4
    checks['no_header_footer'] = not any(re.match(r'word/(header|footer)\d+\.xml', n) for n in z.namelist())
report = {
    'date': '2026-10-20', 'checks': checks, 'all_pass': all(checks.values()),
    'russian_body_word_count_whitespace': ru_word_count,
    'official_sources_verified_with_gstack_browse': [u for u,_ in SOURCES],
    'source_image': '/var/folders/ff/kqw394c93dd7krl0z2hcqqdh0000gn/T/codex-clipboard-8200b162-f1f1-4649-870e-ed480557cc96.jpg',
    'preserved_sources_sha256': preserved,
    'output': str(OUT), 'sha256': hashlib.sha256(OUT.read_bytes()).hexdigest(),
    'visual_review': 'PENDING', 'novice_usability_trial': 'NOT_RUN',
    'gstack_status': 'DONE', 'durable_learning_review': 'No new durable learnings; reused existing paired-document and fontconfig workflow.',
}
(WORK / 'verification.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({'output':str(OUT), 'russian_body_words':ru_word_count,'checks':checks}, ensure_ascii=False))
assert all(checks.values()), 'Document content validation failed'
