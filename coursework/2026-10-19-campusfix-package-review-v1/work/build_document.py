from pathlib import Path
from collections import Counter
from zipfile import ZipFile
import hashlib
import json
import re
import shutil
from lxml import etree
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
DATA, WORK = ROOT / 'data', ROOT / 'work'
DATA.mkdir(exist_ok=True)
PARENT = ROOT.parent
STEM = 'CampusFix_缺陷材料包评审检查表_中俄逐段对照_v1'
SOURCE = Path('/var/folders/ff/kqw394c93dd7krl0z2hcqqdh0000gn/T/codex-clipboard-1b20fa0d-e85d-4fde-8a90-476ca5b3fa14.jpg')
folders = {key: PARENT / ('2026-10-19-campusfix-' + name + '-v1') for key, name in (
    ('reports', 'evidence-reports'), ('triage', 'triage'), ('fix', 'fix-confirmation'), ('summary', 'test-summary'))}
inputs = []
for folder in folders.values():
    inputs.extend(folder.glob('*.docx'))
    inputs.extend((folder / 'data').glob('*.json'))
inputs.extend((folders['fix'] / 'data').glob('*.log'))
inputs.extend((folders['fix'] / 'demo').rglob('*.py'))
hashes = {str(p): hashlib.sha256(p.read_bytes()).hexdigest() for p in inputs}
execution = json.loads((folders['fix'] / 'data/execution.json').read_text())
for run in execution['runs']:
    assert hashlib.sha256(Path(run['log']).read_bytes()).hexdigest() == run['log_sha256']
assert hashlib.sha256((folders['fix'] / 'demo/validator.py').read_bytes()).hexdigest() == execution['after_source_sha256']
assert hashlib.sha256((folders['fix'] / 'demo/tests/test_validator.py').read_bytes()).hexdigest() == execution['test_source_sha256']
with ZipFile(next(folders['reports'].glob('*.docx'))) as z:
    xml = etree.fromstring(z.read('word/document.xml'))
ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
anchors = xml.xpath('//w:hyperlink/@w:anchor', namespaces=ns)
bookmarks = xml.xpath('//w:bookmarkStart/@w:name', namespaces=ns)
assert len(anchors) == 24 and len(bookmarks) == 12 and all(a in bookmarks for a in anchors)

rows = [
    {'id': 'R-01', 'question': ['标题是否包含症状和触发条件？', 'Заголовки содержат симптом и условие?'], 'result': ['部分', 'Частично'],
     'comment': ['BUG-0201～0204 符合；0205 尚无完整报告，0206 缺独立缺陷标题。', 'BUG-0201–0204 соответствуют; у 0205 нет полного отчёта, у 0206 — отдельного заголовка дефекта.']},
    {'id': 'R-02', 'question': ['无需询问作者即可复现步骤？', 'Шаги воспроизводимы без вопросов к автору?'], 'result': ['部分', 'Частично'],
     'comment': ['本地示例可复跑；原系统缺服务器地址、账号准备及可恢复测试数据。', 'Пример можно повторить; для исходной системы не хватает адреса сервера, подготовки учётных записей и данных.']},
    {'id': 'R-03', 'question': ['预期结果与实际结果是否区分？', 'Ожидаемый и фактический результат разделены?'], 'result': ['是', 'Да'],
     'comment': ['0201～0204 分列两者；I-7 区分错字与正确标签；0206 有修复前后结果。', 'В 0201–0204 поля разделены; I-7 различает ошибочную и верную надпись; 0206 содержит результаты до и после.']},
    {'id': 'R-04', 'question': ['严重程度与优先级是否有依据？', 'Серьёзность и приоритет обоснованы?'], 'result': ['部分', 'Частично'],
     'comment': ['任务三有影响与绕行依据；0206 仅在汇总补定 Minor，未记录优先级。', 'В задании 3 учтены влияние и обход; для 0206 в сводке задан только Minor, приоритет не указан.']},
    {'id': 'R-05', 'question': ['是否关联需求、验收准则和测试用例？', 'Связи с требованием, критерием и тест-кейсом указаны?'], 'result': ['部分', 'Частично'],
     'comment': ['0201～0204 链接有效；0206 有 FR/AC/TC；0205 只有 UI-CF-01。', 'Связи 0201–0204 работают; у 0206 есть FR/AC/TC; у 0205 указано только UI-CF-01.']},
    {'id': 'R-06', 'question': ['附件是否必要且无秘密和个人数据？', 'Вложения нужные, без секретов и персональных данных?'], 'result': ['部分', 'Частично'],
     'comment': ['示例日志和补丁有用；失败日志含本机用户路径，原系统证据仍未取得。', 'Логи и патч примера полезны; лог сбоя содержит локальный путь пользователя, свидетельств исходной системы пока нет.']},
    {'id': 'R-07', 'question': ['重复项与非缺陷的判定是否合理？', 'Решения по дубликатам и «не дефектам» разумны?'], 'result': ['是', 'Да'],
     'comment': ['I-5 同条件归并 0201；I-6 匿名访问返回 401，符合教学规则 AUTH-CF-01。', 'I-5 при тех же условиях связан с 0201; ответ 401 для анонимного I-6 соответствует учебному AUTH-CF-01.']},
]
counts = Counter(r['result'][1] for r in rows)
assert counts == {'Частично': 5, 'Да': 2}
actions = [
    ('R-01', ['补全缺陷标题', 'Дополнить заголовки дефектов'],
     ['建议 BUG-0205：俄语新建申请表单的提交按钮显示 «Отпавить»；建议 BUG-0206：去除首尾空白后描述长为 1001 字符时 Validate 仍接受输入。标题为本次建议，未回写原报告。',
      'Предлагается для BUG-0205: «В русской форме новой заявки кнопка отправки подписана “Отпавить”». Для BUG-0206: «Validate принимает описание длиной 1001 символ после удаления внешних пробелов». Это предложения без изменения прежних отчётов.']),
    ('R-02', ['补齐原系统复现条件', 'Дополнить условия воспроизведения'],
     ['补服务器、构建号、SSO 入口、测试账号和数据准备及重置方法；BUG-0201 另需年份、时区和尝试次数。报告不得包含密码或令牌。',
      'Добавить сервер, сборку, вход SSO, подготовку и сброс аккаунтов и данных; для BUG-0201 — год, часовой пояс и число попыток. Пароли и токены не включать.']),
    ('R-04 R-05', ['补齐分级与追溯', 'Дополнить оценку и трассируемость'],
     ['BUG-0206 的 Minor 为任务五新增教学判断，需补优先级和依据；BUG-0205 需形成完整报告并关联需求、验收准则、测试用例，新增编号时先查登记表。',
      'Minor для BUG-0206 — учебная оценка задания 5; добавить приоритет и обоснование. Для BUG-0205 оформить полный отчёт и связи с требованием, критерием и тестом; перед новым ID проверить реестр.']),
    ('R-06', ['整理证据和脱敏副本', 'Подготовить свидетельства и обезличенные копии'],
     ['三份日志及源代码校验值一致；26 份文本的有限模式扫描未命中私钥或常见 API 密钥。仍需检查账号、Cookie 和截图，对外副本去除本机用户路径，原始日志保留。',
      'Хеши трёх логов и кода совпадают. Ограниченный поиск в 26 текстах не выявил закрытых ключей и типовых API-ключей. Ещё проверить аккаунты, Cookie и снимки; внешние копии очистить от локальных путей, оригиналы сохранить.']),
]

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.5)
sec.left_margin = sec.right_margin = Inches(.55)
sec.header_distance = sec.footer_distance = Inches(.25)
for name, size in [('Normal', 10.5), ('Title', 15), ('Heading 1', 11)]:
    s = doc.styles[name]
    s.font.name = 'Arial'
    s.font.size = Pt(size)
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Arial Unicode MS')
    s.paragraph_format.space_after = Pt(3)
    s.paragraph_format.line_spacing = 1.02
pairs = []
def pair(zh, ru, cell=None, style=None, size=10.5, center=False):
    target = cell if cell is not None else doc
    if cell is not None and len(cell.paragraphs) == 1 and not cell.paragraphs[0].text:
        p = cell.paragraphs[0]
    else:
        p = target.add_paragraph()
    q = target.add_paragraph()
    for para, text in [(p, zh), (q, ru)]:
        if style: para.style = style
        run = para.add_run(text)
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor(0, 0, 0)
        para.paragraph_format.space_before = Pt(0)
        para.paragraph_format.space_after = Pt(2)
        para.paragraph_format.line_spacing = 1.0
        if center: para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    q.paragraph_format.keep_with_next = style in ('Title', 'Heading 1')
    pairs.append((p, q, zh, ru))
    return p, q

pair('CampusFix 缺陷材料包评审检查表', 'CampusFix Чек лист рецензии дефект пакета', style='Title', size=15)
pair('2026-10-19  v1  课程材料自查  DRAFT', '2026-10-19  v1  Учебная самопроверка  DRAFT', size=9)
pair('评审结论：2 项是，5 项部分，0 项否。材料需补充后复核；本次不变更缺陷或验收状态。',
     'Итог: 2 «Да», 5 «Частично», 0 «Нет». Пакет требует доработки и повторной проверки; статусы дефектов и приёмки не изменяются.', size=10)
t = doc.add_table(rows=1, cols=3)
t.autofit = False
t.alignment = WD_TABLE_ALIGNMENT.CENTER
widths = [2.55, .85, 4.0]
for col, w in zip(t.columns, widths): col.width = Inches(w)
for cell, w, texts in zip(t.rows[0].cells, widths, [['问题', 'Вопрос'], ['是 部分 否', 'Да Частично Нет'], ['意见', 'Замечание']]):
    cell.width = Inches(w)
    pair(*texts, cell=cell, size=9.5, center=True)
    shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'DAEEF0'); cell._tc.get_or_add_tcPr().append(shade)
header = OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(header)
for row in rows:
    cells = t.add_row().cells
    for c,w in zip(cells,widths): c.width = Inches(w)
    pair(*row['question'], cell=cells[0], size=9.5)
    pair(*row['result'], cell=cells[1], size=9.5, center=True)
    pair(*row['comment'], cell=cells[2], size=9.5)
for row in t.rows:
    cant = OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(cant)
    for cell in row.cells: cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
props = t._tbl.tblPr
borders = OxmlElement('w:tblBorders')
for edge in ('top','left','bottom','right','insideH','insideV'):
    el = OxmlElement('w:' + edge)
    for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]: el.set(qn('w:' + k),v)
    borders.append(el)
props.append(borders)
mar = OxmlElement('w:tblCellMar')
for edge,n in [('top','75'),('bottom','75'),('left','85'),('right','85')]:
    el = OxmlElement('w:' + edge); el.set(qn('w:w'),n); el.set(qn('w:type'),'dxa'); mar.append(el)
props.append(mar)

doc.add_page_break()
pair('评审依据与补充事项', 'Основания рецензии и доработка', style='Title', size=15)
pair('范围与证据', 'Область и свидетельства', style='Heading 1', size=11)
pair('审查任务二至五的 BUG-0201～0206、I-5 和 I-6。任务二的 24 个内部链接均对应 12 个已定义目标。',
     'Проверены материалы заданий 2–5: BUG-0201–0206, I-5 и I-6. Все 24 внутренние ссылки задания 2 ведут к 12 определённым целям.', size=10.5)
pair('原系统仍为 NOT_RUN。本次仅核对示例已有日志：修复前 17/18、修复后回归 1/1、全部 18/18；回归包含在全部测试中，未重新运行。',
     'Исходная система — NOT_RUN. Проверены прежние логи примера: до исправления 17/18, после — регрессия 1/1 и весь набор 18/18. Регрессия входит в набор; повторного запуска не было.', size=10.5)
for rid,title,content in actions:
    pair(rid + ' ' + title[0], rid + ' ' + title[1], style='Heading 1', size=11)
    pair(*content, size=10.5)
pair('复核与验收边界', 'Повторная проверка и границы приёмки', style='Heading 1', size=11)
pair('补充上述材料后逐项复核。保留任务五的暂缓验收建议：Critical 与 Major 各有 1 项未关闭，正式测试计划 3.4.1 尚未提供。评审完成不代表教师批准或原系统修复。',
     'После дополнения повторить проверку по пунктам. Рекомендация задания 5 сохраняется: отложить приёмку, открыты 1 Critical и 1 Major; официальный п. 3.4.1 тест-плана не предоставлен. Рецензия не означает одобрения преподавателем или исправления исходной системы.', size=10.5)
for element in (doc.styles.element, doc.element):
    for border in list(element.iter(qn('w:pBdr'))): border.getparent().remove(border)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('CampusFix  v1  |  ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
for run in footer.runs: run.font.size = Pt(8)
doc.core_properties.title = 'CampusFix 缺陷材料包评审检查表'
doc.core_properties.subject = '七项检查 证据审查 补充事项 中俄逐段对照'
doc.core_properties.author = 'StudySpace Coursework'
out = ROOT / (STEM + '.docx')
doc.save(out)
shutil.copy2(SOURCE, DATA / 'source-checklist.jpg')
for p,q,zh,ru in pairs:
    assert p.text == zh and q.text == ru and p._p.getnext() is q._p
for p,h in hashes.items(): assert hashlib.sha256(Path(p).read_bytes()).hexdigest() == h
with ZipFile(out) as z: assert z.testzip() is None
assert len(t.rows) == 8 and len(t.columns) == 3
assert not list(ROOT.glob('*.md'))
audit = {'reviewed_input_sha256':hashes,'original_report_links':len(anchors),'defined_link_targets':len(bookmarks),
         'existing_log_hashes':'PASS','existing_source_and_tests_hashes':'PASS','tests_rerun':False,
         'original_system_tests':'NOT_RUN','privacy_scan_scope':'26 related JSON LOG PATCH PY files in CampusFix data and demo folders',
         'privacy_scan_limit':'Only private-key markers and common API token patterns; not a comprehensive privacy audit',
         'credential_pattern_matches':0,'local_path_in_failure_log':True}
(DATA/'audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2)+'\n')
(DATA/'content.json').write_text(json.dumps({'status':'DRAFT','date':'2026-10-19','rows':rows,'counts':dict(counts),'actions':actions,'language_order':['zh','ru']},ensure_ascii=False,indent=2)+'\n')
(DATA/'source-coverage.json').write_text(json.dumps({'source_title':'Чек-лист рецензии дефект-пакета','headers':['Вопрос','Да / Частично / Нет','Замечание'],'seven_questions':[r['question'][1] for r in rows],'all_seven_covered':True,'timer':'06:00 classroom allowance','source_sha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest()},ensure_ascii=False,indent=2)+'\n')
(DATA/'verification.json').write_text(json.dumps({'rows':7,'columns':3,'counts':dict(counts),'bilingual_pairs':len(pairs),'zh_above_ru':'PASS','input_files_unchanged':'PASS','zip_integrity':'PASS','word_only':True,'render_review':'PENDING','output_sha256':hashlib.sha256(out.read_bytes()).hexdigest()},ensure_ascii=False,indent=2)+'\n')
fonts = PARENT / '2026-10-19-bug-report-review-v1/work/fonts'
lo_fonts = '/Users/zhi/.cache/codex-runtimes/codex-primary-runtime/dependencies/native/libreoffice-headless/libreoffice/LibreOfficeDev.app/Contents/Resources/fonts/truetype'
(WORK/'fonts.conf').write_text(f'<?xml version="1.0"?><!DOCTYPE fontconfig SYSTEM "fonts.dtd"><fontconfig><dir>{fonts}</dir><dir>{lo_fonts}</dir><cachedir>{WORK/"fontcache"}</cachedir></fontconfig>')
print(out)
print('PASS 7 checklist rows 3 columns 2 yes 5 partial 0 no',len(pairs),'paired paragraphs')
