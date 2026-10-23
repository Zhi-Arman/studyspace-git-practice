from __future__ import annotations

import hashlib
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
PROJECT = ROOT.parents[1]
CONTENT = ROOT / "data/content.json"
OUTPUT = ROOT / "StudySpace_任务三数据模式_中俄双语_v1.docx"
VERIFY = ROOT / "data/verification.json"
data = json.loads(CONTENT.read_text(encoding="utf-8"))


def font(run, *, size=8.6, bold=False, color=None, mono=False):
    run.font.name = "Courier New" if mono else "Arial"
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key, val in (("ascii", run.font.name), ("hAnsi", run.font.name), ("eastAsia", "Heiti SC")):
        rfonts.set(qn("w:" + key), val)


def pair(container, value, *, size=8.6, bold=False, first=False):
    for index, lang in enumerate(("zh", "ru")):
        p = container.paragraphs[0] if first and index == 0 else container.add_paragraph()
        p.paragraph_format.space_after = Pt(2 if lang == "zh" else 0)
        p.paragraph_format.line_spacing = 1.05
        font(p.add_run(value[lang]), size=size, bold=bold)


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def padding(cell, top=65, bottom=65, left=80, right=80):
    tcpr = cell._tc.get_or_add_tcPr()
    mar = tcpr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tcpr.append(mar)
    for side, value in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        node = mar.find(qn("w:" + side))
        if node is None:
            node = OxmlElement("w:" + side)
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def prevent_split(row, *, repeat=False):
    trpr = row._tr.get_or_add_trPr()
    trpr.append(OxmlElement("w:cantSplit"))
    if repeat:
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        trpr.append(header)


def json_text(example):
    return json.dumps(example, ensure_ascii=False, separators=(", ", ": "))


doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width = Inches(16.54)
sec.page_height = Inches(11.69)
sec.left_margin = sec.right_margin = Inches(0.48)
sec.top_margin = Inches(0.43)
sec.bottom_margin = Inches(0.41)
sec.header_distance = sec.footer_distance = Inches(0.20)

normal = doc.styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(8.5)
normal.paragraph_format.space_after = Pt(2)
normal.paragraph_format.line_spacing = 1.04

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(1)
font(p.add_run(data["title"]["zh"]), size=16, bold=True, color=(18, 67, 83))
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(1)
font(p.add_run(data["title"]["ru"]), size=13, bold=True, color=(18, 67, 83))
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(5)
font(p.add_run(f'{data["document_id"]}  ·  {data["version"]}  ·  {data["date"]}  ·  {data["status"]}'), size=8, color=(85, 95, 104))

for lang in ("zh", "ru"):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1 if lang == "zh" else 5)
    font(p.add_run(data["intro"][lang]), size=8.8)

widths = [2.58, 2.55, 1.75, 5.14, 3.56]
assert round(sum(widths), 2) == 15.58
table = doc.add_table(rows=1, cols=5)
table.autofit = False
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = "Table Grid"
for index, width in enumerate(widths):
    table.columns[index].width = Inches(width)
for index, header in enumerate(data["headers"]):
    cell = table.rows[0].cells[index]
    cell.width = Inches(widths[index])
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    shade(cell, "DCEEF0")
    padding(cell, top=75, bottom=75)
    pair(cell, header, size=9.0, bold=True, first=True)
prevent_split(table.rows[0], repeat=True)

for row_index, item in enumerate(data["rows"]):
    row = table.add_row()
    prevent_split(row)
    for col_index, cell in enumerate(row.cells):
        cell.width = Inches(widths[col_index])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        padding(cell)
        if row_index % 2:
            shade(cell, "F4F8F9")
    cell = row.cells[0]
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    font(p.add_run(item["id"]), size=8.8, bold=True, color=(18, 67, 83))
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    font(p.add_run(item["schema"]), size=8.6, bold=True)
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    font(p.add_run(item["endpoint"]), size=7.4, color=(85, 95, 104))
    pair(row.cells[1], item["fields"], size=8.5, first=True)
    pair(row.cells[2], item["required"], size=8.4, first=True)
    pair(row.cells[3], item["constraints"], size=8.3, first=True)
    cell = row.cells[4]
    if item["id"] == "DS-07":
        for index, lang in enumerate(("zh", "ru")):
            p = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
            p.paragraph_format.space_after = Pt(2 if lang == "zh" else 0)
            font(p.add_run(lang + ": "), size=7.7, bold=True)
            font(p.add_run(json_text(item["example"][lang])), size=7.3, mono=True)
    else:
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        font(p.add_run(item["example"] if isinstance(item["example"], str) else json_text(item["example"])), size=7.35, mono=True)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(6)
p.paragraph_format.space_after = Pt(1)
font(p.add_run("说明 / Примечания"), size=9.3, bold=True, color=(18, 67, 83))
for note in data["notes"]:
    for lang in ("zh", "ru"):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        font(p.add_run(note[lang]), size=7.9)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(4)
p.paragraph_format.space_after = Pt(1)
font(p.add_run("依据与交付 / Источники и файл"), size=8.9, bold=True, color=(18, 67, 83))
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(0)
p.paragraph_format.line_spacing = 1.0
font(p.add_run(" · ".join(data["sources"])), size=6.9, color=(85, 95, 104))
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(0)
font(p.add_run(data["delivery_path"]), size=7.2, color=(85, 95, 104))
for lang in ("zh", "ru"):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    font(p.add_run(data["environment"][lang]), size=7.5, color=(85, 95, 104))

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)

check = Document(OUTPUT)
assert len(check.tables) == 1
assert len(check.tables[0].rows) == 1 + len(data["rows"])
assert all(len(row.cells) == 5 for row in check.tables[0].rows)
text = "\n".join(p.text for p in check.paragraphs)
text += "\n" + "\n".join(c.text for row in check.tables[0].rows for c in row.cells)
required = list(data["title"].values()) + list(data["intro"].values())
for entry in data["headers"] + data["notes"]:
    required.extend(entry.values())
for item in data["rows"]:
    required.extend((item["id"], item["schema"], item["endpoint"]))
    for key in ("fields", "required", "constraints"):
        required.extend(item[key].values())
    if item["id"] == "DS-07":
        required.extend(json_text(item["example"][lang]) for lang in ("zh", "ru"))
    else:
        required.append(item["example"] if isinstance(item["example"], str) else json_text(item["example"]))
required.extend(data["sources"])
required.extend((data["delivery_path"], *data["environment"].values()))
missing = [entry for entry in required if entry not in text]
assert not missing, missing
assert all((PROJECT / source).exists() for source in data["sources"])
assert (PROJECT / data["delivery_path"]).resolve() == OUTPUT.resolve()
result = {
    "document": str(OUTPUT),
    "sha256": hashlib.sha256(OUTPUT.read_bytes()).hexdigest(),
    "status": data["status"],
    "schema_count": len(data["rows"]),
    "table_dimensions": [len(check.tables[0].rows), len(check.tables[0].columns)],
    "paired_text_checks": len(required),
    "structural_result": "PASS",
    "business_tests": "NOT_RUN",
}
VERIFY.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps(result, ensure_ascii=False, indent=2))
