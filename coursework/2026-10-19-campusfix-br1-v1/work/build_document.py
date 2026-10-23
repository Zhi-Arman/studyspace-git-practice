from pathlib import Path
from zipfile import ZipFile
import hashlib
import json
import shutil

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / 'data'
WORK = ROOT / 'work'
DATA.mkdir(exist_ok=True)
STEM = 'CampusFix_BR-1_报告人答复与缺陷报告_中俄逐段对照_v1'
SOURCE = Path('/var/folders/ff/kqw394c93dd7krl0z2hcqqdh0000gn/T/codex-clipboard-42f4c1bc-21dd-4e11-935e-ce5bca656805.jpg')

answers = [
    {
        'question_zh': '版本与环境', 'question_ru': 'Версия и среда',
        'answer_zh': 'CampusFix 1.0.3，测试服务器；Firefox 131，Windows 11。',
        'answer_ru': 'CampusFix 1.0.3, тестовый сервер; Firefox 131, Windows 11',
    },
    {
        'question_zh': '具体发生了什么', 'question_ru': 'Что именно происходит',
        'answer_zh': '在 SSO 中输入密码后，页面返回登录表单，没有任何提示信息。',
        'answer_ru': 'После ввода пароля в SSO страница возвращается на форму входа без сообщения',
    },
    {
        'question_zh': '时间与用户', 'question_ru': 'Время и пользователь',
        'answer_zh': '11 月 3 日 09:15；student09。',
        'answer_ru': '3 ноября, 09:15; student09',
    },
    {
        'question_zh': '复现情况', 'question_ru': 'Воспроизводимость',
        'answer_zh': '在 Firefox 中每次都出现；在 Chrome 130 中登录正常。',
        'answer_ru': 'Всегда в Firefox; в Chrome 130 вход работает',
    },
    {
        'question_zh': '预期结果', 'question_ru': 'Ожидаемое',
        'answer_zh': '登录成功，并打开“我的申请”列表。',
        'answer_ru': 'Вход выполнен, открывается список «Мои заявки»',
    },
]

summary = (
    'Windows 11 的 Firefox 131 中，CampusFix 1.0.3 在 SSO 输入密码后无提示返回登录表单。',
    'В CampusFix 1.0.3 в Firefox 131 на Windows 11 после ввода пароля в SSO страница возвращается к форме входа без сообщения.',
)
steps = [
    ('在 Windows 11 中使用 Firefox 131 打开 CampusFix 1.0.3 测试服务器的登录页。',
     'В Windows 11 открыть страницу входа тестового сервера CampusFix 1.0.3 в Firefox 131.'),
    ('进入 SSO 登录流程，使用账号 student09 输入密码并尝试登录。',
     'Перейти к входу через SSO, использовать учётную запись student09, ввести пароль и выполнить вход.'),
    ('观察是否打开“我的申请”列表，或无提示返回登录表单。',
     'Проверить, открывается ли список «Мои заявки» или происходит возврат к форме входа без сообщения.'),
]
outcomes = [
    ('实际结果', 'Фактический результат', answers[1]['answer_zh'], answers[1]['answer_ru']),
    ('预期结果', 'Ожидаемый результат', answers[4]['answer_zh'], answers[4]['answer_ru']),
    ('复现与对照', 'Повторяемость и сравнение', answers[3]['answer_zh'], answers[3]['answer_ru']),
]
open_items = [
    ('补充事件年份、时区、测试服务器 URL 和 SSO 入口地址。',
     'Уточнить год события, часовой пояс, URL тестового сервера и адрес входа SSO.'),
    ('提供一次失败登录的截图或录屏，以及对应的请求日志和重定向链。',
     'Приложить снимок или запись неудачного входа, соответствующие журналы запросов и цепочку перенаправлений.'),
    ('确认两个浏览器在相同服务器、网络和账号下的对照结果，并记录尝试次数与失败次数。',
     'Подтвердить сравнение браузеров на одном сервере, в одной сети и с одной учётной записью; указать число попыток и отказов.'),
]

doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Inches(11), Inches(8.5)
sec.top_margin = sec.bottom_margin = Inches(0.48)
sec.left_margin = sec.right_margin = Inches(0.55)
for name in ('Normal', 'Title', 'Heading 1'):
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Arial Unicode MS')
    style.paragraph_format.line_spacing = 1.02
    style.paragraph_format.space_after = Pt(3)
doc.styles['Normal'].font.size = Pt(10.5)
doc.styles['Title'].font.size = Pt(16)
doc.styles['Heading 1'].font.size = Pt(11)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(6)

paired_paragraphs = []


def pair(zh, ru, cell=None, style=None, bold=False, size=10.5):
    if cell is None:
        p = doc.add_paragraph(style=style)
    else:
        p = cell.paragraphs[0] if not cell.text else cell.add_paragraph()
        if style:
            p.style = style
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.keep_together = True
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(zh)
    r.bold, r.font.size = bold, Pt(size)
    q = doc.add_paragraph(style=style) if cell is None else cell.add_paragraph()
    if cell is not None and style:
        q.style = style
    q.paragraph_format.keep_together = True
    q.paragraph_format.keep_with_next = style in ('Title', 'Heading 1')
    q.paragraph_format.space_after = Pt(5)
    r = q.add_run(ru)
    r.bold, r.font.size = bold, Pt(size)
    paired_paragraphs.append((p, q, zh, ru))


def make_table(rows, widths, headings):
    table = doc.add_table(rows=1, cols=len(widths))
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for col, width in zip(table.columns, widths):
        col.width = Inches(width)
    for cell, labels in zip(table.rows[0].cells, headings):
        pair(*labels, cell=cell, bold=True)
        shade = OxmlElement('w:shd')
        shade.set(qn('w:fill'), 'DAEEF0')
        cell._tc.get_or_add_tcPr().append(shade)
    repeat = OxmlElement('w:tblHeader')
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for values in rows:
        cells = table.add_row().cells
        for cell, values_pair in zip(cells, values):
            pair(*values_pair, cell=cell)
    for row in table.rows:
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    props = table._tbl.tblPr
    margins = OxmlElement('w:tblCellMar')
    for side in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement('w:' + side)
        e.set(qn('w:w'), '85')
        e.set(qn('w:type'), 'dxa')
        margins.append(e)
    props.append(margins)
    borders = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + side)
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), '4')
        e.set(qn('w:color'), 'D9D9D9')
        borders.append(e)
    props.append(borders)
    return table


pair('CampusFix BR-1 报告人的补充答复', 'CampusFix Ответ репортёра по BR-1', style='Title', size=16)
pair('课堂任务  |  BR-1  |  v1  |  DRAFT', 'Учебное задание  |  BR-1  |  v1  |  DRAFT', size=9)
pair('登录问题的已知信息包括运行环境、实际表现、发生时间、复现情况和预期结果。',
     'Для проблемы входа известны окружение, фактическое поведение, время события, повторяемость и ожидаемый результат.')
pair('1 报告人问答', '1 Вопросы и ответы репортёра', style='Heading 1', bold=True, size=11)
source_table = make_table(
    [[(x['question_zh'], x['question_ru']), (x['answer_zh'], x['answer_ru'])] for x in answers],
    [2.30, 7.60], [('问题', 'Вопрос'), ('答复', 'Ответ')],
)

doc.add_page_break()
pair('CampusFix BR-1 登录缺陷报告', 'CampusFix BR-1 Отчёт о дефекте входа', style='Title', size=16)
pair('2 问题摘要', '2 Краткое описание', style='Heading 1', bold=True, size=11)
pair(*summary)
pair('环境：CampusFix 1.0.3 测试服务器；Windows 11；Firefox 131；账号 student09。',
     'Окружение: тестовый сервер CampusFix 1.0.3; Windows 11; Firefox 131; учётная запись student09.', size=10)
pair('3 待复核的复现步骤', '3 Шаги воспроизведения для проверки', style='Heading 1', bold=True, size=11)
for n, (zh, ru) in enumerate(steps, 1):
    pair(f'{n}. {zh}', f'{n}. {ru}', size=10)
pair('预期结果：登录成功，并打开“我的申请”列表。',
     'Ожидаемый результат: успешный вход и открытие списка «Мои заявки».', size=10)
pair('4 仍需补充的信息', '4 Данные для уточнения', style='Heading 1', bold=True, size=11)
for n, (zh, ru) in enumerate(open_items, 1):
    pair(f'{n}. {zh}', f'{n}. {ru}', size=9.5)

for element in (doc.styles.element, doc.element):
    for border in list(element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
doc.core_properties.title = 'CampusFix BR-1 报告人答复与登录缺陷报告'
doc.core_properties.subject = '中文在上 俄语在下 逐部分对应'
doc.core_properties.author = 'StudySpace Coursework'
output = ROOT / (STEM + '.docx')
doc.save(output)
shutil.copy2(SOURCE, DATA / 'source.jpg')

content = {
    'version': 'v1', 'document_date': '2026-10-19', 'status': 'DRAFT',
    'case': 'CampusFix', 'report_id': 'BR-1',
    'source_type': 'USER_SUPPLIED_CLASSROOM_IMAGE',
    'source_title_ru': 'Ответ репортёра по BR-1',
    'source_title_zh': '报告人针对 BR-1 的答复',
    'answers': answers, 'summary': summary, 'steps_for_verification': steps,
    'outcomes': outcomes, 'open_items': open_items,
    'year_of_event': None, 'timezone_of_event': None,
    'independently_reproduced': False, 'root_cause': None,
    'language_order': ['zh', 'ru'], 'deliverable_formats': ['docx'],
}
(DATA / 'content.json').write_text(json.dumps(content, ensure_ascii=False, indent=2) + '\n')
coverage = {
    'source': 'source.jpg', 'source_sha256': hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
    'title': 'Ответ репортёра по BR-1', 'headers': ['Вопрос', 'Ответ'],
    'all_five_rows_preserved': True,
    'row_ids': [x['question_ru'] for x in answers],
    'source_table': {'rows_including_header': 6, 'columns': 2},
    'additional_sections': 'Structured report, steps for verification, remaining questions',
    'earlier_studyspace_examples_modified': False,
}
(DATA / 'source-coverage.json').write_text(json.dumps(coverage, ensure_ascii=False, indent=2) + '\n')

with ZipFile(output) as z:
    assert z.testzip() is None
    assert b'<w:pBdr' not in z.read('word/document.xml')
assert len(doc.tables) == 1
assert len(source_table.rows) == 6 and len(source_table.columns) == 2
for p, q, zh, ru in paired_paragraphs:
    assert p.text == zh and q.text == ru
    assert p._p.getnext() is q._p
for row, answer in zip(source_table.rows[1:], answers):
    assert row.cells[0].paragraphs[0].text == answer['question_zh']
    assert row.cells[0].paragraphs[1].text == answer['question_ru']
    assert row.cells[1].paragraphs[0].text == answer['answer_zh']
    assert row.cells[1].paragraphs[1].text == answer['answer_ru']
check = {
    'source_table': 'PASS', 'all_five_answers': 'PASS',
    'zh_then_ru_adjacency': 'PASS', 'bilingual_pairs': len(paired_paragraphs),
    'docx_zip': 'PASS', 'render_review': 'PENDING',
    'execution_status': 'NOT_RUN', 'deliverable_formats': ['docx'],
    'sha256': hashlib.sha256(output.read_bytes()).hexdigest(),
}
(DATA / 'verification.json').write_text(json.dumps(check, ensure_ascii=False, indent=2) + '\n')

font_source = ROOT.parent / '2026-10-19-bug-report-review-v1/work/fonts'
lo_fonts = '/Users/zhi/.cache/codex-runtimes/codex-primary-runtime/dependencies/native/libreoffice-headless/libreoffice/LibreOfficeDev.app/Contents/Resources/fonts/truetype'
(WORK / 'fonts.conf').write_text(f'<?xml version="1.0"?><!DOCTYPE fontconfig SYSTEM "fonts.dtd"><fontconfig><dir>{font_source}</dir><dir>{lo_fonts}</dir><cachedir>{WORK / "fontcache"}</cachedir></fontconfig>')
print(f'PASS: 5 source rows, {len(paired_paragraphs)} adjacent Chinese-Russian pairs, DOCX integrity')
print(output)
