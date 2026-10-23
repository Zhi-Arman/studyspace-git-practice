from pathlib import Path
import json
import hashlib
import shutil
from zipfile import ZipFile
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = Path(__file__).resolve().parents[1]
STEM = '操作说明改写与验收清单_中俄双语_v1'
DOCX = BASE / (STEM + '.docx')
MD = BASE / (STEM + '.md')
source = Path('/var/folders/ff/kqw394c93dd7krl0z2hcqqdh0000gn/T/codex-clipboard-80a95489-7a30-4e34-86cd-c354e45f4698.jpg')
shutil.copy2(source, BASE / 'data' / 'assignment.jpg')

title = ('操作说明改写与验收清单', 'Редактирование инструкций и чек лист приёмки')
intro = ('完成报表打印说明改写、一页软件指南、十项验收清单和故障提示改写。', 'Работа включает инструкцию по печати отчёта, одностраничное руководство, десять критериев приёмки и редактирование сообщения о сбое.')
meta = ('2026-10-22  v1.0  DRAFT  课程作业', '22.10.2026  v1.0  DRAFT  Учебная работа')

task1 = {
 'heading': ('任务一 报表打印操作说明', 'Задание 1 Инструкция по печати отчёта'),
 'scope': ('教学示例：俄文界面名称为约定名称，发布前须核对实际界面。', 'Учебный пример: названия интерфейса условные; перед публикацией сверьте их с приложением.'),
 'goal': ('目标：生成指定报表并打印一份纸质副本。', 'Цель: сформировать выбранный отчёт и напечатать один бумажный экземпляр.'),
 'pre': ('前提：已登录且有报表权限；已确定报表类型与日期范围；打印机已安装、开机、联机，纸张充足。', 'Условия: выполнен вход, есть доступ к отчётам, определены тип отчёта и период. Принтер установлен, включён, доступен и заправлен бумагой.'),
 'steps': [
  ('打开主菜单“Отчёты”（报表）。结果：显示报表列表。', 'Откройте пункт главного меню «Отчёты». Результат: показан список отчётов.'),
  ('在列表中选择所需报表。结果：打开该报表的参数页面，标题显示所选名称。', 'Выберите нужный отчёт из списка. Результат: открыта форма параметров, в заголовке указано имя выбранного отчёта.'),
  ('填写“Дата начала”（开始日期）和“Дата окончания”（结束日期），确认开始日期不晚于结束日期。结果：日期范围正确且无校验错误。', 'Заполните поля «Дата начала» и «Дата окончания»; начальная дата не должна быть позже конечной. Результат: период задан, ошибок проверки нет.'),
  ('单击“Сформировать”（生成）。结果：出现报表预览；核对标题、日期范围和内容，确认这是要打印的报表。', 'Нажмите «Сформировать». Результат: открыт предпросмотр; проверьте название, период и данные отчёта.'),
  ('在预览页面单击“Печать”（打印）。结果：打开打印对话框。', 'В предпросмотре нажмите «Печать». Результат: открыто диалоговое окно печати.'),
  ('在“Принтер”（打印机）中选择目标设备；将“Копии”（份数）设为 1，并核对页码范围、纸张大小和方向。结果：选项与预览一致，内容未超出页面。', 'В поле «Принтер» выберите устройство, в поле «Копии» укажите 1. Проверьте диапазон страниц, формат бумаги и ориентацию. Результат: параметры верны, содержимое не обрезано.'),
  ('在打印对话框中单击“Печать”（打印）。结果：打印任务进入所选设备的队列；取出纸张并核对全部页码和内容。', 'В диалоговом окне нажмите «Печать». Результат: задание поступило в очередь выбранного принтера; заберите листы и проверьте все страницы и данные.'),
 ],
 'result': ('完成判据：所选打印机输出一份报表；标题、日期范围、内容和页数正确，文字与表格无截断。队列中出现任务本身不能证明打印成功。', 'Критерий завершения: выбранный принтер выдал один экземпляр с верными названием, периодом, данными и числом страниц; текст и таблицы не обрезаны. Наличие задания в очереди само по себе не подтверждает печать.'),
 'trouble': [
  ('看不到报表菜单或提示无权限', 'Нет пункта «Отчёты» или отказано в доступе', '确认登录账号；记录缺少的菜单或完整错误信息，并联系系统管理员申请相应权限。', 'Проверьте учётную запись. Запишите отсутствующий пункт или полный текст ошибки и обратитесь к администратору за доступом.'),
  ('报表为空或日期校验失败', 'Отчёт пуст или период не проходит проверку', '核对报表类型和起止日期；确认期间应有数据后重新生成。仍为空时记录参数并联系管理员核实数据。', 'Проверьте тип отчёта и даты, затем сформируйте отчёт снова. Если за период должны быть данные, но их нет, передайте параметры администратору.'),
  ('打印机不在列表或处于离线状态', 'Принтера нет в списке или он недоступен', '检查设备电源、连接和系统打印机状态；恢复联机后重新打开打印对话框，再选择设备。', 'Проверьте питание, подключение и состояние устройства в системе. После восстановления связи откройте диалог печати заново и выберите принтер.'),
  ('已提交任务但没有纸张输出', 'Задание отправлено, но листы не выходят', '检查打印队列、缺纸或卡纸提示。处理已显示的问题；确认原任务失败或已取消后再提交，避免重复打印。', 'Проверьте очередь, сообщения об отсутствии бумаги или замятии. Устраните указанную причину; отправляйте отчёт повторно только после подтверждения сбоя или отмены старого задания.'),
  ('表格被截断或纸张设置错误', 'Таблица обрезана или неверно задана бумага', '返回打印设置，核对纸张大小与方向；若有缩放选项，调整后先检查预览，再打印。', 'Вернитесь к настройкам, проверьте формат бумаги и ориентацию. Если доступно масштабирование, настройте его и проверьте предпросмотр перед печатью.'),
 ],
 'changes': ('改写要点：删去“通常”“相应的地方”等含糊表述，补全操作入口、参数、按钮、每步结果和故障处理路径。', 'Что исправлено: удалены расплывчатые выражения, указаны путь к функции, параметры, кнопки, результат каждого шага и действия при сбое.'),
}

task2 = {
 'heading': ('任务二 VS Code 查找并替换文本', 'Задание 2 Поиск и замена текста в VS Code'),
 'goal': ('用途：在当前文件中把两处 draft 替换为 final，并保存练习文件。', 'Назначение: заменить два вхождения draft на final в текущем файле и сохранить учебный файл.'),
 'pre': ('前提：macOS，VS Code 英文界面、默认快捷键，可写的练习文件夹。仅操作新建练习文件。', 'Условия: macOS, VS Code с английским интерфейсом и стандартными сочетаниями клавиш; доступная для записи учебная папка. Используйте новый учебный файл.'),
 'steps': [
  ('按 Cmd+N，新建文本文件，输入两行：draft report 与 draft checklist。结果：两行各有一个 draft。', 'Нажмите Cmd+N. Введите две строки: draft report и draft checklist. Результат: по одному draft в каждой строке.'),
  ('按 Cmd+Option+F（Option 即 ⌥）。结果：当前编辑器中出现 Find 和 Replace 两个输入框。', 'Нажмите Cmd+Option+F (Option обозначается ⌥). Результат: в текущем редакторе видны поля Find и Replace.'),
  ('在 Find 中输入 draft，关闭 Regular Expression 和 Find in Selection。结果：两处 draft 被匹配。', 'В поле Find введите draft. Отключите Regular Expression и Find in Selection. Результат: найдены два вхождения draft.'),
  ('在 Replace 中输入 final。结果：替换框显示 final，正文尚未改变。', 'В поле Replace введите final. Результат: поле содержит final, исходный текст ещё не изменён.'),
  ('单击 Replace All 图标；可悬停核对工具提示。结果：正文变为 final report 和 final checklist。', 'Нажмите значок Replace All; название можно проверить во всплывающей подсказке. Результат: строки стали final report и final checklist.'),
  ('按 Esc 关闭查找框，再按 Cmd+S；在练习文件夹中以 replace-practice.txt 保存。结果：文件已保存，标签页无未保存标记。', 'Нажмите Esc, затем Cmd+S. Сохраните файл как replace-practice.txt в учебной папке. Результат: файл сохранён, на вкладке нет индикатора изменений.'),
  ('在当前文件中按 Cmd+F，分别查找 draft 和 final。结果：draft 无匹配，final 有两处匹配。', 'В текущем файле нажмите Cmd+F. По очереди найдите draft и final. Результат: draft не найден, final найден дважды.'),
 ],
 'result': ('最终结果：replace-practice.txt 已保存，恰含 final report 与 final checklist 两行。', 'Итог: сохранён replace-practice.txt ровно с двумя строками: final report и final checklist.'),
 'warn': ('警告：Replace All 会修改当前查找范围的全部匹配项。执行前确认目标文件和匹配数量；误替换后立即按 Cmd+Z，再保存。', 'Предупреждение: Replace All изменяет все совпадения в области поиска. Проверьте файл и число совпадений. При ошибке сразу нажмите Cmd+Z, затем сохраните.'),
 'trouble': ('常见问题：无匹配时核对拼写及查找选项；快捷键失效时使用 Edit → Replace；不能保存时换到可写文件夹。', 'Частые проблемы: нет совпадений — проверьте текст и параметры; не работает сочетание клавиш — выберите Edit → Replace; сохранение недоступно — выберите папку с правом записи.'),
 'source': ('依据：VS Code 官方 Basic editing 与 Default keyboard shortcuts，核对日期 2026-10-22。', 'Источники: официальные Basic editing и Default keyboard shortcuts VS Code, проверено 22.10.2026.'),
}

checklist = [
 ('目标与读者明确：说明适用软件、平台和要完成的任务。', 'Указаны читатель, приложение, платформа и цель.'),
 ('前提完整：账号权限、文件、设备及初始状态可确认。', 'Указаны права, файлы, устройства и начальное состояние.'),
 ('步骤已编号，顺序连续且无遗漏。', 'Шаги пронумерованы последовательно, без пропусков.'),
 ('动作具体：写明单击、选择或输入的对象与内容。', 'Действия конкретны: указано, что нажать, выбрать или ввести.'),
 ('操作入口可定位，菜单、按钮和字段名称与目标界面一致。', 'Путь к функции понятен; названия меню, кнопок и полей совпадают с интерфейсом.'),
 ('没有“通常”“大概”“相应的地方”等含糊说法。', 'Нет выражений «как правило», «скорее всего», «в соответствующем разделе».'),
 ('每个步骤给出可观察或可核对的结果。', 'После каждого шага указан наблюдаемый или проверяемый результат.'),
 ('指南末尾给出明确的最终结果和完成判据。', 'Указаны итог и проверяемый критерий завершения.'),
 ('风险提示放在相关操作之前，说明如何避免或撤销错误。', 'Предупреждение дано до рискованного действия; описаны предотвращение или отмена ошибки.'),
 ('有“如果无法正常工作”部分：列出症状、处理动作和求助时提供的信息。', 'Есть раздел «Если что то не работает»: симптомы, действия и сведения для обращения в поддержку.'),
]

task4 = {
 'heading': ('任务四 适合翻译的故障提示', 'Задание 4 Сообщение о сбое для перевода'),
 'original': 'Если программа вдруг зависла, не паникуйте, просто попробуйте перезапустить её, и, скорее всего, всё заработает как надо.',
 'translation': '如果程序突然卡住了，不要惊慌，只要试着重新启动它，多半一切就会正常工作。',
 'problems': [
  ('口语和多余语气：“вдруг”“не паникуйте”“просто”增加语气，却没有补充操作信息。', 'Разговорность и лишние слова: «вдруг», «не паникуйте», «просто» добавляют эмоцию, но не уточняют действие.'),
  ('不确定性：“попробуйте”“скорее всего”没有说明操作后的判断方法。', 'Неопределённость: «попробуйте» и «скорее всего» не задают способ проверки результата.'),
  ('对象和成功状态不明确：“её”依赖上下文，“всё заработает как надо”未定义何为恢复正常。', 'Неясны объект и результат: «её» зависит от контекста, а «всё заработает как надо» не определяет исправное состояние.'),
 ],
 'rewrite': ('如果应用程序不响应用户操作，请关闭并重新打开该应用程序。重新打开后，检查界面是否响应操作。如果问题仍然存在，请联系技术支持，并提供应用程序名称及问题出现前的操作。', 'Если приложение не отвечает на действия пользователя, закройте и снова откройте приложение. После открытия проверьте, реагирует ли интерфейс на действия. Если проблема сохраняется, обратитесь в службу поддержки и укажите название приложения и действия перед сбоем.'),
 'warning': ('重启前提示：关闭应用程序可能导致未保存的数据丢失。若界面仍可操作，请先保存；若无法保存，请在关闭前记录这一情况。', 'Предупреждение перед перезапуском: закрытие приложения может привести к потере несохранённых данных. Если интерфейс доступен, сохраните данные; если сохранить их невозможно, зафиксируйте это перед закрытием.'),
 'reason': ('改写理由：删除口语与不确定的恢复承诺，分开说明操作、检查和求助，明确对象及成功标准。', 'Обоснование: убраны разговорные слова и необоснованное обещание восстановления. Отдельно описаны действие, проверка и обращение в поддержку; объект и критерий успеха указаны явно.'),
}

SOURCES = [
 ('VS Code Basic editing', 'https://code.visualstudio.com/docs/editing/codebasics#_find-and-replace'),
 ('VS Code Default keyboard shortcuts', 'https://code.visualstudio.com/docs/reference/default-keybindings'),
]

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.6)
sec.left_margin = sec.right_margin = Inches(.7)
for sty in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
 s = doc.styles[sty]
 s.font.name = 'Arial'
 s.font.color.rgb = RGBColor(0, 0, 0)
 rf = s.element.get_or_add_rPr().find(qn('w:rFonts'))
 if rf is None:
  rf = OxmlElement('w:rFonts'); s.element.get_or_add_rPr().insert(0, rf)
 for k in ['ascii', 'hAnsi']: rf.set(qn('w:'+k), 'Arial')
 rf.set(qn('w:eastAsia'), 'Arial Unicode MS')
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.line_spacing = 1.0
doc.styles['Normal'].paragraph_format.space_after = Pt(3)
for sty, size in [('Title',18),('Subtitle',12),('Heading 1',14),('Heading 2',11.5)]:
 doc.styles[sty].font.size = Pt(size)
 doc.styles[sty].paragraph_format.space_after = Pt(4)
 doc.styles[sty].paragraph_format.space_before = Pt(7 if 'Heading' in sty else 0)

md = []
all_pairs = []
next_page = False
def paragraph(text, bold=False, keep=False, after=3, size=None):
 p = doc.add_paragraph()
 p.paragraph_format.keep_with_next = keep
 p.paragraph_format.keep_together = True
 p.paragraph_format.space_after = Pt(after)
 r=p.add_run(text); r.bold=bold
 if size: r.font.size=Pt(size)
 return p

def pair(value, number=None, bold=False, size=None, gap=4):
 all_pairs.append(value)
 prefix = str(number)+'. ' if number else ''
 paragraph(prefix+value[0], bold, True, 1, size)
 paragraph(value[1], bold, False, gap, size)
 md.extend([prefix+value[0], '', value[1], ''])

def heading(value, level=1):
 global next_page
 all_pairs.append(value)
 hp = doc.add_paragraph(value[0], 'Heading '+str(level))
 if next_page:
  hp.paragraph_format.page_break_before = True
  next_page = False
 paragraph(value[1], bold=True, keep=True, after=5, size=11.5 if level==1 else 11)
 md.extend(['#'*(level+1)+' '+value[0], '', value[1], ''])

def page():
 global next_page
 next_page = True

doc.add_paragraph(title[0], 'Title')
doc.add_paragraph(title[1], 'Subtitle')
md.extend(['# '+title[0], '', title[1], ''])
pair(meta, size=9, gap=4)
heading(task1['heading'])
for k in ['scope','goal','pre']: pair(task1[k],gap=4)
heading(('操作步骤', 'Порядок действий'),2)
for i, s in enumerate(task1['steps'],1): pair(s,i,gap=5)

page()
heading(('任务一 打印结果与故障处理', 'Задание 1 Проверка печати и устранение неполадок'))
pair(task1['result'])
heading(('如果无法正常工作', 'Если что то не работает'),2)
for i,(zh,ru,act_zh,act_ru) in enumerate(task1['trouble'],1):
 pair((zh,ru),i,bold=True,gap=2)
 pair((act_zh,act_ru),gap=7)
pair(task1['changes'])

page()
heading(task2['heading'])
for k in ['goal','pre']: pair(task2[k],gap=4)
pair(task2['warn'],gap=5)
for i,s in enumerate(task2['steps'],1): pair(s,i,gap=4)
pair(task2['result'],bold=True,gap=4)
heading(('如果无法正常工作', 'Если что то не работает'),2)
pair(task2['trouble'],gap=4)

page()
heading(('任务三 操作指南验收清单', 'Задание 3 Чек лист приёмки руководства'))
pair(('验收方法：逐条核对指南；满足条件标记“是”，否则标记“否”并记录页码与修改内容。任一项为“否”时先修订，再复核。', 'Порядок: проверьте каждый пункт. Отметьте «да» либо «нет»; при ответе «нет» укажите страницу и исправление. После правок повторите проверку.'))
for i,s in enumerate(checklist,1):
 pair((f'{s[0]}  □ 是  □ 否', f'{s[1]}  □ да  □ нет'), i, gap=10)
pair(('记录栏：检查人________  日期________  不符合项及修改记录________', 'Запись проверки: проверяющий________  дата________  замечания и исправления________'),size=10)

page()
heading(task4['heading'])
heading(('原句及中文翻译', 'Исходное предложение и перевод'),2)
pair((task4['translation'], '«'+task4['original']+'»'))
heading(('原句存在的问题', 'Проблемы исходного предложения'),2)
for i,s in enumerate(task4['problems'],1): pair(s,i,gap=7)
heading(('改写后的提示', 'Отредактированное сообщение'),2)
pair(task4['warning'],gap=6)
pair(task4['rewrite'],gap=7)
pair(task4['reason'],gap=4)
heading(('参考资料', 'Источники'),2)
for name,url in SOURCES:
 paragraph(name+'\n'+url,size=9,after=3)
 md.extend([f'- [{name}]({url})',''])

for border in list(doc.styles.element.xpath('.//w:pBdr')) + list(doc.element.xpath('.//w:pBdr')):
 border.getparent().remove(border)
for grid in list(doc.element.xpath('.//w:docGrid')):
 grid.getparent().remove(grid)
for p in doc.paragraphs:
 el=OxmlElement('w:snapToGrid');el.set(qn('w:val'),'0');p._p.get_or_add_pPr().append(el)
doc.core_properties.title=title[0]
doc.core_properties.subject='四项操作说明课程任务 / Четыре задания по пользовательским инструкциям'
doc.core_properties.author=''
doc.core_properties.keywords='用户指南, 验收清单, 中文, русский'
doc.save(DOCX)
MD.write_text('\n'.join(md)+'\n',encoding='utf-8')
(BASE/'data'/'content.json').write_text(json.dumps({'task1':task1,'task2':task2,'checklist':checklist,'task4':task4},ensure_ascii=False,indent=2)+'\n')
all_text='\n'.join(p.text for p in Document(DOCX).paragraphs)
md_text=MD.read_text()
checks={
 'four_tasks_present':all(x in all_text for x in ['任务一','任务二','任务三','任务四']),
 'printing_steps_each_has_result':len(task1['steps'])==7 and all('结果' in z and 'Результат' in r for z,r in task1['steps']),
 'printing_troubleshooting':len(task1['trouble'])==5,
 'guide_step_count_5_to_8':5<=len(task2['steps'])<=8,
 'guide_all_required_sections':all(task2[k] for k in ['goal','pre','result','warn','trouble']),
 'guide_steps_each_has_result':all('结果' in z and 'Результат' in r for z,r in task2['steps']),
 'checklist_exactly_10':len(checklist)==10,
 'translation_problems_and_rewrite':len(task4['problems'])==3 and bool(task4['rewrite']),
 'all_bilingual_content_in_both_outputs':all(v in all_text and v in md_text for pair_ in all_pairs for v in pair_),
 'sample_ui_explicit': 'условные' in all_text,
 'no_ambiguous_words_in_new_guides':not any(w in '\n'.join(r for z,r in task1['steps']+task2['steps']) for w in ['как правило','скорее всего']),
 'editable_text_no_document_images':len(doc.inline_shapes)==0,
}
with ZipFile(DOCX) as z: checks['docx_zip_integrity']=z.testzip() is None
assert all(checks.values()),checks
report={'date':'2026-10-22','status':'DRAFT','checks':checks,'all_structural_checks_pass':all(checks.values()),'visual_review':'PENDING','guide_one_page':'PENDING','printing_execution':'NOT_RUN','vscode_gui_execution':'NOT_RUN','novice_usability_trial':'NOT_RUN','source_image_sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'outputs':{p.name:{'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in [DOCX,MD]}}
(BASE/'work'/'verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
print(json.dumps({'outputs':[str(DOCX),str(MD)],'checks':checks},ensure_ascii=False,indent=2))
