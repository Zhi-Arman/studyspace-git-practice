from pathlib import Path
from datetime import datetime, timezone
from zipfile import ZipFile
import hashlib
import json
import re

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

ROOT=Path('/Users/zhi/Documents/软件生命周期/studyspace')
OUT=ROOT/'coursework/2026-10-15-deployment-checklist-v1'
DATA=OUT/'data'
DATA.mkdir(parents=True,exist_ok=True)
STEM='StudySpace_部署图检查表_中俄双语_v1'
DOCX=OUT/f'{STEM}.docx'
sources={
 'S1':ROOT/'coursework/2026-10-15-deployment-nodes-v1/StudySpace_节点与制品表_中俄双语_v1.docx',
 'S2':ROOT/'coursework/2026-10-15-environments-protection-v1/StudySpace_环境与防护表_中俄双语_v1.docx',
 'S3':ROOT.parent/'StudySpace_组件交付/StudySpace_组件候选与组件图_中俄双语_v2.md',
 'S4':ROOT.parent/'StudySpace_组件交付/StudySpace_组件图_课堂任务2_中俄双语.png',
 'S5':ROOT/'docs/decisions/zh-CN/0004-docker-https.md',
 'S6':ROOT/'docs/decisions/ru-RU/0004-docker-https.md',
 'S7':ROOT/'compose.yaml',
 'S8':ROOT/'frontend/nginx.conf',
 'S9':ROOT/'.env.example',
 'S10':ROOT/'docs/data/project-decisions.json',
 'S11':ROOT/'AGENTS.md',
 'S12':ROOT.parent/'StudySpace_组件交付/StudySpace_组件图_课堂任务2_中俄双语.drawio',
}
fingerprints={str(p):hashlib.sha256(p.read_bytes()).hexdigest() for p in sources.values()}
headers=[('检查要求','Требование проверки'),('StudySpace 填写内容','Заполнение для StudySpace'),('依据与核对状态','Основания и статус сверки')]
rows=[
 [
  ('1 每个组件放置在节点上','Каждый компонент размещён на узле'),
  ('CMP-001 在 N-01 浏览器执行，其静态制品由 N-03 Nginx 分发；CMP-002 至 CMP-009 位于 N-04 FastAPI 单体。N-05 为 Caddy 入口，N-06 PostgreSQL 存放数据，SQLAlchemy 访问层 CMP-009 仍在 N-04。CMP-010 对应外部依赖 N-07；N-02 是本地开发方式。',
   'CMP-001 исполняется в браузере N-01, статика раздаётся Nginx N-03. CMP-002–CMP-009 находятся в монолите FastAPI N-04. N-05 — вход Caddy, N-06 PostgreSQL хранит данные; слой SQLAlchemy CMP-009 остаётся в N-04. CMP-010 соответствует внешней зависимости N-07; N-02 — локальная разработка.'),
  ('S1 S3 S4\n映射已填写\n待图面核对','S1 S3 S4\nСвязь задана\nСверка ожидается'),
 ],
 [
  ('2 标注关键连接的协议和端口','Подписаны протоколы и порты ключевых связей'),
  ('目标连接：浏览器 → Caddy 为 HTTPS/TCP 443；Caddy → Nginx 为 HTTP/TCP 80；Nginx 的 /api/ → FastAPI 为 HTTP/JSON/TCP 8000；后端 → PostgreSQL 为 PostgreSQL/TCP 5432，5432 沿用连接示例，最终绑定待生产配置确认。单体内部调用不新增网络端口；本地开发为 5173/8000，本地容器入口为 8080。',
   'Целевые связи: браузер → Caddy HTTPS/TCP 443; Caddy → Nginx HTTP/TCP 80; /api/ Nginx → FastAPI HTTP/JSON/TCP 8000; сервер → PostgreSQL PostgreSQL/TCP 5432. Порт 5432 взят из примера, привязки уточняются в конфигурации размещения. Вызовы внутри монолита не требуют портов. Локально: 5173/8000, вход контейнеров 8080.'),
  ('S1 S5 S7\nS8 S9\n端口方案已填写\n待图面核对','S1 S5 S7\nS8 S9\nПорты заданы\nСверка ожидается'),
 ],
 [
  ('3 显示数据库不可越过的网络边界','Показана граница, за которую нельзя выставлять БД'),
  ('目标图应区分公网与服务器内部 Docker 网络。Caddy 是公开 Web 入口；PostgreSQL N-06 及其数据卷放在内部边界内，仅由后端访问，不向公网发布数据库端口。图中不应出现浏览器直连数据库；Nginx 与 FastAPI 的目标生产绑定随部署配置确认。',
   'Целевая схема разделяет публичную сеть и внутреннюю сеть Docker на сервере. Caddy — публичный Web-вход; PostgreSQL N-06 и том данных находятся внутри, доступ к БД имеет сервер приложения без публикации порта в Интернет. Прямой связи браузера с БД нет; производственные привязки Nginx и FastAPI уточняются при размещении.'),
  ('S1 S2 S5\n边界方案已填写\n待图面核对','S1 S2 S5\nГраница задана\nСверка ожидается'),
 ],
 [
  ('4 注明 TLS 终止位置','Указано, где завершается TLS'),
  ('目标 TLS 终止点为 N-05 Caddy：在入口标注 HTTPS 443 和 TLS termination；证书及其持久存储由 Caddy 管理。到 Nginx 的内部段为 HTTP 80，再由 Nginx 转发 /api/ 到 HTTP 8000。图例须说明该方案不是全链路 TLS；当前本地 HTTP 配置尚无 TLS。',
   'Целевая точка завершения TLS — Caddy N-05: у входа обозначаются HTTPS 443 и TLS termination; сертификаты и постоянное хранилище обслуживает Caddy. Далее идут HTTP 80 к Nginx и /api/ на HTTP 8000. Легенда уточняет отсутствие TLS на внутренних сегментах; текущая локальная HTTP-конфигурация без TLS.'),
  ('S1 S5 S8\n终止点已填写\n待图面核对','S1 S5 S8\nТочка задана\nСверка ожидается'),
 ],
 [
  ('5 外部系统与组件图保持一致','Внешние системы отражены так же, как на диаграмме компонентов'),
  ('沿用 CMP-010 可信学校核验来源，在 StudySpace 系统边界外标为 N-07。与 CMP-003 的 IF-04 保持人工核验、来源引用的虚线依赖；不改为自动 SSO 或已有 API。学校确认的名单或平台仍由 OPEN-003 确定，承载形式及端口待确认。',
   'Сохранить CMP-010 «Доверенный источник вуза» как N-07 вне границы StudySpace. IF-04 с CMP-003 остаётся пунктирной зависимостью ручной проверки и ссылки на источник, без замены на автоматический SSO или готовый API. Реестр или платформа определяются в OPEN-003; форма размещения и порты уточняются.'),
  ('S1 S3 S4\n外部依赖一致\n待图面核对','S1 S3 S4\nИсточник согласован\nСверка ожидается'),
 ],
 [
  ('6 提供图名和说明','Есть название и пояснение'),
  ('图名：StudySpace 目标部署图。说明：React 浏览器 → Caddy HTTPS 入口 → Nginx 静态服务及 API 转发 → FastAPI 单体 → 内网 PostgreSQL；学校来源用人工核验虚线连接。图例解释网络边界、协议和目标状态，并单独注明开发 SQLite 路径。',
   'Название: StudySpace Диаграмма целевого развёртывания. Пояснение: браузер React → HTTPS-вход Caddy → Nginx → монолит FastAPI → внутренняя PostgreSQL; источник вуза связан пунктиром ручной проверки. Легенда поясняет границы сети, протоколы, целевой статус и отдельный путь SQLite разработки.'),
  ('S1 S2 S3\n图名说明已填写\n待图面核对','S1 S2 S3\nНазвание задано\nСверка ожидается'),
 ],
]
intro=(
 '六项检查内容已根据 StudySpace 架构填写。组件与节点沿用 CMP 和 N 编号；“已填写”表示设计内容已补全，图面核对状态为待核对，实际部署仍为 NOT_EXECUTED。',
 'Шесть пунктов заполнены по архитектуре StudySpace с сохранением ID CMP и N. «Заполнено» означает подготовленное содержание проекта; сверка изображения ожидается, фактическое размещение остаётся NOT_EXECUTED.'
)
doc=Document()
sec=doc.sections[0]
sec.page_width,sec.page_height=Inches(8.5),Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.55)
sec.left_margin=sec.right_margin=Inches(.6)
for name in ['Normal','Title','Heading 1']:
 st=doc.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
 rf=OxmlElement('w:rFonts');rf.set(qn('w:ascii'),'Arial');rf.set(qn('w:hAnsi'),'Arial');rf.set(qn('w:eastAsia'),'PingFang SC');st.element.get_or_add_rPr().insert(0,rf)
for st in doc.styles:
 for b in list(st.element.iter(qn('w:pBdr'))):b.getparent().remove(b)
doc.styles['Normal'].font.size=Pt(10.5)
doc.styles['Normal'].paragraph_format.space_after=Pt(4)
doc.styles['Normal'].paragraph_format.line_spacing=1.05
doc.styles['Title'].font.size=Pt(18)
doc.styles['Title'].paragraph_format.space_after=Pt(4)
doc.styles['Heading 1'].font.size=Pt(12)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(8)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(4)

def addrun(p,text,size=10.5,bold=False,color='000000',lang='zh-CN'):
 r=p.add_run(text);r.font.name='Arial';r.font.size=Pt(size);r.bold=bold;r.font.color.rgb=RGBColor.from_string(color)
 rp=r._element.get_or_add_rPr();rp.rFonts.set(qn('w:eastAsia'),'PingFang SC');rp.rFonts.set(qn('w:hAnsi'),'Arial')
 le=OxmlElement('w:lang');le.set(qn('w:val'),lang);le.set(qn('w:eastAsia'),'zh-CN');rp.append(le)
 return r

def para(text,size=10.5,style=None,lang='zh-CN'):
 p=doc.add_paragraph(style=style);addrun(p,text,size,lang=lang);return p

def pair(zh,ru,style=None,size=10.5):
 p=para(zh,size,style);p.paragraph_format.keep_with_next=True
 p=para(ru,size,style,'ru-RU');p.paragraph_format.keep_together=True;return p

pair('StudySpace 部署图检查表','StudySpace Проверка диаграммы развёртывания','Title',18)
para('SS-DEP-CHK-01  |  v1  |  2026-10-15  |  DRAFT',9)
pair(*intro)
table=doc.add_table(rows=1,cols=3);table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
widths=[1.65,4.2,1.45]
for col,w in zip(table.columns,widths):col.width=Inches(w)
borders=OxmlElement('w:tblBorders')
for edge in ['top','left','bottom','right','insideH','insideV']:
 el=OxmlElement('w:'+edge)
 for k,v in [('val','single'),('sz','5'),('color','D9D9D9')]:el.set(qn('w:'+k),v)
 borders.append(el)
table._tbl.tblPr.append(borders)
repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)

def fill(cell,texts,ci,header=False):
 cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
 mar=OxmlElement('w:tcMar')
 for edge,w in [('top',95),('bottom',95),('left',80),('right',80)]:
  el=OxmlElement('w:'+edge);el.set(qn('w:w'),str(w));el.set(qn('w:type'),'dxa');mar.append(el)
 cell._tc.get_or_add_tcPr().append(mar)
 for j,text in enumerate(texts):
  p=cell.paragraphs[0] if j==0 else cell.add_paragraph()
  p.paragraph_format.space_after=Pt(4 if j==0 else 0);p.paragraph_format.line_spacing=1.04;p.paragraph_format.keep_with_next=False
  p.alignment=WD_ALIGN_PARAGRAPH.CENTER if header or ci==2 else WD_ALIGN_PARAGRAPH.LEFT
  addrun(p,text,9.5,header,'FFFFFF' if header else '000000','zh-CN' if j==0 else 'ru-RU')
for ci,h in enumerate(headers):
 c=table.rows[0].cells[ci];fill(c,h,ci,True)
 sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'29485D');c._tc.get_or_add_tcPr().append(sh)
for ri,vals in enumerate(rows):
 row=table.add_row();ns=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(ns)
 for ci,val in enumerate(vals):
  c=row.cells[ci];fill(c,val,ci)
  sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F0F4F7' if ri%2==0 else 'FFFFFF');c._tc.get_or_add_tcPr().append(sh)

pair('检查结论','Итог проверки','Heading 1',12)
pair('六项设计内容均已填写，图面验收仍待逐项核对。当前 Compose 为本地 HTTP 候选，生产方案未部署；OPEN-003、OPEN-010 待确认，业务用例保持 NOT_RUN。',
     'Шесть пунктов заполнены, сверка изображения ожидается. Compose — локальный HTTP-вариант, продакшен не развёрнут; OPEN-003 и OPEN-010 уточняются, бизнес-кейсы остаются NOT_RUN.',size=9.5)
para('核对依据 Основания сверки',12,'Heading 1')
para('S1 节点与制品 / Узлы SS-DEP-T1-01; S2 环境与防护 / Среды SS-ENV-T1B-01;\nS3 组件表 / Компоненты SS-ARCH-03; S4 组件图 / Диаграмма компонентов SS-ARCH-02.',9)
para('S5 S6 中俄 / Двуязычная ADR-0004; S7 compose.yaml; S8 frontend/nginx.conf;\nS9 .env.example; S10 docs/data/project-decisions.json.',9)
para('项目目录 / Каталог проекта\n/Users/zhi/Documents/软件生命周期/studyspace',9)
doc.core_properties.title='StudySpace 部署图检查表 Проверка диаграммы развёртывания'
doc.core_properties.subject='部署图六项检查 / Шесть проверок диаграммы развёртывания'
doc.core_properties.author='StudySpace'
doc.save(DOCX)

content={'doc_id':'SS-DEP-CHK-01','version':'v1','date':'2026-10-15','status':'DRAFT','diagram_review':'PENDING','deployment_status':'NOT_EXECUTED','intro':intro,'headers':headers,'rows':rows,'sources':{k:str(p) for k,p in sources.items()},'source_fingerprints':fingerprints}
(DATA/'content.json').write_text(json.dumps(content,ensure_ascii=False,indent=2),encoding='utf-8')
fontconfig=(ROOT/'coursework/2026-10-15-environments-protection-v1/data/fonts.conf').read_text().replace('studyspace-env-fontcache','studyspace-checklist-fontcache')
(DATA/'fonts.conf').write_text(fontconfig,encoding='utf-8')
md=['# StudySpace 部署图检查表','# StudySpace Проверка диаграммы развёртывания','SS-DEP-CHK-01 · v1 · 2026-10-15 · DRAFT',*intro,'| '+' | '.join('<br>'.join(v) for v in headers)+' |','| --- | --- | --- |']
md+=['| '+' | '.join('<br>'.join(v).replace('\n','<br>') for v in row)+' |' for row in rows]
(OUT/f'{STEM}.md').write_text('\n\n'.join(md),encoding='utf-8')
(OUT/'README.md').write_text(f'''# StudySpace 部署图检查表 Проверка диаграммы развёртывания

SS-DEP-CHK-01 · v1 · 2026-10-15 · DRAFT

交付 / Результат [{STEM}.docx]({STEM}.docx)

六项检查内容已填写，图面核对为 PENDING，实际部署 NOT_EXECUTED。中俄内容在同一可编辑单元格内上下对应。
Шесть пунктов заполнены, сверка изображения PENDING, размещение NOT_EXECUTED. Китайский и русский идут попарно в редактируемых ячейках.

内容与来源 / Содержимое и источники data/content.json
文档验证 / Проверка документа data/verification.json
''',encoding='utf-8')
rd=Document(DOCX)
with ZipFile(DOCX) as z:
 xml=etree.fromstring(z.read('word/document.xml'));sty=etree.fromstring(z.read('word/styles.xml'));names=z.namelist()
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
txt='\n'.join(p.text for p in rd.paragraphs)+'\n'+'\n'.join(c.text for r in rd.tables[0].rows for c in r.cells)
tests={
 'six_checks_three_columns':len(rd.tables)==1 and len(rd.tables[0].rows)==7 and len(rd.tables[0].columns)==3,
 'all_bilingual_cells_match':all(rd.tables[0].cell(ri,ci).text=='\n'.join(v) for ri,vs in enumerate([headers,*rows]) for ci,v in enumerate(vs)),
 'chinese_above_russian':all(len(c.paragraphs)==2 and re.search('[\u4e00-\u9fff]',c.paragraphs[0].text) and re.search('[А-Яа-яЁё]',c.paragraphs[1].text) for r in rd.tables[0].rows for c in r.cells),
 'six_original_requirements_present':all(row[0][1] in txt for row in rows),
 'stable_node_ids_retained':all(f'N-{n:02d}' in txt for n in range(1,8)),
 'component_mapping_retained':all(v in txt for v in ['CMP-001','CMP-002','CMP-003','CMP-009','CMP-010','IF-04']),
 'target_protocols_and_ports_present':all(v in txt for v in ['HTTPS/TCP 443','HTTP/TCP 80','HTTP/JSON/TCP 8000','PostgreSQL/TCP 5432']),
 'tls_termination_at_caddy':'TLS termination' in txt and '目标 TLS 终止点为 N-05 Caddy' in txt,
 'database_boundary_explicit':'不向公网发布数据库端口' in txt,
 'manual_school_verification_preserved':'人工核验' in txt and 'OPEN-003' in txt and 'SSO' in txt,
 'graphic_review_not_falsely_passed':'待图面核对' in txt and '图面验收仍待逐项核对' in txt and 'PASS' not in txt,
 'deployment_and_business_not_claimed':all(v in txt for v in ['NOT_EXECUTED','NOT_RUN','OPEN-010']),
 'repeating_header':bool(xml.xpath('//w:tbl/w:tr[1]/w:trPr/w:tblHeader',namespaces=ns)),
 'rows_not_split':len(xml.xpath('//w:tbl/w:tr/w:trPr/w:cantSplit',namespaces=ns))==6,
 'no_fixed_heights_or_screenshots':not xml.xpath('//w:trHeight|//w:drawing',namespaces=ns),
 'no_title_border':not sty.xpath('//w:style[@w:styleId="Title"]//w:pBdr',namespaces=ns),
 'no_headers_footers':not any(re.fullmatch(r'word/(header|footer)\d+\.xml',n) for n in names),
 'existing_files_preserved':all(hashlib.sha256(Path(p).read_bytes()).hexdigest()==h for p,h in fingerprints.items()),
}
report={'document':str(DOCX),'execution_environment':'Local macOS; bundled Codex Python; document verification only','checked_at_utc':datetime.now(timezone.utc).isoformat(),'checks':tests,'all_pass':all(tests.values()),'visual_review':'PENDING','diagram_review':'PENDING','business_tests':'NOT_RUN','source_fingerprints':fingerprints,'docx_sha256':hashlib.sha256(DOCX.read_bytes()).hexdigest()}
(DATA/'verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'document':str(DOCX),'checks':len(tests),'all_pass':report['all_pass'],'failed':[k for k,v in tests.items() if not v]},ensure_ascii=False))
if not report['all_pass']:raise SystemExit(1)
