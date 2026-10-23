from pathlib import Path
from datetime import datetime, timezone
from zipfile import ZipFile
from html import escape
import ast
import hashlib
import json
import math
import re
import shutil
import xml.etree.ElementTree as ET
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path('/Users/zhi/Documents/软件生命周期/studyspace')
OUT=ROOT/'coursework/2026-10-15-sequence-review-v2'
DATA=OUT/'data'
WORK=OUT/'work'
FIG=OUT/'diagrams'
for p in [OUT,DATA,WORK,FIG]:p.mkdir(parents=True,exist_ok=True)
PREV=ROOT/'coursework/2026-10-15-sequence-diagrams-v1'
MSG=ROOT/'coursework/2026-10-15-message-tables-v1'
source=json.loads((PREV/'data/content.json').read_text())
STEM='StudySpace_时序图检查与精简版_中俄双语_v2'
DOCX=OUT/(STEM+'.docx')
DRAWIO=OUT/'StudySpace_双场景精简时序图_中俄双语_v2.drawio'
paths=[PREV/'data/content.json',PREV/'work/build_document.py',PREV/'StudySpace_双场景时序图与说明_中俄双语_v1.docx',PREV/'StudySpace_双场景时序图_中俄双语_v1.drawio',MSG/'StudySpace_任务三消息表_中俄双语_v1.docx',ROOT.parent/'StudySpace_组件交付/StudySpace_组件候选与组件图_中俄双语_v2.md',ROOT/'docs/zh-CN/04_System_Design.md',ROOT/'docs/ru-RU/04_System_Design.md',ROOT/'docs/data/project-decisions.json',ROOT/'coursework/2026-10-13-requirements-v0.2/data/requirements-practice.json',ROOT/'AGENTS.md']
fingerprints={str(p):hashlib.sha256(p.read_bytes()).hexdigest() for p in paths}
shutil.copy2('/var/folders/ff/kqw394c93dd7krl0z2hcqqdh0000gn/T/codex-clipboard-7c8f5582-fe7c-45df-90e9-9f3120162bfd.jpg',DATA/'classroom-sequence-check.jpg')

# Reuse only the existing diagram primitive classes; no previous builder code executes.
tree=ast.parse((PREV/'work/build_document.py').read_text())
classes=[n for n in tree.body if isinstance(n,ast.ClassDef) and n.name in {'Canvas','Sequence'}]
exec(compile(ast.Module(body=classes,type_ignores=[]),'<existing diagram primitives>','exec'),globals())

b=Sequence('01_US002_compact',[
    ('学生','学生','Студент'),('CMP-001','界面','Интерфейс'),
    ('CMP-006','预约事务','Бронирование'),('CMP-009','数据访问','Данные')])
b.note('一个目标 创建一条单座位预约；业务接口为拟定契约','Одна цель: создать одну бронь места; API проектируется')
b.msg('1','学生','CMP-001','选择 seat_id、业务日 D、start/end','Выбрать место, дату и интервал')
b.msg('2','CMP-001','CMP-006','POST /api/v1/bookings + 会话','Передать заявку и сессию')
b.note('CMP-006 协调 CMP-002 校验会话；细节沿用详细版','CMP-006 проверяет сессию через CMP-002; детали в v1')
b.msg('3','CMP-006','CMP-009','开始 tx；加锁并读取当前决策数据','Начать tx, блокировать и прочитать текущие данные')
b.msg('4','CMP-009','CMP-006','返回 tx、锁后 now、资格规则与重叠数据','Вернуть tx, время после блокировок, данные допуска и пересечений',True)
b.note('同一 tx 内由 CMP-006 协调 CMP-003、004 完成业务校验','CMP-006 с CMP-003 и CMP-004 проверяет правила в той же tx')
f=b.frame('[资格或时间规则失败，或学生／座位区间冲突]','[Нет допуска, нарушены правила времени либо есть пересечение]')
b.note('失败先回滚 tx；回滚内部交互见详细版','Сначала откатить tx; внутренние вызовы приведены в v1')
b.msg('5a','CMP-006','CMP-001','403 / 422 / 409 + 原因码','Вернуть отказ и код причины',True)
b.msg('6a','CMP-001','学生','显示拒绝原因或建议重新选择','Показать причину либо предложить другой выбор',True)
b.branch('[资格与规则通过，且无学生／座位区间冲突]','[Допуск и правила соблюдены, пересечений нет]')
b.msg('5b','CMP-006','CMP-009','持久化一条预约与最少事件，COMMIT','Сохранить одну бронь и минимальное событие, COMMIT')
b.msg('6b','CMP-009','CMP-006','提交成功 + booking_id','Вернуть успех фиксации и booking_id',True)
b.msg('7b','CMP-006','CMP-001','201 Created + booking_id、seat、start/end','Вернуть созданную запись',True)
b.msg('8b','CMP-001','学生','显示预约确认','Показать подтверждение',True)
b.end_frame(f)

r=Sequence('02_US004_compact',[
    ('管理员','管理员','Администратор'),('CMP-001','界面','Интерфейс'),
    ('CMP-003','校内审核','Проверка вуза'),('CMP-009','数据访问','Данные'),
    ('CMP-010','学校来源 外部','Источник вуза')])
r.note('前提 已打开启用的 PENDING 账号并取得身份状态快照','Условие: открыт активный аккаунт PENDING и получен снимок данных')
r.msg('1','管理员','CMP-010','人工核验 IF-04 学校来源','Проверить источник IF-04 вручную')
r.msg('2','CMP-010','管理员','返回人工核验结果','Вернуть результат ручной сверки',True)
r.msg('3','管理员','CMP-001','填写决定、来源和 revision','Внести решение, источник и revision')
r.msg('4','CMP-001','CMP-003','POST /api/v1/identity-reviews','Передать заявку на сохранение проверки')
r.note('CMP-003 通过 CMP-002 再次授权；来源由人工核对','CMP-003 повторно авторизует через CMP-002; источник проверяет человек')
r.msg('5','CMP-003','CMP-009','开始 tx；锁定并读取账号快照','Начать tx и прочитать аккаунты под блокировкой')
r.msg('6','CMP-009','CMP-003','返回当前权限、状态和 revision','Вернуть права, статус и revision',True)
f=r.frame('[无权限、来源不可核验，或身份／状态快照已变化]','[Нет прав, источник не проверен либо данные или статус изменились]')
r.note('失败先回滚；没有可信来源时保持 PENDING','При отказе откат; без доверенного источника сохраняется PENDING')
r.msg('7a','CMP-003','CMP-001','403 / 422 / 409 + 原因码','Вернуть ошибку или конфликт',True)
r.branch('[权限与来源有效，且身份／状态快照一致]','[Права и источник действуют, данные и статус совпадают]')
r.msg('7b','CMP-003','CMP-009','保存审核决定和最少事件，COMMIT','Сохранить решение и минимальное событие, COMMIT')
r.msg('8b','CMP-009','CMP-003','提交成功 + VERIFIED / REJECTED','Вернуть сохранённый статус',True)
r.msg('9b','CMP-003','CMP-001','200 OK + 已保存的审核状态','Вернуть статус в интерфейс',True)
r.end_frame(f)

text_bounds=[]
for p in [b,r]:
    # The smaller number of participants allows larger labels at print size.
    for it in p.c.items:
        if it['kind']=='text':
            if it['size']==27:it['size']=30
            elif it['size']==25:it['size']=28
            chinese=bool(re.search('[\u3000-\u303f\u4e00-\u9fff]',it['text']))
            fp='/System/Library/Fonts/STHeiti Light.ttc' if chinese else '/System/Library/Fonts/Supplemental/Arial.ttf'
            font=ImageFont.truetype(fp,it['size'])
            tw=font.getlength(it['text'])
            while tw>2150:
                it['size']-=1;font=ImageFont.truetype(fp,it['size']);tw=font.getlength(it['text'])
            if it['align']=='center':it['x']=max(tw/2+25,min(2175-tw/2,it['x']))
            left=it['x']-tw/2 if it['align']=='center' else it['x']
            text_bounds.append((left,left+tw))
    p.finish()

def export_drawio():
    mxfile=ET.Element('mxfile',host='app.diagrams.net',type='device',version='29.6.6')
    for i,(title,p) in enumerate([('US-002 单座位预约 Бронирование',b),('US-004 校内审核 Проверка вуза',r)],1):
        diagram=ET.SubElement(mxfile,'diagram',id=f'compact-{i}',name=title)
        model=ET.SubElement(diagram,'mxGraphModel',dx='2200',dy=str(p.c.height+70),grid='1',gridSize='10',page='1',pageScale='1',pageWidth='2200',pageHeight=str(p.c.height+70))
        root=ET.SubElement(model,'root');ET.SubElement(root,'mxCell',id='0');ET.SubElement(root,'mxCell',id='1',parent='0')
        for j,it in enumerate(p.c.items,2):
            attrs={'id':str(j),'parent':'1'}
            if it['kind']=='line':
                attrs.update(edge='1',style=f'edgeStyle=none;rounded=0;strokeWidth=2;strokeColor={it["color"]};endArrow={"classic" if it["arrow"] else "none"};'+('dashed=1;dashPattern=6 5;' if it['dash'] else ''))
                cell=ET.SubElement(root,'mxCell',**attrs);g=ET.SubElement(cell,'mxGeometry',relative='1',attrib={'as':'geometry'})
                for end in ['source','target']:
                    n='1' if end=='source' else '2'
                    ET.SubElement(g,'mxPoint',x=str(it['x'+n]),y=str(it['y'+n]),attrib={'as':end+'Point'})
            elif it['kind']=='rect':
                attrs.update(vertex='1',style=f'rounded=0;fillColor={it["fill"]};strokeColor={it["stroke"]};strokeWidth=2;')
                cell=ET.SubElement(root,'mxCell',**attrs);ET.SubElement(cell,'mxGeometry',x=str(it['x']),y=str(it['y']),width=str(it['w']),height=str(it['h']),attrib={'as':'geometry'})
            else:
                w=min(2120,max(220,len(it['text'])*it['size']*.9))
                attrs.update(vertex='1',value=it['text'],style=f'text;html=0;whiteSpace=wrap;overflow=visible;fontSize={it["size"]};align={it["align"]};verticalAlign=middle;strokeColor=none;fillColor={"#FFFFFF" if it["bg"] else "none"};fontColor={it["color"]};')
                cell=ET.SubElement(root,'mxCell',**attrs)
                ET.SubElement(cell,'mxGeometry',x=str(it['x']-w/2 if it['align']=='center' else it['x']),y=str(it['y']-22),width=str(w),height='44',attrib={'as':'geometry'})
    ET.ElementTree(mxfile).write(DRAWIO,encoding='utf-8',xml_declaration=True)
export_drawio()

SHORT={'学生':('学生','Студент'),'管理员':('管理员','Администратор'),'CMP-001':('CMP-001 界面','Интерфейс'),'CMP-006':('CMP-006 预约','Бронирование'),'CMP-003':('CMP-003 审核','Проверка вуза'),'CMP-009':('CMP-009 数据','Данные'),'CMP-010':('CMP-010 来源','Источник вуза')}
DETAIL_B={
'1':(('seat_id；D；start/end','seat_id, D, start/end'),('界面取得单座位与时段','Принят выбор места и времени')),
'2':(('会话；seat_id；D；带时区时刻','Сессия, место, дата, время с часовым поясом'),('发起创建；user_id 取自会话','Создание; user_id берётся из сессии')),
'3':(('tx；user_id；room_id；seat_id','tx, user_id, room_id, seat_id'),('按既有锁序读取一致数据','Прочитать согласованные данные под блокировками')),
'4':(('tx；锁后 now；状态、规则和重叠数据','tx; now после блокировок; статус, правила, пересечения'),('业务组件在同一 tx 内作出校验决定','Бизнес-компонент проверяет правила в той же tx')),
'5a':(('403：资格；422：资源时间；409：重叠','403: допуск; 422: ресурс/время; 409: пересечение'),('先回滚，不创建预约','Сначала откат, без создания')),
'6a':(('原因码；重新选择提示','Код причины; предложение другого выбора'),('学生可改时段后再查询','Студент может выбрать другое время')),
'5b':(('tx；一个座位的预约；最少事件','tx; бронь одного места; минимальное событие'),('预约与事件原子提交','Бронь и событие фиксируются вместе')),
'6b':(('提交结果；booking_id','Результат COMMIT, booking_id'),('提交已确认，才允许响应成功','Фиксация подтверждена, успех допустим')),
'7b':(('HTTP 201；编号、座位、起止时刻','HTTP 201; ID, место, начало/конец'),('界面取得已保存预约','Интерфейс получает сохранённую бронь')),
'8b':(('预约编号与确认时段','ID брони и подтверждённый интервал'),('学生看到预约确认','Студент видит подтверждение')),
}
DETAIL_R={
'1':(('必要学号、姓名；认可渠道','Номер, имя; признанный канал'),('人工核验；OPEN-003 待确认','Ручная сверка; OPEN-003 открыт')),
'2':(('核验结论；来源引用','Итог сверки и ссылка'),('无法核验不得批准 VERIFIED','Без проверки VERIFIED запрещён')),
'3':(('user_id；decision；source_type/ref；revision','ID, решение, источник, revision'),('界面整理申请','Интерфейс формирует заявку')),
'4':(('会话；步骤 3 的审核申请','Сессия и заявка из шага 3'),('服务端重新授权并检查','Сервер повторно авторизует и проверяет')),
'5':(('tx；相关账号与固定锁序','tx; аккаунты, порядок блокировок'),('读取锁内当前状态','Прочитать данные под блокировками')),
'6':(('当前权限、状态、revision','Права, статус, revision'),('比对原身份与状态快照','Сверить с исходным снимком')),
'7a':(('403：权限；422：来源；409：快照变化','403: права; 422: источник; 409: изменённый снимок'),('回滚；无来源保持 PENDING','Откат; без источника остаётся PENDING')),
'7b':(('tx；决定、审核人、时间及来源引用','tx; решение, автор, время и ссылка'),('相符→VERIFIED；真实不符→REJECTED','Совпадение: VERIFIED; несоответствие: REJECTED')),
'8b':(('提交结果；VERIFIED 或 REJECTED','COMMIT; VERIFIED либо REJECTED'),('状态与事件原子提交','Общая фиксация статуса и события')),
'9b':(('HTTP 200；审核状态','HTTP 200; статус проверки'),('显示已保存结果','Показать сохранённый итог')),
}
HEAD6=[('序号','№'),('发送方','От'),('接收方','Кому'),('操作或消息','Операция / сообщение'),('数据','Данные'),('响应或结果','Ответ / результат')]
checkrows=[
 [('1 参与者与 3.3.1 组件及外部系统一致','Участники совпадают с компонентами и внешними системами диаграммы 3.3.1'),
  ('图 1 使用学生、CMP-001、006、009；图 2 使用管理员、CMP-001、003、009 和外部 CMP-010。均来自既有组件基线。','Рис. 1: студент, CMP-001, 006, 009. Рис. 2: администратор, CMP-001, 003, 009 и внешний CMP-010. Все ID из прежней схемы.'),('通过','Пройдено')],
 [('2 每条消息有名称和清晰方向','Каждое сообщение имеет подпись и понятное направление'),
  ('每条箭头标注编号、操作和中俄文字；六列消息表逐条记录发送方、接收方、数据和结果。','Все стрелки имеют номер, операцию и две подписи; таблицы задают отправителя, получателя, данные и результат.'),('通过','Пройдено')],
 [('3 至少一个替代流程','Показан хотя бы один альтернативный поток (ошибка данных, нет прав или недоступность)'),
  ('两图均有 alt。预约覆盖资格、非法时间和重叠；审核覆盖权限、来源不可核验与过期快照。','В обеих схемах alt. Бронь: недопуск, неверное время, пересечение. Проверка: нет прав, источник не проверен, снимок устарел.'),('通过','Пройдено')],
 [('4 有返回码或结果','Есть ответы с кодами или результатами'),
  ('预约返回 201 或 403/422/409；审核返回 200 或 403/422/409。对用户的成功响应在提交确认后返回。','Бронь: 201 либо 403/422/409. Проверка: 200 либо 403/422/409. Успех операции пользователю возвращается после COMMIT.'),('通过','Пройдено')],
 [('5 与用例及验收标准一致','Сценарий согласован с описанием Use Case и критериями приёмки'),
  ('沿用 US-002、US-004 及其正常、异常和并发边界；第 4 节核对验收标准，第 5 节列出需求和计划测试。','Сохранены US-002, US-004 и их границы; раздел 4 сверяет приёмку, раздел 5 указывает требования и плановые тесты.'),('通过','Пройдено')],
 [('6 一个目标且 6–10 条消息','Диаграмма не перегружена: одна цель, 6–10 сообщений'),
  ('每图总计 10 条箭头，包含全部分支及返回。预约正常 8／拒绝 6；审核正常 9／拒绝 7。','В каждой схеме 10 стрелок со всеми ветками и возвратами. Бронь: успех 8, отказ 6; проверка: успех 9, отказ 7.'),('通过','Пройдено')],
]

doc=Document()
def section(landscape=False,new=False):
    s=doc.add_section(WD_SECTION_START.NEW_PAGE) if new else doc.sections[0]
    if new:
        boundary=doc.paragraphs[-1]
        boundary.paragraph_format.space_before=Pt(0)
        boundary.paragraph_format.space_after=Pt(0)
        boundary.paragraph_format.line_spacing=Pt(1)
        boundary.add_run().font.size=Pt(1)
    s.orientation=WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    s.page_width,s.page_height=Inches(11 if landscape else 8.5),Inches(8.5 if landscape else 11)
    s.top_margin=s.bottom_margin=Inches(.4 if landscape else .55)
    s.left_margin=s.right_margin=Inches(.55)
    return s
section()
for name in ['Normal','Title','Heading 1','Heading 2']:
    st=doc.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
    rf=st.element.get_or_add_rPr().get_or_add_rFonts()
    for k,v in [('ascii','Arial'),('hAnsi','Arial'),('eastAsia','Heiti SC')]:rf.set(qn('w:'+k),v)
    for el in list(st.element.iter(qn('w:pBdr'))):el.getparent().remove(el)
doc.styles['Normal'].font.size=Pt(10.5)
doc.styles['Normal'].paragraph_format.line_spacing=1.04
doc.styles['Normal'].paragraph_format.space_after=Pt(4)
for name,size in [('Title',18),('Heading 1',13),('Heading 2',11.5)]:
    doc.styles[name].font.size=Pt(size)
    doc.styles[name].paragraph_format.space_before=Pt(7 if name!='Title' else 0)
    doc.styles[name].paragraph_format.space_after=Pt(3)
MD=[]
def run(p,text,size=None,bold=False,color='000000',lang='zh-CN'):
    r=p.add_run(text);r.font.name='Arial';r.bold=bold;r.font.color.rgb=RGBColor.from_string(color)
    if size:r.font.size=Pt(size)
    pr=r._element.get_or_add_rPr();pr.rFonts.set(qn('w:eastAsia'),'Heiti SC')
    le=OxmlElement('w:lang');le.set(qn('w:val'),lang);pr.append(le)
    return r
def pair(zh,ru,style=None,size=None):
    aftertable=len(doc.element.body)>1 and doc.element.body[-2].tag==qn('w:tbl')
    p=doc.add_paragraph(style=style);run(p,zh,size);p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(2)
    if aftertable:p.paragraph_format.space_before=Pt(5)
    p=doc.add_paragraph(style=style);run(p,ru,size,lang='ru-RU');p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(4);p.paragraph_format.keep_together=True
    MD.extend([zh,ru,''])
    return p
def table(headers,rows,widths,size=9.5,padding=85):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    bd=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge)
        for k,v in [('val','single'),('sz','5'),('color','D9D9D9')]:e.set(qn('w:'+k),v)
        bd.append(e)
    t._tbl.tblPr.append(bd);t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for ri,vals in enumerate([headers,*rows]):
        row=t.rows[0] if ri==0 else t.add_row();row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for ci,(c,txts) in enumerate(zip(row.cells,vals)):
            c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
            for edge,val in [('top',padding),('bottom',padding),('left',75),('right',75)]:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),str(val));e.set(qn('w:type'),'dxa');mar.append(e)
            pr.append(mar);sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'29485D' if ri==0 else ('F1F5F8' if ri%2 else 'FFFFFF'));pr.append(sh)
            for li,text in enumerate(txts):
                p=c.paragraphs[0] if li==0 else c.add_paragraph();p.paragraph_format.space_after=Pt(3 if li==0 and len(txts)>1 else 0);p.paragraph_format.line_spacing=1.02;p.paragraph_format.keep_with_next=False
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER if ri==0 or (ci==0 and len(headers)==6) or ci==len(headers)-1 and len(headers)==3 else WD_ALIGN_PARAGRAPH.LEFT
                run(p,text,size,ri==0,'FFFFFF' if ri==0 else '000000','zh-CN' if li==0 else 'ru-RU')
    MD.extend(['| '+' | '.join('<br>'.join(v) for v in headers)+' |','| '+' | '.join('---' for _ in headers)+' |'])
    MD.extend('| '+' | '.join('<br>'.join(v).replace('\n','<br>') for v in row)+' |' for row in rows);MD.append('')
    return t
def figure(p,zh,ru):
    section(True,True);pair(zh,ru,'Heading 1',12)
    for hp in doc.paragraphs[-2:]:hp.paragraph_format.space_before=Pt(0);hp.paragraph_format.space_after=Pt(2);hp.paragraph_format.line_spacing=1
    q=doc.add_paragraph();q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.paragraph_format.line_spacing=1;q.paragraph_format.space_after=Pt(0)
    shape=q.add_run().add_picture(str(FIG/(p.key+'.png')),width=Inches(9.9));shape._inline.docPr.set('descr',zh+' '+ru)
    MD.append(f'![{zh} {ru}](diagrams/{p.key}.svg)\n')
def msgtable(p,detail,zh,ru):
    section(True,True);pair(zh,ru,'Heading 1',12)
    for hp in doc.paragraphs[-2:]:
        hp.paragraph_format.space_before=Pt(0)
        hp.paragraph_format.space_after=Pt(2)
        hp.paragraph_format.line_spacing=1
    rows=[[(m['number'],),SHORT[m['src']],SHORT[m['dst']],(m['zh'],m['ru']),*detail[m['number']]] for m in p.messages]
    table(HEAD6,rows,[.46,1.12,1.12,2.72,2.48,2.0],9.5,padding=75)
    return rows

pair('StudySpace 时序图检查与精简版','StudySpace Проверка и краткие диаграммы последовательности','Title')
p=doc.add_paragraph();run(p,'SS-SEQ-01  |  v2  |  2026-10-15  |  DRAFT',9)
pair('本版按课件六项标准核查并精简 US-002、US-004，每图均为 10 条可见消息。此前 v1 分别为 25 条与 27 条，不符合本次新增的数量限制，保留作实现细节参考。',
     'По шести критериям проверены и сокращены US-002 и US-004: в каждой схеме 10 видимых сообщений. Подробная v1 с 25 и 27 сообщениями не отвечает новому лимиту и сохранена для деталей реализации.',size=10)
pair('1 六项课堂检查结果','1 Результаты шести проверок','Heading 1')
table([('课堂检查项','Критерий слайда'),('StudySpace 核对内容与证据','Сверка и свидетельства StudySpace'),('结果','Итог')],checkrows,[2.4,4.1,.9],9.5)
pair('通过仅指本版设计材料的核对；业务测试保持 NOT_RUN，教师验收尚未执行。计数包含全部分支的调用与返回箭头。',
     'Пройдено означает сверку проектного материала. Бизнес-тесты NOT_RUN, приёмка преподавателя не выполнена. Подсчёт включает вызовы и возвраты всех веток.',size=9)

figure(b,'图 1 US-002 创建单座位预约 共 10 条消息','Рисунок 1 US-002 Создание одной брони 10 сообщений')
rows_b=msgtable(b,DETAIL_B,'2 图 1 配套消息表','2 Таблица сообщений рисунка 1')
figure(r,'图 2 US-004 校内身份审核 共 10 条消息','Рисунок 2 US-004 Проверка вуза 10 сообщений')
rows_r=msgtable(r,DETAIL_R,'3 图 2 配套消息表','3 Таблица сообщений рисунка 2')

section(False,True)
pair('4 用例与验收标准核对','4 Сверка Use Case и критериев приёмки','Heading 1')
pair('4.1 US-002 创建单座位预约','4.1 US-002 Создание одной брони','Heading 2')
pair('目标是创建一条预约；前提是学生已登录、已选择已有座位，系统仍在服务端复核会话、启用状态和 VERIFIED。查询结果不保证提交时空闲，user_id 从会话取得。成功仅在预约与创建事件一起提交后返回 201。',
     'Цель — одна бронь. Студент вошёл и выбрал существующее место; сервер повторно проверяет сессию, активность и VERIFIED. Поиск не гарантирует доступность при отправке; user_id берётся из сессии. Ответ 201 возможен после общего COMMIT записи и события.')
pair('对应既有验收标准：合法申请只创建一条并返回 ID；未核验、同学生或同座位重叠均拒绝；[start,end) 边界相邻及多条非重叠预约允许。校验失败回滚，不能把查询后再写入当作并发保证。',
     'Прежняя приёмка: допустимая заявка создаёт одну запись с ID; недопуск и пересечения студента либо места запрещены. Для [start,end) соседние и несколько непересекающихся записей разрешены. Отказ требует отката; обычные поиск и запись не доказывают конкурентную корректность.')
pair('时间仍按 Asia/Shanghai 业务日、UTC 存储，默认 08:00–22:00、30 分钟粒度，今天至今天+6，不跨所选业务日，时长硬上限 24 小时且默认窗口最多 14 小时。起点严格晚于锁后实际 now；PostgreSQL 锁序 User → Room → Seat，SQLite 使用 BEGIN IMMEDIATE。',
     'День Asia/Shanghai, хранение UTC; по умолчанию 08:00–22:00, шаг 30 минут, от сегодня до сегодня+6 без выхода за день. Предел 24 часа, при обычных часах 14. Начало позже фактического now после блокировок; PostgreSQL: User → Room → Seat, SQLite: BEGIN IMMEDIATE.',size=10)
pair('4.2 US-004 校内身份审核','4.2 US-004 Проверка принадлежности к вузу','Heading 2')
pair('目标是保存一项人工审核决定。图从管理员已打开启用的 PENDING 账号并取得身份状态快照开始；读取待审核列表是前置步骤。CMP-010 为外部人工核验来源，具体来源仍由 OPEN-003 确认，不假设已有自动 API。',
     'Цель — сохранить одно решение ручной проверки. В начале открыт активный аккаунт PENDING и получен снимок; загрузка списка является предварительным шагом. CMP-010 — внешний источник ручной сверки, реальный канал в OPEN-003, без готового автоматического API.')
pair('对应既有验收标准：可信来源与身份相符才可保存 VERIFIED；真实核验不符保存 REJECTED；无可信来源不批准，保持 PENDING。保存时重新授权并比较锁内快照；身份变更或停用竞争导致快照失效时回滚，不能取得预约资格。',
     'Прежняя приёмка: достоверное совпадение позволяет VERIFIED, несоответствие — REJECTED; без источника остаётся PENDING. При сохранении повторно авторизовать и сравнить снимок под блокировкой. Изменение идентификаторов или отключение делает снимок устаревшим: откат без допуска.')
pair('状态与审核事件一起提交。事件仅含 user_id、审核人、event_at、决定和来源类型引用，不复制原始学号、姓名或证明影像。REJECTED 可申请复核；只有启用且 VERIFIED 的 STUDENT 能预约。',
     'Статус и событие фиксируются вместе. Событие содержит user_id, автора, event_at, решение, тип и ссылку источника без копий номера, имени или изображений. После REJECTED возможна повторная заявка; бронирует только активный STUDENT с VERIFIED.',size=10)

doc.add_page_break()
pair('5 抽象层次与追溯说明','5 Уровень детализации и трассируемость','Heading 1')
pair('图 1 的业务控制仍由 CMP-006 承担；CMP-009 只承担锁、数据读取、持久化与事务支持。CMP-002 认证、CMP-003 资格、CMP-004 资源和 CMP-008 最少审计事件属于未展开的内部协作，不转移给数据访问组件。',
     'На рисунке 1 бизнес-решения остаются у CMP-006; CMP-009 отвечает за блокировки, данные, хранение и транзакцию. Аутентификация CMP-002, допуск CMP-003, ресурсы CMP-004 и минимальный аудит CMP-008 — нераскрытые внутренние вызовы, не переданные слою данных.')
pair('图 2 保留管理员与 CMP-010 的人工信息交互。CMP-003 负责审核决定，CMP-002 再次授权、CMP-008 准备最少事件的内部细节沿用 v1。两图的合并保存调用表示同一事务的工作单元，成功返回必须晚于 COMMIT 确认。',
     'Рисунок 2 сохраняет ручной обмен администратора с CMP-010. Решение принимает CMP-003, повторная авторизация CMP-002 и минимальное событие CMP-008 раскрыты в v1. Объединённое сохранение означает работу одной транзакции; успех следует после подтверждения COMMIT.')
pair('本版合并相同保存路径的不同决定参数：相符为 VERIFIED，真实不符为 REJECTED；失败分支按原因分别返回 403、422 或 409。锁超时和未知数据库异常仍须回滚并按实际错误处理，不一律归为 409。',
     'Одинаковый путь сохранения объединяет параметры VERIFIED при совпадении и REJECTED при достоверном несоответствии. Отказ возвращает 403, 422 либо 409 по причине. Тайм-аут и неизвестные сбои БД требуют отката и отдельной обработки, а не всегда 409.')
table([('用例','Сценарий'),('需求与规则','Требования и правила'),('计划测试','Плановые тесты')],[
 [('US-002','US-002'),('SP-FR-008、SP-FR-009、SP-NFR-004；DEC-016','SP-FR-008, SP-FR-009, SP-NFR-004; DEC-016'),('TC-007、TC-008、TC-017、TC-025、TC-031、TC-033','TC-007, TC-008, TC-017, TC-025, TC-031, TC-033')],
 [('US-004','US-004'),('SP-FR-022、SP-IF-002、SP-DR-001；OPEN-003','SP-FR-022, SP-IF-002, SP-DR-001; OPEN-003'),('TC-015、TC-035、TC-038','TC-015, TC-035, TC-038')],
],[.9,3.8,2.7],9.5)
pair('图和表均为 DRAFT 设计。API 路径、错误码与审核 revision 是拟定契约；内部调用仍在 FastAPI 模块化单体内。业务实现未完成，业务测试 NOT_RUN，材料结构和视觉检查不替代教师验收。',
     'Диаграммы и таблицы — проект DRAFT. Пути API, коды и revision являются предлагаемыми контрактами; внутренние вызовы остаются в модульном монолите FastAPI. Реализация не завершена, тесты NOT_RUN; проверка материала не заменяет приёмку преподавателя.',size=10)
pair('依据为本次六项检查课件、既有 SS-SEQ-01 v1 图与任务三消息表、SS-ARCH-03 v2 组件基线、DOC-04 0.3.0、课堂需求登记 v0.2 与 DEC-016。本次新增 v2，原详细版及业务需求文件保持不变。',
     'Основания: слайд шести проверок, схемы SS-SEQ-01 v1 и таблицы задания 3, компоненты SS-ARCH-03 v2, DOC-04 0.3.0, учебный реестр v0.2 и DEC-016. Добавлена v2 без изменения подробной версии и требований.',size=9)
pair('交付位置为 studyspace/coursework/2026-10-15-sequence-review-v2。执行环境为本地 macOS、Codex 内置 Python 与文档渲染工具；可编辑 draw.io 含两个场景页签。',
     'Каталог: studyspace/coursework/2026-10-15-sequence-review-v2. Среда: локальная macOS, встроенные Python и рендеринг Codex; редактируемый draw.io содержит две вкладки.',size=9)

doc.core_properties.title='StudySpace 时序图检查与精简版 Проверка диаграмм'
doc.core_properties.subject='六项检查 两张十消息时序图及配套消息表'
doc.core_properties.author='StudySpace'
doc.save(DOCX)
(OUT/(STEM+'.md')).write_text('\n\n'.join(MD),encoding='utf-8')
content={'doc_id':'SS-SEQ-01','version':'v2','date':'2026-10-15','status':'DRAFT','business_tests':'NOT_RUN','checklist':checkrows,'old_message_counts':{'US-002':25,'US-004':27},'diagrams':[{'id':'US-002','key':b.key,'goal_zh':'创建一条单座位预约','goal_ru':'Создать одну бронь места','lanes':b.lanes,'messages':b.messages,'rows':rows_b,'paths':{'success':['1','2','3','4','5b','6b','7b','8b'],'failure':['1','2','3','4','5a','6a']}},{'id':'US-004','key':r.key,'goal_zh':'保存一项人工校内审核决定','goal_ru':'Сохранить одно решение проверки вуза','lanes':r.lanes,'messages':r.messages,'rows':rows_r,'paths':{'success':['1','2','3','4','5','6','7b','8b','9b'],'failure':['1','2','3','4','5','6','7a']}}],'source_fingerprints':fingerprints}
(DATA/'content.json').write_text(json.dumps(content,ensure_ascii=False,indent=2),encoding='utf-8')
shutil.copy2(PREV/'data/fonts.conf',DATA/'fonts.conf')
rd=Document(DOCX)
with ZipFile(DOCX) as z:
    xml=ET.fromstring(z.read('word/document.xml'));names=z.namelist()
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
texts='\n'.join(p.text for p in rd.paragraphs)+'\n'+'\n'.join(c.text for t in rd.tables for row in t.rows for c in row.cells)
cmps=set(re.findall(r'CMP-\d{3}',(ROOT.parent/'StudySpace_组件交付/StudySpace_组件候选与组件图_中俄双语_v2.md').read_text()))
check={
'six_slide_criteria':len(rd.tables[0].rows)==7,
'two_scenarios_preserved':[d['id'] for d in content['diagrams']]==['US-002','US-004'],
'ten_visible_messages_each':all(len(p.messages)==10 for p in [b,r]),
'ten_actual_arrow_primitives_each':all(sum(1 for it in p.c.items if it['kind']=='line' and it['arrow'])==10 for p in [b,r]),
'every_normal_and_error_path_within_limit':all(6<=len(path)<=10 for d in content['diagrams'] for path in d['paths'].values()),
'participants_match_existing_components':all(k in cmps or k in {'学生','管理员'} for p in [b,r] for k,_,_ in p.lanes),
'all_messages_have_direction_and_two_labels':all(m['src'] in {x[0] for x in p.lanes} and m['dst'] in {x[0] for x in p.lanes} and m['zh'] and re.search('[А-Яа-яЁё]',m['ru']) for p in [b,r] for m in p.messages),
'both_diagrams_have_alt':bool(b.frames) and bool(r.frames),
'both_diagrams_have_coded_returns':all(any(m['return_message'] and re.search(r'200|201|403|422|409',m['zh']) for m in p.messages) for p in [b,r]),
'two_matching_ten_row_message_tables':len(rd.tables[1].rows)==11 and len(rd.tables[2].rows)==11 and len(rd.tables[1].columns)==len(rd.tables[2].columns)==6,
'message_table_order_matches_diagrams':all(t.cell(i+1,0).text==m['number'] and t.cell(i+1,3).text=='\n'.join((m['zh'],m['ru'])) for t,p in [(rd.tables[1],b),(rd.tables[2],r)] for i,m in enumerate(p.messages)),
'bilingual_table_cells':all(len(c.paragraphs)==2 for ti,t in enumerate(rd.tables) for ri,row in enumerate(t.rows) for ci,c in enumerate(row.cells) if not (ti in {1,2} and ri>0 and ci==0)),
'commits_precede_success':all(next(i for i,m in enumerate(p.messages) if 'COMMIT' in m['zh'])<next(i for i,m in enumerate(p.messages) if m['return_message'] and ('201' in m['zh'] or '200' in m['zh'])) for p in [b,r]),
'manual_external_source_preserved':'OPEN-003' in texts and '人工' in r.messages[0]['zh'] and 'CMP-010' in texts,
'business_policy_retained':all(s in texts for s in ['[start,end)','多条非重叠预约','User → Room → Seat','BEGIN IMMEDIATE']),
'business_tests_not_claimed_run':'NOT_RUN' in texts and '拟定契约' in texts,
'previous_excess_count_disclosed':all(s in texts for s in ['25 条','27 条','不符合本次新增的数量限制']),
'two_embedded_diagrams':len(rd.inline_shapes)==2,
'figures_fit_landscape':all(p.c.height/2200*9.9<6.95 for p in [b,r]),
'all_figure_text_within_horizontal_bounds':all(left>=20 and right<=2180 for left,right in text_bounds),
'two_editable_drawio_tabs':len(ET.parse(DRAWIO).findall('diagram'))==2,
'all_source_files_preserved':all(hashlib.sha256(Path(p).read_bytes()).hexdigest()==h for p,h in fingerprints.items()),
'no_headers_or_footers':not any(re.fullmatch(r'word/(header|footer)\d+\.xml',n) for n in names),
}
report={'checks':check,'check_count':len(check),'all_pass':all(check.values()),'visual_review':'PENDING','business_tests':'NOT_RUN','message_counts':{'US-002':len(b.messages),'US-004':len(r.messages)},'source_fingerprints':fingerprints,'docx_sha256':hashlib.sha256(DOCX.read_bytes()).hexdigest(),'environment':'Local macOS; bundled Codex Python; design material checks only','timestamp_utc':datetime.now(timezone.utc).isoformat()}
(DATA/'verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
(OUT/'README.md').write_text(f'''# StudySpace 时序图检查与精简版
# StudySpace Проверка и краткие диаграммы последовательности

SS-SEQ-01 · v2 · 2026-10-15 · DRAFT

主文档 / Документ [{STEM}.docx]({STEM}.docx)
可编辑图源 / Редактируемые схемы [{DRAWIO.name}]({DRAWIO.name})

六项课件检查、两个各 10 条箭头的简图及各 10 行配套六列消息表。原详细 v1 为 25/27 条，仅作细节参考。
Шесть проверок, две схемы по 10 стрелок и две таблицы по 10 строк. Подробная v1 с 25/27 сообщениями остаётся для деталей.

内容 data/content.json；检查 data/verification.json；业务测试 NOT_RUN。
Содержание data/content.json; проверки data/verification.json; бизнес-тесты NOT_RUN.
''',encoding='utf-8')
print(json.dumps({'docx':str(DOCX),'checks':len(check),'all_pass':report['all_pass'],'failed':[k for k,v in check.items() if not v],'messages':report['message_counts'],'figure_heights':[b.c.height,r.c.height]},ensure_ascii=False))
if not report['all_pass']:raise SystemExit(1)
