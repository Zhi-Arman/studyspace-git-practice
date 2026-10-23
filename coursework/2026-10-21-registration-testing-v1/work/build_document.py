from pathlib import Path
import json
import hashlib
import shutil
from zipfile import ZipFile
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.text.paragraph import Paragraph
from docx.table import Table

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / 'data'
DATA.mkdir(parents=True, exist_ok=True)
STEM = '网店注册表单_测试设计与缺陷报告_中俄双语_v1'
FONTDIR=ROOT/'work'/'fonts'
FONTDIR.mkdir(exist_ok=True)
for font in ['Arial.ttf','Arial Bold.ttf','Arial Unicode.ttf']:
    shutil.copy2(ROOT.parent/'2026-10-19-bug-report-review-v1'/'work'/'fonts'/font,FONTDIR/font)
font_config = ('<?xml version="1.0"?><!DOCTYPE fontconfig SYSTEM "fonts.dtd"><fontconfig>'
               f'<dir>{FONTDIR}</dir>'
               '<dir>/Users/zhi/.cache/codex-runtimes/codex-primary-runtime/dependencies/native/libreoffice-headless/libreoffice/LibreOfficeDev.app/Contents/Resources/fonts/truetype</dir>'
               f'<cachedir>{ROOT}/work/fontcache</cachedir></fontconfig>')
(ROOT/'work'/'fonts.conf').write_text(font_config)

def pair(zh, ru):
    return [zh, ru]

requirements = [
    ['REG-R01', *pair('姓名：必填，长度 2–30 个字符，包含端点。', 'Имя: обязательное, от 2 до 30 символов включительно.')],
    ['REG-R02', *pair('Email：必填，含 @，域名部分含点。', 'Email: обязательный, содержит @ и точку в доменной части.')],
    ['REG-R03', *pair('密码：必填，长度 8–20 个字符，至少一位数字。', 'Пароль: обязательный, 8–20 символов, минимум одна цифра.')],
    ['REG-R04', *pair('年龄：必填，18–100 的整数，包含端点。', 'Возраст: обязательный, целое число от 18 до 100 включительно.')],
    ['REG-R05', *pair('注册成功：显示“Регистрация завершена”并向填写的邮箱发送邮件。', 'Успех: сообщение «Регистрация завершена» и письмо на указанный адрес.')],
]
tasks = [
    pair('任务 1：为姓名、密码、年龄填写等价类和边界值表。', 'Задание 1: заполнить таблицу классов эквивалентности и граничных значений для имени, пароля и возраста.'),
    pair('任务 2：至少 5 个用例，含 1 个正向和至少 3 个负向；每例含 ID、名称、前置条件、步骤及预期。', 'Задание 2: не менее 5 кейсов, один позитивный и не менее трёх негативных; ID, название, предусловия, шаги и ожидаемый результат.'),
    pair('任务 3：按 18 岁被拒绝的情境写缺陷报告，判定严重程度和优先级并说明原因。', 'Задание 3: оформить баг-репорт об отказе при возрасте 18 лет; определить серьёзность и приоритет с обоснованием.'),
    pair('任务 4：评价首页拼写错误、超大头像导致崩溃、所有用户无法支付这三项错误。', 'Задание 4: оценить опечатку на главной, сбой при большом аватаре и неработающую у всех кнопку оплаты.'),
]

equivalence = [
    [pair('姓名', 'Имя'),
     pair('有效：N-V，2 ≤ L ≤ 30。\n无效：N-E，空值 L=0；N-S，L=1；N-L，L≥31。\nL 为字符数。', 'Допустимый: N-V, 2 ≤ L ≤ 30.\nНедопустимые: N-E, пусто L=0; N-S, L=1; N-L, L≥31.\nL — число символов.'),
     pair('长度：1 拒绝；2、3、29、30 接受；31 拒绝。\n必填另测空值。', 'Длины: 1 — отказ; 2, 3, 29, 30 — допуск; 31 — отказ.\nОтдельно проверить пустое поле.')],
    [pair('密码', 'Пароль'),
     pair('有效：P-V，8 ≤ L ≤ 20 且数字数 D≥1。\n无效：P-E，空值；P-S，1≤L≤7；P-L，L≥21；P-D，8≤L≤20 但 D=0。', 'Допустимый: P-V, 8 ≤ L ≤ 20 и число цифр D≥1.\nНедопустимые: P-E, пусто; P-S, 1≤L≤7; P-L, L≥21; P-D, 8≤L≤20, но D=0.'),
     pair('含数字时，长度：7 拒绝；8、9、19、20 接受；21 拒绝。\n合法长度下：0 位数字拒绝，1 位数字接受。', 'При наличии цифры длины: 7 — отказ; 8, 9, 19, 20 — допуск; 21 — отказ.\nПри допустимой длине: 0 цифр — отказ, 1 цифра — допуск.')],
    [pair('年龄', 'Возраст'),
     pair('有效：A-V，整数 18≤A≤100。\n无效：A-E，空值；A-S，整数 A≤17；A-L，整数 A≥101；A-F，非整数数值（18.5）；A-T，非数值（abc）。', 'Допустимый: A-V, целое 18≤A≤100.\nНедопустимые: A-E, пусто; A-S, целое A≤17; A-L, целое A≥101; A-F, дробь (18.5); A-T, нечисловое значение (abc).'),
     pair('整数：17 拒绝；18、19、99、100 接受；101 拒绝。\n18.5、abc 和空值均拒绝。', 'Целые: 17 — отказ; 18, 19, 99, 100 — допуск; 101 — отказ.\n18.5, abc и пустое поле отклоняются.')],
]

common = pair('C0 前置条件：测试环境的注册页可用；未登录；每次使用尚未注册、可接收邮件的受控测试邮箱。用例互相独立，参数组逐组重置表单。',
              'Предусловия C0: доступна страница регистрации тестового стенда; пользователь не вошёл; для каждого запуска подготовлен новый контролируемый адрес с доступом к почте. Кейсы независимы, перед каждым набором форма сбрасывается.')
baseline = pair('基线 B：姓名 Анна，Email 为本次受控地址，密码 Password1，年龄 25。下文只替换指定字段，其他字段保持 B。',
                'База B: имя Анна, email — контролируемый адрес текущего запуска, пароль Password1, возраст 25. Ниже меняется только указанное поле, остальные остаются как в B.')
good = pair('允许注册，显示“Регистрация завершена”，受控邮箱收到注册邮件。', 'Регистрация разрешена; показано «Регистрация завершена»; на контролируемый адрес приходит письмо.')

cases = []
def case(num, zh, ru, typ, values, steps, expected, coverage):
    cases.append({'id': f'REG-TC-{num:02}', 'title': pair(zh, ru), 'type': typ, 'preconditions': pair('C0；使用基线 B。', 'C0; использовать базу B.'), 'values': values, 'steps': steps, 'expected': expected, 'coverage': coverage, 'actual': 'NOT_RUN'})

submit = pair('1 打开并填写表单。2 点击“注册”。3 检查页面和受控邮箱。', '1 Открыть и заполнить форму. 2 Нажать «Регистрация». 3 Проверить страницу и контролируемую почту.')
reject_steps = pair('1 填写 B 并替换指定字段。2 提交；如输入被控件阻止，记录该现象。3 对每组数据重置后重复。', '1 Заполнить B, заменив указанное поле. 2 Отправить; если ввод блокируется, зафиксировать это. 3 Сбросить форму и повторить для каждого набора.')

case(1, '正常注册', 'Обычная регистрация', 'POSITIVE', pair('使用 B 的全部有效数据。', 'Все допустимые данные B.'), submit, good, ['REG-R01','REG-R02','REG-R03','REG-R04','REG-R05'])
case(2, '有效下边界', 'Допустимые нижние границы', 'POSITIVE', pair('姓名 Ли（2）；密码 Passwor1（8）；年龄 18。', 'Имя Ли (2); пароль Passwor1 (8); возраст 18.'), submit, good, ['N-V','P-V','A-V','REG-BR-01'])
case(3, '有效上边界及边界邻值', 'Допустимые верхние и соседние значения', 'POSITIVE', pair('独立参数组：① A×30 / A×19+1 / 100；② A×3 / A×8+1 / 19；③ A×29 / A×18+1 / 99。顺序为姓名 / 密码 / 年龄。', 'Независимые наборы имя / пароль / возраст: ① A×30 / A×19+1 / 100; ② A×3 / A×8+1 / 19; ③ A×29 / A×18+1 / 99.'), pair('每组用新测试邮箱填写表单并提交，检查页面和该邮箱。', 'Для каждого набора взять новый тестовый адрес, заполнить и отправить форму; проверить страницу и почту.'), good, ['N-V','P-V','A-V'])
case(4, '姓名必填', 'Обязательность имени', 'NEGATIVE', pair('姓名留空，其余为 B。', 'Имя пустое, остальные поля как в B.'), reject_steps, pair('拒绝注册并提示姓名必填；无成功提示、无注册成功邮件。', 'Отказ с указанием обязательности имени; нет сообщения об успехе и письма об успешной регистрации.'), ['N-E'])
case(5, '姓名长度越界', 'Недопустимая длина имени', 'NEGATIVE', pair('两组：姓名 A（1）；姓名 A×31（31）。', 'Два набора: имя A (1); имя A×31 (31).'), reject_steps, pair('两组均拒绝，提示姓名长度应为 2–30；无成功提示及成功邮件。', 'Оба набора отклоняются: имя должно содержать 2–30 символов; нет сообщения и письма об успехе.'), ['N-S','N-L'])
case(6, '密码必填', 'Обязательность пароля', 'NEGATIVE', pair('密码留空，其余为 B。', 'Пароль пустой, остальные поля как в B.'), reject_steps, pair('拒绝注册并提示密码必填；无成功提示及成功邮件。', 'Отказ с указанием обязательности пароля; нет сообщения и письма об успехе.'), ['P-E'])
case(7, '密码长度越界', 'Недопустимая длина пароля', 'NEGATIVE', pair('两组密码：A×6+1（7）；A×20+1（21），均含一位数字。', 'Два пароля: A×6+1 (7); A×20+1 (21), каждый содержит одну цифру.'), reject_steps, pair('两组均拒绝，提示密码长度应为 8–20；无成功提示及成功邮件。', 'Оба набора отклоняются: длина пароля должна быть 8–20; нет сообщения и письма об успехе.'), ['P-S','P-L'])
case(8, '密码缺少数字', 'Пароль без цифры', 'NEGATIVE', pair('密码 Password（8 个字符，0 位数字）。', 'Пароль Password (8 символов, 0 цифр).'), reject_steps, pair('拒绝注册并提示至少一位数字；无成功提示及成功邮件。', 'Отказ: пароль должен содержать минимум одну цифру; нет сообщения и письма об успехе.'), ['P-D'])
case(9, '年龄必填', 'Обязательность возраста', 'NEGATIVE', pair('年龄留空，其余为 B。', 'Возраст пустой, остальные поля как в B.'), reject_steps, pair('拒绝注册并提示年龄必填；无成功提示及成功邮件。', 'Отказ с указанием обязательности возраста; нет сообщения и письма об успехе.'), ['A-E'])
case(10, '年龄数值越界', 'Возраст вне диапазона', 'NEGATIVE', pair('两组年龄：17；101。', 'Два значения возраста: 17; 101.'), reject_steps, pair('两组均拒绝，提示年龄应为 18–100；无成功提示及成功邮件。', 'Оба значения отклоняются: возраст должен быть 18–100; нет сообщения и письма об успехе.'), ['A-S','A-L'])
case(11, '年龄不是整数', 'Возраст не является целым числом', 'NEGATIVE', pair('两组年龄：18.5；abc。', 'Два значения возраста: 18.5; abc.'), reject_steps, pair('两组均不得成功注册，提示输入 18–100 的整数或由控件阻止非法输入；不得静默取整。', 'Успешная регистрация невозможна: сообщение о целом числе 18–100 или блокировка ввода; без скрытого округления.'), ['A-F','A-T'])
case(12, 'Email 必填及格式', 'Обязательность и формат email', 'NEGATIVE', pair('三组 Email：空值；testmail.ru（无 @）；test@mailru（域名无点）。', 'Три email: пусто; testmail.ru (нет @); test@mailru (нет точки в домене).'), reject_steps, pair('分别提示 Email 必填或格式不正确；均不得成功注册或发送成功邮件。', 'Соответственно указать обязательность email или ошибку формата; без успешной регистрации и письма об успехе.'), ['REG-R02'])

bug = [
    [pair('编号与状态', 'ID и статус'), pair('REG-BR-01；DRAFT；题设情境报告，尚未实测。', 'REG-BR-01; DRAFT; учебная ситуация, фактическая проверка не выполнялась.')],
    [pair('标题', 'Заголовок'), pair('有效年龄 18 被拒绝，提示年龄必须大于 18', 'Допустимый возраст 18 отклоняется с требованием возраста больше 18')],
    [pair('模块与环境', 'Модуль и окружение'), pair('网店注册；网站 2.3；Chrome（具体版本未给出）；Windows 11。', 'Регистрация; сайт 2.3; Chrome (версия не указана); Windows 11.')],
    [pair('前置条件', 'Предусловия'), pair('注册页打开；未登录；执行前确认 test@mail.ru 未注册。', 'Форма открыта; вход не выполнен; перед запуском проверить, что test@mail.ru не зарегистрирован.')],
    [pair('复现步骤', 'Шаги воспроизведения'), pair('1 打开注册页。2 输入姓名 Ли、Email test@mail.ru、密码 Password1、年龄 18。3 点击注册并观察。', '1 Открыть регистрацию. 2 Ввести имя Ли, email test@mail.ru, пароль Password1, возраст 18. 3 Нажать регистрацию и проверить результат.')],
    [pair('预期结果', 'Ожидаемый результат'), pair('18 满足 REG-R04 的含端点范围。应注册成功，显示“Регистрация завершена”，并向邮箱发送邮件。', '18 входит в диапазон REG-R04 включительно. Регистрация успешна, показано «Регистрация завершена», отправлено письмо.')],
    [pair('实际结果', 'Фактический результат'), pair('题设：页面显示“Возраст должен быть больше 18”（年龄必须大于 18），未显示成功消息。账号是否创建、邮件是否发送需另行核实。', 'По условию: показано «Возраст должен быть больше 18», сообщения об успехе нет. Создание аккаунта и отправка письма требуют проверки.')],
    [pair('严重程度与理由', 'Серьёзность и причина'), pair('S2 Major：合法的 18 岁用户注册受阻；影响边界人群。', 'S2 Major: регистрация допустимых пользователей 18 лет заблокирована; затронута граничная группа.')],
    [pair('优先级与理由', 'Приоритет и причина'), pair('P1 High：合法用户注册受阻，应尽快修复。', 'P1 High: требуется быстро устранить блокировку регистрации допустимых пользователей.')],
    [pair('证据与复现率', 'Свидетельства и частота'), pair('题设第 7/8 页；无实际错误截图或日志；复现率未测量。', 'Ситуация на слайде 7/8; нет реального снимка или журналов; частота не измерена.')],
]

priorities = [
    [pair('首页“Регистрация”拼写错误', 'Опечатка в слове «Регистрация» на главной'), pair('S4 Trivial\n轻微', 'S4 Trivial\nКосметическая'), pair('P3 Low\n低', 'P3 Low\nНизкий'), pair('仅影响文字呈现，题设没有功能阻断。随常规内容修订处理；若影响关键活动，可调整优先级。', 'Влияет на текст, функциональная блокировка не указана. Исправить при обычной редактуре; при важной кампании приоритет может меняться.')],
    [pair('上传超过 50 MB 的头像时崩溃，要求最大 5 MB', 'Аварийное завершение при аватаре более 50 МБ при лимите 5 МБ'), pair('S2 Major\n严重', 'S2 Major\nСерьёзная'), pair('P2 Normal\n中', 'P2 Normal\nОбычный'), pair('超限文件应被校验拒绝，不能导致崩溃。触发数据不符合要求且有使用 ≤5 MB 文件的替代方式，故优先级低于核心流程阻断。若影响全站或造成数据丢失，需升级。', 'Файл сверх лимита следует отклонять, а не аварийно завершать приложение. Сценарий вне допустимого ввода; можно использовать файл ≤5 МБ. При сбое всего сервиса или потере данных оценку повысить.')],
    [pair('支付按钮对所有用户失效', 'Кнопка оплаты не работает у всех'), pair('S1 Blocker\n阻断', 'S1 Blocker\nБлокирующая'), pair('P0 Urgent\n紧急', 'P0 Urgent\nСрочный'), pair('所有用户均不能完成核心交易，直接阻断订单收入；立即处理。题设未提供替代支付路径。', 'Все пользователи лишены возможности завершить ключевую покупку, что блокирует выручку; требуется немедленное исправление. Альтернативный способ оплаты не указан.')],
]

meta = {
    'document_id': 'REG-TEST-01', 'version':'1.0', 'date':'2026-10-21',
    'status':'DRAFT', 'execution_status':'NOT_RUN',
    'subject':'Internet shop registration, separate classroom scenario',
    'source_images': ['codex-clipboard-23e620b7-caa9-4763-a5d6-614b9f2805e1.jpg','codex-clipboard-30a1771c-f141-4e15-9913-dd8783dfffe8.jpg'],
    'tasks':tasks, 'requirements':requirements, 'equivalence':equivalence,
    'common':common, 'baseline':baseline, 'cases':cases, 'bug':bug, 'priorities':priorities,
}
(DATA/'content.json').write_text(json.dumps(meta,ensure_ascii=False,indent=2)+'\n')
for filename in meta['source_images']:
    shutil.copy2(Path('/var/folders/ff/kqw394c93dd7krl0z2hcqqdh0000gn/T')/filename, DATA/filename)

doc = Document()
sec=doc.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.6)
sec.left_margin=sec.right_margin=Inches(.65)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    s=doc.styles[name]
    s.font.name='Arial'; s.font.color.rgb=RGBColor(0,0,0)
    fonts=s.element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn('w:eastAsia'),'Arial Unicode MS')
    fonts.set(qn('w:ascii'),'Arial'); fonts.set(qn('w:hAnsi'),'Arial')
    s.paragraph_format.line_spacing=Pt(14)
    s.paragraph_format.space_before=Pt(0)
    s.paragraph_format.space_after=Pt(3)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Title'].font.size=Pt(20)
doc.styles['Title'].paragraph_format.line_spacing=Pt(27)
doc.styles['Subtitle'].font.size=Pt(13)
doc.styles['Heading 1'].font.size=Pt(15)
doc.styles['Heading 1'].paragraph_format.line_spacing=Pt(20)
doc.styles['Heading 2'].font.size=Pt(11.5)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(7)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(7)
# The bundled default template has a decorative Title paragraph border.
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)

inventory=[]
def para(text, style=None, size=None, keep=False, bold=False, parent=doc):
    p=parent.add_paragraph(style=style)
    p.paragraph_format.keep_together=True
    p.paragraph_format.keep_with_next=keep
    r=p.add_run(text); r.bold=bold
    if size:r.font.size=Pt(size)
    inventory.append(text)
    return p

def paired(texts, style=None, size=None, parent=doc):
    para(texts[0],style,size,True,parent=parent)
    para(texts[1],None,size,parent=parent)

def title(zh,ru):
    paired(pair(zh,ru),'Heading 1')

def table(headers,rows,widths,size=10.5):
    t=doc.add_table(rows=1,cols=len(headers))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for values in [headers,*rows]:
        row=t.rows[0] if values is headers else t.add_row()
        for i,(cell,v) in enumerate(zip(row.cells,values)):
            cell.width=Inches(widths[i]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text=''
            vals = v if isinstance(v,list) else [v]
            for j,txt in enumerate(vals):
                p=cell.paragraphs[0] if j==0 else cell.add_paragraph()
                p.paragraph_format.space_after=Pt(2)
                p.paragraph_format.line_spacing=Pt(13.5)
                p.paragraph_format.keep_together=True
                p.paragraph_format.keep_with_next=(j<len(vals)-1)
                r=p.add_run(txt); r.font.size=Pt(size)
                inventory.append(txt)
                if values is headers:r.bold=True; r.font.color.rgb=RGBColor(255,255,255)
            tcpr=cell._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'254D68' if values is headers else ('F2F6F9' if len(t.rows)%2==0 else 'FFFFFF'));tcpr.append(shade)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                x=OxmlElement('w:'+side);x.set(qn('w:w'),'80');x.set(qn('w:type'),'dxa');margins.append(x)
            tcpr.append(margins)
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        if values is headers:trpr.append(OxmlElement('w:tblHeader'))
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
    t._tbl.tblPr.append(borders)
    return t

paired(pair('网店注册表单测试设计与缺陷报告','Тестирование регистрации интернет магазина'),'Title')
paired(pair('REG-TEST-01  版本 1.0  2026年10月21日','REG-TEST-01  Версия 1.0  21 октября 2026 года'),size=10.5)
paired(pair('本练习完成四项任务，使用等价类和边界值设计注册测试。18 岁属于合法范围，应允许注册。','Работа охватывает четыре задания по классам эквивалентности и границам. Возраст 18 лет допустим, регистрация должна быть разрешена.'))
paired(pair('文档状态 DRAFT；测试用例 NOT_RUN；REG-BR-01 为题设情境报告。','Статус документа DRAFT; тест-кейсы NOT_RUN; REG-BR-01 составлен по учебной ситуации.'),size=10.5)
title('题目要求与表单规则','Задания и правила формы')
for task in tasks:paired(task,size=10.5)
table([pair('规则 ID','ID правила'),pair('内容','Содержание')],[[r[0],r[1:]] for r in requirements],[1.05,6.15],10.5)

doc.add_page_break()
title('任务 1 等价类与边界值','Задание 1 Классы эквивалентности и границы')
paired(pair('边界采用下界附近和上界附近各三个值；输入非法值时，其他字段均有效。','Проверяются по три значения около нижней и верхней границ; при недопустимом значении остальные поля корректны.'),size=10.5)
table([pair('字段','Поле'),pair('等价类','Классы эквивалентности'),pair('边界值及预期','Граничные значения и результат')],equivalence,[.85,3.45,2.9],10.5)
paired(pair('长度规则不增加“只能字母”等额外限制；全空格姓名、前后空格处理和特殊 Unicode 计数未规定，需另行澄清。','Не добавляются ограничения вроде «только буквы». Имя из пробелов, обрезка пробелов и подсчёт сложных Unicode-символов требуют уточнения.'),size=10.5)
paired(pair('记法 A×n 表示重复 n 次大写字母 A；A×n+1 表示在 n 个 A 后追加数字 1，符号本身不输入。','A×n означает n букв A; A×n+1 — n букв A и затем цифра 1. Обозначения не вводятся буквально.'),size=10.5)

doc.add_page_break()
title('任务 2 测试用例','Задание 2 Тест кейсы')
paired(common,size=10.5);paired(baseline,size=10.5)
paired(pair('以下共 12 个用例：3 个正向、9 个负向。参数组分别执行，不将多项非法输入合并提交。所有实际结果均为 NOT_RUN。','Всего 12 кейсов: 3 позитивных, 9 негативных. Наборы выполняются отдельно, разные ошибки не объединяются в одной отправке. Фактический результат каждого — NOT_RUN.'),size=10.5)
paired(pair('Email 正向用例使用受控测试邮箱。题设地址 test@mail.ru 仅保留在缺陷报告中，不向该地址进行实际操作。','В позитивных кейсах используется контролируемая почта. Адрес test@mail.ru сохранён только в баг-репорте; реальные действия с ним не выполняются.'),size=10.5)
paired(pair('页面错误文案除题设原文外只要求表达对应校验含义；邮件到达时限由测试环境约定，题目未给出具体秒数。','Кроме приведённой ошибки точные тексты не заданы: сообщение должно объяснять соответствующее правило. Срок доставки письма задаётся стендом, в условии он не указан.'),size=10.5)
for idx,c in enumerate(cases):
    if idx%2==0:
        doc.add_page_break()
        title(f'任务 2 测试用例 {idx+1} 至 {idx+2}',f'Задание 2 Тест кейсы {idx+1}–{idx+2}')
    para(f"{c['id']}  {c['title'][0]} / {c['title'][1]}",'Heading 2')
    paired(pair('类型与前置：'+('正向；' if c['type']=='POSITIVE' else '负向；')+c['preconditions'][0],
                'Тип и предусловия: '+('позитивный; ' if c['type']=='POSITIVE' else 'негативный; ')+c['preconditions'][1]),size=10.5)
    paired(pair('数据：'+c['values'][0],'Данные: '+c['values'][1]),size=10.5)
    paired(pair('步骤：'+c['steps'][0],'Шаги: '+c['steps'][1]),size=10.5)
    paired(pair('预期：'+c['expected'][0],'Ожидается: '+c['expected'][1]),size=10.5)
    para('实际 / Фактически: NOT_RUN',size=10)

doc.add_page_break()
title('任务 3 缺陷报告','Задание 3 Баг репорт')
table([pair('字段','Поле'),pair('报告内容','Содержание отчёта')],bug,[1.3,5.9],10.5)
paired(pair('复测建议：在隔离环境以受控邮箱验证 17 拒绝、18 与 19 接受、99 与 100 接受、101 拒绝，并检查成功提示及邮件。可能涉及比较运算的边界，但根因需查代码确认。','Повторная проверка на изолированном стенде с контролируемой почтой: 17 — отказ; 18, 19, 99, 100 — допуск; 101 — отказ; проверить сообщение и письмо. Возможна ошибка сравнения, но причина устанавливается по коду.'),size=10.5)

doc.add_page_break()
title('任务 4 严重程度与优先级','Задание 4 Серьёзность и приоритет')
paired(pair('严重程度描述影响，优先级描述处理先后。以下采用本作业的工作分级，数值越小越严重或越紧急；分级可随影响证据调整。','Серьёзность описывает влияние, приоритет — порядок исправления. Ниже используется рабочая шкала этой работы: меньшее число означает больший ущерб или срочность; оценка уточняется по свидетельствам.'),size=10.5)
paired(pair('严重程度：S1 Blocker 核心流程全面阻断；S2 Major 关键功能失败或崩溃；S3 Minor 局部功能轻微异常；S4 Trivial 外观或文字问题。','Серьёзность: S1 Blocker — полная блокировка ключевого процесса; S2 Major — отказ важной функции или аварийное завершение; S3 Minor — небольшой локальный дефект; S4 Trivial — оформление или текст.'),size=10.5)
paired(pair('优先级：P0 Urgent 立即；P1 High 尽快；P2 Normal 常规计划；P3 Low 可延后。','Приоритет: P0 Urgent — немедленно; P1 High — как можно скорее; P2 Normal — по плану; P3 Low — можно отложить.'),size=10.5)
table([pair('错误','Дефект'),pair('严重程度','Серьёзность'),pair('优先级','Приоритет'),pair('判定理由','Обоснование')],priorities,[1.6,1.25,1.15,3.2],10.5)
paired(pair('结论：优先处理全员支付失败；合法年龄被拒绝尽快修复；超大头像崩溃纳入正常修复计划；拼写问题可延后。','Порядок: сначала отказ оплаты у всех; затем отказ при допустимом возрасте; авария с большим аватаром — по плану; опечатку можно отложить.'),size=10.5)

out=ROOT/(STEM+'.docx')
# Use break-before headings so a full page cannot push a break-only paragraph
# onto a new blank page.
for element in list(doc.element.body):
    if element.tag==qn('w:p') and element.xpath('.//w:br[@w:type="page"]'):
        following=element.getnext()
        if following is not None and following.tag==qn('w:p'):
            Paragraph(following,doc).paragraph_format.page_break_before=True
            element.getparent().remove(element)
doc.core_properties.title='网店注册表单测试设计与缺陷报告'
doc.core_properties.subject='Классы эквивалентности, граничные значения, тест-кейсы и баг-репорт'
doc.core_properties.author=''
doc.save(out)

# Markdown is an editable companion with the same substantive content.
md=['# 网店注册表单测试设计与缺陷报告 / Тестирование регистрации интернет магазина','',
    'REG-TEST-01 · 1.0 · 2026-10-21 · DRAFT · NOT_RUN','']
def mp(texts): md.extend([texts[0],'',texts[1],''])
def mt(headers,rows):
    def cell(v):return '<br>'.join(v).replace('\n','<br>') if isinstance(v,list) else v.replace('\n','<br>')
    md.append('| '+' | '.join(cell(v) for v in headers)+' |')
    md.append('| '+' | '.join('---' for _ in headers)+' |')
    for row in rows:md.append('| '+' | '.join(cell(v) for v in row)+' |')
    md.append('')
for task in tasks:mp(task)
mt([['规则 ID','ID'],['内容','Содержание']],[[r[0],r[1:]] for r in requirements])
md.extend(['## 1 等价类与边界值 / Классы эквивалентности и границы',''])
mt([['字段','Поле'],['等价类','Классы'],['边界值','Границы']],equivalence)
mp(pair('A×n 表示 n 个 A；A×n+1 表示 n 个 A 后追加数字 1。','A×n — n букв A; A×n+1 — n букв A и затем цифра 1.'))
md.extend(['## 2 测试用例 / Тест кейсы','']);mp(common);mp(baseline)
for c in cases:
    md.extend([f"### {c['id']} {c['title'][0]} / {c['title'][1]}",''])
    md.extend([f"{c['type']} · NOT_RUN",''])
    for key,zh,ru in [('preconditions','前置条件','Предусловия'),('values','数据','Данные'),('steps','步骤','Шаги'),('expected','预期','Ожидается')]:mp(pair(zh+'：'+c[key][0],ru+': '+c[key][1]))
md.extend(['## 3 缺陷报告 / Баг репорт',''])
mt([['字段','Поле'],['报告内容','Содержание']],bug)
md.extend(['## 4 严重程度与优先级 / Серьёзность и приоритет',''])
mp(pair('严重程度：S1 阻断、S2 严重、S3 轻微功能问题、S4 外观文字问题；优先级：P0 立即、P1 尽快、P2 常规、P3 可延后。','Серьёзность: S1 Blocker, S2 Major, S3 Minor, S4 Trivial; приоритет: P0 Urgent, P1 High, P2 Normal, P3 Low.'))
mt([['错误','Дефект'],['严重程度','Серьёзность'],['优先级','Приоритет'],['理由','Обоснование']],priorities)
(ROOT/(STEM+'.md')).write_text('\n'.join(md)+'\n')

# Export in document order so both editable formats contain every final note.
md=[]
for element in doc.element.body:
    if element.tag==qn('w:p'):
        p=Paragraph(element,doc)
        if not p.text.strip():continue
        prefix={'Title':'# ','Heading 1':'## ','Heading 2':'### '}.get(p.style.name,'')
        md.extend([prefix+p.text,''])
    elif element.tag==qn('w:tbl'):
        t=Table(element,doc)
        for n,row in enumerate(t.rows):
            md.append('| '+' | '.join(c.text.replace('\n','<br>').replace('|','\\|') for c in row.cells)+' |')
            if n==0:md.append('| '+' | '.join('---' for _ in row.cells)+' |')
        md.append('')
(ROOT/(STEM+'.md')).write_text('\n'.join(md)+'\n')

# Reopen the deliverable; verify the normalized inventory and coverage counts.
reopened=Document(out)
all_text='\n'.join([p.text for p in reopened.paragraphs]+[c.text for t in reopened.tables for r in t.rows for c in r.cells])
checks={
    'inventory_all_present': all(v in all_text for v in inventory),
    'four_tasks': len(tasks)==4,
    'three_equivalence_fields': len(equivalence)==3,
    'template_three_columns': len(reopened.tables[1].columns)==3 and len(reopened.tables[1].rows)==4,
    'twelve_cases': len(cases)==12 and len({c['id'] for c in cases})==12,
    'positive_minimum': sum(c['type']=='POSITIVE' for c in cases)>=1,
    'negative_minimum': sum(c['type']=='NEGATIVE' for c in cases)>=3,
    'all_cases_complete': all(all(c.get(k) for k in ['id','title','preconditions','steps','expected']) for c in cases),
    'all_cases_not_run': all(c['actual']=='NOT_RUN' for c in cases),
    'equivalence_classes_covered': set(['N-V','N-E','N-S','N-L','P-V','P-E','P-S','P-L','P-D','A-V','A-E','A-S','A-L','A-F','A-T']).issubset({v for c in cases for v in c['coverage']}),
    'situation_exact_data': all(v in all_text for v in ['Ли','test@mail.ru','Password1','Возраст должен быть больше 18','Windows 11']),
    'three_triage_items':len(priorities)==3,
    'docx_zip_integrity': ZipFile(out).testzip() is None,
}
verification={'checks':checks,'all_pass':all(checks.values()),'case_counts':{'positive':3,'negative':9,'total':12},'table_dimensions':[[len(t.rows),len(t.columns)] for t in reopened.tables], 'inventory_items':len(inventory), 'visual_review':'PENDING','business_execution':'NOT_RUN','sha256':hashlib.sha256(out.read_bytes()).hexdigest()}
(DATA/'verification.json').write_text(json.dumps(verification,ensure_ascii=False,indent=2)+'\n')
assert verification['all_pass'],verification
print(json.dumps({'output':str(out),'checks':checks,'inventory_items':len(inventory)},ensure_ascii=False))
