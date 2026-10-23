from pathlib import Path
from zipfile import ZipFile
import hashlib
import json
import shutil

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
DATA, WORK = ROOT / 'data', ROOT / 'work'
DATA.mkdir(exist_ok=True)
STEM = 'CampusFix_任务三缺陷分诊_中俄逐段对照_v1'
SOURCE = Path('/var/folders/ff/kqw394c93dd7krl0z2hcqqdh0000gn/T/codex-clipboard-a61fbc9b-28d2-4207-826b-7bccc3929cc7.jpg')
PREVIOUS = ROOT.parent / '2026-10-19-campusfix-evidence-reports-v1/data/content.json'
previous_bytes = PREVIOUS.read_bytes()
previous = json.loads(previous_bytes)

rows = [
    {
        'evidence': 'I-1', 'bug': 'BUG-0201', 'decision': 'DEFECT',
        'decision_text': ['缺陷 待复核', 'Дефект Проверить'],
        'severity': 'S2', 'severity_text': ['S2 重大', 'S2 Значительная'],
        'priority': 'P1', 'priority_text': ['P1 高', 'P1 Высокий'],
        'reason': ['违反 AC-CF-01，Firefox 登录受阻；Chrome 可绕行，未达全站阻断。',
                   'Нарушен AC-CF-01: вход в Firefox недоступен; есть обход через Chrome, общего отказа нет.'],
    },
    {
        'evidence': 'I-2', 'bug': 'BUG-0202', 'decision': 'DEFECT',
        'decision_text': ['缺陷', 'Дефект'],
        'severity': 'S3', 'severity_text': ['S3 一般', 'S3 Умеренная'],
        'priority': 'P2', 'priority_text': ['P2 普通', 'P2 Обычный'],
        'reason': ['违反 AC-CF-02，保存空描述污染数据；有效申请仍可提交。',
                   'Нарушен AC-CF-02: сохраняется пустое описание; корректные заявки отправляются.'],
    },
    {
        'evidence': 'I-3', 'bug': 'BUG-0203', 'decision': 'DEFECT',
        'decision_text': ['缺陷', 'Дефект'],
        'severity': 'S1', 'severity_text': ['S1 严重', 'S1 Критическая'],
        'priority': 'P0', 'priority_text': ['P0 立即', 'P0 Немедленно'],
        'reason': ['违反 AC-CF-03，普通学生可读他人详情，跨越数据访问边界。',
                   'Нарушен AC-CF-03: студент читает чужие подробности, нарушена изоляция данных.'],
    },
    {
        'evidence': 'I-4', 'bug': 'BUG-0204', 'decision': 'DEFECT',
        'decision_text': ['缺陷', 'Дефект'],
        'severity': 'S3', 'severity_text': ['S3 一般', 'S3 Умеренная'],
        'priority': 'P2', 'priority_text': ['P2 普通', 'P2 Обычный'],
        'reason': ['违反 AC-CF-04，筛选与结果不符；回到首页筛选可绕行。',
                   'Нарушен AC-CF-04: неверный список; обход — повторить фильтр со страницы 1.'],
    },
    {
        'evidence': 'I-5', 'bug': None, 'decision': 'DUPLICATE', 'duplicate_of': 'BUG-0201',
        'decision_text': ['重复 归并 BUG-0201', 'Дубликат к BUG-0201'],
        'severity': 'S2', 'severity_text': ['S2 沿用主单', 'S2 По основной'],
        'priority': 'P1', 'priority_text': ['P1 随主单', 'P1 По основной'],
        'reason': ['与主单环境、步骤及现象一致，无新触发条件；证据并入主单。',
                   'Совпадают среда, шаги и симптом; новых условий нет. Улику добавить в основную запись.'],
    },
    {
        'evidence': 'I-6', 'bug': None, 'decision': 'NOT_A_DEFECT',
        'decision_text': ['非缺陷', 'Не дефект'],
        'severity': None, 'severity_text': ['不适用', 'Не применимо'],
        'priority': None, 'priority_text': ['不排缺陷期', 'Не планируется'],
        'reason': ['未登录访问返回 401 并引导登录，符合补充规则 AUTH-CF-01。',
                   'Без сеанса ответ 401 ведёт ко входу, как предусмотрено правилом AUTH-CF-01.'],
    },
    {
        'evidence': 'I-7', 'bug': 'BUG-0205', 'decision': 'DEFECT', 'new_bug_proposed': True,
        'decision_text': ['缺陷 拟建 BUG-0205', 'Дефект Создать BUG-0205'],
        'severity': 'S4', 'severity_text': ['S4 轻微', 'S4 Незначительная'],
        'priority': 'P3', 'priority_text': ['P3 低', 'P3 Низкий'],
        'reason': ['按钮错字违反 UI-CF-01；提交正常，无数据或访问权限影响。',
                   'Опечатка нарушает UI-CF-01; отправка работает, данные и права не затронуты.'],
    },
]

supplements = [
    {
        'evidence': 'I-5',
        'title': ['另一份相同 SSO 登录失败报告', 'Повторный отчёт о том же сбое входа SSO'],
        'content': ['教学设定：另一份报告同样使用 CampusFix 1.0.3、Windows 11、Firefox 131 和 student09；输入 SSO 密码后无提示返回登录表单，Chrome 130 正常，与 BUG-0201 的最小步骤及可见结果一致。',
                    'Учебное условие: второй отчёт относится к CampusFix 1.0.3, Windows 11, Firefox 131 и student09. После ввода пароля SSO происходит возврат без сообщения; в Chrome 130 вход работает. Минимальные шаги и видимый результат совпадают с BUG-0201.'],
        'handling': ['归并到 BUG-0201，保留 I-5 为附加证据；不独立排期。若复核发现不同触发条件或原因，再拆分。',
                     'Связать с BUG-0201, сохранить I-5 как дополнительную улику и не планировать отдельно. Если проверка выявит иные условия или причину, разделить записи.'],
    },
    {
        'evidence': 'I-6',
        'title': ['匿名访问受保护资源被拒绝', 'Отказ анонимному пользователю в доступе к защищённому ресурсу'],
        'content': ['教学设定：未登录时请求申请列表，服务器返回 401，界面引导到登录页，响应不含申请数据。补充约定 AUTH-CF-01 明确要求此行为；它与 I-3 中已登录学生越权读取他人申请是不同条件。',
                    'Учебное условие: запрос списка заявок без входа получает 401, интерфейс направляет ко входу, данных заявок в ответе нет. Дополнительное правило AUTH-CF-01 требует именно этого. Условия отличаются от I-3, где вошедший студент читает чужую заявку.'],
        'handling': ['判为非缺陷并记录符合预期；严重程度与修复优先级均不适用。',
                     'Признать ожидаемым поведением и не дефектом; серьёзность и приоритет исправления не применяются.'],
    },
    {
        'evidence': 'I-7',
        'title': ['提交按钮俄语标签存在拼写错误', 'Опечатка в русской надписи кнопки отправки'],
        'content': ['教学设定：新建申请的按钮显示 «Отпавить»，UI-CF-01 规定应为 «Отправить»。使用有效表单点击按钮仍能成功创建申请；未观察到数据错误。',
                    'Учебное условие: в форме новой заявки кнопка подписана «Отпавить», а UI-CF-01 задаёт «Отправить». Нажатие с корректными полями успешно создаёт заявку; ошибок данных не наблюдается.'],
        'handling': ['建议分配新编号 BUG-0205，修正标签并复查俄语界面；严重程度 S4，优先级 P3。',
                     'Предложить новый ID BUG-0205, исправить надпись и проверить русский интерфейс; серьёзность S4, приоритет P3.'],
    },
]

severity_scale = [
    ('S1', ['严重', 'Критическая'], ['数据机密性、完整性或全局关键服务受到严重破坏。', 'Серьёзно нарушены конфиденциальность, целостность или общая доступность ключевого сервиса.']),
    ('S2', ['重大', 'Значительная'], ['受支持环境中的关键流程不可用，存在局部绕行办法。', 'Ключевой сценарий недоступен в поддерживаемой среде; возможен ограниченный обход.']),
    ('S3', ['一般', 'Умеренная'], ['局部功能或数据质量不符合要求，主要流程仍可完成。', 'Локальная функция или качество данных нарушены; основной сценарий остаётся доступным.']),
    ('S4', ['轻微', 'Незначительная'], ['文字或外观问题，不影响功能、数据和权限。', 'Ошибка текста или оформления без влияния на функцию, данные и права.']),
]
priority_scale = [
    ('P0', ['立即处置', 'Немедленно'], ['优先控制风险，修复并验证后再决定该版本能否发布。', 'Сначала ограничить риск; решение о выпуске — после исправления и проверки.']),
    ('P1', ['高', 'Высокий'], ['优先安排修复，发布前验证关键流程。', 'Исправлять в первую очередь и проверить ключевой сценарий до выпуска.']),
    ('P2', ['普通', 'Обычный'], ['按计划与迭代容量安排处理。', 'Планировать с учётом очереди и возможностей итерации.']),
    ('P3', ['低', 'Низкий'], ['进入后续文字或界面改进计划。', 'Включить в последующее улучшение текста или интерфейса.']),
]

doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Inches(11), Inches(8.5)
sec.top_margin = sec.bottom_margin = Inches(0.43)
sec.left_margin = sec.right_margin = Inches(0.55)
for name in ('Normal', 'Title', 'Heading 1'):
    s = doc.styles[name]
    s.font.name = 'Arial'
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Arial Unicode MS')
    s.paragraph_format.line_spacing = 1.0
    s.paragraph_format.space_after = Pt(3)
doc.styles['Normal'].font.size = Pt(10)
doc.styles['Title'].font.size = Pt(15)
doc.styles['Heading 1'].font.size = Pt(11)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(5)
pairs = []


def pair(zh, ru, cell=None, style=None, size=10, bold=False):
    p = doc.add_paragraph(style=style) if cell is None else (cell.paragraphs[0] if not cell.text else cell.add_paragraph())
    if style:
        p.style = style
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.keep_together = True
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(zh)
    r.bold, r.font.size = bold, Pt(size)
    q = doc.add_paragraph(style=style) if cell is None else cell.add_paragraph()
    if style:
        q.style = style
    q.paragraph_format.keep_with_next = style in ('Title', 'Heading 1')
    q.paragraph_format.keep_together = True
    q.paragraph_format.space_before = Pt(0)
    q.paragraph_format.space_after = Pt(3)
    r = q.add_run(ru)
    r.bold, r.font.size = bold, Pt(size)
    pairs.append((p, q, zh, ru))


def table(headers, widths, values, size=9.5):
    t = doc.add_table(rows=1, cols=len(widths))
    t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    for cell, text in zip(t.rows[0].cells, headers):
        pair(*text, cell=cell, bold=True, size=size)
        sh = OxmlElement('w:shd')
        sh.set(qn('w:fill'), 'DAEEF0')
        cell._tc.get_or_add_tcPr().append(sh)
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for data_row in values:
        cells = t.add_row().cells
        for cell, value in zip(cells, data_row):
            if isinstance(value, str):
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(3)
                p.add_run(value).font.size = Pt(size)
            else:
                pair(*value, cell=cell, size=size)
    for row in t.rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    props = t._tbl.tblPr
    margins = OxmlElement('w:tblCellMar')
    for side in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement('w:' + side)
        e.set(qn('w:w'), '55')
        e.set(qn('w:type'), 'dxa')
        margins.append(e)
    props.append(margins)
    borders = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + side)
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), '4')
        e.set(qn('w:color'), 'D9D9D9')
        borders.append(e)
    props.append(borders)
    return t


pair('CampusFix 任务三 缺陷分诊', 'CampusFix Задание 3 Разбор и триаж', style='Title', size=15)
pair('课堂示例  |  v1  |  DRAFT', 'Учебные примеры  |  v1  |  DRAFT', size=9)
pair('I-1～I-4 沿用任务二，I-5～I-7 为补充教学设定。下表为初步分诊建议；真实环境测试状态仍为 NOT_RUN。',
     'I-1–I-4 взяты из задания 2; I-5–I-7 — дополнительные учебные условия. Решения предварительные; реальное тестирование — NOT_RUN.', size=9)
triage_table = table([
    ['证据与缺陷', 'Улика и дефект'], ['判断', 'Решение'], ['严重程度', 'Серьёзность'],
    ['优先级', 'Приоритет'], ['理由', 'Обоснование'],
], [1.10, 1.77, 1.26, 1.23, 4.54], [
    [x['evidence'] + ('\n' + x['bug'] if x['bug'] and not x.get('new_bug_proposed') else ''),
     x['decision_text'], x['severity_text'], x['priority_text'], x['reason']] for x in rows
], size=9.3)

doc.add_page_break()
pair('补充证据与判断依据', 'Дополнительные улики и основания решений', style='Title', size=15)
pair('以下 I-5～I-7 为补充教学示例，用于说明重复、非缺陷和轻微缺陷的判定。',
     'Учебные I-5–I-7 показывают, как отличить дубликат, ожидаемое поведение и незначительный дефект.', size=9.5)
for item in supplements:
    pair(item['evidence'] + ' ' + item['title'][0], item['evidence'] + ' ' + item['title'][1], style='Heading 1', bold=True, size=11)
    pair(*item['content'], size=10)
    pair('处理：' + item['handling'][0], 'Решение: ' + item['handling'][1], size=10)
pair('建议顺序', 'Рекомендуемая очередь', style='Heading 1', bold=True, size=10.5)
p = doc.add_paragraph('P0  BUG-0203  →  P1  BUG-0201  →  P2  BUG-0202 / BUG-0204  →  P3  BUG-0205')
p.runs[0].font.size = Pt(10)
pair('I-5 随 BUG-0201 合并处理，不独立排期；I-6 按预期行为记录。复测若改变影响范围，再调整分级与优先级。',
     'I-5 обрабатывается вместе с BUG-0201 без отдельного планирования; I-6 фиксируется как ожидаемое поведение. Если проверка изменит оценку влияния, пересмотреть уровни.', size=9)

doc.add_page_break()
pair('分级标准与处理顺序', 'Критерии оценки и порядок обработки', style='Title', size=15)
pair('本页定义仅用于本课堂练习。严重程度衡量影响，优先级决定处理顺序，两者分别判断。',
     'Эта шкала используется в упражнении. Серьёзность описывает влияние, приоритет задаёт порядок обработки; оценки независимы.', size=9.5)
pair('严重程度', 'Серьёзность', style='Heading 1', bold=True, size=10.5)
table([['级别', 'Уровень'], ['含义', 'Значение'], ['判定依据', 'Критерий']], [0.7, 1.9, 7.3],
      [[code, label, definition] for code, label, definition in severity_scale], size=9.5)
pair('优先级', 'Приоритет', style='Heading 1', bold=True, size=10.5)
table([['级别', 'Уровень'], ['处理安排', 'Порядок'], ['安排依据', 'Основание']], [0.7, 1.9, 7.3],
      [[code, label, definition] for code, label, definition in priority_scale], size=9.5)

for element in (doc.styles.element, doc.element):
    for border in list(element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
doc.core_properties.title = 'CampusFix 任务三 缺陷分诊'
doc.core_properties.subject = 'I-1 至 I-7 判断 严重程度 优先级 理由 中俄逐段对照'
doc.core_properties.author = 'StudySpace Coursework'
out = ROOT / (STEM + '.docx')
doc.save(out)
shutil.copy2(SOURCE, DATA / 'source-task3.jpg')

content = {
    'version': 'v1', 'date': '2026-10-19', 'status': 'DRAFT',
    'source_context': 'Continuation of user-authorized CampusFix classroom examples',
    'prior_authorization': previous['authorization'],
    'prior_source_sha256': hashlib.sha256(previous_bytes).hexdigest(),
    'rows': rows, 'supplements': supplements,
    'severity_scale': severity_scale, 'priority_scale': priority_scale,
    'actual_business_tests': 'NOT_RUN', 'deliverables': ['docx'],
    'language_order': ['zh', 'ru'], 'external_issue_mutations': False,
}
(DATA / 'content.json').write_text(json.dumps(content, ensure_ascii=False, indent=2) + '\n')
with ZipFile(out) as z:
    assert z.testzip() is None
assert len(triage_table.rows) == 8 and len(triage_table.columns) == 5
assert [x['evidence'] for x in rows] == [f'I-{n}' for n in range(1, 8)]
assert [(x['evidence'], x['bug']) for x in rows[:4]] == [(x['evidence_id'], x['bug_id']) for x in previous['cases']]
assert rows[4]['duplicate_of'] == rows[0]['bug']
assert rows[4]['severity'] == rows[0]['severity'] and rows[4]['priority'] == rows[0]['priority']
assert rows[5]['severity'] is None and rows[5]['priority'] is None
assert len({x['bug'] for x in rows if x['bug']}) == 5
for p, q, zh, ru in pairs:
    assert p.text == zh and q.text == ru and p._p.getnext() is q._p
assert PREVIOUS.read_bytes() == previous_bytes
assert not list(ROOT.glob('*.md'))
coverage = {
    'source_title': 'Задание 3: разбор (триаж)',
    'headers': ['Улика / дефект', 'Решение (дефект / дубликат / не дефект)', 'Серьёзность', 'Приоритет', 'Обоснование'],
    'rows': [x['evidence'] for x in rows],
    'original_bug_ids': [x['bug'] for x in rows[:4]],
    'supplemental_teaching_evidence': ['I-5', 'I-6', 'I-7'],
    'source_sha256': hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
    'timer_15_00': 'Classroom time allowance only',
}
(DATA / 'source-coverage.json').write_text(json.dumps(coverage, ensure_ascii=False, indent=2) + '\n')
check = {
    'task_table': 'PASS', 'columns': 5, 'filled_rows': 7,
    'preserved_prior_bug_ids': 'PASS', 'duplicate_target': 'BUG-0201',
    'non_defect_has_no_repair_priority': 'PASS',
    'decision_counts': {key: sum(x['decision'] == key for x in rows) for key in ('DEFECT', 'DUPLICATE', 'NOT_A_DEFECT')},
    'bilingual_pairs': len(pairs), 'zh_above_ru': 'PASS', 'docx_zip': 'PASS',
    'previous_task_data_unchanged': 'PASS', 'render_review': 'PENDING',
    'actual_business_tests': 'NOT_RUN', 'word_only': True,
    'sha256': hashlib.sha256(out.read_bytes()).hexdigest(),
}
(DATA / 'verification.json').write_text(json.dumps(check, ensure_ascii=False, indent=2) + '\n')
fonts = ROOT.parent / '2026-10-19-bug-report-review-v1/work/fonts'
lo_fonts = '/Users/zhi/.cache/codex-runtimes/codex-primary-runtime/dependencies/native/libreoffice-headless/libreoffice/LibreOfficeDev.app/Contents/Resources/fonts/truetype'
(WORK / 'fonts.conf').write_text(f'<?xml version="1.0"?><!DOCTYPE fontconfig SYSTEM "fonts.dtd"><fontconfig><dir>{fonts}</dir><dir>{lo_fonts}</dir><cachedir>{WORK / "fontcache"}</cachedir></fontconfig>')
print(f'PASS: 7 triage rows, 5 columns, consistent original IDs, valid duplicate target, {len(pairs)} Chinese-Russian pairs')
print(out)
