from datetime import datetime
from pathlib import Path
from zipfile import ZipFile
import hashlib
import json
import shutil

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
DATA, WORK = ROOT / 'data', ROOT / 'work'
STEM = 'CampusFix_任务四修复与确认_中俄逐段对照_v1'
SOURCE = Path('/var/folders/ff/kqw394c93dd7krl0z2hcqqdh0000gn/T/codex-clipboard-99cb6576-3cf9-4e99-ad37-91b2265b6e49.jpg')
execution = json.loads((DATA / 'execution.json').read_text())
before, regression, after = execution['runs']
assert (before['tests_run'], before['passed'], before['failures'], before['exit_code']) == (18, 17, 1, 1)
assert (regression['tests_run'], regression['passed'], regression['exit_code']) == (1, 1, 0)
assert (after['tests_run'], after['passed'], after['exit_code']) == (18, 18, 0)

steps = [
    [
        ['运行测试，确认 1001 字符用例失败。实测：18 个测试中 17 个通过、1 个失败，退出码 1；失败项为 TC-CF-05。',
         'Запустить тесты и подтвердить падение на 1001 символе. Получено: 17 из 18 тестов прошли, 1 не прошёл, код выхода 1; упал TC-CF-05.'],
    ],
    [
        ['查找并修复 Validate 中的错误：上限条件由 > MAX_DESCRIPTION_LENGTH + 1 改为 > MAX_DESCRIPTION_LENGTH，正确拒绝超过 1000 字符的文本。',
         'Найти и исправить ошибку в Validate: заменить > MAX_DESCRIPTION_LENGTH + 1 на > MAX_DESCRIPTION_LENGTH, чтобы отклонять текст длиннее 1000 символов.'],
    ],
    [
        ['运行回归测试和全部测试。实测：1001 字符回归 1/1 通过，全部测试 18/18 通过，两次退出码均为 0；测试文件未改动。',
         'Запустить регрессионный тест и все тесты. Получено: регрессия на 1001 символе — 1/1, весь набор — 18/18, оба кода выхода 0; файл тестов не изменён.'],
    ],
    [
        ['填写确认记录：本文件末页 CONF-CF-01 已记录执行环境、时间、命令、结果、证据与确认范围。',
         'Заполнить запись подтверждения: на последней странице CONF-CF-01 указаны среда, время, команды, результаты, свидетельства и область подтверждения.'],
    ],
]

doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Inches(11), Inches(8.5)
sec.top_margin = sec.bottom_margin = Inches(0.45)
sec.left_margin = sec.right_margin = Inches(0.55)
for name in ('Normal', 'Title', 'Heading 1'):
    s = doc.styles[name]
    s.font.name = 'Arial'
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Arial Unicode MS')
    s.paragraph_format.line_spacing = 1.0
    s.paragraph_format.space_after = Pt(3)
doc.styles['Normal'].font.size = Pt(10)
doc.styles['Title'].font.size = Pt(15)
doc.styles['Heading 1'].font.size = Pt(10.5)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(4)
pairs = []


def pair(zh, ru, cell=None, style=None, size=10, bold=False):
    p = doc.add_paragraph(style=style) if cell is None else (cell.paragraphs[0] if not cell.text else cell.add_paragraph())
    if style:
        p.style = style
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.keep_together = True
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(zh)
    r.bold, r.font.size = bold, Pt(size)
    q = doc.add_paragraph(style=style) if cell is None else cell.add_paragraph()
    if style:
        q.style = style
    q.paragraph_format.keep_with_next = style in ('Title', 'Heading 1')
    q.paragraph_format.keep_together = True
    q.paragraph_format.space_before = Pt(0)
    q.paragraph_format.space_after = Pt(4)
    r = q.add_run(ru)
    r.bold, r.font.size = bold, Pt(size)
    pairs.append((p, q, zh, ru))


def table(headers, widths, values, size=10):
    t = doc.add_table(rows=1, cols=len(widths))
    t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    for cell, value in zip(t.rows[0].cells, headers):
        pair(*value, cell=cell, bold=True, size=size)
        sh = OxmlElement('w:shd')
        sh.set(qn('w:fill'), 'DAEEF0')
        cell._tc.get_or_add_tcPr().append(sh)
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for data_row in values:
        cells = t.add_row().cells
        for cell, value in zip(cells, data_row):
            if isinstance(value, str):
                cell.paragraphs[0].add_run(value).font.size = Pt(size)
            else:
                pair(*value, cell=cell, size=size)
    for row in t.rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    props = t._tbl.tblPr
    mar = OxmlElement('w:tblCellMar')
    for side in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement('w:' + side)
        e.set(qn('w:w'), '75')
        e.set(qn('w:type'), 'dxa')
        mar.append(e)
    props.append(mar)
    borders = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + side)
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), '4')
        e.set(qn('w:color'), 'D9D9D9')
        borders.append(e)
    props.append(borders)
    return t


def code_line(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.font.name = 'Courier New'
    r.font.size = Pt(9)


pair('CampusFix 任务四 修复与确认', 'CampusFix Задание 4 Исправление и подтверждение', style='Title', size=15)
pair('独立教学示例实测  |  v1  |  文档 DRAFT  |  示例验证 PASS',
     'Проверка отдельного учебного примера  |  v1  |  Документ DRAFT  |  Пример PASS', size=9)
pair('本次验证对象为自建 Python Validate 示例，已实际执行测试并保存日志；原 CampusFix 程序未提供，其测试状态为 NOT_RUN。',
     'Проверен отдельный учебный Validate на Python; тесты реально выполнены, журналы сохранены. Исходный CampusFix не предоставлен, его тесты имеют статус NOT_RUN.', size=9.5)
task_table = table([['步骤', 'Шаг'], ['完成内容与结果', 'Что сделано и результат']], [0.65, 9.25],
                   [[str(i), item[0]] for i, item in enumerate(steps, 1)], size=10.5)
pair('结论：1001 字符从错误接受变为正确拒绝；1000 字符仍被接受，18 个测试全部通过。',
     'Вывод: 1001 символ теперь корректно отклоняется, 1000 символов принимается; все 18 тестов прошли.', size=10)

doc.add_page_break()
pair('错误原因与回归结果', 'Причина ошибки и результаты регрессии', style='Title', size=15)
pair('规则与原因', 'Правило и причина', style='Heading 1', bold=True, size=10.5)
pair('FR-CF-02 要求去除首尾空白后保留 1～1000 个字符。教学故障版本把拒绝条件写成长度 > 1001，因此恰好 1001 个字符被错误放行。该示例按 Unicode 码点计数。',
     'FR-CF-02 допускает от 1 до 1000 символов после удаления пробелов по краям. Учебная ошибочная версия отклоняет только длину > 1001 и пропускает ровно 1001 символ. В примере считаются кодовые точки Unicode.', size=10)
pair('修复前的判断', 'Условие до исправления', bold=True, size=9.5)
code_line('if len(text) > MAX_DESCRIPTION_LENGTH + 1:')
pair('修复后的判断', 'Условие после исправления', bold=True, size=9.5)
code_line('if len(text) > MAX_DESCRIPTION_LENGTH:')
pair('只修改一行校验逻辑，并清理过时说明；测试文件与测试数据保持不变。',
     'Изменена одна строка логики проверки и удалено устаревшее пояснение; тесты и тестовые данные остались прежними.', size=9.5)
table([['字符数', 'Длина'], ['预期', 'Ожидается'], ['修复前实测', 'До исправления'], ['修复后实测', 'После исправления']],
      [0.8, 2.7, 3.2, 3.2], [
          ['999', ['接受', 'Принять'], ['接受', 'Принято'], ['接受', 'Принято']],
          ['1000', ['接受', 'Принять'], ['接受', 'Принято'], ['接受', 'Принято']],
          ['1001', ['拒绝', 'Отклонить'], ['错误接受', 'Ошибочно принято'], ['拒绝', 'Отклонено']],
          ['1002', ['拒绝', 'Отклонить'], ['拒绝', 'Отклонено'], ['拒绝', 'Отклонено']],
      ], size=9.5)
pair('测试覆盖', 'Покрытие тестами', style='Heading 1', bold=True, size=10.5)
pair('18 个测试覆盖：1/999/1000/1001/1002/5000 长度；空串与纯空白；None、整数和 bytes；中文与俄文；先去首尾空白再计数及保留内部空白。',
     '18 тестов покрывают длины 1/999/1000/1001/1002/5000, пустую строку и пробельные символы, None, число и bytes, китайский и русский текст, подсчёт после обрезки краёв и сохранение внутренних пробелов.', size=9.5)

doc.add_page_break()
pair('确认记录 CONF-CF-01', 'Запись подтверждения CONF-CF-01', style='Title', size=15)
after_time = datetime.fromisoformat(after['started_at']).strftime('%Y-%m-%d %H:%M:%S')
record_values = [
    [['对象与关联', 'Объект и связи'],
     ['demo/validator.py 的 Validate；BUG-0206；FR-CF-02 / AC-CF-05 / TC-CF-05。',
      'Validate в demo/validator.py; BUG-0206; FR-CF-02 / AC-CF-05 / TC-CF-05.']],
    [['执行环境与时间', 'Среда и время'],
     [f"Python {execution['python_version']}；macOS arm64；{after_time}，Asia/Shanghai（UTC+08:00）。",
      f"Python {execution['python_version']}; macOS arm64; {after_time}, Asia/Shanghai (UTC+08:00)."]],
    [['失败确认', 'Подтверждение сбоя'],
     ['修复前：17 通过、1 失败，18 个测试；退出码 1。证据：data/before-full-suite.log。',
      'До исправления: 17 прошли, 1 упал, всего 18; код выхода 1. Журнал: data/before-full-suite.log.']],
    [['回归确认', 'Подтверждение регрессии'],
     ['TC-CF-05 通过，1/1；退出码 0。证据：data/after-regression.log。',
      'TC-CF-05 прошёл, 1/1; код выхода 0. Журнал: data/after-regression.log.']],
    [['全部测试', 'Полный набор'],
     ['18/18 通过，0 失败；退出码 0。证据：data/after-full-suite.log。',
      '18/18 прошли, 0 отказов; код выхода 0. Журнал: data/after-full-suite.log.']],
    [['确认结论', 'Итог подтверждения'],
     ['独立示例修复确认 PASS；原 CampusFix 与 StudySpace 业务测试仍为 NOT_RUN。',
      'Исправление отдельного примера подтверждено: PASS; бизнес-тесты исходных CampusFix и StudySpace остаются NOT_RUN.']],
    [['执行记录', 'Запись выполнения'],
     ['Codex 自动执行，结果来自真实进程输出。执行命令、代码与日志 SHA-256 见 data/execution.json；代码差异见 data/validator.patch。',
      'Автоматически выполнено Codex; результаты получены из реального вывода процессов. Команды и SHA-256 кода и журналов: data/execution.json; изменение: data/validator.patch.']],
]
record_table = table([['字段', 'Поле'], ['记录内容', 'Запись']], [2.0, 7.9], record_values, size=9.5)
pair('复跑位置：本交付目录下 demo 文件夹；使用执行记录中的 Python 解释器。',
     'Для повторного запуска перейти в папку demo этого материала и использовать интерпретатор из записи выполнения.', size=9)
code_line('python -m unittest tests.test_validator.ValidateTests.test_tc_cf_05_rejects_1001_characters -v')
code_line('python -m unittest discover -s tests -v')

for element in (doc.styles.element, doc.element):
    for b in list(element.iter(qn('w:pBdr'))):
        b.getparent().remove(b)
doc.core_properties.title = 'CampusFix 任务四 修复与确认'
doc.core_properties.subject = '独立教学示例 1001 字符回归测试与确认记录 中俄逐段对照'
doc.core_properties.author = 'StudySpace Coursework'
out = ROOT / (STEM + '.docx')
doc.save(out)
shutil.copy2(SOURCE, DATA / 'source-task4.jpg')

content = {
    'version': 'v1', 'status': 'DRAFT', 'task': 4,
    'scope': execution['scope'], 'authorization': execution['authorization'],
    'steps': steps, 'bug_id': 'BUG-0206', 'confirmation_id': 'CONF-CF-01',
    'requirement': 'FR-CF-02', 'criterion': 'AC-CF-05', 'regression_case': 'TC-CF-05',
    'confirmation': record_values, 'execution': execution,
    'language_order': ['zh', 'ru'], 'deliverables': ['docx'],
}
(DATA / 'content.json').write_text(json.dumps(content, ensure_ascii=False, indent=2) + '\n')
with ZipFile(out) as z:
    assert z.testzip() is None
assert len(task_table.rows) == 5 and len(task_table.columns) == 2
assert len(record_table.rows) == 8 and len(record_table.columns) == 2
for p, q, zh, ru in pairs:
    assert p.text == zh and q.text == ru and p._p.getnext() is q._p
for run in execution['runs']:
    assert hashlib.sha256(Path(run['log']).read_bytes()).hexdigest() == run['log_sha256']
assert not list(ROOT.glob('*.md'))
coverage = {
    'title': 'Задание 4: исправление и подтверждение',
    'source_steps_ru': [
        'Запустить тесты и подтвердить падение на 1001 символе',
        'Найти и исправить ошибку в Validate',
        'Запустить регрессионный тест и все тесты',
        'Заполнить запись подтверждения (ниже)',
    ],
    'all_four_steps_completed_in_demonstrator': True,
    'confirmation_form_visible_in_image': False,
    'confirmation_form_created_for_this_exercise': 'CONF-CF-01',
    'source_sha256': hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
    'original_system_test_status': 'NOT_RUN',
    'timer_20_00': 'Classroom time allowance only',
}
(DATA / 'source-coverage.json').write_text(json.dumps(coverage, ensure_ascii=False, indent=2) + '\n')
check = {
    'image_step_coverage': 'PASS', 'steps': 4, 'columns': 2,
    'before': {'tests': 18, 'passed': 17, 'failed': 1},
    'after_regression': {'tests': 1, 'passed': 1},
    'after_full_suite': {'tests': 18, 'passed': 18},
    'log_hash_integrity': 'PASS', 'bilingual_pairs': len(pairs), 'zh_above_ru': 'PASS',
    'docx_zip': 'PASS', 'word_only': True, 'render_review': 'PENDING',
    'demonstrator': 'PASS', 'original_system': 'NOT_RUN',
    'sha256': hashlib.sha256(out.read_bytes()).hexdigest(),
}
(DATA / 'verification.json').write_text(json.dumps(check, ensure_ascii=False, indent=2) + '\n')
fonts = ROOT.parent / '2026-10-19-bug-report-review-v1/work/fonts'
lo_fonts = '/Users/zhi/.cache/codex-runtimes/codex-primary-runtime/dependencies/native/libreoffice-headless/libreoffice/LibreOfficeDev.app/Contents/Resources/fonts/truetype'
(WORK / 'fonts.conf').write_text(f'<?xml version="1.0"?><!DOCTYPE fontconfig SYSTEM "fonts.dtd"><fontconfig><dir>{fonts}</dir><dir>{lo_fonts}</dir><cachedir>{WORK / "fontcache"}</cachedir></fontconfig>')
print(f'PASS: 4 task steps, real fail-then-pass evidence, confirmation record, {len(pairs)} Chinese-Russian pairs')
print(out)
