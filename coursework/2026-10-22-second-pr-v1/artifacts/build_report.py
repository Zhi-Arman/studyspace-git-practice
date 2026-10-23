from pathlib import Path
import json
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
E=json.loads((ROOT/'artifacts/execution.json').read_text())
TITLE='StudySpace_第二个PR_中俄双语_v1'
pages=[]
def page(zh,ru):
    blocks=[];pages.append(dict(zh=zh,ru=ru,blocks=blocks));return blocks
def pair(b,zh,ru):b.append(dict(type='pair',zh=zh,ru=ru))
def code(b,text):b.append(dict(type='code',text=text))
def table(b,heads,rows,widths):b.append(dict(type='table',heads=heads,rows=rows,widths=widths))

b=page('StudySpace 第二个 Pull Request','StudySpace Второй pull request')
pair(b,'SS-GIT-PR2-01 · v1.0 · 2026-10-22 · Asia/Shanghai · 文档 DRAFT',
       'SS-GIT-PR2-01 · v1.0 · 22.10.2026 · Asia/Shanghai · Документ DRAFT')
pair(b,'任务五已完成：沿用已解决冲突的 docs/conflict-b 分支，新开 PR #4，留下一个具体审查意见，完成一次修正，以 COMMENT 确认后合并至 main。',
       'Задание 5 выполнено: сохранённая ветка docs/conflict-b после разрешения конфликта опубликована в новом PR #4; оставлено одно конкретное замечание, выполнено одно исправление, затем подтверждение COMMENT и слияние в main.')
code(b,E['pr_url'])
table(b,['编号 / Код','图片要求 / Требование','实际结果 / Результат'],[
['PR2-01','发布解决冲突的分支\nОпубликовать ветку с разрешённым конфликтом','原 B 分支普通推送，保留解决提交\nОбычный push ветки B; коммит решения сохранён'],
['PR2-02','创建指向 main 的 PR\nОткрыть PR в main','新 PR #4，目标 main\nНовый PR #4, цель main'],
['PR2-03','描述冲突、方案与检查\nОписать конфликт, решение и проверки','PR 正文记录 A/B 同行冲突、-W -n 和实测\nВ PR описаны конфликт строки A/B, -W -n и проверки'],
['PR2-04','审查无标记及文本正确性\nПроверить маркеры и корректность текста','命令和双语说明核对；扫描无匹配\nКоманда и языковые пары проверены; маркеров нет'],
['PR2-05','按意见修正一次\nВнести одно исправление по замечанию','新增明确扫描命令和退出码解释\nДобавлены команда поиска и трактовка кодов'],
['PR2-06','确认后合并\nСлить после подтверждения','修正提交上的 COMMENT 后完成合并\nСлито после COMMENT на исправленном коммите'],
],[1.9,7.1,8.4])
pair(b,'前序任务四已经合并实际冲突解决 PR #3。本题不改写旧历史：继续使用原 B 分支，在解决提交之后补充说明并完成新的审查修正流程；本轮没有重新制造冲突。',
       'В задании 4 фактическое решение конфликта уже слито через PR #3. История не переписана: в прежней ветке B после коммита решения добавлена запись и проведён новый цикл рецензии; новый конфликт не создавался.')

b=page('2 一条意见与一次修正','2 Одно замечание и одно исправление')
pair(b,'审查者确认命令仍同时保留 A 的 -W 与 B 的 -n，当前受跟踪 Markdown 没有冲突标记；指出文档只列出 diff 与索引检查，缺少扫描已提交文本的可复现命令及退出码说明。',
       'Рецензент подтвердил сохранение -W из A и -n из B, отсутствие маркеров в Markdown под контролем Git. Замечание: в инструкции есть проверки diff и индекса, но нет воспроизводимого поиска по уже закоммиченному тексту и объяснения кодов возврата.')
pair(b,'作者角色在 docs/index.md 增加一处成对修正：明确扫描所有受跟踪 Markdown，解释 1 为无匹配、0 为发现标记、大于 1 为命令错误。只提交一次修正，未更改已解决的命令行。',
       'Роль автора внесла одну парную правку в docs/index.md: поиск во всех Markdown под контролем Git; код 1 — совпадений нет, 0 — маркеры найдены, больше 1 — ошибка команды. Создан один исправляющий коммит, итоговая команда не изменена.')
code(b,"git grep -n -E '^(<{7}|={7}|>{7})( |$)' -- '*.md'")
pair(b,'该命令只检查常见的行首冲突标记，不替代人工语义审查。扫描返回 1 是预期通过结果，不能将命令错误当成通过。',
       'Команда ищет обычные маркеры в начале строки и не заменяет смысловую рецензию. Код 1 здесь является ожидаемым успешным результатом; ошибку команды нельзя считать успехом.')
table(b,['检查 / Проверка','真实结果 / Фактический результат'],[
['git diff --check','初稿及修正后均通过\nПройдено до и после исправления'],
['git ls-files -u','索引无未解决项\nНеразрешённых записей нет'],
['受跟踪 Markdown 标记扫描\nПоиск маркеров в Markdown','无匹配，预期退出码 1\nСовпадений нет, ожидаемый код 1'],
['依赖检查 / Зависимости','现有隔离环境 uv pip check 通过\nuv pip check в изолированной среде пройден'],
['严格文档构建 / Строгая сборка','初稿与修正后各运行一次；均退出 0、警告 0\nДва запуска: до и после правки; оба код 0, предупреждений 0'],
],[6.6,10.8])
code(b,E['build_command'])
pair(b,'修正复核后，在修正后的确切提交上留下独立的一条确认 COMMENT，并将原审查线程标为已解决，随后合并 PR。线上最终文件树与实际构建过的修正树一致。',
       'После повторной проверки на точном исправленном коммите оставлен отдельный подтверждающий COMMENT, исходная ветка обсуждения закрыта как решённая, затем PR слит. Итоговое дерево GitHub совпадает с реально собранным исправленным деревом.')

b=page('3 提交记录与交付证据','3 Коммиты и свидетельства')
pair(b,f'执行时间 {E["executed_at"]}；最终核验 {len(E["checks"])}/{len(E["checks"])} 项通过。A/B 均在 main，与 GitHub 一致，工作区干净；所有练习分支保留。',
       f'Время {E["executed_at"]}; пройдены {len(E["checks"])}/{len(E["checks"])} итоговых проверок. A/B находятся на main и совпадают с GitHub, рабочие деревья чистые; все учебные ветки сохранены.')
table(b,['记录 / Запись','提交 / Коммит'],[
['原冲突解决 / Прежнее решение',E['resolution']],
['PR #4 初稿 / Начальная версия',E['initial']],
['一次修正 / Одно исправление',E['corrected']],
['最终合并 / Итоговое слияние',E['merge']],
],[5.6,11.8])
pair(b,'实际 GitHub 账号为 Zhi-Arman。B（俞智文）为作者角色，曹浩然为审查者角色，仅作用户已授权的自动化课堂演示；未声称两位成员本人操作。审查和确认都是 COMMENT，不是独立 APPROVE。',
       'Фактический аккаунт GitHub — Zhi-Arman. B (俞智文) — роль автора, 曹浩然 — роль рецензента в разрешённой автоматической учебной демонстрации; личное выполнение участниками не заявляется. Рецензия и подтверждение — COMMENT, не независимый APPROVE.')
pair(b,'已核验 Git 历史、发布内容、标记和文档构建。真人独立批准仍为 NOT_VERIFIED，应用业务测试仍为 NOT_RUN；没有声称课程或业务验收。',
       'Проверены история Git, опубликованный текст, маркеры и сборка документации. Независимое одобрение студентом — NOT_VERIFIED, бизнес-тесты приложения — NOT_RUN; приёмка курса или приложения не заявляется.')
pair(b,'本轮材料位于下列目录；README 与 Word 内容对应。artifacts 保留初稿、审查、修正 diff、两次构建日志、确认及最终历史。',
       'Материалы находятся в каталоге ниже; содержание README и Word соответствует. В artifacts сохранены начальная версия, рецензия, diff правки, два журнала сборки, подтверждение и итоговая история.')
code(b,'/Users/zhi/Documents/软件生命周期/studyspace/coursework/\n2026-10-22-second-pr-v1/\n\nREADME.md\n'+TITLE+'.docx\nartifacts/execution.json\nartifacts/review-posted.json\nartifacts/correction.diff\nartifacts/corrected-sphinx-build.log\nartifacts/confirmation-posted.json\nartifacts/history-final.txt')
doc=Document();s=doc.sections[0]
s.page_width=Cm(21);s.page_height=Cm(29.7);s.top_margin=Cm(1.6);s.bottom_margin=Cm(1.6);s.left_margin=Cm(1.8);s.right_margin=Cm(1.8)
for name in ('Normal','Title','Subtitle','Heading 1'):
    sty=doc.styles[name];sty.font.name='Arial';sty.font.color.rgb=RGBColor(0,0,0)
    sty._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Hiragino Sans GB')
n=doc.styles['Normal'];n.font.size=Pt(10.5);n.paragraph_format.space_after=Pt(4);n.paragraph_format.line_spacing=1.03
doc.styles['Title'].font.size=Pt(20);doc.styles['Subtitle'].font.size=Pt(13);doc.styles['Heading 1'].font.size=Pt(16)
cs=doc.styles.add_style('Code',1);cs.font.size=Pt(8.5);cs.font.name='Courier New';cs.paragraph_format.space_after=Pt(6);cs.paragraph_format.line_spacing=1
doc.core_properties.title='StudySpace 第二个PR Второй pull request';doc.core_properties.author='StudySpace'
def paragraph(text,style=None,ru=False):
    p=doc.add_paragraph(style=style)
    if ru:p.paragraph_format.space_after=Pt(6)
    for i,line in enumerate(text.split('\n')):
        if i:p.add_run().add_break()
        p.add_run(line)
    return p
md=[]
for pi,pag in enumerate(pages):
    h=paragraph(pag['zh'],'Title' if pi==0 else 'Heading 1')
    if pi:h.paragraph_format.page_break_before=True
    paragraph(pag['ru'],'Subtitle');md.extend([('#' if pi==0 else '##')+' '+pag['zh'],'',pag['ru'],''])
    for block in pag['blocks']:
        if block['type']=='pair':
            paragraph(block['zh']);paragraph(block['ru'],ru=True);md.extend([block['zh'],'',block['ru'],''])
        elif block['type']=='code':
            paragraph(block['text'],'Code');md.extend(['```text',block['text'],'```',''])
        else:
            t=doc.add_table(rows=1,cols=len(block['heads']));t.style='Table Grid';t.autofit=False
            for j,width in enumerate(block['widths']):t.columns[j].width=Cm(width)
            for i,values in enumerate([block['heads']]+block['rows']):
                row=t.rows[0] if i==0 else t.add_row();row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
                if i==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
                for j,value in enumerate(values):
                    cell=row.cells[j];cell.width=Cm(block['widths'][j]);cell.text=value
                    for p in cell.paragraphs:
                        p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1
                        for r in p.runs:r.font.size=Pt(9.5);r.bold=(i==0)
                    if i==0:
                        sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EEF4');cell._tc.get_or_add_tcPr().append(sh)
            doc.add_paragraph().paragraph_format.space_after=Pt(0)
            md.append('| '+' | '.join(x.replace('\n','<br>') for x in block['heads'])+' |');md.append('| '+' | '.join('---' for _ in block['heads'])+' |')
            md.extend('| '+' | '.join(x.replace('\n','<br>') for x in row)+' |' for row in block['rows']);md.append('')
for root in (doc._element,doc.styles.element):
    for border in root.xpath('.//w:pBdr'):border.getparent().remove(border)
for p in list(doc.paragraphs)+[p for t in doc.tables for row in t.rows for c in row.cells for p in c.paragraphs]:
    for r in p.runs:
        rf=r._element.get_or_add_rPr().get_or_add_rFonts()
        for attr in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):rf.attrib.pop(qn('w:'+attr),None)
        for key in ('ascii','hAnsi'):rf.set(qn('w:'+key),'Courier New' if p.style.name=='Code' else 'Arial')
        rf.set(qn('w:eastAsia'),'Hiragino Sans GB')
doc.save(ROOT/(TITLE+'.docx'))


md.extend(['## 文件入口 / Входные файлы','',f'[可编辑 Word / Редактируемый Word]({TITLE}.docx)',
           '[已合并 PR #4 / Слитый PR #4]('+E['pr_url']+')',
           '[审查意见 / Замечание]('+E['review_comment_url']+')',
           '[确认 / Подтверждение]('+E['confirmation_url']+')',
           '[完整检查 / Все проверки](artifacts/execution.json)',
           '[修正差异 / Diff правки](artifacts/correction.diff)',
           '[构建日志 / Журнал сборки](artifacts/corrected-sphinx-build.log)',
           '[最终历史 / История](artifacts/history-final.txt)',''])
(ROOT/'README.md').write_text('\n'.join(md))
(ROOT/'artifacts/report-content.json').write_text(json.dumps(pages,ensure_ascii=False,indent=2)+'\n')
previous=ROOT.parent/'2026-10-22-resolve-conflict-v1/artifacts/fonts.conf'
(ROOT/'artifacts/fonts.conf').write_text(previous.read_text().replace(str(previous.parent/'font-cache'),str(ROOT/'artifacts/font-cache')))
print(ROOT/(TITLE+'.docx'))
