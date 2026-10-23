from pathlib import Path
import json
import hashlib
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[1]
V = json.loads((ROOT / 'artifacts/pre-add-verification.json').read_text())
TITLE = '任务二项目材料组成与存储策略_中俄双语_v1'
pages = []
def page(zh, ru):
    p={'zh':zh,'ru':ru,'blocks':[]};pages.append(p);return p['blocks']
def pair(b,zh,ru):b.append({'type':'pair','zh':zh,'ru':ru})
def code(b,text):b.append({'type':'code','text':text})
def table(b,headers,rows,widths):b.append({'type':'table','headers':headers,'rows':rows,'widths':widths})

b=page('项目材料组成与存储策略','Состав проекта и политика хранения')
pair(b,'任务 2  版本 1  日期 2026-10-21  状态 DRAFT  范围 LOCAL_ONLY',
       'Задача 2  Версия 1  Дата 21.10.2026  Статус DRAFT  Область LOCAL_ONLY')
pair(b,'已使用现有 StudySpace 自有材料准备独立源包 project-docs。37 份材料与 6 个新建管理文件构成 43 个文件；相对路径、脱敏和存储规则已检查，文档严格构建通过。',
       'Подготовлен отдельный project-docs из собственных материалов StudySpace. В нём 37 исходных и 6 новых служебных файлов, всего 43. Проверены относительные пути, обезличивание и правила хранения; строгая сборка документации пройдена.')
table(b,['编号\nКод','图片要求\nТребование на фото','执行结果\nРезультат'],[
 ['MAT-01','取主题 3.1–3.4 的源材料\nВозьмите исходники материалов тем 3.1–3.4','已取自有源包；3.1–3.3 映射待核对\nСобственный комплект; номера 3.1–3.3 требуют сверки'],
 ['MAT-02','没有自有包时使用起始文件夹\nИспользуйте стартовую папку, если собственного пакета нет','采用自有包，无需起始文件夹\nЕсть собственный комплект'],
 ['MAT-03','保留页面与图片之间的相对路径\nСохраните относительные пути между страницами и изображениями','PASS\n52 处本地相对链接有效\n52 относительные ссылки действуют'],
 ['MAT-04','不要移动其他项目已有的 .git\nНе перемещайте уже существующую .git из другого проекта','PASS\n包内没有 .git；原仓库未移动\n.git не перенесён, исходный репозиторий сохранён'],
 ['MAT-05','首次 add 前去除秘密和个人数据\nУберите секреты и персональные данные до первого add','PASS 按所列规则检查\nПроверка по указанным правилам PASS\n本包未执行 git add'],
 ['MAT-06','创建 README 和 .gitignore 并确定存储策略\nСоздайте README и .gitignore; определите политику хранения','PASS\n三份文件已建立，19 项规则通过\nФайлы созданы, 19 проверок правил пройдены'],
],[0.65,3.25,3.0])
pair(b,'3.4.3、3.4.4 的对应关系可从已有记录确认；当前材料不能证明已覆盖教师主题 3.1–3.4 的全部内容。具体课程编号按 SOURCE_REGISTER.md 登记。',
       'Соответствие 3.4.3 и 3.4.4 подтверждается прежними записями. Текущий комплект не доказывает полноту всех материалов преподавателя тем 3.1–3.4. Соответствие номерам записано в SOURCE_REGISTER.md.')

b=page('2 文件结构与存储规则','2 Структура и правила хранения')
code(b,'''project-docs/
  README.md
  .gitignore
  STORAGE_POLICY.md
  SOURCE_REGISTER.md
  sources-manifest.json
  requirements-docs.in
  requirements-docs.txt
  build-site.sh
  docs/
    conf.py
    index.md
    user-guide.md
    admin-guide.md
    document-register.md
    release-notes.md
    images/development-chain.svg
    _static/coursework.css
    materials/''')
pair(b,'复制时完整保留 docs 层级。页面仍使用 images/development-chain.svg 和 materials 下的下载路径，未替换为绝对路径。两份 Word 下载材料也在副本内保留。',
       'Структура docs сохранена. Страницы используют images/development-chain.svg и загрузки внутри materials, без замены на абсолютные пути. Два Word-файла для загрузки также сохранены в копии.')
table(b,['类型\nКатегория','存储策略\nПравило'],[
 ['可编辑源与依赖锁\nИсходники и lock-файлы','跟踪 .md、.svg、.css、.py、构建脚本与带哈希依赖锁\nХранить .md, .svg, .css, .py, скрипт сборки и зависимости с хешами'],
 ['既有 Word 与历史证据\nWord и исторические записи','跟踪两份已处理 Word 与两份白名单历史日志；每个 Word 小于 10 MiB\nХранить два обработанных Word и два разрешённых исторических журнала; Word до 10 MiB'],
 ['秘密与运行数据\nСекреты и рабочие данные','忽略 .env、密钥、个人数据、数据库和备份；不复制源 .git\nИсключать .env, ключи, личные данные, БД и копии; не переносить .git'],
 ['依赖与生成结果\nЗависимости и результаты','忽略 .venv、node_modules、site、_build、根 artifacts、缓存与临时文件\nИсключать .venv, node_modules, site, _build, корневую artifacts, кеш и временные файлы'],
],[2.35,4.55])

b=page('3 检查结果与使用步骤','3 Проверки и порядок использования')
pair(b,'副本已替换 37 处本机用户路径和 4 处已知成员姓名，并清除 Word 作者元数据。历史来源清单已明确标注为脱敏前记录；当前副本的哈希以根 sources-manifest.json 为准。',
       'В копии заменены 37 локальных пользовательских путей и 4 известных имени участников, удалены метаданные авторов Word. Прежний манифест помечен как запись до обезличивания; хеши текущих копий находятся в sources-manifest.json корня.')
table(b,['检查项\nПроверка','实际结果\nРезультат'],[
 ['源材料与副本\nИсходники и копии','37 份来源逐文件校验，源文件哈希未变\n37 источников проверены, исходные хеши сохранены'],
 ['链接与忽略规则\nСсылки и исключения','52 处相对链接有效；19 项忽略规则通过；43 个文件均在候选清单\n52 ссылки действуют, 19 правил пройдены, 43 файла в списке кандидатов'],
 ['秘密与个人信息检查\nСекреты и персональные данные','已知姓名、用户路径、邮箱及所列令牌/私钥模式无剩余匹配\nНет оставшихся совпадений известных имён, путей, почты, указанных токенов и ключей'],
 ['文档严格构建\nСтрогая сборка','Sphinx 退出 0，警告 0；仅在临时目录生成网站\nSphinx: код 0, предупреждений 0; сайт создан во временной папке'],
 ['Git 边界\nОбласть Git','包内无 .git，本轮未执行 add、commit 或线上推送\nВ пакете нет .git; add, commit и онлайн-push не выполнялись'],
],[2.35,4.55])
pair(b,'从 project-docs 目录阅读 README.md、STORAGE_POLICY.md 与 SOURCE_REGISTER.md。A 继续维护 docs/user-guide.md，B 维护 docs/admin-guide.md；首次加入 Git 前再次查看候选文件。',
       'В project-docs прочитайте README.md, STORAGE_POLICY.md и SOURCE_REGISTER.md. A продолжает docs/user-guide.md, B — docs/admin-guide.md. Перед первым добавлением в Git снова проверьте список файлов.')
pair(b,'已有文档环境可用时运行下列构建命令。DOCS_PYTHON 只指定本机已有环境，不随包保存环境或账号信息。',
       'При наличии среды документации используйте команды ниже. DOCS_PYTHON задаёт существующую локальную среду, сама среда и данные аккаунтов в комплект не входят.')
code(b,'''# 在 project-docs 目录 / Из каталога project-docs
DOCS_PYTHON=/path/to/docs-venv/bin/python bash build-site.sh
# 加入选定练习仓库后检查 / После помещения в учебный репозиторий
git status --short
git ls-files --others --exclude-standard
git diff --cached --name-only''')
pair(b,'扫描是针对随包文本与 Word XML 的规则检查，不保证识别所有可能的秘密。历史测试结果不转为本轮结果；应用业务、真实账号权限与课堂验收均未验证。',
       'Это проверка текста и XML Word по заданным правилам, без гарантии обнаружения любых возможных секретов. Исторические тесты не становятся результатами текущего запуска; бизнес-функции, права реальных аккаунтов и приёмка не проверялись.')

(ROOT/'artifacts/report-content.json').write_text(json.dumps(pages,ensure_ascii=False,indent=2)+'\n')
md=[]
for i,p in enumerate(pages):
    md += [('# ' if not i else '## ')+p['zh'],'',p['ru'],'']
    for x in p['blocks']:
        if x['type']=='pair':md.extend([x['zh'],'',x['ru'],''])
        elif x['type']=='code':md.extend(['```text',x['text'],'```',''])
        else:
            clean=lambda s:s.replace('\n','<br>')
            md.extend(['| '+' | '.join(map(clean,x['headers']))+' |','| '+' | '.join('---' for _ in x['headers'])+' |'])
            md.extend('| '+' | '.join(map(clean,row))+' |' for row in x['rows']);md.append('')
(ROOT/(TITLE+'.md')).write_text('\n'.join(md))

doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.left_margin=sec.right_margin=Inches(.8);sec.top_margin=sec.bottom_margin=Inches(.63)
for name,size in [('Normal',10.5),('Title',19),('Subtitle',12.5),('Heading 1',16)]:
    s=doc.styles[name];s.font.name='Arial';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
    s._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Hiragino Sans GB')
    s.paragraph_format.space_after=Pt(4);s.paragraph_format.line_spacing=1.04
for border in list(doc.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)

def paragraph(text,style=None):return doc.add_paragraph(text,style)
for i,p in enumerate(pages):
    if i:doc.add_page_break()
    paragraph(p['zh'],'Title' if not i else 'Heading 1');paragraph(p['ru'],'Subtitle')
    for x in p['blocks']:
        if x['type']=='pair':
            paragraph(x['zh']).paragraph_format.keep_with_next=True
            para=paragraph(x['ru'])
            for r in para.runs:r.font.color.rgb=RGBColor.from_string('3D4650')
        elif x['type']=='code':
            para=doc.add_paragraph();para.paragraph_format.line_spacing=1
            for n,line in enumerate(x['text'].splitlines()):
                if n:para.add_run().add_break()
                r=para.add_run(line);r.font.name='Courier New';r.font.size=Pt(9)
        else:
            t=doc.add_table(rows=1,cols=len(x['headers']));t.autofit=False
            for j,w in enumerate(x['widths']):t.columns[j].width=Inches(w)
            borders=OxmlElement('w:tblBorders')
            for side in ['top','left','bottom','right','insideH','insideV']:
                el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            t._tbl.tblPr.append(borders)
            for n,rowvals in enumerate([x['headers'],*x['rows']]):
                row=t.rows[0] if n==0 else t.add_row()
                pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
                if n==0:pr.append(OxmlElement('w:tblHeader'))
                for j,val in enumerate(rowvals):
                    cell=row.cells[j];cell.width=Inches(x['widths'][j]);cell.text=val
                    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    props=cell._tc.get_or_add_tcPr();margins=OxmlElement('w:tcMar')
                    for side in ['top','bottom','left','right']:
                        el=OxmlElement('w:'+side);el.set(qn('w:w'),'75');el.set(qn('w:type'),'dxa');margins.append(el)
                    props.append(margins)
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'203E54' if n==0 else ('F2F6F8' if n%2 else 'FFFFFF'));props.append(shade)
                    for pp in cell.paragraphs:
                        pp.paragraph_format.space_after=Pt(0);pp.paragraph_format.line_spacing=1.02
                        for r in pp.runs:
                            r.font.size=Pt(9.3)
                            if n==0:r.bold=True;r.font.color.rgb=RGBColor(255,255,255)
            paragraph('').paragraph_format.space_after=Pt(0)
allps=list(doc.paragraphs)+[p for t in doc.tables for row in t.rows for c in row.cells for p in c.paragraphs]
for p in allps:
    for r in p.runs:r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Hiragino Sans GB')
doc.core_properties.title='项目材料组成与存储策略';doc.core_properties.author='';doc.core_properties.last_modified_by=''
out=ROOT/(TITLE+'.docx');doc.save(out)

writeup=f'''# 任务二 项目材料组成

Задача 2 Состав проекта

2026-10-21 · v1 · LOCAL_ONLY

已整理本地自有材料包：43 个文件、52 处相对链接、19 项忽略规则检查通过；文档严格构建退出 0、警告 0。源材料未改，源项目 .git 未复制或移动，本包未执行 git add。

Подготовлен собственный локальный комплект: 43 файла, 52 относительные ссылки и 19 проверок исключений пройдены. Строгая сборка: код 0, предупреждений 0. Исходники сохранены, .git не копировался и не переносился; git add в пакете не выполнялся.

- [中俄 Word 报告 / Отчёт Word]({TITLE}.docx)
- [同内容 Markdown / Тот же текст Markdown]({TITLE}.md)
- [项目材料包 README / README комплекта](project-docs/README.md)
- [存储策略 / Политика хранения](project-docs/STORAGE_POLICY.md)
- [材料来源与编号对应 / Источники и темы](project-docs/SOURCE_REGISTER.md)
- [当前副本哈希 / Хеши копий](project-docs/sources-manifest.json)
- [本轮检查 / Текущая проверка](artifacts/pre-add-verification.json)
- [本轮构建结果 / Текущая сборка](artifacts/current-build.json)

3.4.3 和 3.4.4 对应关系沿用已有记录；3.1–3.3 的具体编号映射缺少教师课程索引，保持待核对。所列扫描属于规则检查，不能保证识别所有可能的秘密。既有历史测试与业务 NOT_RUN 状态没有改变。

Соответствие 3.4.3 и 3.4.4 сохраняется по имеющимся записям; соответствие 3.1–3.3 требует индекса преподавателя. Проверка основана на правилах и не гарантирует обнаружение всех возможных секретов. Исторические результаты и состояние бизнес-тестов NOT_RUN не изменялись.
'''
(ROOT/'README.md').write_text(writeup)
reopened=Document(out);joined='\n'.join([p.text for p in reopened.paragraphs]+[c.text for t in reopened.tables for row in t.rows for c in row.cells])
expected=[]
for p in pages:
    expected.extend([p['zh'],p['ru']])
    for b in p['blocks']:
        if b['type']=='pair':expected.extend([b['zh'],b['ru']])
        elif b['type']=='code':expected.append(b['text'])
        else:expected.extend(b['headers']+[c for row in b['rows'] for c in row])
assert all(s in joined for s in expected)
checks={'all_inventory_text_present':True,'inventory_items':len(expected),
        'editable_tables':[(len(t.rows),len(t.columns)) for t in reopened.tables],
        'six_requirement_ids':all(f'MAT-{n:02}' in joined for n in range(1,7)),
        'visual_review':'PENDING','sha256':{out.name:hashlib.sha256(out.read_bytes()).hexdigest()},
        'source_documents_require_visual_review':['project-docs/docs/materials/StudySpace_管理员指南_中俄双语_v1.docx','project-docs/docs/materials/submission/StudySpace_管理员手册_中俄双语_v0.2.docx']}
(ROOT/'artifacts/report-verification.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2)+'\n')
print(json.dumps({'report':str(out),'inventory_items':len(expected),'tables':checks['editable_tables']},ensure_ascii=False))
