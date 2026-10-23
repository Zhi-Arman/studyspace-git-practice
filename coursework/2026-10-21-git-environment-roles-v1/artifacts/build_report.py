from pathlib import Path
import hashlib
import json
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[1]
E = json.loads((ROOT / 'artifacts/execution.json').read_text())
TITLE = 'Git环境准备与双人角色分工_中俄双语_v1'
pages = []

def page(title_zh, title_ru):
    p = {'title_zh': title_zh, 'title_ru': title_ru, 'blocks': []}
    pages.append(p)
    return p['blocks']

def pair(blocks, zh, ru):
    blocks.append({'type': 'pair', 'zh': zh, 'ru': ru})

def code(blocks, text):
    blocks.append({'type': 'code', 'text': text})

b = page('Git环境准备与双人角色分工', 'Подготовка среды Git и распределение ролей')
pair(b, '版本 1  日期 2026-10-21  文档状态 DRAFT  执行状态 LOCAL_DEMO',
        'Версия 1  Дата 21.10.2026  Документ DRAFT  Выполнение LOCAL_DEMO')
pair(b, '独立学习目录和 A/B 两个本地仓库已准备，Git 命令与文件分工检查通过。选定线上平台为 GitHub；当前练习通过本地裸仓库交换提交。',
        'Подготовлены отдельная учебная папка и два локальных репозитория A/B. Проверены команды Git и распределение файлов. Выбрана платформа GitHub; сейчас обмен коммитами идёт через локальный bare-репозиторий.')
b.append({'type': 'table', 'headers': ['编号\nКод', '图片要求\nТребование на фото', '完成情况\nРезультат'], 'widths': [0.55, 3.40, 2.95], 'rows': [
    ['GIT-01', '打开独立学习文件夹\nОткройте отдельную учебную папку', 'PASS\n已新建本练习目录\nСоздана отдельная папка'],
    ['GIT-02', '在终端检查 Git\nПроверьте Git в терминале', 'PASS\nGit 2.54.0 可运行\nGit 2.54.0 работает'],
    ['GIT-03', '选择 GitHub、GitLab 或教学服务器\nВыберите GitHub, GitLab или учебный сервер', '已选 GitHub\nВыбран GitHub\n线上仓库 NOT_CREATED'],
    ['GIT-04', 'A 创建初始仓库，B 随后克隆\nA создаёт исходный репозиторий; B клонирует его позднее', 'LOCAL_DEMO PASS\nA 建库并首次提交后 B 克隆\nB клонировал после первого коммита A'],
    ['GIT-05', 'A 编辑 user-guide.md，B 编辑 admin-guide.md\nA работает с user-guide.md; B с admin-guide.md', 'PASS\n两个角色分别提交指定文件\nКаждая роль изменила свой файл'],
    ['GIT-06', '每个人亲自检查命令\nПроверяйте команды от своего имени, а не только смотрите на партнёра', '本地角色检查 PASS\nПроверка локальных ролей PASS\n真人独立操作 NOT_VERIFIED'],
]})
pair(b, '图中建议用时 10 分钟。本地演示身份只用于验证流程；双方本人操作与课堂确认由参与者完成。',
        'На фото отведено 10 минут. Учебные подписи используются для проверки процесса; личное выполнение и подтверждение в аудитории выполняют участники.')

b = page('2 目录与角色', '2 Папки и роли')
pair(b, '本练习位于 StudySpace 的 coursework 下，使用独立 Git 元数据。既有项目文件未改动，原项目仓库未创建提交。',
        'Упражнение находится в coursework проекта StudySpace и имеет отдельные данные Git. В исходном репозитории проекта коммиты не создавались и файлы не изменялись.')
code(b, '''2026-10-21-git-environment-roles-v1/
  participant-a/       # A 的工作区 / каталог A
    user-guide.md
    admin-guide.md
  participant-b/       # B 的克隆 / клон B
    user-guide.md
    admin-guide.md
  teaching-origin.git/ # 本地裸仓库 / локальный bare
  artifacts/           # 命令记录 / журналы
  README.md''')
pair(b, 'A 负责创建初始仓库和 user-guide.md；B 克隆后负责 admin-guide.md。两份文件都存在于两个工作区中，分工通过各自的后续提交体现。',
        'A создаёт исходный репозиторий и ведёт user-guide.md. B клонирует его и ведёт admin-guide.md. Оба файла есть в обоих рабочих каталогах; распределение подтверждено отдельными последующими коммитами.')
pair(b, '本地身份分别为 Participant A (local exercise) 和 Participant B (local exercise)，邮箱使用 example.invalid。配置仅写入练习仓库，真实参与者应设置自己的本地提交身份。',
        'Локальные подписи: Participant A (local exercise) и Participant B (local exercise), адреса в example.invalid. Настройки записаны только в учебные репозитории. Участники задают собственные локальные подписи.')
pair(b, 'A 负责用户指南中的编辑、保存和预览步骤；B 负责仓库准备、同步与常见问题。每个人先确认当前目录与身份，再检查自己负责的文件差异。',
        'A описывает редактирование, сохранение и предпросмотр. B описывает подготовку репозитория, синхронизацию и частые затруднения. Каждый сначала проверяет каталог и подпись, затем diff своего файла.')
pair(b, '本轮属于 Git 文档协作练习。StudySpace 原有业务状态与团队长期角色不因本练习改变。',
        'Это упражнение по совместной работе с документацией в Git. Оно не изменяет состояние бизнес-функций StudySpace и постоянные роли команды.')

b = page('3 已执行的协作流程', '3 Выполненный порядок работы')
pair(b, '执行顺序为 A 首次提交、B 克隆、A 修改用户指南、B 同步并修改管理员指南、A 再同步。以下列出实际操作摘要；完整参数与退出码保存在 artifacts/commands.log。',
        'Порядок: первый коммит A, клонирование B, правка руководства пользователя A, получение B и правка руководства администратора, затем получение A. Ниже основные команды; полные параметры и коды выхода в artifacts/commands.log.')
pair(b, '准备阶段先创建 teaching-origin.git 裸仓库，再初始化 participant-a 的 main 分支。A 首次提交包含 README.md、两份指南和 .gitignore，随后推送到本地 origin。',
        'Сначала создан bare-репозиторий teaching-origin.git и ветка main в participant-a. Первый коммит A содержит README.md, два руководства и .gitignore, затем отправлен в локальный origin.')
code(b, '''git init --bare --initial-branch=main teaching-origin.git
git init --initial-branch=main participant-a
# 在 A 工作区内 / В рабочем каталоге A
git add README.md user-guide.md admin-guide.md .gitignore
git commit -m "docs: initialize bilingual guide practice"
git remote add origin ../teaching-origin.git
git push -u origin main
# 回到练习根目录后 / Из корня упражнения
git clone --no-hardlinks teaching-origin.git participant-b''')
pair(b, '首次初始化和克隆已完成，请在新空白目录中复现这些步骤。为保持本地路径明确，实际记录中的 origin 与 clone 参数使用绝对路径。',
        'Инициализация и клонирование уже выполнены; воспроизводите их в новой пустой папке. В журнале для origin и clone использованы абсолютные пути.')
code(b, '''# A 保存 user-guide.md 后 / A сохраняет user-guide.md
git add user-guide.md
git commit -m "docs: add user guide editing steps"
git push origin main
# B 先同步再保存 admin-guide.md / B получает A и правит файл
git pull --ff-only origin main
git add admin-guide.md
git commit -m "docs: add admin guide repository checks"
git push origin main
# A 最后同步 / A получает итоговый коммит B
git pull --ff-only origin main''')

b = page('4 检查结果与个人复核', '4 Результаты и личная проверка')
pair(b, f'执行日期 2026-10-21，终端返回 {E["git_version"]}。本轮 {len(E["checks"])} 项本地仓库检查全部通过；两个工作区各有 3 个提交，最终提交一致。',
        f'Дата выполнения 21.10.2026; терминал: {E["git_version"]}. Пройдены все {len(E["checks"])} проверок локальных репозиториев. В каждом рабочем каталоге 3 коммита; итоговый коммит совпадает.')
b.append({'type': 'table', 'headers': ['检查项\nПроверка', '实际结果\nФактический результат'], 'widths': [3.1, 3.8], 'rows': [
    ['仓库独立性\nИзоляция репозиториев', 'A/B 各有独立 Git 根目录\nУ A/B отдельные корни Git'],
    ['克隆顺序\nПорядок клонирования', 'B 克隆时 HEAD 等于 A 的初始提交\nHEAD B при клонировании равен первому коммиту A'],
    ['指定文件分工\nРаспределение файлов', 'A 后续提交仅 user-guide.md；B 仅 admin-guide.md\nA изменил только user-guide.md, B только admin-guide.md'],
    ['同步与完整性\nСинхронизация и целостность', 'A、B、origin 的 main 一致；fsck 通过\nmain A, B и origin совпадает; fsck пройден'],
    ['工作区\nРабочее дерево', 'A/B 的 status --porcelain 均无输出\nstatus --porcelain у A/B пуст'],
    ['线上与本人操作\nОнлайн и личное выполнение', 'GitHub NOT_CREATED；真人 NOT_VERIFIED\nGitHub NOT_CREATED; личное выполнение NOT_VERIFIED'],
]})
pair(b, '双方分别在自己的工作区执行以下命令。检查仓库根目录与身份，确认 main 分支、两份指南和提交记录；完成同步后工作区应干净。',
        'Каждый выполняет команды в своём рабочем каталоге. Проверьте корень, подпись, ветку main, два руководства и историю. После синхронизации рабочее дерево должно быть чистым.')
code(b, '''pwd
git --version
git rev-parse --show-toplevel
git config --local --get user.name
git config --local --get user.email
git branch --show-current
git status --short
git log --oneline -3
git rev-parse HEAD''')
pair(b, f'本轮最终提交前 12 位为 {E["b_edit_commit"][:12]}。A、B 分别自行复核并记录姓名、日期与输出后，才能确认“本人独立执行”。',
        f'Первые 12 символов итогового коммита: {E["b_edit_commit"][:12]}. Для подтверждения личного выполнения A и B самостоятельно проверяют команды и записывают имя, дату и вывод.')
pair(b, '选定 GitHub 尚需线上建库与访问权限设置；本地 origin 只接受本机练习提交。VS Code 预览、真人双账号权限和应用业务测试本轮未验证。',
        'Для выбранного GitHub ещё требуются создание онлайн-репозитория и права доступа. Локальный origin обслуживает только упражнение на этом компьютере. Предпросмотр VS Code, права двух реальных аккаунтов и бизнес-тесты приложения сейчас не проверялись.')
pages[1]['blocks'].append(pages[3]['blocks'].pop())

(ROOT / 'artifacts/content.json').write_text(json.dumps(pages, ensure_ascii=False, indent=2) + '\n')

md = []
for i, page_data in enumerate(pages):
    md += [('#' if i == 0 else '##') + ' ' + page_data['title_zh'], '', page_data['title_ru'], '']
    for block in page_data['blocks']:
        if block['type'] == 'pair':
            md += [block['zh'], '', block['ru'], '']
        elif block['type'] == 'code':
            md += ['```bash', block['text'], '```', '']
        else:
            cell = lambda s: s.replace('\n', '<br>')
            md += ['| ' + ' | '.join(map(cell, block['headers'])) + ' |', '| ' + ' | '.join(['---'] * len(block['headers'])) + ' |']
            md += ['| ' + ' | '.join(map(cell, row)) + ' |' for row in block['rows']]
            md += ['']
(ROOT / (TITLE + '.md')).write_text('\n'.join(md), encoding='utf-8')

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(0.63)
sec.left_margin = sec.right_margin = Inches(0.8)
styles = doc.styles
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1']:
    s = styles[name]
    s.font.name = 'Arial'
    s._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Hiragino Sans GB')
    s.font.color.rgb = RGBColor.from_string('000000')
styles['Normal'].font.size = Pt(10.5)
styles['Normal'].paragraph_format.space_after = Pt(4)
styles['Normal'].paragraph_format.line_spacing = 1.08
styles['Title'].font.size = Pt(20)
styles['Title'].paragraph_format.space_after = Pt(5)
styles['Subtitle'].font.size = Pt(13)
styles['Subtitle'].paragraph_format.space_after = Pt(12)
styles['Heading 1'].font.size = Pt(17)
styles['Heading 1'].paragraph_format.space_after = Pt(5)
for el in list(doc.styles.element.iter(qn('w:pBdr'))):
    el.getparent().remove(el)

def add_text(text, style=None, size=None, color=None):
    p = doc.add_paragraph(text, style=style)
    if size:
        for r in p.runs:
            r.font.size = Pt(size)
    if color:
        for r in p.runs:
            r.font.color.rgb = RGBColor.from_string(color)
    return p

def add_table(block):
    table = doc.add_table(rows=1, cols=len(block['headers']))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for idx, width in enumerate(block['widths']):
        table.columns[idx].width = Inches(width)
    pr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        el = OxmlElement('w:' + side)
        el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '4'); el.set(qn('w:color'), 'D9D9D9')
        borders.append(el)
    pr.append(borders)
    for idx, row_values in enumerate([block['headers'], *block['rows']]):
        row = table.rows[0] if idx == 0 else table.add_row()
        if idx == 0:
            repeat = OxmlElement('w:tblHeader'); row._tr.get_or_add_trPr().append(repeat)
        trpr = row._tr.get_or_add_trPr()
        prevent_split = OxmlElement('w:cantSplit'); trpr.append(prevent_split)
        for j, value in enumerate(row_values):
            cell = row.cells[j]; cell.width = Inches(block['widths'][j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for side in ['top', 'left', 'bottom', 'right']:
                edge = OxmlElement('w:' + side); edge.set(qn('w:w'), '85'); edge.set(qn('w:type'), 'dxa'); margins.append(edge)
            tcpr.append(margins)
            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), '203E54' if idx == 0 else ('F2F6F8' if idx % 2 else 'FFFFFF')); tcpr.append(shade)
            cell.text = value
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.04
                for r in p.runs:
                    r.font.size = Pt(9.5)
                    if idx == 0:
                        r.font.bold = True; r.font.color.rgb = RGBColor(255, 255, 255)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)

for i, p in enumerate(pages):
    if i:
        doc.add_page_break()
    add_text(p['title_zh'], 'Title' if not i else 'Heading 1')
    add_text(p['title_ru'], 'Subtitle')
    for block in p['blocks']:
        if block['type'] == 'pair':
            para = add_text(block['zh']); para.paragraph_format.keep_with_next = True
            add_text(block['ru'], color='3D4650')
        elif block['type'] == 'code':
            para = doc.add_paragraph()
            para.paragraph_format.space_after = Pt(7)
            para.paragraph_format.line_spacing = 1.0
            for idx, line in enumerate(block['text'].splitlines()):
                if idx: para.add_run().add_break()
                r = para.add_run(line); r.font.name = 'Courier New'; r.font.size = Pt(9.5)
        else:
            add_table(block)
doc.core_properties.title = 'Git环境准备与双人角色分工'
doc.core_properties.subject = 'Подготовка среды Git и распределение ролей'
doc.core_properties.author = ''
doc.core_properties.keywords = 'Git, Chinese, Russian, LOCAL_DEMO'
all_paragraphs = list(doc.paragraphs)
all_paragraphs += [p for table in doc.tables for row in table.rows for cell in row.cells for p in cell.paragraphs]
for p in all_paragraphs:
    for el in list(p._p.iter(qn('w:pBdr'))):
        el.getparent().remove(el)
    for r in p.runs:
        r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Hiragino Sans GB')
out = ROOT / (TITLE + '.docx')
doc.save(out)

readme = f'''# Git 环境与角色练习

Подготовка среды Git и роли

2026-10-21 · v1 · LOCAL_DEMO · 平台选择 GitHub / Выбран GitHub

本地演示已完成，{len(E['checks'])} 项仓库检查通过；线上仓库 NOT_CREATED，两位同学本人执行 NOT_VERIFIED。

Локальное упражнение завершено, пройдены {len(E['checks'])} проверок. Онлайн-репозиторий NOT_CREATED, личное выполнение двух студентов NOT_VERIFIED.

## 交付文件 / Материалы

- [中俄双语 Word 报告 / Отчёт Word]({TITLE}.docx)
- [同内容 Markdown / Тот же текст Markdown]({TITLE}.md)
- [A 用户指南 / Руководство пользователя A](participant-a/user-guide.md)
- [B 管理员指南 / Руководство администратора B](participant-b/admin-guide.md)
- [本地仓库检查结果 / Результаты проверок](artifacts/execution.json)
- [完整命令记录 / Журнал команд](artifacts/commands.log)

## 打开与复核 / Открытие и проверка

学习根目录 / Корень упражнения:

```text
{ROOT}
```

A 在 VS Code 打开 participant-a，B 打开 participant-b。双方分别在终端切换到自己的文件夹，执行报告第 4 节命令，记录本人姓名、日期与终端结果。当前配置身份为本地演示角色，请用 git config --local 设置本人身份。配置变化只影响后续提交。

A открывает participant-a в VS Code, B — participant-b. Каждый выполняет команды раздела 4 в своём каталоге и записывает имя, дату и результат. Текущие подписи учебные; задайте собственную подпись через git config --local. Это влияет только на будущие коммиты.

当前 origin 指向绝对路径 {ROOT / 'teaching-origin.git'}，只在这台电脑有效。不要把本地裸仓库当成已创建的 GitHub 仓库。若整体移动目录，应在 A/B 各仓库执行 git remote set-url origin 后接新裸仓库路径。

Текущий origin указывает на абсолютный путь локального teaching-origin.git и работает на этом компьютере. Онлайн-репозиторий GitHub ещё не создан. При переносе всей папки обновите origin у A/B командой git remote set-url origin с новым путём к bare-репозиторию.

## 来源与执行 / Источники и выполнение

任务来源为附图“Готовность среды и роли”，含 6 个要点，建议 10 分钟。原图保存在 artifacts/assignment.jpg；Git 版本以当前终端实测为准。指南是本练习的双语文件，未替换既有课程材料。

Источник — фото «Готовность среды и роли»: 6 пунктов, 10 минут. Фото сохранено в artifacts/assignment.jpg; версия Git проверена в терминале. Руководства созданы для этого упражнения и не заменяют прежние материалы.

本机自动执行了 A/B 两种角色的命令，不能替代两名同学本人操作和教师确认。GitHub 建库、真实账号授权、VS Code 预览和应用业务测试本轮未执行。

Команды обеих ролей выполнены автоматически на одном компьютере. Это не заменяет личную работу двух студентов и подтверждение преподавателя. Создание GitHub, права реальных аккаунтов, предпросмотр VS Code и бизнес-тесты не выполнялись.
'''
(ROOT / 'README.md').write_text(readme, encoding='utf-8')

reopened = Document(out)
texts = [p.text for p in reopened.paragraphs]
texts.extend(cell.text for table in reopened.tables for row in table.rows for cell in row.cells)
joined = '\n'.join(texts)
expected = []
for p in pages:
    expected += [p['title_zh'], p['title_ru']]
    for block in p['blocks']:
        if block['type'] == 'pair': expected += [block['zh'], block['ru']]
        elif block['type'] == 'code': expected += [block['text']]
        else: expected += block['headers'] + [cell for row in block['rows'] for cell in row]
checks = [
    {'name': 'all inventory text present in DOCX', 'pass': all(s in joined for s in expected)},
    {'name': 'two editable tables with expected dimensions', 'pass': [(len(t.rows), len(t.columns)) for t in reopened.tables] == [(7, 3), (7, 2)]},
    {'name': 'six original requirement IDs', 'pass': all(f'GIT-{i:02}' in joined for i in range(1, 7))},
    {'name': 'each report text pair contains Chinese and Cyrillic', 'pass': all(re.search('[\u4e00-\u9fff]', b['zh']) and re.search('[А-Яа-яЁё]', b['ru']) for p in pages for b in p['blocks'] if b['type'] == 'pair')},
    {'name': 'local and human execution boundaries present', 'pass': all(x in joined for x in ['LOCAL_DEMO', 'NOT_CREATED', 'NOT_VERIFIED'])},
    {'name': 'Markdown contains all six requirement IDs', 'pass': all(f'GIT-{i:02}' in '\n'.join(md) for i in range(1, 7))},
    {'name': 'no retained PDF companion', 'pass': not list(ROOT.glob('*.pdf'))},
]
report = {'checks': checks, 'all_pass': all(c['pass'] for c in checks), 'inventory_text_count': len(expected), 'visual_review': 'PENDING', 'sha256': {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in [out, ROOT / (TITLE + '.md')]}}
(ROOT / 'artifacts/document-verification.json').write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n')
assert report['all_pass'], report
print(json.dumps({'docx': str(out), 'structural_checks': len(checks), 'inventory_text_count': len(expected)}, ensure_ascii=False))
