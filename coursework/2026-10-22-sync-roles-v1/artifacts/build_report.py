from pathlib import Path
import json
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
E=json.loads((ROOT/'artifacts/online-verification.json').read_text())
L=json.loads((ROOT/'artifacts/verification.json').read_text())
TITLE='StudySpace_同步与角色分工_中俄双语_v1'
pages=[]
def page(zh,ru):
    b=[];pages.append(dict(zh=zh,ru=ru,blocks=b));return b
def pair(b,zh,ru): b.append(dict(type='pair',zh=zh,ru=ru))
def table(b,heads,rows,widths): b.append(dict(type='table',heads=heads,rows=rows,widths=widths))
def code(b,text): b.append(dict(type='code',text=text))

b=page('StudySpace 同步与角色分工','StudySpace Синхронизация и роли')
pair(b,'SS-GIT-SYNC-01 · v1.0 · 2026-10-22 · Asia/Shanghai · 文档 DRAFT',
       'SS-GIT-SYNC-01 · v1.0 · 22.10.2026 · Asia/Shanghai · Документ DRAFT')
pair(b,'已在 GitHub 创建私有共享仓库，延续实践 4.1.1 的本地练习历史。A、B、审查工作区已同步到线上 main，工作目录与暂存区均干净；角色、语言和协作规则已记录。',
       'Создан частный общий репозиторий GitHub с локальной историей практики 4.1.1. Каталоги A, B и рецензента синхронизированы с main GitHub; рабочие деревья и индекс чистые. Роли, язык и правила записаны.')
code(b,E['repository']['html_url'])
table(b,['编号\nКод','图片要求\nТребование на фото','实际结果\nФактический результат'],[
['SYNC-01','检查 Git 和所选平台访问\nПроверьте Git и доступ к выбранной платформе',L['git_version']+'；GitHub 所有者读写通过\nGit работает; доступ владельца на чтение и запись проверен'],
['SYNC-02','打开实践 4.1.1 的共享仓库\nОткройте общий репозиторий из практики 4.1.1','保留原 3 次练习提交，新增规则提交并上传私有仓库\nСохранены 3 учебных коммита; добавлены правила и отправлены в частный репозиторий'],
['SYNC-03','检查工作目录干净且 main 最新\nПроверьте чистоту рабочей папки и актуальность main','A、B、审查目录 CLEAN；main 与线上一致\nA, B и рецензент CLEAN; main совпадает с GitHub'],
['SYNC-04','指定作者 A、作者 B 和审查者\nНазначьте автора A, автора B и рецензента','A 支浩然；B 俞智文；审查 曹浩然\nA 支浩然; B 俞智文; рецензент 曹浩然'],
['SYNC-05','约定 commit 与 PR 语言\nСогласуйте язык коммитов и сообщений PR','英语；正文文档保持中俄对应\nАнглийский; документы на китайском и русском'],
['SYNC-06','记录禁止强推与删分支的规则\nЗафиксируйте правила без force push и без удаления веток','规则已记录；自动删分支已关闭\nПравила записаны; автоудаление веток выключено'],
],[1.9,7.3,8.2])
pair(b,'平台证据限于仓库所有者账号。三名同学的个人账号权限和本人执行尚未核验，成员共同确认状态为 NOT_VERIFIED。',
       'Платформа проверена от имени владельца. Права личных аккаунтов трёх студентов, личное выполнение и совместное подтверждение имеют статус NOT_VERIFIED.')

b=page('2 角色与沟通语言','2 Роли и язык общения')
pair(b,'分工沿用 StudySpace 已有成员顺序，并用于本次 Git 练习。长期课程岗位继续保留；不将自动化操作记作同学本人执行。',
       'Роли опираются на существующий порядок участников StudySpace и относятся к этому упражнению Git. Постоянные учебные обязанности сохраняются; автоматизация не считается личным выполнением студентов.')
table(b,['角色\nРоль','成员\nУчастник','职责与交付\nОбязанности и результат'],[
['作者 A\nАвтор A','支浩然','维护 user-guide.md，核对差异并提交自己的 PR\nВедёт user-guide.md, проверяет diff и создаёт свой PR'],
['作者 B\nАвтор B','俞智文','维护 admin-guide.md，核对差异并提交自己的 PR\nВедёт admin-guide.md, проверяет diff и создаёт свой PR'],
['审查者\nРецензент','曹浩然','检查双语对应、修改范围、冲突处理与验证记录\nПроверяет языковые пары, состав изменений, конфликты и проверки'],
['协调人\nКоординатор','吕嘉伟','维护分工、执行顺序及问题记录\nВедёт роли, порядок выполнения и вопросы'],
],[3.2,2.8,11.4])
pair(b,'提交信息及 PR 标题、正文统一使用英语，与既有练习历史保持一致。文档正文仍采用中文在上、俄语在下；技术标识保留原名。',
       'Сообщения коммитов, заголовки и описания PR ведутся на английском, как в истории упражнения. В документах китайский текст расположен над русским; технические идентификаторы сохраняются.')
pair(b,'提交类型采用 feat、fix、docs、test、refactor、chore。本轮实际新增提交如下，由本机已有 zhi 身份执行，未替换为同学姓名。',
       'Типы коммитов: feat, fix, docs, test, refactor, chore. Ниже фактический новый коммит с существующей локальной подписью zhi, без подмены именем студента.')
code(b,'docs: define synchronization roles and collaboration rules\n'+E['online_main'])
pair(b,'后续 PR 使用以下英语结构。它是填写格式，本轮没有创建 PR，也没有实际评审或批准。',
       'Для следующего PR используется английская структура ниже. Это шаблон; в этом запуске PR не создавался, рецензия и утверждение не выполнялись.')
code(b,'Title: docs: clarify user guide editing steps\nChange: Describe the actual changes.\nValidation: List only completed checks and results.\nReviewer: Cao Haoran')

b=page('3 协作规则与同步步骤','3 Правила и порядок синхронизации')
pair(b,'规则完整版本保存在本地及线上 COLLABORATION.md。本轮新增规则适用于后续协作，旧指南中直接推送 main 的示例属于上一轮本地演示。',
       'Полная версия правил находится в локальном и онлайн-файле COLLABORATION.md. Новые правила действуют для дальнейшей работы; примеры прямого push в main в старых руководствах относятся к прежней локальной демонстрации.')
for zh,ru in [
('1 禁止强推，包括 --force、--force-with-lease 和非快进 +refspec。推送被拒绝时停止并与伙伴检查差异。','1 Запрещён принудительный push, включая --force, --force-with-lease и +refspec без fast-forward. При отказе остановитесь и проверьте diff с партнёром.'),
('2 禁止删除本地与远程分支，PR 合并后也保留分支。已实测 GitHub delete_branch_on_merge 为 false。','2 Запрещено удалять локальные и удалённые ветки; после слияния PR они сохраняются. На GitHub проверено delete_branch_on_merge = false.'),
('3 工作前确认目录干净，再 fetch 和 pull --ff-only。已有修改先保存并协调，不用 reset --hard 或 clean 清场。','3 Перед fetch и pull --ff-only убедитесь в чистоте каталога. Сохраните и согласуйте существующие правки; не очищайте их через reset --hard или clean.'),
('4 使用 feature/<topic>、fix/<topic> 或 docs/<topic>，以小范围 PR 交给审查者。作者不得自行批准。','4 Используйте feature/<topic>, fix/<topic> или docs/<topic> и небольшие PR для рецензента. Автор не утверждает собственный PR.'),
('5 冲突处理先确认内容，解决后普通提交，不改写共享历史。审查记录只写实际运行的检查与结果。','5 Сначала согласуйте содержание при конфликте, затем создайте обычный коммит без переписывания общей истории. Записывайте только реально выполненные проверки.')]: pair(b,zh,ru)
pair(b,'本轮已克隆线上 A/B 工作区，并在审查目录 fetch 后以快进方式 pull。后续在各自目录复核：',
       'Созданы онлайн-клоны A/B; в каталоге рецензента выполнены fetch и pull с fast-forward. Для повторной проверки в своём каталоге:')
code(b,'git status --short\ngit fetch --no-prune origin\ngit switch main\ngit pull --ff-only origin main\ngit rev-list --left-right --count main...origin/main\ngit rev-parse HEAD\ngit ls-remote origin refs/heads/main')
pair(b,'预期：status 无输出，ahead/behind 为 0 0，HEAD 与 ls-remote 的 main 提交号一致。若目录不干净或出现分叉，应停止同步并协调。',
       'Ожидается: status пуст, ahead/behind равен 0 0, HEAD совпадает с main из ls-remote. При правках или расхождении истории остановите синхронизацию и согласуйте действия.')

b=page('4 实测证据与交付位置','4 Проверки и расположение материалов')
pair(b,f'执行时间 {E["executed_at"]}。本地共享仓库检查 {len(L["checks"])}/{len(L["checks"])} 通过，线上仓库及同步检查 {len(E["checks"])}/{len(E["checks"])} 通过。',
       f'Время выполнения {E["executed_at"]}. Пройдены {len(L["checks"])}/{len(L["checks"])} локальных проверок и {len(E["checks"])}/{len(E["checks"])} проверок онлайн-репозитория и синхронизации.')
table(b,['对象\nОбъект','实测结果\nРезультат'],[
['GitHub 仓库\nРепозиторий GitHub','private = true；default_branch = main；所有者 pull/push = true\nprivate = true; default_branch = main; pull/push владельца = true'],
['author-a / author-b / reviewer','CLEAN；ahead 0 / behind 0；Git 对象完整性检查通过\nCLEAN; ahead 0 / behind 0; проверка целостности Git пройдена'],
['线上 main\nmain GitHub',E['online_main']],
['上传范围\nОтправленный состав','仅 5 个文档类文件和 4 次练习提交\nТолько 5 файлов документации и 4 учебных коммита'],
['StudySpace 根仓库\nКорневой репозиторий StudySpace','仍为 e2c21ee；已跟踪内容未改；原有未跟踪文件保留\nОстался e2c21ee; отслеживаемые файлы не изменены; прежние неотслеживаемые файлы сохранены'],
],[5.5,11.9])
pair(b,'线上读写与同步已验证。尚未添加同学账号或验证各成员权限；本人操作、团队共同确认和教师验收为 NOT_VERIFIED。服务器分支保护未配置；禁止强推和删除属于已记录的协作约定。业务测试 NOT_RUN。',
       'Чтение, запись и синхронизация GitHub проверены. Личные аккаунты студентов ещё не добавлены; их права, личное выполнение, совместное подтверждение и приёмка преподавателем — NOT_VERIFIED. Серверная защита веток не настроена; запрет force push и удаления записан как правило команды. Бизнес-тесты NOT_RUN.')
pair(b,'交付根目录与主要文件如下。author-a、author-b、reviewer 是本轮线上工作区；shared-work 保存上传准备。原来的本地 A/B 与裸仓库仍保留。',
       'Ниже корень материалов и основные файлы. author-a, author-b и reviewer — текущие онлайн-каталоги; shared-work содержит подготовку отправки. Прежние локальные A/B и bare-репозиторий сохранены.')
code(b,str(ROOT)+'\n\nREADME.md\n'+TITLE+'.docx\nCOLLABORATION.md\nauthor-a/  author-b/  reviewer/  shared-work/\nartifacts/task.jpg\nartifacts/verification.json\nartifacts/online-verification.json')
pair(b,'来源为本次任务原图、StudySpace 的 AGENTS.md 与 docs/data/project-decisions.json 中 DEC-017、上次练习仓库，以及本次终端和 GitHub 实测记录。',
       'Источники: фото задания, AGENTS.md StudySpace, DEC-017 в docs/data/project-decisions.json, предыдущий учебный репозиторий и текущие результаты терминала и GitHub.')

doc=Document()
sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.6);sec.bottom_margin=Cm(1.6);sec.left_margin=Cm(1.8);sec.right_margin=Cm(1.8)
for name in ('Normal','Title','Subtitle','Heading 1','Heading 2'):
    s=doc.styles[name];s.font.name='Arial';s.font.color.rgb=RGBColor(0,0,0)
    s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Hiragino Sans GB')
normal=doc.styles['Normal'];normal.font.size=Pt(10.5)
normal.paragraph_format.space_after=Pt(4)
normal.paragraph_format.line_spacing=1.08
doc.styles['Title'].font.size=Pt(20)
doc.styles['Subtitle'].font.size=Pt(13)
doc.styles['Heading 1'].font.size=Pt(16)
doc.core_properties.title='StudySpace 同步与角色分工 Синхронизация и роли'
doc.core_properties.author='StudySpace'

def para(text,ru=False,style=None,size=None):
    p=doc.add_paragraph(style=style)
    if ru:p.paragraph_format.space_after=Pt(7)
    for i,line in enumerate(text.split('\n')):
        if i:p.add_run().add_break()
        r=p.add_run(line)
        if size:r.font.size=Pt(size)
        if style=='Code':r.font.name='Consolas'
    return p

cs=doc.styles.add_style('Code',1);cs.font.name='Consolas';cs.font.size=Pt(8.5)
cs.paragraph_format.space_after=Pt(6)
cs.paragraph_format.line_spacing=1.0

md=[]
for pi,pag in enumerate(pages):
    if pi:doc.add_page_break()
    para(pag['zh'],style='Title' if pi==0 else 'Heading 1')
    para(pag['ru'],style='Subtitle')
    md.extend([('#' if pi==0 else '##')+' '+pag['zh'],'',pag['ru'],''])
    for block in pag['blocks']:
        if block['type']=='pair':
            para(block['zh']);para(block['ru'],ru=True)
            md.extend([block['zh'],'',block['ru'],''])
        elif block['type']=='code':
            para(block['text'],style='Code')
            md.extend(['```text',block['text'],'```',''])
        else:
            t=doc.add_table(rows=1, cols=len(block['heads']));t.style='Table Grid';t.autofit=False
            for idx,w in enumerate(block['widths']): t.columns[idx].width=Cm(w)
            for rownum,values in enumerate([block['heads']]+block['rows']):
                row=t.rows[0] if rownum==0 else t.add_row()
                trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
                if rownum==0:trpr.append(OxmlElement('w:tblHeader'))
                for idx,value in enumerate(values):
                    cell=row.cells[idx];cell.width=Cm(block['widths'][idx])
                    cell.text=value
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after=Pt(3)
                        p.paragraph_format.space_before=Pt(3)
                        p.paragraph_format.line_spacing=1.05
                        for r in p.runs:r.font.size=Pt(9.5);r.bold=(rownum==0)
                    if rownum==0:
                        sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EEF4');cell._tc.get_or_add_tcPr().append(sh)
            doc.add_paragraph().paragraph_format.space_after=Pt(0)
            md.append('| '+' | '.join(x.replace('\n','<br>') for x in block['heads'])+' |')
            md.append('| '+' | '.join('---' for _ in block['heads'])+' |')
            md.extend('| '+' | '.join(x.replace('\n','<br>') for x in row)+' |' for row in block['rows'])
            md.append('')
for parent in (doc._element,doc.styles.element):
    for border in parent.xpath('.//w:pBdr'):
        border.getparent().remove(border)
for p in list(doc.paragraphs)+[p for t in doc.tables for row in t.rows for c in row.cells for p in c.paragraphs]:
    for r in p.runs:
        rf=r._element.get_or_add_rPr().get_or_add_rFonts()
        for attr in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):
            rf.attrib.pop(qn('w:'+attr),None)
        rf.set(qn('w:ascii'),'Courier New' if p.style.name=='Code' else 'Arial')
        rf.set(qn('w:hAnsi'),'Courier New' if p.style.name=='Code' else 'Arial')
        rf.set(qn('w:eastAsia'),'Hiragino Sans GB')
doc.save(ROOT/(TITLE+'.docx'))
md.extend(['## 文件入口 Входные файлы','',f'[可编辑 Word Редактируемый Word]({TITLE}.docx)',
           '[协作规则 Правила](COLLABORATION.md)',
           '[线上共享仓库 Общий репозиторий]('+E['repository']['html_url']+')',
           '[本地检查 Локальные проверки](artifacts/verification.json)',
           '[线上检查 Онлайн проверки](artifacts/online-verification.json)',''])
(ROOT/'README.md').write_text('\n'.join(md))
(ROOT/'artifacts/report-content.json').write_text(json.dumps(pages,ensure_ascii=False,indent=2)+'\n')
print(ROOT/(TITLE+'.docx'))
