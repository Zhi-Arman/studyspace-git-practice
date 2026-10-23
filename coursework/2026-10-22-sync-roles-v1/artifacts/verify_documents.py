from pathlib import Path
from datetime import datetime
from zoneinfo import ZoneInfo
from zipfile import ZipFile
import hashlib
import json
import re
import subprocess
from docx import Document
from lxml import etree

ROOT=Path(__file__).resolve().parents[1]
path=ROOT/'StudySpace_同步与角色分工_中俄双语_v1.docx'
content=json.loads((ROOT/'artifacts/report-content.json').read_text())
doc=Document(path)
alltext=''.join(p.text for p in doc.paragraphs)+''.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
compact=lambda s:re.sub(r'\s+','',s)
flat=compact(alltext)
expected=[]
for pag in content:
    expected.extend([pag['zh'],pag['ru']])
    for b in pag['blocks']:
        if b['type']=='pair':expected.extend([b['zh'],b['ru']])
        elif b['type']=='code':expected.append(b['text'])
        else:
            expected.extend(b['heads'])
            expected.extend(x for row in b['rows'] for x in row)
assert all(compact(x) in flat for x in expected)
with ZipFile(path) as z:
    assert z.testzip() is None
    for n in z.namelist():
        if n.endswith('.xml'):etree.fromstring(z.read(n))
    assert b'<w:pBdr' not in z.read('word/document.xml')
    assert b'<w:pBdr' not in z.read('word/styles.xml')
md=(ROOT/'README.md').read_text()
for target in re.findall(r'\]\(([^)]+)\)',md):
    if not target.startswith(('https://','http://')):assert (ROOT/target).exists(),target
assert (ROOT/'COLLABORATION.md').read_bytes()==(ROOT/'author-a/COLLABORATION.md').read_bytes()
assert len(re.findall(r'^\| SYNC-0[1-6] \|',md,re.M))==6
assert subprocess.check_output(['git','rev-list','--count','HEAD'],cwd=ROOT/'author-a',text=True).strip()=='4'
for role in ('author-a','author-b','reviewer','shared-work'):
    assert subprocess.check_output(['git','status','--porcelain=v1','--untracked-files=all'],cwd=ROOT/role,text=True)==''
online=json.loads((ROOT/'artifacts/online-verification.json').read_text())
assert all(c['passed'] for c in online['checks'])
images=sorted((ROOT/'artifacts/render-delivery').glob('page-*.png'))
assert len(images)==4,len(images)
result=dict(checked_at=datetime.now(ZoneInfo('Asia/Shanghai')).isoformat(),
            docx_sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
            docx_structure='PASS',bilingual_content_parity='PASS',source_text_items=len(expected),
            markdown_links='PASS',image_task_items=6,canonical_rules_match_uploaded_file=True,
            online_clone_worktrees='CLEAN',pages=len(images),
            rendering='Bundled LibreOffice with task-local FONTCONFIG_FILE',
            all_pages_visually_reviewed=False)
(ROOT/'artifacts/document-verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(result,ensure_ascii=False,indent=2))
