from pathlib import Path
from datetime import datetime
from zoneinfo import ZoneInfo
from zipfile import ZipFile
import hashlib
import json
import re
import subprocess
from docx import Document
from lxml import etree

ROOT=Path(__file__).resolve().parents[1]
path=ROOT/'StudySpace_第一个PR_中俄双语_v1.docx'
doc=Document(path)
content=json.loads((ROOT/'artifacts/report-content.json').read_text())
compact=lambda s:re.sub(r'\s+','',s)
flat=compact(''.join(p.text for p in doc.paragraphs)+''.join(c.text for t in doc.tables for row in t.rows for c in row.cells))
expected=[]
for pag in content:
    expected.extend([pag['zh'],pag['ru']])
    for b in pag['blocks']:
        if b['type']=='pair':expected.extend([b['zh'],b['ru']])
        elif b['type']=='code':expected.append(b['text'])
        else:
            expected.extend(b['heads']);expected.extend(x for row in b['rows'] for x in row)
assert all(compact(x) in flat for x in expected)
with ZipFile(path) as z:
    assert z.testzip() is None
    for name in z.namelist():
        if name.endswith('.xml'):etree.fromstring(z.read(name))
    assert b'<w:pBdr' not in z.read('word/document.xml')
    assert b'<w:pBdr' not in z.read('word/styles.xml')
md=(ROOT/'README.md').read_text()
for target in re.findall(r'\]\(([^)]+)\)',md):
    if not target.startswith(('http://','https://')):assert (ROOT/target).exists(),target
assert len(re.findall(r'^\| PR-0[1-6] \|',md,re.M))==6
E=json.loads((ROOT/'artifacts/execution.json').read_text())
assert E['merged'] and E['branch_preserved'] and all(c['passed'] for c in E['checks'])
parents=subprocess.check_output(['git','rev-list','--parents','-n','1',E['merge_commit']],cwd=ROOT/'author-a',text=True).split()
assert parents==[E['merge_commit'],E['base'],E['corrected_head']]
assert E['actual_actor']=='Zhi-Arman' and E['review_event']=='COMMENT' and E['independent_approval']=='NOT_VERIFIED'
assert '单账号自动化角色演示' in md and 'COMMENT' in md and 'NOT_VERIFIED' in md
images=sorted((ROOT/'artifacts/render-delivery').glob('page-*.png'))
assert len(images)==3,len(images)
result=dict(checked_at=datetime.now(ZoneInfo('Asia/Shanghai')).isoformat(),docx_sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
            zip_xml_integrity='PASS',bilingual_content_parity='PASS',source_text_items=len(expected),markdown_links='PASS',
            six_task_items_present=True,merged_commit_parents_verified=True,role_demo_disclosed=True,
            pages=len(images),all_pages_visually_reviewed=False)
(ROOT/'artifacts/document-verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(result,ensure_ascii=False,indent=2))
