from pathlib import Path
import json
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
E=json.loads((ROOT/'artifacts/execution.json').read_text())
TITLE='StudySpace_第一个PR_中俄双语_v1'
pages=[]
def page(zh,ru):
    blocks=[];pages.append(dict(zh=zh,ru=ru,blocks=blocks));return blocks
def pair(b,zh,ru):b.append(dict(type='pair',zh=zh,ru=ru))
def code(b,text):b.append(dict(type='code',text=text))
def table(b,heads,rows,widths):b.append(dict(type='table',heads=heads,rows=rows,widths=widths))

b=page('StudySpace 第一个 Pull Request','StudySpace Первый pull request')
pair(b,'SS-GIT-PR-01 · v1.0 · 2026-10-22 · Asia/Shanghai · 文档 DRAFT',
       'SS-GIT-PR-01 · v1.0 · 22.10.2026 · Asia/Shanghai · Документ DRAFT')
pair(b,'PR #1 已完成创建、具体审查意见、修正、确认和合并；docs/preview 分支保留。按用户授权采用单账号自动化角色演示，未取得真人 B 的独立批准。',
       'PR #1 создан, получил конкретное замечание, исправлен, проверен и слит; ветка docs/preview сохранена. По разрешению пользователя выполнена автоматическая демонстрация ролей одним аккаунтом; независимого одобрения реального B нет.')
code(b,E['pr_url'])
table(b,['编号\nКод','图片要求\nТребование на фото','实际结果\nФактический результат'],[
['PR-01','A 从最新 main 创建 docs/preview\nA создаёт docs/preview от актуального main','PASS；起点 '+E['base'][:7]+'\nPASS; исходный коммит '+E['base'][:7]],
['PR-02','在 user-guide.md 做一处小澄清\nОдно небольшое уточнение в user-guide.md','将推送示例改为当前练习分支\nПример push исправлен на ветку упражнения'],
['PR-03','提交并发布分支\nСоздать коммит и опубликовать ветку','初始提交 '+E['initial_head'][:7]+' 已推送\nНачальный коммит '+E['initial_head'][:7]+' отправлен'],
['PR-04','PR 写明目的、修改、验证及限制\nPR с целью, изменениями, проверкой и ограничениями','已创建 PR #1；四部分完整\nPR #1 создан; все четыре части заполнены'],
['PR-05','B 留下一条具体审查意见\nB оставляет одно конкретное замечание','已提交 COMMENT 审查与一条行内意见\nДобавлены COMMENT и одно замечание в diff'],
['PR-06','A 修正，B 确认并合并\nA исправляет, B подтверждает, PR объединяют','修正 '+E['corrected_head'][:7]+'；B 角色 COMMENT 确认；已合并\nИсправление '+E['corrected_head'][:7]+'; подтверждение роли B через COMMENT; PR слит'],
],[1.8,7.3,8.3])
pair(b,'本轮课堂角色为 A 支浩然、B 俞智文。根据本图，B 在本轮承担审查；长期分工及上一轮记录保留。GitHub 实际操作账号始终为 Zhi-Arman。',
       'Учебные роли: A 支浩然, B 俞智文. По этому фото B выполняет рецензию; постоянные роли и прежние записи сохраняются. Все действия GitHub фактически выполнены аккаунтом Zhi-Arman.')

b=page('2 审查意见与修正记录','2 Замечание и исправление')
pair(b,'首个提交仅把 git push origin main 改为 git push -u origin docs/preview。B 角色在独立工作目录检查差异后，指出必须先检查当前分支。',
       'Первый коммит заменил только git push origin main на git push -u origin docs/preview. Роль B проверила diff в отдельном каталоге и указала необходимость проверки текущей ветки.')
pair(b,'具体意见：新的 push 指向 docs/preview，但 commit 仍发生在当前分支。应在 git add 前执行 git branch --show-current，中俄两段均说明预期输出 docs/preview；输出不符时停止，避免在 main 提交却推送另一分支。',
       'Замечание: новый push отправляет docs/preview, но commit создаётся в текущей ветке. До git add выполните git branch --show-current. На обоих языках укажите ожидаемый вывод docs/preview и остановку при несовпадении, чтобы не создать коммит в main и отправить другую ветку.')
pair(b,'A 角色增加中俄对应说明和分支检查命令，普通提交后再次推送。核心命令及条件如下；复核已运行其中的分支检查，未重复执行示例提交。',
       'Роль A добавила парные инструкции и команду проверки ветки, создала обычный коммит и отправила его. Ниже основные команды и условия. Команда проверки ветки выполнена; пример коммита повторно не запускался.')
code(b,'先检查分支名，必须为 docs/preview，否则停止。\nПроверьте ветку: ожидается docs/preview; иначе остановитесь.\n\ngit branch --show-current\ngit add user-guide.md\ngit commit -m "docs: add user guide editing steps"\ngit push -u origin docs/preview\ngit status --short')
pair(b,'B 角色拉取修正提交，核对中俄停止条件、命令顺序、唯一修改文件、干净工作区与线上提交号后，留下第二条 COMMENT 作为自动化确认；原行内线程已解决。',
       'Роль B получила исправление и проверила условия остановки на обоих языках, порядок команд, единственный файл, чистоту каталога и совпадение с GitHub. Второй COMMENT фиксирует автоматическое подтверждение; исходная ветка обсуждения разрешена.')
table(b,['记录\nЗапись','提交或状态\nКоммит или состояние'],[
['初始提交\nПервый коммит',E['initial_head']],
['修正提交\nИсправление',E['corrected_head']],
['合并提交\nКоммит слияния',E['merge_commit']],
['审查与确认\nРецензия и подтверждение','2 次 COMMENTED；1 条行内问题已解决\n2 COMMENTED; 1 замечание в diff разрешено'],
],[5.5,11.9])

b=page('3 合并验证与交付文件','3 Проверка слияния и материалы')
pair(b,f'执行时间 {E["executed_at"]}。最终 {len(E["checks"])}/{len(E["checks"])} 项流程与仓库检查通过。GitHub 返回 merged=true，PR 已关闭并标记为已合并。',
       f'Время выполнения {E["executed_at"]}. Пройдены {len(E["checks"])}/{len(E["checks"])} итоговых проверок процесса и репозиториев. GitHub вернул merged=true; PR закрыт со статусом слияния.')
table(b,['检查\nПроверка','实际结果\nРезультат'],[
['合并前提交校验\nПроверка перед слиянием','PR 头等于已确认的修正提交；只修改 user-guide.md\nHEAD PR равен проверенному исправлению; изменён только user-guide.md'],
['历史与分支\nИстория и ветки','采用 merge；两个普通提交保留，远程和本地 docs/preview 均保留\nИспользован merge; два обычных коммита и docs/preview локально и на GitHub сохранены'],
['A 与 B 工作目录\nКаталоги A и B','均在 main，等于线上合并提交；工作区干净；git fsck 通过\nОба в main, совпадают с GitHub; каталоги чистые; git fsck пройден'],
['StudySpace 根仓库\nКорень StudySpace','原 main 与已跟踪文件未改动\nИсходные main и отслеживаемые файлы не изменены'],
],[5.1,12.3])
pair(b,'验证范围为 Git、PR 记录与文档内容。VS Code 视觉预览及业务测试为 NOT_RUN；真人 B 独立批准为 NOT_VERIFIED。单账号确认与合并方式由用户明确选择。',
       'Проверены Git, записи PR и содержание документа. Визуальный предпросмотр VS Code и бизнес-тесты — NOT_RUN; независимое одобрение реального B — NOT_VERIFIED. Пользователь явно выбрал подтверждение и слияние одним аккаунтом.')
pair(b,'报告和原始证据保存在以下目录。A/B 两个工作目录均为私有课堂仓库的独立克隆；原练习材料保留。',
       'Отчёт и исходные свидетельства находятся ниже. Каталоги A/B — отдельные клоны частного учебного репозитория; прежние материалы сохранены.')
code(b,str(ROOT)+'\n\nREADME.md\n'+TITLE+'.docx\nauthor-a/\nreviewer-b/\nartifacts/task.jpg\nartifacts/start.json\nartifacts/review-posted.json\nartifacts/correction.json\nartifacts/confirmation-posted.json\nartifacts/review-threads-final.json\nartifacts/merge-response.json\nartifacts/pr-final.json\nartifacts/execution.json')
pair(b,'来源为本次任务图片、上一轮共享仓库、本次实际 GitHub PR 与终端输出。审查、修正、确认和合并均保留时间、账号与提交证据。',
       'Источники: фото задания, общий репозиторий предыдущего этапа, фактический PR GitHub и вывод терминала. Для рецензии, исправления, подтверждения и слияния сохранены время, аккаунт и коммиты.')

doc=Document();s=doc.sections[0]
s.page_width=Cm(21);s.page_height=Cm(29.7)
s.top_margin=Cm(1.6);s.bottom_margin=Cm(1.6);s.left_margin=Cm(1.8);s.right_margin=Cm(1.8)
for name in ('Normal','Title','Subtitle','Heading 1'):
    sty=doc.styles[name];sty.font.name='Arial';sty.font.color.rgb=RGBColor(0,0,0)
    sty._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Hiragino Sans GB')
n=doc.styles['Normal'];n.font.size=Pt(10.5);n.paragraph_format.space_after=Pt(4);n.paragraph_format.line_spacing=1.05
doc.styles['Title'].font.size=Pt(20);doc.styles['Subtitle'].font.size=Pt(13);doc.styles['Heading 1'].font.size=Pt(16)
cs=doc.styles.add_style('Code',1);cs.font.size=Pt(8.5);cs.font.name='Courier New';cs.paragraph_format.space_after=Pt(6);cs.paragraph_format.line_spacing=1
doc.core_properties.title='StudySpace 第一个 Pull Request Первый pull request';doc.core_properties.author='StudySpace'
def paragraph(text,style=None,ru=False):
    p=doc.add_paragraph(style=style)
    if ru:p.paragraph_format.space_after=Pt(7)
    for i,line in enumerate(text.split('\n')):
        if i:p.add_run().add_break()
        p.add_run(line)
    return p
md=[]
for pi,pag in enumerate(pages):
    heading=paragraph(pag['zh'],'Title' if pi==0 else 'Heading 1')
    if pi:heading.paragraph_format.page_break_before=True
    paragraph(pag['ru'],'Subtitle')
    md.extend([('#' if pi==0 else '##')+' '+pag['zh'],'',pag['ru'],''])
    for block in pag['blocks']:
        if block['type']=='pair':
            paragraph(block['zh']);paragraph(block['ru'],ru=True)
            md.extend([block['zh'],'',block['ru'],''])
        elif block['type']=='code':
            paragraph(block['text'],'Code');md.extend(['```text',block['text'],'```',''])
        else:
            t=doc.add_table(rows=1,cols=len(block['heads']));t.style='Table Grid';t.autofit=False
            for j,width in enumerate(block['widths']):t.columns[j].width=Cm(width)
            for i,values in enumerate([block['heads']]+block['rows']):
                row=t.rows[0] if i==0 else t.add_row()
                row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
                if i==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
                for j,value in enumerate(values):
                    cell=row.cells[j];cell.width=Cm(block['widths'][j]);cell.text=value
                    for p in cell.paragraphs:
                        p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1
                        for r in p.runs:r.font.size=Pt(9.5);r.bold=(i==0)
                    if i==0:
                        sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EEF4');cell._tc.get_or_add_tcPr().append(sh)
            doc.add_paragraph().paragraph_format.space_after=Pt(0)
            md.append('| '+' | '.join(x.replace('\n','<br>') for x in block['heads'])+' |')
            md.append('| '+' | '.join('---' for _ in block['heads'])+' |')
            md.extend('| '+' | '.join(x.replace('\n','<br>') for x in row)+' |' for row in block['rows']);md.append('')
for root in (doc._element,doc.styles.element):
    for border in root.xpath('.//w:pBdr'):border.getparent().remove(border)
for p in list(doc.paragraphs)+[p for t in doc.tables for row in t.rows for c in row.cells for p in c.paragraphs]:
    for r in p.runs:
        rf=r._element.get_or_add_rPr().get_or_add_rFonts()
        for attr in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):rf.attrib.pop(qn('w:'+attr),None)
        rf.set(qn('w:ascii'),'Courier New' if p.style.name=='Code' else 'Arial');rf.set(qn('w:hAnsi'),'Courier New' if p.style.name=='Code' else 'Arial');rf.set(qn('w:eastAsia'),'Hiragino Sans GB')
doc.save(ROOT/(TITLE+'.docx'))
md.extend(['## 文件入口 Входные файлы','',f'[双语 Word Двуязычный Word]({TITLE}.docx)',
           '[PR #1]('+E['pr_url']+')',
           '[初始审查 Рецензия]('+E['initial_review_url']+')',
           '[修正确认 Подтверждение]('+E['confirmation_url']+')',
           '[完整检查记录 Проверки](artifacts/execution.json)',''])
(ROOT/'README.md').write_text('\n'.join(md))
(ROOT/'artifacts/report-content.json').write_text(json.dumps(pages,ensure_ascii=False,indent=2)+'\n')
fonts=(PROJECT_FONTS:=ROOT.parent/'2026-10-22-sync-roles-v1/artifacts/fonts.conf').read_text()
fonts=fonts.replace(str(PROJECT_FONTS.parent/'font-cache'),str(ROOT/'artifacts/font-cache'))
(ROOT/'artifacts/fonts.conf').write_text(fonts)
print(ROOT/(TITLE+'.docx'))
