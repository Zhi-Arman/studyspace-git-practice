from pathlib import Path
from datetime import datetime, timezone
import json
import re
import hashlib
from zipfile import ZipFile
from lxml import etree
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path('/Users/zhi/Documents/软件生命周期/studyspace/coursework/2026-10-15-scrum-kanban-v1')
OUT = ROOT / '任务二_Scrum冲刺规划与Kanban看板设计_中俄双语_v1.docx'
PAGES = []
CURRENT = []

def bi(ru, zh):
    return {'ru': ru, 'zh': zh}

def item(kind, value, **kwargs):
    CURRENT.append({'kind': kind, 'value': value, **kwargs})

def para(ru, zh, **kwargs):
    item('p', bi(ru, zh), **kwargs)

def head(ru, zh):
    item('h', bi(ru, zh))

def table(headers, rows, widths, size=10.5, centered=()):
    item('table', {'headers': headers, 'rows': rows, 'widths': widths, 'size': size, 'centered': centered})

def page(landscape=False):
    global CURRENT
    if CURRENT:
        PAGES.append({'landscape': landscape, 'items': CURRENT})
        CURRENT = []

# Page 1
item('title', bi('Задание 2 Планирование спринта Scrum и доска Kanban', '任务二 Scrum冲刺规划与Kanban看板设计'))
para('Практическая работа по жизненному циклу программного обеспечения. Версия 1.0. Дата 15.10.2026.',
     '软件生命周期课程实践作业。版本 1.0。日期 2026年10月15日。', small=True)
head('1 Кейс Scrum Заказ еды в столовой', '1 案例一 Scrum食堂订餐')
para('Продукт — мобильное приложение для предварительного заказа еды в университетской столовой. Пять разработчиков, спринт две недели, историческая скорость всей команды 20 очков за спринт. Эталон входа по логину и паролю — 3 очка.',
     '产品为大学食堂提前订餐手机应用。开发人员5名，冲刺两周，整个团队的历史速度为每次冲刺20个故事点；账号密码登录为3点基准。')
head('1 1 Распределение ролей', '1 1 角色分配')
para('Учебные обозначения A–E: каждый участвует в разработке; A также владелец продукта, B также Scrum-мастер.',
     '课堂代号A—E均参与开发，A兼任产品负责人，B兼任Scrum Master。', small=True)
table([bi('Код', '代号'), bi('Роль и ответственность', '角色与职责'), bi('Основной вклад', '主要开发工作')], [
    ['A', bi('Владелец продукта и разработчик: порядок бэклога, ценность, критерии приёмки.', '产品负责人兼开发者，负责排序、价值与验收条件。'), bi('US01 Вход', 'US01 登录')],
    ['B', bi('Scrum-мастер и разработчик: события Scrum и устранение препятствий.', 'Scrum Master兼开发者，负责活动与排障协调。'), bi('US02 Меню', 'US02 菜单')],
    ['C', bi('Разработчик. Корзина, расчёт суммы и взаимное ревью.', '开发者，负责购物车、金额计算及交叉评审。'), bi('US03 Корзина', 'US03 购物车')],
    ['D', bi('Разработчик. Оформление заказа и целостность данных.', '开发者，负责预约下单与数据一致性。'), bi('US04 Заказ', 'US04 预约下单')],
    ['E', bi('Разработчик. Список заказов столовой и интеграционные проверки.', '开发者，负责食堂订单列表及集成检查。'), bi('US06 Заказы столовой', 'US06 食堂订单')],
], [0.65, 4.55, 1.9], centered=(0,))
head('1 2 Цель спринта одним предложением', '1 2 一句话冲刺目标')
para('За две недели обеспечить студенту путь от входа и выбора блюд до заказа на выбранное время с оплатой при получении, а сотруднику столовой — просмотр поступившего заказа.',
     '在两周内让学生完成从登录、选餐到预约下单并到店付款的流程，同时让食堂工作人员能够查看收到的订单。')
page()

# Page 2
head('1 3 Исходный бэклог продукта', '1 3 原始产品待办列表')
para('Сохраняется порядок 12 историй из раздаточного материала. По заданию оцениваются только US01–US06; остальные истории остаются в продукте для будущих спринтов.',
     '保留讲义12条故事的原始顺序。本次按题目要求只估算US01—US06，其余留待后续冲刺。', small=True)
STORIES = [
    ('US01','Как студент, я хочу войти по логину и паролю, чтобы видеть свой профиль.','作为学生，我希望使用账号和密码登录，以查看个人资料。','Высокий','高'),
    ('US02','Как студент, я хочу видеть меню на сегодня, чтобы выбрать блюда.','作为学生，我希望查看今日菜单，以选择菜品。','Высокий','高'),
    ('US03','Как студент, я хочу добавить блюда в корзину, чтобы собрать заказ.','作为学生，我希望把菜品加入购物车，以组成订单。','Высокий','高'),
    ('US04','Как студент, я хочу оформить заказ на время, чтобы забрать его без очереди.','作为学生，我希望预约取餐时间并下单，以便免排队取餐。','Высокий','高'),
    ('US05','Как студент, я хочу оплатить заказ картой в приложении.','作为学生，我希望在应用内使用银行卡支付订单。','Высокий','高'),
    ('US06','Как сотрудник столовой, я хочу видеть список входящих заказов.','作为食堂工作人员，我希望查看收到的订单列表。','Высокий','高'),
    ('US07','Как студент, я хочу видеть историю моих заказов.','作为学生，我希望查看自己的历史订单。','Средний','中'),
    ('US08','Как студент, я хочу повторить прошлый заказ одной кнопкой.','作为学生，我希望一键重复之前的订单。','Средний','中'),
    ('US09','Как студент, я хочу получать уведомление, когда заказ готов.','作为学生，我希望在订单准备好时收到通知。','Средний','中'),
    ('US10','Как студент, я хочу оставить отзыв о блюде.','作为学生，我希望对菜品发表评论。','Низкий','低'),
    ('US11','Как студент, я хочу фильтровать меню по аллергенам.','作为学生，我希望按过敏原筛选菜单。','Низкий','低'),
    ('US12','Как администратор, я хочу видеть статистику продаж за неделю.','作为管理员，我希望查看一周的销售统计。','Низкий','低'),
]
table([bi('ID','编号'),bi('Пользовательская история','用户故事'),bi('Приоритет','优先级')],
      [[s[0],bi(s[1],s[2]),bi(s[3],s[4])] for s in STORIES], [0.7,5.35,1.05],size=10.5,centered=(0,2))
page()

# Page 3
head('1 4 Оценка методом покера планирования', '1 4 规划扑克故事点估算')
para('Шкала — 1, 2, 3, 5, 8, 13. Разработчики независимо выбирают карты, одновременно раскрывают их, обсуждают крайние оценки и повторяют голосование до согласования. Оценка относительная: сложность, объём и неопределённость; очки не переводятся в часы.',
     '采用1、2、3、5、8、13。开发者独立选牌并同时亮牌，讨论最大与最小估值的理由，再次投票直至达成共识。故事点比较复杂度、工作量与不确定性，不直接换算工时。')
para('Таблица — модель учебного раунда, а не протокол реально проведённой встречи. Карты указаны в порядке A, B, C, D, E; итог — согласованная оценка, не среднее арифметическое.',
     '下表为课堂模拟轮次，不代表已经举行过团队会议。投牌顺序为A、B、C、D、E；最终点数代表讨论后的共识，不是简单平均值。', small=True)
ESTIMATES = [
    ('US01','3 3 3 3 3',3,'Эталон: форма входа, проверка данных, сессия и ошибки входа.','基准故事：登录表单、身份校验、会话与错误提示。'),
    ('US02','2 3 3 5 3',3,'Готовый источник меню; список, цены, доступность и пустое состояние. Сложность близка к входу.','已有菜单数据源；展示菜品、价格、可售状态与空状态，复杂度接近登录。'),
    ('US03','3 5 5 5 8',5,'Количество, удаление, сумма и согласованность состояния корзины дают больше комбинаций, чем вход.','数量修改、删除、总价及购物车状态一致性使组合情况多于登录。'),
    ('US04','5 5 8 5 5',5,'Выбор допустимого времени, создание заказа и защита от повторной отправки; без онлайн-оплаты.','校验取餐时间、创建订单并避免重复提交；不包含在线支付。'),
    ('US05','8 13 8 8 13',8,'Платёжный провайдер, подтверждение и повторные уведомления, сбой/отмена и согласованность статуса.','支付服务对接、回调与重复回调、失败或取消以及订单支付状态一致性。'),
    ('US06','3 3 5 3 3',3,'Роль сотрудника и простой список с номером, блюдами и временем. Без кухни, уведомлений и сложных фильтров.','员工权限与订单编号、菜品、时间的简单列表；不扩展厨房流程、通知或复杂筛选。'),
]
table([bi('ID','故事'),bi('Карты A–E','A—E投牌'),bi('Итог','点数'),bi('Обоснование','估算依据')],
      [[a,b,str(c),bi(d,e)] for a,b,c,d,e in ESTIMATES], [0.65,1.15,0.55,4.75],size=10.5,centered=(0,1,2))
head('1 5 Допущения для оценки', '1 5 估算所用假设')
para('Учебный вариант предполагает готовые учётные записи, источник меню и заданные столовой интервалы выдачи. Регистрация, редактирование меню и расчёт мощности кухни не входят в эти истории.',
     '规划假设已有账号、菜单数据来源及食堂给定的取餐时段；账号注册、菜单维护和厨房产能算法不包含在上述故事内。')
para('Для первой версии допускается оплата при получении. US05 сохраняет высокий приоритет. Если столовая требует предоплату, состав спринта и оценки нужно пересмотреть до принятия плана.',
     '首版假设允许到店付款，US05仍保持高优先级。若食堂必须预付款，应在正式采用计划前重新调整冲刺范围和估算。')
page()

# Page 4
head('1 6 Бэклог спринта и критерии приёмки', '1 6 冲刺待办列表与验收条件')
para('Выбираются US01, US02, US03, US04 и US06: 3 + 3 + 5 + 5 + 3 = 19 очков ≤ 20. Остаётся 1 очко незагруженной ёмкости. Скорость 20 относится ко всей команде, её не умножают на пять разработчиков.',
     '选择US01、US02、US03、US04、US06：3＋3＋5＋5＋3＝19点，不超过20点，保留1点未分配容量。20点是全队速度，不乘以5名开发人员。')
SPRINT = [
    ('US01\n3 / A','Форма входа; проверка пользователя; сессия; тесты.','登录表单、身份校验、会话及测试。','Верные данные открывают профиль; неверные дают понятную ошибку; без входа профиль недоступен.','正确凭据可进入个人资料；错误凭据提示明确；未登录不能访问资料。'),
    ('US02\n3 / B','Чтение меню; список блюд; цена и доступность.','读取菜单、展示菜品、价格及可售状态。','Отображается меню выбранного текущего дня; пустое меню и сбой загрузки объяснены; недоступное блюдо нельзя выбрать.','展示当前日期菜单；空菜单和加载失败有提示；不可售菜品不能选择。'),
    ('US03\n5 / C','Добавление; количество; удаление; итоговая сумма.','添加菜品、修改数量、删除及计算总价。','Изменение количества меняет сумму; нулевое количество удаляет позицию; отрицательное количество не допускается.','数量变化时总价同步更新；数量为零则删除；不允许负数量。'),
    ('US04\n5 / D','Выбор времени; серверная проверка; запись заказа.','选择时间、服务端校验并保存订单。','Непустая корзина и допустимое время создают один заказ с номером; повторная отправка не создаёт дубль; оплата при получении указана явно.','非空购物车和有效时段可生成唯一订单号；重复提交不重复建单；明确显示到店付款。'),
    ('US06\n3 / E','Проверка роли; список; отображение деталей заказа.','员工权限校验、订单列表及订单明细展示。','Сотрудник видит созданный заказ, состав и время; студент не получает доступ к списку всех заказов.','工作人员能看到新订单、菜品及取餐时间；学生无权访问全体订单列表。'),
]
table([bi('ID SP Код','编号 点数\n牵头者'),bi('Работы','拆分工作'),bi('Проверяемый результат','可检验的结果')],
      [[a,bi(b,c),bi(d,e)] for a,b,c,d,e in SPRINT], [1.0,2.35,3.75],size=10.5,centered=(0,))
para('Ведущий координирует историю; реализация, ревью и тестирование выполняются совместно. Каждая история включает проверку и интеграцию в своей оценке, поэтому задачи не прибавляют очки повторно.',
     '牵头人协调该故事，开发、评审和测试由团队协作完成。每个故事的估算已包含测试与集成，拆分任务不重复累计故事点。', small=True)
head('1 7 Почему выбран такой объём', '1 7 范围选择理由')
para('Первые четыре истории дают 16 очков. Следующая US05 требует 8 и превысила бы предел: 16 + 8 = 24. Предлагается согласовать с владельцем продукта перенос US05 и взять US06 за 3 очка, чтобы столовая получила созданные заказы. Это решение зависит от допущения об оплате при получении.',
     '前四条故事合计16点，再加入US05的8点会达到24点并超限。因此本方案建议与产品负责人确认暂缓US05，加入3点的US06，使食堂能接收学生订单；该选择以允许到店付款为前提。')
para('US05 остаётся первым кандидатом следующего спринта; US07–US12 пока не выбираются. Дополнительная ёмкость в 1 очко не считается отдельной историей или обещанием дополнительной функции.',
     'US05保留为下一冲刺的首要候选；US07—US12本次不选。剩余1点容量不构成单独故事或额外功能承诺。')
page()

# Page 5
head('1 8 План на две недели', '1 8 两周实施安排')
para('План предполагает 10 рабочих дней. Разработка и проверка идут параллельно; поздние дни предназначены для интеграции и исправлений, а не для первого запуска тестов.',
     '按两周共10个工作日规划。开发和测试交错进行，后期用于集成与修复，测试从早期开始。', small=True)
table([bi('Срок','时间'),bi('Действия и ожидаемый результат','工作及预期结果')],[
    [bi('День 1','第1天'),bi('Планирование до 2 ч: цель, оценки, критерии, контракты данных и тестовые примеры; уточнить оплату при получении.','最多2小时规划：确定目标、估算、验收条件、数据接口及测试样例；澄清到店付款假设。')],
    [bi('Дни 2–3','第2—3天'),bi('US01 и US02; каркасы корзины и заказа; ранние тесты; E готовит проверку роли и интерфейс списка.','完成登录和菜单，推进购物车与订单骨架、早期测试；E准备员工权限及列表接口。')],
    [bi('Дни 4–6','第4—6天'),bi('US03 и US04; US06 по согласованному контракту; ежедневное объединение, ревью и проверки.','完成购物车与预约下单；按约定接口实现食堂订单；每日集成、评审及验证。')],
    [bi('Дни 7–9','第7—9天'),bi('Сквозной путь US01→US02→US03→US04→US06; отрицательные случаи, исправления и регрессия.','联调登录→菜单→购物车→下单→食堂查看；验证异常情况、修复并回归。')],
    [bi('День 10','第10天'),bi('Проверить DoD; обзор до 1 ч; ретроспектива до 45 мин; обновить бэклог по обратной связи.','检查DoD；最多1小时评审，最多45分钟回顾；据反馈更新待办列表。')],
], [1.05,6.05],size=10.5)
para('Каждый рабочий день — Daily Scrum 15 минут. На нём разработчики проверяют продвижение к цели и меняют план; при риске срыва обсуждают объём с владельцем продукта, сохраняя качество.',
     '每个工作日开展15分钟Daily Scrum，开发者检查目标进展并调整计划；若目标有风险，与产品负责人协商范围，同时保持质量。', small=True)
head('1 9 Definition of Done пять пунктов', '1 9 完成定义 共五项')
DOD = [
    bi('1 Все критерии приёмки истории проверены; результаты и примеры зафиксированы.', '1 故事的全部验收条件已验证，并记录结果及样例。'),
    bi('2 Код прошёл взаимное ревью и включён в общий собираемый вариант приложения.', '2 代码完成交叉评审，并已集成到可构建的应用版本。'),
    bi('3 Модульные, интеграционные и относящиеся к истории сквозные тесты проходят; блокирующих и критических дефектов нет.', '3 相关单元、集成与端到端测试通过，无阻塞或严重缺陷。'),
    bi('4 Проверены права студента и сотрудника, ошибочные данные и повторная отправка; пароли не хранятся и не записываются в журналы открытым текстом.', '4 验证学生与员工权限、无效输入及重复提交；密码不以明文保存或写入日志。'),
    bi('5 Инкремент запускается в учебной тестовой среде; инструкция и известные ограничения обновлены, путь заказа воспроизводим.', '5 增量可在课堂测试环境启动；操作说明与已知限制已更新，订餐流程可复现。'),
]
for x in DOD:
    item('p',x,small=True)
para('Это условия будущего завершения. Работа считается выполненной только при наличии фактических результатов; незавершённая история не приносит частичных очков скорости.',
     '以上是后续工作的完成条件；须有实际证据才能认定完成，未完成故事不计部分速度点数。', small=True)
page()

# Pages 6 and 7
KANBAN = [
 ('KB01','P1','—','Как пользователь, я хочу загрузить фото графика, чтобы начать оцифровку.','作为用户，我希望上传函数图像照片，以开始数字化。','JPG/PNG открывается в предпросмотре; повреждённый файл объяснён.','JPG/PNG可预览；损坏文件有明确提示。'),
 ('KB02','P1','KB01','Как пользователь, я хочу обрезать и повернуть фото, чтобы оставить нужный график.','作为用户，我希望裁剪并旋转照片，以保留所需图像区域。','Область и поворот видны до применения; исходник сохранён.','应用前可预览区域和角度，并保留原图。'),
 ('KB03','P1','KB02','Как пользователь, я хочу улучшить контраст и убрать шум, чтобы линия была различима.','作为用户，我希望增强对比度并去噪，使曲线更清晰。','Показаны исходник и обработка; настройки можно отменить.','可对照原图与处理结果，支持撤销设置。'),
 ('KB04','P1','KB03','Как пользователь, я хочу подавить сетку и подписи, чтобы они не стали точками функции.','作为用户，我希望抑制网格与文字，避免它们被识别为函数点。','На контрольном примере сетка удалена без разрыва линии; есть просмотр маски.','控制样例中网格被去除且曲线连续，可查看掩膜。'),
 ('KB05','P1','KB03','Как пользователь, я хочу найти оси координат, чтобы определить систему отсчёта.','作为用户，我希望识别坐标轴，以确定坐标系。','Оси и начало координат показаны; положение можно исправить.','显示坐标轴和原点，并可手动修正位置。'),
 ('KB06','P1','KB05','Как пользователь, я хочу задать значения делений, чтобы перевести пиксели в координаты.','作为用户，我希望设置刻度值，将像素位置转换为坐标。','По двум отметкам на каждой оси контрольные точки дают ожидаемые значения.','每轴用两个标定点校准后，控制点能得到预期坐标。'),
 ('KB07','P1','KB04 KB06','Как пользователь, я хочу выделить кривую и получить точки x, y, чтобы анализировать функцию.','作为用户，我希望提取曲线并得到x、y点集，以分析函数。','Точки показаны таблицей и наложением на фото; другие линии исключены.','以表格和叠加图展示点集，并排除其他线条。'),
 ('KB08','P1','KB07','Как пользователь, я хочу исправить и удалить ошибочные точки, чтобы улучшить данные.','作为用户，我希望修改或删除误识别点，以改善数据。','Добавление, изменение и удаление отражаются на графике; доступна отмена.','增删改点同步反映到图中，并可撤销。'),
 ('KB09','P1','KB08','Как пользователь, я хочу получить приближённую формулу, чтобы использовать функцию в расчётах.','作为用户，我希望获得近似公式，以便用于计算。','Для линейной или квадратичной модели показаны коэффициенты и область данных.','线性或二次模型能输出系数及数据适用区间。'),
 ('KB10','P1','KB09','Как пользователь, я хочу сравнить формулу с точками, чтобы оценить качество приближения.','作为用户，我希望比较公式与点集，以判断近似质量。','Показаны наложение кривой, среднеквадратичная ошибка и предупреждение о плохом соответствии.','展示拟合曲线叠加、均方根误差及拟合不佳提示。'),
 ('KB11','P1','KB07','Как пользователь, я хочу выгрузить точки в CSV, чтобы работать с ними в таблицах.','作为用户，我希望将点集导出为CSV，以便在表格软件中使用。','CSV содержит x и y; десятичный разделитель и порядок совпадают с предпросмотром.','CSV含x、y列，小数格式及顺序与预览一致。'),
 ('KB12','P2','KB09 KB10','Как пользователь, я хочу сохранить формулу и параметры, чтобы воспроизвести результат.','作为用户，我希望保存公式与参数，以便复现结果。','TXT содержит формулу, коэффициенты, диапазон и ошибку; файл открывается.','TXT包含公式、系数、区间及误差，文件可正常打开。'),
 ('KB13','P1','KB01','Как пользователь, я хочу понятные сообщения о непригодном фото, чтобы исправить входные данные.','作为用户，我希望收到照片不适用的明确提示，以调整输入。','Размытое фото или недостаток контрольных точек не приводят к молчаливому неверному результату.','模糊图像或标定点不足时明确提示，不静默输出错误结果。'),
 ('KB14','P1','—','Как тестировщик, я хочу эталонный набор изображений и точек, чтобы проверять точность изменений.','作为测试者，我希望建立带已知点集的图像样例，以验证变更的准确性。','Есть прямая, парабола и зашумлённое фото; эталон, допуск и сценарий проверки записаны.','包含直线、抛物线及噪声图片，并记录参考点、容差和检查步骤。'),
]
head('2 Кейс Kanban График по фотографии', '2 案例二 Kanban照片中的函数图像数字化')
para('Программа получает фотографию графика функции и возвращает набор координатных точек и приближённую формулу. Цель — путь от загрузки фото до проверяемого цифрового результата.',
     '程序接收函数图像照片，输出坐标点集及近似公式，目标是完成从照片上传到可核对数字结果的流程。')
para('Для начальной версии предполагаются одна выбранная кривая, линейные шкалы осей и модели первой/второй степени. Логарифмические шкалы, несколько кривых и сложные модели требуют отдельных историй. Эти границы предложены для планирования.',
     '首版规划假设处理一条选定曲线、线性坐标轴和一次或二次拟合模型；对数坐标、多曲线及复杂模型需另立故事。以上为本方案提出的范围假设。', small=True)
head('2 1 Бэклог продукта 14 элементов', '2 1 产品待办列表 共14项')
para('P1 — основной поток и достоверность; P2 — удобство воспроизведения. Зависимости задают порядок готовности к работе, а не календарные сроки. Общие проверки качества входят в каждый элемент.',
     'P1表示核心流程与结果可靠性，P2表示复现便利性。依赖关系约束条目何时可开始，不代表日历工期；每项均包含相关质量检查。', small=True)
KH = [bi('ID P\nЗавис.','编号 级别\n依赖'),bi('История или задача','用户故事或任务'),bi('Критерий приёмки','验收条件')]
def kb_rows(subset):
    return [[a+'\n'+b+'\n'+c,bi(d,e),bi(f,g)] for a,b,c,d,e,f,g in subset]
table(KH,kb_rows(KANBAN[:7]),[0.9,3.45,2.75],size=10.5,centered=(0,))
page()
head('2 1 Бэклог продукта продолжение', '2 1 产品待办列表 续')
table(KH,kb_rows(KANBAN[7:]),[0.9,3.45,2.75],size=10.5,centered=(0,))
para('Первый ценный путь: KB01→KB02→KB03, затем очистка KB04 и калибровка KB05→KB06, после чего KB07→KB08→KB09→KB10; KB11 даёт экспорт точек. KB13 и KB14 поддерживают достоверность на всём пути.',
     '首条价值路径为KB01→KB02→KB03，随后进行KB04清理与KB05→KB06标定，再完成KB07→KB08→KB09→KB10；KB11提供点集导出。KB13和KB14持续支持过程可靠性。', small=True)
page()

# Page 8 landscape
head('2 2 Доска Kanban и ограничения WIP', '2 2 Kanban看板与WIP限制')
para('Ниже — учебный снимок возможного состояния, а не сведения о фактически выполненной разработке. Все 14 карточек размещены ровно один раз; готовность зависимостей соблюдена.',
     '下表为可行状态的课堂示意，不代表实际开发进度。14张卡片各出现一次，进入准备和执行阶段的依赖均已满足。',small=True)
BOARD = [
    (bi('Бэклог\nБез лимита','产品待办\n无上限'), ['KB07 Точки\n点集提取','KB08 Правка точек\n修正点集','KB09 Формула\n近似公式','KB10 Качество\n拟合质量','KB11 CSV\n点集导出','KB12 Сохранение\n保存结果']),
    (bi('Готово к работе\nЛимит 3','准备就绪\n队列上限3'), ['KB14 Эталоны\n参考样例集']),
    (bi('Разработка\nWIP 2','开发中\nWIP 2'), ['KB13 Ошибки\n异常提示']),
    (bi('Ревью\nWIP 1','评审中\nWIP 1'), ['KB06 Калибровка\n坐标标定']),
    (bi('Проверка\nWIP 1','测试中\nWIP 1'), ['KB04 Сетка\n去除网格']),
    (bi('Готово\nБез лимита','已完成\n无上限'), ['KB01 Загрузка\n图片上传','KB02 Обрезка\n裁剪旋转','KB03 Очистка\n图像预处理','KB05 Оси\n坐标轴识别']),
]
item('board', BOARD)
para('Направление потока: Бэклог → Готово к работе → Разработка → Ревью → Проверка → Готово. В снимке: очередь 1/3, разработка 1/2, ревью 1/1, проверка 1/1; общий активный WIP = 3/4.',
     '流向：产品待办→准备就绪→开发中→评审中→测试中→已完成。示意占用：准备1/3、开发1/2、评审1/1、测试1/1；全流程活动WIP为3/4。')
head('2 3 Почему выбраны такие лимиты', '2 3 WIP限制的理由')
para('Для расчёта начальных лимитов предполагается пять участников: двое могут вести разработку, один — ревью, один — проверку, пятый помогает с анализом и препятствиями; роли можно менять. Размер команды Kanban в условии не задан, поэтому это отдельное допущение.',
     '初始限额假设有5名可协作成员：2人开发、1人评审、1人测试，第5人协助分析与排障，职责可流动。题目未指定Kanban团队规模，因此这是独立规划假设。',small=True)
para('Очередь 3 обеспечивает небольшой запас подготовленных задач. Разработка 2 ограничивает переключения, ревью 1 и проверка 1 делают узкие места заметными. Общий WIP 4 ограничивает все начатые, но незавершённые карточки, включая ожидание и блокировки.',
     '准备队列上限3，用于保留少量可拉取任务；开发上限2减少切换，评审和测试各上限1便于发现瓶颈。活动WIP总上限4，统计所有已开始但未完成的卡片，包括等待和阻塞。',small=True)
page(landscape=True)

# Page 9
head('2 4 Явные правила движения карточек', '2 4 卡片流转规则')
POLICIES = [
    bi('1 В «Готово к работе» попадает элемент с понятным результатом, критерием приёмки и завершёнными зависимостями. Входной бэклог и эта очередь ещё не входят в активный WIP.', '1 条目须明确结果和验收条件、满足依赖，才可进入准备队列；产品待办及准备队列不计活动WIP。'),
    bi('2 Начало — переход в «Разработку». Карточку вытягивают только при свободном месте в целевой колонке и общем WIP менее 4; приоритет учитывается среди готовых элементов.', '2 进入开发列即为开始。目标列有空位且活动WIP小于4时才能拉取新工作；在可开始条目中按优先级选择。'),
    bi('3 Ревью требует готового изменения и самопроверки; проверка — пройденного ревью и воспроизводимого примера. «Готово» означает выполненные критерии, пройденные проверки, интеграцию и обновлённую инструкцию.', '3 进入评审需提交完整变更并自检；进入测试需通过评审且有可复现场景；完成须满足验收条件、通过检查、集成并更新说明。'),
    bi('4 Когда следующая колонка заполнена, карточка остаётся в текущей и продолжает занимать её WIP. Участники помогают завершить и проверить начатое, прежде чем вытягивать новую работу.', '4 下游列已满时，卡片留在当前列并继续占用其限额；团队先协助完成或检查已有工作，再拉取新条目。'),
    bi('5 Заблокированная карточка остаётся на месте с причиной и ответственным за снятие блока. Дефект внутри проверки исправляют в той же карточке; возврат в разработку допускается при соблюдении лимитов.', '5 阻塞卡片留在原列，标注原因及解除阻塞负责人；测试发现的缺陷在同一卡片中修复，回退开发也须遵守限额。'),
]
for x in POLICIES:
    item('p',x,small=True)
head('2 5 Проверка и настройка потока', '2 5 流程检查与限额调整')
para('Ежедневно команда просматривает доску справа налево, отмечает возраст активных карточек и снимает блокировки. Для каждого элемента фиксируются начало и завершение; анализируются WIP, пропускная способность, возраст и время цикла.',
     '团队每日从右向左检查，关注活动卡片年龄并清除阻塞；逐项记录开始与完成时间，观察活动WIP、吞吐量、工作项年龄和周期时间。',small=True)
para('Начальное SLE как проверяемое допущение: 85% небольших готовых к работе элементов завершаются не более чем за 5 рабочих дней от начала разработки. Это прогноз без исторической статистики; после накопления данных его и WIP пересматривают. Крупные элементы дробят до начала, сохраняя проверяемый результат.',
     '初始服务水平预期SLE暂设为：85%的小型可实施条目从开始开发到完成不超过5个工作日。该预测尚无历史数据支撑，积累数据后连同WIP一起调整；大条目开始前须拆成可检验的结果。',small=True)
head('3 Источники', '3 参考依据')
para('Исходные условия — раздаточный материал и слайд практического задания, предоставленные в виде двух фотографий. Числовые оценки и лимиты в решении — учебный план.',
     '题目条件来自用户提供的讲义与课堂课件两张图片；本解答中的估值和限额为课程规划方案。',small=True)
para('[1] Schwaber K., Sutherland J. The Scrum Guide. 2020. Роли, планирование спринта, Sprint Backlog и Definition of Done.',
     '[1] Scrum指南，2020年版。用于核对角色、冲刺规划、冲刺待办与完成定义。',small=True)
item('link',{'label':'https://scrumguides.org/scrum-guide.html','url':'https://scrumguides.org/scrum-guide.html'})
para('[2] The Kanban Guide. May 2025. Визуализация потока, контроль WIP, вытягивание, SLE и метрики потока.',
     '[2] Kanban指南，2025年5月版。用于核对流程可视化、WIP控制、拉取、SLE与流动指标。',small=True)
item('link',{'label':'https://kanbanguides.org/the-kanban-guide/2025.5/','url':'https://kanbanguides.org/the-kanban-guide/2025.5/'})
para('Дата обращения к обоим руководствам 15.10.2026.', '两份指南查阅日期均为2026年10月15日。',small=True)
page()

# Keep the board on one landscape page and place its rationale in portrait prose.
board_page = PAGES[7]
rationale_index = next(i for i,b in enumerate(board_page['items']) if b['kind']=='h' and b['value']['ru'].startswith('2 3'))
PAGES[8]['items'] = board_page['items'][rationale_index:] + PAGES[8]['items']
board_page['items'] = board_page['items'][:rationale_index]
# Keep one continuous Kanban backlog table with repeating headers.
kb_first = next(b for b in PAGES[5]['items'] if b['kind']=='table')
kb_second = next(b for b in PAGES[6]['items'] if b['kind']=='table')
kb_first['value']['rows'].extend(kb_second['value']['rows'])
value_path = PAGES[6]['items'][-1]
PAGES[5]['items'].insert(PAGES[5]['items'].index(kb_first), value_path)
PAGES[6]['items'] = []

# Persist normalized inventory before document creation.
ROOT.joinpath('data/content.json').write_text(json.dumps(PAGES,ensure_ascii=False,indent=2),encoding='utf-8')
ROOT.joinpath('data/fonts.conf').write_text('''<?xml version="1.0"?>
<!DOCTYPE fontconfig SYSTEM "urn:fontconfig:fonts.dtd">
<fontconfig>
<dir>/System/Library/Fonts</dir><dir>/System/Library/Fonts/Supplemental</dir>
<dir>/Library/Fonts</dir><dir>/Users/zhi/Library/Fonts</dir>
<cachedir>/private/tmp/scrum-kanban-fontcache</cachedir>
<match target="pattern"><test name="family"><string>sans-serif</string></test><edit name="family" mode="prepend"><string>Arial</string><string>PingFang SC</string></edit></match>
</fontconfig>''',encoding='utf-8')

doc = Document()
doc.core_properties.title = '任务二 Scrum冲刺规划与Kanban看板设计'
doc.core_properties.subject = '中俄双语软件生命周期课程实践'
doc.core_properties.author = ''
doc.core_properties.last_modified_by = ''
doc.core_properties.version = '1.0'
doc.core_properties.created = datetime(2026,10,15,0,0,tzinfo=timezone.utc)
doc.core_properties.modified = datetime(2026,10,15,0,0,tzinfo=timezone.utc)

def font(style,size,bold=False):
    style.font.name='Arial'
    style.font.size=Pt(size)
    style.font.bold=bold
    style.font.color.rgb=RGBColor(0,0,0)
    rf=style.element.get_or_add_rPr().get_or_add_rFonts()
    for k,v in [('ascii','Arial'),('hAnsi','Arial'),('eastAsia','PingFang SC'),('cs','Arial')]:
        rf.set(qn('w:'+k),v)

for name,size,bold in [('Normal',11,False),('Title',16,True),('Heading 1',13.5,True),('Heading 2',12,True)]:
    font(doc.styles[name],size,bold)
    pf=doc.styles[name].paragraph_format
    pf.space_after=Pt(6)
    pf.space_before=Pt(0 if name=='Normal' else 8)
    pf.line_spacing=1.08
    pf.widow_control=True
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
for sty in doc.styles:
    for border in sty.element.findall('.//'+qn('w:pBdr')):
        border.getparent().remove(border)

def section_setup(sec,landscape):
    sec.orientation=WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    sec.page_width=Inches(11 if landscape else 8.5)
    sec.page_height=Inches(8.5 if landscape else 11)
    sec.top_margin=Inches(.62)
    sec.bottom_margin=Inches(.62)
    sec.left_margin=Inches(.7)
    sec.right_margin=Inches(.7)
    sec.header_distance=Inches(.2)
    sec.footer_distance=Inches(.2)
    for grid in sec._sectPr.findall(qn('w:docGrid')):
        sec._sectPr.remove(grid)

def run_text(p,text,size=None,bold=None,color=None):
    r=p.add_run(text)
    r.font.name='Arial'
    rf=r._element.get_or_add_rPr().get_or_add_rFonts()
    rf.set(qn('w:eastAsia'),'PingFang SC')
    if size:
        r.font.size=Pt(size)
    if bold is not None:
        r.bold=bold
    r.font.color.rgb=RGBColor.from_string(color or '000000')
    return r

def write_value(p,val,size=11,bold=False,color='000000'):
    if isinstance(val,dict):
        run_text(p,val['ru'],size,bold,color)
        run_text(p,'\n'+val['zh'],size,bold,color)
    else:
        run_text(p,str(val),size,bold,color)

TABLE_DIMS=[]
def draw_table(val,board=False):
    headers=val['headers']; rows=val['rows']; widths=val['widths']; size=val.get('size',10.5)
    t=doc.add_table(rows=1,cols=len(headers))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    for c,w in zip(t.columns,widths):
        c.width=Inches(w)
    props=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for name in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+name);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    props.append(borders)
    for ri,values in enumerate([headers]+rows):
        row=t.rows[0] if ri==0 else t.add_row()
        pr=row._tr.get_or_add_trPr()
        no_split=OxmlElement('w:cantSplit');pr.append(no_split)
        if ri==0:
            rep=OxmlElement('w:tblHeader');pr.append(rep)
        for ci,(cell,value,w) in enumerate(zip(row.cells,values,widths)):
            cell.width=Inches(w)
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=cell._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for edge,n in [('top',75),('bottom',75),('left',90),('right',90)]:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),str(n));e.set(qn('w:type'),'dxa');margins.append(e)
            cp.append(margins)
            shade=OxmlElement('w:shd')
            shade.set(qn('w:fill'),'DCE6F1' if ri==0 else ('F4F7FA' if ri%2==0 else 'FFFFFF'))
            cp.append(shade)
            p=cell.paragraphs[0]
            p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.line_spacing=1.03
            p.paragraph_format.keep_with_next=False
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER if ci in val.get('centered',()) or ri==0 else WD_ALIGN_PARAGRAPH.LEFT
            if board and ri>0:
                for n,card in enumerate(value):
                    cp_=p if n==0 else cell.add_paragraph()
                    cp_.paragraph_format.space_after=Pt(7)
                    cp_.paragraph_format.line_spacing=1.0
                    write_value(cp_,card,size=size)
            else:
                write_value(p,value,size=size,bold=ri==0)
    TABLE_DIMS.append([len(rows)+1,len(headers)])
    p=doc.add_paragraph()
    p.paragraph_format.space_after=Pt(1)
    p.paragraph_format.space_before=Pt(0)
    p.paragraph_format.line_spacing=1
    run_text(p,'',size=2)
    return t

def hyperlink(val):
    p=doc.add_paragraph()
    p.paragraph_format.space_after=Pt(4)
    rel=doc.part.relate_to(val['url'],'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rel)
    r=OxmlElement('w:r');rp=OxmlElement('w:rPr')
    c=OxmlElement('w:color');c.set(qn('w:val'),'19558A');rp.append(c)
    s=OxmlElement('w:sz');s.set(qn('w:val'),'19');rp.append(s)
    r.append(rp);text=OxmlElement('w:t');text.text=val['label'];r.append(text);h.append(r);p._p.append(h)

prev=False
for pi,pg in enumerate(PAGES):
    landscape=pg['landscape']
    if pi==0:
        section_setup(doc.sections[0],landscape)
    elif landscape!=prev:
        section_setup(doc.add_section(WD_SECTION_START.NEW_PAGE),landscape)
    for block in pg['items']:
        kind=block['kind'];v=block['value']
        if kind in ('p','h','title'):
            p=doc.add_paragraph(style='Title' if kind=='title' else ('Heading 1' if kind=='h' else 'Normal'))
            if kind=='h':
                p.paragraph_format.keep_with_next=True
            p.paragraph_format.keep_together=True
            snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');p._p.get_or_add_pPr().append(snap)
            write_value(p,v,size=16 if kind=='title' else (13.5 if kind=='h' else (10.5 if block.get('small') else 11)),bold=kind in ('title','h'))
        elif kind=='table':
            draw_table(v)
        elif kind=='board':
            draw_table({'headers':[x[0] for x in v],'rows':[[x[1] for x in v]],'widths':[1.6]*6,'size':10.5},board=True)
        elif kind=='link':
            hyperlink(v)
    prev=landscape

doc.save(OUT)
loaded=Document(OUT)
with ZipFile(OUT) as z:
    xml=z.read('word/document.xml')
    root=etree.fromstring(xml)
    full='\n'.join(root.xpath('//w:t/text()',namespaces={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}))

def all_strings(obj):
    if isinstance(obj,dict):
        if 'ru' in obj and 'zh' in obj:
            yield obj['ru'];yield obj['zh']
        else:
            for k,v in obj.items():
                if k not in ('kind','widths','size','centered','landscape','url'):
                    yield from all_strings(v)
    elif isinstance(obj,(list,tuple)):
        for v in obj:yield from all_strings(v)
    elif isinstance(obj,str):yield obj

required=list(all_strings(PAGES))
missing=[s for s in required if s not in full]
actual_dims=[[len(t.rows),len(t.columns)] for t in loaded.tables]
assert not missing,missing
assert actual_dims==TABLE_DIMS,(actual_dims,TABLE_DIMS)
assert sum(x[2] for x in ESTIMATES if x[0] in ['US01','US02','US03','US04','US06'])==19
assert all(x[2] in [1,2,3,5,8,13] for x in ESTIMATES)
assert len(KANBAN)>=10 and len(DOD)==5 and len(STORIES)==12
assert sum(len(x[1]) for x in BOARD)==14
assert len(set(re.findall(r'KB\d{2}',' '.join(' '.join(x[1]) for x in BOARD))))==14
assert not any('header' in n or 'footer' in n for n in ZipFile(OUT).namelist())
report={
    'artifact':str(OUT),'date':'2026-10-15','environment':'Codex local desktop bundled Python and LibreOffice',
    'status':'STRUCTURAL_PASS','required_text_items':len(required),'missing_text':missing,
    'tables':actual_dims,'stories_original':12,'scrum_estimated':6,'sprint_story_ids':['US01','US02','US03','US04','US06'],
    'sprint_points':19,'capacity':20,'dod_items':5,'kanban_items':14,
    'board_each_item_once':True,'board_snapshot_is_illustrative':True,'active_wip':3,'active_wip_limit':4,
    'planned_pages':9,'headers_footers':False,'visual_review':'PENDING',
    'sha256':hashlib.sha256(OUT.read_bytes()).hexdigest()
}
ROOT.joinpath('data/verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False,indent=2))
