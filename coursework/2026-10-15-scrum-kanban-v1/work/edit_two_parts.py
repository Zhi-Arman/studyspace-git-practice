from pathlib import Path
from copy import deepcopy
from datetime import datetime, timezone
from zipfile import ZipFile
import ast
import json
import re
import hashlib
from lxml import etree
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path('/Users/zhi/Documents/软件生命周期/studyspace/coursework/2026-10-15-scrum-kanban-v1')
BEFORE=ROOT/'work/before-two-parts'
SOURCE=BEFORE/'任务二_Scrum冲刺规划与Kanban看板设计_中俄双语_v1.docx'
OUT=ROOT/'任务一与任务二_Scrum冲刺规划与Kanban看板设计_中俄双语_v2.docx'
original_inventory=json.loads((BEFORE/'content.json').read_text())

def bi(ru,zh):return {'ru':ru,'zh':zh}
def paragraph(ru,zh,small=False):return {'kind':'p','value':bi(ru,zh),'small':small}
def heading(ru,zh,major=False):return {'kind':'h','value':bi(ru,zh),'major':major}

summary=[
 heading('Задание 2 Комплексная практика Scrum и Kanban','任务二 课件综合练习与Kanban规划',True),
 paragraph('Эта часть объединяет задание 2 первого изображения и весь второй слайд: бэклог Kanban, а также роли, цель, оценки, выбор спринта и Definition of Done для Scrum. Для одного и того же кейса столовой используется согласованная учебная модель из задания 1.',
           '本部分合并第一张图的任务二与第二张课件的完整内容：除Kanban待办与看板外，还包含Scrum角色、目标、估算、冲刺选择及完成定义。同一食堂案例采用与任务一一致的课堂模拟方案。',True),
 heading('2 1 Scrum для кейса столовой по второму слайду','2 1 第二张课件中的食堂订餐Scrum方案'),
 paragraph('Роли: A — владелец продукта и разработчик; B — Scrum-мастер и разработчик; C, D, E — разработчики. Все пять человек участвуют в реализации; A отвечает за порядок бэклога и ценность, B помогает организовать события и снять препятствия.',
           '角色：A为产品负责人兼开发者，B为Scrum Master兼开发者，C、D、E为开发者。五人共同参与实现；A负责待办排序与价值，B负责活动组织与排障协调。',True),
 paragraph('Цель спринта: за две недели обеспечить студенту путь от входа и выбора блюд до заказа на выбранное время с оплатой при получении, а сотруднику столовой — просмотр поступившего заказа.',
           '冲刺目标：在两周内让学生完成从登录、选餐到预约下单并到店付款的流程，同时让食堂工作人员能够查看收到的订单。',True),
 paragraph('Покер планирования: карты 1, 2, 3, 5, 8, 13; эталон входа 3 очка. Ниже повторены модельный первый раунд и итог после обсуждения; это учебная оценка, а не протокол встречи. Обоснования приведены в разделе 1 4.',
           '规划扑克采用1、2、3、5、8、13，登录基准为3点。下表重列模拟首轮投牌及讨论后点数；这是一份课堂估算，不是会议记录，具体依据见1 4节。',True),
 {'kind':'table','value':{
     'headers':[bi('ID','编号'),bi('История','用户故事'),bi('Очки','点数'),bi('Карты A–E','A—E投牌')],
     'rows':[
         ['US01',bi('Вход по логину и паролю','账号密码登录'),'3','3 3 3 3 3'],
         ['US02',bi('Меню на сегодня','今日菜单'),'3','2 3 3 5 3'],
         ['US03',bi('Корзина заказа','购物车'),'5','3 5 5 5 8'],
         ['US04',bi('Заказ на выбранное время','预约取餐时间并下单'),'5','5 5 8 5 5'],
         ['US05',bi('Оплата картой в приложении','应用内银行卡支付'),'8','8 13 8 8 13'],
         ['US06',bi('Список заказов для столовой','食堂收到的订单列表'),'3','3 3 5 3 3'],
     ],'widths':[.65,3.35,.65,2.45],'size':10.5,'centered':[0,2,3]
 }},
 paragraph('Спринт две недели, ёмкость всей команды 20 очков. Выбраны US01, US02, US03, US04, US06: 3 + 3 + 5 + 5 + 3 = 19 ≤ 20. Ведущие соответственно A, B, C, D, E; остаётся 1 очко незагруженной ёмкости.',
           '冲刺两周，全队容量20点。选择US01、US02、US03、US04、US06，合计3＋3＋5＋5＋3＝19点，不超过20点；牵头者依次为A、B、C、D、E，剩余1点未分配容量。',True),
 paragraph('US05 переносится: вместе с первыми четырьмя историями получилось бы 24 очка. US06 замыкает получение заказа столовой. Решение опирается на допущение об оплате при получении; при обязательной предоплате план нужно пересмотреть. Подробные работы и двухнедельный план — в разделах 1 6–1 8.',
           'US05暂缓，因为与前四条合计会达到24点；US06使食堂能够接收订单。该方案以允许到店付款为假设，若必须预付则应调整。任务拆分与两周实施安排见1 6至1 8节。',True),
 heading('2 1 1 Definition of Done пять пунктов','2 1 1 完成定义 共五项'),
]
# Preserve the same five DoD requirements rather than introducing conflicting rules.
dod_source=original_inventory[4]['items']
dod_index=next(i for i,b in enumerate(dod_source) if b['kind']=='h' and b['value']['ru'].startswith('1 9'))
summary.extend(deepcopy(dod_source[dod_index+1:dod_index+6]))

# Inventory all changes before editing the existing Word document.
changes={
 'date':'2026-10-15','version':'2.0','output':str(OUT),
 'part_1':'First image Task 1 Scrum only',
 'part_2':'First image Task 2 Kanban plus all requirements on the second image',
 'new_blocks':summary,
 'preserved_source_inventory':str(BEFORE/'content.json'),
 'heading_mapping':{
   '1 Кейс Scrum Заказ еды в столовой':bi('Задание 1 Планирование спринта Scrum','任务一 Scrum冲刺规划'),
   '2 Кейс Kanban График по фотографии':bi('2 2 Kanban Оцифровка графика по фотографии','2 2 Kanban照片中的函数图像数字化'),
   '2 1 Бэклог продукта 14 элементов':bi('2 2 1 Бэклог продукта 14 элементов','2 2 1 产品待办列表 共14项'),
   '2 2 Доска Kanban и ограничения WIP':bi('2 3 Доска Kanban и ограничения WIP','2 3 Kanban看板与WIP限制'),
   '2 3 Почему выбраны такие лимиты':bi('2 4 Почему выбраны такие лимиты','2 4 WIP限制的理由'),
   '2 4 Явные правила движения карточек':bi('2 5 Явные правила движения карточек','2 5 卡片流转规则'),
   '2 5 Проверка и настройка потока':bi('2 6 Проверка и настройка потока','2 6 流程检查与限额调整'),
   '3 Источники':bi('Источники','参考依据')
 }
}
(ROOT/'data/two-parts-content.json').write_text(json.dumps(changes,ensure_ascii=False,indent=2))

doc=Document(SOURCE)
original_doc=Document(SOURCE)
table_text_before=[[ [c.text for c in row.cells] for row in t.rows] for t in original_doc.tables]
helper_names={'font','run_text','write_value','draw_table','all_strings'}
tree=ast.parse((ROOT/'work/build_document.py').read_text())
helper_tree=ast.Module(body=[n for n in tree.body if isinstance(n,ast.FunctionDef) and n.name in helper_names],type_ignores=[])
exec(compile(helper_tree,str(ROOT/'work/build_document.py'),'exec'))
TABLE_DIMS=[]

def replace_pair(p,value,style=None,size=13.5):
    p.clear()
    if style:p.style=style
    p.paragraph_format.keep_together=True
    p.paragraph_format.keep_with_next=bool(style and style.startswith('Heading'))
    write_value(p,value,size=size,bold=bool(style and (style.startswith('Heading') or style=='Title')))

kanban_anchor=next(p for p in doc.paragraphs if p.text.startswith('2 Кейс Kanban'))
title=next(p for p in doc.paragraphs if p.style.name=='Title')
replace_pair(title,bi('Задания 1 и 2 Планирование Scrum и доска Kanban','任务一与任务二 Scrum冲刺规划与Kanban看板设计'),'Title',16)
date_para=next(p for p in doc.paragraphs if p.text.startswith('Практическая работа по'))
replace_pair(date_para,bi('Практическая работа по жизненному циклу программного обеспечения. Версия 2.0. Дата 15.10.2026.','软件生命周期课程实践作业。版本2.0。日期2026年10月15日。'),'Normal',10.5)

for p in list(doc.paragraphs):
    first=p.text.split('\n',1)[0]
    if first in changes['heading_mapping']:
        major=first.startswith('1 Кейс Scrum')
        replace_pair(p,changes['heading_mapping'][first],'Heading 1' if major else 'Heading 2')
    elif p.style.name=='Heading 1' and re.match(r'^1 \d',first):
        p.style='Heading 2'
    # Source description is retained as a bibliographic note with clear section coverage.
    if first.startswith('Исходные условия — раздаточный материал'):
        replace_pair(p,bi('Первое изображение задаёт два задания; второй слайд объединяет практику Scrum и Kanban. В задании 1 приведено решение первой части первого изображения; задание 2 объединяет его вторую часть и весь второй слайд. Оценки и лимиты — учебный план.','第一张图片包含两个任务，第二张课件汇总Scrum与Kanban练习。本文任务一对应首图的任务一；任务二合并首图的任务二及第二张课件的全部要求。估算和限额为课堂规划。'),'Normal',10.5)

# Insert actual Scrum answers into Part 2, ahead of the preserved Kanban section.
for block in summary:
    old_children=set(doc._element.body)
    if block['kind']=='table':
        draw_table(block['value'])
    else:
        major=block.get('major',False)
        p=doc.add_paragraph(style='Heading 1' if major else ('Heading 2' if block['kind']=='h' else 'Normal'))
        if major:p.paragraph_format.page_break_before=True
        p.paragraph_format.keep_together=True
        p.paragraph_format.keep_with_next=block['kind']=='h'
        snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');p._p.get_or_add_pPr().append(snap)
        write_value(p,block['value'],size=13.5 if block['kind']=='h' else (10.5 if block.get('small') else 11),bold=block['kind']=='h')
    new_children=[e for e in doc._element.body if e not in old_children]
    for e in new_children:
        kanban_anchor._p.addprevious(e)

for name in ('Title','Heading 1','Heading 2'):
    doc.styles[name].font.color.rgb=RGBColor(0,0,0)
doc.core_properties.title='任务一与任务二 Scrum冲刺规划与Kanban看板设计'
doc.core_properties.version='2.0'
doc.core_properties.modified=datetime(2026,10,15,0,0,tzinfo=timezone.utc)
doc.save(OUT)

loaded=Document(OUT)
with ZipFile(OUT) as z:
    x=etree.fromstring(z.read('word/document.xml'))
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
text='\n'.join(x.xpath('//w:t/text()',namespaces=ns))
all_tables=[[[c.text for c in row.cells] for row in t.rows] for t in loaded.tables]
assert len(all_tables)==8
assert all_tables[:5]==table_text_before[:5]
assert all_tables[6:]==table_text_before[5:]
assert all_tables[5][1:]==[[a,b['ru']+'\n'+b['zh'],c,d] for a,b,c,d in summary[6]['value']['rows']]
all_text=list(all_strings(summary))+list(all_strings(list(changes['heading_mapping'].values())))
assert not [s for s in all_text if s not in text]
major_headings=[p.text for p in loaded.paragraphs if p.style.name=='Heading 1']
assert major_headings==[
 'Задание 1 Планирование спринта Scrum\n任务一 Scrum冲刺规划',
 'Задание 2 Комплексная практика Scrum и Kanban\n任务二 课件综合练习与Kanban规划']
paras=[p.text for p in loaded.paragraphs]
part1_index=next(i for i,s in enumerate(paras) if s.startswith('Задание 1 '))
part2_index=next(i for i,s in enumerate(paras) if s.startswith('Задание 2 '))
assert part1_index<part2_index
assert not any('Kanban' in s for s in paras[part1_index:part2_index])
assert all(k in '\n'.join(paras[part2_index:]) for k in ['2 1 Scrum','2 2 Kanban','2 3 Доска','2 1 1 Definition of Done'])
assert len(x.xpath('//w:tbl/w:tr[1]/w:trPr/w:tblHeader',namespaces=ns))==8
report={
 'artifact':str(OUT),'source_document':str(SOURCE),'date':'2026-10-15','version':'2.0',
 'environment':'Codex local desktop bundled Python and LibreOffice',
 'status':'STRUCTURAL_PASS','major_parts':2,'part1_scope':'First image Task 1 only',
 'part2_scope':'First image Task 2 plus all of second image',
 'original_tables_preserved':True,'table_count':8,
 'task1_estimated_stories':6,'task2_estimated_stories':6,
 'sprint_points_both_parts':19,'capacity':20,'dod_count_each_part':5,
 'kanban_items':14,'board_preserved':True,'visual_review':'PENDING',
 'sha256':hashlib.sha256(OUT.read_bytes()).hexdigest(),
 'source_sha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest()
}
(ROOT/'data/two-parts-verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
print(json.dumps(report,ensure_ascii=False,indent=2))
