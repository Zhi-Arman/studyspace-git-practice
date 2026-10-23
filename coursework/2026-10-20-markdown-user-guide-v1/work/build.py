from pathlib import Path
import hashlib
import json
import re
from zipfile import ZipFile

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = Path(__file__).resolve().parents[1]
BASE.mkdir(parents=True, exist_ok=True)
DOCX = BASE / 'VSCode_Markdown初学者操作指南_中俄双语.docx'
MD = BASE / 'user-guide.md'

TITLE = ('VS Code Markdown 初学者操作指南', 'Руководство по Markdown в VS Code для начинающих')
INTRO = (
    '读者：第一次在 VS Code 中使用 Markdown 的学生。先完成任务 A，再完成任务 B，即可创建、保存并预览自己的第一个 Markdown 文件。',
    'Читатель: студент, который впервые работает с Markdown в VS Code. Выполните задачу A, затем задачу B, чтобы создать, сохранить и просмотреть свой первый Markdown-файл.'
)
NOTE = (
    '菜单和命令采用英文界面名称；中文或俄文界面可选择对应功能。以下操作使用 VS Code 内置功能，无需安装额外扩展。',
    'Названия меню и команд приведены для английского интерфейса; в китайском или русском интерфейсе выбирайте соответствующие пункты. Используются встроенные функции VS Code, дополнительные расширения не нужны.'
)
SAMPLE = '# 我的第一份 Markdown 文档\n\nЭто мой первый документ Markdown.'
TASKS = [
    {
        'id': 'A',
        'heading': ('任务 A 打开学习文件夹并创建和保存文件', 'Задача A Открыть учебную папку создать и сохранить файл'),
        'goal': ('在学习文件夹中创建并保存 user-guide.md。', 'Создать и сохранить user-guide.md в учебной папке.'),
        'initial': (
            '已安装并启动 VS Code；已准备好学习文件夹，例如 markdown-practice，且知道其位置；文件夹中尚无 user-guide.md。',
            'VS Code установлен и запущен. Учебная папка, например markdown-practice, уже существует, её расположение известно. Файла user-guide.md в ней пока нет.'
        ),
        'steps': [
            ('选择“File → Open Folder…”（文件 → 打开文件夹）。', 'Выберите File → Open Folder… (Файл → Открыть папку).'),
            ('选中学习文件夹并确认打开；若询问是否信任，对自己创建的文件夹选择信任。左侧 Explorer（资源管理器）应显示该文件夹。', 'Выберите учебную папку и подтвердите открытие. Если появится запрос доверия, подтвердите его для созданной вами папки. Её имя появится слева в Explorer (Проводник).'),
            ('在左侧 Explorer 中右键单击学习文件夹，选择“New File…”（新建文件）。', 'В Explorer щёлкните правой кнопкой по учебной папке и выберите New File… (Создать файл).'),
            ('输入 user-guide.md 并按 Enter；保留 .md 扩展名。', 'Введите user-guide.md и нажмите Enter. Сохраните расширение .md.'),
            ('在编辑区输入下方示例；“#”后保留一个空格，不要输入示例外围的三反引号。', 'Введите в редакторе пример ниже. После # оставьте один пробел; тройные обратные кавычки вокруг примера вводить не нужно.'),
            ('选择“File → Save”（文件 → 保存），或按 macOS 的 Cmd+S、Windows/Linux 的 Ctrl+S。', 'Выберите File → Save (Файл → Сохранить) или нажмите Cmd+S в macOS, Ctrl+S в Windows/Linux.'),
            ('检查左侧学习文件夹下存在 user-guide.md，文件名不是 user-guide.md.txt，且标签页没有未保存标记。', 'Убедитесь, что в учебной папке виден user-guide.md, а не user-guide.md.txt, и на вкладке нет индикатора несохранённых изменений.'),
        ],
        'result': ('学习文件夹中已有保存好的 user-guide.md，文件包含示例标题和正文。', 'В учебной папке сохранён user-guide.md с заголовком и текстом из примера.'),
    },
    {
        'id': 'B',
        'heading': ('任务 B 打开 Markdown 预览', 'Задача B Открыть предпросмотр Markdown'),
        'goal': ('在编辑区旁打开预览，并确认标题和正文正确显示。', 'Открыть предпросмотр рядом с редактором и проверить отображение заголовка и текста.'),
        'initial': ('任务 A 已完成；user-guide.md 已保存并在编辑区打开，文件中已有示例标题和正文。', 'Задача A выполнена. Файл user-guide.md сохранён и открыт в редакторе; он содержит заголовок и текст из примера.'),
        'steps': [
            ('单击 user-guide.md 的编辑区，使该文件成为当前活动文件。', 'Щёлкните в редакторе user-guide.md, чтобы сделать этот файл активным.'),
            ('按 macOS 的 Cmd+Shift+P、Windows/Linux 的 Ctrl+Shift+P，打开命令面板。', 'Откройте палитру команд: Cmd+Shift+P в macOS или Ctrl+Shift+P в Windows/Linux.'),
            ('输入并选择“Markdown: Open Preview to the Side”，按 Enter；编辑区旁应出现预览窗格。', 'Введите и выберите Markdown: Open Preview to the Side, затем нажмите Enter. Рядом с редактором появится предпросмотр.'),
            ('比较源文和预览：第一行应显示为较大的标题，不显示开头的“#”；俄文正文应作为普通段落显示。', 'Сравните исходный текст с предпросмотром: первая строка отображается крупным заголовком без начального #, а русский текст — обычным абзацем.'),
            ('在源文正文末尾添加“ 预览成功 / Предпросмотр работает.”并保存，确认预览也显示这段新增文字。', 'Добавьте в конец абзаца исходного текста « 预览成功 / Предпросмотр работает.» и сохраните файл. Убедитесь, что добавленный текст появился в предпросмотре.'),
        ],
        'result': ('源文与预览并排显示；标题、正文和新增文字均正确显示，修改后的文件已保存。', 'Исходный текст и предпросмотр видны рядом. Заголовок, абзац и добавленный текст отображаются правильно; изменённый файл сохранён.'),
    },
]
SOURCES = [
    ('VS Code 官方文档 Markdown 预览', 'Официальная документация VS Code о предпросмотре Markdown', 'https://code.visualstudio.com/docs/languages/markdown#_markdown-preview'),
    ('VS Code 官方文档 文件与文件夹', 'Официальная документация VS Code о файлах и папках', 'https://code.visualstudio.com/docs/editing/getting-started/editor-tutorial#_open-a-folder-in-vs-code'),
]

# Markdown and Word are generated from the same bilingual content.
lines = [f'# {TITLE[0]}', '', TITLE[1], '', INTRO[0], '', INTRO[1], '', NOTE[0], '', NOTE[1], '']
labels = [('goal', '目标 / Цель'), ('initial', '初始状态 / Начальное состояние')]
for task in TASKS:
    lines += [f'## {task["heading"][0]}', '', task['heading'][1], '']
    for key, label in labels:
        lines += [f'**{label}**', '', task[key][0], '', task[key][1], '']
    lines += ['### 操作步骤 / Порядок действий', '']
    for i, (zh, ru) in enumerate(task['steps'], 1):
        lines += [f'{i}. {zh}', f'   {ru}', '']
        if task['id'] == 'A' and i == 5:
            lines += ['   ```markdown'] + ['   ' + s if s else '   ' for s in SAMPLE.splitlines()] + ['   ```', '']
    lines += ['**结果 / Итог**', '', task['result'][0], '', task['result'][1], '']
lines += ['## 参考资料 / Источники', '', '核对日期：2026-10-20。 / Дата проверки: 20.10.2026.', '']
for zh, ru, url in SOURCES:
    lines += [f'- [{zh} / {ru}]({url})']
MD.write_text('\n'.join(lines) + '\n', encoding='utf-8')

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.65)
sec.left_margin = sec.right_margin = Inches(.7)
for sty in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2'):
    s = doc.styles[sty]
    s.font.name = 'Arial'
    s.font.color.rgb = RGBColor(0, 0, 0)
    rpr = s.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.insert(0, rf)
    rf.set(qn('w:eastAsia'), 'Arial Unicode MS')
    rf.set(qn('w:ascii'), 'Arial')
    rf.set(qn('w:hAnsi'), 'Arial')
normal = doc.styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = 1.02
normal.paragraph_format.space_after = Pt(3)
doc.styles['Title'].font.size = Pt(19)
doc.styles['Title'].paragraph_format.space_after = Pt(4)
doc.styles['Subtitle'].font.size = Pt(12)
doc.styles['Subtitle'].paragraph_format.space_after = Pt(8)
doc.styles['Heading 1'].font.size = Pt(13)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(9)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(3)
doc.styles['Heading 2'].font.size = Pt(11)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(6)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(3)

def p(text, *, bold=False, after=3, keep=False, indent=False, size=None):
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(after)
    par.paragraph_format.keep_with_next = keep
    if indent:
        par.paragraph_format.left_indent = Inches(.19)
    run = par.add_run(text)
    run.bold = bold
    if size:
        run.font.size = Pt(size)
    return par

def pair(value, *, label=None, step=None):
    prefix = f'{step}. ' if step else ''
    if label:
        zh_label, ru_label = label.split(' / ')
        par = p('', after=1, keep=True)
        par.add_run(zh_label + '：').bold = True
        par.add_run(value[0])
        par = p('', after=3.5)
        par.add_run(ru_label + ': ').bold = True
        par.add_run(value[1])
    else:
        p(prefix + value[0], after=1, keep=True)
        p(value[1], after=3.5, indent=bool(step))

doc.add_paragraph(TITLE[0], 'Title')
doc.add_paragraph(TITLE[1], 'Subtitle')
pair(INTRO)
for task in TASKS:
    doc.add_paragraph(task['heading'][0], 'Heading 1')
    p(task['heading'][1], bold=True, after=5, keep=True)
    for key, label in labels:
        pair(task[key], label=label)
    doc.add_paragraph('操作步骤 Порядок действий', 'Heading 2')
    for i, value in enumerate(task['steps'], 1):
        pair(value, step=i)
        if task['id'] == 'A' and i == 5:
            par = p(SAMPLE, after=5, size=10.5)
            par.paragraph_format.left_indent = Inches(.19)
            par.paragraph_format.keep_together = True
    pair(task['result'], label='结果 / Итог')

pair(NOTE)
doc.add_paragraph('参考资料 Источники', 'Heading 2')
p('核对日期 2026-10-20 / Дата проверки 20.10.2026', size=9, after=3)
for zh, ru, url in SOURCES:
    p(zh + ' / ' + ru, size=9, after=1, keep=True)
    p(url, size=9, after=3)
doc.core_properties.title = TITLE[0]
doc.core_properties.subject = 'Markdown 文件创建保存与预览 / Создание сохранение и предпросмотр Markdown'
doc.core_properties.author = ''
doc.core_properties.keywords = 'Markdown, VS Code, 中文, русский'
for border in list(doc.styles.element.xpath('.//w:pBdr')) + list(doc.element.xpath('.//w:pBdr')):
    border.getparent().remove(border)
for grid in list(doc.element.xpath('.//w:docGrid')):
    grid.getparent().remove(grid)
for par in doc.paragraphs:
    grid = OxmlElement('w:snapToGrid')
    grid.set(qn('w:val'), '0')
    par._p.get_or_add_pPr().append(grid)
doc.save(DOCX)

with ZipFile(DOCX) as z:
    assert z.testzip() is None
    xml = z.read('word/document.xml').decode()
all_text = '\n'.join(par.text for par in Document(DOCX).paragraphs)
checks = {
    'reader_is_beginner': INTRO[0] in all_text and INTRO[1] in all_text,
    'required_task_A': TASKS[0]['heading'][0] in all_text,
    'selected_task_B_is_preview': TASKS[1]['heading'][0] in all_text,
    'step_counts_5_to_8': all(5 <= len(t['steps']) <= 8 for t in TASKS),
    'all_bilingual_fields_in_word': all(v in all_text for t in TASKS for k in ('goal','initial','result') for v in t[k]),
    'all_bilingual_steps_in_word': all(v in all_text for t in TASKS for pair_ in t['steps'] for v in pair_),
    'all_bilingual_steps_in_markdown': all(v in MD.read_text() for t in TASKS for pair_ in t['steps'] for v in pair_),
    'sample_is_identical': SAMPLE in all_text and SAMPLE.splitlines()[0] in MD.read_text(),
    'required_filename_present': 'user-guide.md' in all_text,
    'docx_no_header_footer': not any(n.startswith(('word/header', 'word/footer')) for n in ZipFile(DOCX).namelist()),
    'no_placeholder_text': not bool(re.search(r'待填写|TODO|TBD|Lorem ipsum', all_text)),
}
assert all(checks.values()), checks
report = {
    'date': '2026-10-20',
    'source_image': '/var/folders/ff/kqw394c93dd7krl0z2hcqqdh0000gn/T/codex-clipboard-9de64f4f-c9b1-4eb5-9431-7f7cb82fb4ea.jpg',
    'source_sha256': hashlib.sha256(Path('/var/folders/ff/kqw394c93dd7krl0z2hcqqdh0000gn/T/codex-clipboard-9de64f4f-c9b1-4eb5-9431-7f7cb82fb4ea.jpg').read_bytes()).hexdigest(),
    'step_counts': {t['id']: len(t['steps']) for t in TASKS},
    'checks': checks,
    'all_pass': all(checks.values()),
    'visual_review': 'PENDING',
    'novice_usability_trial': 'NOT_RUN',
    'outputs': {f.name: {'sha256':hashlib.sha256(f.read_bytes()).hexdigest(), 'bytes':f.stat().st_size} for f in (MD,DOCX)},
}
(BASE / 'work' / 'verification.json').write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n')
print(json.dumps({'outputs':[str(MD),str(DOCX)],'checks':checks,'step_counts':report['step_counts']}, ensure_ascii=False, indent=2))
