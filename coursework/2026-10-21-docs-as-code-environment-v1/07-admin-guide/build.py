from pathlib import Path
import json
import hashlib
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT=Path(__file__).resolve().parent
PACK=OUT.parent
DOCX=PACK/'StudySpace_管理员指南_中俄双语_v1.docx'
for d in ['data','artifacts']: (OUT/d).mkdir(exist_ok=True)
blocks=[]
def h(n,zh,ru): blocks.append({'type':'heading','level':n,'zh':zh,'ru':ru})
def p(zh,ru): blocks.append({'type':'pair','zh':zh,'ru':ru})
def code(value,lang='bash'): blocks.append({'type':'code','text':value,'lang':lang})
def table(headers,rows,widths):blocks.append({'type':'table','headers':headers,'rows':rows,'widths':widths})

h(2,'1 概述与架构','1 Обзор и архитектура')
p('本指南供 StudySpace 课程小组在独立学习副本中安装、配置、诊断和演练恢复。应用版本 0.1.0，需求版本 0.3.0；正式文档 DRAFT，38 条需求仍 PROPOSED，38 个业务用例仍 NOT_RUN。',
  'Руководство для учебной группы StudySpace описывает установку, настройку, диагностику и восстановление отдельной учебной копии. Приложение 0.1.0, требования 0.3.0; документ DRAFT, 38 требований PROPOSED и 38 бизнес-кейсов NOT_RUN.')
p('架构为 React 浏览器界面、FastAPI 模块化单体及同步 SQLAlchemy。开发使用 SQLite。本地 Compose 为 HTTP 候选，不含 Caddy 或 PostgreSQL；目标部署选用 Caddy 与内网 PostgreSQL，仍未执行。依据 R01、R02、R03、R04、R09。',
  'Архитектура: React в браузере, монолит FastAPI и синхронный SQLAlchemy; разработка на SQLite. Локальный Compose — вариант HTTP без Caddy/PostgreSQL. Целевая схема с Caddy и внутренним PostgreSQL ещё не развёрнута. Основания R01, R02, R03, R04, R09.')
p('当前只提供首页、health、ready、Swagger 与 OpenAPI。认证、管理员创建、校内审核、资源、预约及统计业务未实现；没有默认管理员账号或密码。依据 NFR-003、ADR-0001 与 R06。',
  'Сейчас доступны главная страница, health, ready, Swagger и OpenAPI. Аутентификация, создание администратора, проверка вуза, ресурсы, бронь и статистика не реализованы; нет стандартного аккаунта или пароля администратора. NFR-003, ADR-0001, R06.')

h(2,'2 环境前提与访问权限','2 Среда и права доступа')
p('全部命令从以下学习包根目录执行。只操作 00-project 副本和本包内文件；本指南不要求修改原 StudySpace 项目或生产服务器。',
  'Все команды выполняются из корня учебного пакета ниже. Изменяются только 00-project и файлы внутри пакета; исходный проект и продакшен для этих процедур не нужны.')
code("cd '/Users/zhi/Documents/软件生命周期/studyspace/coursework/2026-10-21-docs-as-code-environment-v1'")
p('需要本地目录读写权限、Python ≥3.13 且 <3.15、uv、Node 22.x 的 22.13+ 或 ≥24、npm 与 Git。包根 .venv 已安装 46 个后端依赖，前端 node_modules 未打包。首次下载需要网络；版本与安装依据 R07、R08、R10。',
  'Нужны права чтения/записи, Python ≥3.13 и <3.15, uv, Node 22.x от 22.13 либо ≥24, npm и Git. В корневой .venv есть 46 пакетов сервера; node_modules не включён. Для первой загрузки нужна сеть. Основания R07, R08, R10.')
p('服务只绑定回环地址。检查 8000 和 5173 是否可用，不停止其他项目的进程。Docker 仅用于可选的配置检查；容器运行和目标 HTTPS 部署需另具备相应环境。依据 ADR-0004、NFR-003。',
  'Сервисы привязаны к loopback. Проверить 8000 и 5173, не останавливая чужие процессы. Docker нужен только для необязательной проверки конфигурации; контейнеры и HTTPS требуют отдельной среды. ADR-0004, NFR-003.')

h(2,'3 安装流程','3 Процедура установки')
p('AG-P01。状态 PARTIAL_VERIFIED：后端安装及包兼容性有 R10 证据；本次未重装前端或运行完整安装流程。依据 NFR-001、NFR-008、ADR-0001。',
  'AG-P01. PARTIAL_VERIFIED: установка и совместимость пакетов сервера подтверждены R10; полная установка и клиент в этой работе не выполнялись. NFR-001, NFR-008, ADR-0001.')
h(3,'3.1 前置条件','3.1 Предусловия')
p('先完成第 2 章；确认 00-project/backend/uv.lock 和 frontend/package-lock.json 存在，工作目录正确，网络可用或所需包已缓存。已有 .env 不覆盖。',
  'Выполнить главу 2; проверить backend/uv.lock и frontend/package-lock.json внутри 00-project, рабочий каталог и сеть либо кэш. Имеющийся .env не заменять.')
h(3,'3.2 操作步骤','3.2 Действия')
p('检查工具，按锁文件同步到学习包根 .venv，安装前端依赖，再仅在 .env 不存在时复制示例。所有缓存也放在本包内。',
  'Проверить инструменты, синхронизировать корневую .venv по lock-файлу, установить клиент и скопировать пример только при отсутствии .env. Кэш также хранится в пакете.')
code('python3 --version\nuv --version\nnode --version\nnpm --version\nUV_PROJECT_ENVIRONMENT="$PWD/.venv" UV_CACHE_DIR="$PWD/07-admin-guide/artifacts/uv-cache" \\\n  uv sync --project 00-project/backend --python 3.13 --frozen --no-python-downloads\nnpm_config_cache="$PWD/07-admin-guide/artifacts/npm-cache" \\\n  npm --prefix 00-project/frontend ci --no-audit --no-fund\nif [ ! -e 00-project/.env ]; then cp 00-project/.env.example 00-project/.env; fi\nUV_CACHE_DIR="$PWD/07-admin-guide/artifacts/uv-cache" uv pip check --python .venv/bin/python')
h(3,'3.3 预期结果','3.3 Ожидаемый результат')
p('各安装命令退出 0；.venv/bin/python、frontend/node_modules 存在，uv pip check 报包兼容，锁文件未变。保存版本及安装日志，不保存凭据。R10 记录原阶段后端安装通过；前端本次 NOT_RUN。',
  'Коды завершения 0; существуют .venv/bin/python и node_modules клиента, uv pip check подтверждает совместимость, lock-файлы сохранены. Записать версии и журнал без секретов. В R10 подтверждён прежний сервер; клиент сейчас NOT_RUN.')
h(3,'3.4 失败处理','3.4 При неуспехе')
p('Python 版本不符时选择已安装的 3.13；DNS 或下载失败时检查网络或缓存后重试相同锁定安装。不删除锁文件、不全局安装、不用自动升级依赖绕过故障；保留原 .env。',
  'При неверной версии выбрать установленный Python 3.13; при DNS/загрузке проверить сеть или кэш и повторить ту же установку. Не удалять lock-файлы, не ставить глобально и не обновлять зависимости для обхода ошибки; сохранить .env.')

h(2,'4 配置参数与敏感信息','4 Параметры и секреты')
p('至少配置以下 9 项。前四项是应用或代理环境变量，后两项控制安装工具，最后三项是显式 CLI 或配置文件参数；不把它们伪装为应用已支持的环境变量。依据 R05、R07、R08、R09。',
  'Ниже 9 параметров. Первые четыре — переменные приложения/прокси, следующие два — инструменты установки, последние три — CLI или настройки файла, а не вымышленные переменные приложения. Основания R05, R07, R08, R09.')
parameters=[
 ('APP_ENV','应用 / Приложение','development；允许 test、production / Допустимы test, production','00-project/.env / R05'),
 ('DATABASE_URL','应用 / Приложение','默认绝对 SQLite 路径；自定义目录须存在 / Абсолютный SQLite, каталог должен существовать','00-project/.env / R05'),
 ('CORS_ORIGINS','应用 / Приложение','JSON 列表，回环 5173；Compose 使用 8080 / JSON, loopback 5173; Compose 8080','00-project/.env / R05 R09'),
 ('API_PROXY_TARGET','Vite 代理 / Прокси','http://127.0.0.1:8000','00-project/.env / R07'),
 ('UV_PROJECT_ENVIRONMENT','安装工具 / Инструмент','学习包根 .venv / Корневая .venv','安装命令 / Команда установки'),
 ('UV_CACHE_DIR','安装工具 / Инструмент','本包 07-admin-guide/artifacts/uv-cache / Кэш в пакете','安装命令 / Команда установки'),
 ('Uvicorn --host / --port','后端 CLI / Сервер','127.0.0.1 / 8000','启动命令 / Команда запуска'),
 ('Vite server.host / server.port','前端配置 / Клиент','127.0.0.1 / 5173；strictPort=true','frontend/vite.config.ts / R07'),
 ('Compose frontend ports','容器配置 / Контейнер','127.0.0.1:8080:80；仅本地 HTTP / Локальный HTTP','compose.yaml / R09'),
]
table(['参数 / Параметр','类别 / Тип','值与含义 / Значение','位置与依据 / Где'],parameters,[4.1,2.8,7.2,3.7])
p('最小本地示例：保留应用默认 SQLite 路径。自定义 DATABASE_URL 时使用绝对路径，并先创建仅学习副本使用的父目录。APP_ENV=production 只是一项取值，不会自动补齐认证、HTTPS 或部署配置。',
  'Минимальный пример сохраняет стандартный SQLite. Для DATABASE_URL задать абсолютный путь и заранее создать каталог учебной копии. APP_ENV=production само по себе не включает авторизацию, HTTPS или готовый продакшен.')
code('APP_ENV=development\nCORS_ORIGINS=["http://127.0.0.1:5173","http://localhost:5173"]\nAPI_PROXY_TARGET=http://127.0.0.1:8000\n# Target-only placeholder; not active / 仅目标占位，未启用\n# DATABASE_URL=postgresql+psycopg://APP_USER:REPLACE_WITH_SECRET@db.invalid:5432/studyspace')
blocks[-1]['lang']='dotenv'
p('REPLACE_WITH_SECRET 是安全占位符，db.invalid 是示意主机，不是可用服务。实际密钥由受限环境配置提供；不写入指南、Git、日志或 VITE_* 客户端变量。当前开发骨架不需要生产密码。依据 NFR-003、NFR-007、ADR-0004。',
  'REPLACE_WITH_SECRET — заглушка, db.invalid — пример хоста, не сервис. Реальные секреты задаются ограниченной конфигурацией среды, без публикации в руководстве, Git, логах и VITE_*. Локальный каркас не требует производственного пароля. NFR-003, NFR-007, ADR-0004.')

h(2,'5 本地启动与停止流程','5 Локальный запуск и остановка')
p('AG-P02。状态 READY_NOT_RUN；以下命令适配包根 .venv，不直接调用期待 backend/.venv 的原 dev.sh。依据 R05、R06、R07、ADR-0001。',
  'AG-P02. READY_NOT_RUN; команды учитывают корневую .venv вместо dev.sh, ожидающего backend/.venv. Основания R05, R06, R07, ADR-0001.')
h(3,'5.1 前置条件','5.1 Предусловия')
p('AG-P01 完成；已核对 .env 不含外部数据库地址；8000 与 5173 空闲。打开两个本地终端，每个终端先进入第 2 章的学习包根目录。',
  'AG-P01 выполнен; в .env нет внешней БД, 8000 и 5173 свободны. Открыть два терминала и в обоих перейти в корень пакета из главы 2.')
h(3,'5.2 操作步骤','5.2 Действия')
p('终端 A 启动后端；终端 B 启动前端。打开 http://127.0.0.1:5173，随后执行第 6 章检查。正常停止时在两个终端各按 Ctrl+C。',
  'В терминале A запустить сервер, в B — клиент. Открыть http://127.0.0.1:5173 и выполнить главу 6. Для остановки нажать Ctrl+C в каждом терминале.')
code('# Terminal A / 终端 A\n.venv/bin/python -m uvicorn app.main:app --app-dir 00-project/backend \\\n  --host 127.0.0.1 --port 8000\n# Terminal B / 终端 B\nnpm --prefix 00-project/frontend run dev')
h(3,'5.3 预期结果','5.3 Ожидаемый результат')
p('后端监听回环 8000，Vite 监听回环 5173；首页加载且基础接口检查通过。Ctrl+C 后进程退出；本次未启动服务，预期结果不标为实测通过。',
  'Сервер слушает loopback 8000, Vite — 5173; главная и базовые API доступны. После Ctrl+C процессы завершаются. В этой работе сервисы не запускались, ожидание не считается фактом.')
h(3,'5.4 失败处理','5.4 При неуспехе')
p('缺依赖先返回 AG-P01；端口占用先确认占用者，不结束无关进程；就绪失败按 D03 处理。只停止本人启动的服务，保留数据库及日志，不用删除卷或删库作为通用修复。',
  'При нехватке зависимостей вернуться к AG-P01; при занятом порте установить владельца, не завершать чужие процессы; readiness — D03. Останавливать только свои сервисы, сохранять БД и журналы, не удалять тома или БД для общего исправления.')

h(2,'6 运行正常性检查流程','6 Проверка работоспособности')
p('AG-P03。状态 PARTIAL_VERIFIED：R11 的已有 TestClient 检查 8 passed、1 warning；真实 HTTP、页面与代理流程本次 NOT_RUN。依据 NFR-007、ADR-0001、R06。',
  'AG-P03. PARTIAL_VERIFIED: в R11 существующие тесты TestClient — 8 passed, 1 warning; реальные HTTP, страница и прокси сейчас NOT_RUN. NFR-007, ADR-0001, R06.')
h(3,'6.1 前置条件','6.1 Предусловия')
p('检查目标必须是学习服务；HTTP 命令要求 AG-P02 正在运行。没有启动服务时，改用下方已有基础测试，不把进程内结果标为网络部署通过。',
  'Проверять только учебные сервисы; HTTP-команды требуют запущенного AG-P02. Без сервиса использовать существующие тесты ниже, не приравнивая TestClient к сетевому развёртыванию.')
h(3,'6.2 操作步骤','6.2 Действия')
code('curl --fail --show-error http://127.0.0.1:8000/api/v1/health\ncurl --fail --show-error http://127.0.0.1:8000/api/v1/ready\ncurl --fail --show-error http://127.0.0.1:5173/api/v1/health\ncurl --fail --show-error http://127.0.0.1:5173/api/v1/ready\n# Isolated alternative / 不启动服务时的隔离检查\nPYTHONDONTWRITEBYTECODE=1 .venv/bin/python -m pytest \\\n  00-project/backend/tests/test_operations.py \\\n  00-project/backend/tests/test_config.py -q -p no:cacheprovider')
p('浏览器打开 /api/v1/docs；检查 /api/v1/openapi.json 仅列 health 与 ready。查看前端首页和语言切换，记录实际现象。HTTP 检查应同时覆盖后端和 Vite 代理地址。',
  'Открыть /api/v1/docs; проверить, что /api/v1/openapi.json содержит только health и ready. Посмотреть главную и переключение языков, записать факт. HTTP проверяется на сервере и через Vite.')
h(3,'6.3 预期结果','6.3 Ожидаемый результат')
p('health 返回 200，JSON 为 {"status":"ok","service":"StudySpace","version":"0.1.0"}。ready 返回 200 并增加 database=ok；仅证明 SELECT 1 连接，不证明业务表或预约可用。故障 ready 应为脱敏 503。',
  'health: 200 и JSON {"status":"ok","service":"StudySpace","version":"0.1.0"}. ready: 200 и database=ok; это только SELECT 1, не бизнес-таблицы или бронь. При сбое ready ожидается обезличенный 503.')
h(3,'6.4 失败处理','6.4 При неуспехе')
p('连接拒绝检查 AG-P02 与端口；health 成功但 ready 503 检查数据库目录及权限；代理失败按 D04。记录状态码及脱敏日志，处理后重复同一检查。测试中的 httpx 弃用提示是既有警告，不掩盖失败断言。',
  'При отказе соединения проверить AG-P02 и порты; health работает, ready 503 — каталог и права БД; прокси — D04. Сохранить код и безопасный журнал, затем повторить проверку. Прежнее предупреждение httpx не заменяет результаты тестов.')

h(2,'7 故障诊断场景','7 Диагностика неисправностей')
diagnostics=[
 ('D01', '安装失败或 Python 不符\nСбой установки или Python', '检查 Python 3.13、锁文件、DNS 与缓存\nПроверить Python 3.13, lock-файлы, DNS и кэш', '按 AG-P01 使用相同锁文件重试，不全局升级\nПовторить AG-P01 без глобального обновления', 'NFR-001 / R08'),
 ('D02', '端口已占用\nПорт занят', 'lsof -nP -iTCP:8000 -sTCP:LISTEN；5173 同理\nТа же проверка для 5173', '确认所有者；停止自己的旧实例或等待端口释放\nОпределить владельца, остановить только свой экземпляр', 'R07 / AG-P02'),
 ('D03', 'health 200 但 ready 503\nhealth 200, ready 503', '检查实际 SQLite 父目录和写权限；不输出含密钥的 URL\nКаталог и права SQLite, без URL с секретом', '修正学习副本路径和权限，重测 ready；不删库\nИсправить путь/права и повторить ready, БД не удалять', 'NFR-007 / R05 R06'),
 ('D04', '页面 API 不通或代理 502\nAPI недоступен или proxy 502', '先测 8000，再测 5173；核对 API_PROXY_TARGET\nСначала 8000, затем 5173, проверить API_PROXY_TARGET', '修正本地目标并重启自己启动的 Vite，再测代理\nИсправить target, перезапустить свой Vite и проверить', 'R07 / AG-P03'),
 ('D05', '跨域请求被浏览器拒绝\nБраузер отклоняет cross-origin', '核对准确 Origin 与 CORS JSON；普通同源代理不需跨域放行\nСверить Origin и JSON CORS; same-origin не требует CORS', '只加入所需本地来源并重启后端，避免无条件通配\nДобавить нужный локальный Origin и перезапустить сервер', 'NFR-003 / R05 R06'),
 ('D06', '恢复输出已存在或完整性失败\nФайл восстановления есть или повреждён', '检查新输出名称、源文件及演练日志\nПроверить новые имена, исходник и журнал', '不覆盖源或旧备份；保留失败候选，换新文件名或可靠备份重试\nНе заменять исходник/копии; сохранить сбойный файл и выбрать новую цель', 'NFR-005 NFR-007 / R13'),
]
table(['ID 与现象 / Симптом','检查 / Проверка','处理与复核 / Действия','依据 / Источник'],[(row[0]+' '+row[1],row[2],row[3],row[4]) for row in diagnostics],[4.2,5.2,6,2.4])
p('业务路径返回 404 时，先核对该功能是否实现。当前 OpenAPI 只有基础端点，不能靠改代理或填虚构管理员账号让预约业务可用。每次处理后回到原失败检查复核，并保留时间、命令、状态码及脱敏结果。',
  'При 404 бизнес-пути сначала проверить наличие реализации. OpenAPI сейчас содержит только базовые API; прокси и выдуманный администратор не включат бронь. После исправления повторить исходную проверку и сохранить время, команду, код и безопасный результат.')

h(2,'8 备份恢复与回退流程','8 Резервирование восстановление и возврат')
p('AG-P04。本地 SQLite 新文件恢复演练状态 PASS_DEMO_ONLY，依据 R12、R13；没有激活恢复文件，没有执行生产或业务数据恢复。依据 NFR-005、NFR-007、ADR-0002。',
  'AG-P04. Восстановление SQLite в новый файл: PASS_DEMO_ONLY, R12/R13. Кандидат не активирован, продакшен и бизнес-данные не восстанавливались. NFR-005, NFR-007, ADR-0002.')
h(3,'8.1 前置条件','8.1 Предусловия')
p('仅对本学习包内 SQLite 文件操作。停止本人启动的服务并确认无写入者；源数据库存在且完整，两个输出路径互异且尚不存在。先保存学习副本当前 DATABASE_URL 配置；源文件和原配置均保留。',
  'Только SQLite внутри учебного пакета. Остановить свои сервисы и убедиться в отсутствии записывающих процессов. Исходник существует и исправен, два разных выходных файла отсутствуют. Сохранить прежний DATABASE_URL учебной копии; исходник и настройки не заменяются.')
h(3,'8.2 操作步骤','8.2 Действия')
p('运行随包工具，以 SQLite backup API 生成一致性备份，再从备份恢复到新文件。下例源库是两条合成教学记录；重跑须选择尚不存在的输出名。实际学习数据恢复时，将 --source 改为已停写的学习库，仍保持新输出。',
  'Инструмент использует SQLite backup API, затем восстанавливает копию в новый файл. В примере две синтетические учебные записи; при повторе нужны новые имена. Для учебной БД заменить --source на остановленный исходник, сохраняя новые выходные файлы.')
code('.venv/bin/python 07-admin-guide/sqlite-backup-restore.py \\\n  --source 07-admin-guide/artifacts/recovery-demo/source-demo.sqlite \\\n  --backup 07-admin-guide/artifacts/recovery-demo/backup-new-01.sqlite \\\n  --restore 07-admin-guide/artifacts/recovery-demo/restored-new-01.sqlite')
p('工具在源、备份和恢复候选上执行 integrity_check，并确认源主文件哈希未变。检查需要的表与记录；确认仅为当前无业务表骨架或教学库后，可把学习副本 .env 的 DATABASE_URL 指向恢复候选的绝对路径，再按 AG-P02、AG-P03 重启并复核。此切换本次未执行。',
  'Инструмент выполняет integrity_check исходника, копии и кандидата, проверяя неизменность хеша исходного файла. Сверить нужные таблицы и записи. Только для текущего каркаса без бизнес-таблиц или учебной БД можно задать абсолютный путь кандидата в DATABASE_URL учебного .env, перезапустить AG-P02 и проверить AG-P03. Сейчас переключение не выполнялось.')
h(3,'8.3 预期结果','8.3 Ожидаемый результат')
p('输出 backup_integrity=ok、restored_integrity=ok、source_unchanged=true；原文件不变，备份及恢复候选独立存在。本次演练另核对两条合成记录完整一致，并验证重复输出和包外路径均被拒绝。activation 与 domain_acceptance 为 NOT_RUN。',
  'Ожидаются backup_integrity=ok, restored_integrity=ok, source_unchanged=true; исходник сохранён, копия и кандидат отдельны. В упражнении также совпали обе записи; повторные имена и пути вне пакета отклонены. activation и domain_acceptance — NOT_RUN.')
h(3,'8.4 失败处理与回退','8.4 При неуспехе и возврат')
p('完整性检查、数据核对或启动检查失败时，不激活候选。若已切换学习配置，先停服务，再还原此前保存的 DATABASE_URL，按 AG-P02、AG-P03 启动和复核。保留原库、备份和失败候选，不覆盖原数据，不用自动删库修复。',
  'При сбое целостности, сверки или запуска кандидат не активировать. Если учебная конфигурация уже переключена, остановить сервисы, вернуть сохранённый DATABASE_URL и повторить AG-P02/AG-P03. Сохранить исходник, копию и сбойный кандидат; без перезаписи или автоматического удаления БД.')
p('目标策略为每日备份保留最近 7 份；恢复后先运行到期数据清理再开放。现阶段未实现业务清理与迁移，因此本工具仅支持无真实学生数据的本地演练，不能据此开放恢复的真实业务数据。生产 PostgreSQL 的迁移、回滚、备份与恢复仍需独立实现和验证。',
  'Целевая политика: семь последних ежедневных копий и очистка истёкших данных перед открытием восстановленной системы. Бизнес-очистка и миграции ещё не реализованы, поэтому инструмент только для локальных упражнений без реальных данных студентов. PostgreSQL, миграция, откат и восстановление требуют отдельной реализации и проверки.')

h(2,'9 限制','9 Ограничения')
p('已验证范围：R10 的后端依赖安装、R11 的 8 项进程内基础测试，以及 R12 的两条合成记录恢复。Word 排版检查与这些运行证据分别记录。完整前端安装、服务启动、真实 HTTP/浏览器、容器运行及恢复后配置切换本次未执行。',
  'Проверены прежняя установка сервера R10, восемь базовых тестов TestClient R11 и восстановление двух синтетических записей R12. Проверка Word ведётся отдельно. Полная установка клиента, сервисы, реальные HTTP/браузер, контейнеры и переключение восстановленного файла сейчас не выполнялись.')
p('学校来源 OPEN-003、部署位置和域名 OPEN-010 仍待补；教师起始包与原始意见清单仍缺失。3.4.3 指南身份已由用户确认。健康或恢复演练通过不改变业务 TC 的 NOT_RUN，也不代表验收或生产可用。',
  'Источник вуза OPEN-003, сервер и домен OPEN-010 остаются открытыми; архив преподавателя и исходные замечания отсутствуют. Руководство 3.4.3 подтверждено пользователем. Проверка health или упражнение восстановления не меняют бизнес-TC NOT_RUN и не означают приёмку.')

h(2,'10 变更历史','10 История изменений')
table(['版本 / Версия','日期 / Дата','变更 / Изменение','状态 / Статус'],[
 ('0.1','2026-10-21','从既有架构形成 8 章指南，增加参数、诊断、恢复演练及追溯\n8 глав по архитектуре, параметры, диагностика, восстановление и трассировка','DRAFT')
],[2,2.5,10.8,2.5])
h(3,'10.1 依据与证据索引','10.1 Источники и свидетельства')
refs=[
 ('R01','NFR-001、003、005、007、008 需求登记 / Реестр требований','00-project/docs/data/requirements.json'),
 ('R02','ADR-0001 模块化单体 / Модульный монолит','00-project/docs/decisions/zh-CN/0001-simple-monolith.md'),
 ('R03','ADR-0002 数据库边界 / Границы БД','00-project/docs/decisions/zh-CN/0002-sqlite-postgresql.md'),
 ('R04','ADR-0004 目标部署 / Целевое размещение','00-project/docs/decisions/zh-CN/0004-docker-https.md'),
 ('R05','运行配置 / Конфигурация','00-project/backend/app/config.py'),
 ('R06','健康与就绪 / Health и readiness','00-project/backend/app/main.py'),
 ('R07','前端代理与端口 / Прокси и порты','00-project/frontend/vite.config.ts'),
 ('R08','Python 与锁定依赖 / Python и зависимости','00-project/backend/pyproject.toml'),
 ('R09','本地 Compose / Локальный Compose','00-project/compose.yaml'),
 ('R10','第 1 部分环境证据 / Среда части 1','artifacts/environment-check.json'),
 ('R11','第 2 部分基础测试 / Базовые тесты части 2','06-architecture-guide/artifacts/foundation-tests.log'),
 ('R12','本次恢复演练 / Текущее упражнение','07-admin-guide/artifacts/recovery-verification.json'),
 ('R13','随包恢复工具 / Инструмент восстановления','07-admin-guide/sqlite-backup-restore.py'),
]
for rid,label,path in refs:
    blocks.append({'type':'reference','id':rid,'label':label,'path':path})
p('ADR 与需求均有配对俄文版本：ADR 见 docs/decisions/ru-RU，需求见 docs/ru-RU/14_Requirements_Register.md；以上路径均相对于本学习包根目录的 00-project。NFR 描述表示目标与约束，不自动表示已实现。',
  'Парные русские ADR находятся в docs/decisions/ru-RU, реестр — docs/ru-RU/14_Requirements_Register.md внутри 00-project. NFR задают цели и ограничения, а не автоматически реализованные свойства.')

content={'doc_id':'SS-DAC-ADMIN-01','version':'0.1','status':'DRAFT','date':'2026-10-21','blocks':blocks,'procedures':['AG-P01','AG-P02','AG-P03','AG-P04'],'parameters':parameters,'diagnostics':diagnostics,'references':refs}
(OUT/'data/content.json').write_text(json.dumps(content,ensure_ascii=False,indent=2)+'\n')

doc=Document();s=doc.sections[0]
s.page_width,s.page_height=Cm(21),Cm(29.7)
s.left_margin,s.right_margin=Cm(1.6),Cm(1.6)
s.top_margin,s.bottom_margin=Cm(1.4),Cm(1.5);s.footer_distance=Cm(.65)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 st=doc.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
 st.font.size=Pt(10 if name=='Normal' else 17 if name=='Title' else 13 if name=='Heading 1' else 11)
 st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'PingFang SC')
 st.paragraph_format.space_after=Pt(4);st.paragraph_format.line_spacing=1.02
 if name.startswith('Heading'):st.paragraph_format.space_before=Pt(4)
footer=s.footer.paragraphs[0];footer.alignment=2
footer.add_run('SS-DAC-ADMIN-01  DRAFT  ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.add_paragraph('StudySpace 管理员指南','Title')
doc.add_paragraph('StudySpace Руководство администратора','Subtitle')
doc.add_paragraph('SS-DAC-ADMIN-01 · v0.1 · DRAFT · 2026-10-21 · Asia/Shanghai')
md=['# StudySpace 管理员指南 / Руководство администратора\n\nSS-DAC-ADMIN-01 · v0.1 · DRAFT · 2026-10-21 · Asia/Shanghai\n']
for b in blocks:
 if b['type']=='heading':
  num=int(b['zh'].split()[0].split('.')[0])
  if b['level']==2 and num in [3,4,5,6,7,8,9]:doc.add_page_break()
  doc.add_paragraph(b['zh'],'Heading 1' if b['level']==2 else 'Heading 2')
  doc.add_paragraph(b['ru'],'Heading 2' if b['level']==2 else 'Heading 3')
  md.append('#'*b['level']+' '+b['zh']+' / '+b['ru']+'\n')
 elif b['type']=='pair':
  first=doc.add_paragraph(b['zh']);first.paragraph_format.keep_with_next=True
  doc.add_paragraph(b['ru'])
  md.append(b['zh']+'\n\n'+b['ru']+'\n')
 elif b['type']=='code':
  for line in b['text'].splitlines():
   p0=doc.add_paragraph(line);p0.paragraph_format.space_after=Pt(1);p0.paragraph_format.line_spacing=1.0
   for run in p0.runs:run.font.name='Courier New';run.font.size=Pt(8.6)
  md.append('```'+b['lang']+'\n'+b['text']+'\n```\n')
 elif b['type']=='reference':
  doc.add_paragraph(b['id']+' '+b['label']+'\n'+b['path'])
  md.append(f"- [{b['id']} {b['label']}]({b['path']})")
 elif b['type']=='table':
  t=doc.add_table(rows=1,cols=len(b['headers']));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
  for i,w in enumerate(b['widths']):t.columns[i].width=Cm(w)
  for i,value in enumerate(b['headers']):t.rows[0].cells[i].text=value
  header=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(header)
  for values in b['rows']:
   cells=t.add_row().cells
   for i,value in enumerate(values):cells[i].text=value
  for ri,row in enumerate(t.rows):
   keep=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(keep)
   for i,cell in enumerate(row.cells):
    cell.width=Cm(b['widths'][i]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    props=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
    for edge in ['top','bottom','left','right']:
     e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    props.append(borders)
    padding=OxmlElement('w:tcMar')
    for edge in ['top','bottom','left','right']:
     e=OxmlElement('w:'+edge);e.set(qn('w:w'),'70');e.set(qn('w:type'),'dxa');padding.append(e)
    props.append(padding)
    if ri==0:
     fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'E7EDF3');props.append(fill)
    for p0 in cell.paragraphs:
     p0.paragraph_format.space_after=Pt(2);p0.paragraph_format.line_spacing=1.0
     for run in p0.runs:
      run.font.size=Pt(9);run.font.name='Arial';run.bold=(ri==0)
      run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'PingFang SC')
  md.append('| '+' | '.join(b['headers'])+' |\n| '+' | '.join('---' for _ in b['headers'])+' |')
  for row in b['rows']:md.append('| '+' | '.join(x.replace('|','\\|').replace('\n','<br>') for x in row)+' |')
  md.append('')
for element in [doc.styles.element,doc.element]:
 for border in list(element.iter(qn('w:pBdr'))):border.getparent().remove(border)
doc.core_properties.title='StudySpace 管理员指南'
doc.core_properties.author='StudySpace 课程小组'
doc.save(DOCX)
(PACK/'admin-guide.md').write_text('\n'.join(md)+'\n',encoding='utf-8')
source_records=[]
for rid,label,path in refs:
 p0=PACK/path;assert p0.exists(),p0
 source_records.append({'id':rid,'path':str(p0),'sha256':hashlib.sha256(p0.read_bytes()).hexdigest()})
for path in ['00-project/backend/uv.lock','00-project/frontend/package-lock.json','00-project/backend/app/database.py','00-project/frontend/package.json','00-project/.env.example','00-project/scripts/setup.sh','00-project/scripts/dev.sh']:
 p0=PACK/path;source_records.append({'id':'SUPPORT','path':str(p0),'sha256':hashlib.sha256(p0.read_bytes()).hexdigest()})
(OUT/'artifacts/source-manifest.json').write_text(json.dumps(source_records,ensure_ascii=False,indent=2)+'\n')
print(json.dumps({'markdown':str(PACK/'admin-guide.md'),'word':str(DOCX),'main_sections':8,'extra_sections':2,'procedures':4,'parameters':len(parameters),'diagnostics':len(diagnostics)},ensure_ascii=False))
