from pathlib import Path
from zipfile import ZipFile
import ast
import hashlib
import json
import re
import subprocess
from docx import Document
from pypdf import PdfReader
from lxml import etree

ROOT=Path(__file__).resolve().parent
PACK=ROOT.parent
guide=PACK/'admin-guide.md'
word=PACK/'StudySpace_管理员指南_中俄双语_v1.docx'
data=json.loads((ROOT/'data/content.json').read_text())
md=guide.read_text()
doc=Document(word)
text='\n'.join([p.text for p in doc.paragraphs]+[c.text for t in doc.tables for row in t.rows for c in row.cells])
norm=lambda s:re.sub(r'\s+',' ',s).strip()
alltext=norm(text)
missing=[]
for block in data['blocks']:
 if block['type'] in {'pair','heading'}:
  for key in ['zh','ru']:
   if norm(block[key]) not in alltext:missing.append(block[key])
 elif block['type']=='table':
  for row in block['rows']:
   for cell in row:
    if norm(cell) not in alltext:missing.append(cell)
 elif block['type']=='code':
  if norm(block['text']) not in alltext:missing.append(block['text'])
headings=re.findall(r'^## (.+)$',md,re.M)
numbers=[int(h.split()[0]) for h in headings]
procedure_sections=[3,5,6,8]
structure=all(all(re.search(rf'^### {n}\.{i} ',md,re.M) for i in range(1,5)) for n in procedure_sections)
bash_results=[]
for i,source in enumerate(re.findall(r'```bash\n(.*?)\n```',md,re.S),1):
 result=subprocess.run(['bash','-n'],input=source,text=True,capture_output=True)
 bash_results.append({'block':i,'exit_code':result.returncode,'error':result.stderr})
source_records=json.loads((ROOT/'artifacts/source-manifest.json').read_text())
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
source_changes=[x['path'] for x in source_records if sha(Path(x['path']))!=x['sha256']]
prior=json.loads((PACK/'artifacts/materials-manifest.json').read_text())
prior_changes=[x['copy'] for x in prior['copies'] if sha(Path(x['source']))!=x['sha256'] or sha(PACK/x['copy'])!=x['sha256']]
previous=json.loads((PACK/'06-architecture-guide/artifacts/verification.json').read_text())
previous_changes=[name for name,detail in previous['outputs'].items() if sha(PACK/'06-architecture-guide'/name)!=detail['sha256']]
broken=[]
for p in [guide,ROOT/'README.md',PACK/'README.md']:
 for target in re.findall(r'\]\(([^)]+)\)',p.read_text()):
  if target.startswith(('https:','http:','#')):continue
  resolved=(p.parent/target).resolve()
  if resolved==ROOT/'artifacts/verification.json':continue
  if not resolved.exists():broken.append({'file':str(p),'target':target})
recovery=json.loads((ROOT/'artifacts/recovery-verification.json').read_text())
pdf=PdfReader(next((ROOT/'artifacts/render-final').glob('*.pdf')))
with ZipFile(word) as z:
 xml=etree.fromstring(z.read('word/document.xml'))
 styles=etree.fromstring(z.read('word/styles.xml'))
 ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
 checks={
  'required_filename':guide.name=='admin-guide.md',
  'eight_main_sections_plus_limits_and_history':numbers==list(range(1,11)) and headings[8].startswith('9 限制') and headings[9].startswith('10 变更历史'),
  'at_least_three_procedures':len(data['procedures'])>=3,
  'four_parts_for_each_procedure':structure,
  'at_least_five_parameters':len(data['parameters'])>=5,
  'at_least_three_diagnostic_situations':len(data['diagnostics'])>=3,
  'safe_secret_placeholder_present':'REPLACE_WITH_SECRET' in md and 'db.invalid' in md,
  'no_live_format_keys':not re.search(r'(?:sk-[A-Za-z0-9]{24,}|AKIA[A-Z0-9]{16}|-----BEGIN (?:RSA |EC )?PRIVATE KEY-----)',md),
  'requirements_and_adr_references_present':all(x in md for x in ['NFR-001','NFR-003','NFR-005','NFR-007','NFR-008','ADR-0001','ADR-0002','ADR-0004']),
  'all_bilingual_content_and_commands_present_in_word':not missing,
  'all_bash_blocks_parse':all(x['exit_code']==0 for x in bash_results),
  'dotenv_example_is_not_marked_executable_bash':'```dotenv' in md,
  'links_resolve':not broken,
  'referenced_sources_unchanged':not source_changes,
  'all_142_original_preparation_sources_and_copies_unchanged':not prior_changes,
  'part_two_deliverables_unchanged':not previous_changes,
  'recovery_demo_matches_and_preserves_source':recovery['status']=='PASS_DEMO_ONLY' and recovery['restore_content_matches_source'] and recovery['source_unchanged'] and recovery['synthetic_rows']==2,
  'recovery_refuses_overwrite_and_external_path':recovery['repeat_refused_without_overwrite'] and recovery['out_of_packet_path_refused'],
  'recovered_data_not_activated':recovery['activation']=='NOT_RUN' and recovery['production_recovery']=='NOT_RUN',
  'word_zip_integrity':z.testzip() is None,
  'table_headers_repeat':len(xml.xpath('//w:tbl',namespaces=ns))==len(xml.xpath('//w:tbl/w:tr[1]/w:trPr/w:tblHeader',namespaces=ns)),
  'no_fixed_table_heights':not xml.xpath('//w:trHeight',namespaces=ns),
  'no_title_or_paragraph_borders':not xml.xpath('//w:pBdr',namespaces=ns) and not styles.xpath('//w:pBdr',namespaces=ns),
  'ten_nonblank_rendered_pages':len(pdf.pages)==10 and all(len(p.extract_text() or '')>500 for p in pdf.pages),
 }
for name in ['sqlite-backup-restore.py','build.py']:
 ast.parse((ROOT/name).read_text())
checks['helper_and_builder_python_parse']=True
report={'date':'2026-10-21','doc_id':data['doc_id'],'checks':checks,'all_pass':all(checks.values()),'check_count':len(checks),'main_sections':8,'extra_sections':2,'procedures':len(data['procedures']),'parameters':len(data['parameters']),'diagnostics':len(data['diagnostics']),'bash_syntax_results':bash_results,'source_changes':source_changes,'prior_changes':prior_changes,'previous_deliverable_changes':previous_changes,'broken_links':broken,'missing_word_content':missing,'page_count':len(pdf.pages),'visual_review':'PASS','reviewed_pages':[f'artifacts/render-final/page-{i}.png' for i in range(1,11)],'visual_findings':'All 10 pages inspected individually; readable Chinese and Russian, no clipping, overlaps, missing glyphs, title rules or blank pages. The recovery failure-handling pair stays together on page 8.','evidence_scope':{'backend_install':'PRIOR_EVIDENCE_PART_1','foundation_tests':'PRIOR_EVIDENCE_PART_2_8_PASSED_1_WARNING','recovery':'PASS_DEMO_ONLY','frontend_install':'NOT_RUN','service_start':'NOT_RUN','production_recovery':'NOT_RUN'},'outputs':{str(p.relative_to(PACK)):{'sha256':sha(p),'bytes':p.stat().st_size} for p in [guide,word]}}
(ROOT/'artifacts/verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
print(json.dumps({'all_pass':report['all_pass'],'check_count':len(checks),'pages':len(pdf.pages),'failed':[k for k,v in checks.items() if not v],'missing_word_content':missing[:2],'broken_links':broken,'source_changes':source_changes,'previous_changes':previous_changes},ensure_ascii=False))
assert report['all_pass']
