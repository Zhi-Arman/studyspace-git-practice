"""Create the bilingual task-five report from real build and browser records."""
from pathlib import Path
import hashlib
import json
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent
ART=ROOT/'artifacts'
initial=json.loads((ART/'initial-build.json').read_text())
final=json.loads((ART/'final-build.json').read_text())
browser=json.loads((ART/'browser-check.json').read_text())
checks=browser['checks']
assert initial['exit_code']==final['exit_code']==0
assert final['warning_bytes']==0 and not final['changed_document_sources']
assert all(c['passed'] for c in checks)
byid={c['id']:c for c in checks}
blocks=[]
def h(zh,ru):blocks.append({'type':'heading','zh':zh,'ru':ru})
def p(zh,ru):blocks.append({'type':'pair','zh':zh,'ru':ru})
def table(headers,rows,widths):blocks.append({'type':'table','headers':headers,'rows':rows,'widths':widths})
def code(text):blocks.append({'type':'code','text':text})

h('1 检查结论与任务范围','1 Результат и область проверки')
p('图中第 5 部分已完成。首次及最终 Sphinx 严格构建均退出 0、警告 0；从目录逐一点击四个正文页面，完成 8 项链接或跳转检查，包含图片及标题。命令、表格、警告经过浏览器指标检查及截图查看。13 项浏览器检查通过。',
  'Часть 5 выполнена. Первичная и итоговая строгие сборки Sphinx завершились с кодом 0 без предупреждений. Из оглавления открыты четыре страницы; выполнены 8 проверок ссылок или переходов, включая изображение и заголовок. Команды, таблицы и предупреждения проверены по метрикам браузера и снимкам. Пройдены 13 проверок браузера.')
p('检查对象为前一部分 project-docs/site 的本地文档站。执行者为 Codex，2026-10-21，Asia/Shanghai，macOS arm64、Python 3.13.12、Sphinx 8.2.3、MyST 4.0.1、gstack 自带 Chromium；桌面视口 1280×900。应用命令没有在本轮执行。',
  'Объект — локальная документация project-docs/site из предыдущей части. Исполнитель Codex, 21.10.2026, Asia/Shanghai, macOS arm64, Python 3.13.12, Sphinx 8.2.3, MyST 4.0.1 и Chromium gstack; viewport 1280×900. Команды приложения сейчас не выполнялись.')
table(['图片要求 / Требование','执行与证据 / Проверка и свидетельство'],[
 ['使用 -W 构建并记录结果和退出码\nСборка с -W, результат и код','初测与复测均保存 JSON、完整日志及警告文件\nJSON, журналы и предупреждения обеих сборок'],
 ['从目录打开每页\nОткрыть каждую страницу из оглавления','L01–L04：实际点击四页\nРеальные переходы на четыре страницы'],
 ['至少 5 处，包含图片和标题\nМинимум 5, включая изображение и заголовок','L01–L08：8 项；L05 标题，L06–07 图片\n8 проверок; L05 заголовок, L06–07 изображение'],
 ['检查命令 表格 警告可读性\nЧитаемость команд таблиц предупреждений','R01–R03：指标和截图\nМетрики и просмотр снимков'],
 ['记录错误 修复 复测\nОшибки исправления повторный результат','1 项检查工具超时已修复；网站未发现新缺陷\nИсправлен тайм-аут инструмента, новых дефектов сайта нет'],
 ['构建成功不等于命令验证成功\nСборка не равна проверке команд приложения','应用操作与业务用例保持 NOT_RUN\nОперации и бизнес-кейсы остаются NOT_RUN']
],[6.4,11.2])

h('2 构建结果与页面检查','2 Сборка и страницы')
p('以下命令在前一部分 project-docs 根目录执行，使用其独立虚拟环境。警告日志分别写入本部分 artifacts/initial-warnings.log 和 final-warnings.log；完整绝对命令及时间见对应 build.json。',
  'Команда выполнялась из project-docs с его отдельной средой. Предупреждения сохранены в artifacts/initial-warnings.log и final-warnings.log этой части; полные команды и время — в build.json.')
code('.venv/bin/python -m sphinx -b html -n -W --keep-going -E -a -N '+chr(92)+'\n  -w ../../09-site-build-check/artifacts/final-warnings.log docs site')
p('命令以相对日志路径表示；实际执行 JSON 保存了对应的完整绝对路径。下方是实际结果。没有屏蔽警告，使用 -E -a 重读并重写全部页面。',
  'Путь журнала показан относительно каталога; JSON содержит полный фактический путь. Ниже фактические результаты. Предупреждения не подавлялись, -E -a перечитывали и переписывали все страницы.')
table(['执行 / Запуск','时间 Asia/Shanghai / Время','退出码 / Код','结果 / Результат'],[
 ['初测 / Первичный',initial['started_at'],'0','build succeeded.\n警告 0 / Предупреждений 0'],
 ['复测 / Повторный',final['started_at'],'0','build succeeded.\n警告 0 / Предупреждений 0']
],[3.0,6.0,2.2,6.4])
table(['ID','目录页面 / Страница','实际结果 / Фактический результат'],[
 ['L01','user-guide.html\nVS Code 用户指南 / Руководство VS Code','从首页目录点击到达，双语标题正确\nПереход из оглавления, верный заголовок'],
 ['L02','admin-guide.html\n管理员手册 / Администратор','从首页目录点击到达，九章及历史可见\nПереход из оглавления, девять глав и история'],
 ['L03','release-notes.html\n文档发布说明 / Примечания','从首页目录点击到达，版本说明可见\nПереход из оглавления, версия видна'],
 ['L04','document-register.html\n文档登记 / Реестр','从首页目录点击到达，登记表可见\nПереход из оглавления, реестр виден']
],[1.3,7.0,9.3])

h('3 链接图片与标题检查','3 Ссылки изображение и заголовок')
p('本节记录真实点击或资源打开结果。L01–L04 已覆盖全部目录页；另检查标题锚点、图片加载、图片直接打开，以及搜索结果点击。',
  'Здесь фактические клики и переходы к ресурсам. L01–L04 охватывают все страницы оглавления; дополнительно проверены якорь, загрузка и открытие изображения, поиск с переходом к результату.')
rows=[]
for cid in ['L01','L02','L03','L04','L05','L06','L07','L08']:
 c=byid[cid]
 results={
 'L01':'PASS：目标 URL 和标题正确\nВерны URL и заголовок',
 'L02':'PASS：目标 URL 和标题正确\nВерны URL и заголовок',
 'L03':'PASS：目标 URL 和标题正确\nВерны URL и заголовок',
 'L04':'PASS：目标 URL 和标题正确\nВерны URL и заголовок',
 'L05':'PASS：片段标识改变，7 更新与回滚位于视口顶部\nИзменён hash, заголовок 7 вверху viewport',
 'L06':'PASS：complete=true，1000×160，有替代文本\nИзображение загружено, alt присутствует',
 'L07':'PASS：图片 URL 返回 200，SVG 包含 SQLite\nHTTP 200, SVG содержит SQLite',
 'L08':'PASS：搜索 SQLite 后实际点击返回管理员手册\nПоиск SQLite и клик ведут к администратору'}
 rows.append([cid,c['zh']+'\n'+c['ru'],results[cid]])
table(['ID','检查 / Проверка','结果 / Результат'],rows,[1.3,8.0,8.3])
p('每项的目标、落地 URL、浏览器返回及截图路径保存在 artifacts/browser-check.json。图片源是本站 _images/development-chain.svg；标题为管理员手册第 7 章。没有访问站外个人账户。',
  'Цели, URL результата, вывод браузера и снимки в artifacts/browser-check.json. Изображение — _images/development-chain.svg, заголовок — глава 7 администратора. Сторонние аккаунты не использовались.')

h('4 可读性与错误复测','4 Читаемость ошибки и повторная проверка')
cmds=byid['R01']['result'];tables=byid['R02']['result'];warnings=byid['R03']['result']
table(['ID','对象 / Объект','实际观察 / Наблюдение'],[
 ['R01',f'{len(cmds)} 个命令块 / Блоков команд','字体不小于 13px；换行后无横向截断；截图查看通过\nШрифт ≥13px, без обрезки, снимок просмотрен'],
 ['R02',f'{len(tables)} 张表格 / Таблиц','字体不小于 13px，内容未超出表格，表头与单元格可辨\nШрифт ≥13px, нет переполнения, заголовки и ячейки различимы'],
 ['R03',f'{len(warnings)} 处中文警告及对应俄文 / Предупреждений с переводом','安装、恢复、更新前警告均可读；字体不小于 15px\nПредупреждения установки, восстановления, обновления читаются, ≥15px'],
 ['C01','控制台 / Консоль','清空此前缓冲后，本地流程无控制台错误\nПосле очистки буфера нет локальных ошибок']
],[1.3,5.7,10.6])
table(['编号与分类 / ID и тип','错误与处理 / Ошибка и исправление','复测 / Повторный результат'],[
 ['CHK-01\n检查工具 / Инструмент','独立 SVG 已 HTTP 200，但 gstack snapshot -i 超时，错误为 The operation timed out.。保留初测记录；直接 SVG 改用 DOM、HTTP 状态及截图核对，HTML 页面继续快照；为命令添加 30 秒上限。\nSVG вернул HTTP 200, но snapshot -i завершился тайм-аутом. Сохранён первичный журнал; SVG проверяется через DOM, HTTP и снимок, HTML через snapshot. Командам задан лимит 30 секунд.','PASS：L06、L07 及整轮 13 项通过；最终构建退出 0。\nL06/L07 и все 13 проверок пройдены, сборка 0.'],
 ['SITE-NEW\n网站缺陷 / Дефекты сайта','本轮未发现需修改文档站源码的新缺陷。\nНовых дефектов исходников сайта не обнаружено.','前后文档源文件哈希未变；初测与复测均警告 0。\nХеши исходников сохранены, обе сборки без предупреждений.']
],[3.0,10.0,4.6])

h('5 证据与未执行范围','5 Свидетельства и невыполненная область')
p('本轮网站未新增缺陷，检查工具问题已排除。最终另运行既有交付检查：88 项通过，250 处本地链接、资源、锚点无断链，22 份原材料哈希一致。该静态检查与本轮 8 项浏览器链接及资源检查分别记录。',
  'Новых дефектов сайта нет, проблема инструмента устранена. Итоговая прежняя проверка поставки: 88 проверок, 250 локальных ссылок/ресурсов/якорей без разрывов, хеши 22 материалов совпали. Статические проверки записаны отдельно от 8 проверок ссылок и ресурсов в браузере.')
table(['工作 / Работа','状态 / Статус','边界 / Граница'],[
 ['Sphinx 构建与目录页面\nСборка Sphinx и страницы','PASS','仅静态文档站\nТолько документация'],
 ['命令 表格 警告可读性\nЧитаемость','PASS','查看显示，不执行应用命令\nПросмотр, без выполнения команд приложения'],
 ['安装 启停 备份 恢复 更新 回滚\nУстановка запуск копирование восстановление обновление откат','NOT_RUN','本轮未运行应用操作\nОперации приложения сейчас не выполнялись'],
 ['38 个业务用例与团队验收\n38 бизнес-кейсов и приёмка','NOT_RUN / PENDING','文档成功不改变业务状态\nДокументация не меняет статус бизнеса']
],[8.0,3.2,6.4])
p('核心证据均在本目录 artifacts：initial-build.json、initial-build.log、initial-warnings.log；final-build.json、final-build.log、final-warnings.log；browser-check-initial.json（含超时）、browser-check.json（复测）；delivery-verification.json；source-before.json。',
  'Основные свидетельства в artifacts: initial-build.json/log и initial-warnings.log; final-build.json/log и final-warnings.log; browser-check-initial.json с тайм-аутом, browser-check.json с повторной проверкой; delivery-verification.json; source-before.json.')
p('截图 00–11 分别覆盖首页、四正文页、标题、图片、搜索、命令、表格和警告。此报告是 Codex 文档检查记录，未冒充团队成员签字或教师验收。教师原始起始包及意见清单仍缺失。',
  'Снимки 00–11 охватывают главную, четыре страницы, заголовок, изображение, поиск, команды, таблицы и предупреждения. Это проверка документации Codex, не подпись команды или приёмка преподавателя. Исходный архив и замечания пока отсутствуют.')

md=['# StudySpace 网站与构建检查\n\nПроверка сайта и сборки StudySpace\n',
 'SS-DAC-CHECK-05 · v1.0 · DRAFT · 2026-10-21 · Asia/Shanghai\n']
for b in blocks:
 if b['type']=='heading':md.append('## '+b['zh']+' / '+b['ru']+'\n')
 elif b['type']=='pair':md.append(b['zh']+'\n\n'+b['ru']+'\n')
 elif b['type']=='code':md.append('```bash\n'+b['text']+'\n```\n')
 else:
  md.append('| '+' | '.join(b['headers'])+' |\n| '+' | '.join('---' for _ in b['headers'])+' |')
  md.extend('| '+' | '.join(x.replace('|','\\|').replace('\n','<br>') for x in row)+' |' for row in b['rows'])
  md.append('')
name='StudySpace_任务五网站与构建检查_中俄双语_v1'
(ROOT/(name+'.md')).write_text('\n'.join(md)+'\n')
(ART/'report-content.json').write_text(json.dumps(blocks,ensure_ascii=False,indent=2)+'\n')

doc=Document();sec=doc.sections[0]
sec.page_width,sec.page_height=Cm(21),Cm(29.7)
sec.left_margin=sec.right_margin=Cm(1.7);sec.top_margin=sec.bottom_margin=Cm(1.5)
for name0,size in [('Normal',10),('Title',18),('Subtitle',13),('Heading 1',13),('Heading 2',11)]:
 s=doc.styles[name0];s.font.name='Arial';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
 s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Hiragino Sans GB')
 s.paragraph_format.space_after=Pt(4);s.paragraph_format.line_spacing=1.04
 if name0.startswith('Heading'):s.paragraph_format.space_before=Pt(7)
doc.add_paragraph('StudySpace 网站与构建检查','Title')
doc.add_paragraph('Проверка сайта и сборки StudySpace','Subtitle')
doc.add_paragraph('SS-DAC-CHECK-05 · v1.0 · DRAFT · 2026-10-21 · Asia/Shanghai')
for b in blocks:
 if b['type']=='heading':
  doc.add_paragraph(b['zh'],'Heading 1');doc.add_paragraph(b['ru'],'Heading 2')
 elif b['type']=='pair':
  para=doc.add_paragraph(b['zh']);para.paragraph_format.keep_with_next=True
  doc.add_paragraph(b['ru'])
 elif b['type']=='code':
  for line in b['text'].splitlines():
   para=doc.add_paragraph(line);para.paragraph_format.space_after=Pt(1)
   for r in para.runs:r.font.name='Courier New';r.font.size=Pt(8)
 else:
  t=doc.add_table(rows=1,cols=len(b['headers']));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
  for i,w in enumerate(b['widths']):t.columns[i].width=Cm(w)
  for i,v in enumerate(b['headers']):t.rows[0].cells[i].text=v
  t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
  for row in b['rows']:
   cells=t.add_row().cells
   for i,v in enumerate(row):cells[i].text=v
  for ri,row in enumerate(t.rows):
   row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
   for i,c in enumerate(row.cells):
    c.width=Cm(b['widths'][i]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    pr=c._tc.get_or_add_tcPr();edges=OxmlElement('w:tcBorders')
    for edge in ['top','bottom','left','right']:
     el=OxmlElement('w:'+edge)
     for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]:el.set(qn('w:'+k),v)
     edges.append(el)
    pr.append(edges);pad=OxmlElement('w:tcMar')
    for edge in ['top','bottom','left','right']:
     el=OxmlElement('w:'+edge);el.set(qn('w:w'),'65');el.set(qn('w:type'),'dxa');pad.append(el)
    pr.append(pad)
    if ri==0:
     el=OxmlElement('w:shd');el.set(qn('w:fill'),'E5EDF3');pr.append(el)
    for para in c.paragraphs:
     para.paragraph_format.space_after=Pt(2);para.paragraph_format.line_spacing=1.0
     for r in para.runs:
      r.font.name='Arial';r.font.size=Pt(9);r.bold=ri==0
      r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Hiragino Sans GB')
  doc.add_paragraph()
footer=sec.footer.paragraphs[0];footer.alignment=2;footer.add_run('SS-DAC-CHECK-05  DRAFT  ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
for tree in [doc.styles.element,doc.element]:
 for el in list(tree.iter(qn('w:pBdr'))):el.getparent().remove(el)
doc.core_properties.title='StudySpace 网站与构建检查'
doc.core_properties.author='StudySpace 课程小组'
doc.save(ROOT/(name+'.docx'))
print(json.dumps({'markdown':str(ROOT/(name+'.md')),'word':str(ROOT/(name+'.docx')),'browser_checks':len(checks)},ensure_ascii=False))
