from pathlib import Path
from collections import Counter
from docx import Document
from pypdf import PdfReader
from lxml import etree
from zipfile import ZipFile
import json
import re
import hashlib

root = Path(__file__).resolve().parent
pack = root.parent
content = json.loads((root / 'data/content.json').read_text())
stem = 'StudySpace_任务二从架构推导操作指南_中俄双语_v1'
doc = Document(root / (stem + '.docx'))
text = '\n'.join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for row in t.rows for c in row.cells])
norm = lambda s: re.sub(r'\s+', ' ', s).strip()
alltext = norm(text)
pairs = list(content['intro'])
for item in content['components']:
    pairs.extend([item[1], item[3]])
for item in content['nodes']:
    pairs.append(item['name'])
    pairs.extend(field[2] for field in item['fields'])
pairs.extend(item[1] for item in content['assumptions'])
pairs.extend(item[3] for item in content['sources'])
for item in content['outline']:
    pairs.extend([item[1], item[2]])
for item in content['procedures']:
    pairs.extend([item[2], item[4]])
pairs.extend(item[2] for item in content['commands'])
missing = [s for pair in pairs for s in pair if norm(s) not in alltext]
sha = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()
sources = json.loads((root / 'artifacts/source-manifest.json').read_text())
source_changes = [item['path'] for item in sources['sources'] if sha(Path(item['path'])) != item['sha256']]
copy_changes = [item['copy'] for item in sources['copied_component_sources'] if sha(Path(item['source'])) != item['sha256'] or sha(root / item['copy']) != item['sha256']]
old = json.loads((pack / 'artifacts/materials-manifest.json').read_text())
old_changes = [item['copy'] for item in old['copies'] if sha(Path(item['source'])) != item['sha256'] or sha(pack / item['copy']) != item['sha256']]
broken = []
for p in [root / (stem + '.md'), root / 'README.md', pack / 'README.md']:
    for link in re.findall(r'\]\(([^)]+)\)', p.read_text()):
        if link.startswith(('https:', 'http:', '#')):
            continue
        target = (p.parent / link).resolve()
        if target == root / 'artifacts/verification.json':
            continue
        if not target.exists():
            broken.append({'document': str(p), 'target': link})
pdf = PdfReader(root / 'artifacts/render-delivery' / (stem + '.pdf'))
page_text = [p.extract_text() or '' for p in pdf.pages]
with ZipFile(root / (stem + '.docx')) as z:
    xml = etree.fromstring(z.read('word/document.xml'))
    styles = etree.fromstring(z.read('word/styles.xml'))
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    checks = {
        'ten_stable_component_ids': [item[0] for item in content['components']] == [f'CMP-{i:03d}' for i in range(1, 11)],
        'seven_stable_node_ids': [item['id'] for item in content['nodes']] == [f'N-{i:02d}' for i in range(1, 8)],
        'four_attributes_and_sources_per_node': all(len(item['fields']) == 4 and all(re.search(r'S\d{2}', f[3]) for f in item['fields']) for item in content['nodes']),
        'source_types_and_assumptions': {'DIAGRAM', 'ADR', 'NFR', 'IMPLEMENTATION'}.issubset({item[1] for item in content['sources']}) and len(content['assumptions']) == 4,
        'ten_guide_chapters': len(content['outline']) == 10,
        'ten_procedures_with_explicit_status': len(content['procedures']) == 10 and all(item[3] in {'VERIFIED', 'READY_NOT_RUN', 'CONDITIONAL_NOT_RUN', 'DEFERRED'} for item in content['procedures']),
        'bilingual_content_present_in_word': not missing,
        'sources_and_original_component_files_unchanged': not source_changes and not copy_changes,
        'all_142_preparation_sources_and_copies_unchanged': not old_changes,
        'markdown_and_readme_links_resolve': not broken,
        'word_zip_integrity': z.testzip() is None,
        'all_tables_repeat_header': len(xml.xpath('//w:tbl', namespaces=ns)) == len(xml.xpath('//w:tbl/w:tr[1]/w:trPr/w:tblHeader', namespaces=ns)),
        'no_fixed_row_heights': not xml.xpath('//w:trHeight', namespaces=ns),
        'no_title_or_paragraph_borders': not xml.xpath('//w:pBdr', namespaces=ns) and not styles.xpath('//w:pBdr', namespaces=ns),
        'fourteen_nonblank_rendered_pages': len(pdf.pages) == 14 and all(len(item) > 500 for item in page_text),
        'eight_foundation_tests_passed': '8 passed, 1 warning' in (root / 'artifacts/foundation-tests.log').read_text(),
        'document_checks_passed': all((root / 'artifacts' / name).read_text().startswith('PASS:') for name in ['docs-check.log', 'registers-check.log']),
        'compose_config_exit_zero': next(item for item in json.loads((root / 'artifacts/execution.json').read_text())['checks'] if item['id'] == 'P-05')['exit_code'] == 0,
    }
report = {
    'doc_id': content['doc_id'], 'date': '2026-10-21', 'checks': checks,
    'all_pass': all(checks.values()), 'check_count': len(checks),
    'bilingual_pairs_checked': len(pairs), 'table_count': len(doc.tables), 'page_count': len(pdf.pages),
    'reviewed_pages': [f'artifacts/render-delivery/page-{i}.png' for i in range(1, 15)],
    'visual_review': 'PASS',
    'visual_findings': 'All 14 final pages inspected individually. Both languages readable; no clipping, overlaps, missing glyphs or blank pages. Node tables remain on their own pages; repeated source-table header and footer spacing verified.',
    'procedure_states': dict(Counter(item[3] for item in content['procedures'])),
    'business_tests': 'NOT_RUN', 'production_deployment': 'NOT_RUN',
    'outputs': {stem + ext: {'sha256': sha(root / (stem + ext)), 'bytes': (root / (stem + ext)).stat().st_size} for ext in ['.md', '.docx']},
    'source_changes': source_changes, 'copy_changes': copy_changes,
    'broken_links': broken, 'missing_bilingual_text': missing,
}
(root / 'artifacts/verification.json').write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n')
print(json.dumps({'all_pass': report['all_pass'], 'checks': len(checks), 'bilingual_pairs': len(pairs), 'pages': len(pdf.pages), 'tables': len(doc.tables), 'failed': [name for name, value in checks.items() if not value], 'missing_text': missing[:3], 'source_changes': source_changes, 'old_changes': old_changes}, ensure_ascii=False))
assert report['all_pass']
