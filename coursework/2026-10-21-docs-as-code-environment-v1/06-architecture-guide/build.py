from pathlib import Path
from datetime import datetime
from zoneinfo import ZoneInfo
import hashlib
import json
import shutil
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).resolve().parent
PACK = OUT.parent
ROOT = PACK.parents[1]
STEM = 'StudySpace_任务二从架构推导操作指南_中俄双语_v1'
for d in ['data', 'artifacts', 'sources']:
    (OUT / d).mkdir(exist_ok=True)

def pair(zh, ru): return (zh, ru)
def digest(p): return hashlib.sha256(p.read_bytes()).hexdigest()

# Retain the prior classroom component sources inside this same learning folder.
component_root = ROOT.parent / 'StudySpace_组件交付'
source_records = []
for filename in ['StudySpace_组件候选与组件图_中俄双语_v2.md', 'StudySpace_组件图_课堂任务2_中俄双语.drawio', 'StudySpace_组件图_课堂任务2_中俄双语.png']:
    src = component_root / filename
    dst = OUT / 'sources' / filename
    if not dst.exists(): shutil.copy2(src, dst)
    assert digest(src) == digest(dst)
    source_records.append({'source': str(src), 'copy': str(dst.relative_to(OUT)), 'sha256': digest(src)})

sources = [
 ('S01','DIAGRAM', 'sources/StudySpace_组件图_课堂任务2_中俄双语.drawio', pair('既有组件图及数字草图，CMP 和 IF 标识来源','Прежняя диаграмма и черновик, источник ID CMP и IF')),
 ('S02','DIAGRAM', 'sources/StudySpace_组件候选与组件图_中俄双语_v2.md', pair('统一组件表 SS-ARCH-03 v2，10 个组件的名称、职责与边界','Единая таблица SS-ARCH-03 v2, имена, функции и границы 10 компонентов')),
 ('S03','DIAGRAM', '../02-nodes-environments/StudySpace_节点与制品表_中俄双语_v1.md', pair('N-01 至 N-07 节点、制品、端口及环境；配套环境表同目录','Узлы N-01–N-07, артефакты, порты и среды; таблица сред в том же каталоге')),
 ('S04','DIAGRAM', '../01-deployment/部署图说明.md', pair('本次学习包整理的部署图 DRAFT；不是既有原图或已部署拓扑','Учебный DRAFT развёртывания, не прежний оригинал и не фактическая топология')),
 ('S05','ADR', '../00-project/docs/decisions/README.md', pair('ADR-0001 模块化单体；0002 开发 SQLite、目标 PostgreSQL；0003 座位 MVP；0004 Compose 与 Caddy','ADR-0001 монолит, 0002 SQLite и PostgreSQL, 0003 MVP мест, 0004 Compose и Caddy')),
 ('S06','NFR', '../00-project/docs/data/requirements.json', pair('NFR-001 可维护，003 安全，004 事务，005 迁移，006 界面，007 诊断，008 追溯；均 PROPOSED','NFR-001 сопровождение, 003 безопасность, 004 транзакции, 005 перенос, 006 интерфейс, 007 диагностика, 008 трассировка; PROPOSED')),
 ('S07','IMPLEMENTATION', '../00-project/compose.yaml', pair('本地 HTTP 容器候选、回环端口、SQLite 卷、健康检查','Локальный HTTP Compose, loopback-порты, том SQLite, healthcheck')),
 ('S08','IMPLEMENTATION', '../00-project/frontend/nginx.conf', pair('/healthz、静态目录、/api/ → backend:8000；frontend/Dockerfile 为配套构建来源','/healthz, статика, /api/ → backend:8000; сборка в frontend/Dockerfile')),
 ('S09','IMPLEMENTATION', '../00-project/scripts/dev.sh', pair('本地服务端口、启动前占用检查和进程停止；配套 frontend/vite.config.ts 及 setup.sh','Порты, проверка занятости и остановка; также frontend/vite.config.ts и setup.sh')),
 ('S10','IMPLEMENTATION', '../00-project/backend/app/config.py', pair('APP_ENV、DATABASE_URL、CORS_ORIGINS 与默认绝对 SQLite 路径；配套 database.py','Настройки APP_ENV, DATABASE_URL, CORS_ORIGINS и абсолютный SQLite; также database.py')),
 ('S11','IMPLEMENTATION', '../00-project/backend/app/main.py', pair('health、ready、Swagger、OpenAPI；就绪只执行 SELECT 1','health, ready, Swagger, OpenAPI; readiness выполняет только SELECT 1')),
 ('S12','IMPLEMENTATION', '../00-project/backend/pyproject.toml', pair('Python ≥3.13 且 <3.15、包依赖及 postgres 可选项；实际版本由 uv.lock 固定','Python ≥3.13 и <3.15, зависимости и postgres extra; версии в uv.lock')),
 ('S13','IMPLEMENTATION', '../00-project/docs/data/project-decisions.json', pair('DEC 决策输入；OPEN-003 核验来源、OPEN-010 位置和域名仍 PARTIAL；配套系统设计 DOC-04','Решения DEC, OPEN-003 и OPEN-010 остаются PARTIAL; также проект DOC-04')),
 ('S14','EVIDENCE', '../artifacts/environment-check.json', pair('第 1 部分实测环境与 46 个包；指南身份另见 guide-3.4.3-confirmation.json','Среда части 1 и 46 пакетов; подтверждение руководства в guide-3.4.3-confirmation.json')),
 ('S15','EVIDENCE', 'artifacts/foundation-tests.log', pair('本次已有基础测试 8 passed、1 warning；仅 TestClient 与临时 SQLite','Текущие существующие тесты 8 passed, 1 warning; TestClient и временный SQLite')),
 ('S16','EVIDENCE', 'artifacts/compose-config.log', pair('本次 docker compose config --quiet 退出码 0；未构建或启动容器','docker compose config --quiet, код 0; без сборки и запуска')),
]

components = [
 ('CMP-001',pair('学生与管理员界面','Интерфейс студента и администратора'),'N-01 / N-02 / N-03',pair('基础界面已存在，业务界面待实现','Основа интерфейса есть, бизнес-экраны впереди')),
 ('CMP-002',pair('账号认证与权限控制','Аутентификация и права доступа'),'N-04',pair('设计，待实现','Проект, не реализован')),
 ('CMP-003',pair('校内身份审核与预约资格','Проверка принадлежности к вузу и допуск'),'N-04',pair('设计，依赖 N-07','Проект, зависит от N-07')),
 ('CMP-004',pair('房间座位与开放规则管理','Помещения места и правила работы'),'N-04',pair('设计，待实现','Проект, не реализован')),
 ('CMP-005',pair('位置筛选与座位可用性查询','Поиск помещений и свободных мест'),'N-04',pair('设计，待实现','Проект, не реализован')),
 ('CMP-006',pair('预约事务与冲突控制','Транзакции бронирования и конфликты'),'N-04',pair('设计，待实现','Проект, не реализован')),
 ('CMP-007',pair('预约记录与基础统计','Записи бронирований и показатели'),'N-04',pair('设计，待实现','Проект, не реализован')),
 ('CMP-008',pair('审计与数据生命周期','Аудит и жизненный цикл данных'),'N-04',pair('设计，待实现','Проект, не реализован')),
 ('CMP-009',pair('数据访问与持久化','Доступ к данным и хранение'),'N-04 / N-06',pair('SQLAlchemy 基础已存在，业务模型和迁移待实现','Основа SQLAlchemy есть, модели и миграции впереди')),
 ('CMP-010',pair('可信学校核验来源','Доверенный источник вуза'),'N-07',pair('外部依赖，OPEN-003','Внешняя зависимость, OPEN-003')),
]

# Each of the four required attributes carries its own provenance.
nodes = [
 {'id':'N-01','name':pair('学生或管理员终端','Устройство студента или администратора'),'fields':[
  ('软件','ПО',pair('浏览器运行 CMP-001；React 制品由前端服务分发。浏览器兼容与 375px 布局是 NFR-006 目标，尚未完成验收。','Браузер выполняет CMP-001, получая React от клиента. Совместимость и 375px — цель NFR-006 без приёмки.'),'S02 S03 S06 NFR-006'),
  ('访问','Доступ',pair('开发 HTTP 127.0.0.1:5173；容器候选 HTTP 127.0.0.1:8080；目标为域名 HTTPS 443。业务登录和管理员创建仍待实现。','Разработка HTTP 127.0.0.1:5173, контейнеры HTTP 127.0.0.1:8080, цель — домен HTTPS 443. Бизнес-вход и создание администратора впереди.'),'S03 S07 S09 S05 ADR-0004'),
  ('参数','Параметры',pair('基础界面语言 zh-CN、ru-RU、en-US；/api 同源代理；终端不直连数据库。目标域名不填虚构值。','Языки основы zh-CN, ru-RU, en-US; /api через тот же origin; без прямого доступа к БД. Целевой домен не выдуман.'),'S02 S09 S13 NFR-006 A-01'),
  ('检查','Проверка',pair('启动前端后查看首页、语言切换和 API 状态；代理 health/ready 应返回基础契约。当前可安排验证，本次浏览器流程 NOT_RUN。','После запуска проверить страницу, языки и API; health/ready через прокси должны соответствовать контракту. Проверка возможна, браузерный сценарий NOT_RUN.'),'S09 S11 NFR-006 P-06'),
 ]},
 {'id':'N-02','name':pair('本地开发工作站','Локальная рабочая станция'),'fields':[
  ('软件','ПО',pair('Python ≥3.13 且 <3.15、uv；Node 22.13+ 的 22.x 或 ≥24、npm、Git。学习 .venv 已安装后端依赖；前端依赖不在快照中。','Python ≥3.13 и <3.15, uv; Node 22.x от 22.13 либо ≥24, npm, Git. В учебной .venv есть зависимости сервера; node_modules в снимок не входит.'),'S09 S12 S14 NFR-001'),
  ('访问','Доступ',pair('本机终端与学习目录读写权限；常规开发服务只绑定 127.0.0.1:5173/8000。当前验证用学习 .venv 与临时测试数据。','Терминал и права на учебный каталог; dev-сервисы на 127.0.0.1:5173/8000. Текущая проверка использует учебную .venv и временные данные.'),'S09 S10 S15'),
  ('参数','Параметры',pair('APP_ENV 默认 development；DATABASE_URL 默认绝对 SQLite 路径；CORS_ORIGINS 为 JSON 列表；API_PROXY_TARGET 默认 http://127.0.0.1:8000。','APP_ENV по умолчанию development; DATABASE_URL — абсолютный SQLite; CORS_ORIGINS — JSON; API_PROXY_TARGET — http://127.0.0.1:8000.'),'S09 S10'),
  ('检查','Проверка',pair('检查版本、锁依赖、文档一致性；已有配置与健康测试 8/8 通过。学习环境在包根 .venv，而 dev.sh 期待 backend/.venv 和 node_modules，不能直接套用。','Проверить версии, зависимости и документы; тесты настроек и health 8/8. Учебная среда в корне пакета .venv; dev.sh ожидает backend/.venv и node_modules.'),'S09 S14 S15 P-01 P-02 P-03 P-04'),
 ]},
 {'id':'N-03','name':pair('前端静态服务','Сервис статического клиента'),'fields':[
  ('软件','ПО',pair('前端构建镜像 node:22-alpine；运行镜像 nginx:stable-alpine。此为 Dockerfile 配置，未实测镜像内部版本。','Сборка node:22-alpine, исполнение nginx:stable-alpine. Это конфигурация Dockerfile, версии внутри образов не проверены.'),'S03 S08'),
  ('访问','Доступ',pair('本地候选 127.0.0.1:8080 → 容器 80；容器 /api/ → backend:8000；目标入口经 N-05。','Локально 127.0.0.1:8080 → 80 контейнера; /api/ → backend:8000; целевой вход через N-05.'),'S07 S08 S05 ADR-0004'),
  ('参数','Параметры',pair('静态目录 /usr/share/nginx/html；/assets 缓存 immutable，首页 no-cache；上游连接超时 5s、读取超时 15s。','Статика /usr/share/nginx/html; /assets immutable, главная no-cache; таймаут соединения 5s, чтения 15s.'),'S08'),
  ('检查','Проверка',pair('配置静态核查已通过；容器运行后 GET /healthz 期望 200 ok，再测 /api/v1/health 与 ready。需 Docker 运行环境和镜像构建，运行验证 NOT_RUN。','Статическая конфигурация проверена; после запуска GET /healthz ожидается 200 ok, затем health и ready через /api. Нужны Docker и сборка; запуск NOT_RUN.'),'S08 S11 S16 P-05 P-07'),
 ]},
 {'id':'N-04','name':pair('后端应用服务','Сервис приложения'),'fields':[
  ('软件','ПО',pair('Python 3.13 镜像、uv、Uvicorn、FastAPI、Pydantic、同步 SQLAlchemy。CMP-002 至 CMP-009 共用单体进程；业务模块待实现。','Образ Python 3.13, uv, Uvicorn, FastAPI, Pydantic, синхронный SQLAlchemy. CMP-002–CMP-009 в одном монолите; бизнес-модули впереди.'),'S02 S03 S05 ADR-0001 S12'),
  ('访问','Доступ',pair('本地绑定 127.0.0.1:8000；容器内监听 0.0.0.0:8000，由 Nginx 转发。镜像使用 studyspace 非 root 用户。基础接口不提供业务认证证明。','Снаружи локально 127.0.0.1:8000, внутри 0.0.0.0:8000, доступ через Nginx. Пользователь образа studyspace, не root. Базовые API не доказывают авторизацию.'),'S07 S11 backend/Dockerfile NFR-003'),
  ('参数','Параметры',pair('Compose 中 APP_ENV=development；SQLite /app/backend/data/studyspace.db 与 studyspace-data 卷；CORS 允许本地 8080。目标生产配置另行准备。','В Compose APP_ENV=development; SQLite /app/backend/data/studyspace.db и том studyspace-data; CORS для локального 8080. Продакшен-конфигурация впереди.'),'S07 S10 S05 ADR-0004'),
  ('检查','Проверка',pair('health 200 仅证明进程接口；ready 200 database=ok 仅证明 SELECT 1，故障返回脱敏 503；Swagger 与 OpenAPI 可核对。本次进程内测试通过，真实容器 HTTP NOT_RUN。','health 200 проверяет liveness; ready 200 database=ok — только SELECT 1, сбой даёт обезличенный 503; проверить Swagger/OpenAPI. TestClient пройден, HTTP контейнера NOT_RUN.'),'S11 S15 NFR-007 P-04 P-07'),
 ]},
 {'id':'N-05','name':pair('目标 HTTPS 入口','Целевой вход HTTPS'),'fields':[
  ('软件','ПО',pair('目标选用 Caddy；当前 Compose 没有此服务，也没有完成的生产 Caddy 配置。','Выбран Caddy; в текущем Compose сервиса и готовой продакшен-конфигурации нет.'),'S03 S05 ADR-0004 S07'),
  ('访问','Доступ',pair('目标公网入口 80/443；HTTPS 443 在 Caddy 终止，随后内部 HTTP 80 到 Nginx。域名及服务器由 OPEN-010 补充。','Целевой вход 80/443; TLS на 443 завершается в Caddy, далее HTTP 80 к Nginx. Домен и сервер — OPEN-010.'),'S04 S05 ADR-0004 S13 A-01'),
  ('参数','Параметры',pair('待提供域名、DNS、上游地址、可写持久证书存储和端口规则；具体值未配置，不生成假证书、账号或公网地址。','Нужны домен, DNS, upstream, постоянное хранилище сертификатов и правила портов; значения не заданы, сертификаты и адрес не выдумываются.'),'S05 ADR-0004 S13 A-01'),
  ('检查','Проверка',pair('前提满足后检查证书信任、域名、有效期、重定向和 /api 转发；不能只凭 443 可连接判定合格。当前 DEFERRED，NOT_RUN。','После выполнения условий проверить доверие сертификату, домен, срок, редирект и API; одного соединения 443 недостаточно. Сейчас DEFERRED, NOT_RUN.'),'S05 ADR-0004 NFR-003 P-08'),
 ]},
 {'id':'N-06','name':pair('目标数据库服务','Целевой сервис БД'),'fields':[
  ('软件','ПО',pair('目标 PostgreSQL，具体服务端版本未确定；Python postgres 可选项预留 psycopg。业务表、Alembic 和数据库迁移尚未实现。','Цель PostgreSQL, версия сервера не задана; postgres extra предусматривает psycopg. Бизнес-таблицы, Alembic и миграции не реализованы.'),'S05 ADR-0002 S12 NFR-005 A-02'),
  ('访问','Доступ',pair('仅后端经内部网络访问；5432 为目标连接示例，不向公网发布。账号按权限最小化设计，真实凭据未提供。','Доступ только с сервера по внутренней сети; 5432 — целевой пример, без публикации в Интернет. Минимальные права, реальные реквизиты отсутствуют.'),'S03 S04 S05 ADR-0004 NFR-003 A-02'),
  ('参数','Параметры',pair('需确认 DATABASE_URL 的主机、端口、库名和密钥来源、持久卷与备份位置；目标每日备份保留 7 份，恢复后先运行到期清理。','Нужно определить хост, порт, БД и источник секрета DATABASE_URL, том и резервные копии; цель — 7 ежедневных копий и очистка после восстановления.'),'S10 S13 NFR-005 NFR-007 A-02'),
  ('检查','Проверка',pair('后续分别验证受限连接、迁移与回滚、备份恢复及并发事务。SELECT 1 不能代替这些测试。当前实现与环境前提不足，DEFERRED，NOT_RUN。','Позже отдельно проверить соединение, миграцию и откат, восстановление и конкурентные транзакции. SELECT 1 их не заменяет. Сейчас DEFERRED, NOT_RUN.'),'S05 ADR-0002 S06 NFR-004 NFR-005 P-09'),
 ]},
 {'id':'N-07','name':pair('可信学校核验来源','Доверенный источник вуза'),'fields':[
  ('软件','ПО',pair('学校确认的名单或平台；承载软件和版本由学校提供，OPEN-003 尚为 PARTIAL，不假定 SSO 或自动 API。','Подтверждённый вузом реестр или платформа; ПО и версии сообщает вуз. OPEN-003 PARTIAL, SSO/API не предполагаются.'),'S02 S03 S13 A-03'),
  ('访问','Доступ',pair('管理员通过学校授权渠道人工核验；与 CMP-003 的 IF-04 为人工依赖，网络协议与端口未定义。无可信来源不得批准 VERIFIED。','Администратор проверяет вручную по разрешённому вузом каналу; IF-04 к CMP-003 — ручная зависимость, протокол и порт не заданы. Без источника VERIFIED не разрешается.'),'S01 S02 S13 NFR-003 A-03'),
  ('参数','Параметры',pair('确认来源负责人、适用范围、更新周期和可复核的来源类型引用；具体值待学校提供。日志不复制学号姓名或证明影像。','Уточнить владельца, область, обновление и проверяемую ссылку типа источника; значения сообщает вуз. Логи без копий номера, имени и изображений документов.'),'S13 S06 NFR-003 NFR-007 A-03'),
  ('检查','Проверка',pair('可先静态核对来源需求和隐私规则；真实来源有效性、人工审核流程和服务端权限待来源及实现就绪后验证。当前 DEFERRED，业务 TC 保持 NOT_RUN。','Сейчас возможна сверка требований к источнику и приватности; реальная проверка источника, процесса и прав требует источника и реализации. DEFERRED, бизнес-TC NOT_RUN.'),'S02 S06 S13 NFR-003 P-10'),
 ]},
]

outline = [
 ('G-01',pair('适用范围与架构','Область и архитектура'),pair('读者、当前骨架、CMP 和 N 清单、开发与目标部署区别','Аудитория, каркас, CMP и N, различия разработки и цели'),'S01 S02 S03'),
 ('G-02',pair('本地环境准备','Подготовка среды'),pair('Python、Node、uv、锁文件、独立学习目录与路径','Python, Node, uv, lock-файлы, учебный каталог и пути'),'N-02 S12 S14'),
 ('G-03',pair('参数与访问权限','Параметры и права'),pair('APP_ENV、DATABASE_URL、CORS、代理；凭据管理与最小权限','APP_ENV, DATABASE_URL, CORS, прокси; секреты и минимальные права'),'N-02 N-04 S10 NFR-003'),
 ('G-04',pair('本地启动与停止','Локальный запуск и остановка'),pair('安装前提、端口占用、127.0.0.1、停止自己启动的进程','Зависимости, занятость портов, 127.0.0.1, остановка своих процессов'),'N-01 N-02 S09'),
 ('G-05',pair('诊断与故障定位','Диагностика сбоев'),pair('health、ready、503、CORS、代理与脱敏日志；不以删库处理故障','health, ready, 503, CORS, прокси и безопасные логи; без удаления БД'),'N-03 N-04 S11 NFR-007'),
 ('G-06',pair('本地容器运行','Локальные контейнеры'),pair('Compose 配置、构建、健康检查、端口、持久卷与停止','Compose, сборка, healthcheck, порты, тома и остановка'),'N-03 N-04 S07 S08'),
 ('G-07',pair('目标 HTTPS 部署','Целевой HTTPS'),pair('Caddy、域名、证书、边界；前提未满足时保留待执行','Caddy, домен, сертификаты и границы; до готовности отложено'),'N-05 ADR-0004 A-01'),
 ('G-08',pair('目标数据库迁移与恢复','Миграция и восстановление'),pair('PostgreSQL、迁移回滚、备份恢复、并发与生命周期','PostgreSQL, миграция, откат, копии, восстановление и конкуренция'),'N-06 ADR-0002 NFR-004 NFR-005'),
 ('G-09',pair('学校来源与审核前提','Источник вуза и проверка'),pair('可信渠道、人工核验、角色权限；业务实现后补操作步骤','Канал, ручная проверка, права; шаги после реализации'),'N-07 OPEN-003 NFR-003'),
 ('G-10',pair('验证记录与维护','Проверки и сопровождение'),pair('命令、预期、实际结果、版本、来源、未执行项和中俄同步','Команды, ожидания, результаты, версии, источники и синхронизация'),'S06 NFR-001 NFR-008 S14 S15 S16'),
]

procedures = [
 ('P-01','N-02',pair('版本与锁定依赖','Версии и lock-зависимости'),'VERIFIED',pair('第 1 部分 46 包安装与兼容性通过；本次作为既有证据引用 S14','В части 1 установлены 46 пакетов, совместимость пройдена; ссылка на прежнее S14')),
 ('P-02','N-02',pair('双语文档结构','Структура документов'),'VERIFIED',pair('本次运行 check_docs.py，通过；只证明结构','Сейчас check_docs.py пройден; только структура')),
 ('P-03','N-02',pair('登记表同步','Синхронизация реестров'),'VERIFIED',pair('本次 generate_registers.py --check 通过','Сейчас generate_registers.py --check пройден')),
 ('P-04','N-02 / N-04',pair('配置与基础 API 诊断','Настройки и диагностика API'),'VERIFIED',pair('本次已有测试 8 passed、1 warning，TestClient 与隔离 SQLite，无网络监听','Сейчас 8 passed, 1 warning; TestClient, отдельный SQLite, без сетевого сервера')),
 ('P-05','N-03 / N-04',pair('Compose 静态配置','Статическая конфигурация Compose'),'VERIFIED',pair('本次 config --quiet 退出 0；不需要容器启动','Сейчас config --quiet код 0; без запуска контейнеров')),
 ('P-06','N-01 / N-02',pair('页面和开发代理','Страница и dev-прокси'),'READY_NOT_RUN',pair('安装学习副本前端依赖并选定虚拟环境启动方式后可测；本次未启动','После установки клиента и выбора запуска учебной среды; сейчас не запускалось')),
 ('P-07','N-03 / N-04',pair('容器健康与 API 转发','Контейнеры и прокси API'),'CONDITIONAL_NOT_RUN',pair('需要 Docker 守护进程、镜像及端口；本次没有构建或启动','Нужны Docker daemon, образы и порты; сборка и запуск не выполнялись')),
 ('P-08','N-05',pair('HTTPS 和证书','HTTPS и сертификаты'),'DEFERRED',pair('OPEN-010、生产配置、域名与证书前提未满足，NOT_RUN','OPEN-010, конфигурация, домен и сертификат не готовы; NOT_RUN')),
 ('P-09','N-06',pair('迁移并发和恢复','Миграция конкуренция восстановление'),'DEFERRED',pair('PostgreSQL、业务模型与迁移未就绪，NOT_RUN','PostgreSQL, модели и миграции не готовы; NOT_RUN')),
 ('P-10','N-07',pair('真实校内核验','Реальная проверка вуза'),'DEFERRED',pair('OPEN-003 与审核实现未就绪，NOT_RUN','OPEN-003 и реализация проверки не готовы; NOT_RUN')),
]

intro = [
 pair('本成果面向 StudySpace 课程小组，完成实践 3.4.4 第 2 部分：把既有架构转化为操作指南的依据、节点操作要求和章节计划。复用既有组件及节点编号，3.4.3 指南对应关系已由用户确认。','Материал для учебной группы StudySpace выполняет часть 2 практики 3.4.4: превращает архитектуру в основания, требования к операциям узлов и план руководства. Прежние ID сохранены; связь руководства с 3.4.3 подтверждена пользователем.'),
 pair('结论：已整理 10 个组件、7 个节点；每个节点均列出软件、访问、参数、检查及来源。当前可验证基础配置和诊断；HTTPS、PostgreSQL 与学校来源保留实施前提。','Итог: 10 компонентов и 7 узлов; для каждого заданы ПО, доступ, параметры, проверка и источники. Сейчас проверяются настройки и диагностика; HTTPS, PostgreSQL и источник вуза требуют предпосылок.'),
 pair('当前本地 Compose 使用 SQLite，不含 Caddy 或 PostgreSQL。CMP-002 至 CMP-009 属于同一模块化单体；CMP-009 数据访问代码在 N-04，目标数据库数据存放 N-06。N-02 直接运行与 N-03/N-04 容器运行是可选方式。','Текущий локальный Compose использует SQLite без Caddy/PostgreSQL. CMP-002–CMP-009 образуют один монолит; код доступа CMP-009 расположен в N-04, целевые данные — N-06. Прямой запуск N-02 и контейнеры N-03/N-04 являются вариантами.'),
 pair('本次在学习副本运行静态检查及已有基础测试，未访问生产服务器。正式文档 DRAFT，需求 PROPOSED，38 个业务用例仍 NOT_RUN。','В учебной копии выполнены статические проверки и существующие базовые тесты, без продакшена. Документы DRAFT, требования PROPOSED, 38 бизнес-кейсов NOT_RUN.'),
]
assumptions = [
 ('A-01',pair('N-05 域名、服务器、证书目录及生产上游待 OPEN-010；仅保留字段，未赋予虚构值。','N-05: домен, сервер, сертификаты и upstream зависят от OPEN-010; поля без вымышленных значений.')),
 ('A-02',pair('N-06 数据库版本、账号、网络绑定和备份存储尚未配置；5432 只作为已有目标示例。','N-06: версия БД, аккаунт, привязки и копии не настроены; 5432 — прежний целевой пример.')),
 ('A-03',pair('N-07 平台、授权渠道、负责人及更新周期待学校确认；管理员人工核验方向已定，不假定自动 API。','N-07: платформу, канал, владельца и обновление определяет вуз; ручная проверка выбрана, API не предполагается.')),
 ('A-04',pair('P-06 的启动步骤需按学习副本路径适配。学习包 .venv 已存在，00-project/backend/.venv 与 frontend/node_modules 未随快照复制。','P-06 требует адаптации путей: .venv есть в корне пакета, но backend/.venv и frontend/node_modules в снимок не копировались.')),
]

commands = [
 ('P-02','PYTHONDONTWRITEBYTECODE=1 .venv/bin/python 00-project/scripts/check_docs.py',pair('预期与本次实际：PASS 双语元数据、章节、ID 与追溯引用。','Ожидание и факт: PASS метаданных, разделов, ID и ссылок.')),
 ('P-03','PYTHONDONTWRITEBYTECODE=1 .venv/bin/python 00-project/scripts/generate_registers.py --check',pair('预期与本次实际：PASS 生成登记表与 JSON 一致。','Ожидание и факт: PASS соответствия реестров JSON.')),
 ('P-04','PYTHONDONTWRITEBYTECODE=1 .venv/bin/python -m pytest 00-project/backend/tests/test_operations.py 00-project/backend/tests/test_config.py -q -p no:cacheprovider',pair('本次另指定 artifacts/pytest-temp 为隔离目录。实际 8 passed、1 warning；health 不建库，ready 仅 SELECT 1，故障 503 脱敏，另检查 CORS、OpenAPI 和配置。','Для текущего запуска задан отдельный artifacts/pytest-temp. Факт: 8 passed, 1 warning; health без БД, ready SELECT 1, обезличенный 503, CORS, OpenAPI и настройки.')),
 ('P-05','docker compose -f 00-project/compose.yaml config --quiet',pair('预期与本次实际：退出码 0；无输出是此命令正常行为。没有容器构建或启动。','Ожидание и факт: код 0; отсутствие вывода нормально. Сборка и запуск контейнеров не выполнялись.')),
]

content = {'doc_id':'SS-DAC-T2-01','version':'1.0','status':'DRAFT','date':'2026-10-21','intro':intro,'components':components,'nodes':nodes,'sources':sources,'assumptions':assumptions,'outline':outline,'procedures':procedures,'commands':commands,'source_copies':source_records}
(OUT/'data/content.json').write_text(json.dumps(content,ensure_ascii=False,indent=2)+'\n')

# One content model drives both editable formats.
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin, section.bottom_margin = Cm(1.5), Cm(1.5)
section.left_margin, section.right_margin = Cm(1.6), Cm(1.6)
section.footer_distance = Cm(0.65)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    style=doc.styles[name]
    style.font.name='Arial'
    style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'PingFang SC')
    style.font.color.rgb=RGBColor(0,0,0)
    style.font.size=Pt(10 if name=='Normal' else 18 if name=='Title' else 14 if name=='Heading 1' else 11)
    style.paragraph_format.space_after=Pt(4)
    style.paragraph_format.line_spacing=1.05
    if name in ['Heading 1','Heading 2']:
        style.paragraph_format.space_before=Pt(4)
doc.styles['Normal'].paragraph_format.widow_control=True
footer=section.footer.paragraphs[0]
footer.alignment=2
footer.add_run('SS-DAC-T2-01  DRAFT  ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
md=['# StudySpace 从架构推导操作指南\n\n# StudySpace Руководство на основе архитектуры\n\nSS-DAC-T2-01 · v1.0 · DRAFT · 2026-10-21 · Asia/Shanghai\n']

def paragraph(text,style=None):
    p=doc.add_paragraph(text,style)
    return p
def paired(text):
    for s in text: paragraph(s)
    md.append(text[0]+'\n\n'+text[1]+'\n')
def heading(zh,ru,number):
    paragraph(f'{number} {zh}', 'Heading 1')
    paragraph(f'{number} {ru}', 'Heading 2')
    md.append(f'## {number} {zh}\n\n## {number} {ru}\n')
def page(): doc.add_page_break()
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,w in enumerate(widths):t.columns[i].width=Cm(w)
    for i,x in enumerate(headers):t.rows[0].cells[i].text=x
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells=t.add_row().cells
        for i,x in enumerate(row):cells[i].text=str(x)
    for r_idx,row in enumerate(t.rows):
        cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
        for i,c in enumerate(row.cells):
            c.width=Cm(widths[i]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr=c._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            tcPr.append(borders)
            margin=OxmlElement('w:tcMar')
            for edge in ['top','bottom','left','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'75');e.set(qn('w:type'),'dxa');margin.append(e)
            tcPr.append(margin)
            if r_idx==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7EDF3');tcPr.append(shade)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(1)
                p.paragraph_format.line_spacing=1.0
                for run in p.runs:
                    run.font.size=Pt(9.4);run.bold=(r_idx==0)
                    run.font.name='Arial';run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'PingFang SC')
    md.append('| '+' | '.join(x.replace('\n','<br>') for x in headers)+' |\n| '+' | '.join('---' for _ in headers)+' |')
    for row in rows: md.append('| '+' | '.join(str(x).replace('|','\\|').replace('\n','<br>') for x in row)+' |')
    md.append('')

paragraph('StudySpace 从架构推导操作指南','Title')
paragraph('StudySpace Руководство на основе архитектуры','Subtitle')
paragraph('SS-DAC-T2-01 · v1.0 · DRAFT · 2026-10-21 · Asia/Shanghai')
heading('目的和完成范围','Цель и область результата','1')
for p in intro: paired(p)
table(['图片要求\nТребование','本成果位置\nРаздел результата'],[
 ('列出组件和节点\nПеречислить компоненты и узлы','第 2 至 6 节，10 CMP 与 7 N\nРазделы 2–6, 10 CMP и 7 N'),
 ('软件 访问 参数 检查\nПО доступ параметры проверка','第 3 至 6 节，各节点 4 项\nРазделы 3–6, четыре поля узла'),
 ('注明依据\nУказать источники','各字段来源标识及第 7 节\nID у каждого поля и раздел 7'),
 ('拟定章节计划\nПлан руководства','第 8 节，10 个章节\nРаздел 8, десять глав'),
 ('标明当前可验证操作\nДоступные проверки','第 9 至 10 节，P-01 至 P-10\nРазделы 9–10, P-01–P-10'),
 ],[7,10.8])

page();heading('组件和节点映射','Компоненты и узлы','2')
paired(pair('组件提取自 S01/S02，节点提取自 S03。保留组件设计状态，CMP-001 在终端执行、N-03 分发；CMP-009 的数据库访问代码在 N-04，N-06 仅承载目标数据库。','Компоненты из S01/S02, узлы из S03. Сохранены статусы проекта; CMP-001 выполняется на терминале и раздаётся N-03; код CMP-009 в N-04, N-06 хранит целевую БД.'))
table(['组件\nКомпонент','名称\nНазвание','节点与实现状态\nУзел и реализация'],[(cid,'\n'.join(name),loc+'\n'+'\n'.join(status)) for cid,name,loc,status in components],[2.1,7.2,8.5])
paired(pair('9 个内部组件与 1 个外部来源。CMP-002 至 CMP-008 尚未实现业务功能，不能写成独立已运行微服务。','9 внутренних компонентов и один внешний источник. Бизнес-функции CMP-002–CMP-008 не реализованы; они не являются запущенными микросервисами.'))

for number, subset, title in [('3',nodes[:2],pair('终端与开发工作站','Терминал и рабочая станция')),('4',nodes[2:4],pair('本地容器节点','Локальные контейнерные узлы')),('5',nodes[4:6],pair('目标入口和数据库','Целевой вход и БД')),('6',nodes[6:],pair('学校来源与待定参数','Источник вуза и параметры'))]:
    page();heading(*title,number)
    for node in subset:
        if node != subset[0]:
            page()
        paragraph(node['id']+' '+node['name'][0],'Heading 2');paragraph(node['name'][1],'Heading 2')
        md.append('### '+node['id']+' '+node['name'][0]+' / '+node['name'][1]+'\n')
        table(['要求\nТребование','操作内容与信息来源\nДействие и источник'],[(zh+'\n'+ru,body[0]+'\n'+body[1]+'\n来源 / Источник '+src) for zh,ru,body,src in node['fields']],[2.2,15.6])
    if number=='6':
        paragraph('假设与待定输入','Heading 2');paragraph('Допущения и входные условия','Heading 2')
        md.append('### 假设与待定输入 / Допущения и входные условия\n')
        for aid,body in assumptions:
            paired((aid+' '+body[0],aid+' '+body[1]))

page();heading('来源登记与使用规则','Источники и правила использования','7')
paired(pair('DIAGRAM 表示图和架构表；ADR 表示已选设计决策；NFR 表示质量需求，仍为 PROPOSED；IMPLEMENTATION 表示可检查配置或代码；EVIDENCE 表示实际执行记录。A-01 至 A-04 为待定输入或操作适配说明，不充当已实现证据。','DIAGRAM — схемы и таблицы, ADR — выбранные решения, NFR — требования качества PROPOSED, IMPLEMENTATION — код и настройки, EVIDENCE — результаты выполнения. A-01–A-04 — открытые условия и адаптация, а не доказательства реализации.'))
table(['ID','类别\nТип','依据和用途\nОснование и применение'],[(sid,kind,'\n'.join(desc)) for sid,kind,path,desc in sources],[1.3,3.1,13.4])
paired(pair('Markdown 下方提供可点击来源路径；Word 以来源编号追溯，完整绝对路径与 SHA-256 在 artifacts/source-manifest.json。S01/S02 是既有课堂组件成果；本次不对未提供的 3.3.2 原始题目或教师批准另作推定。','В Markdown ниже приведены ссылки; Word ссылается по ID. Абсолютные пути и SHA-256 в artifacts/source-manifest.json. S01/S02 — прежние учебные материалы; принадлежность к непредоставленному исходному заданию 3.3.2 и согласование преподавателя не предполагаются.'))
# Detailed links are appendix material in Markdown; source IDs remain readable in Word.
for sid,kind,path,desc in sources:md.append(f'- [{sid} {desc[0]}]({path}) — {kind}')
md.append('')

page();heading('操作指南章节计划','План глав руководства','8')
paired(pair('下表为后续完整指南的章节计划；G-07 至 G-09 保留前提和待执行状态。已确认的 3.4.3 Markdown 指南用于文档编辑，不代替 StudySpace 部署及运维指南。','Ниже план полного руководства; G-07–G-09 сохраняют условия и отложенный статус. Подтверждённое руководство Markdown 3.4.3 помогает редактировать документы, но не заменяет руководство эксплуатации StudySpace.'))
table(['章节\nГлава','计划内容\nПланируемое содержание','节点与依据\nУзлы и основание'],[(gid+'\n'+'\n'.join(title),'\n'.join(body),src) for gid,title,body,src in outline],[4.3,9.8,3.7])

page();heading('当前可验证的流程','Проверяемые сейчас процедуры','9')
paired(pair('VERIFIED 表示有明确执行证据；READY_NOT_RUN 表示现有源码可安排本地验证但本次未运行；CONDITIONAL_NOT_RUN 表示还需运行环境；DEFERRED 表示缺少外部输入或实现。P-01 引用前一阶段证据，其余 VERIFIED 为本次执行。','VERIFIED — есть свидетельства выполнения; READY_NOT_RUN — локальная проверка возможна, сейчас не выполнена; CONDITIONAL_NOT_RUN — нужна среда; DEFERRED — нет реализации или входных данных. P-01 использует прежнее свидетельство, остальные VERIFIED выполнены сейчас.'))
table(['流程与节点\nПроцедура и узел','操作\nДействие','状态与本次证据\nСтатус и свидетельство'],[(pid+'\n'+node,'\n'.join(title),status+'\n'+'\n'.join(body)) for pid,node,title,status,body in procedures],[3.3,5.3,9.2])

page();heading('复核步骤和证据边界','Повторная проверка и границы свидетельств','10')
paired(pair('在学习包根目录执行下列命令，工作目录为 /Users/zhi/Documents/软件生命周期/studyspace/coursework/2026-10-21-docs-as-code-environment-v1。','Команды выполняются из корня учебного пакета /Users/zhi/Documents/软件生命周期/studyspace/coursework/2026-10-21-docs-as-code-environment-v1.'))
for pid,command,result in commands:
    paragraph(pid,'Heading 2');md.append('### '+pid+'\n')
    p=paragraph(command);p.paragraph_format.space_after=Pt(5)
    for run in p.runs:run.font.name='Courier New';run.font.size=Pt(9)
    md.append('```bash\n'+command+'\n```\n');paired(result)
paired(pair('测试仅使用学习副本及临时 SQLite；health 不创建数据库，ready 不创建业务表。Compose 配置检查只验证解析，HTTPS、数据库迁移、并发、备份恢复和真实校内核验均未执行。','Тесты используют учебную копию и временный SQLite; health не создаёт БД, ready не создаёт бизнес-таблиц. Compose проверяет только конфигурацию; HTTPS, миграции, конкуренция, восстановление и реальная проверка вуза не выполнялись.'))
paired(pair('已有 FastAPI TestClient 发出 httpx 弃用提示，8 项测试均通过，依赖未修改。每项正式操作应记录前提、权限、参数、步骤、预期、实际结果、证据和停止方式；未实现的业务命令不编造。','Существующий FastAPI TestClient предупреждает об устаревании httpx; 8 тестов пройдены, зависимости не менялись. Процедуры должны фиксировать условия, права, параметры, шаги, ожидания, результаты, свидетельства и остановку; нереализованные команды не выдумываются.'))
paired(pair('可复核文件：本目录 artifacts/ 下的 foundation-tests.log、docs-check.log、registers-check.log、compose-config.log、execution.json 及 source-manifest.json。Word 排版检查记录另存 verification.json。','Свидетельства в artifacts/: foundation-tests.log, docs-check.log, registers-check.log, compose-config.log, execution.json и source-manifest.json. Проверка Word записывается отдельно в verification.json.'))

doc.core_properties.title='StudySpace 从架构推导操作指南'
doc.core_properties.subject='实践 3.4.4 第 2 部分 中文俄语对照'
doc.core_properties.author='StudySpace 课程小组'
doc.core_properties.comments='DRAFT; not business acceptance'
for element in [doc.styles.element, doc.element]:
    for border in list(element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
doc.save(OUT/(STEM+'.docx'))
(OUT/(STEM+'.md')).write_text('\n'.join(md)+'\n',encoding='utf-8')

fingerprints=[]
for sid,kind,path,desc in sources:
    p=(OUT/path).resolve()
    assert p.exists(),p
    fingerprints.append({'id':sid,'type':kind,'path':str(p),'sha256':digest(p)})
for rel in ['backend/Dockerfile','frontend/Dockerfile','frontend/vite.config.ts','backend/app/database.py','scripts/setup.sh','backend/uv.lock','frontend/package-lock.json']:
    p=PACK/'00-project'/rel
    fingerprints.append({'id':'SUPPORT','type':'IMPLEMENTATION','path':str(p),'sha256':digest(p)})
(OUT/'artifacts/source-manifest.json').write_text(json.dumps({'sources':fingerprints,'copied_component_sources':source_records},ensure_ascii=False,indent=2)+'\n')
print(json.dumps({'docx':str(OUT/(STEM+'.docx')),'markdown':str(OUT/(STEM+'.md')),'components':len(components),'nodes':len(nodes),'node_fields':sum(len(n['fields']) for n in nodes),'chapters':len(outline),'procedures':len(procedures)},ensure_ascii=False))
