"""Render the same structured administrator content into an editable DOCX."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
data = json.loads((ROOT/'artifacts/admin-content.json').read_text())
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(1.7)
sec.top_margin, sec.bottom_margin = Cm(1.5), Cm(1.6)
sec.footer_distance = Cm(.7)
for name, size in [('Normal',10),('Title',19),('Subtitle',13),('Heading 1',14),('Heading 2',11.5),('Heading 3',10.5),('Heading 4',10)]:
    s=doc.styles[name]
    s.font.name='Arial'
    s.font.size=Pt(size)
    s.font.color.rgb=RGBColor(0,0,0)
    s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Hiragino Sans GB')
    s.paragraph_format.space_after=Pt(4)
    s.paragraph_format.line_spacing=1.05
    if name.startswith('Heading'):
        s.paragraph_format.space_before=Pt(8)
        s.paragraph_format.keep_with_next=True
footer=sec.footer.paragraphs[0]
footer.alignment=2
footer.add_run('SS-DAC-ADMIN-01  v0.2  DRAFT   ')
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)
doc.add_paragraph('StudySpace 管理员手册','Title')
doc.add_paragraph('Руководство администратора StudySpace','Subtitle')
doc.add_paragraph('SS-DAC-ADMIN-01 · v0.2 · DRAFT · 2026-10-21 · Asia/Shanghai')
doc.add_paragraph('文档版本 0.2  系统版本 0.1.0  需求版本 0.3.0  说明部分验证')
doc.add_paragraph('Документ 0.2  Система 0.1.0  Требования 0.3.0  Проверено частично')

for b in data['blocks']:
    kind=b['type']
    if kind=='heading':
        level=b['level']
        p=doc.add_paragraph(b['zh'],'Heading '+str(level-1))
        p.paragraph_format.keep_with_next=True
        p=doc.add_paragraph(b['ru'],'Heading '+str(min(level,4)))
        p.paragraph_format.space_before=Pt(1)
    elif kind=='pair':
        p=doc.add_paragraph(b['zh']);p.paragraph_format.keep_with_next=True
        doc.add_paragraph(b['ru'])
    elif kind=='code':
        lines=b['text'].splitlines()
        for i,line in enumerate(lines):
            p=doc.add_paragraph(line)
            p.paragraph_format.space_after=Pt(1)
            p.paragraph_format.line_spacing=1.0
            p.paragraph_format.keep_with_next=i<len(lines)-1
            for r in p.runs:r.font.name='Courier New';r.font.size=Pt(8)
    elif kind=='image':
        doc.add_paragraph('React → Vite /api → FastAPI → SQLAlchemy → SQLite')
    elif kind=='reference':
        doc.add_paragraph(b['id']+' '+b['label'])
        p=doc.add_paragraph('docs/materials/'+b['path'])
        for r in p.runs:r.font.size=Pt(8)
    elif kind=='table':
        table=doc.add_table(rows=1, cols=len(b['headers']))
        table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
        widths=b.get('widths',[17.6/len(b['headers'])]*len(b['headers']))
        factor=17.6/sum(widths)
        widths=[w*factor for w in widths]
        for i,w in enumerate(widths):table.columns[i].width=Cm(w)
        for i,value in enumerate(b['headers']):table.rows[0].cells[i].text=value
        head=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(head)
        for row in b['rows']:
            cells=table.add_row().cells
            for i,value in enumerate(row):cells[i].text=value
        for ri,row in enumerate(table.rows):
            trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
            for i,cell in enumerate(row.cells):
                cell.width=Cm(widths[i]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                pr=cell._tc.get_or_add_tcPr()
                edges=OxmlElement('w:tcBorders')
                for edge in ['top','bottom','left','right']:
                    el=OxmlElement('w:'+edge)
                    for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]:el.set(qn('w:'+k),v)
                    edges.append(el)
                pr.append(edges)
                margins=OxmlElement('w:tcMar')
                for edge in ['top','bottom','left','right']:
                    el=OxmlElement('w:'+edge);el.set(qn('w:w'),'75');el.set(qn('w:type'),'dxa');margins.append(el)
                pr.append(margins)
                if ri==0:
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E5EDF3');pr.append(shade)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(2)
                    p.paragraph_format.line_spacing=1.0
                    for run in p.runs:
                        run.font.name='Arial';run.font.size=Pt(9);run.bold=ri==0
                        run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Hiragino Sans GB')
        doc.add_paragraph()
for tree in [doc.styles.element,doc.element]:
    for el in list(tree.iter(qn('w:pBdr'))):el.getparent().remove(el)
doc.core_properties.title='StudySpace 管理员手册'
doc.core_properties.author='StudySpace 课程小组'
target=ROOT/'StudySpace_管理员手册_中俄双语_v0.2.docx'
doc.save(target)
print(target)
