"""Create the classroom comparison table for existing ADR-0001 and ADR-0002."""

import hashlib
import json
import re
import sys
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "coursework/2026-10-14-architecture-decisions-v1"
DATA = OUT / "data"
NAME = "StudySpace_两份ADR架构决策记录_中俄双语_v1"
DEST = OUT / (NAME + ".docx")
SKILL = Path(
    "/Users/zhi/.codex/plugins/cache/openai-primary-runtime/documents/26.1013.11914/skills/documents"
)
sys.path.insert(0, str(SKILL / "scripts"))
from table_geometry import (  # noqa: E402 — bundled helper path is added above.
    apply_table_geometry,
    audit_docx_tables,
    column_widths_from_weights,
    section_content_width_dxa,
)

if DEST.exists() and "--replace-task-output" not in sys.argv:
    raise SystemExit("Deliverable exists. Inspect it before replacing task output.")


def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()


def bi(zh, ru):
    return {"zh-CN": zh, "ru-RU": ru}


paths = [
    "AGENTS.md",
    "docs/data/project-decisions.json",
    "docs/data/requirements.json",
    "docs/zh-CN/04_System_Design.md",
    "docs/ru-RU/04_System_Design.md",
    "docs/zh-CN/16_Requirements_Clarification.md",
    "docs/ru-RU/16_Requirements_Clarification.md",
    "docs/decisions/README.md",
    "docs/decisions/zh-CN/0001-simple-monolith.md",
    "docs/decisions/ru-RU/0001-simple-monolith.md",
    "docs/decisions/zh-CN/0002-sqlite-postgresql.md",
    "docs/decisions/ru-RU/0002-sqlite-postgresql.md",
    "coursework/2026-10-13-requirements-v0.2/data/requirements-practice.json",
]
paths += [
    str(p.relative_to(ROOT)) for p in (ROOT / "coursework").rglob("*.docx") if OUT not in p.parents
]
baseline = {p: sha(ROOT / p) for p in paths}
reqs = {x["id"]: x for x in json.loads((ROOT / paths[12]).read_text())["requirements"]}
decisions = {
    x["id"]: x
    for x in json.loads((ROOT / "docs/data/project-decisions.json").read_text())["decisions"]
}

fields = [
    ("name_status", bi("名称与状态", "Название и статус")),
    (
        "context",
        bi("背景\n需求 约束 驱动因素", "Контекст\nтребования ограничения силы"),
    ),
    ("decision", bi("决策", "Решение")),
    ("consequences", bi("后果\n优点与缺点", "Последствия\nплюсы и минусы")),
    ("alternatives", bi("替代方案及放弃原因", "Альтернативы и причины отказа")),
]

adrs = [
    dict(
        id="ADR-0001",
        decision_status="ACCEPTED",
        document_status="DRAFT",
        requirement_ids=[
            "SP-C-005",
            "SP-C-006",
            "SP-NFR-001",
            "SP-IF-001",
            "SP-FR-004",
            "SP-NFR-004",
        ],
        decision_ids=["DEC-001", "DEC-017"],
        fields={
            "name_status": [
                bi("前后端分离的模块化单体", "Клиент и модульный монолитный сервер"),
                bi(
                    "技术决策 ACCEPTED\n文档状态 DRAFT",
                    "Техническое решение ACCEPTED\nДокумент DRAFT",
                ),
            ],
            "context": [
                bi(
                    "四人课程团队需协作交付查询、校内审核和预约；系统要求三语界面、统一 API、可重复构建及服务端权限和事务校验，开发效率与运行复杂度影响选择。",
                    "Команда из четырёх участников создаёт поиск, проверку вуза и бронирование; нужны три языка, единый API, воспроизводимая сборка, серверные права и транзакции при приемлемой сложности разработки и запуска.",
                )
            ],
            "decision": [
                bi(
                    "沿用 React／TypeScript 浏览器客户端和 FastAPI 后端，经 JSON API 通信；后端按认证、审核、资源、预约和统计划分内部模块，保持单体部署。",
                    "Сохранить браузерный клиент React/TypeScript и сервер FastAPI с JSON API; внутри сервера выделить аутентификацию, проверку, ресурсы, бронирование и статистику при монолитном развёртывании.",
                ),
                bi(
                    "Pydantic 校验输入，同步 SQLAlchemy 访问数据库；权限和预约冲突由服务端事务最终校验，uv/npm 锁定依赖。",
                    "Pydantic проверяет входы, синхронный SQLAlchemy обращается к БД; права и конфликты окончательно проверяются серверной транзакцией, зависимости фиксируются uv/npm.",
                ),
            ],
            "consequences": [
                bi(
                    "优点 运行单元较少，业务模块可共享事务，便于四人团队开发、测试和交接。",
                    "Плюсы Меньше единиц запуска, общая транзакция для бизнес-модулей, проще разработка, проверка и передача небольшой командой.",
                ),
                bi(
                    "缺点 前后端需同步 API 契约；后端故障可影响全部业务，模块间依赖需控制，独立扩缩容能力受限。",
                    "Минусы Нужно согласовывать API клиента и сервера; сбой сервера затрагивает весь бизнес, зависимости модулей требуют контроля, отдельное масштабирование ограничено.",
                ),
            ],
            "alternatives": [
                bi(
                    "微服务 可独立部署和扩缩容，但增加接口、部署及跨服务一致性成本，本次课程范围不采纳。",
                    "Микросервисы дают отдельное развёртывание и масштабирование, но усложняют интерфейсы, запуск и согласованность между сервисами; для текущего курса не выбраны.",
                ),
                bi(
                    "单一全栈框架 可减少前后端衔接工作，但需改造已建 React/FastAPI 骨架，本次保留现有技术基础。",
                    "Один full-stack фреймворк уменьшает связующую работу, но требует переделки каркаса React/FastAPI; в этом выпуске сохранена имеющаяся основа.",
                ),
            ],
        },
    ),
    dict(
        id="ADR-0002",
        decision_status="ACCEPTED",
        document_status="DRAFT",
        requirement_ids=[
            "SP-C-005",
            "SP-C-001",
            "SP-C-002",
            "SP-NFR-004",
            "SP-NFR-005",
            "SP-TR-001",
            "SP-DR-003",
        ],
        decision_ids=["DEC-015", "DEC-016"],
        fields={
            "name_status": [
                bi(
                    "SQLite 开发与 PostgreSQL 目标部署",
                    "SQLite для разработки и PostgreSQL для целевого развёртывания",
                ),
                bi(
                    "技术决策 ACCEPTED\n文档状态 DRAFT",
                    "Техническое решение ACCEPTED\nДокумент DRAFT",
                ),
            ],
            "context": [
                bi(
                    "开发环境需便于本地启动；预约要求阻止同学生或同座位的时段重叠，并保持事务原子性、UTC 时间语义和迁移后数据关联，目标数据库需独立验证。",
                    "Нужен простой локальный запуск; бронь требует запрета пересечений студента и места, атомарности, семантики UTC и сохранения связей при миграции, целевая БД проверяется отдельно.",
                )
            ],
            "decision": [
                bi(
                    "开发默认 SQLite，目标部署候选为 PostgreSQL，通过 SQLAlchemy 和 DATABASE_URL 统一访问与配置；首个业务实体前引入 Alembic。",
                    "В разработке использовать SQLite, кандидат для целевого развёртывания PostgreSQL; SQLAlchemy и DATABASE_URL задают общий доступ и конфигурацию, Alembic вводится до первой бизнес-сущности.",
                ),
                bi(
                    "在两库分别验证迁移、回滚和预约并发；使用 [start,end) 区间，同一学生可有多条非重叠预约，创建与取消在服务端事务内完成。",
                    "На обеих БД отдельно проверить миграцию, откат и конкурентное бронирование; интервалы [start,end), несколько непересекающихся броней студента допустимы, создание и отмена выполняются серверной транзакцией.",
                ),
            ],
            "consequences": [
                bi(
                    "优点 简化本地开发，同时保留目标数据库迁移路径和统一数据访问层。",
                    "Плюсы Простой локальный запуск с сохранением пути к целевой БД и общего слоя доступа к данным.",
                ),
                bi(
                    "缺点 需维护两库的迁移与测试；ORM 不替代事务、约束和时区验证，上线前还需备份恢复演练。",
                    "Минусы Нужны миграции и проверки обеих БД; ORM не заменяет проверку транзакций, ограничений и времени, до запуска требуется восстановление из резервной копии.",
                ),
            ],
            "alternatives": [
                bi(
                    "全程仅 SQLite 配置简单，但不覆盖已选择的 PostgreSQL 目标环境及其并发、迁移验证，因此不作为唯一方案。",
                    "Только SQLite упрощает настройку, но не покрывает выбранную целевую среду PostgreSQL и её проверки конкуренции и миграции; не выбран единственным вариантом.",
                ),
                bi(
                    "全程仅 PostgreSQL 可减少两库差异，但为每位成员的基础本地运行增加数据库服务依赖，本阶段保留 SQLite 开发入口。",
                    "Только PostgreSQL сокращает различия БД, но добавляет зависимость от сервиса БД для базового локального запуска каждого участника; на этом этапе сохранён вход через SQLite.",
                ),
            ],
        },
    ),
]

packet = dict(
    doc_id="SS-ADR-T4-01",
    version="v1",
    date="2026-10-14",
    document_status="DRAFT",
    languages=["zh-CN", "ru-RU"],
    task=bi("任务 4 两份 ADR", "Задание 4 два ADR"),
    source_image="/var/folders/ff/kqw394c93dd7krl0z2hcqqdh0000gn/T/codex-clipboard-00a74dfa-771f-488b-a1d9-b982bcf7a2e2.jpg",
    source_image_sha256=sha(
        Path(
            "/var/folders/ff/kqw394c93dd7krl0z2hcqqdh0000gn/T/codex-clipboard-00a74dfa-771f-488b-a1d9-b982bcf7a2e2.jpg"
        )
    ),
    source_fingerprints=baseline,
    fields=[{"key": k, "label": v} for k, v in fields],
    adrs=adrs,
    source_versions={"engineering": "0.3.0", "classroom": "0.2"},
    business_tests="NOT_RUN",
    deployment_status="NOT_DEPLOYED",
    review_conditions=bi(
        "学校核验来源 OPEN-003、课程条件 OPEN-008 和部署环境 OPEN-010 明确，或业务模型、负载及模块边界变化时复查；迁移、并发、恢复和业务验收分别保存实际证据。",
        "Пересмотреть после определения источника вуза OPEN-003, условий курса OPEN-008 и среды OPEN-010 либо изменения модели, нагрузки и границ модулей; миграция, конкуренция, восстановление и бизнес-приёмка требуют отдельных свидетельств.",
    ),
)
for adr in adrs:
    assert set(adr["requirement_ids"]) <= reqs.keys()
    assert set(adr["decision_ids"]) <= decisions.keys()

OUT.mkdir(parents=True, exist_ok=True)
DATA.mkdir(exist_ok=True)
(DATA / "two-adrs-v1.json").write_text(json.dumps(packet, ensure_ascii=False, indent=2) + "\n")
(DATA / "fonts.conf").write_text("""<?xml version="1.0"?>
<!DOCTYPE fontconfig SYSTEM "urn:fontconfig:fonts.dtd">
<fontconfig>
 <dir>/System/Library/Fonts</dir><dir>/System/Library/Fonts/Supplemental</dir>
 <dir>/Library/Fonts</dir><dir>/Users/zhi/Library/Fonts</dir>
 <cachedir>/private/tmp/studyspace-adr-fontcache</cachedir>
 <match target="pattern"><test name="family"><string>sans-serif</string></test><edit name="family" mode="prepend"><string>Arial</string><string>PingFang SC</string></edit></match>
</fontconfig>
""")

doc = Document()
section = doc.sections[0]
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width, section.page_height = Inches(11), Inches(8.5)
section.left_margin = section.right_margin = Inches(0.6)
section.top_margin = section.bottom_margin = Inches(0.55)
grid = section._sectPr.find(qn("w:docGrid"))
if grid is not None:
    section._sectPr.remove(grid)
for st in doc.styles:
    for e in list(st._element.iter(qn("w:pBdr"))):
        e.getparent().remove(e)
    for font in st._element.iter(qn("w:rFonts")):
        for a in ["asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"]:
            font.attrib.pop(qn("w:" + a), None)
for name in ["Normal", "Title", "Heading 1", "Heading 2"]:
    st = doc.styles[name]
    st.font.name = "Arial"
    st.font.color.rgb = RGBColor(0, 0, 0)
    st._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "PingFang SC")
doc.styles["Normal"].font.size = Pt(11)
doc.styles["Normal"].paragraph_format.space_after = Pt(5)
doc.styles["Normal"].paragraph_format.line_spacing = Pt(14)
doc.styles["Title"].font.size = Pt(19)
doc.styles["Title"].paragraph_format.space_after = Pt(3)
doc.styles["Heading 1"].font.size = Pt(12)
doc.styles["Heading 1"].paragraph_format.space_before = Pt(9)
doc.styles["Heading 1"].paragraph_format.space_after = Pt(4)


def format_p(p, size=11, bold=False, center=False):
    p.paragraph_format.line_spacing = Pt(14)
    snap = OxmlElement("w:snapToGrid")
    snap.set(qn("w:val"), "0")
    p._p.get_or_add_pPr().append(snap)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    for run in p.runs:
        run.font.name = "Arial"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor(0, 0, 0)
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "PingFang SC")
    return p


title = doc.add_paragraph("StudySpace 两份架构决策记录", "Title")
format_p(title, size=19)
format_p(doc.add_paragraph("StudySpace Две записи архитектурных решений"), size=13)
format_p(
    doc.add_paragraph("任务 4  Задание 4  |  SS-ADR-T4-01  |  v1  |  2026-10-14"),
    size=9,
)
format_p(
    doc.add_paragraph(
        "两份 ADR 分别记录应用架构和数据库策略；ACCEPTED 表示项目已选定技术方案，DRAFT 表示课程文档仍待正式评审。\nДве ADR фиксируют архитектуру и стратегию БД; ACCEPTED означает выбранное техническое решение, DRAFT — документ до формального рассмотрения."
    ),
    size=10,
)

table = doc.add_table(rows=1, cols=3)
for c, text in zip(table.rows[0].cells, ["字段\nПоле", "ADR-0001", "ADR-0002"], strict=False):
    c.text = text
table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
for key, label in fields:
    cells = table.add_row().cells
    cells[0].text = label["zh-CN"] + "\n" + label["ru-RU"]
    for c, adr in zip(cells[1:], adrs, strict=False):
        for i, pair in enumerate(adr["fields"][key]):
            p = c.paragraphs[0] if i == 0 else c.add_paragraph()
            p.text = pair["zh-CN"] + "\n" + pair["ru-RU"]
            p.paragraph_format.space_after = Pt(4 if i < len(adr["fields"][key]) - 1 else 0)
        if key == "context":
            c.paragraphs[-1].paragraph_format.space_after = Pt(4)
            c.add_paragraph(
                "需求 / Требования\n" + ", ".join(adr["requirement_ids"])
            ).paragraph_format.space_after = Pt(0)

apply_table_geometry(
    table,
    column_widths_from_weights([15, 42.5, 42.5], section_content_width_dxa(section)),
    cell_margins_dxa={"top": 95, "bottom": 95, "start": 110, "end": 110},
)
borders = OxmlElement("w:tblBorders")
for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
    e = OxmlElement("w:" + edge)
    for k, v in [("val", "single"), ("sz", "5"), ("color", "D9D9D9")]:
        e.set(qn("w:" + k), v)
    borders.append(e)
table._tbl.tblPr.append(borders)
for i, row in enumerate(table.rows):
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    for j, c in enumerate(row.cells):
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        fill = OxmlElement("w:shd")
        fill.set(qn("w:fill"), "DDEBF7" if i == 0 else ("F4F7FA" if i % 2 == 0 else "FFFFFF"))
        c._tc.get_or_add_tcPr().append(fill)
        for p in c.paragraphs:
            format_p(p, size=10.5 if j == 0 else 11, bold=(i == 0 or j == 0), center=(i == 0))
            if p.text.startswith("需求 / Требования"):
                format_p(p, size=10)
            p.paragraph_format.keep_with_next = False
            if i == 0:
                p.paragraph_format.space_after = Pt(0)

doc.add_paragraph("复查条件  Условия пересмотра", "Heading 1")
format_p(
    doc.add_paragraph(
        packet["review_conditions"]["zh-CN"] + "\n" + packet["review_conditions"]["ru-RU"]
    ),
    size=10,
)
doc.add_paragraph("来源  Источники", "Heading 1")
format_p(
    doc.add_paragraph(
        "既有 ADR-0001、ADR-0002；DOC-04 系统设计 0.3.0；DEC-001、DEC-015、DEC-016、DEC-017；课堂需求登记 v0.2。\nИсходные ADR-0001 и ADR-0002; DOC-04 проект системы 0.3.0; DEC-001, DEC-015, DEC-016, DEC-017; учебный реестр требований v0.2."
    ),
    size=9,
)
doc.core_properties.title = "StudySpace 两份架构决策记录 Две ADR"
doc.core_properties.subject = "任务4 五字段三列ADR对照表"
doc.core_properties.author = "StudySpace"
doc.core_properties.language = "zh-CN; ru-RU"
doc.save(DEST)

md = [
    "# StudySpace 两份架构决策记录",
    "# StudySpace Две записи архитектурных решений",
    "任务 4 / Задание 4 · SS-ADR-T4-01 · v1 · 2026-10-14 · DRAFT",
    "",
    "| 字段 / Поле | ADR-0001 | ADR-0002 |",
    "| --- | --- | --- |",
]
for key, label in fields:
    contents = [(label["zh-CN"] + " / " + label["ru-RU"]).replace("\n", "<br>")]
    for adr in adrs:
        content = "<br><br>".join(
            (p["zh-CN"] + "<br>" + p["ru-RU"]).replace("\n", "<br>") for p in adr["fields"][key]
        )
        if key == "context":
            content += "<br><br>需求 / Требования<br>" + ", ".join(adr["requirement_ids"])
        contents.append(content)
    md.append("| " + " | ".join(contents) + " |")
md += [
    "",
    "## 复查条件 Условия пересмотра",
    "",
    packet["review_conditions"]["zh-CN"],
    packet["review_conditions"]["ru-RU"],
    "",
    "## 来源 Источники",
    "",
    *paths[:13],
    "",
]
(OUT / (NAME + ".md")).write_text("\n".join(md))
(OUT / "README.md").write_text(f"""# StudySpace 两份 ADR 课堂作业 Две учебные ADR

2026-10-14 · v1 · SS-ADR-T4-01 · DRAFT

交付 / Результат [{NAME}.docx]({NAME}.docx)

按任务4的五字段三列表格填写ADR-0001和ADR-0002，中文与俄语在同一单元格配对。
В пяти полях и трёх колонках задания4 заполнены ADR-0001 и ADR-0002; китайский и русский тексты объединены в каждой ячейке.

两项技术决定沿用ACCEPTED，正式文档DRAFT；迁移、并发、恢复和业务验收保持未执行状态。既有项目文件保持原样。
Технические решения сохранены как ACCEPTED, документ DRAFT; миграция, конкуренция, восстановление и бизнес-приёмка не объявляются выполненными. Исходные файлы сохранены.

来源与内容 / Источники и содержимое data/two-adrs-v1.json
验证 / Проверка data/verification.json
""")

checks = []


def check(name, value, detail=None):
    checks.append({"name": name, "passed": bool(value), "detail": detail})


check("two_existing_adr_ids", [a["id"] for a in adrs] == ["ADR-0001", "ADR-0002"])
check(
    "all_five_slide_fields",
    all(set(a["fields"]) == {k for k, _ in fields} for a in adrs),
)
check(
    "all_cells_bilingual",
    all(
        set(p) == {"zh-CN", "ru-RU"} and all(p.values())
        for a in adrs
        for v in a["fields"].values()
        for p in v
    ),
)
check(
    "requirements_and_decisions_exist",
    all(
        set(a["requirement_ids"]) <= reqs.keys() and set(a["decision_ids"]) <= decisions.keys()
        for a in adrs
    ),
)
check(
    "accepted_status_matches_original_adrs",
    all("| Decision-Status | ACCEPTED |" in (ROOT / p).read_text() for p in paths[8:12]),
)
check(
    "latest_confirmed_four_person_team",
    "四人" in adrs[0]["fields"]["context"][0]["zh-CN"]
    and decisions["DEC-001"]["status"] == "CONFIRMED",
)
check(
    "alembic_boundary_matches_dec015",
    "Alembic" in adrs[1]["fields"]["decision"][0]["zh-CN"]
    and "Alembic" in decisions["DEC-015"]["description"]["zh-CN"],
)
check(
    "dec016_multiple_nonoverlapping_bookings",
    "多条非重叠预约" in adrs[1]["fields"]["decision"][1]["zh-CN"]
    and "несколько непересекающихся" in adrs[1]["fields"]["decision"][1]["ru-RU"],
)
check(
    "editable_three_column_five_field_table",
    len(doc.tables) == 1 and len(table.rows) == 6 and all(len(r.cells) == 3 for r in table.rows),
)
check(
    "all_source_content_in_docx",
    all(
        p[lang] in "\n".join(c.text for r in table.rows for c in r.cells)
        for a in adrs
        for v in a["fields"].values()
        for p in v
        for lang in ["zh-CN", "ru-RU"]
    ),
)
check("matching_table_geometry", audit_docx_tables(DEST) == 0)
check(
    "existing_documents_and_sources_unchanged",
    all(sha(ROOT / p) == h for p, h in baseline.items()),
    baseline,
)
with ZipFile(DEST) as z:
    xml = z.read("word/document.xml").decode()
    check(
        "repeat_header_and_no_row_split",
        "<w:tblHeader" in xml and xml.count("<w:cantSplit") == 6,
    )
    check(
        "no_header_footer_or_title_border",
        not any(re.match(r"word/(header|footer)\d+\.xml$", n) for n in z.namelist())
        and "<w:pBdr" not in xml
        and "<w:pBdr" not in z.read("word/styles.xml").decode(),
    )
check(
    "draft_document_without_business_pass_claim",
    packet["document_status"] == "DRAFT" and packet["business_tests"] == "NOT_RUN",
)
report = dict(
    date="2026-10-14",
    document=str(DEST),
    working_directory=str(ROOT),
    runtime="Codex bundled runtime 26.1013.11914",
    sha256=sha(DEST),
    checks=checks,
    check_count=len(checks),
    all_pass=all(x["passed"] for x in checks),
    visual_review="PENDING",
    business_tests="NOT_RUN",
)
(DATA / "verification.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
print(
    json.dumps(
        {
            "document": str(DEST),
            "all_pass": report["all_pass"],
            "checks": len(checks),
            "failed": [c["name"] for c in checks if not c["passed"]],
        },
        ensure_ascii=False,
    )
)
if not report["all_pass"]:
    raise SystemExit(1)
