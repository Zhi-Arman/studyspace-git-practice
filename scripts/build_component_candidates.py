"""Build and structurally verify the bilingual component-candidate coursework table."""
from __future__ import annotations

import hashlib
import json
import re
import sys
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

SKILL = Path('/Users/zhi/.codex/plugins/cache/openai-primary-runtime/documents/26.1013.11914/skills/documents')
sys.path.insert(0, str(SKILL / 'scripts'))
from table_geometry import apply_table_geometry, audit_docx_tables, column_widths_from_weights, section_content_width_dxa

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'coursework/2026-10-14-component-candidates-v1'
DATA = OUT / 'data'
NAME = 'StudySpace_组件候选表_中俄双语_v1'
DEST = OUT / (NAME + '.docx')
if DEST.exists() and '--replace-task-output' not in sys.argv:
    raise SystemExit('The deliverable already exists. Inspect it before replacing it.')

REQ_SOURCE = ROOT / 'coursework/2026-10-13-requirements-v0.2/data/requirements-practice.json'
STORY_SOURCE = ROOT / 'coursework/2026-10-14-product-backlog-v1/data/backlog-v1.json'
REQUIREMENTS = json.loads(REQ_SOURCE.read_text())
BACKLOG = json.loads(STORY_SOURCE.read_text())
REQ = {r['id']: r for r in REQUIREMENTS['requirements']}
STORY = {s['id']: s for s in BACKLOG['items']}
sha = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()
protected = [REQ_SOURCE, STORY_SOURCE, ROOT/'docs/data/project-decisions.json',
             ROOT/'docs/data/requirements.json', ROOT/'AGENTS.md',
             *sorted((ROOT/'coursework/2026-10-14-product-backlog-v1').glob('*.docx'))]
before = {str(p.relative_to(ROOT)): sha(p) for p in protected}

def bi(zh, ru):
    return {'zh-CN': zh, 'ru-RU': ru}

def fr(*numbers):
    return [f'SP-FR-{n:03}' for n in numbers]

rows = [
    dict(id='CMP-001', name=bi('三语浏览器界面', 'Браузерный интерфейс на трёх языках'),
         responsibility=bi('向学生和管理员提供查询、预约、取消及审核的中俄英操作界面，并通过统一业务 API 显示结果。',
                           'Предоставляет студентам и администраторам интерфейс поиска, бронирования, отмены и проверки на китайском, русском и английском языках с результатами через единый бизнес API.'),
         requirement_ids=['SP-IF-001', 'SP-NFR-006', 'SP-C-004'],
         scenario_ids=['US-001', 'US-002', 'US-003', 'US-004', 'US-005', 'US-006', 'US-015'],
         display_requirements='SP-IF-001\nSP-NFR-006, SP-C-004', display_scenarios='US-001…006, US-015', boundary='INTERNAL'),
    dict(id='CMP-002', name=bi('认证与权限', 'Аутентификация и права доступа'),
         responsibility=bi('负责学生注册、登录退出、首位管理员创建及账号角色管理，并在服务端检查访问资格。',
                           'Обеспечивает регистрацию, вход и выход, создание первого администратора и управление ролями с проверкой допуска на сервере.'),
         requirement_ids=fr(1, 2, 3, 4, 26, 31)+['SP-NFR-009', 'SP-NFR-010'],
         scenario_ids=['US-007', 'US-008', 'US-004', 'EN-001'],
         display_requirements='SP-FR-001…004\nSP-FR-026, SP-FR-031\nSP-NFR-009…010', display_scenarios='US-007, US-008\nUS-004, EN-001', boundary='INTERNAL'),
    dict(id='CMP-003', name=bi('校内审核与账号资格', 'Проверка вуза и допуск аккаунта'),
         responsibility=bi('依据可信学校来源记录校内审核决定，并在账号停用或身份变更时联动撤销未结束预约。',
                           'Фиксирует решение о принадлежности к вузу по доверенному источнику и атомарно отменяет незавершённые брони при отключении аккаунта или смене идентификаторов.'),
         requirement_ids=fr(20, 21, 22)+['SP-BR-003', 'SP-IF-002', 'SP-DR-001'],
         scenario_ids=['US-004'],
         display_requirements='SP-FR-020…022\nSP-BR-003, SP-IF-002\nSP-DR-001', display_scenarios='US-004', boundary='INTERNAL'),
    dict(id='CMP-004', name=bi('资源与可用性', 'Ресурсы и доступность'),
         responsibility=bi('维护房间、座位和开放规则，并按位置、日期与时段返回可预约的单座位。',
                           'Ведёт помещения, места и правила открытия и возвращает доступные отдельные места по расположению, дате и интервалу.'),
         requirement_ids=fr(5, 6, 7, 14, 15, 16, 18)+['SP-NFR-002', 'SP-C-003'],
         scenario_ids=['US-001', 'US-005', 'US-009'],
         display_requirements='SP-FR-005…007\nSP-FR-014…016, 018\nSP-NFR-002, SP-C-003', display_scenarios='US-001, US-005, US-009', boundary='INTERNAL'),
    dict(id='CMP-005', name=bi('预约事务', 'Транзакции бронирования'),
         responsibility=bi('在事务内创建或幂等取消单座位预约、阻止学生及座位时段重叠，并处理资源停用或开放缩减后的受影响预约。',
                           'В транзакции создаёт или идемпотентно отменяет бронь одного места, запрещает пересечения студента и места и отменяет затронутые записи при отключении ресурса или сокращении часов.'),
         requirement_ids=fr(8, 9, 10, 11, 12, 17)+['SP-BR-002', 'SP-NFR-004', 'SP-C-001', 'SP-C-002'],
         scenario_ids=['US-002', 'US-003', 'US-005', 'US-010'],
         display_requirements='SP-FR-008…012, 017\nSP-BR-002, SP-NFR-004\nSP-C-001…002', display_scenarios='US-002, US-003\nUS-005, US-010', boundary='INTERNAL'),
    dict(id='CMP-006', name=bi('预约记录与统计', 'Записи и статистика бронирований'),
         responsibility=bi('在权限与保留期范围内查询预约记录，并按同一业务日期和 as_of 快照计算创建数、取消数及未结束数。',
                           'Возвращает записи в пределах прав и срока хранения и вычисляет числа созданных, отменённых и незавершённых броней для одной деловой даты и снимка as_of.'),
         requirement_ids=fr(13, 19, 23, 24, 25)+['SP-BR-004', 'SP-DR-002'],
         scenario_ids=['US-006'],
         display_requirements='SP-FR-013, SP-FR-019\nSP-FR-023…025\nSP-BR-004, SP-DR-002', display_scenarios='US-006', boundary='INTERNAL'),
    dict(id='CMP-007', name=bi('数据存储与恢复', 'Хранение и восстановление данных'),
         responsibility=bi('持久化最少必要的业务及审核记录，并执行数据保留、身份清理、数据库迁移和备份恢复策略。',
                           'Хранит минимально необходимые бизнес записи и аудит и выполняет политику сроков хранения, удаления идентификаторов, миграции БД и резервного восстановления.'),
         requirement_ids=['SP-NFR-003', 'SP-NFR-005', 'SP-DR-001', 'SP-DR-002', 'SP-DR-003', 'SP-C-005', 'SP-TR-001'],
         scenario_ids=['EN-003', 'EN-004'],
         display_requirements='SP-NFR-003, SP-NFR-005\nSP-DR-001…003\nSP-C-005, SP-TR-001', display_scenarios='EN-003, EN-004', boundary='INTERNAL'),
    dict(id='CMP-008', name=bi('可信学校核验来源', 'Доверенный источник вуза'),
         responsibility=bi('向管理员提供经学校确认且可验证的校内身份依据，供内部审核模块引用。',
                           'Предоставляет администратору признанное вузом и проверяемое основание принадлежности студента к вузу для ссылки во внутреннем модуле проверки.'),
         requirement_ids=['SP-IF-002', 'SP-FR-022', 'SP-BR-003'],
         scenario_ids=['US-004'], display_requirements='SP-IF-002, SP-FR-022\nSP-BR-003',
         display_scenarios='US-004\nOPEN-003', boundary='EXTERNAL'),
]

notes = [
    bi('划分边界 内部指由 StudySpace 团队维护的模块，外部指学校提供的可信核验来源；数据库属于内部基础设施，单独运行不改变其归属。',
       'Граница Внутренние модули ведёт команда StudySpace, внешний источник проверки предоставляет вуз; отдельный процесс БД остаётся внутренней инфраструктурой.'),
    bi('协作方式 内部组件沿用模块化单体，界面调用业务 API；资源模块提供配置，预约模块在服务端事务内作最终校验，账号或资源变更复用预约取消逻辑。',
       'Взаимодействие Внутренние компоненты образуют модульный монолит, интерфейс вызывает бизнес API; ресурсы дают настройки, бронь окончательно проверяется серверной транзакцией, изменения аккаунтов и ресурсов используют общую отмену.'),
    bi('预约规则 每条预约仅含一个座位，同一学生的未取消预约不得重叠，相邻时段及多条非重叠预约允许，沿用 DEC-016。',
       'Правило Каждая бронь содержит одно место, неотменённые интервалы студента не пересекаются, соседние интервалы и несколько непересекающихся броней допустимы по DEC-016.'),
    bi('外部前提 OPEN-003 的真实核验渠道尚未确定，可采用学校认可的名单或系统供人工核验；不预设已有 API，来源未知或不可验证时不能批准 VERIFIED。',
       'Внешнее условие Реальный канал OPEN-003 не определён; признанный вузом список или система может использоваться для ручной проверки, API не предполагается, неизвестный либо непроверяемый источник не даёт VERIFIED.'),
    bi('版本范围 本表覆盖 MVP 的全部功能需求，提醒、QR 签到、违约记录及实际使用率仍属未来范围 SP-FR-027…030。',
       'Объём Таблица покрывает все функциональные требования MVP; напоминания, QR регистрация присутствия, нарушения и фактическая загрузка остаются будущим объёмом SP-FR-027…030.'),
]

metadata = dict(doc_id='SS-COMP-01', version='v1', updated='2026-10-14', status='DRAFT',
                component_status='PROPOSED', languages=['zh-CN', 'ru-RU'],
                source_course_version='0.2', source_engineering_version='0.3.0',
                source_image='/var/folders/ff/kqw394c93dd7krl0z2hcqqdh0000gn/T/codex-clipboard-551be475-70f2-4e9d-a39d-ad2f92ac89ef.jpg',
                source_image_sha256=sha(Path('/var/folders/ff/kqw394c93dd7krl0z2hcqqdh0000gn/T/codex-clipboard-551be475-70f2-4e9d-a39d-ad2f92ac89ef.jpg')),
                source_fingerprints=before, boundary_counts={'INTERNAL': 7, 'EXTERNAL': 1},
                implementation_claim='NONE', business_test_status='NOT_RUN', rows=rows, notes=notes)

all_fr = {r['id'] for r in REQUIREMENTS['requirements'] if r['type']=='FR' and r['scope']=='MVP'}
linked = {rid for row in rows for rid in row['requirement_ids']}
assert all_fr <= linked, sorted(all_fr-linked)
for row in rows:
    assert set(row['requirement_ids']) <= REQ.keys(), row['id']
    assert set(row['scenario_ids']) <= STORY.keys(), row['id']

OUT.mkdir(parents=True, exist_ok=True)
DATA.mkdir(exist_ok=True)
(DATA/'component-candidates-v1.json').write_text(json.dumps(metadata, ensure_ascii=False, indent=2)+'\n')
(DATA/'fonts.conf').write_text('''<?xml version="1.0"?>
<!DOCTYPE fontconfig SYSTEM "urn:fontconfig:fonts.dtd">
<fontconfig>
  <dir>/System/Library/Fonts</dir>
  <dir>/System/Library/Fonts/Supplemental</dir>
  <dir>/Library/Fonts</dir>
  <dir>/Users/zhi/Library/Fonts</dir>
  <cachedir>/private/tmp/studyspace-component-fontcache</cachedir>
  <match target="pattern"><test name="family"><string>sans-serif</string></test><edit name="family" mode="prepend"><string>Arial</string><string>PingFang SC</string></edit></match>
</fontconfig>
''')

doc = Document()
s = doc.sections[0]
s.orientation = WD_ORIENT.LANDSCAPE
s.page_width, s.page_height = Inches(11), Inches(8.5)
s.left_margin = s.right_margin = Inches(.65)
s.top_margin = s.bottom_margin = Inches(.55)
grid = s._sectPr.find(qn('w:docGrid'))
if grid is not None:
    s._sectPr.remove(grid)
for st in doc.styles:
    for element in list(st._element.iter(qn('w:pBdr'))):
        element.getparent().remove(element)
    for font in st._element.iter(qn('w:rFonts')):
        for attribute in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']:
            font.attrib.pop(qn('w:'+attribute), None)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    st = doc.styles[name]
    st.font.name = 'Arial'
    st.font.color.rgb = RGBColor(0, 0, 0)
    st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'PingFang SC')
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing = 1.12
doc.styles['Title'].font.size = Pt(19)
doc.styles['Title'].paragraph_format.space_after = Pt(4)
doc.styles['Heading 1'].font.size = Pt(13)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(10)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(5)

def para(text, style=None, size=None):
    p = doc.add_paragraph(text, style)
    snap = OxmlElement('w:snapToGrid')
    snap.set(qn('w:val'), '0')
    p._p.get_or_add_pPr().append(snap)
    if size:
        for run in p.runs:
            run.font.size = Pt(size)
    return p

para('StudySpace 组件候选表', 'Title')
para('StudySpace Кандидаты компонентов', size=14)
para('任务 1  Задание 1  |  SS-COMP-01  |  v1  |  2026-10-14  |  DRAFT', size=9)
para('识别 MVP 的 7 个内部组件和 1 个外部依赖，按课堂四列表格说明职责及需求来源。\nВыделены 7 внутренних компонентов и 1 внешняя зависимость MVP с ответственностью и источниками по четырём колонкам задания.', size=10.5)

headers = ['候选组件\nКандидат', '职责 一句话\nОтветственность одним предложением',
           '来源 需求或场景 ID\nИсточник ID требований или сценариев', '内部／外部\nВнутренний / внешний']
t = doc.add_table(rows=1, cols=4)
for cell, text in zip(t.rows[0].cells, headers):
    cell.text = text
hdr = OxmlElement('w:tblHeader')
t.rows[0]._tr.get_or_add_trPr().append(hdr)

for row in rows:
    cells = t.add_row().cells
    cells[0].text = f"{row['id']}\n{row['name']['zh-CN']}\n{row['name']['ru-RU']}"
    cells[1].text = row['responsibility']['zh-CN']+'\n'+row['responsibility']['ru-RU']
    cells[2].text = row['display_requirements']+'\n'+row['display_scenarios']
    cells[3].text = '内部\nВнутренний' if row['boundary']=='INTERNAL' else '外部\nВнешний\nOPEN-003'

widths = column_widths_from_weights([18, 43, 28, 11], section_content_width_dxa(s))
apply_table_geometry(t, widths, cell_margins_dxa={'top': 90, 'bottom': 90, 'start': 110, 'end': 110})
pr = t._tbl.tblPr
borders = OxmlElement('w:tblBorders')
for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
    el = OxmlElement('w:'+edge)
    for key, value in [('val', 'single'), ('sz', '5'), ('color', 'D9D9D9')]:
        el.set(qn('w:'+key), value)
    borders.append(el)
pr.append(borders)
for i, row in enumerate(t.rows):
    trpr = row._tr.get_or_add_trPr()
    trpr.append(OxmlElement('w:cantSplit'))
    for j, cell in enumerate(row.cells):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shading = OxmlElement('w:shd')
        shading.set(qn('w:fill'), 'DDEBF7' if i==0 else ('F4F7FA' if i%2==0 else 'FFFFFF'))
        cell._tc.get_or_add_tcPr().append(shading)
        for p in cell.paragraphs:
            snap = OxmlElement('w:snapToGrid')
            snap.set(qn('w:val'), '0')
            p._p.get_or_add_pPr().append(snap)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = Pt(14)
            p.paragraph_format.keep_with_next = False
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i==0 or j==3 else WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.name = 'Arial'
                run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'PingFang SC')
                run.font.size = Pt(10 if i==0 or j==2 else 11)
                run.font.bold = i==0
                run.font.color.rgb = RGBColor(0, 0, 0)

para('边界与范围说明  Границы и объём', 'Heading 1').paragraph_format.page_break_before = True
for item in notes:
    para(item['zh-CN']+'\n'+item['ru-RU'], size=10.5)
para('来源  Источники', 'Heading 1')
para('课堂需求登记 v0.2 requirements-practice.json；用户故事 backlog-v1.json；系统设计 0.3.0 DOC-04；DEC-016；OPEN-003。\nУчебный реестр v0.2 requirements-practice.json; истории backlog-v1.json; проект системы 0.3.0 DOC-04; DEC-016; OPEN-003.', size=10)
doc.core_properties.title = 'StudySpace 组件候选表 Кандидаты компонентов'
doc.core_properties.subject = '任务 1 四列组件候选表 Задание 1 Кандидаты компонентов'
doc.core_properties.author = 'StudySpace'
doc.core_properties.language = 'zh-CN; ru-RU'
doc.save(DEST)

md = ['# StudySpace 组件候选表', '# StudySpace Кандидаты компонентов',
      '任务 1 / Задание 1 · SS-COMP-01 · v1 · 2026-10-14 · DRAFT', '',
      '| 候选组件 / Кандидат | 职责 一句话 / Ответственность одним предложением | 来源 需求或场景 ID / Источник ID требований или сценариев | 内部或外部 / Внутренний или внешний |',
      '| --- | --- | --- | --- |']
for row in rows:
    md.append('| '+' | '.join([
        row['id']+'<br>'+row['name']['zh-CN']+'<br>'+row['name']['ru-RU'],
        row['responsibility']['zh-CN']+'<br>'+row['responsibility']['ru-RU'],
        ', '.join(row['requirement_ids'])+'<br>'+', '.join(row['scenario_ids']),
        '内部 / Внутренний' if row['boundary']=='INTERNAL' else '外部 / Внешний — OPEN-003'
    ])+' |')
md += ['', '## 边界与范围说明 Границы и объём', '']
for n in notes:
    md.extend([n['zh-CN'], n['ru-RU'], ''])
md += ['## 来源 Источники', '', str(REQ_SOURCE.relative_to(ROOT)), str(STORY_SOURCE.relative_to(ROOT)), 'docs/zh-CN/04_System_Design.md · docs/ru-RU/04_System_Design.md · DEC-016 · OPEN-003', '']
(OUT/(NAME+'.md')).write_text('\n'.join(md))
(OUT/'README.md').write_text(f'''# StudySpace 组件候选课堂作业 Учебная работа о компонентах

2026-10-14 · v1 · SS-COMP-01 · DRAFT

交付 / Результат [{NAME}.docx]({NAME}.docx)

按课件任务1的四列填写8个候选，含7个内部组件和1个外部核验来源，中文与俄语在同一单元格配对。
В четырёх колонках задания1 заполнены8 кандидатов:7 внутренних компонентов и1 внешний источник проверки; китайский и русский тексты объединены в каждой ячейке.

需求与故事沿用现有SP及US/EN编号，外部前提OPEN-003未闭合；候选状态PROPOSED，业务测试NOT_RUN。
Сохранены ID требований SP и историй US/EN; условие OPEN-003 открыто; кандидаты PROPOSED, бизнес-тесты NOT_RUN.

数据 / Данные data/component-candidates-v1.json
验证 / Проверка data/verification.json
''')

checks = []
def check(name, condition, detail=None):
    checks.append(dict(name=name, passed=bool(condition), detail=detail))

check('eight_unique_candidates', len(rows)==8 and len({r['id'] for r in rows})==8)
check('seven_internal_one_external', sum(r['boundary']=='INTERNAL' for r in rows)==7 and sum(r['boundary']=='EXTERNAL' for r in rows)==1)
check('bilingual_names_and_responsibilities', all(set(r['name'])=={'zh-CN','ru-RU'} and set(r['responsibility'])=={'zh-CN','ru-RU'} for r in rows))
check('one_sentence_per_language', all(r['responsibility']['zh-CN'].count('。')==1 and r['responsibility']['ru-RU'].count('.')==1 for r in rows))
check('all_requirement_ids_exist', linked <= REQ.keys())
check('all_scenario_ids_exist', all(set(r['scenario_ids']) <= STORY.keys() for r in rows))
check('all_mvp_functional_requirements_covered', all_fr <= linked, {'covered': len(all_fr & linked), 'total': len(all_fr), 'missing': sorted(all_fr-linked)})
check('future_requirements_not_in_mvp_candidates', all(REQ[rid]['scope']=='MVP' for rid in linked))
check('external_gate_preserved', rows[-1]['boundary']=='EXTERNAL' and 'OPEN-003' in rows[-1]['display_scenarios'] and any('不能批准 VERIFIED' in n['zh-CN'] for n in notes))
check('nonoverlapping_multiple_bookings_preserved', any('多条非重叠预约允许' in n['zh-CN'] and 'несколько непересекающихся' in n['ru-RU'] for n in notes))
check('editable_four_column_table', len(doc.tables)==1 and len(t.rows)==9 and all(len(r.cells)==4 for r in t.rows))
check('exact_table_geometry', audit_docx_tables(DEST)==0)
check('all_candidate_text_in_docx', all(r['name'][lang] in '\n'.join(c.text for row in t.rows for c in row.cells) and r['responsibility'][lang] in '\n'.join(c.text for row in t.rows for c in row.cells) for r in rows for lang in ['zh-CN','ru-RU']))
check('no_source_or_prior_document_modified', all(sha(ROOT/p)==h for p,h in before.items()), before)
with ZipFile(DEST) as z:
    xml = z.read('word/document.xml').decode()
    check('repeated_header_and_unsplit_rows', '<w:tblHeader' in xml and xml.count('<w:cantSplit')==9)
    check('no_header_footer_or_page_numbers', not any(re.match(r'word/(header|footer)\d+\.xml$', name) for name in z.namelist()) and 'PAGE' not in xml)
check('draft_status_and_no_business_pass_claim', metadata['status']=='DRAFT' and metadata['component_status']=='PROPOSED' and metadata['business_test_status']=='NOT_RUN')

report = dict(date='2026-10-14', runtime='Codex bundled runtime 26.1013.11914',
              working_directory=str(ROOT), document=str(DEST), sha256=sha(DEST), checks=checks,
              all_pass=all(c['passed'] for c in checks), check_count=len(checks),
              visual_review='PENDING', business_tests='NOT_RUN')
(DATA/'verification.json').write_text(json.dumps(report, ensure_ascii=False, indent=2)+'\n')
print(json.dumps({'document': str(DEST), 'checks': len(checks), 'all_pass': report['all_pass'], 'mvp_fr_covered': len(all_fr), 'failed': [c['name'] for c in checks if not c['passed']]}, ensure_ascii=False))
if not report['all_pass']:
    raise SystemExit(1)
