import hashlib
import json
from pathlib import Path
from zipfile import ZipFile

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "coursework/2026-10-14-product-backlog-v1"
p = json.loads((OUT / "data/backlog-v1.json").read_text())
checks = []


def ck(name, value):
    checks.append(dict(name=name, passed=bool(value)))


files = [d["path"] for d in p["documents"]]
expected_count = 4 if "classroom_task4" in p else 3
ck(
    "files in requested consecutive order",
    len(files) == expected_count
    and [d["order"] for d in p["documents"]] == list(range(1, expected_count + 1))
    and all(name.startswith(f"{i:02}_") for i, name in enumerate(files, 1)),
)
ck(
    "active Word deliverables match registry",
    {f.name for f in OUT.glob("*.docx")} == set(files),
)
ck(
    "retired documents and Markdown removed",
    all(
        not (OUT / d["path"]).exists() and not (OUT / d["path"]).with_suffix(".md").exists()
        for d in p["retired_documents"]
    ),
)
ck(
    "old retained names removed",
    all(not (OUT / d["old"]).exists() for d in p["preserved_files"]),
)
historical_hash_matches = {
    d["new"]: hashlib.sha256((OUT / d["new"]).read_bytes()).hexdigest() == d["sha256"]
    for d in p["preserved_files"]
}
if "classroom_task4" in p:
    # The task-4 snapshot records the current user-edited files, not the earlier generation.
    current_baseline = json.loads((OUT / p["classroom_task4"]["source"]).read_text())[
        "preserved_docx"
    ]
    ck(
        "01 to 03 byte identical to task4 starting snapshot",
        all(
            hashlib.sha256((OUT / f).read_bytes()).hexdigest() == h
            for f, h in current_baseline.items()
        ),
    )
else:
    ck("retained documents byte identical", all(historical_hash_matches.values()))
ck(
    "manifest matches active files",
    (OUT / "manifest.txt").read_text().splitlines() == files,
)
ck(
    "stable document IDs retained",
    [d["id"] for d in p["documents"]][:3] == ["SS-PB-02", "SS-PB-03", "SS-PB-05"]
    and len({d["id"] for d in p["documents"]}) == len(files),
)
req = json.loads(
    (ROOT / "coursework/2026-10-13-requirements-v0.2/data/requirements-practice.json").read_text()
)
req_ids = {r["id"] for r in req["requirements"]}
legacy = {s["id"]: s for s in req["stories"]}
records = p["classroom_task3"]
ck(
    "six stable core story IDs",
    [s["id"] for s in records] == [f"US-{i:03}" for i in range(1, 7)],
)
ck(
    "original bilingual story text preserved",
    all(s["story"] == legacy[s["id"]]["story"] for s in records),
)
ck(
    "all stories have valid requirement links",
    all(s["requirement_ids"] and set(s["requirement_ids"]) <= req_ids for s in records),
)
ck(
    "all stories have valid epic links",
    all(s["epic_id"] in {e["id"] for e in p["epics"]} for s in records),
)
ck(
    "one completion sentence in each language",
    all(
        s["done"]["zh-CN"].count("。") == 1 and s["done"]["ru-RU"].count(".") == 1 for s in records
    ),
)
ck("all core stories remain MVP", all(s["release"] == "MVP" for s in records))
ck(
    "engineering baselines unchanged",
    all(
        hashlib.sha256((ROOT / f).read_bytes()).hexdigest() == h
        for f, h in p["source_fingerprints"].items()
    ),
)
d = Document(OUT / files[2])
tables = d.tables
ck(
    "editable five-column Chinese and Russian tables",
    len(tables) == 2 and all(len(t.columns) == 5 and len(t.rows) == 7 for t in tables),
)
ck(
    "same six row IDs in both languages",
    all([row.cells[0].text for row in t.rows[1:]] == [s["id"] for s in records] for t in tables),
)
ck(
    "same requirement links in both languages",
    all(tables[0].rows[i].cells[3].text == tables[1].rows[i].cells[3].text for i in range(1, 7)),
)
ck(
    "table content matches source stories",
    all(
        t.rows[i].cells[1].text == records[i - 1]["story"][lang]
        and t.rows[i].cells[4].text == records[i - 1]["done"][lang]
        for t, lang in zip(tables, ["zh-CN", "ru-RU"], strict=False)
        for i in range(1, 7)
    ),
)
alltext = (
    "\n".join(a.text for a in d.paragraphs)
    + "\n"
    + "\n".join(c.text for t in tables for r in t.rows for c in r.cells)
)
status_text = alltext
if "classroom_task4" in p:
    # The user removed explanatory paragraphs from document 03 before task 4.
    # Keep that file intact; check the status disclosure in the active packet instead.
    status_text += "\n" + (OUT / "README.md").read_text()
ck(
    "draft and unrun business status disclosed in active packet",
    "DRAFT" in status_text and "NOT_RUN" in status_text and "SS-PB-05" in alltext,
)
with ZipFile(OUT / files[2]) as z:
    xml = z.read("word/document.xml").decode()
    styles = z.read("word/styles.xml").decode()
    ck(
        "repeat table headers and no fixed row heights",
        xml.count("<w:tblHeader") == 2 and "<w:trHeight" not in xml,
    )
    ck("no title style borders", "<w:pBdr" not in styles)
ck(
    "no obsolete delivery links in README",
    not any(s["path"] in (OUT / "README.md").read_text() for s in p["retired_documents"]),
)
result = dict(
    date="2026-10-14",
    environment=str(ROOT),
    layout_version=2,
    all_pass=all(x["passed"] for x in checks),
    check_count=len(checks),
    checks=checks,
    failed=[x["name"] for x in checks if not x["passed"]],
    historical_generation_hash_matches=historical_hash_matches,
    visual_review="PENDING",
)
(OUT / "data/verification.json").write_text(json.dumps(result, ensure_ascii=False, indent=2))
print(
    json.dumps(
        dict(
            all_pass=result["all_pass"],
            check_count=len(checks),
            failed=result["failed"],
            files=files,
        ),
        ensure_ascii=False,
        indent=2,
    )
)
raise SystemExit(0 if result["all_pass"] else 1)
