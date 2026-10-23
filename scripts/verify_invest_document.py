"""Validate the classroom table, bilingual inventory and requirement preservation."""
from pathlib import Path
import hashlib
import json
from zipfile import ZipFile
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT/'coursework/2026-10-14-product-backlog-v1'
DATA = OUT/'data'
p = json.loads((DATA/'backlog-v1.json').read_text())
i = json.loads((DATA/'invest-v1.json').read_text())
doc = Document(OUT/i['path'])
text = '\n'.join(a.text for a in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
checks = []
def ck(name, value): checks.append({'name': name, 'passed': bool(value)})
sha = lambda f: hashlib.sha256(f.read_bytes()).hexdigest()
ck('01 to 03 documents byte identical', all(sha(OUT/f)==h for f,h in i['preserved_docx'].items()))
ck('baseline backlog story items unchanged', hashlib.sha256(json.dumps(p['items'],ensure_ascii=False,sort_keys=True).encode()).hexdigest()==i['baseline_story_items_sha256'])
ck('four documents in consecutive manifest order', [d['order'] for d in p['documents']]==[1,2,3,4] and (OUT/'manifest.txt').read_text().splitlines()==[d['path'] for d in p['documents']])
ck('five stable audit parent IDs', [r['id'] for r in i['rows']]==[f'US-{n:03}' for n in range(1,6)])
ck('all six judgments in every row', all(set(r['grades'])==set('INVEST') and set(r['grades'].values())<=set('PCF') for r in i['rows']))
ck('slide ID mapping is explicit', [r['slide_id'] for r in i['rows']]==[f'US-{n:02}' for n in range(1,6)])
ck('editable nine-column table matches slide', len(doc.tables[0].columns)==9 and len(doc.tables[0].rows)==6)
ck('table cell judgments match inventory', all([c.text for c in doc.tables[0].rows[n+1].cells[:7]]==[r['id'],*[r['grades'][k] for k in 'INVEST']] for n,r in enumerate(i['rows'])))
ck('bilingual corrections and result cells', all(doc.tables[0].rows[n+1].cells[7].text==r['fix']['zh-CN']+'\n'+r['fix']['ru-RU'] and doc.tables[0].rows[n+1].cells[8].text==r['result_display'] for n,r in enumerate(i['rows'])))
originals={s['id']:s for s in p['classroom_task3']}
ck('baseline originals preserved and rationale bilingual', all(r['original_story']==originals[r['id']]['story'] and all(r['rationale'][lang] in text for lang in i['languages']) for r in i['rows']))
ck('eleven unique refinement and split candidates', len(i['rewritten_stories'])==11 and len({s['id'] for s in i['rewritten_stories']})==11)
ck('every proposed story and acceptance is bilingual and present', all(s[f][lang] and s[f][lang] in text for s in i['rewritten_stories'] for f in ['story','acceptance','title'] for lang in i['languages']))
req = json.loads((ROOT/'coursework/2026-10-13-requirements-v0.2/data/requirements-practice.json').read_text())
reqids = {r['id'] for r in req['requirements']}
ck('candidate requirement IDs are valid', all(set(s['requirement_ids'])<=reqids and s['requirement_ids'] for s in i['rewritten_stories']))
items = {s['id']: s for s in p['items']}
coverage = {}
for row in i['rows']:
    links = {r for s in i['rewritten_stories'] if s['parent_id']==row['id'] for r in s['requirement_ids']}
    links |= {r for s in i['support_tasks'] if s['parent_id']==row['id'] for r in s['requirement_ids']}
    if row['id']=='US-004': links |= set(items['US-007']['requirement_ids']) | set(items['EN-001']['requirement_ids'])
    coverage[row['id']] = set(row['requirement_ids']) <= links
ck('no parent requirement lost by splitting', all(coverage.values()))
ck('draft proposed and unrun status retained', i['status']=='DRAFT' and all(s['status']=='PROPOSED' and s['implementation_status']=='NOT_IMPLEMENTED' and s['test_status']=='NOT_RUN' and s['estimate_state']=='NOT_ESTIMATED' for s in i['rewritten_stories']) and all(x in text for x in ['DRAFT','PROPOSED','NOT_IMPLEMENTED','NOT_RUN']))
ck('verification source gate retained bilingually', 'OPEN-003' in text and '真实审核必须' in text and 'реальная проверка требует' in text)
ck('multiple nonoverlapping bookings policy retained', '相邻及多条非重叠预约允许' in text and 'несколько непересекающихся записей разрешены' in text)
ck('atomic deactivation and rollback preserved', all(term in text for term in ['同一事务','失败全部回滚','重新启用','одной транзакции','откатывает всё','не восстанавливает']))
ck('six criteria meanings and rules present bilingually', len(i['criteria'])==6 and all(c[f][lang] in text for c in i['criteria'] for f in ['meaning','rule'] for lang in i['languages']))
ck('04 registry and source inventory aligned', p['documents'][-1]['id']==i['doc_id'] and p['documents'][-1]['path']==i['path'] and p['classroom_task4']['rewritten_ids']==[s['id'] for s in i['rewritten_stories']])
with ZipFile(OUT/i['path']) as z:
    xml = z.read('word/document.xml').decode(); styles = z.read('word/styles.xml').decode()
    ck('repeat headers and expandable rows', xml.count('<w:tblHeader')==len(doc.tables) and '<w:trHeight' not in xml)
    ck('title has no decorative border', '<w:pBdr' not in styles)
ck('no internal citation tokens or placeholders', not any(s in text for s in ['oai-mem-citation','turn0search','TODO','lorem ipsum']))
result = {'date': '2026-10-14', 'environment': str(ROOT), 'all_pass': all(x['passed'] for x in checks), 'check_count': len(checks), 'checks': checks, 'failed': [x['name'] for x in checks if not x['passed']], 'requirement_coverage_by_parent': coverage, 'visual_review': 'PENDING', 'business_test_status': 'NOT_RUN', 'output_sha256': sha(OUT/i['path'])}
(DATA/'invest-verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2))
print(json.dumps({k:result[k] for k in ['all_pass','check_count','failed']},ensure_ascii=False))
raise SystemExit(0 if result['all_pass'] else 1)
