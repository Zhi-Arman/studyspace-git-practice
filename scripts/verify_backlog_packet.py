import hashlib
import json
from collections import Counter
from pathlib import Path
from zipfile import ZipFile

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "coursework/2026-10-14-product-backlog-v1"
p = json.loads((OUT / "data/backlog-v1.json").read_text())
if p.get("layout_version") == 2:
    import subprocess
    import sys

    raise SystemExit(
        subprocess.call([sys.executable, str(ROOT / "scripts/verify_classroom_packet.py")])
    )
source = json.loads(
    (ROOT / "coursework/2026-10-13-requirements-v0.2/data/requirements-practice.json").read_text()
)
eng = json.loads((ROOT / "docs/data/requirements.json").read_text())
checks = []


def ck(name, ok, detail=""):
    checks.append(dict(name=name, pass_=bool(ok), detail=detail))


items = {x["id"]: x for x in p["items"]}
reqs = {x["id"]: x for x in source["requirements"]}
ck(
    "four separate DOCX outputs",
    len(p["documents"]) == 4 and all((OUT / x["path"]).exists() for x in p["documents"]),
)
ck(
    "22 unique ordered backlog items",
    len(items) == 22 and sorted(x["order"] for x in items.values()) == list(range(1, 23)),
)
ck(
    "15 user stories and 7 enablers",
    Counter(x["kind"] for x in items.values()) == Counter(USER_STORY=15, ENABLER=7),
)
ck(
    "7 epics each item assigned exactly once",
    len(p["epics"]) == 7
    and Counter(i for e in p["epics"] for i in e["item_ids"]) == Counter(items.keys()),
)
ck(
    "all 60 candidate IDs covered",
    len(p["coverage"]) == 60
    and {r["requirement_id"] for r in p["coverage"]} == set(reqs)
    and all(r["item_ids"] for r in p["coverage"]),
)
ck(
    "coverage links point to real supporting items",
    all(
        all(r["requirement_id"] in items[i]["requirement_ids"] for i in r["item_ids"])
        for r in p["coverage"]
    ),
)
ck(
    "all 38 engineering requirements preserved",
    len(eng["requirements"]) == 38
    and {x["id"] for x in eng["requirements"]}
    <= {i for r in p["coverage"] for i in r["legacy_ids"]},
)
ck(
    "stable original US001 to US006 text IDs retained",
    all(
        items[x["id"]]["story"] == x["story"]
        and items[x["id"]]["role"] == x["role"]
        and items[x["id"]]["acceptance"] == x["acceptance"]
        for x in source["stories"]
    ),
)
ck(
    "dependencies exist and precede dependents",
    all(
        d in items and items[d]["order"] < x["order"]
        for x in items.values()
        for d in x["dependencies"]
    ),
)
ck(
    "18 MVP and 4 future items",
    Counter(x["release"] for x in items.values()) == Counter(MVP=18, R1=2, R2=2),
)
ck(
    "future source requirements excluded from MVP",
    all(
        (reqs[r]["scope"] == "FUTURE") == (x["release"] != "MVP")
        for x in items.values()
        for r in x["requirement_ids"]
    ),
)
ck(
    "56 MVP candidate requirements retained",
    sum(reqs[r["requirement_id"]]["scope"] == "MVP" for r in p["coverage"]) == 56,
)
ck(
    "MVP seat verification and concurrency retained",
    all(
        any(r in x["requirement_ids"] and x["release"] == "MVP" for x in items.values())
        for r in ["SP-C-001", "SP-FR-008", "SP-FR-022", "SP-IF-002", "SP-NFR-004"]
    ),
)
ck(
    "all items bilingual with three acceptance conditions",
    all(
        all(
            x[f].get(language)
            for language in ["zh-CN", "ru-RU"]
            for f in ["title", "role", "story", "note"]
        )
        and len(x["acceptance"]) == 3
        and all(a.get(language) for a in x["acceptance"] for language in ["zh-CN", "ru-RU"])
        for x in items.values()
    ),
)
ck(
    "three requirement-based personas valid",
    len(p["personas"]) == 3
    and all(
        x["basis"] == "REQUIREMENTS_BASED_HYPOTHESIS_NOT_INTERVIEW"
        and set(x["reqs"]) <= set(reqs)
        and set(x["item_ids"]) <= set(items)
        for x in p["personas"]
    ),
)
ck(
    "story maps cover all user stories",
    set(items) - {x["id"] for x in items.values() if x["kind"] == "ENABLER"}
    == {
        s
        for m in ["student_map", "admin_map"]
        for a in p[m]
        for rel in ["MVP", "R1", "R2"]
        for s in a[rel]
    }
    | {"US-015"},
)
ck(
    "source baseline files unchanged",
    all(
        hashlib.sha256((ROOT / f).read_bytes()).hexdigest() == h
        for f, h in p["source_fingerprints"].items()
    ),
)
ck(
    "no fictional approval test or implementation status",
    all(
        x["status"] == "DRAFT"
        and x["implementation_status"] == "NOT_IMPLEMENTED"
        and x["test_status"] == "NOT_RUN"
        and x["estimate_state"] == "PROVISIONAL"
        for x in items.values()
    ),
)
ck(
    "external gates preserved",
    all(
        q in {a for x in items.values() for a in x["external_gates"]}
        for q in ["OPEN-003", "OPEN-008", "OPEN-010"]
    ),
)
ck(
    "provisional estimate scale",
    all(x["estimate"] in [3, 5, 8] for x in items.values()),
)
for d in p["documents"]:
    f = OUT / d["path"]
    doc = Document(f)
    text = (
        "\n".join(a.text for a in doc.paragraphs)
        + "\n"
        + "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    )
    ck(
        d["id"] + " bilingual metadata",
        d["id"] in text
        and "DRAFT" in text
        and "2026-10-14" in text
        and any("\u4e00" <= c <= "\u9fff" for c in text)
        and any("\u0400" <= c <= "\u04ff" for c in text),
    )
    with ZipFile(f) as z:
        xml = z.read("word/document.xml").decode()
        ck(
            d["id"] + " editable tables and repeat headers",
            bool(doc.tables) and xml.count("<w:tblHeader") == len(doc.tables),
        )
    ck(
        d["id"] + " no fixed row heights",
        all("<w:trHeight" not in r._tr.xml for t in doc.tables for r in t.rows),
    )
    ck(
        d["id"] + " no citation tokens",
        not any(a in text for a in ["turn0search", "oai-mem-citation", "TODO", "lorem ipsum"]),
    )

result = dict(
    environment=str(ROOT),
    date="2026-10-14",
    checks=checks,
    all_pass=all(x["pass_"] for x in checks),
    check_count=len(checks),
    failed=[x["name"] for x in checks if not x["pass_"]],
    counts=dict(
        requirements=len(reqs),
        engineering=len(eng["requirements"]),
        items=len(items),
        personas=len(p["personas"]),
    ),
    story_points=dict(
        Counter(
            {
                rel: sum(x["estimate"] for x in items.values() if x["release"] == rel)
                for rel in ["MVP", "R1", "R2"]
            }
        )
    ),
    visual_review="PENDING",
)
(OUT / "data/verification.json").write_text(json.dumps(result, ensure_ascii=False, indent=2))
print(
    json.dumps(
        {k: result[k] for k in ["all_pass", "check_count", "failed", "counts", "story_points"]},
        ensure_ascii=False,
        indent=2,
    )
)
raise SystemExit(0 if result["all_pass"] else 1)
