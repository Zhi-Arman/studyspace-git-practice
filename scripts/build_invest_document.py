"""Create the next numbered classroom document without rewriting tasks 1 to 3."""
from pathlib import Path
import hashlib
import json
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'coursework/2026-10-14-product-backlog-v1'
DATA = OUT / 'data'
FILENAME = '04_StudySpace_INVEST检查与故事改写_中俄双语_v1.docx'
DEST = OUT / FILENAME
if DEST.exists():
    raise SystemExit('Document 04 already exists. Inspect it before regenerating.')
packet_path = DATA / 'backlog-v1.json'
packet = json.loads(packet_path.read_text())
originals = {s['id']: s for s in packet['classroom_task3']}
sha = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()
bi = lambda zh, ru: {'zh-CN': zh, 'ru-RU': ru}

inventory = {
    'doc_id': 'SS-PB-06', 'order': 4, 'version': 'v1', 'date': '2026-10-14',
    'status': 'DRAFT', 'implementation_status': 'NOT_IMPLEMENTED', 'test_status': 'NOT_RUN',
    'languages': ['zh-CN', 'ru-RU'],
    'path': FILENAME,
    'source_image': '/var/folders/ff/kqw394c93dd7krl0z2hcqqdh0000gn/T/codex-clipboard-6a5c1b44-346e-440f-bce1-b82a440da359.png',
    'method_source': {
        'author': 'Bill Wake', 'title': 'INVEST in Good Stories, and SMART Tasks',
        'date': '2003-08-17', 'accessed': '2026-10-14',
        'url': 'https://xp123.com/invest-in-good-stories-and-smart-tasks/',
        'verification': 'gstack browse HTTP 200; page text inspected; no console errors'
    },
    'preserved_docx': {d['path']: sha(OUT/d['path']) for d in packet['documents']},
    'baseline_story_items_sha256': hashlib.sha256(json.dumps(packet['items'], ensure_ascii=False, sort_keys=True).encode()).hexdigest(),
    'source_packet_sha256_before_append': sha(packet_path),
    'grading': {
        'P': bi('通过', 'Соответствует'),
        'C': bi('需澄清', 'Уточнить'),
        'F': bi('需修改', 'Изменить')
    },
    'criteria': [
        {'code': 'I', 'english': 'Independent', 'meaning': bi('独立性', 'Независимость'), 'rule': bi('业务目标不重叠，依赖明确，可用预置数据单独验证。依赖仍可能限制交付顺序。', 'Цели не дублируются, зависимости явны; историю можно проверить с заранее подготовленными данными. Зависимости могут ограничивать порядок поставки.')},
        {'code': 'N', 'english': 'Negotiable', 'meaning': bi('可协商性', 'Обсуждаемость'), 'rule': bi('保留用户目标，界面与实现细节可协商；已确认的权限和业务约束继续适用。', 'Цель пользователя сохраняется, интерфейс и реализация обсуждаются; согласованные права и бизнес-ограничения остаются обязательными.')},
        {'code': 'V', 'english': 'Valuable', 'meaning': bi('用户价值', 'Ценность для пользователя'), 'rule': bi('产生学生或管理员能够观察到的业务结果；按业务结果拆分，覆盖所需各层。', 'Есть наблюдаемый результат для студента или администратора; разбиение идёт по бизнес-результату и включает необходимые слои.')},
        {'code': 'E', 'english': 'Estimable', 'meaning': bi('可估算性', 'Оцениваемость'), 'rule': bi('范围、输入和异常足够明确，团队可以讨论估算；P不表示团队已完成估算。', 'Объём, входы и ошибки понятны для обсуждения оценки командой; P не означает, что команда уже дала оценку.')},
        {'code': 'S', 'english': 'Small', 'meaning': bi('足够小', 'Небольшой размер'), 'rule': bi('围绕一个可验收结果；是否适合一次迭代仍需团队按实际能力确认。', 'Один принимаемый результат; пригодность для одной итерации команда подтверждает с учётом своих возможностей.')},
        {'code': 'T', 'english': 'Testable', 'meaning': bi('可测试性', 'Проверяемость'), 'rule': bi('有明确的正常、异常或边界结果；可测试不代表测试已执行或通过。', 'Заданы результаты нормы, ошибок или границ; проверяемость не означает выполнение или успех тестов.')}
    ],
    'rows': [], 'rewritten_stories': [], 'support_tasks': []
}

def audit(parent, grades, fix, result, rationale):
    inventory['rows'].append({
        'id': parent, 'slide_id': 'US-'+parent[-2:],
        'original_story': originals[parent]['story'],
        'requirement_ids': originals[parent]['requirement_ids'],
        'grades': dict(zip('INVEST', grades)), 'fix': bi(*fix),
        'result': bi(*result), 'rationale': bi(*rationale)
    })

audit('US-001', 'CPPPPP',
      ('明确位置、日期、区间；预置资源与占用记录，独立验证查询。', 'Уточнить расположение, дату и интервал; проверять поиск с подготовленными ресурсами и записями.'),
      ('US-001-R\n按位置和时段查询可预约座位', 'US-001-R\nНайти доступные места по расположению и времени'),
      ('I需澄清：依赖资格和资源配置，但查询本身不应创建预约。E、S、T通过：单一查询目标及已有排除条件足够清晰。', 'I требует уточнения: есть зависимости от допуска и ресурсов, но поиск не создаёт бронь. E, S, T соответствуют: один результат поиска и ясные условия исключения.'))
audit('US-002', 'CPPCPP',
      ('直接按座位与区间预约；补充重叠、边界及并发结果，复核估算。', 'Бронировать по месту и интервалу; уточнить пересечения, границы и гонку, пересмотреть оценку.'),
      ('US-002-R\n创建一条单座位预约', 'US-002-R\nСоздать одно бронирование одного места'),
      ('I需澄清：依赖账号与资源；查询界面不必成为预约接口的强制前置步骤。E需澄清：事务和并发验证影响工作量。S、T通过：仍是一个预约结果，已有可观察边界。', 'I требует уточнения: нужны аккаунт и ресурс, но экран поиска не обязан быть шагом перед вызовом бронирования. E требует уточнения из-за транзакций и гонки. S, T соответствуют: один результат и проверяемые границы.'))
audit('US-003', 'CPPPPP',
      ('限定本人、开始前；明确开始时刻拒绝与重复取消不重复计数。', 'Ограничить своей записью до начала; отклонять в момент начала, повторно не увеличивать счётчик.'),
      ('US-003-R\n开始前取消本人预约', 'US-003-R\nОтменить свою запись до начала'),
      ('I需澄清：已有预约是输入，可预置记录单独验证。E、S、T通过：用户、时间边界及取消结果已明确。管理员取消仍由US-010承接。', 'I требует уточнения: существующая бронь — вход, который можно подготовить заранее. E, S, T соответствуют: ясны владелец, граница времени и результат. Отмена администратором остаётся в US-010.'))
audit('US-004', 'FPPCFC',
      ('拆开审核、停用、身份变更、角色维护；OPEN-003闭合前不能真实通过审核。', 'Разделить проверку, отключение, смену идентификаторов и ролей; без OPEN-003 нет реального подтверждения.'),
      ('US-004-A 审核学生\nUS-004-B 停用账号\nUS-004-C 变更后重新核验\nUS-004-D 维护角色', 'US-004-A Проверить студента\nUS-004-B Отключить аккаунт\nUS-004-C Повторная проверка\nUS-004-D Назначить роль'),
      ('I、S需修改：身份审核与多种账号生命周期操作混合。E、T需澄清：OPEN-003的学校来源尚未落实；状态规则可用合成数据验证，真实资格核验仍需来源和授权。', 'I, S требуют изменения: смешаны проверка и разные операции жизненного цикла аккаунта. E, T требуют уточнения: источник в OPEN-003 не определён; правила состояния проверяются на синтетических данных, реальная проверка требует источника и разрешения.'))
audit('US-005', 'FPPCFP',
      ('拆开房间、座位、开放时间和可用性减少；停用与联动取消必须原子完成。', 'Разделить помещения, места, часы и уменьшение доступности; отключение и отмены должны быть атомарными.'),
      ('US-005-A 维护房间\nUS-005-B 维护座位\nUS-005-C 设置开放时间\nUS-005-D 安全减少可用性', 'US-005-A Вести помещения\nUS-005-B Вести места\nUS-005-C Настроить часы\nUS-005-D Уменьшить доступность'),
      ('I、S需修改：目录维护、日历配置和预约补偿有不同业务结果。E需澄清：整体工作量受组合场景影响。T通过：唯一编码、例外优先、原子取消及回滚均有可观察结果。', 'I, S требуют изменения: каталог, расписание и отмены дают разные результаты. E требует уточнения из-за сочетания сценариев. T соответствует: уникальность кодов, приоритет исключения, атомарная отмена и откат наблюдаемы.'))

result_display = {
    'US-001': 'US-001-R 查询空座 / Найти места',
    'US-002': 'US-002-R 创建预约 / Создать бронь',
    'US-003': 'US-003-R 本人取消 / Своя отмена',
    'US-004': 'US-004-A 审核 / Проверить\nUS-004-B 停用 / Отключить\nUS-004-C 重核 / Перепроверить\nUS-004-D 角色 / Назначить роль',
    'US-005': 'US-005-A 房间 / Помещения\nUS-005-B 座位 / Места\nUS-005-C 开放时间 / Часы\nUS-005-D 安全缩减 / Сократить доступ'
}
for row in inventory['rows']:
    row['result_display'] = result_display[row['id']]

def story(sid, parent, title, zh, ru, reqs, ac_zh, ac_ru):
    inventory['rewritten_stories'].append({
        'id': sid, 'parent_id': parent, 'title': bi(*title), 'story': bi(zh, ru),
        'requirement_ids': reqs, 'acceptance': bi(ac_zh, ac_ru),
        'status': 'PROPOSED', 'release': 'MVP', 'implementation_status': 'NOT_IMPLEMENTED',
        'test_status': 'NOT_RUN', 'estimate_state': 'NOT_ESTIMATED'
    })

story('US-001-R', 'US-001', ('查询可预约座位', 'Найти доступные места'),
      '作为已核验学生，我希望按位置、日期和起止时间查询可预约座位，以便选择适合的自习地点。',
      'Как подтверждённый студент, я хочу искать доступные места по расположению, дате и времени начала и конца, чтобы выбрать подходящее место для самостоятельной учёбы.',
      originals['US-001']['requirement_ids'],
      '给定启用资源及占用记录，只返回匹配筛选条件、开放且与目标区间不重叠的座位；相邻区间可用，停用资源被排除；查询不创建预约。',
      'При заданных активных ресурсах и записях возвращаются только подходящие открытые места без пересечения с интервалом; соседние интервалы допустимы, отключённые ресурсы исключены; поиск не создаёт бронь.')
story('US-002-R', 'US-002', ('创建单座位预约', 'Создать бронь одного места'),
      '作为启用且已核验的学生，我希望为一个座位提交起止时间并获得预约编号，以便确定我的学习安排。',
      'Как активный подтверждённый студент, я хочу указать место и время начала и конца и получить ID брони, чтобы зафиксировать свой план учёбы.',
      originals['US-002']['requirement_ids'],
      '合法提交只创建一条单座位记录并返回ID；同学生或同座位重叠拒绝，相邻及多条非重叠预约允许；同座位同区间并发仅一个请求成功，失败不留部分写入。',
      'Допустимая заявка создаёт одну запись одного места и возвращает ID; пересечения студента или места отклоняются, соседние и несколько непересекающихся записей разрешены; при гонке на одно место и интервал успешен один запрос, сбой не оставляет частичной записи.')
story('US-003-R', 'US-003', ('取消本人预约', 'Отменить свою бронь'),
      '作为预约所有者，我希望在预约开始前取消这条预约，以便把不再需要的时段释放给其他学生。',
      'Как владелец брони, я хочу отменить её до начала, чтобы освободить ненужный интервал для других студентов.',
      originals['US-003']['requirement_ids'],
      '仅本人且now < start时取消成功并释放时段；他人记录及now >= start拒绝；重复取消不重复计数，其他预约保持不变。管理员取消沿用US-010。',
      'Отмена освобождает интервал только для владельца при now < start; чужая запись и now >= start отклоняются; повтор не увеличивает счётчик, другие брони сохраняются. Административная отмена остаётся в US-010.')

story('US-004-A', 'US-004', ('审核学生身份', 'Проверить студента'),
      '作为管理员，我希望依据学校认可的可信来源审核一名待核验学生并保存决定，以便能够追溯该账号的校内资格。',
      'Как администратор, я хочу проверить ожидающего студента по доверенному источнику, признанному вузом, и сохранить решение, чтобы допуск аккаунта можно было проследить.',
      ['SP-BR-003', 'SP-DR-001', 'SP-FR-022', 'SP-IF-002', 'SP-NFR-003'],
      '真实可信来源有效时记录审核人、时间、决定和来源引用；按核验结果保存VERIFIED或REJECTED。来源未定或不可验证时不得通过；不保存证明影像。角色、启停与核验状态分别判断。',
      'При действительном доверенном источнике сохраняются проверяющий, время, решение и ссылка; результатом становится VERIFIED либо REJECTED. Неизвестный или непроверяемый источник не даёт подтверждение; изображения доказательств не хранятся. Роль, активность и проверка учитываются отдельно.')
story('US-004-B', 'US-004', ('停用学生账号', 'Отключить аккаунт студента'),
      '作为管理员，我希望停用不再允许使用系统的学生账号并取消其未结束预约，以便撤销该账号的预约资格。',
      'Как администратор, я хочу отключить аккаунт студента, которому больше нельзя пользоваться системой, и отменить его незавершённые брони, чтобы отозвать допуск к бронированию.',
      ['SP-BR-003', 'SP-FR-020', 'SP-NFR-003'],
      '停用和未结束预约取消在同一事务完成；并发预约不能绕过停用；失败全部回滚，其他账号的预约保持。',
      'Отключение и отмена незавершённых броней выполняются в одной транзакции; конкурентная бронь не обходит отключение; сбой откатывает всё, записи других аккаунтов сохраняются.')
story('US-004-C', 'US-004', ('身份变更后重新核验', 'Повторно проверить после смены идентификаторов'),
      '作为管理员，我希望学生身份字段变更后恢复待核验状态并取消其未结束预约，以便原来的审核结果不会被错误沿用。',
      'Как администратор, я хочу возвращать статус ожидания и отменять незавершённые брони после изменения идентификаторов студента, чтобы прежнее подтверждение не использовалось ошибочно.',
      ['SP-BR-003', 'SP-FR-021', 'SP-NFR-003'],
      '变更学号或身份字段时，修改、PENDING重置与未结束预约取消原子完成；原VERIFIED失效，重新核验前拒绝预约。',
      'Изменение номера или идентификаторов, возврат в PENDING и отмена незавершённых броней атомарны; прежний VERIFIED недействителен, до повторной проверки бронирование запрещено.')
story('US-004-D', 'US-004', ('维护账号角色', 'Назначить роль аккаунта'),
      '作为管理员，我希望按授权策略调整账号角色并保存审计记录，以便职责变化时访问权限保持正确。',
      'Как администратор, я хочу менять роль аккаунта по разрешённой политике и сохранять аудит, чтобы права доступа соответствовали изменению обязанностей.',
      ['SP-FR-004', 'SP-FR-031', 'SP-NFR-003'],
      '学生直接提交角色修改被拒绝；管理员的授权变更保存审计，且不自动改变启停或校内核验状态；未核验账号即使设为STUDENT仍不得预约。',
      'Запрос изменения роли студентом отклоняется; разрешённое изменение администратором аудируется и не меняет автоматически активность или проверку вуза; неподтверждённый STUDENT по-прежнему не бронирует.')
inventory['support_tasks'].append({
    'id': 'US-004-T1', 'parent_id': 'US-004', 'requirement_ids': ['SP-FR-026'],
    'text': bi('首位管理员通过显式交互CLI创建，不提供默认密码。作为支撑任务保留，不能包装为学生用户价值；注册需求SP-FR-001继续由US-007承担，通用授权SP-FR-004继续由EN-001承担。',
               'Первый администратор создаётся явной интерактивной CLI без исходного пароля. Это вспомогательная задача; регистрация SP-FR-001 остаётся в US-007, общее разграничение SP-FR-004 — в EN-001.'),
    'status': 'PROPOSED', 'test_status': 'NOT_RUN'
})

story('US-005-A', 'US-005', ('维护房间信息', 'Вести данные помещения'),
      '作为管理员，我希望维护房间的位置、名称和类型，以便学生能够按准确的信息寻找学习地点。',
      'Как администратор, я хочу вести расположение, название и тип помещения, чтобы студенты искали место по точным данным.',
      ['SP-FR-006', 'SP-FR-014', 'SP-DR-002'],
      '管理员更新的信息可被查询，类型使用CLASSROOM或STUDY_ROOM；学生修改拒绝，历史数据保留；任何停用操作必须同时满足US-005-D的原子规则。',
      'Изменения администратора видны в поиске, тип — CLASSROOM либо STUDY_ROOM; изменение студентом отклоняется, история сохраняется; любое отключение также выполняет атомарные правила US-005-D.')
story('US-005-B', 'US-005', ('维护座位编码', 'Вести коды мест'),
      '作为管理员，我希望维护房间内的座位编码，以便学生能够识别并预约具体座位。',
      'Как администратор, я хочу вести коды мест в помещении, чтобы студент мог распознать и забронировать конкретное место.',
      ['SP-FR-006', 'SP-FR-015', 'SP-DR-002'],
      '同房间重复编码拒绝，不同房间相同编码允许；容量等于启用座位数，学生修改拒绝；任何停用操作必须同时满足US-005-D的原子规则。',
      'Повтор кода внутри помещения запрещён, между помещениями допустим; ёмкость равна числу активных мест, изменение студентом отклоняется; любое отключение также выполняет атомарные правила US-005-D.')
story('US-005-C', 'US-005', ('设置开放时间', 'Настроить часы работы'),
      '作为管理员，我希望设置常规开放时间和指定日期例外，以便查询与预约使用实际开放窗口。',
      'Как администратор, я хочу задавать обычные часы и исключения на конкретные даты, чтобы поиск и бронирование использовали действующие интервалы работы.',
      ['SP-FR-006', 'SP-FR-016', 'SP-BR-004', 'SP-DR-002'],
      '日期例外覆盖常规时间，不合法窗口拒绝，查询使用新配置；任何缩短窗口的变更必须同时满足US-005-D，不能先缩减再另行取消预约。',
      'Исключение даты заменяет обычные часы, недопустимый интервал отклоняется, поиск использует новую настройку; любое сокращение также выполняет US-005-D, нельзя сначала сократить часы, а затем отдельно отменять брони.')
story('US-005-D', 'US-005', ('安全减少资源可用性', 'Безопасно уменьшить доступность ресурса'),
      '作为管理员，我希望停用房间或座位、缩短开放时间时同步取消受影响的未结束预约，以便学生不会保留不可使用的预约。',
      'Как администратор, я хочу при отключении помещения или места либо сокращении часов одновременно отменять затронутые незавершённые брони, чтобы у студентов не оставались недоступные бронирования.',
      ['SP-BR-004', 'SP-DR-002', 'SP-FR-014', 'SP-FR-015', 'SP-FR-016', 'SP-FR-017', 'SP-NFR-004'],
      '配置变更、受影响预约取消及原因记录在同一事务完成；未受影响记录保持，失败全部回滚；重新启用或恢复开放时间不恢复旧预约。分别覆盖停用房间、停用座位和缩短窗口。',
      'Изменение конфигурации, отмена затронутых броней и запись причины выполняются в одной транзакции; остальные записи сохраняются, сбой откатывает всё; включение или восстановление часов не восстанавливает прежние брони. Проверяются помещение, место и сокращение часов.')

(DATA/'invest-v1.json').write_text(json.dumps(inventory, ensure_ascii=False, indent=2))
doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Cm(29.7), Cm(21)
sec.top_margin, sec.bottom_margin = Cm(1.15), Cm(1.15)
sec.left_margin, sec.right_margin = Cm(1.6), Cm(1.6)
sec.footer_distance = Cm(.45)
for n in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    style = doc.styles[n]
    style.font.name, style.font.color.rgb = 'Arial', RGBColor(0, 0, 0)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for key, family in [('ascii', 'Arial'), ('hAnsi', 'Arial'), ('eastAsia', 'PingFang SC')]:
        fonts.set(qn('w:'+key), family)
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(4)
    snap = OxmlElement('w:snapToGrid'); snap.set(qn('w:val'), '0')
    style.element.get_or_add_pPr().append(snap)
for e in doc.styles.element.xpath('.//w:pBdr'):
    e.getparent().remove(e)
doc.styles['Normal'].font.size = Pt(10.5)
doc.styles['Normal'].paragraph_format.line_spacing = Pt(14)
doc.styles['Normal'].paragraph_format.space_after = Pt(2.5)
doc.styles['Title'].font.size = Pt(19)
doc.styles['Subtitle'].font.size = Pt(12)
doc.styles['Heading 1'].font.size = Pt(13)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(4)
doc.styles['Heading 2'].font.size = Pt(11)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(5)
md = []

def p(text, style=None, size=None):
    para = doc.add_paragraph(text, style)
    if size:
        for run in para.runs: run.font.size = Pt(size)
    md.append(text)
    return para

def pair(x, style=None, size=None):
    p(x['zh-CN'], style, size)
    p(x['ru-RU'], None if style == 'Heading 2' else style, size)

def page_heading(text):
    para = p(text, 'Heading 1')
    para.paragraph_format.page_break_before = True

def table(headers, rows, widths, font=10):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for col, w in zip(t.columns, widths): col.width = Cm(w)
    for c, text in zip(t.rows[0].cells, headers): c.text = text
    for row in rows:
        for c, text in zip(t.add_row().cells, row): c.text = text
    for ri, row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        if ri == 0: row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for ci, c in enumerate(row.cells):
            c.width = Cm(widths[ci]); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr = c._tc.get_or_add_tcPr()
            mar = OxmlElement('w:tcMar')
            for side in ['top', 'bottom', 'left', 'right']:
                e = OxmlElement('w:'+side); e.set(qn('w:w'), '65' if side in ['top', 'bottom'] else '85'); e.set(qn('w:type'), 'dxa'); mar.append(e)
            pr.append(mar)
            borders = OxmlElement('w:tcBorders')
            for side in ['top', 'bottom', 'left', 'right']:
                e = OxmlElement('w:'+side); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), 'D9D9D9'); borders.append(e)
            pr.append(borders)
            fill = OxmlElement('w:shd'); fill.set(qn('w:fill'), 'DCE8F1' if ri == 0 else ('F5F7FA' if ri % 2 == 0 else 'FFFFFF')); pr.append(fill)
            for para in c.paragraphs:
                para.paragraph_format.space_after = Pt(1)
                para.paragraph_format.line_spacing = Pt(12.5)
                if ci < 7 and len(headers) == 9: para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:
                    run.font.size = Pt(font); run.font.bold = ri == 0
    md.append('| ' + ' | '.join(x.replace('\n', '<br>') for x in headers) + ' |')
    md.append('| ' + ' | '.join('---' for _ in headers) + ' |')
    md.extend('| ' + ' | '.join(x.replace('\n', '<br>') for x in row) + ' |' for row in rows)
    return t

p('StudySpace INVEST检查与故事改写', 'Title')
p('StudySpace Проверка INVEST и уточнение историй', 'Subtitle')
p('排序 04 / Порядок 04  |  Doc-ID SS-PB-06  |  v1  |  2026-10-14  |  DRAFT', size=9)
p('任务4 INVEST检查  Задание 4 Проверка INVEST', 'Heading 1')
pair(bi('审核与资源维护需要拆分，查询、预约和取消需要明确依赖。下表评估第03份文档的US-001至US-005；课件US-01至US-05依次对应这些稳定ID。US-006不在本次五行范围内。',
        'Проверку и управление ресурсами нужно разделить; для поиска, бронирования и отмены уточнить зависимости. Оценены US-001…US-005 документа 03, соответствующие US-01…US-05 слайда. US-006 вне пяти строк этого задания.'), size=9.5)
p('P 通过 / соответствует    C 需澄清 / уточнить    F 需修改 / изменить', size=9)
table(['ID', 'I', 'N', 'V', 'E', 'S', 'T', '需要修改什么\nЧто исправить', '结果 新故事\nРезультат новые истории'],
      [[r['id'], *[r['grades'][k] for k in 'INVEST'], r['fix']['zh-CN']+'\n'+r['fix']['ru-RU'], r['result_display']] for r in inventory['rows']],
      [1.8, .7, .7, .7, .7, .7, .7, 10.0, 8.5], font=9.5)
pair(bi('N、V均通过：五条故事都有用户目标和可观察价值，未锁定界面与实现技术。E、S是需求层面的判断，团队估算及迭代容量仍需确认。P不代表功能实现或测试通过，业务测试NOT_RUN。',
        'N и V соответствуют: у всех пяти историй есть цель и наблюдаемая ценность, интерфейс и технология не зафиксированы. E и S оцениваются по требованиям; оценка команды и ёмкость итерации ещё требуют подтверждения. P не означает реализацию или успех теста; бизнес-тесты NOT_RUN.'), size=8.5)

page_heading('1 查询预约与取消  Поиск бронирование и отмена')
for parent in ['US-001', 'US-002', 'US-003']:
    row = next(r for r in inventory['rows'] if r['id'] == parent)
    s = next(s for s in inventory['rewritten_stories'] if s['parent_id'] == parent)
    p(s['id']+'  '+s['title']['zh-CN']+'  '+s['title']['ru-RU'], 'Heading 2')
    pair(bi('判断理由：'+row['rationale']['zh-CN'], 'Обоснование: '+row['rationale']['ru-RU']), size=9.5)
    pair(s['story'])
    pair(bi('计划验收：'+s['acceptance']['zh-CN'], 'Плановая приёмка: '+s['acceptance']['ru-RU']), size=9.5)
    p('需求 / Требования  '+', '.join(s['requirement_ids']), size=8.5)

def split_page(parent, heading):
    page_heading(heading)
    row = next(r for r in inventory['rows'] if r['id'] == parent)
    pair(bi('判断理由：'+row['rationale']['zh-CN'], 'Обоснование: '+row['rationale']['ru-RU']), size=9.5)
    for s in inventory['rewritten_stories']:
        if s['parent_id'] != parent: continue
        p(s['id']+'  '+s['title']['zh-CN']+'  '+s['title']['ru-RU'], 'Heading 2')
        pair(s['story'])
        pair(bi('计划验收：'+s['acceptance']['zh-CN'], 'Плановая приёмка: '+s['acceptance']['ru-RU']), size=9.5)
        p('需求 / Требования  '+', '.join(s['requirement_ids']), size=8.5)

split_page('US-004', '2 学生审核与账号资格  Проверка студента и допуск аккаунта')
pair(bi('外部前提：OPEN-003尚未闭合。合成数据只用于状态逻辑验证，真实审核必须有学校认可的来源与授权，不能将测试数据当作真实核验。',
        'Внешнее условие: OPEN-003 ещё не решён. Синтетические данные подходят только для логики состояния; реальная проверка требует признанного вузом источника и разрешения, тестовые данные не подтверждают реальный допуск.'), size=9)
pair(inventory['support_tasks'][0]['text'], size=9)

split_page('US-005', '3 资源与开放时间  Ресурсы и часы работы')
pair(bi('依赖与粒度：房间可预置以单测座位，资源与预约可预置以单测开放变更。US-005-D统一表达减少可用性的完整结果，原子规则同时适用于A、B、C的相关变更；不能把配置修改与取消预约拆成两个独立发布步骤。团队仍需评估每条的实际工作量。',
        'Зависимости и размер: помещения можно подготовить для проверки мест, ресурсы и брони — для проверки часов. US-005-D задаёт полный результат уменьшения доступности; его атомарные правила также действуют для соответствующих изменений A, B, C. Изменение конфигурации и отмены нельзя выпускать отдельными шагами. Команда ещё оценивает фактический объём каждой истории.'), size=9)

page_heading('4 判断口径与编号追溯  Критерии и трассируемость ID')
table(['项\nКритерий', '含义\nСмысл', '本次判断口径\nПравило оценки'],
      [[c['code']+'\n'+c['english'], c['meaning']['zh-CN']+'\n'+c['meaning']['ru-RU'], c['rule']['zh-CN']+'\n'+c['rule']['ru-RU']] for c in inventory['criteria']],
      [3.2, 5.2, 18.1], font=10)
pair(bi('编号追溯：US-001-R至US-003-R是原故事的改写候选，US-004-A至D、US-005-A至D是父故事的拆分候选。原US-001至US-006、原需求ID及第01至03份文档保留；这些候选不增加新的业务规则，不计为额外的父故事工作量。',
        'Трассируемость: US-001-R…US-003-R — кандидаты уточнения исходных историй, US-004-A…D и US-005-A…D — кандидаты разбиения. Исходные US-001…US-006, ID требований и документы 01…03 сохранены. Кандидаты не вводят новых бизнес-правил и не учитываются дополнительно к объёму родительских историй.'), size=10)
pair(bi('候选故事的验收条件用于后续实现验证，状态为PROPOSED，未实现NOT_IMPLEMENTED，业务测试NOT_RUN。拆分有助于独立验证，但共享授权、数据和事务仍需集成验证；不能据此宣称所有依赖已消除或全部INVEST项已最终通过。',
        'Условия приёмки кандидатов предназначены для будущей реализации; статус PROPOSED, реализация NOT_IMPLEMENTED, бизнес-тесты NOT_RUN. Разбиение помогает отдельной проверке, но общие права, данные и транзакции требуют интеграционной проверки; все зависимости и критерии INVEST ещё не считаются окончательно закрытыми.'), size=10)
p('来源 Источники', 'Heading 2')
p('课堂任务4 INVEST检查表 / Слайд задания 4 Проверка INVEST', size=9)
p('第03份用户故事待办表；data/backlog-v1.json；课堂requirements-practice.json v0.2；DEC-016；OPEN-003。\nДокумент 03 с бэклогом; data/backlog-v1.json; requirements-practice.json v0.2; DEC-016; OPEN-003.', size=9)
p('Bill Wake, INVEST in Good Stories, and SMART Tasks, 2003-08-17. 查阅 / Дата обращения 2026-10-14.\nhttps://xp123.com/invest-in-good-stories-and-smart-tasks/', size=9)
paragraphs=doc.paragraphs
for n,a in enumerate(paragraphs):
    if a.style.name=='Heading 2' and a.text.startswith('US-'):
        for b in paragraphs[n:]:
            b.paragraph_format.keep_with_next=not b.text.startswith('需求 / Требования')
            if b.text.startswith('需求 / Требования'): break
method_section=False
for a in paragraphs:
    if a.text.startswith('4 判断口径'): method_section=True
    if method_section and a.style.name=='Normal':
        a.paragraph_format.line_spacing=Pt(12.5)
        a.paragraph_format.space_after=Pt(1.5)
method_heading=next(a for a in doc.paragraphs if a.text.startswith('4 判断口径'))
for a in list(doc.paragraphs):
    if a.text.startswith(('候选故事的验收条件','Условия приёмки кандидатов')):
        method_heading._p.addprevious(a._p)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('StudySpace  04  |  ').font.size = Pt(8)
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
doc.core_properties.title = 'StudySpace INVEST检查与故事改写 Проверка INVEST'
doc.core_properties.author = 'StudySpace'
doc.save(DEST)
DEST.with_suffix('.md').write_text('\n\n'.join(md))

packet['documents'].append(dict(id='SS-PB-06', path=FILENAME, order=4, version='v1', status='DRAFT', languages=['zh-CN', 'ru-RU']))
packet['classroom_task4'] = {'source': 'data/invest-v1.json', 'audited_ids': [r['id'] for r in inventory['rows']], 'rewritten_ids': [s['id'] for s in inventory['rewritten_stories']], 'baseline_replaced': False}
packet_path.write_text(json.dumps(packet, ensure_ascii=False, indent=2))
(OUT/'manifest.txt').write_text('\n'.join(d['path'] for d in packet['documents'])+'\n')
(OUT/'README.md').write_text(
    '# StudySpace 课堂材料 / Учебные материалы\n\n2026-10-14 · v1 · DRAFT\n\n'+
    '\n'.join(f"{d['order']}. [{d['path']}]({d['path']})" for d in packet['documents'])+
    '\n\n01 用户画像 / Профили пользователей\n02 用户故事地图与版本计划 / Карта историй и выпуски\n03 六条核心用户故事五列表格 / Пять колонок и шесть основных историй\n04 任务4 INVEST九列检查表与故事改写 / Девять колонок INVEST и уточнение историй задания 4\n\n'+
    '编号按01至04续排；第01至03份文档内容保持不变。原文档、故事及需求ID保留。04按课件五行检查US-001至US-005，改写与拆分候选见data/invest-v1.json，未替换原待办故事。\n'+
    'Порядок продолжен от01 до04; содержимое документов01…03 не изменено. Сохранены ID документов, историй и требований. Документ04 проверяет US-001…US-005 по пяти строкам слайда; кандидаты уточнения и разбиения записаны в data/invest-v1.json и не заменяют исходные истории.\n\n'+
    '验收条件是计划，业务测试NOT_RUN；OPEN-003的真实学校核验前提仍未闭合。结构检查记录见data/verification.json及data/invest-verification.json。\n'+
    'Приёмка плановая, бизнес-тесты NOT_RUN; реальный источник проверки вуза в OPEN-003 ещё не определён. Структурная проверка: data/verification.json и data/invest-verification.json.\n\n'+
    '此前按要求删除的旧产品待办综合文档继续保持删除。\nРанее удалённые по запросу общие документы бэклога не восстановлены.\n')
print(json.dumps({'created': str(DEST), 'audit_rows': len(inventory['rows']), 'rewritten_stories': len(inventory['rewritten_stories']), 'preserved_prior_documents': len(inventory['preserved_docx'])}, ensure_ascii=False))
